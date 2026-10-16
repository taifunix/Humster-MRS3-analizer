"""Workbook of tick-tested v4 strategies: parameters, training estimate and real MRS3 results per window.

Usage: local_best_strategies_xlsx.py <tick results csv> <equity features csv> <output xlsx> [<trainwin results csv>]
Sheets: Best (score > 20), Admitted (score > 10), All tested, By pair. Forward results are real tick
tests of the 28 days after the cut; training columns come from the Source v6 summed member curves
(proxy) and, when given, from real tick tests on the training window.
"""
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def orders(text):
    out = {}
    for k, token in enumerate(text.split(), 1):
        shift, ma = re.match(r"(\d+)bp/MA(\d+)", token).groups()
        out[f"{k} Shift, %"] = int(shift) / 100
        out[f"{k} MA"] = int(ma)
    return out


def build(res, feat, trainwin):
    df = res.merge(feat[["strategy", "train_days", "ret_full30", "ulcer", "maxdd", "float_mean", "er", "up7", "new_ath_w7"]],
                   on="strategy", how="left")
    o = pd.DataFrame([orders(t) for t in df.orders])
    df = pd.concat([df.reset_index(drop=True), o], axis=1)
    df["test_start"] = df.t_date
    df["test_end"] = (pd.to_datetime(df.t_date) + pd.to_timedelta(28, unit="D")).dt.date.astype(str)
    group = {"ADMIT r1": "Допущена", "ADMIT r2": "Допущена (2-я)", "REJECT lo": "Отсеяна (< 5)", "REJECT mid": "Отсеяна (5–10)"}
    cols = {
        "strategy": "Стратегия", "symbol": "Пара", "timeframe": "ТФ", "group": "Группа", "score": "Оценка",
        "order_count": "ORD", "close_ma": "Close",
        "1 Shift, %": "1 Shift, %", "1 MA": "1 MA", "2 Shift, %": "2 Shift, %", "2 MA": "2 MA", "3 Shift, %": "3 Shift, %", "3 MA": "3 MA",
        "lot_mode": "Lot mode",
        "train_days": "Обуч: дней", "ret_full30": "Обуч: PnL/30", "maxdd": "Обуч: DD", "ulcer": "Обуч: Ulcer",
        "float_mean": "Обуч: плав. убыток ср.", "er": "Обуч: гладкость",
        "tw_ret30": "Обуч тик: PnL/30", "tw_dd": "Обуч тик: DD", "tw_trades": "Обуч тик: Trades",
        "test_start": "Тест: Start", "test_end": "Тест: End",
        "pnl_pct": "Тест: PnL, %", "ret30": "Тест: PnL/30", "dd": "Тест: DD", "q": "Тест: PnL DD5/30",
        "ulcer_fwd": "Тест: Ulcer", "upnl_dd": "Тест: Max uPnL DD", "trades": "Тест: Trades", "win_rate": "Тест: W/R",
        "proxy_ret30": "Прогноз: PnL/30", "proxy_dd_fwd": "Прогноз: DD",
    }
    df["order_count"] = df.orders.str.split().str.len()
    df["lot_mode"] = "EQUAL"
    df["group"] = df.group.map(group).fillna(df.group)
    df = df.rename(columns={"ulcer_x": "ulcer", "ulcer_y": "ulcer_fwd"}) if "ulcer_x" in df else df
    if trainwin is not None:
        tw = trainwin.rename(columns={"ret30": "tw_ret30", "dd": "tw_dd", "trades": "tw_trades"})[["strategy", "tw_ret30", "tw_dd", "tw_trades"]]
        df = df.merge(tw, on="strategy", how="left")
    keep = [k for k in cols if k in df.columns]
    return df[keep].rename(columns=cols)


def by_pair(table):
    a = table[table["Группа"] == "Допущена"]
    g = a.groupby("Пара")
    out = pd.DataFrame({
        "Стратегий": g.size(), "ТФ": g["ТФ"].agg(lambda v: ", ".join(sorted(set(v)))),
        "Срезов": g["Тест: Start"].nunique(), "Оценка мед.": g["Оценка"].median(),
        "Тест: PnL/30 ср.": g["Тест: PnL/30"].mean(), "Тест: PnL/30 мин.": g["Тест: PnL/30"].min(),
        "Тест: DD ср.": g["Тест: DD"].mean(), "Тест: DD макс.": g["Тест: DD"].max(),
        "Убыточных": g["Тест: PnL, %"].agg(lambda v: int((v < 0).sum())),
        "Оценка > 20": g["Оценка"].agg(lambda v: int((v > 20).sum())),
    })
    return out.sort_values("Тест: PnL/30 ср.", ascending=False).reset_index()


def style(ws, frame):
    head = PatternFill("solid", fgColor="FFD9E1F2")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), head
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    for i, name in enumerate(frame.columns, 1):
        width = 44 if name == "Стратегия" else max(9, min(16, len(str(name)) + 2))
        ws.column_dimensions[get_column_letter(i)].width = width
        fmt = "0.00" if frame[name].dtype.kind == "f" else None
        if name.endswith("Shift, %"):
            fmt = "0.00"
        if fmt:
            for cell in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell[0].number_format = fmt
    green, red = PatternFill("solid", fgColor="FFE2EFDA"), PatternFill("solid", fgColor="FFFCE4D6")
    if "Тест: PnL, %" in frame.columns:
        col = list(frame.columns).index("Тест: PnL, %") + 1
        for row in ws.iter_rows(min_row=2):
            v = row[col - 1].value
            if isinstance(v, (int, float)) and np.isfinite(v):
                row[col - 1].fill = green if v >= 0 else red


def main():
    res, feat, out = pd.read_csv(sys.argv[1]), pd.read_csv(sys.argv[2]), Path(sys.argv[3])
    trainwin = pd.read_csv(sys.argv[4]) if len(sys.argv) > 4 else None
    res = res.rename(columns={"ulcer": "ulcer_fwd"})
    table = build(res, feat, trainwin)
    table = table.sort_values(["Пара", "ТФ", "Тест: Start", "Группа"]).reset_index(drop=True)
    sheets = {
        "Лучшие (оценка > 20)": table[(table["Оценка"] > 20) & table["Группа"].isin(["Допущена"])],
        "Допущенные (> 10)": table[table["Группа"].str.startswith("Допущена")],
        "Все протестированные": table,
        "По парам": by_pair(table),
    }
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        for name, frame in sheets.items():
            frame.to_excel(xw, sheet_name=name, index=False)
            style(xw.sheets[name], frame)
    print({k: len(v) for k, v in sheets.items()}, "->", out)


if __name__ == "__main__":
    main()
