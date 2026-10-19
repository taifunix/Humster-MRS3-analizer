"""Export the optimizer base lot: from the pair registry (LOT_MODEL_V2, default) or from Bybit minute files (V1).

Registry mode (pair registry stage 4, docs/specs/2026-10-18-pair-registry.md): each (pair, date) cell takes the
registry window ending on the last weekday before that date; its K and step must equal the export's. The workbook
records the model of its values (defined name MRS3_BASE_LOT_MODEL; none means LOT_MODEL_V1): a run of the other model
refuses to add to an existing history, only `--rebuild-history` replaces all ten dates (ADR-0079: V1 and V2 values
are never mixed). Registry windows of past days come from `python -m mrs3.pair_registry liquidity --days 14`."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_DOWN, ROUND_HALF_EVEN, localcontext
import argparse
from copy import copy
import os
from pathlib import Path
import re
import stat
import sys
import tempfile
from typing import Callable, Iterable
import zipfile

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName

if __package__ in {None, ""}:  # Allow ``python scripts/calculate_bybit_base_lots.py``.
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from mrs3.portfolio.minute_capacity import (  # noqa: E402
    MinuteCapacityError,
    backfill_missing_days,
    calculate_minute_capacity,
    fetch_bybit_trade_archive,
    resolve_liquidity_window,
)


DEFAULT_WORKBOOK = Path(r"D:\SHARE\!MN\hamster\MRS-Analizer\Input\bybit_tradfi_liquidity.xlsx")
DEFAULT_DATA_ROOT = Path(r"D:\SHARE\!MN\hamster\hb\tester\data\bybit")
DEFAULT_SHEET = "Actual"
DEFAULT_K = Decimal("9")
DEFAULT_ROUND_DOWN_USDT = Decimal("10")
DEFAULT_PUBLICATION_LAG_HOURS = 6
DEFAULT_WORKERS = 16
HISTORY_DAYS = 10
AVERAGE_COLUMN = 3
HISTORY_FIRST_COLUMN = 4
PREVIOUS_HISTORY_FIRST_COLUMN = 3
AVERAGE_HEADER = "Среднее 10д, USDT"
LEGACY_ACTUALIZATION_HEADER = "\u0414\u0430\u0442\u0430 \u0430\u043a\u0442\u0443\u0430\u043b\u0438\u0437\u0430\u0446\u0438\u0438"
LOT_NUMBER_FORMAT = "#,##0"
FRACTIONAL_NUMBER_FORMAT = "#,##0.##########"
ERROR_FILL = PatternFill(fill_type="solid", fgColor="FFC7CE")
_CLEAR_FILL = PatternFill(fill_type=None)
_SYMBOL = re.compile(r"[A-Z0-9]+\Z", re.ASCII)
MODEL_NAME = "MRS3_BASE_LOT_MODEL"
MINUTE_FILES_MODEL = "LOT_MODEL_V1"
REGISTRY_MODEL = "LOT_MODEL_V2"


def recorded_model(workbook) -> str:
    """The model of the workbook's base-lot values; a workbook without the marker holds LOT_MODEL_V1 values."""
    marker = workbook.defined_names.get(MODEL_NAME)
    return str(marker.attr_text).strip('"') if marker is not None and marker.attr_text else MINUTE_FILES_MODEL


class BaseLotExportError(ValueError):
    """A user-facing input, data, or publication error."""


@dataclass(frozen=True, slots=True)
class ExportRow:
    symbol: str
    actualization_date: date
    v25: Decimal
    a15: Decimal
    base_lot_usdt: Decimal
    content_digest: str
    file_hashes: tuple[tuple[str, str], ...]


@dataclass(frozen=True, slots=True)
class ExportResult:
    workbook_path: Path
    rows_written: int
    rows: tuple[ExportRow, ...]
    errors: tuple[ExportError, ...]
    actualization_date: date
    skipped_existing_date: bool = False


@dataclass(frozen=True, slots=True)
class ExportError:
    symbol: str
    actualization_date: date
    message: str


_ISO_DATE = re.compile(r"\d{4}-\d{2}-\d{2}\Z", re.ASCII)


def _date_header(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    # Excel tables store header cells as text; a strict ISO date is the same header.
    if isinstance(value, str) and _ISO_DATE.match(value.strip()):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


def _number_format_for_value(value: object) -> str:
    if isinstance(value, str):
        return "General"
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        try:
            if Decimal(str(value)) != Decimal(str(value)).to_integral_value():
                return FRACTIONAL_NUMBER_FORMAT
        except InvalidOperation:
            pass
    return LOT_NUMBER_FORMAT


def _calculation_anchor(target_date: date, publication_lag_hours: int) -> datetime:
    return datetime.combine(target_date, time(), timezone.utc) + timedelta(hours=publication_lag_hours)


def _history_columns(sheet, first_column: int = HISTORY_FIRST_COLUMN) -> tuple[date, ...] | None:
    headers = tuple(_date_header(sheet.cell(1, col).value) for col in range(first_column, first_column + HISTORY_DAYS))
    if not any(value is not None for value in headers):
        return None
    if any(value is None for value in headers) or tuple(sorted(headers)) != headers or len(set(headers)) != HISTORY_DAYS:
        raise BaseLotExportError("Actual sheet has an incomplete or unordered base-lot date history")
    if any((headers[index + 1] - headers[index]).days != 1 for index in range(HISTORY_DAYS - 1)):
        raise BaseLotExportError("Actual sheet date history must contain ten consecutive dates")
    return headers  # type: ignore[return-value]


def _migrate_previous_layout(sheet) -> bool:
    """Move a C:L date history (layout before the average column) one column right to D:M."""
    # A date in C1 marks the previous layout; in the current one C1 is the average header.
    if _date_header(sheet.cell(1, PREVIOUS_HISTORY_FIRST_COLUMN).value) is None:
        return False
    _history_columns(sheet, PREVIOUS_HISTORY_FIRST_COLUMN)
    last = PREVIOUS_HISTORY_FIRST_COLUMN + HISTORY_DAYS - 1
    if any(sheet.cell(row, last + 1).value is not None for row in range(1, sheet.max_row + 1)):
        raise BaseLotExportError("cannot move the date history to D:M: column M contains existing data")
    from openpyxl.utils.cell import range_boundaries

    for row in range(1, sheet.max_row + 1):
        for col in range(last, PREVIOUS_HISTORY_FIRST_COLUMN - 1, -1):
            _copy_cell(sheet.cell(row, col), sheet.cell(row, col + 1))
        average = sheet.cell(row, AVERAGE_COLUMN)
        average.value = None
        average.comment = None
        average.fill = _CLEAR_FILL
        average.number_format = LOT_NUMBER_FORMAT if row > 1 else "General"
    _shift_column_dimensions(sheet, PREVIOUS_HISTORY_FIRST_COLUMN, last)
    if sheet.auto_filter.ref:
        min_col = range_boundaries(sheet.auto_filter.ref)[0]
        for column_filter in sheet.auto_filter.filterColumn:
            if PREVIOUS_HISTORY_FIRST_COLUMN <= min_col + column_filter.colId <= last:
                column_filter.colId += 1
    return True


def _shift_column_dimensions(sheet, first: int, last: int) -> None:
    """Move column widths of first..last one column right; first keeps its own width.

    openpyxl keeps grouped <col min max> entries under the first letter, so every entry that
    touches first..last+1 is split into single columns before shifting; ranges never overlap.
    """
    from openpyxl.utils.cell import column_index_from_string, get_column_letter
    from openpyxl.worksheet.dimensions import ColumnDimension

    per_column: dict[int, tuple] = {}
    for key in list(sheet.column_dimensions.keys()):
        dimension = sheet.column_dimensions[key]
        low = dimension.min or column_index_from_string(key)
        high = dimension.max or low
        if high < first or low > last + 1:
            continue
        for column in range(low, high + 1):
            per_column[column] = (dimension.width, dimension.customWidth, copy(dimension._style))
        del sheet.column_dimensions[key]
    shifted = dict(per_column)
    for column in range(last, first - 1, -1):
        if column in per_column:
            shifted[column + 1] = per_column[column]
    for column, (width, custom, style) in sorted(shifted.items()):
        letter = get_column_letter(column)
        dimension = ColumnDimension(sheet, index=letter, width=width, customWidth=custom)
        dimension._style = style
        sheet.column_dimensions[letter] = dimension


def _check_average_column(sheet) -> None:
    value = sheet.cell(1, AVERAGE_COLUMN).value
    if value is not None and value != AVERAGE_HEADER:
        raise BaseLotExportError(f"column C is reserved for the 10-day average; it holds {value!r}")


def _average_lot(values: Iterable[object], step: Decimal) -> int | Decimal | None:
    numbers = [Decimal(str(v)) for v in values if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool)]
    if not numbers:
        return None
    with localcontext() as context:
        context.prec = 28
        mean = sum(numbers, Decimal(0)) / Decimal(len(numbers))
        floored = (mean / step).to_integral_value(rounding=ROUND_DOWN) * step
    return int(floored) if floored == floored.to_integral_value() else floored


def _expected_averages(sheet, symbol_rows: dict[str, tuple[int, ...]], step: Decimal) -> dict[int, object]:
    return {
        row: _average_lot((sheet.cell(row, col).value for col in range(HISTORY_FIRST_COLUMN, HISTORY_FIRST_COLUMN + HISTORY_DAYS)), step)
        for rows in symbol_rows.values() for row in rows
    }


def _averages_stale(sheet, symbol_rows: dict[str, tuple[int, ...]], step: Decimal) -> bool:
    if sheet.cell(1, AVERAGE_COLUMN).value != AVERAGE_HEADER:
        return True
    return any(sheet.cell(row, AVERAGE_COLUMN).value != value for row, value in _expected_averages(sheet, symbol_rows, step).items())


def _write_averages(sheet, symbol_rows: dict[str, tuple[int, ...]], step: Decimal) -> None:
    sheet.cell(1, AVERAGE_COLUMN).value = AVERAGE_HEADER
    for row, value in _expected_averages(sheet, symbol_rows, step).items():
        cell = sheet.cell(row, AVERAGE_COLUMN)
        cell.value = value
        cell.comment = None
        cell.number_format = LOT_NUMBER_FORMAT


def _copy_cell(source, target) -> None:
    target.value = source.value
    target._style = copy(source._style)
    target.number_format = source.number_format
    target.comment = copy(source.comment) if source.comment else None


def _symbols_listed_by(sheet, symbols: tuple[str, ...], symbol_rows: dict[str, tuple[int, ...]], target_date: date) -> tuple[str, ...]:
    eligible = []
    for symbol in symbols:
        listing_dates = []
        has_unknown_listing = False
        for row in symbol_rows[symbol]:
            value = sheet.cell(row, 2).value
            if isinstance(value, datetime):
                listing_dates.append(value.date())
            elif isinstance(value, date):
                listing_dates.append(value)
            else:
                has_unknown_listing = True
        if has_unknown_listing or not listing_dates or any(value <= target_date for value in listing_dates):
            eligible.append(symbol)
    return tuple(eligible)


def _symbols_to_fill(sheet, symbols: tuple[str, ...], symbol_rows: dict[str, tuple[int, ...]], target_date: date, column: int) -> tuple[str, ...]:
    """Pairs listed by the date whose cell in this date column holds no numeric lot (blank, ERROR or text)."""
    listed = set(_symbols_listed_by(sheet, symbols, symbol_rows, target_date))
    return tuple(
        symbol for symbol, row_numbers in symbol_rows.items()
        if symbol in listed and any(
            isinstance(sheet.cell(row, column).value, bool)
            or not isinstance(sheet.cell(row, column).value, (int, float, Decimal))
            for row in row_numbers
        )
    )


def _validate_legacy_history(sheet) -> None:
    """Legacy layout: C = base lot, D = actualization date; the dates go to D:M and C becomes the average."""
    if sheet.cell(1, AVERAGE_COLUMN).value is None or sheet.cell(1, AVERAGE_COLUMN + 1).value != LEGACY_ACTUALIZATION_HEADER:
        raise BaseLotExportError("Actual sheet is not in the recognized legacy base-lot layout")
    if any(sheet.cell(row, col).value is not None for row in range(1, sheet.max_row + 1) for col in range(AVERAGE_COLUMN + 2, HISTORY_FIRST_COLUMN + HISTORY_DAYS)):
        raise BaseLotExportError("cannot initialize date history: columns E:M contain existing data")


def _clear_output_filters(sheet, first_changed_col: int = HISTORY_FIRST_COLUMN, last_changed_col: int = HISTORY_FIRST_COLUMN + HISTORY_DAYS - 1) -> None:
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    auto_filter = sheet.auto_filter
    min_col = range_boundaries(auto_filter.ref)[0] if auto_filter.ref else 1
    retained_filters = []
    for column_filter in auto_filter.filterColumn:
        worksheet_col = min_col + column_filter.colId
        if first_changed_col <= worksheet_col <= last_changed_col:
            continue
        column_filter.colId = worksheet_col - 1
        retained_filters.append(column_filter)
    auto_filter.filterColumn = retained_filters
    auto_filter.ref = f"A1:{get_column_letter(max(sheet.max_column, HISTORY_FIRST_COLUMN + HISTORY_DAYS - 1))}{sheet.max_row}"


def _initialize_history(sheet, dates: tuple[date, ...]) -> None:
    _validate_legacy_history(sheet)
    last_row = sheet.max_row
    template_style = copy(sheet.cell(1, AVERAGE_COLUMN)._style) if sheet.cell(1, AVERAGE_COLUMN).has_style else None
    legacy_styles = [copy(sheet.cell(row, AVERAGE_COLUMN)._style) if sheet.cell(row, AVERAGE_COLUMN).has_style else None for row in range(2, last_row + 1)]
    legacy_width = sheet.column_dimensions["C"].width
    for offset, target_date in enumerate(dates):
        col = HISTORY_FIRST_COLUMN + offset
        letter = sheet.cell(1, col).column_letter
        sheet.column_dimensions[letter].width = legacy_width
        for row, style in enumerate(legacy_styles, start=2):
            target = sheet.cell(row, col)
            target.value = None
            target.comment = None
            if style is not None:
                target._style = copy(style)
            target.number_format = LOT_NUMBER_FORMAT
        header = sheet.cell(1, col)
        if template_style is not None:
            header._style = copy(template_style)
        header.value = target_date
        header.comment = None
        header.number_format = "yyyy-mm-dd"
    # C changes meaning (legacy base lot -> average), so its old criteria go too.
    _clear_output_filters(sheet, AVERAGE_COLUMN)


def _reset_history(
    sheet,
    dates: tuple[date, ...],
    previous_dates: tuple[date, ...],
    *,
    preserve_existing: bool = True,
) -> None:
    previous_columns = {}
    if preserve_existing:
        for old_column, old_date in enumerate(previous_dates, start=HISTORY_FIRST_COLUMN):
            previous_columns[old_date] = [
                (sheet.cell(row, old_column).value, copy(sheet.cell(row, old_column)._style), copy(sheet.cell(row, old_column).comment) if sheet.cell(row, old_column).comment else None)
                for row in range(2, sheet.max_row + 1)
            ]
    for col, target_date in enumerate(dates, start=HISTORY_FIRST_COLUMN):
        header = sheet.cell(1, col)
        header.value = target_date
        header.number_format = "yyyy-mm-dd"
        if preserve_existing and target_date in previous_columns:
            for row, (value, style, comment) in enumerate(previous_columns[target_date], start=2):
                cell = sheet.cell(row, col)
                cell.value = value
                cell._style = copy(style)
                cell.comment = comment
                cell.number_format = _number_format_for_value(value)
            continue
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row, col)
            cell.value = None
            cell.fill = _CLEAR_FILL
            cell.comment = None
            cell.number_format = LOT_NUMBER_FORMAT
    _clear_output_filters(sheet)


def _rotate_history(sheet, new_date: date) -> None:
    for row in range(1, sheet.max_row + 1):
        for col in range(HISTORY_FIRST_COLUMN, HISTORY_FIRST_COLUMN + HISTORY_DAYS - 1):
            _copy_cell(sheet.cell(row, col + 1), sheet.cell(row, col))
        newest = sheet.cell(row, HISTORY_FIRST_COLUMN + HISTORY_DAYS - 1)
        if row == 1:
            newest.value = new_date
            newest.number_format = "yyyy-mm-dd"
        else:
            newest.value = None
            newest.fill = _CLEAR_FILL
            newest.comment = None
            newest.number_format = LOT_NUMBER_FORMAT


def _decimal(value: object, field: str, *, positive: bool = False, nonnegative: bool = False) -> Decimal:
    if isinstance(value, (bool, float)):
        raise BaseLotExportError(f"{field} must be an exact decimal")
    try:
        result = value if isinstance(value, Decimal) else Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError) as error:
        raise BaseLotExportError(f"{field} must be an exact decimal") from error
    if not result.is_finite() or (positive and result <= 0) or (nonnegative and result < 0):
        raise BaseLotExportError(f"{field} has an invalid value")
    return result


def calculate_base_lot(
    v25: object,
    a15: object,
    *,
    k: object = DEFAULT_K,
    round_down_usdt: object = DEFAULT_ROUND_DOWN_USDT,
) -> Decimal:
    """Return ``floor_to_step(K * V25 * A15)`` using Decimal arithmetic."""
    turnover = _decimal(v25, "v25", positive=True)
    activity = _decimal(a15, "a15", nonnegative=True)
    coefficient = _decimal(k, "k", positive=True)
    step = _decimal(round_down_usdt, "round_down_usdt", positive=True)
    if step != step.to_integral_value():
        raise BaseLotExportError("round_down_usdt must be a whole number")
    precision = max(28, sum(len(value.as_tuple().digits) for value in (coefficient, turnover, activity)) + 8)
    with localcontext() as context:
        context.prec = precision
        context.rounding = ROUND_HALF_EVEN
        raw = coefficient * turnover * activity
    value_digits = int("".join(str(digit) for digit in raw.as_tuple().digits) or "0")
    step_digits = int("".join(str(digit) for digit in step.as_tuple().digits) or "0")
    value_exponent = raw.as_tuple().exponent
    step_exponent = step.as_tuple().exponent
    if value_exponent >= step_exponent:
        numerator = value_digits * 10 ** (value_exponent - step_exponent)
        denominator = step_digits
    else:
        numerator = value_digits
        denominator = step_digits * 10 ** (step_exponent - value_exponent)
    quotient = numerator // denominator
    with localcontext() as context:
        context.prec = max(28, len(str(quotient)) + len(step.as_tuple().digits) + 8)
        context.rounding = ROUND_HALF_EVEN
        return Decimal(quotient) * step


def _parse_symbol(value: object, row_number: int, sheet_name: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise BaseLotExportError(f"{sheet_name} row {row_number}: symbol must be uppercase ASCII letters/digits")
    symbol = value.strip()
    if not _SYMBOL.fullmatch(symbol):
        raise BaseLotExportError(f"{sheet_name} row {row_number}: symbol must be uppercase ASCII letters/digits")
    return symbol


def _validate_ordinary_workbook(workbook, path: Path) -> None:
    if path.suffix.lower() != ".xlsx":
        raise BaseLotExportError("only ordinary .xlsx workbooks are supported")
    for sheet in workbook.worksheets:
        if getattr(sheet, "_charts", ()) or getattr(sheet, "_images", ()) or getattr(sheet, "tables", {}):
            raise BaseLotExportError("only ordinary xlsx workbooks without charts, images, or tables are supported")


def _validate_xlsx_package(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as package:
            names = package.namelist()
    except (OSError, zipfile.BadZipFile) as error:
        raise BaseLotExportError(f"cannot inspect workbook package {path}: {error}") from error
    unsupported_prefixes = (
        "xl/charts/", "xl/media/", "xl/tables/", "xl/pivottables/", "xl/pivotcache/",
        "xl/embeddings/", "xl/slicers/", "xl/slicercaches/", "xl/ctrlprops/",
    )
    unsupported = any(name.lower().startswith(unsupported_prefixes) for name in names)
    unsupported = unsupported or any(
        name.lower().startswith("xl/drawings/")
        and not name.lower().endswith(".vml")
        and not (name.lower().startswith("xl/drawings/_rels/") and "vmldrawing" in name.lower())
        for name in names
    )
    unsupported = unsupported or any(name.lower().endswith("vbaproject.bin") for name in names)
    if unsupported:
        raise BaseLotExportError("only ordinary xlsx workbooks without charts, images, or tables are supported")


def _read_symbol_rows(path: Path, sheet_name: str) -> tuple[tuple[str, ...], dict[str, tuple[int, ...]]]:
    _validate_xlsx_package(path)
    try:
        workbook = load_workbook(path, read_only=False, data_only=False)
    except (OSError, ValueError) as error:
        raise BaseLotExportError(f"cannot open workbook {path}: {error}") from error
    try:
        _validate_ordinary_workbook(workbook, path)
        if sheet_name not in workbook.sheetnames:
            raise BaseLotExportError(f"workbook is missing sheet {sheet_name!r}")
        sheet = workbook[sheet_name]
        symbols: list[str] = []
        rows: dict[str, list[int]] = {}
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, max_col=1, values_only=True), start=2):
            value = row[0] if row else None
            if value is None or (isinstance(value, str) and not value.strip()):
                continue
            symbol = _parse_symbol(value, row_number, sheet_name)
            if symbol not in rows:
                symbols.append(symbol)
                rows[symbol] = []
            rows[symbol].append(row_number)
        return tuple(symbols), {symbol: tuple(indices) for symbol, indices in rows.items()}
    finally:
        workbook.close()


def _ensure_data(
    data_root: Path,
    symbol: str,
    days: Iterable,
    *,
    workers: int,
    fetch_day: Callable[[str, object], bytes],
) -> None:
    try:
        result = backfill_missing_days(
            data_root,
            symbol,
            days,
            fetch_day=fetch_day,
            enabled=True,
            max_workers=workers,
        )
    except (MinuteCapacityError, OSError, ValueError) as error:
        raise BaseLotExportError(f"{symbol}: cannot validate Bybit minute data: {error}") from error
    if result.missing or result.failed:
        failed = ", ".join(f"{day}: {reason}" for day, reason in result.failed.items())
        missing = ", ".join(day.isoformat() for day in result.missing)
        detail = "; ".join(item for item in (failed, f"missing {missing}" if missing else "") if item)
        raise BaseLotExportError(f"{symbol}: official Bybit data unavailable ({detail})")


def _atomic_save_workbook(workbook, target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(dir=target.parent, prefix=f".{target.name}.", suffix=".tmp", delete=False) as stream:
            temporary = Path(stream.name)
        workbook.save(temporary)
        try:
            mode = stat.S_IMODE(target.stat().st_mode)
        except FileNotFoundError:
            mode = None
        if mode is not None:
            os.chmod(temporary, mode)
        os.replace(temporary, target)
    except OSError as error:
        raise BaseLotExportError(f"cannot atomically save workbook {target}: {error}") from error
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def run_export(
    workbook_path: str | Path = DEFAULT_WORKBOOK,
    data_root: str | Path = DEFAULT_DATA_ROOT,
    *,
    sheet_name: str = DEFAULT_SHEET,
    anchor_created_at: datetime | None = None,
    k: object = DEFAULT_K,
    round_down_usdt: object = DEFAULT_ROUND_DOWN_USDT,
    publication_lag_hours: int = DEFAULT_PUBLICATION_LAG_HOURS,
    workers: int = DEFAULT_WORKERS,
    rebuild_history: bool = False,
    fetch_day: Callable[[str, object], bytes] | None = None,
    lot_source: Callable[[str, date], ExportRow] | None = None,
    model: str | None = None,
) -> ExportResult:
    """Calculate and publish a rolling ten-date base-lot history. `lot_source(symbol, date)` supplies each cell
    (the pair registry, `registry_lot_source`); without it LOT_MODEL_V1 reads the minute files. `model` is the model
    of the values (default: LOT_MODEL_V2 with a `lot_source`, else LOT_MODEL_V1)."""
    model = model or (REGISTRY_MODEL if lot_source is not None else MINUTE_FILES_MODEL)
    workbook_path = Path(workbook_path)
    data_root = Path(data_root)
    if not isinstance(sheet_name, str) or not sheet_name:
        raise BaseLotExportError("sheet_name must be non-empty text")
    if type(workers) is not int or not 1 <= workers <= 16:
        raise BaseLotExportError("workers must be an integer from 1 through 16")
    if type(publication_lag_hours) is not int or not 0 <= publication_lag_hours <= 48:
        raise BaseLotExportError("publication_lag_hours must be an integer from 0 through 48")
    coefficient = _decimal(k, "k", positive=True)
    step = _decimal(round_down_usdt, "round_down_usdt", positive=True)
    if step != step.to_integral_value():
        raise BaseLotExportError("round_down_usdt must be a whole number")
    if anchor_created_at is None:
        anchor_created_at = datetime.now(timezone.utc)
    if not isinstance(anchor_created_at, datetime) or anchor_created_at.tzinfo is None or anchor_created_at.utcoffset() is None:
        raise BaseLotExportError("anchor_created_at must be timezone-aware")
    anchor_created_at = anchor_created_at.astimezone(timezone.utc)
    actualization_date = anchor_created_at.date()
    try:
        source_stat = workbook_path.stat()
        source_signature = source_stat.st_mtime_ns, source_stat.st_size
    except OSError as error:
        raise BaseLotExportError(f"cannot inspect workbook {workbook_path}: {error}") from error
    symbols, symbol_rows = _read_symbol_rows(workbook_path, sheet_name)
    if not symbols:
        raise BaseLotExportError(f"{sheet_name} sheet contains no symbols")

    try:
        workbook = load_workbook(workbook_path, data_only=False)
    except (OSError, ValueError) as error:
        raise BaseLotExportError(f"cannot open workbook {workbook_path}: {error}") from error
    try:
        _validate_ordinary_workbook(workbook, workbook_path)
        if sheet_name not in workbook.sheetnames:
            raise BaseLotExportError(f"workbook is missing sheet {sheet_name!r}")
        sheet = workbook[sheet_name]
        migrated = _migrate_previous_layout(sheet)
        history_dates = _history_columns(sheet)
        if history_dates is not None and recorded_model(workbook) != model and not rebuild_history:
            raise BaseLotExportError(
                f"the sheet holds {recorded_model(workbook)} base lots and this run computes {model}: replace all ten "
                "dates with --rebuild-history (registry windows of past days: python -m mrs3.pair_registry liquidity "
                "--days 14), or use the matching --source; workbook unchanged")
        if history_dates is not None:
            _check_average_column(sheet)
        current_date_present = history_dates is not None and actualization_date in history_dates
        calculation_symbols: dict[date, tuple[str, ...]] = {}
        history_action = "initialize"
        if rebuild_history:
            if history_dates is not None and actualization_date < history_dates[-1]:
                raise BaseLotExportError("anchor date is older than the newest date in the workbook history")
            if history_dates is None:
                _validate_legacy_history(workbook[sheet_name])
                target_dates = tuple(actualization_date - timedelta(days=offset) for offset in range(HISTORY_DAYS - 1, -1, -1))
                history_action = "initialize"
            elif current_date_present:
                target_dates = history_dates
                history_action = "rebuild"
            else:
                target_dates = tuple(actualization_date - timedelta(days=offset) for offset in range(HISTORY_DAYS - 1, -1, -1))
                history_action = "rebuild"
            calculation_symbols = {
                target_date: _symbols_listed_by(sheet, symbols, symbol_rows, target_date)
                for target_date in target_dates
            }
        elif current_date_present:
            for target_date in history_dates:
                column = HISTORY_FIRST_COLUMN + history_dates.index(target_date)
                needs_fill = _symbols_to_fill(sheet, symbols, symbol_rows, target_date, column)
                if needs_fill:
                    calculation_symbols[target_date] = needs_fill
            if not calculation_symbols and not migrated and not _averages_stale(sheet, symbol_rows, step):
                return ExportResult(workbook_path, 0, (), (), actualization_date, True)
            target_dates = history_dates
            history_action = "retry"
        else:
            if history_dates is None:
                _validate_legacy_history(workbook[sheet_name])
                target_dates = tuple(actualization_date - timedelta(days=offset) for offset in range(HISTORY_DAYS - 1, -1, -1))
                history_action = "initialize"
                calculation_symbols = {
                    target_date: _symbols_listed_by(sheet, symbols, symbol_rows, target_date)
                    for target_date in target_dates
                }
            else:
                gap = (actualization_date - history_dates[-1]).days
                if gap <= 0:
                    raise BaseLotExportError("anchor date is older than the newest date in the workbook history")
                if gap == 1:
                    target_dates = (*history_dates[1:], actualization_date)
                    history_action = "rotate"
                    calculation_symbols[actualization_date] = _symbols_listed_by(sheet, symbols, symbol_rows, actualization_date)
                else:
                    target_dates = tuple(actualization_date - timedelta(days=offset) for offset in range(HISTORY_DAYS - 1, -1, -1))
                    history_action = "rebuild"
                    for target_date in target_dates:
                        if target_date not in history_dates:
                            calculation_symbols[target_date] = _symbols_listed_by(sheet, symbols, symbol_rows, target_date)
                        else:
                            old_column = HISTORY_FIRST_COLUMN + history_dates.index(target_date)
                            needs_fill = _symbols_to_fill(sheet, symbols, symbol_rows, target_date, old_column)
                            if needs_fill:
                                calculation_symbols[target_date] = needs_fill
                if history_action == "rotate":
                    for target_date in target_dates[:-1]:
                        old_column = HISTORY_FIRST_COLUMN + history_dates.index(target_date)
                        needs_fill = _symbols_to_fill(sheet, symbols, symbol_rows, target_date, old_column)
                        if needs_fill:
                            calculation_symbols[target_date] = needs_fill
    finally:
        workbook.close()

    fetch = fetch_day or fetch_bybit_trade_archive
    rows: list[ExportRow] = []
    errors: list[ExportError] = []
    for target_date, write_symbols in calculation_symbols.items():
        calculation_anchor = _calculation_anchor(target_date, publication_lag_hours)
        days = resolve_liquidity_window(calculation_anchor, publication_lag_hours)
        for symbol in write_symbols:
            try:
                if lot_source is not None:
                    rows.append(lot_source(symbol, target_date))
                    continue
                _ensure_data(data_root, symbol, days, workers=workers, fetch_day=fetch)
                capacity = calculate_minute_capacity(
                    data_root,
                    symbol,
                    round_down_usdt=step,
                    publication_lag_hours=publication_lag_hours,
                    lot_model=True,
                    anchor_created_at=calculation_anchor,
                )
                features = capacity.lot_features
                if features is None:
                    raise BaseLotExportError("liquidity features are unavailable")
                rows.append(ExportRow(symbol, target_date, features.v25, features.a15, calculate_base_lot(features.v25, features.a15, k=coefficient, round_down_usdt=step), features.content_digest, features.file_hashes))
            except (BaseLotExportError, MinuteCapacityError, OSError, ValueError, ArithmeticError) as error:
                message = str(error)
                prefix = f"{symbol}: "
                errors.append(ExportError(symbol, target_date, message[len(prefix):] if message.startswith(prefix) else message))

    if rebuild_history and errors:
        first_error = errors[0]
        raise BaseLotExportError(
            f"history rebuild failed for {len(errors)} symbol/date cell(s) "
            f"(first: {first_error.symbol} {first_error.actualization_date}: {first_error.message}); "
            "workbook unchanged"
        )

    try:
        workbook = load_workbook(workbook_path, data_only=False)
    except (OSError, ValueError) as error:
        raise BaseLotExportError(f"cannot open workbook {workbook_path}: {error}") from error
    try:
        current_stat = workbook_path.stat()
        current_signature = current_stat.st_mtime_ns, current_stat.st_size
        if current_signature != source_signature:
            raise BaseLotExportError("workbook changed while the export was running; no values were published")
        _validate_ordinary_workbook(workbook, workbook_path)
        if sheet_name not in workbook.sheetnames:
            raise BaseLotExportError(f"workbook is missing sheet {sheet_name!r}")
        sheet = workbook[sheet_name]
        _migrate_previous_layout(sheet)
        if history_action == "initialize":
            _initialize_history(sheet, target_dates)
        elif history_action == "rebuild":
            _reset_history(sheet, target_dates, history_dates, preserve_existing=not rebuild_history)
        elif history_action == "rotate":
            _rotate_history(sheet, actualization_date)
            _clear_output_filters(sheet)
        elif history_action == "retry":
            for target_date in calculation_symbols:
                column = HISTORY_FIRST_COLUMN + history_dates.index(target_date)
                _clear_output_filters(sheet, column, column)
        values = {(row.symbol, row.actualization_date): row.base_lot_usdt for row in rows}
        error_values = {(error.symbol, error.actualization_date): error.message for error in errors}
        if history_action in {"initialize", "rebuild"}:
            columns_for_date = {target_date: HISTORY_FIRST_COLUMN + offset for offset, target_date in enumerate(target_dates)}
        else:
            columns_for_date = {target_date: HISTORY_FIRST_COLUMN + target_dates.index(target_date) for target_date in calculation_symbols}
        for target_date, write_symbols in calculation_symbols.items():
            column = columns_for_date[target_date]
            if history_action in {"initialize", "rebuild"}:
                sheet.cell(1, column).value = target_date
                sheet.cell(1, column).number_format = "yyyy-mm-dd"
            for symbol, row_numbers in symbol_rows.items():
                if symbol not in write_symbols:
                    continue
                for row_number in row_numbers:
                    cell = sheet.cell(row=row_number, column=column)
                    key = (symbol, target_date)
                    if key in values:
                        cell.value = values[key]
                        cell.fill = _CLEAR_FILL
                        cell.number_format = LOT_NUMBER_FORMAT
                    else:
                        cell.value = f"ERROR: {symbol}: {error_values[key]}"
                        cell.number_format = "General"
                        cell.fill = ERROR_FILL
        for column in range(HISTORY_FIRST_COLUMN, HISTORY_FIRST_COLUMN + HISTORY_DAYS):
            header = sheet.cell(1, column)
            if isinstance(header.value, str) and _date_header(header.value) is not None:
                header.value = _date_header(header.value)
                header.number_format = "yyyy-mm-dd"
        _write_averages(sheet, symbol_rows, step)
        _clear_output_filters(sheet, 1, 0)  # empty range: keep every criterion, extend the filter over A:M and beyond
        workbook.defined_names[MODEL_NAME] = DefinedName(MODEL_NAME, attr_text=f'"{model}"')
        _atomic_save_workbook(workbook, workbook_path)
    finally:
        workbook.close()

    successful_cell_count = sum(len(symbol_rows[row.symbol]) for row in rows)
    return ExportResult(workbook_path, successful_cell_count, tuple(rows), tuple(errors), actualization_date)


def _anchor(value: str) -> datetime:
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError as error:
        raise argparse.ArgumentTypeError("anchor must be an ISO-8601 timestamp") from error
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        raise argparse.ArgumentTypeError("anchor must include a timezone")
    return parsed


def registry_lot_source(config_path: str | Path, *, k: object = DEFAULT_K,
                        round_down_usdt: object = DEFAULT_ROUND_DOWN_USDT) -> Callable[[str, date], ExportRow]:
    """Cells from the pair registry: the READY window ending on the last weekday before the cell's date."""
    import duckdb

    from mrs3.pair_registry.liquidity import last_weekday_before, liquidity_on
    from mrs3.pair_registry.store import DATABASE, RegistryError, load_registry_settings

    try:
        settings = load_registry_settings(Path(config_path))
    except RegistryError as error:
        raise BaseLotExportError("pair registry settings are unavailable") from error
    database = settings.root / DATABASE
    if not database.is_file():
        raise BaseLotExportError("pair registry not found: run the registry sync first")
    coefficient = _decimal(k, "k", positive=True)
    step = _decimal(round_down_usdt, "round_down_usdt", positive=True)
    if settings.liquidity_k != coefficient or settings.liquidity_round_down != step:
        raise BaseLotExportError("the registry's liquidity K or step differ from the export's; workbook unchanged")

    def source(symbol: str, target_date: date) -> ExportRow:
        window_end = last_weekday_before(target_date)
        try:
            with duckdb.connect(str(database), read_only=True) as con:
                row = liquidity_on(con, "bybit", symbol, window_end)
        except duckdb.Error as error:              # a sync or a Panel write holds the registry: an error cell
            raise BaseLotExportError("pair registry is busy: run the export again") from error
        if row is None:
            raise BaseLotExportError(f"no READY registry window ending {window_end.isoformat()}")
        if Decimal(str(row["k"])) != coefficient or Decimal(str(row["step"])) != step:
            raise BaseLotExportError("registry K or step differ from the export's")
        return ExportRow(symbol, target_date, Decimal(str(row["v25"])), Decimal(str(row["a15"])),
                         Decimal(str(row["base_lot"])), str(row["inputs_sha256"]),
                         tuple((str(day), str(digest)) for day, digest in row["window_files"]))

    return source


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", choices=("registry", "minute-files"), default="registry",
                        help="registry: LOT_MODEL_V2 windows of the pair registry; minute-files: LOT_MODEL_V1")
    parser.add_argument("--config", type=Path, default=Path(__file__).resolve().parents[1] / "config.local.json",
                        help="local config with pair_registry (registry source)")
    parser.add_argument("--input", dest="workbook_path", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--data-root", type=Path, default=DEFAULT_DATA_ROOT)
    parser.add_argument("--sheet", dest="sheet_name", default=DEFAULT_SHEET)
    parser.add_argument("--anchor-created-at", type=_anchor, default=None)
    parser.add_argument("--k", default=DEFAULT_K)
    parser.add_argument("--round-down-usdt", default=DEFAULT_ROUND_DOWN_USDT)
    parser.add_argument("--publication-lag-hours", type=int, default=DEFAULT_PUBLICATION_LAG_HOURS)
    parser.add_argument("--workers", type=int, default=DEFAULT_WORKERS)
    parser.add_argument(
        "--rebuild-history",
        action="store_true",
        help="recalculate all ten date columns; keep the workbook unchanged if any cell fails",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    try:
        options = vars(build_parser().parse_args(argv))
        source, config = options.pop("source"), options.pop("config")
        if source == "registry":
            options["lot_source"] = registry_lot_source(config, k=options["k"], round_down_usdt=options["round_down_usdt"])
        result = run_export(**options)
    except Exception as error:  # CLI boundary: every failure is explicit and nonzero.
        print(f"ERROR: {type(error).__name__}: {error}")
        return 1
    if result.skipped_existing_date:
        print(f"SKIPPED: {result.actualization_date} already exists in the date history")
        return 0
    for error in result.errors:
        print(f"ERROR: {error.symbol} {error.actualization_date}: {error.message}")
    if result.errors:
        print(f"FAILED: {len(result.errors)} symbol/date cell(s) failed; wrote {result.rows_written} base lots to {result.workbook_path}")
        return 1
    print(f"OK: wrote {result.rows_written} dated base lots to {result.workbook_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
