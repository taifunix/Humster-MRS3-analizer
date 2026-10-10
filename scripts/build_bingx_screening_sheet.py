"""Create or top up the «Скрининг BingX» sheet of the liquidity registry.

The sheet has the same form as «Скрининг»; it lists only the registry pairs
that exist on BingX, with the BingX ticker (e.g. PANWUSDT, AMDSTOCKUSDT ->
AMDUSDT) and the BingX listing date. Other columns are left for the operator.
Contract: docs/specs/2026-10-10-bingx-data-downloader.md.
"""

from __future__ import annotations

import argparse
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
import shutil
import sys
from typing import Mapping, Sequence

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

if __package__ in {None, ""}:  # Allow ``python scripts/build_bingx_screening_sheet.py``.
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from mrs3.bingx_downloader import BingXClient, BingXDownloadError, normalize_ticker, resolve_contracts  # noqa: E402
from mrs3.screener.errors import ScreenerEvaluationError  # noqa: E402
from mrs3.screener.registry import _atomic_save_workbook  # noqa: E402

SOURCE_SHEET = "Скрининг"
TARGET_SHEET = "Скрининг BingX"
PAIR_COLUMN = "Пара"
LISTING_COLUMN = "Листинг"
DEFAULT_REGISTRY = Path(__file__).resolve().parents[1] / "input" / "bybit_tradfi_liquidity.xlsx"


class BingXSheetError(RuntimeError):
    pass


@dataclass(slots=True)
class SheetReport:
    added: int = 0
    existing: int = 0
    missing: list[str] = field(default_factory=list)
    backup: Path | None = None


def _header(sheet) -> list[object]:
    return [cell.value for cell in sheet[1]]


def _frozen_cell(sheet) -> str | None:
    """Freeze cell from the pane split; ``freeze_panes`` returns the scroll position instead."""
    pane = sheet.sheet_view.pane
    if pane is None or pane.state not in {"frozen", "frozenSplit"}:
        return None
    columns, rows = int(pane.xSplit or 0), int(pane.ySplit or 0)
    if not columns and not rows:
        return None
    return f"{get_column_letter(columns + 1)}{rows + 1}"


def _copy_form(source, target) -> None:
    """Carry over what copy_worksheet does not: filter, freeze, validations, conditional formats."""
    target.auto_filter.ref = source.auto_filter.ref
    target.freeze_panes = _frozen_cell(source)
    for validation in source.data_validations.dataValidation:
        target.add_data_validation(copy(validation))
    for formatting in source.conditional_formatting:
        for rule in formatting.rules:
            target.conditional_formatting.add(str(formatting.sqref), rule)


def _new_target(workbook, source):
    target = workbook.copy_worksheet(source)
    target.title = TARGET_SHEET
    for row in target.iter_rows(min_row=2):
        for cell in row:
            cell.value = None
    _copy_form(source, target)
    active = workbook.active
    workbook.move_sheet(target, offset=workbook.sheetnames.index(SOURCE_SHEET) + 1 - workbook.sheetnames.index(TARGET_SHEET))
    workbook.active = active  # keep the sheet the workbook opened on
    return target


def build_bingx_screening_sheet(path: Path, contracts: Sequence[Mapping[str, object]]) -> SheetReport:
    path = Path(path)
    try:
        workbook = load_workbook(path)
    except (OSError, ValueError) as error:
        raise BingXSheetError(f"cannot open registry, is it open in Excel? ({error})") from None
    if SOURCE_SHEET not in workbook.sheetnames:
        raise BingXSheetError(f"registry has no {SOURCE_SHEET!r} sheet")
    source = workbook[SOURCE_SHEET]
    header = _header(source)
    try:
        pair_column = header.index(PAIR_COLUMN) + 1
        listing_column = header.index(LISTING_COLUMN) + 1
    except ValueError:
        raise BingXSheetError(f"{SOURCE_SHEET!r} sheet must have {PAIR_COLUMN!r} and {LISTING_COLUMN!r} columns") from None

    pairs = [
        str(value).strip()
        for value in (source.cell(row=index, column=pair_column).value for index in range(2, source.max_row + 1))
        if value is not None and str(value).strip()
    ]
    resolved, _errors = resolve_contracts(pairs, contracts)

    created = TARGET_SHEET not in workbook.sheetnames
    if not created:
        target = workbook[TARGET_SHEET]
        if _header(target)[: len(header)] != header:
            raise BingXSheetError(f"{TARGET_SHEET!r} sheet header differs from {SOURCE_SHEET!r}")
    else:
        target = _new_target(workbook, source)

    present = {}
    last_row = 1  # last row with any operator value, so appended pairs never land on it
    for index in range(2, target.max_row + 1):
        values = [target.cell(row=index, column=column).value for column in range(1, len(header) + 1)]
        if any(value is not None and str(value).strip() for value in values):
            last_row = index
        value = values[pair_column - 1]
        if value is not None and str(value).strip():
            present[str(value).strip()] = index

    report = SheetReport()
    for pair in pairs:
        try:
            contract = resolved.get(normalize_ticker(pair))
        except BingXDownloadError:
            contract = None
        if contract is None:
            report.missing.append(pair)
            continue
        if contract.file_ticker in present:
            report.existing += 1
            continue
        last_row += 1
        for column in range(1, len(header) + 1):
            cell = target.cell(row=last_row, column=column)
            if not cell.has_style:
                template = source.cell(row=2, column=column)
                cell._style = copy(template._style)
        target.cell(row=last_row, column=pair_column, value=contract.file_ticker)
        if contract.launch_ms > 0:  # unknown launchTime: leave «Листинг» empty, never 1970-01-01
            listing = datetime.fromtimestamp(contract.launch_ms / 1000, timezone.utc).date()
            listing_cell = target.cell(row=last_row, column=listing_column, value=datetime(listing.year, listing.month, listing.day))
            if listing_cell.number_format in ("General", "yyyy-mm-dd h:mm:ss"):
                listing_cell.number_format = "yyyy-mm-dd"

        present[contract.file_ticker] = last_row
        report.added += 1

    if report.added == 0 and not created:
        return report  # Nothing new: leave the file untouched.
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    report.backup = path.with_name(f"{path.stem}.backup-{stamp}{path.suffix}")
    shutil.copy2(path, report.backup)
    try:
        _atomic_save_workbook(workbook, path)
    except ScreenerEvaluationError as error:
        raise BingXSheetError(str(error)) from None
    return report


def main(argv: Sequence[str] | None = None, *, contracts: Sequence[Mapping[str, object]] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--registry", type=Path, default=DEFAULT_REGISTRY, help="liquidity registry .xlsx")
    args = parser.parse_args(argv)
    try:
        if contracts is None:
            contracts = BingXClient().contracts()
        report = build_bingx_screening_sheet(args.registry, contracts)
    except (BingXDownloadError, BingXSheetError) as error:
        print(f"Ошибка: {error}", file=sys.stderr)
        return 2
    print(f"лист: {TARGET_SHEET}")
    print(f"добавлено пар: {report.added}")
    print(f"уже были в листе: {report.existing}")
    print(f"нет на BingX ({len(report.missing)}): {', '.join(report.missing)}")
    if report.backup is not None:
        print(f"резервная копия: {report.backup.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
