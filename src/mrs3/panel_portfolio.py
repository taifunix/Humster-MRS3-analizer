"""Small, local-only Panel adapter for the Portfolio Optimizer Panel flows.

The adapter owns HTTP-facing validation and lifecycle plumbing.  Portfolio
algorithms remain injected (and, by default, are the package primitives).
Stage 2 accepts only an injected local tester service and committed packages.
"""

from __future__ import annotations

from collections.abc import Callable, Mapping, Sequence
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import asyncio
import base64
import copy
import gzip
import hashlib
import io
import inspect
import json
import math
import os
from pathlib import Path
import re
import shutil
import tempfile
import threading
import time
from typing import Any
from uuid import uuid4

import duckdb
import pandas as pd

from .audit import normalize_xlsx_workbook, write_audit_workbook
from .panel_jobs import PanelJobError, PanelJobRegistry
from .panel_testing import mrs3_tester_config_template
from .runner.files import file_fingerprint, read_stable_file
from .runner.results import CORE_METRICS, ResultParseError, WizardResult, load_wizard_results
from .portfolio.config import (
    POLICY_VERSION,
    SCHEMA_VERSION,
    PortfolioConfigError,
    load_portfolio_config,
    migrate_portfolio_config_document,
)
from .config import load_duckdb_import_settings
from .performance_v2_optimizer import prepare_current_optimizer_inputs
from .portfolio.input import apply_finalist_cutoff, read_current_finalists
from .portfolio.adapter import (
    CAMPAIGN_CONTRACT_VERSION,
    CAMPAIGN_SEARCH_MODE,
    CAMPAIGN_WEIGHTED_ALGO_VERSION,
    CampaignContractError,
    validate_campaign_contract,
)


STAGES = (
    "VALIDATE_SNAPSHOT",
    "LOAD_FINALISTS",
    "SELECT_CANDIDATES",
    "GENERATE_VARIANTS",
    "VALIDATE_VARIANTS",
    "BUILD_WORKBOOK",
    "PUBLISH_RESULTS",
)
STAGE2_STAGES = ("PREPARE", "FILL_PREBUILT", "START", "READBACK")
_TERMINAL = frozenset({"COMMITTED", "CANCELLED", "FAILED"})
_SNAPSHOT_SCHEMA = "portfolio-campaign-input-v1"
_SNAPSHOT_NAME = "campaign-input.json.gz"
_SNAPSHOT_MAX_UNCOMPRESSED = 512 * 1024 * 1024
_SNAPSHOT_MIGRATION_BACKUP = ".panel-jobs.snapshot-migration.bak"
_SNAPSHOT_MIGRATION_RESERVE = 64 * 1024 * 1024
PORTFOLIO_SNAPSHOT_UNSERIALIZABLE = "PORTFOLIO_SNAPSHOT_UNSERIALIZABLE"
_GEOMETRY_FIELDS = ("timeframe", "close_ma_len", "order_count", "strategy_orders")
_SECRET = re.compile(r"(?:password|passwd|secret|token|credential|api[_-]?key|private[_-]?key)", re.I)
_PATH = re.compile(
    r"(?<![\w])[A-Za-z]:[\\/][^\s,;)]*"
    r"|(?<![\w])\\\\[^\s,;)]*"
    r"|(?<![\w/:])/(?![/\s])(?=[A-Za-z0-9_.-])[^\s,;)]*"
    r"|(?<![\w])(?:[^\\/\s,;)]+[\\/])+\.\.(?:[\\/][^\s,;)]*)?"
    r"|(?<![\w])\.\.(?:[\\/][^\s,;)]*)+"
)
_WEIGHTED_TEMPLATE_PATH = Path(__file__).resolve().parents[2] / "templates" / "strategies" / "portfolio-weighted-mrs" / "base.json"
SUMMARY_HEADERS = ("Key", "Value")
FINALIST_HEADERS = (
    "Campaign ID", "Strategy ID", "Result ID", "Pair", "Direction", "User Status", "User Rank",
    "Effective Maximum", "Selection Status", "Selection Reason",
)
PORTFOLIO_HEADERS = (
    "Campaign ID", "Candidate ID", "Profile", "Scheduling Position", "Scheduling Key ID",
    "Scheduling Score (Individual/Margin; Not Portfolio PnL)", "Member Count", "Pair Count", "Limiter",
    "Maximum Individual DD %", "Minimum Free Margin Reserve %", "Maximum Account MM Load %", "Gate Result",
    "Blocking Reasons", "User Decision", "User Test Priority", "User Comment",
    "Search Mode", "Evaluations", "Pretest PnL USDT", "Pretest DD USDT", "Pretest DD %",
    "Pretest Recovery", "Pretest Reserve USDT", "Pretest Reserve %", "Sizing k",
    "Tested Size USDT", "Actual Size USDT", "Capacity Basis", "Pretest Period", "Refinement",
    "Sizing k1", "Corrective Reduction Applied", "Daily PRETEST Rank", "Final PRETEST Rank",
)
MEMBER_HEADERS = (
    "Campaign ID", "Candidate ID", "Profile", "Member Ordinal", "Strategy ID", "Result ID", "Pair",
    "Direction", "User Rank", "Scalar %", "Quantity", "Leverage", "Notional USDT", "Estimated Individual DD USDT",
    "Estimated Individual DD %", "Liquidity Scalar Ceiling %", "Calculated Initial Margin USDT", "Gate Result", "Reasons",
    "Capacity Status", "7d Available Days", "7d Mean Minute Turnover USDT", "7d Position Cap USDT",
    "5d Available Days", "5d Mean Minute Turnover USDT", "5d Analytic Cap USDT", "Spread Status",
    "Mean p95 Spread bps", "Sizing Digest", "Capacity Digest", "Reference Digest",
    "Finalist", "Side", "x USDT", "C USDT", "q", "Priority", "max_balance", "Source Scale", "Hold90", "Warnings",
)
EXCLUDED_HEADERS = (
    "Campaign ID", "Scope", "Object ID", "Pair", "Direction", "Profile", "Stage", "Gate Result",
    "Portfolio Reason", "Message",
)
METADATA_HEADERS = ("Key", "Value")
PROFILE_STATUS_HEADERS = (
    "Campaign ID", "Profile", "Status", "Max Candidates", "Evaluations", "Evaluation Budget",
    "Pretest Period", "Coverage %", "Blockers", "Metric Basis", "Joint Metrics",
)
WEIGHTED_SUMMARY_HEADERS = ("Key", "Value")
# Stage 1 workbook uses operator-facing columns; legacy sheets remain unchanged.
WEIGHTED_VARIANT_HEADERS = (
    "№", "ID", "Профиль", "Поз.", "Состав", "Банк\nнасыщ., USDT", "Целевой\nбанк, USDT",
    "Мин. банк DD\nистории, USDT", "Банк DD P95\nстресса, USDT", "Банк\nлимитов, USDT", "PnL 30д,\nUSDT", "MaxDD SUM,\nUSDT", "DD истории,\n%",
    "CDaR 20%\nUSDT · ц/н", "CDaR 10%\nUSDT · ц/н", "IM\nUSDT · ц/н", "MM\nUSDT · ц/н", "Номинал\nпортф., USDT",
)
WEIGHTED_MEMBER_HEADERS = (
    "№ вар.", "ID", "Профиль", "Пара", "Сторона", "Strategy ID", "Result ID", "Позиция,\nUSDT",
    "Множитель\nпары, %", "Max balance,\nUSDT", "X, %", "Y, %", "Z, %", "W, %", "Плечо", "Лимит ликв.,\nUSDT", "TF", "User Rank",
    "Source PnL,\nUSDT", "Source MaxDD,\nUSDT", "Source MaxDD,\n%", "ORD_N", "Исп. ликв.\nx/C, %", "Инд. MaxDD,\nUSDT",
)
WEIGHTED_FINALIST_HEADERS = (
    "Campaign\nID", "Strategy\nID", "Result\nID", "Пара", "Сторона",
    "User\nstatus", "User\nrank", "Лимит", "Статус\nотбора", "Причина\nотбора",
)
WEIGHTED_EXCLUDED_HEADERS = (
    "Campaign\nID", "Объект", "Object\nID", "Пара", "Сторона",
    "Профиль", "Этап", "Результат", "Причина", "Сообщение",
)


class PortfolioPanelError(ValueError):
    """A typed, client-safe Panel error."""

    def __init__(self, code: str, message: str | None = None, *, status: int = 400, field_errors: Sequence[Mapping[str, str]] = ()) -> None:
        self.code = code
        self.status = status
        self.field_errors = [dict(item) for item in field_errors]
        super().__init__(message or code)


class _PortfolioProgressReporter:
    """Small in-memory progress state; it deliberately stores no event history."""

    def __init__(self, clock: Callable[[], float] = time.monotonic) -> None:
        self._clock = clock
        self._lock = threading.RLock()
        self._states: dict[str, dict[str, Any]] = {}

    def start(self, job_id: str) -> None:
        now = self._clock()
        with self._lock:
            self._states[job_id] = {
                "started": now,
                "last_emit": 0.0,
                "last_update": now,
                "last_heartbeat": now,
                "last_persist": None,
                "substage": None,
                "unit": "",
                "completed": 0,
                "total": None,
                "detail": "",
                "inconsistent": False,
            }

    @staticmethod
    def _event(event: Mapping[str, Any]) -> dict[str, Any]:
        detail = str(event.get("detail", "")).encode("utf-8")[:160].decode("utf-8", "ignore")
        total = event.get("total")
        return {
            "substage": str(event.get("substage", "")),
            "unit": str(event.get("unit", "")),
            "completed": max(0, int(event.get("completed", 0))),
            "total": None if total is None else max(0, int(total)),
            "detail": detail,
        }

    def _snapshot(self, state: Mapping[str, Any], now: float) -> dict[str, Any]:
        elapsed = max(0.0, now - float(state["started"]))
        completed = int(state["completed"])
        total = state["total"]
        reliable = not state["inconsistent"] and isinstance(total, int) and total > 0 and 2 <= completed <= total and elapsed >= 2.0
        eta = (elapsed / completed) * (total - completed) if reliable else None
        return {
            "substage": state["substage"],
            "unit": state["unit"],
            "completed": completed,
            "total": total if not state["inconsistent"] else None,
            "detail": state["detail"],
            "elapsed_seconds": elapsed,
            "eta_seconds": max(0.0, eta) if eta is not None else None,
            "heartbeat_age_seconds": max(0.0, now - float(state["last_heartbeat"])),
            "last_update_age_seconds": max(0.0, now - float(state["last_update"])),
            "indeterminate": bool(state["inconsistent"] or not (isinstance(total, int) and total > 0)),
        }

    def emit(self, job_id: str, event: Mapping[str, Any], *, force: bool = False) -> tuple[dict[str, Any], bool]:
        now = self._clock()
        parsed = self._event(event)
        with self._lock:
            state = self._states.get(job_id)
            if state is None:
                self.start(job_id)
                state = self._states[job_id]
            changed_substage = parsed["substage"] != state["substage"]
            if not force and not changed_substage and now - float(state["last_emit"]) < 0.25:
                return self._snapshot(state, now), False
            previous_total = state["total"]
            if previous_total is not None and parsed["total"] is not None and parsed["total"] != previous_total:
                state["inconsistent"] = True
            if parsed["total"] is not None and parsed["completed"] > parsed["total"]:
                state["inconsistent"] = True
            state.update(parsed)
            state["last_emit"] = now
            state["last_update"] = now
            state["last_heartbeat"] = now
            last_persist = state["last_persist"]
            persist = bool(
                force
                or (changed_substage and (last_persist is None or now - float(last_persist) >= 2.0))
            )
            if persist:
                state["last_persist"] = now
            return self._snapshot(state, now), persist

    def heartbeat(self, job_id: str) -> tuple[dict[str, Any] | None, bool]:
        now = self._clock()
        with self._lock:
            state = self._states.get(job_id)
            if state is None:
                return None, False
            state["last_heartbeat"] = now
            last_persist = state["last_persist"]
            persist = last_persist is not None and now - float(last_persist) >= 10.0
            if persist:
                state["last_persist"] = now
            return self._snapshot(state, now), persist

    def snapshot(self, job_id: str) -> dict[str, Any] | None:
        now = self._clock()
        with self._lock:
            state = self._states.get(job_id)
            return self._snapshot(state, now) if state is not None else None

    def stop(self, job_id: str) -> None:
        with self._lock:
            self._states.pop(job_id, None)


def _now() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def _digest(value: bytes | str) -> str:
    return hashlib.sha256(value if isinstance(value, bytes) else value.encode("utf-8")).hexdigest()


def _report_folder_snapshot(folder: Path) -> dict[str, tuple[int, int, str]]:
    if not folder.exists():
        return {}
    if folder.is_symlink() or not folder.is_dir():
        raise PortfolioPanelError("PORTFOLIO_JOB_STAGE2_RESULT_INVALID", "tester report folder is invalid", status=500)
    snapshot: dict[str, tuple[int, int, str]] = {}
    for path in folder.iterdir():
        if path.is_symlink() or not path.is_file():
            continue
        fingerprint = file_fingerprint(path)
        if fingerprint is not None:
            snapshot[path.name] = fingerprint
    return snapshot


def _fresh_report_fingerprint(
    folder: Path,
    report_name: str,
    before: Mapping[str, tuple[int, int, str]],
    start_wall_ns: int,
) -> tuple[int, int, str] | None:
    if Path(report_name).name != report_name or not report_name.casefold().endswith(".html"):
        raise PortfolioPanelError("PORTFOLIO_JOB_STAGE2_RESULT_INVALID", "tester report name is invalid", status=500)
    if folder.is_symlink() or not folder.is_dir():
        return None
    now_ns = time.time_ns()
    fresh: dict[str, tuple[int, int, str]] = {}
    for path in folder.iterdir():
        if not path.name.casefold().endswith(".html"):
            continue
        if path.is_symlink() or not path.is_file():
            raise PortfolioPanelError("PORTFOLIO_JOB_STAGE2_RESULT_INVALID", "tester report is not a regular file", status=500)
        fingerprint = file_fingerprint(path)
        if fingerprint is None:
            continue
        if fingerprint[0] > now_ns:
            raise PortfolioPanelError("PORTFOLIO_JOB_STAGE2_RESULT_INVALID", "tester report timestamp is invalid", status=500)
        if fingerprint[0] >= start_wall_ns and before.get(path.name) != fingerprint:
            fresh[path.name] = fingerprint
    if not fresh:
        return None
    if set(fresh) != {report_name}:
        raise PortfolioPanelError("PORTFOLIO_JOB_STAGE2_RESULT_INVALID", "tester report set is ambiguous", status=500)
    return fresh[report_name]


def _validate_pretest_period(value: Any) -> None:
    if not isinstance(value, Mapping) or set(value) != {"start_utc", "end_utc"}:
        raise ValueError("pretest period shape is invalid")
    parsed: dict[str, datetime] = {}
    for key in ("start_utc", "end_utc"):
        text = value[key]
        if not isinstance(text, str):
            raise ValueError("pretest period timestamp is invalid")
        try:
            timestamp = datetime.fromisoformat(text.replace("Z", "+00:00"))
        except ValueError as error:
            raise ValueError("pretest period timestamp is invalid") from error
        if timestamp.tzinfo is None or timestamp.utcoffset() != timedelta(0):
            raise ValueError("pretest period timestamp is not UTC")
        timestamp = timestamp.astimezone(timezone.utc)
        if timestamp.hour or timestamp.minute or timestamp.second or timestamp.microsecond:
            raise ValueError("pretest period timestamp is not a UTC day boundary")
        if text != timestamp.isoformat(timespec="seconds").replace("+00:00", "Z"):
            raise ValueError("pretest period timestamp is not canonical")
        parsed[key] = timestamp
    if parsed["start_utc"] >= parsed["end_utc"] or parsed["end_utc"] - parsed["start_utc"] < timedelta(days=1):
        raise ValueError("pretest period is too short")


def _candidate_required_bank(candidate: Mapping[str, Any]) -> Decimal:
    metrics = candidate.get("metrics")
    raw = metrics.get("required_bank_usdt") if isinstance(metrics, Mapping) else None
    if isinstance(raw, bool) or raw is None:
        raise ValueError("candidate required bank is invalid")
    try:
        bank = Decimal(str(raw))
    except (InvalidOperation, TypeError, ValueError) as error:
        raise ValueError("candidate required bank is invalid") from error
    if not bank.is_finite() or bank <= 0:
        raise ValueError("candidate required bank is invalid")
    return bank


def _stage2_material(candidate: Mapping[str, Any]) -> dict[str, Any]:
    """Render the exact tester inputs and receipt from one committed candidate."""
    candidate_id = candidate.get("candidate_id")
    if (
        not isinstance(candidate_id, str)
        or re.fullmatch(r"[0-9a-f]{64}", candidate_id) is None
        or candidate.get("identity") != candidate_id
    ):
        raise ValueError("candidate id is unsafe")
    required_bank = _candidate_required_bank(candidate)
    payloads = candidate.get("strategy_payloads")
    if isinstance(payloads, (str, bytes)) or not isinstance(payloads, Sequence) or len(payloads) < 2:
        raise ValueError("candidate strategies are invalid")
    strategy_jsons: dict[str, str] = {}
    sizing_by_name: dict[str, dict[str, Any]] = {}
    max_balance_by_name: dict[str, Any] = {}
    for payload in payloads:
        if not isinstance(payload, Mapping):
            raise ValueError("candidate strategy wrapper is invalid")
        facts = payload.get("facts")
        if not isinstance(facts, Mapping) or any(
            isinstance(facts.get(key), bool) or facts.get(key) is None for key in ("B", "C", "q", "x")
        ):
            raise ValueError("candidate sizing evidence is incomplete")
        try:
            numeric = {key: Decimal(str(facts[key])) for key in ("B", "C", "q", "x")}
        except (InvalidOperation, TypeError, ValueError) as error:
            raise ValueError("candidate sizing evidence is invalid") from error
        if any(not value.is_finite() or value <= 0 for value in numeric.values()) or numeric["B"] != required_bank:
            raise ValueError("candidate sizing evidence is invalid")
        strategy = payload.get("strategy")
        basic = strategy.get("basic") if isinstance(strategy, Mapping) else None
        name = strategy.get("name") if isinstance(strategy, Mapping) else None
        if (
            not isinstance(strategy, Mapping)
            or not isinstance(basic, Mapping)
            or not isinstance(name, str)
            or re.fullmatch(r"[A-Z0-9_]{1,128}", name) is None
            or name in strategy_jsons
            or isinstance(basic.get("max_balance"), bool)
            or basic.get("max_balance") is None
        ):
            raise ValueError("candidate strategy is invalid")
        try:
            max_balance = Decimal(str(basic["max_balance"]))
        except (InvalidOperation, TypeError, ValueError) as error:
            raise ValueError("candidate max_balance is invalid") from error
        if not max_balance.is_finite() or max_balance <= 0:
            raise ValueError("candidate max_balance is invalid")
        strategy_jsons[name] = _json(strategy) + "\n"
        sizing_by_name[name] = {key: facts[key] for key in ("B", "C", "q", "x")}
        max_balance_by_name[name] = basic["max_balance"]
    if len(strategy_jsons) < 2:
        raise ValueError("candidate strategies are invalid")

    period = candidate.get("pretest_period")
    _validate_pretest_period(period)
    start = datetime.fromisoformat(period["start_utc"].replace("Z", "+00:00")).astimezone(timezone.utc)
    end = datetime.fromisoformat(period["end_utc"].replace("Z", "+00:00")).astimezone(timezone.utc) - timedelta(days=1)
    if start.date() > end.date():
        raise ValueError("tester period is invalid")
    template = json.loads(mrs3_tester_config_template().read_text(encoding="utf-8"))
    if (
        not isinstance(template, dict)
        or type(template.get("use_runs")) is not bool
        or template.get("use_runs") is not False
        or type(template.get("parameter_mining")) is not list
        or template.get("parameter_mining") != []
    ):
        raise ValueError("tester template is invalid")
    initial_balance: int | float
    if required_bank == required_bank.to_integral_value():
        initial_balance = int(required_bank)
    else:
        initial_balance = float(required_bank)
        if not math.isfinite(initial_balance) or Decimal(str(initial_balance)) != required_bank:
            raise ValueError("tester balance is invalid")
    template.update({
        "name_comment": candidate_id,
        "StartDate": start.date().isoformat(),
        "EndDate": end.date().isoformat(),
        "InitialBalance": initial_balance,
        "single_mode": False,
        "UpdateData": False,
    })
    tester_config_json = _json(template) + "\n"
    members = candidate.get("members")
    if isinstance(members, (str, bytes)) or not isinstance(members, Sequence) or any(
        not isinstance(member, Mapping) for member in members
    ):
        raise ValueError("candidate members are invalid")
    receipt = {
        "candidate_order": candidate.get("order"),
        "profile": candidate.get("profile"),
        "member_identities": [
            {key: member.get(key) for key in ("symbol", "side", "strategy_id", "result_id")}
            for member in members
        ],
        "tester_config_sha256": hashlib.sha256(tester_config_json.encode("utf-8")).hexdigest(),
        "strategy_manifest": [
            {
                "filename": f"{name}.json",
                "size": len(payload.encode("utf-8")),
                "sha256": hashlib.sha256(payload.encode("utf-8")).hexdigest(),
            }
            for name, payload in sorted(strategy_jsons.items())
        ],
        "expected_sizing": [
            {"strategy_name": name, **sizing_by_name[name]} for name in sorted(strategy_jsons)
        ],
        "expected_max_balance": [
            {"strategy_name": name, "max_balance": max_balance_by_name[name]} for name in sorted(strategy_jsons)
        ],
    }
    return {
        "candidate_id": candidate_id,
        "expected_names": list(strategy_jsons),
        "pretest_period": _plain(period),
        "tester_config_json": tester_config_json,
        "strategy_jsons": strategy_jsons,
        "receipt": receipt,
        "required_bank_usdt": format(required_bank, "f"),
    }


def _load_weighted_template() -> tuple[dict[str, Any], str]:
    try:
        raw = _WEIGHTED_TEMPLATE_PATH.read_bytes()
        template = json.loads(raw.decode("utf-8"))
        if not isinstance(template, Mapping):
            raise ValueError("template must be an object")
        frozen = _plain(template)
        json.dumps(frozen, ensure_ascii=False, sort_keys=True, allow_nan=False)
        return frozen, _digest(raw)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, TypeError, ValueError):
        raise PortfolioPanelError("PORTFOLIO_TEMPLATE_INVALID", "portfolio strategy template is invalid", status=422) from None


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def _snapshot_normalize(value: Any, path: str = "$") -> Any:
    """Normalize only values accepted by the snapshot wire contract."""
    if value is None or isinstance(value, (bool, str)):
        return value
    if isinstance(value, Decimal):
        if not value.is_finite():
            raise PortfolioPanelError(PORTFOLIO_SNAPSHOT_UNSERIALIZABLE, f"snapshot field {path} is not finite", status=422)
        return format(value, "f")
    if type(value) is int:
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise PortfolioPanelError(PORTFOLIO_SNAPSHOT_UNSERIALIZABLE, f"snapshot field {path} is not finite", status=422)
        return value
    if isinstance(value, Mapping):
        result: dict[str, Any] = {}
        for key, item in value.items():
            if not isinstance(key, str):
                raise PortfolioPanelError(PORTFOLIO_SNAPSHOT_UNSERIALIZABLE, f"snapshot field {path} has a non-string key", status=422)
            result[key] = _snapshot_normalize(item, f"{path}.{key}")
        return result
    if isinstance(value, (list, tuple)):
        return [_snapshot_normalize(item, f"{path}[{index}]") for index, item in enumerate(value)]
    raise PortfolioPanelError(PORTFOLIO_SNAPSHOT_UNSERIALIZABLE, f"snapshot field {path} has an unsupported value", status=422)


def _snapshot_bytes(campaign: Mapping[str, Any]) -> tuple[bytes, bytes, str]:
    """Return canonical JSON, deterministic gzip, and the raw digest."""
    normalized = _snapshot_normalize(campaign)
    raw = json.dumps(normalized, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")
    output = io.BytesIO()
    with gzip.GzipFile(fileobj=output, mode="wb", compresslevel=6, mtime=0) as stream:
        stream.write(raw)
    return raw, output.getvalue(), hashlib.sha256(raw).hexdigest()


def _geometry_identity(row: Any) -> tuple[str, str, int, int] | None:
    if not isinstance(row, Mapping):
        return None
    symbol, side = row.get("symbol"), row.get("side")
    strategy_id, result_id = row.get("strategy_id"), row.get("result_id")
    if (
        not isinstance(symbol, str) or not symbol.strip()
        or not isinstance(side, str) or side.strip().upper() not in {"LONG", "SHORT"}
        or type(strategy_id) is not int or type(result_id) is not int
    ):
        return None
    return symbol.strip().upper(), side.strip().upper(), strategy_id, result_id


def _repair_legacy_geometry(campaign: Mapping[str, Any], *, strict: bool = False) -> dict[str, Any]:
    """Copy only missing typed geometry from one exact finalist identity."""
    repaired = copy.deepcopy(campaign)
    rows = repaired.get("weighted_input_rows") if isinstance(repaired, Mapping) else None
    finalists = repaired.get("finalists") if isinstance(repaired, Mapping) else None
    if isinstance(rows, (str, bytes)) or not isinstance(rows, Sequence) or isinstance(finalists, (str, bytes)) or not isinstance(finalists, Sequence):
        return repaired
    finalist_rows = tuple(row for row in finalists if isinstance(row, Mapping))
    geometry_sensitive = any(any(field in row for field in _GEOMETRY_FIELDS) for row in finalist_rows)
    if not strict and not geometry_sensitive:
        return repaired
    by_identity: dict[tuple[str, str, int, int], list[Mapping[str, Any]]] = {}
    for finalist in finalist_rows:
        identity = _geometry_identity(finalist)
        if identity is not None:
            by_identity.setdefault(identity, []).append(finalist)
    for index, row in enumerate(rows):
        if not isinstance(row, Mapping):
            continue
        missing = [field for field in _GEOMETRY_FIELDS if field not in row]
        identity = _geometry_identity(row)
        if not missing:
            continue
        matches = by_identity.get(identity, []) if identity is not None else []
        if len(matches) != 1:
            raise PortfolioPanelError(
                "PORTFOLIO_INPUT_GEOMETRY_INVALID",
                f"weighted_input_rows[{index}] has {len(matches)} matching finalists",
                status=422,
            )
        finalist = matches[0]
        for field in _GEOMETRY_FIELDS:
            if field in row and field in finalist and row[field] != finalist[field]:
                raise PortfolioPanelError(
                    "PORTFOLIO_INPUT_GEOMETRY_INVALID",
                    f"weighted_input_rows[{index}].{field} conflicts with finalist geometry",
                    status=422,
                )
            if field in missing:
                if field not in finalist:
                    raise PortfolioPanelError(
                        "PORTFOLIO_INPUT_GEOMETRY_INVALID",
                        f"finalist geometry is missing {field}",
                        status=422,
                    )
                row[field] = copy.deepcopy(finalist[field])
    return repaired


def _plain(value: Any) -> Any:
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, Mapping):
        return {str(key): _plain(item) for key, item in value.items()}
    if isinstance(value, (tuple, list)):
        return [_plain(item) for item in value]
    return value


def _get(value: Any, key: str, default: Any = None) -> Any:
    if isinstance(value, Mapping):
        return value.get(key, default)
    return getattr(value, key, default)


def _invoke(callback: Callable[..., Any], *values: Any) -> Any:
    try:
        parameters = tuple(inspect.signature(callback).parameters.values())
    except (TypeError, ValueError):
        return callback(values[0])
    if any(parameter.kind == parameter.VAR_POSITIONAL for parameter in parameters):
        return callback(*values)
    count = len(tuple(parameter for parameter in parameters if parameter.kind in (parameter.POSITIONAL_ONLY, parameter.POSITIONAL_OR_KEYWORD)))
    return callback(*values[:count])


def _decimal(value: Any, field: str, *, positive: bool = True) -> str:
    if isinstance(value, bool) or value is None:
        raise PortfolioPanelError("PORTFOLIO_CAMPAIGN_INVALID", "campaign fields are invalid", status=422, field_errors=[{"field": field, "code": "INVALID_NUMBER", "message": "must be a finite decimal"}])
    try:
        result = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        result = Decimal("NaN")
    if not result.is_finite() or (positive and result <= 0):
        raise PortfolioPanelError("PORTFOLIO_CAMPAIGN_INVALID", "campaign fields are invalid", status=422, field_errors=[{"field": field, "code": "INVALID_NUMBER", "message": "must be a finite positive decimal"}])
    digits = list(result.as_tuple().digits)
    exponent = result.as_tuple().exponent
    while digits and digits[-1] == 0 and exponent < 0:
        digits.pop()
        exponent += 1
    if max(-exponent, 0) > 12 or max(len(digits) + exponent, 0) > 26:
        raise PortfolioPanelError("PORTFOLIO_CAMPAIGN_INVALID", "campaign fields are invalid", status=422, field_errors=[{"field": field, "code": "INVALID_NUMBER", "message": "must fit DECIMAL(38,12)"}])
    text = format(result, "f")
    return text.rstrip("0").rstrip(".") if "." in text else text


def _integer(value: Any, field: str, *, nonnegative: bool = False) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or (value < 0 if nonnegative else value <= 0):
        raise PortfolioPanelError("PORTFOLIO_CAMPAIGN_INVALID", "campaign fields are invalid", status=422, field_errors=[{"field": field, "code": "INVALID_INTEGER", "message": "must be a valid integer"}])
    return value


def _redact_text(value: Any) -> str:
    text = str(value)
    text = _SECRET.sub("<redacted-secret>", text)
    text = _PATH.sub("<redacted-path>", text)
    return text[:2048]


def _safe_cell(value: Any, *, key: str = "") -> Any:
    if value is None:
        return None
    if _SECRET.search(key) or key.casefold().endswith(("_path", "_root", "_dir")):
        return "<redacted-secret>" if _SECRET.search(key) else "<redacted-path>"
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    if isinstance(value, Mapping):
        return _redact_text(_json({str(name): _safe_cell(item, key=str(name)) for name, item in value.items()}))
    if isinstance(value, (tuple, list)):
        return _redact_text(_json([_safe_cell(item, key=key) for item in value]))
    if isinstance(value, str):
        text = "<redacted-path>" if _PATH.search(value) else _redact_text(value)
        # Excel treats several leading characters as formula input.  Keep all
        # source text as literal cell content, including untrusted messages.
        if text.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
            return "'" + text
        return text
    return _redact_text(value)


def _weighted_payload_pairs(variant: Any) -> tuple[tuple[Any, Any], ...]:
    members = _get(variant, "members", ())
    payloads = _get(variant, "strategy_payloads", ())
    if isinstance(members, (str, bytes)) or not isinstance(members, Sequence):
        members = ()
    if isinstance(payloads, (str, bytes)) or not isinstance(payloads, Sequence):
        payloads = ()
    payload_by_identity: dict[tuple[str, str], Any] = {}
    payload_by_symbol: dict[str, Any] = {}
    tagged_symbols: set[str] = set()
    ambiguous_symbols: set[str] = set()
    invalid_identities: set[tuple[str, str]] = set()
    for payload in payloads:
        if not isinstance(payload, Mapping):
            continue
        basic = _get(_get(payload, "strategy", {}), "basic", {})
        symbol = _get(basic, "symbol")
        side = _get(_get(payload, "strategy", {}), "side", _get(payload, "side"))
        if not isinstance(symbol, str) or not symbol.strip():
            continue
        symbol = symbol.strip().upper()
        if not isinstance(side, str) or side.strip().upper() not in {"LONG", "SHORT"}:
            if symbol in payload_by_symbol:
                ambiguous_symbols.add(symbol)
                payload_by_symbol.pop(symbol, None)
            elif symbol not in ambiguous_symbols:
                payload_by_symbol[symbol] = payload
            continue
        identity = (symbol, side.strip().upper())
        tagged_symbols.add(symbol)
        if identity in payload_by_identity:
            invalid_identities.add(identity)
            payload_by_identity.pop(identity, None)
            continue
        if identity not in invalid_identities:
            payload_by_identity[identity] = payload
    member_identities: dict[tuple[str, str], int] = {}
    for member in members:
        symbol = _get(member, "symbol", _get(member, "pair"))
        side = _get(member, "side", _get(member, "direction"))
        if isinstance(symbol, str) and isinstance(side, str) and side.strip().upper() in {"LONG", "SHORT"}:
            identity = (symbol.strip().upper(), side.strip().upper())
            member_identities[identity] = member_identities.get(identity, 0) + 1
    pairs = []
    for member in members:
        payload = None
        symbol = _get(member, "symbol", _get(member, "pair"))
        side = _get(member, "side", _get(member, "direction"))
        identity = (
            symbol.strip().upper(),
            side.strip().upper(),
        ) if isinstance(symbol, str) and isinstance(side, str) and side.strip().upper() in {"LONG", "SHORT"} else None
        if identity is not None and member_identities.get(identity) == 1 and identity not in invalid_identities:
            payload = payload_by_identity.get(identity)
        if payload is None and isinstance(symbol, str) and isinstance(side, str):
            symbol_key = symbol.strip().upper()
            if side.strip().upper() == "LONG" and symbol_key not in tagged_symbols and member_identities.get(identity, 0) == 1 and sum(count for (candidate_symbol, _candidate_side), count in member_identities.items() if candidate_symbol == symbol_key) == 1 and symbol_key not in ambiguous_symbols:
                payload = payload_by_symbol.get(symbol_key)
        try:
            raw_x = _get(member, "x_usdt")
            if raw_x is None and isinstance(payload, Mapping):
                raw_x = _get(_get(payload, "facts", {}), "x")
            positive = Decimal(str(raw_x)) > 0
        except (InvalidOperation, TypeError, ValueError):
            positive = False
        if not positive:
            payload = None
        pairs.append((member, payload))
    return tuple(pairs)


def _weighted_value(item: Any, *keys: str, default: Any = "UNKNOWN") -> Any:
    for key in keys:
        value = _get(item, key, None)
        if value is not None:
            return value
    return default


def _weighted_number(item: Any, *keys: str) -> Any:
    value = _weighted_value(item, *keys)
    if (isinstance(value, str) and value == "UNKNOWN") or isinstance(value, bool):
        return "UNKNOWN"
    try:
        number = Decimal(str(value))
    except (ArithmeticError, TypeError, ValueError):
        return "UNKNOWN"
    return value if number.is_finite() and number > 0 else "UNKNOWN"


def _weighted_identity(item: Any) -> tuple[Any, ...] | None:
    if not isinstance(item, Mapping):
        return None
    symbol = _get(item, "symbol", _get(item, "pair"))
    side = _get(item, "side", _get(item, "direction"))
    strategy_id = _get(item, "strategy_id", _get(item, "strategyId"))
    result_id = _get(item, "result_id", _get(item, "resultId"))
    if not isinstance(symbol, str) or not isinstance(side, str) or strategy_id is None or result_id is None:
        return None
    return (symbol.strip().upper(), side.strip().upper(), strategy_id, result_id)


def _weighted_entry_order_percentages(payload: Any, side: Any) -> tuple[Decimal, ...] | str:
    strategy = _get(payload, "strategy", {})
    mrs3 = _get(strategy, "mrs3", {})
    normalized_side = str(side).upper()
    if normalized_side not in {"LONG", "SHORT"}:
        return "UNKNOWN"
    order_key = "ma_long" if normalized_side == "LONG" else "ma_short"
    orders = _get(mrs3, order_key, ())
    if isinstance(orders, (str, bytes)) or not isinstance(orders, Sequence) or not orders:
        return "UNKNOWN"
    lots: list[Decimal] = []
    for order in orders:
        raw = _get(order, "lot_x")
        try:
            lot = Decimal(str(raw))
        except (InvalidOperation, TypeError, ValueError):
            return "UNKNOWN"
        if not lot.is_finite() or lot <= 0:
            return "UNKNOWN"
        lots.append(lot)
    total = sum(lots, Decimal(0))
    return tuple(lot / total * Decimal(100) for lot in lots) if total > 0 else "UNKNOWN"


def _weighted_source_evidence(campaign: Mapping[str, Any]) -> dict[tuple[Any, ...], dict[str, Any]]:
    sources: dict[tuple[Any, ...], dict[str, Any]] = {}
    for key in ("finalists", "weighted_input_rows"):
        rows = _get(campaign, key, ())
        if isinstance(rows, (str, bytes)) or not isinstance(rows, Sequence):
            continue
        for row in rows:
            identity = _weighted_identity(row)
            if identity is None:
                continue
            merged = sources.setdefault(identity, {})
            if isinstance(row, Mapping):
                for name, value in row.items():
                    if name in merged and merged[name] not in (None, value):
                        merged["__invalid_source_evidence"] = True
                    elif name not in merged or merged[name] is None:
                        merged[name] = value
    return sources


def _weighted_source_maxdd_sum(campaign: Mapping[str, Any], members: Sequence[Mapping[str, Any]]) -> Decimal | str:
    sources = _weighted_source_evidence(campaign)
    total = Decimal(0)
    included = False
    for member in members:
        try:
            # position_usdt is copied from the executable payload's facts.x above.
            x = Decimal(str(_get(member, "position_usdt")))
        except (InvalidOperation, TypeError, ValueError):
            return "UNKNOWN"
        if not x.is_finite() or x <= 0:
            return "UNKNOWN"
        included = True
        source = sources.get(_weighted_identity(member))
        if source is None or source.get("__invalid_source_evidence"):
            return "UNKNOWN"
        metrics = _get(source, "metrics", {})
        dd_raw = _weighted_value(source, "max_drawdown", "max_drawdown_usdt", "source_max_drawdown", default=_weighted_value(metrics, "max_drawdown", "max_drawdown_usdt", default=None))
        initial_raw = _weighted_value(source, "source_initial_balance", "initial_balance", "result_initial_balance", default=_weighted_value(metrics, "source_initial_balance", "initial_balance", default=None))
        try:
            dd = Decimal(str(dd_raw))
            initial = Decimal(str(initial_raw))
        except (InvalidOperation, TypeError, ValueError):
            return "UNKNOWN"
        if not dd.is_finite() or dd < 0 or not initial.is_finite() or initial <= 0:
            return "UNKNOWN"
        total += dd * x / initial
    return total if included else "UNKNOWN"


def _weighted_decimal_or_unknown(value: Any, *, nonnegative: bool = False) -> Decimal | str:
    try:
        number = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return "UNKNOWN"
    if not number.is_finite() or (number < 0 if nonnegative else number <= 0):
        return "UNKNOWN"
    return number


def _validate_frozen_campaign(campaign: Any) -> None:
    try:
        validate_campaign_contract(campaign)
    except CampaignContractError as error:
        raise PortfolioPanelError(error.code, error.code, status=422) from error


def _portfolio_search_workers(root: Path) -> int:
    """Read the existing importer width; it is scheduling-only campaign input."""
    return load_duckdb_import_settings(root / "config.local.json").workers


class PortfolioPanelService:
    """Server-owned portfolio service with injected package and tester boundaries."""

    def __init__(
        self,
        root: str | Path,
        config_path: str | Path | None = None,
        *,
        registry: PanelJobRegistry | None = None,
        finalists_reader: Callable[..., Any] | None = None,
        finalists_loader: Callable[..., Any] | None = None,
        cutoff_selector: Callable[..., Any] = apply_finalist_cutoff,
        variant_generator: Callable[..., Any] | None = None,
        variant_validator: Callable[..., Any] | None = None,
        workbook_builder: Callable[..., Any] | None = None,
        lock: threading.RLock | None = None,
        optimizer_input_preparer: Callable[..., Any] | None = None,
        local_testing_service_provider: Callable[[], Any] | None = None,
    ) -> None:
        self.root = Path(root).resolve()
        raw_config = Path(config_path) if config_path is not None else self.root / "portfolio_optimizer.local.json"
        self.config_path = raw_config if raw_config.is_absolute() else self.root / raw_config
        self.config_path = self.config_path.resolve()
        self.registry = registry or PanelJobRegistry(self.root / ".panel-jobs.json")
        self.finalists_reader = finalists_loader or finalists_reader or read_current_finalists
        self._uses_production_finalists_reader = finalists_reader is None and finalists_loader is None
        self.optimizer_input_preparer = optimizer_input_preparer or prepare_current_optimizer_inputs
        self.cutoff_selector = cutoff_selector
        self.variant_generator = variant_generator
        self.variant_validator = variant_validator
        self.workbook_builder = workbook_builder
        self._lock = lock or threading.RLock()
        self._local_testing_service_provider = local_testing_service_provider
        self._cancel_events: dict[str, threading.Event] = {}
        self._threads: dict[str, threading.Thread] = {}
        self._results_root = self.root / ".portfolio-results"
        self._staging_root = self.root / ".portfolio-staging"
        self._progress_reporter = _PortfolioProgressReporter()
        self._progress_state_lock = threading.RLock()
        self._progress_generations: dict[str, int] = {}
        self._active_progress: dict[str, int] = {}

    def _settings_raw(self, *, strict: bool = False) -> tuple[str, dict[str, Any] | None, bytes | None]:
        try:
            raw = self.config_path.read_bytes()
        except FileNotFoundError:
            return "MISSING", None, None
        except OSError as error:
            if strict:
                raise PortfolioPanelError(
                    "PORTFOLIO_SETTINGS_UNAVAILABLE",
                    "portfolio settings are unavailable",
                    status=500,
                ) from error
            return "INVALID", None, None
        try:
            document = json.loads(raw.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError) as error:
            if strict:
                raise PortfolioPanelError(
                    "PORTFOLIO_SETTINGS_UNAVAILABLE",
                    "portfolio settings are unavailable",
                    status=500,
                ) from error
            return "INVALID", None, raw
        if not isinstance(document, dict):
            return "INVALID", None, raw
        try:
            document, _ = migrate_portfolio_config_document(document)
        except PortfolioConfigError:
            return "INVALID", document, raw
        schema = document.get("schema_version")
        policy = document.get("policy_version")
        if schema != SCHEMA_VERSION or policy != POLICY_VERSION:
            return "UNSUPPORTED_SCHEMA", document, raw
        try:
            load_portfolio_config(self.config_path)
        except PortfolioConfigError as error:
            if strict:
                raise PortfolioPanelError(
                    "PORTFOLIO_SETTINGS_UNAVAILABLE",
                    "portfolio settings are unavailable",
                    status=500,
                ) from error
            return "INVALID", document, raw
        return "READY", document, raw

    def settings_get(self) -> dict[str, Any]:
        with self._lock:
            try:
                return self._settings_get()
            except PortfolioPanelError:
                raise
            except (OSError, UnicodeError, json.JSONDecodeError, PortfolioConfigError) as error:
                raise PortfolioPanelError(
                    "PORTFOLIO_SETTINGS_UNAVAILABLE",
                    "portfolio settings are unavailable",
                    status=500,
                ) from error

    def _settings_get(self) -> dict[str, Any]:
        state, document, raw = self._settings_raw(strict=True)
        return {
            "state": state,
            "document": document,
            "digest": _digest(raw) if raw is not None and document is not None else None,
            "schema_version": document.get("schema_version") if document else None,
            "policy_version": document.get("policy_version") if document else None,
        }

    @staticmethod
    def _encode_document(document: Mapping[str, Any]) -> bytes:
        try:
            return (json.dumps(document, ensure_ascii=False, indent=2, allow_nan=False) + "\n").encode("utf-8")
        except (TypeError, ValueError) as error:
            raise PortfolioPanelError("CONFIG_INVALID", "portfolio settings are invalid", status=422) from error

    @staticmethod
    def _replace_bytes(path: Path, payload: bytes) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary: Path | None = None
        try:
            with tempfile.NamedTemporaryFile("wb", dir=path.parent, delete=False) as handle:
                handle.write(payload)
                handle.flush()
                os.fsync(handle.fileno())
                temporary = Path(handle.name)
            os.replace(temporary, path)
            temporary = None
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)

    @staticmethod
    def _snapshot_link(path: Path) -> bool:
        if path.is_symlink() or bool(getattr(path, "is_junction", lambda: False)()):
            return True
        try:
            return bool(os.stat(path, follow_symlinks=False).st_file_attributes & 0x400)
        except (AttributeError, FileNotFoundError, OSError):
            return False

    def _snapshot_path(self, campaign_id: str, *, require_existing: bool = False) -> Path:
        if not isinstance(campaign_id, str) or re.fullmatch(r"campaign-[0-9a-f]{32}", campaign_id) is None:
            raise PortfolioPanelError("PORTFOLIO_SNAPSHOT_UNAVAILABLE", "campaign snapshot is unavailable", status=500)
        root = self._results_root
        directory = root / campaign_id
        path = directory / _SNAPSHOT_NAME
        if root.exists() and self._snapshot_link(root):
            raise PortfolioPanelError("PORTFOLIO_SNAPSHOT_UNAVAILABLE", "campaign snapshot is unavailable", status=500)
        for component in (directory, path):
            if component.exists() and self._snapshot_link(component):
                raise PortfolioPanelError("PORTFOLIO_SNAPSHOT_UNAVAILABLE", "campaign snapshot is unavailable", status=500)
        if require_existing and (not path.is_file() or not path.resolve(strict=True).is_file()):
            raise PortfolioPanelError("PORTFOLIO_SNAPSHOT_UNAVAILABLE", "campaign snapshot is unavailable", status=500)
        try:
            expected_root = root.resolve(strict=False)
            resolved_parent = directory.resolve(strict=False)
            if os.path.commonpath((str(expected_root), str(resolved_parent))) != str(expected_root):
                raise ValueError
        except (OSError, ValueError):
            raise PortfolioPanelError("PORTFOLIO_SNAPSHOT_UNAVAILABLE", "campaign snapshot is unavailable", status=500) from None
        return path

    def _write_campaign_snapshot(self, campaign: Mapping[str, Any]) -> dict[str, Any]:
        campaign_id = campaign.get("campaign_id")
        path = self._snapshot_path(campaign_id)
        try:
            raw, compressed, digest = _snapshot_bytes(campaign)
            path.parent.mkdir(parents=True, exist_ok=True)
            if self._snapshot_link(path.parent):
                raise OSError("snapshot directory is not regular")
            self._replace_bytes(path, compressed)
            if path.read_bytes() != compressed:
                raise OSError("snapshot readback mismatch")
            timestamp = _now()
            return {
                "schema": _SNAPSHOT_SCHEMA,
                "campaign_id": campaign_id,
                "path": path.relative_to(self.root).as_posix(),
                "sha256": digest,
                "compressed_size": len(compressed),
                "uncompressed_size": len(raw),
                "state": "available",
                "created_at_utc": timestamp,
                "updated_at_utc": timestamp,
            }
        except PortfolioPanelError:
            raise
        except (OSError, TypeError, ValueError):
            raise PortfolioPanelError("PORTFOLIO_SNAPSHOT_UNAVAILABLE", "campaign snapshot is unavailable", status=500) from None

    def _hydrate_campaign(self, saved: Mapping[str, Any], runtime: Mapping[str, Any], *, allow_terminal: bool = False, expected_campaign_id: str | None = None) -> dict[str, Any]:
        if saved.get("state") in _TERMINAL and not allow_terminal:
            raise PortfolioPanelError("PORTFOLIO_SNAPSHOT_TERMINAL_STATE", "terminal Campaign cannot be retried", status=409)
        if not isinstance(runtime, Mapping):
            _validate_frozen_campaign(None)
        embedded = runtime.get("campaign") if isinstance(runtime, Mapping) else None
        descriptor = runtime.get("campaign_snapshot") if isinstance(runtime, Mapping) else None
        if isinstance(embedded, Mapping):
            campaign = _repair_legacy_geometry(_plain(dict(embedded)))
            if expected_campaign_id is not None and campaign.get("campaign_id") != expected_campaign_id:
                raise PortfolioPanelError("PORTFOLIO_SNAPSHOT_UNAVAILABLE", "campaign snapshot is unavailable", status=500)
            _validate_frozen_campaign(campaign)
            return campaign
        if isinstance(runtime, Mapping) and "campaign" in runtime:
            _validate_frozen_campaign(embedded)
        if not isinstance(descriptor, Mapping) or descriptor.get("schema") != _SNAPSHOT_SCHEMA:
            raise PortfolioPanelError("PORTFOLIO_SNAPSHOT_UNAVAILABLE", "campaign snapshot is unavailable", status=500)
        campaign_id = descriptor.get("campaign_id")
        if (
            not isinstance(campaign_id, str)
            or (expected_campaign_id is not None and campaign_id != expected_campaign_id)
        ):
            raise PortfolioPanelError("PORTFOLIO_SNAPSHOT_UNAVAILABLE", "campaign snapshot is unavailable", status=500)
        try:
            path = self._snapshot_path(campaign_id, require_existing=True)
            expected_relative = path.relative_to(self.root).as_posix()
            if descriptor.get("path") != expected_relative:
                raise ValueError("snapshot path mismatch")
            compressed_size = descriptor.get("compressed_size")
            uncompressed_size = descriptor.get("uncompressed_size")
            digest = descriptor.get("sha256")
            if (
                type(compressed_size) is not int or compressed_size < 1
                or type(uncompressed_size) is not int or uncompressed_size < 1 or uncompressed_size > _SNAPSHOT_MAX_UNCOMPRESSED
                or not isinstance(digest, str) or re.fullmatch(r"[0-9a-f]{64}", digest) is None
                or path.stat().st_size != compressed_size
            ):
                raise ValueError("snapshot descriptor is invalid")
            with path.open("rb") as handle, gzip.GzipFile(fileobj=handle, mode="rb") as stream:
                raw = stream.read(uncompressed_size + 1)
            if len(raw) != uncompressed_size or hashlib.sha256(raw).hexdigest() != digest:
                raise ValueError("snapshot digest mismatch")
            campaign = json.loads(raw.decode("utf-8"))
            if not isinstance(campaign, dict) or campaign.get("campaign_id") != campaign_id:
                raise ValueError("snapshot campaign mismatch")
            campaign = _repair_legacy_geometry(campaign)
            _validate_frozen_campaign(campaign)
            return campaign
        except PortfolioPanelError:
            raise
        except (OSError, EOFError, gzip.BadGzipFile, UnicodeDecodeError, json.JSONDecodeError, TypeError, ValueError):
            raise PortfolioPanelError("PORTFOLIO_SNAPSHOT_UNAVAILABLE", "campaign snapshot is unavailable", status=500) from None

    @property
    def _snapshot_migration_backup(self) -> Path:
        return self.registry.journal.with_name(_SNAPSHOT_MIGRATION_BACKUP)

    @staticmethod
    def _compact_campaign_binding(campaign: Mapping[str, Any]) -> dict[str, Any]:
        return {
            key: campaign[key]
            for key in ("campaign_id", "input_digest", "config_digest")
            if isinstance(campaign.get(key), str)
        }

    @staticmethod
    def _compact_runtime(runtime: Mapping[str, Any], campaign: Mapping[str, Any]) -> dict[str, Any]:
        compact = copy.deepcopy(dict(runtime))
        compact["campaign"] = PortfolioPanelService._compact_campaign_binding(campaign)
        compact.pop("campaign_snapshot", None)
        for key in ("weighted_input_rows", "finalists", "actions", "equity", "progress"):
            compact.pop(key, None)
        diagnostics = compact.get("diagnostics")
        if isinstance(diagnostics, list):
            compact["diagnostics"] = [
                {
                    "severity": str(item.get("severity", "ERROR")),
                    "code": str(item.get("code", "PORTFOLIO_JOB_FAILED")),
                    "message": _redact_text(item.get("message", ""))[:1024],
                }
                for item in diagnostics if isinstance(item, Mapping)
            ]
            while len(_json(compact["diagnostics"]).encode("utf-8")) > 8192 and compact["diagnostics"]:
                compact["diagnostics"].pop()
        return compact

    def _ensure_snapshot_migration_backup(self, journal_bytes: bytes) -> Path:
        backup = self._snapshot_migration_backup
        if backup.exists():
            if backup.is_file() and not backup.is_symlink():
                return backup
            raise PermissionError("snapshot migration backup is not a regular file")
        backup.parent.mkdir(parents=True, exist_ok=True)
        temporary: Path | None = None
        try:
            with tempfile.NamedTemporaryFile("wb", dir=backup.parent, delete=False) as handle:
                handle.write(journal_bytes)
                handle.flush()
                os.fsync(handle.fileno())
                temporary = Path(handle.name)
            if backup.exists():
                return backup
            os.replace(temporary, backup)
            temporary = None
            return backup
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)

    def _migrate_legacy_campaigns_locked(self) -> bool | None:
        """Move legacy embedded inputs out of the journal before recovery."""
        candidates: list[tuple[str, dict[str, Any], dict[str, Any]]] = []
        journal_bytes = self.registry.journal.read_bytes() if self.registry.journal.exists() else b""
        compressed_total = 0
        has_active_embedded = False
        for job_id, job in self.registry.jobs.items():
            if job.get("kind") != "portfolio.stage1":
                continue
            runtime = job.get("runtime")
            if not isinstance(runtime, Mapping) or not isinstance(runtime.get("campaign"), Mapping):
                continue
            if isinstance(runtime.get("campaign_snapshot"), Mapping):
                continue
            campaign = dict(runtime["campaign"])
            state = job.get("state")
            if state in {"QUEUED", "RUNNING", "CANCELLING"}:
                has_active_embedded = True
            if state in {"FAILED", "CANCELLED"}:
                candidates.append((job_id, job, campaign))
                continue
            if state not in {"QUEUED", "RUNNING", "CANCELLING", "COMMITTED"}:
                continue
            try:
                repaired = _repair_legacy_geometry(campaign, strict=True)
                _raw, compressed, _digest_value = _snapshot_bytes(repaired)
            except (PortfolioPanelError, TypeError, ValueError, OverflowError):
                continue
            compressed_total += len(compressed)
            candidates.append((job_id, job, repaired))
        if not candidates and not has_active_embedded:
            return False
        try:
            free = shutil.disk_usage(self.registry.journal.parent).free
        except OSError:
            return None
        required = 2 * len(journal_bytes) + compressed_total + _SNAPSHOT_MIGRATION_RESERVE
        if free < required:
            return None
        if not candidates:
            return False
        self._ensure_snapshot_migration_backup(journal_bytes)
        migrated = False
        for job_id, original_job, campaign in candidates:
            runtime = original_job.get("runtime")
            if not isinstance(runtime, Mapping):
                continue
            job_copy = copy.deepcopy(original_job)
            try:
                if original_job.get("state") in {"FAILED", "CANCELLED"}:
                    job_copy["runtime"] = self._compact_runtime(runtime, campaign)
                else:
                    descriptor = self._write_campaign_snapshot(campaign)
                    compact_runtime = copy.deepcopy(dict(runtime))
                    compact_runtime.pop("campaign", None)
                    compact_runtime["campaign_snapshot"] = descriptor
                    for key in ("weighted_input_rows", "finalists", "actions", "equity", "progress"):
                        compact_runtime.pop(key, None)
                    job_copy["runtime"] = compact_runtime
                self.registry.jobs[job_id] = job_copy
                migrated = True
            except (PortfolioPanelError, OSError, TypeError, ValueError, OverflowError):
                continue
        return migrated

    def _verify_migrated_registry(self) -> None:
        persisted = self.registry._load()
        if set(persisted) != set(self.registry.jobs):
            raise OSError("snapshot migration changed job IDs")
        for job_id, job in persisted.items():
            if not self.registry._valid_saved_job(job):
                raise OSError("snapshot migration produced an invalid job")
            runtime = job.get("runtime") if isinstance(job.get("runtime"), Mapping) else {}
            descriptor = runtime.get("campaign_snapshot") if isinstance(runtime, Mapping) else None
            if job.get("state") == "COMMITTED" and isinstance(descriptor, Mapping) and descriptor.get("state") == "available":
                expected = job.get("campaign_id") if isinstance(job.get("campaign_id"), str) else descriptor.get("campaign_id")
                self._hydrate_campaign(job, runtime, allow_terminal=True, expected_campaign_id=expected)

    def startup_recover(self) -> None:
        """Migrate legacy Campaigns, then perform one idempotent restart recovery."""
        with self.registry.lock:
            original_jobs = copy.deepcopy(self.registry.jobs)
            backup = self._snapshot_migration_backup
            try:
                migrated = self._migrate_legacy_campaigns_locked()
                if migrated is None:
                    return
                recovered = self.registry.recover_interrupted()
                if recovered:
                    retry = self._migrate_legacy_campaigns_locked()
                    if retry is None:
                        return
                    migrated = migrated or retry
                if migrated:
                    self.registry._save()
                for job_id in tuple(self.registry.jobs):
                    self._cleanup_terminal_snapshot(job_id)
                self._verify_migrated_registry()
                if backup.exists():
                    try:
                        backup.unlink()
                    except PermissionError:
                        return
            except PermissionError:
                self.registry.jobs = original_jobs
                return
            except (OSError, PortfolioPanelError, TypeError, ValueError):
                self.registry.jobs = original_jobs
                try:
                    self.registry._save()
                except OSError:
                    pass

    def settings_put(self, payload: Mapping[str, Any]) -> dict[str, Any]:
        with self._lock:
            return self._settings_put(payload)

    def _settings_put(self, payload: Mapping[str, Any]) -> dict[str, Any]:
        if not isinstance(payload, Mapping) or set(payload) != {"expected_digest", "document"} or not isinstance(payload.get("document"), Mapping):
            raise PortfolioPanelError("CONFIG_INVALID", "expected_digest and full document are required", status=422)
        state, current, raw = self._settings_raw()
        if state != "READY":
            code = "CONFIG_UNSUPPORTED_SCHEMA" if state == "UNSUPPORTED_SCHEMA" else "CONFIG_INVALID"
            raise PortfolioPanelError(code, "portfolio settings are read-only", status=422)
        expected = payload.get("expected_digest")
        current_digest = _digest(raw) if raw is not None and current is not None else None
        if expected != current_digest:
            raise PortfolioPanelError("CONFIG_CHANGED", "portfolio settings changed", status=409)
        try:
            document, _ = migrate_portfolio_config_document(payload["document"])
        except PortfolioConfigError as error:
            raise PortfolioPanelError("CONFIG_INVALID", "portfolio settings are invalid", status=422) from error
        encoded = self._encode_document(document)
        temporary: Path | None = None
        try:
            with tempfile.NamedTemporaryFile("wb", dir=self.config_path.parent, delete=False) as handle:
                handle.write(encoded)
                handle.flush()
                os.fsync(handle.fileno())
                temporary = Path(handle.name)
            try:
                load_portfolio_config(temporary)
            except PortfolioConfigError as error:
                raise PortfolioPanelError("CONFIG_INVALID", "portfolio settings are invalid", status=422) from error
            previous = raw
            os.replace(temporary, self.config_path)
            temporary = None
            try:
                readback = self.config_path.read_bytes()
                if _digest(readback) != _digest(encoded):
                    raise OSError("settings readback digest mismatch")
                load_portfolio_config(self.config_path)
            except Exception as error:
                if previous is not None:
                    self._replace_bytes(self.config_path, previous)
                else:
                    self.config_path.unlink(missing_ok=True)
                raise PortfolioPanelError("CONFIG_WRITE_FAILED", "portfolio settings could not be verified") from error
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)
        return self._settings_get()

    def _config(self) -> tuple[Any, bytes, dict[str, Any]]:
        state, document, raw = self._settings_raw()
        if state != "READY" or document is None or raw is None:
            code = "CONFIG_UNSUPPORTED_SCHEMA" if state == "UNSUPPORTED_SCHEMA" else "CONFIG_INVALID"
            raise PortfolioPanelError(code, "portfolio settings are not ready", status=422)
        try:
            config = load_portfolio_config(self.config_path)
            if self.config_path.read_bytes() != raw:
                raise PortfolioPanelError("CONFIG_CHANGED", "portfolio settings changed", status=409)
            return config, raw, document
        except PortfolioConfigError as error:
            raise PortfolioPanelError("CONFIG_INVALID", "portfolio settings are invalid", status=422) from error

    @staticmethod
    def _path_from_config(root: Path, value: Any) -> Path:
        path = Path(str(value))
        return (root / path).resolve() if not path.is_absolute() else path.resolve()

    def readiness(self) -> dict[str, Any]:
        with self._lock:
            try:
                return self._readiness()
            except PortfolioPanelError:
                raise
            except (OSError, UnicodeError, json.JSONDecodeError, PortfolioConfigError) as error:
                raise PortfolioPanelError(
                    "PORTFOLIO_SETTINGS_UNAVAILABLE",
                    "portfolio settings are unavailable",
                    status=500,
                ) from error

    def _readiness(self) -> dict[str, Any]:
        state, document, raw = self._settings_raw(strict=True)
        config_digest = _digest(raw) if raw is not None and document is not None else None
        stage1: list[str] = []
        pairs: list[str] = []
        finalists: dict[str, int] = {}
        if state != "READY":
            stage1.append(f"SETTINGS_{state}")
        if state == "READY" and document is not None:
            inputs = document.get("inputs") if isinstance(document.get("inputs"), Mapping) else {}
            database = self._path_from_config(self.root, inputs.get("performance_db", ""))
            if not database.is_file():
                stage1.append("PERFORMANCE_DB_UNAVAILABLE")
            else:
                try:
                    with duckdb.connect(str(database), read_only=True) as connection:
                        pair_rows = connection.execute("select distinct symbol, side from selection_runs order by symbol, side").fetchall()
                    pair_keys = tuple((str(symbol).strip().upper(), str(side).strip().upper()) for symbol, side in pair_rows if str(side).strip().upper() in {"LONG", "SHORT"})
                    # Readiness only needs current finalist metadata.  Keep
                    # large action/equity series out of this hot path.
                    rows = _invoke(self.finalists_reader, database, pair_keys, False)
                    for row in rows:
                        symbol = row.get("symbol") if isinstance(row, Mapping) else None
                        side = row.get("side") if isinstance(row, Mapping) else None
                        if not isinstance(symbol, str) or not isinstance(side, str) or side.strip().upper() not in {"LONG", "SHORT"}:
                            continue
                        pair = f"{symbol.strip().upper()}|{side.strip().upper()}"
                        finalists[pair] = finalists.get(pair, 0) + 1
                    pairs = sorted(finalists)
                except Exception:
                    pairs = []
                    finalists = {}
                    stage1.append("FINALISTS_UNAVAILABLE")
        if not pairs and state == "READY" and "PERFORMANCE_DB_UNAVAILABLE" not in stage1:
            stage1.append("NO_FINALISTS")
        active = self.active_job()
        if active is not None:
            stage1.append("PORTFOLIO_JOB_ACTIVE")
        stage1 = list(dict.fromkeys(stage1))
        return {
            "stage1": {"enabled": not stage1, "blockers": stage1},
            "stage2": {"enabled": False, "blockers": ["PORTFOLIO_JOB_STAGE2_NOT_AUTHORIZED"]},
            "settings_state": state,
            "schema_version": document.get("schema_version") if document else None,
            "policy_version": document.get("policy_version") if document else None,
            "config_digest": config_digest,
            "available_pairs": pairs,
            "current_finalists": finalists,
        }

    def _snapshot_finalists(self, document: Mapping[str, Any], launch: Mapping[str, Any]) -> tuple[tuple[dict[str, Any], ...], tuple[dict[str, Any], ...]]:
        inputs = document.get("inputs") if isinstance(document.get("inputs"), Mapping) else {}
        database = self._path_from_config(self.root, inputs.get("performance_db", ""))
        pairs = tuple(
            (pair["pair"], side)
            for pair in launch["pairs"]
            for side, maximum in (("LONG", pair["max_finalist_long"]), ("SHORT", pair["max_finalist_short"]))
            if maximum > 0
        )
        try:
            # Campaign snapshots retain the finalist evidence needed by search.
            if self._uses_production_finalists_reader:
                metadata_rows = _invoke(self.finalists_reader, database, pairs, False)
                metadata_result_ids: list[int] = []
                for index, row in enumerate(metadata_rows or ()):
                    if not isinstance(row, Mapping) or type(row.get("result_id")) is not int:
                        raise PortfolioPanelError(
                            "PORTFOLIO_FINALISTS_UNAVAILABLE",
                            f"portfolio finalist metadata result_id is invalid at row {index}",
                            status=422,
                        )
                    metadata_result_ids.append(row["result_id"])
                result_ids = tuple(dict.fromkeys(metadata_result_ids))
                if result_ids:
                    self.optimizer_input_preparer(
                        database,
                        result_ids,
                        workers=_portfolio_search_workers(self.root),
                    )
            loaded = _invoke(self.finalists_reader, database, pairs, True)
            finalists_list = []
            weighted_input_rows = []
            weighted_fields = (
                "symbol", "side", "strategy_id", "result_id", "imported_at_utc",
                "effective_start_utc", "effective_end_utc", "report_start_utc", "report_end_utc",
                "initial_balance", "source_initial_balance", "total_pnl", "source_pnl", "max_drawdown", "max_drawdown_usdt", "max_drawdown_pct", "optimizer_source_metadata", "source_provenance",
                "timeframe", "close_ma_len", "order_count", "strategy_orders",
            )
            for index, row in enumerate(loaded or ()):
                if not isinstance(row, Mapping):
                    raise PortfolioPanelError("PORTFOLIO_FINALISTS_UNAVAILABLE", "portfolio finalist row is unavailable", status=422)
                prepared_source = getattr(row, "_optimizer_prepared", None)
                item = _plain(dict(row))
                missing_identity = next(
                    (field for field in ("symbol", "side", "strategy_id", "result_id") if field not in item),
                    None,
                )
                if missing_identity is not None:
                    raise PortfolioPanelError(
                        "PORTFOLIO_FINALISTS_UNAVAILABLE",
                        "portfolio finalist identity is unavailable",
                        status=422,
                        field_errors=[{
                            "field": f"finalists[{index}].{missing_identity}",
                            "code": "MISSING_FINALIST_IDENTITY",
                            "message": f"finalist {missing_identity} identity is unavailable",
                        }],
                    )
                if (
                    not isinstance(item.get("symbol"), str) or not item["symbol"].strip()
                    or not isinstance(item.get("side"), str) or item["side"].strip().upper() not in {"LONG", "SHORT"}
                    or type(item.get("strategy_id")) is not int
                    or type(item.get("result_id")) is not int
                ):
                    raise PortfolioPanelError("PORTFOLIO_FINALISTS_UNAVAILABLE", "portfolio finalist identity is unavailable", status=422)
                item["symbol"] = item["symbol"].strip().upper()
                item["side"] = item["side"].strip().upper()

                def series(*aliases: str) -> Any:
                    for alias in aliases:
                        value = item.get(alias)
                        if isinstance(value, Sequence) and not isinstance(value, (str, bytes)):
                            return value
                    field = aliases[0]
                    raise PortfolioPanelError(
                        "PORTFOLIO_FINALISTS_UNAVAILABLE",
                        "portfolio finalist series is unavailable",
                        status=422,
                        field_errors=[{
                            "field": f"finalists[{index}].{field}",
                            "code": "MISSING_FINALIST_SERIES",
                            "message": f"finalist {field} series is unavailable",
                        }],
                    )

                weighted_row = {
                    **{field: item[field] for field in weighted_fields if field in item},
                    "actions": series("actions", "action_series", "minute_actions"),
                    "equity": series("equity", "equity_series"),
                }
                if prepared_source is not None:
                    weighted_row["_prepared_cycles"] = _plain(prepared_source.cycles)
                weighted_input_rows.append(weighted_row)
                for field in ("actions", "action_series", "minute_actions", "equity", "equity_series"):
                    item.pop(field, None)
                finalists_list.append(item)
            finalists = tuple(finalists_list)
            try:
                json.dumps((finalists, tuple(weighted_input_rows)), ensure_ascii=False, sort_keys=True, allow_nan=False)
            except (TypeError, ValueError) as error:
                raise PortfolioPanelError("PORTFOLIO_FINALISTS_UNAVAILABLE", "portfolio finalists are not JSON serializable", status=422) from error
        except PortfolioPanelError:
            raise
        except Exception as error:
            raise PortfolioPanelError("PORTFOLIO_FINALISTS_UNAVAILABLE", "portfolio finalists are unavailable", status=422) from error
        return finalists, tuple(weighted_input_rows)

    def _package_variant_generator(
        self,
        selected: Sequence[Mapping[str, Any]],
        campaign: Mapping[str, Any],
        _profiles: Sequence[Mapping[str, Any]],
        progress_callback: Callable[[Mapping[str, Any]], Any] | None = None,
    ) -> dict[str, Any]:
        """Run the package-owned current-facts adapter for one frozen Campaign."""
        from .portfolio.adapter import run_portfolio_adapter

        try:
            adapter_kwargs = {"workspace_root": self.root, "workers": _portfolio_search_workers(self.root)}
            if progress_callback is not None:
                adapter_kwargs["progress_callback"] = progress_callback
            result = run_portfolio_adapter(selected, campaign, **adapter_kwargs)
        except Exception as error:
            raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "portfolio adapter failed", status=500) from error
        return {
            "variants": result.variants,
            "blockers": list(result.blockers),
            "excluded": result.excluded,
            "warnings": result.warnings,
            "status": getattr(result, "status", "PASS"),
        }

    @staticmethod
    def _normalise_campaign(payload: Mapping[str, Any], config: Any, config_digest: str) -> dict[str, Any]:
        if not isinstance(payload, Mapping):
            raise PortfolioPanelError("PORTFOLIO_CAMPAIGN_INVALID", "campaign fields are invalid", status=422)
        allowed = {"pairs", "profiles", "expected_config_digest"}
        if set(payload) != allowed:
            raise PortfolioPanelError("PORTFOLIO_CAMPAIGN_INVALID", "campaign fields are invalid", status=422)
        if payload.get("expected_config_digest") != config_digest:
            raise PortfolioPanelError("CONFIG_CHANGED", "portfolio settings changed", status=409)
        raw_pairs = payload.get("pairs")
        raw_profiles = payload.get("profiles")
        if not isinstance(raw_pairs, list) or not raw_pairs or not isinstance(raw_profiles, list) or not raw_profiles:
            raise PortfolioPanelError("PORTFOLIO_CAMPAIGN_INVALID", "pairs and profiles are required", status=422)
        pairs: list[dict[str, Any]] = []
        seen_symbols: set[str] = set()
        for index, value in enumerate(raw_pairs):
            if not isinstance(value, Mapping) or set(value) != {"pair", "max_finalist_long", "max_finalist_short"}:
                raise PortfolioPanelError("PORTFOLIO_CAMPAIGN_INVALID", "campaign fields are invalid", status=422, field_errors=[{"field": f"pairs[{index}]", "code": "INVALID_PAIR", "message": "pair fields are invalid"}])
            symbol = value.get("pair")
            if not isinstance(symbol, str) or not symbol.strip():
                raise PortfolioPanelError("PORTFOLIO_CAMPAIGN_INVALID", "campaign fields are invalid", status=422, field_errors=[{"field": f"pairs[{index}].pair", "code": "INVALID_PAIR", "message": "pair must be unique"}])
            symbol = symbol.strip().upper()
            if symbol in seen_symbols:
                raise PortfolioPanelError("PORTFOLIO_CAMPAIGN_INVALID", "campaign fields are invalid", status=422, field_errors=[{"field": f"pairs[{index}].pair", "code": "INVALID_PAIR", "message": "pair must be unique"}])
            seen_symbols.add(symbol)
            max_finalist_long = _integer(value["max_finalist_long"], f"pairs[{index}].max_finalist_long", nonnegative=True)
            max_finalist_short = _integer(value["max_finalist_short"], f"pairs[{index}].max_finalist_short", nonnegative=True)
            if max_finalist_long == 0 and max_finalist_short == 0:
                invalid_field = f"pairs[{index}].max_finalist_long"
                raise PortfolioPanelError(
                    "PORTFOLIO_CAMPAIGN_INVALID",
                    "campaign finalist limits are invalid",
                    status=422,
                    field_errors=[
                        {
                            "field": invalid_field,
                            "code": "FINALIST_LIMIT_RANGE",
                            "message": "at least one direction must be enabled",
                        }
                    ],
                )
            pairs.append({"pair": symbol, "max_finalist_long": max_finalist_long, "max_finalist_short": max_finalist_short})
        profiles: list[dict[str, Any]] = []
        seen_profiles: set[str] = set()
        configured = getattr(config, "profiles", {})
        for index, value in enumerate(raw_profiles):
            if not isinstance(value, Mapping):
                raise PortfolioPanelError("PORTFOLIO_CAMPAIGN_INVALID", "campaign fields are invalid", status=422)
            keys = set(value)
            legacy = keys & {"equity_usdt", "max_balance_usdt"}
            if legacy:
                raise PortfolioPanelError(
                    "PORTFOLIO_CAMPAIGN_INVALID",
                    "legacy profile fields are unsupported",
                    status=422,
                    field_errors=[{
                        "field": f"profiles[{index}].{field}",
                        "code": "LEGACY_PROFILE_FIELD_UNSUPPORTED",
                        "message": "use bank_available_usdt instead",
                    } for field in sorted(legacy)],
                )
            required = {"profile_id", "max_candidates"}
            allowed = required | {"bank_available_usdt"}
            if not required.issubset(keys) or not keys.issubset(allowed):
                raise PortfolioPanelError("PORTFOLIO_CAMPAIGN_INVALID", "campaign fields are invalid", status=422)
            profile_id = value.get("profile_id")
            if not isinstance(profile_id, str) or profile_id not in configured or profile_id in seen_profiles:
                raise PortfolioPanelError("PORTFOLIO_CAMPAIGN_INVALID", "campaign profile is invalid", status=422, field_errors=[{"field": f"profiles[{index}].profile_id", "code": "UNKNOWN_PROFILE", "message": "profile is unknown"}])
            seen_profiles.add(profile_id)
            bank_available = value.get("bank_available_usdt") if "bank_available_usdt" in value else None
            max_candidates = value.get("max_candidates")
            if type(max_candidates) is not int or not 1 <= max_candidates <= 50:
                raise PortfolioPanelError(
                    "PORTFOLIO_CAMPAIGN_INVALID",
                    "campaign fields are invalid",
                    status=422,
                    field_errors=[{
                        "field": f"profiles[{index}].max_candidates",
                        "code": "MAX_CANDIDATES_RANGE",
                        "message": "max_candidates must be between 1 and 50",
                    }],
                )
            profiles.append({
                "profile_id": profile_id,
                "bank_available_usdt": _decimal(bank_available, f"profiles[{index}].bank_available_usdt") if bank_available is not None else None,
                "max_candidates": max_candidates,
            })
        launch = {"pairs": pairs, "profiles": profiles}
        launch["selected_pairs"] = [
            [pair["pair"], side]
            for pair in pairs
            for side, maximum in (("LONG", pair["max_finalist_long"]), ("SHORT", pair["max_finalist_short"]))
            if maximum > 0
        ]
        launch["maximums"] = {
            f"{pair['pair']}|{side}": maximum
            for pair in pairs
            for side, maximum in (("LONG", pair["max_finalist_long"]), ("SHORT", pair["max_finalist_short"]))
            if maximum > 0
        }
        return launch

    def _campaign_by_id(self, campaign_id: str) -> tuple[dict[str, Any], dict[str, Any]]:
        try:
            jobs = self.registry.list()
        except (PanelJobError, KeyError, OSError, TypeError, ValueError) as error:
            raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500) from error
        for saved in jobs:
            if saved.get("kind") != "portfolio.stage1":
                continue
            try:
                runtime = self.registry.runtime(saved["job_id"])
            except (PanelJobError, KeyError, OSError, TypeError, ValueError) as error:
                raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500) from error
            if not isinstance(runtime, Mapping):
                raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500)
            campaign = runtime.get("campaign")
            descriptor = runtime.get("campaign_snapshot")
            binding = campaign if isinstance(campaign, Mapping) else descriptor
            if isinstance(binding, Mapping) and binding.get("campaign_id") == campaign_id:
                return saved, runtime
        raise PortfolioPanelError("PORTFOLIO_CAMPAIGN_NOT_FOUND", "campaign is not available", status=404)

    def _project_orphan(self, saved: Mapping[str, Any]) -> dict[str, Any]:
        """Release a persisted nonterminal job after its worker has disappeared."""
        job_id = saved.get("job_id")
        if saved.get("kind") not in {"portfolio.stage1", "portfolio.stage2"} or not isinstance(job_id, str) or saved.get("state") not in {"QUEUED", "RUNNING", "CANCELLING"} or job_id in self._threads:
            return dict(saved)
        projected = {**saved, "state": "FAILED", "phase": "FAILED", "error": {"code": "INTERRUPTED"}}
        try:
            self.registry.sync(job_id, {"state": "FAILED", "phase": "FAILED", "error": {"code": "INTERRUPTED"}})
        except Exception:
            jobs = getattr(self.registry, "jobs", None)
            lock = getattr(self.registry, "lock", None)
            if isinstance(jobs, dict) and lock is not None:
                with lock:
                    current = jobs.get(job_id)
                    if isinstance(current, dict) and current.get("state") not in _TERMINAL:
                        current.update(state="FAILED", phase="FAILED", error={"code": "INTERRUPTED"})
                        try:
                            self.registry._save()
                        except Exception:
                            pass
        return projected

    @staticmethod
    def _project_state(saved: Mapping[str, Any]) -> str:
        state = saved.get("state")
        if state == "COMMITTED":
            return "SUCCEEDED"
        if state == "CANCELLING":
            return "CANCEL_REQUESTED"
        if state == "FAILED" and isinstance(saved.get("error"), Mapping) and saved["error"].get("code") == "INTERRUPTED":
            return "INTERRUPTED"
        return str(state)

    def _public_job(self, saved: Mapping[str, Any], runtime: Mapping[str, Any] | None = None) -> dict[str, Any]:
        runtime = runtime or {}
        stage2 = saved.get("kind") == "portfolio.stage2"
        campaign = runtime.get("campaign") if isinstance(runtime.get("campaign"), Mapping) else {}
        if not campaign and isinstance(runtime.get("campaign_snapshot"), Mapping):
            campaign = runtime["campaign_snapshot"]
        if stage2 and not campaign:
            binding = runtime.get("stage2") if isinstance(runtime.get("stage2"), Mapping) else {}
            campaign = {"campaign_id": binding.get("campaign_id"), "input_digest": binding.get("input_digest"), "config_digest": binding.get("config_digest")}
        state = self._project_state(saved)
        completed = int(runtime.get("completed_stages", 0) or 0)
        stage_names = STAGE2_STAGES if stage2 else STAGES
        if state == "SUCCEEDED":
            completed = len(stage_names)
        stage: dict[str, Any] = {"index": min(completed, len(stage_names) - 1), "name": stage_names[min(completed, len(stage_names) - 1)], "status": "SUCCEEDED" if state == "SUCCEEDED" else ("RUNNING" if state in {"RUNNING", "CANCEL_REQUESTED"} else state), "completed": int(runtime.get("stage_completed", 0) or 0)}
        if state == "SUCCEEDED":
            stage.update(completed=1, total=1, percent=100)
        elif isinstance(runtime.get("stage_total"), int) and runtime["stage_total"] >= 0:
            stage["total"] = runtime["stage_total"]
        if isinstance(runtime.get("stage_percent"), int):
            stage["percent"] = runtime["stage_percent"]
        overall = 100 if state == "SUCCEEDED" else min(99, (completed * 100) // len(stage_names))
        if state in {"CANCELLED", "FAILED", "INTERRUPTED"}:
            overall = (completed * 100) // len(stage_names)
        diagnostics = saved.get("diagnostics") if isinstance(saved.get("diagnostics"), list) else runtime.get("diagnostics", [])
        diagnostics = [
            {"severity": str(item.get("severity", "ERROR")), "code": str(item.get("code", "PORTFOLIO_JOB_FAILED")), "message": _redact_text(item.get("message", ""))}
            for item in diagnostics if isinstance(item, Mapping)
        ]
        journal = runtime.get("journal", [])
        journal = [
            {
                "timestamp_utc": str(item.get("timestamp_utc", "")),
                "stage": str(item.get("stage", "")),
                "severity": str(item.get("severity", "INFO")),
                "code": str(item.get("code", "PORTFOLIO_JOB_FAILED")),
                "text": _redact_text(item.get("text", "")),
                "counters": _plain(item.get("counters", {})),
            }
            for item in journal if isinstance(item, Mapping)
        ]
        current_digest = None
        try:
            if self.config_path.is_file():
                current_digest = _digest(self.config_path.read_bytes())
        except OSError:
            pass
        frozen_digest = campaign.get("config_digest")
        live_progress = self._progress_reporter.snapshot(str(saved.get("job_id")))
        persisted_progress = runtime.get("optimizer_progress")
        progress = live_progress or (dict(persisted_progress) if isinstance(persisted_progress, Mapping) else None)
        result = {"job_id": saved.get("job_id"), "campaign_id": campaign.get("campaign_id"), "kind": "TESTER_SUBMISSION" if stage2 else "STAGE1_CALCULATION", "status": state, "stage": stage, "overall_percent": overall, "counters": _plain(runtime.get("counters", {})), "progress": _plain(progress) if progress is not None else None, "diagnostics": diagnostics, "journal": journal, "input_digest": campaign.get("input_digest"), "config_digest": frozen_digest, "settings_changed_since_freeze": bool(frozen_digest and current_digest != frozen_digest), "created_at": saved.get("created_at_utc"), "started_at": runtime.get("started_at"), "finished_at": runtime.get("finished_at")}
        if stage2 and state == "SUCCEEDED" and isinstance(runtime.get("stage2_result"), Mapping):
            result["result"] = _plain(runtime["stage2_result"])
        return result

    def job(self, job_id: str) -> dict[str, Any]:
        try:
            saved = self.registry.get(job_id)
        except PanelJobError as error:
            if error.code == "NOT_FOUND":
                raise PortfolioPanelError("PORTFOLIO_JOB_NOT_FOUND", "job is not available", status=404) from error
            raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500) from error
        except KeyError as error:
            raise PortfolioPanelError("PORTFOLIO_JOB_NOT_FOUND", "job is not available", status=404) from error
        except (OSError, TypeError, ValueError) as error:
            raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500) from error
        if not isinstance(saved, Mapping):
            raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500)
        saved = self._project_orphan(saved)
        if saved.get("kind") not in {"portfolio.stage1", "portfolio.stage2"}:
            raise PortfolioPanelError("PORTFOLIO_JOB_NOT_FOUND", "job is not available", status=404)
        try:
            runtime = self.registry.runtime(job_id)
        except (PanelJobError, KeyError, OSError, TypeError, ValueError) as error:
            raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500) from error
        if not isinstance(runtime, Mapping):
            raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500)
        return self._public_job(saved, runtime)

    def active_job(self) -> dict[str, Any] | None:
        try:
            jobs = self.registry.list()
        except (PanelJobError, KeyError, OSError, TypeError, ValueError) as error:
            raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500) from error
        for saved in jobs:
            if saved.get("kind") not in {"portfolio.stage1", "portfolio.stage2"}:
                continue
            saved = self._project_orphan(saved)
            if saved.get("state") in _TERMINAL:
                continue
            try:
                runtime = self.registry.runtime(saved["job_id"])
            except (PanelJobError, KeyError, OSError, TypeError, ValueError) as error:
                raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500) from error
            if not isinstance(runtime, Mapping):
                raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500)
            return self._public_job(saved, runtime)
        return None

    def submit_campaign(self, payload: Mapping[str, Any]) -> dict[str, Any]:
        with self._lock, self.registry.lock:
            config, raw, document = self._config()
            config_digest = _digest(raw)
            launch = self._normalise_campaign(payload, config, config_digest)
            input_digest = _digest(_json(launch))
            active_jobs: list[tuple[dict[str, Any], Mapping[str, Any]]] = []
            for saved in self.registry.list():
                if saved.get("kind") != "portfolio.stage1":
                    continue
                saved = self._project_orphan(saved)
                if saved.get("state") in _TERMINAL:
                    continue
                runtime = self.registry.runtime(saved["job_id"])
                frozen = runtime.get("campaign") if isinstance(runtime.get("campaign"), Mapping) else runtime.get("campaign_snapshot", {})
                active_jobs.append((saved, frozen))
                if frozen.get("input_digest") == input_digest and frozen.get("config_digest") == config_digest:
                    raise PortfolioPanelError("PORTFOLIO_JOB_ACTIVE_DUPLICATE", "an identical campaign is active", status=409)
            if active_jobs:
                raise PortfolioPanelError("PORTFOLIO_JOB_BUSY", "portfolio optimizer is busy", status=409)
            strategy_template, strategy_template_digest = _load_weighted_template()
            finalists, weighted_input_rows = self._snapshot_finalists(document, launch)
            campaign_id = f"campaign-{uuid4().hex}"
            frozen_raw = raw
            try:
                if json.loads(raw.decode("utf-8")) != _plain(document):
                    frozen_raw = self._encode_document(document)
            except (UnicodeDecodeError, json.JSONDecodeError, AttributeError):
                raise PortfolioPanelError("CONFIG_INVALID", "portfolio settings are invalid", status=422) from None
            campaign = {
                "campaign_id": campaign_id,
                "created_at_utc": _now(),
                "input_digest": input_digest,
                "config_digest": config_digest,
                "frozen_config_digest": _digest(frozen_raw),
                "config_bytes": base64.b64encode(frozen_raw).decode("ascii"),
                "config_document": _plain(document),
                "launch": _plain(launch),
                "finalists": finalists,
                "weighted_input_rows": weighted_input_rows,
                "strategy_template": strategy_template,
                "strategy_template_digest": strategy_template_digest,
                "campaign_contract_version": CAMPAIGN_CONTRACT_VERSION,
                "search_mode": CAMPAIGN_SEARCH_MODE,
                "weighted_algo_version": CAMPAIGN_WEIGHTED_ALGO_VERSION,
                "versions": {
                    "campaign_contract_version": CAMPAIGN_CONTRACT_VERSION,
                    "search_mode": CAMPAIGN_SEARCH_MODE,
                    "weighted_algo_version": CAMPAIGN_WEIGHTED_ALGO_VERSION,
                    "schema_version": SCHEMA_VERSION,
                    "policy_version": POLICY_VERSION,
                    "algorithm_versions": _plain(document.get("algorithm_versions", {})),
                },
            }
            _validate_frozen_campaign(campaign)
            snapshot = self._write_campaign_snapshot(campaign)
            saved = None
            submission_key = f"portfolio:{campaign_id}"

            def discard_created_job() -> None:
                candidates: list[str] = []
                if isinstance(saved, Mapping) and isinstance(saved.get("job_id"), str):
                    candidates.append(saved["job_id"])
                else:
                    try:
                        candidates.extend(
                            str(item["job_id"])
                            for item in self.registry.list()
                            if item.get("idempotency_key") == submission_key and item.get("state") == "QUEUED"
                        )
                    except Exception:
                        return
                for candidate in candidates:
                    try:
                        self.registry.discard_queued(candidate)
                    except Exception:
                        pass
                try:
                    self._snapshot_path(campaign_id).unlink(missing_ok=True)
                except (OSError, PortfolioPanelError):
                    pass

            try:
                saved = self.registry.submit("portfolio.stage1", {"campaign_id": campaign_id, "input_digest": input_digest, "config_digest": config_digest}, submission_key, ("portfolio_optimizer",))
                self.registry.reserve_runtime(saved["job_id"], "campaign_snapshot", snapshot)
                self.registry.reserve_runtime(saved["job_id"], "campaign", campaign)
            except PanelJobError as error:
                discard_created_job()
                code = "PORTFOLIO_JOB_BUSY" if error.code in {"RESOURCE_BUSY", "JOB_CAPACITY_EXHAUSTED"} else error.code
                raise PortfolioPanelError(code, "portfolio optimizer is busy" if code == "PORTFOLIO_JOB_BUSY" else code, status=409 if code == "PORTFOLIO_JOB_BUSY" else 400) from error
            except Exception as error:
                discard_created_job()
                raise PortfolioPanelError("PORTFOLIO_JOB_START_FAILED", "portfolio job could not start", status=503) from error
            event = threading.Event()
            self._cancel_events[saved["job_id"]] = event
            try:
                worker = threading.Thread(target=self._run, args=(saved["job_id"],), name="mrs3-portfolio-stage1", daemon=True)
                self._threads[saved["job_id"]] = worker
                worker.start()
            except BaseException as error:
                self._threads.pop(saved["job_id"], None)
                self._cancel_events.pop(saved["job_id"], None)
                self.registry.discard_queued(saved["job_id"])
                try:
                    self._snapshot_path(campaign_id).unlink(missing_ok=True)
                except (OSError, PortfolioPanelError):
                    pass
                raise PortfolioPanelError("PORTFOLIO_JOB_START_FAILED", "portfolio job could not start", status=503) from error
        return {"campaign_id": campaign_id, "job_id": saved["job_id"], "status": "QUEUED", "input_digest": input_digest, "config_digest": config_digest}

    def _sync_runtime(self, job_id: str, **values: Any) -> None:
        runtime = self.registry.runtime(job_id)
        runtime.update(values)
        self.registry.sync(job_id, {"state": self.registry.get(job_id)["state"]}, runtime=runtime)

    @staticmethod
    def _progress_persisted(snapshot: Mapping[str, Any]) -> dict[str, Any]:
        return {
            key: _plain(snapshot.get(key))
            for key in (
                "substage", "unit", "completed", "total", "detail",
                "elapsed_seconds", "eta_seconds", "heartbeat_age_seconds",
                "last_update_age_seconds", "indeterminate",
            )
        }

    def _activate_progress(self, job_id: str) -> int:
        with self._progress_state_lock:
            generation = self._progress_generations.get(job_id, 0) + 1
            self._progress_generations[job_id] = generation
            self._active_progress[job_id] = generation
            return generation

    def _deactivate_progress(self, job_id: str) -> None:
        with self._progress_state_lock:
            self._active_progress.pop(job_id, None)

    def _persist_progress(self, job_id: str, generation: int, snapshot: Mapping[str, Any]) -> None:
        """Persist progress only while this generation is active and nonterminal."""
        # Lock order is progress-state -> registry; registry-locked paths never acquire progress-state.
        with self._progress_state_lock:
            if self._active_progress.get(job_id) != generation:
                return
            with self.registry.lock:
                job = self.registry.jobs.get(job_id)
                if not isinstance(job, Mapping) or job.get("state") in _TERMINAL:
                    return
                runtime = job.get("runtime") if isinstance(job.get("runtime"), Mapping) else {}
                self.registry.sync(
                    job_id,
                    {"state": job["state"]},
                    runtime={**dict(runtime), "optimizer_progress": self._progress_persisted(snapshot)},
                )

    def _progress_callback(self, job_id: str, generation: int | None = None) -> Callable[[Mapping[str, Any]], None]:
        if generation is None:
            with self._progress_state_lock:
                generation = self._active_progress.get(job_id)

        def callback(event: Mapping[str, Any]) -> None:
            total = event.get("total") if isinstance(event, Mapping) else None
            completed = event.get("completed", 0) if isinstance(event, Mapping) else 0
            force = isinstance(total, int) and total > 0 and isinstance(completed, int) and completed >= total
            snapshot, persist = self._progress_reporter.emit(job_id, event, force=force)
            if persist and generation is not None:
                try:
                    self._persist_progress(job_id, generation, snapshot)
                except (PanelJobError, KeyError, OSError, TypeError, ValueError):
                    pass
        return callback

    def _progress_heartbeat(self, job_id: str) -> None:
        with self._progress_state_lock:
            generation = self._active_progress.get(job_id)
        snapshot, persist = self._progress_reporter.heartbeat(job_id)
        if persist and snapshot is not None and generation is not None:
            try:
                self._persist_progress(job_id, generation, snapshot)
            except (PanelJobError, KeyError, OSError, TypeError, ValueError):
                pass

    def _progress_heartbeat_loop(self, job_id: str, stop: threading.Event) -> None:
        while not stop.wait(10.0):
            self._progress_heartbeat(job_id)

    def _append_journal(self, job_id: str, *, stage: str, severity: str, code: str, text: Any) -> None:
        code = code if code.startswith("PORTFOLIO_JOB_") else f"PORTFOLIO_JOB_{code}"
        runtime = self.registry.runtime(job_id)
        entries = runtime.get("journal") if isinstance(runtime.get("journal"), list) else []
        entries.append({"timestamp_utc": _now(), "stage": stage, "severity": severity, "code": code, "text": _redact_text(text), "counters": _plain(runtime.get("counters", {}))})
        runtime["journal"] = entries[-200:]
        self.registry.sync(job_id, {"state": self.registry.get(job_id)["state"]}, runtime=runtime)

    def _set_progress(self, job_id: str, *, stage_index: int, completed: int, stage_total: int | None = None, stage_percent: int | None = None, stage_completed: int = 0) -> None:
        runtime = self.registry.runtime(job_id)
        runtime.update({"stage_index": stage_index, "completed_stages": completed, "stage_completed": stage_completed})
        if stage_total is not None:
            runtime["stage_total"] = stage_total
        else:
            runtime.pop("stage_total", None)
        if stage_percent is not None:
            runtime["stage_percent"] = stage_percent
        else:
            runtime.pop("stage_percent", None)
        self.registry.sync(job_id, {"state": self.registry.get(job_id)["state"], "progress": {"current": completed, "total": len(STAGES), "unit": "stages"}}, runtime=runtime)

    def _cancelled(self, job_id: str) -> bool:
        event = self._cancel_events.get(job_id)
        try:
            state = self.registry.get(job_id)["state"]
        except PanelJobError:
            return True
        return bool(event and event.is_set()) or state == "CANCELLING"

    def _cancel_finish(
        self,
        job_id: str,
        published: Path | None = None,
        previous: bytes | None = None,
        published_executables: Path | None = None,
        previous_executables: bytes | None = None,
        staged_executables: Path | None = None,
    ) -> None:
        if published is not None:
            self._rollback_result(published, previous)
        if published_executables is not None:
            self._rollback_result(published_executables, previous_executables)
        runtime = self.registry.runtime(job_id)
        candidate = runtime.get("staged_workbook_path")
        if isinstance(candidate, str):
            try:
                self._staging_path(Path(candidate), runtime["campaign"]["campaign_id"]).unlink(missing_ok=True)
            except (KeyError, OSError, PortfolioPanelError):
                pass
        if staged_executables is not None:
            try:
                staged_executables.unlink(missing_ok=True)
            except OSError:
                pass
        stage_index = runtime.get("stage_index")
        stage = STAGES[stage_index] if isinstance(stage_index, int) and 0 <= stage_index < len(STAGES) else "CANCELLED"
        self._append_journal(job_id, stage=stage, severity="INFO", code="CANCELLED", text="portfolio calculation cancelled")
        try:
            state = self.registry.get(job_id)["state"]
            if state in {"QUEUED", "RUNNING"}:
                state = self.registry.cancel(job_id)["state"]
            if state == "CANCELLING":
                self.registry.transition(job_id, "CANCELLED", phase="CANCELLED")
        except PanelJobError:
            pass
        self._sync_runtime(job_id, finished_at=_now())

    def _cleanup_terminal_snapshot(self, job_id: str) -> None:
        """Delete only a terminal Campaign's private input file."""
        if job_id in self._threads:
            return
        with self.registry.lock:
            job = self.registry.jobs.get(job_id)
            if not isinstance(job, Mapping) or job.get("state") not in {"CANCELLED", "FAILED"}:
                return
            runtime = job.get("runtime")
            descriptor = runtime.get("campaign_snapshot") if isinstance(runtime, Mapping) else None
            if not isinstance(descriptor, dict):
                return
            if descriptor.get("state") == "deleted":
                return
            path_name = descriptor.get("path")
            for other_id, other in self.registry.jobs.items():
                if other_id == job_id or other.get("state") in _TERMINAL:
                    continue
                other_runtime = other.get("runtime")
                other_descriptor = other_runtime.get("campaign_snapshot") if isinstance(other_runtime, Mapping) else None
                if isinstance(other_descriptor, Mapping) and other_descriptor.get("path") == path_name:
                    return
            descriptor["state"] = "deleting"
            descriptor["updated_at_utc"] = _now()
            try:
                self.registry._save()
            except OSError:
                return
            try:
                campaign_id = descriptor.get("campaign_id")
                path = self._snapshot_path(campaign_id)
                path.unlink(missing_ok=True)
            except PermissionError:
                return
            except (OSError, PortfolioPanelError):
                return
            descriptor["state"] = "deleted"
            descriptor["updated_at_utc"] = _now()
            try:
                path.parent.rmdir()
            except OSError:
                pass
            try:
                self.registry._save()
            except OSError:
                pass

    def _staging_path(self, path: Path, campaign_id: str) -> Path:
        expected_parent = self._staging_root / campaign_id
        if not path.is_absolute() or path.name != "stage1.xlsx" or path.parent != expected_parent:
            raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "staged workbook is outside the server staging root", status=500)
        if path.is_symlink() or expected_parent.is_symlink() or self._staging_root.is_symlink():
            raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "staged workbook is outside the server staging root", status=500)
        try:
            resolved = path.resolve(strict=True)
            if resolved.parent != expected_parent.resolve(strict=True):
                raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "staged workbook is outside the server staging root", status=500)
        except FileNotFoundError as error:
            raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "staged workbook is missing", status=500) from error
        return resolved

    def _staging_target(self, campaign_id: str) -> Path:
        """Create the server-owned staging directory without following links."""
        root = self._staging_root
        if root.exists() and root.is_symlink():
            raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "staging root is unsafe", status=500)
        try:
            root.mkdir(parents=True, exist_ok=True)
            if root.resolve(strict=True) != root:
                raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "staging root is unsafe", status=500)
            directory = root / campaign_id
            if directory.exists() and directory.is_symlink():
                raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "staging root is unsafe", status=500)
            directory.mkdir(exist_ok=True)
            if directory.resolve(strict=True) != directory:
                raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "staging root is unsafe", status=500)
        except OSError as error:
            raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "staging root is unsafe", status=500) from error
        return directory / "stage1.xlsx"

    def _result_path(self, campaign_id: str) -> Path:
        if self._results_root.exists() and self._results_root.is_symlink():
            raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "result root is unsafe", status=500)
        directory = self._results_root / campaign_id
        if directory.exists() and directory.is_symlink():
            raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "result root is unsafe", status=500)
        directory.mkdir(parents=True, exist_ok=True)
        if directory.resolve(strict=True) != directory:
            raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "result root is unsafe", status=500)
        return directory / "stage1.xlsx"

    def _stage1_executables_path(self, campaign_id: str, *, create: bool = False) -> Path:
        if not isinstance(campaign_id, str) or re.fullmatch(r"campaign-[0-9a-f]{32}", campaign_id) is None:
            raise PortfolioPanelError("PORTFOLIO_STAGE1_EXECUTABLES_UNAVAILABLE", "stage 1 executables are unavailable", status=409)
        root = self._results_root
        directory = root / campaign_id
        path = directory / "stage1-executables.json"
        if root.is_symlink() or directory.is_symlink() or path.is_symlink():
            raise PortfolioPanelError("PORTFOLIO_STAGE1_EXECUTABLES_UNAVAILABLE", "stage 1 executables are unavailable", status=409)
        try:
            if create:
                directory.mkdir(parents=True, exist_ok=True)
            if root.resolve(strict=True) != root or directory.resolve(strict=True) != directory:
                raise OSError("artifact path is outside the results root")
            if path.exists() and not path.is_file():
                raise OSError("artifact path is not a file")
        except OSError as error:
            raise PortfolioPanelError("PORTFOLIO_STAGE1_EXECUTABLES_UNAVAILABLE", "stage 1 executables are unavailable", status=409) from error
        return path

    def _staging_executables_path(self, campaign_id: str) -> Path:
        if not isinstance(campaign_id, str) or re.fullmatch(r"campaign-[0-9a-f]{32}", campaign_id) is None:
            raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "staging root is unsafe", status=500)
        root = self._staging_root
        directory = root / campaign_id
        path = directory / "stage1-executables.json"
        if root.is_symlink() or directory.is_symlink() or path.is_symlink():
            raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "staging root is unsafe", status=500)
        try:
            if root.resolve(strict=True) != root or directory.resolve(strict=True) != directory:
                raise OSError("staging path is outside the staging root")
            if path.exists() and not path.is_file():
                raise OSError("staged artifact is not a file")
        except OSError as error:
            raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "staging root is unsafe", status=500) from error
        return path

    @staticmethod
    def _validate_stage1_executable_candidate(candidate: Mapping[str, Any]) -> None:
        if (
            candidate.get("search_mode") != CAMPAIGN_SEARCH_MODE
            or type(candidate.get("limiter_L")) is not int
            or candidate["limiter_L"] != 0
            or not isinstance(candidate.get("candidate_id"), str)
            or not candidate["candidate_id"]
            or not isinstance(candidate.get("identity"), str)
            or not candidate["identity"]
            or not isinstance(candidate.get("profile"), str)
            or not candidate["profile"]
        ):
            raise ValueError("candidate is not executable under the off-only policy")
        _validate_pretest_period(candidate.get("pretest_period"))
        required_bank = _candidate_required_bank(candidate)
        metrics = candidate.get("metrics", {})
        if isinstance(metrics, Mapping) and "limiter_L" in metrics and (type(metrics["limiter_L"]) is not int or metrics["limiter_L"] != 0):
            raise ValueError("candidate metrics use a nonzero limiter")
        payloads = candidate.get("strategy_payloads")
        if isinstance(payloads, (str, bytes)) or not isinstance(payloads, Sequence) or len(payloads) < 2:
            raise ValueError("candidate does not contain two executable strategies")
        members = candidate.get("members")
        if isinstance(members, (str, bytes)) or not isinstance(members, Sequence) or len(members) != len(payloads):
            raise ValueError("candidate members do not match its executable strategies")
        names: set[str] = set()
        identities: set[tuple[str, str]] = set()
        payload_identities = []
        for payload in payloads:
            if not isinstance(payload, Mapping):
                raise ValueError("strategy payload is invalid")
            strategy = payload.get("strategy")
            basic = strategy.get("basic") if isinstance(strategy, Mapping) else None
            mrs = strategy.get("mrs") if isinstance(strategy, Mapping) else None
            account = payload.get("account")
            facts = payload.get("facts")
            symbol = basic.get("symbol") if isinstance(basic, Mapping) else None
            side = payload.get("side")
            name = strategy.get("name") if isinstance(strategy, Mapping) else None
            if (
                not isinstance(symbol, str) or not symbol or symbol != symbol.strip().upper()
                or not isinstance(side, str) or side not in {"LONG", "SHORT"}
                or not isinstance(name, str) or not name
                or not isinstance(mrs, Mapping) or type(mrs.get("position_priority")) is not int or mrs["position_priority"] != 1
                or not isinstance(account, Mapping) or type(account.get("open_positions_limiter")) is not int or account["open_positions_limiter"] != 0
                or not isinstance(facts, Mapping)
            ):
                raise ValueError("strategy payload violates the off-only executable contract")
            try:
                bank = Decimal(str(facts.get("B")))
                x = Decimal(str(facts.get("x")))
            except (InvalidOperation, TypeError, ValueError) as error:
                raise ValueError("strategy sizing evidence is invalid") from error
            if (
                not bank.is_finite() or bank <= 0 or bank != required_bank
                or not x.is_finite() or x <= 0 or name in names or (symbol, side) in identities
            ):
                raise ValueError("strategy payload identity or positive size is invalid")
            names.add(name)
            identities.add((symbol, side))
            payload_identities.append((symbol, side))
        member_identities = []
        strategy_ids: set[int] = set()
        result_ids: set[int] = set()
        for member in members:
            if not isinstance(member, Mapping) or set(member) != {"symbol", "side", "strategy_id", "result_id"}:
                raise ValueError("candidate member identity is invalid")
            symbol, side = member["symbol"], member["side"]
            strategy_id, result_id = member["strategy_id"], member["result_id"]
            if (
                not isinstance(symbol, str) or not symbol or symbol != symbol.strip().upper()
                or not isinstance(side, str) or side not in {"LONG", "SHORT"}
                or type(strategy_id) is not int or strategy_id <= 0 or strategy_id in strategy_ids
                or type(result_id) is not int or result_id <= 0 or result_id in result_ids
                or (symbol, side) in member_identities
            ):
                raise ValueError("candidate member identity is ambiguous")
            strategy_ids.add(strategy_id)
            result_ids.add(result_id)
            member_identities.append((symbol, side))
        if member_identities != payload_identities:
            raise ValueError("candidate members do not match executable strategy order")

    @classmethod
    def _build_stage1_executables(cls, campaign: Mapping[str, Any], variants: Sequence[Any]) -> dict[str, Any]:
        candidates = []
        for order, variant in enumerate(variants):
            if not isinstance(variant, Mapping):
                continue
            metrics = variant.get("metrics")
            if isinstance(metrics, Mapping) and "limiter_L" in metrics and (type(metrics["limiter_L"]) is not int or metrics["limiter_L"] != 0):
                continue
            candidate_id = variant.get("candidate_id", variant.get("identity"))
            identity = variant.get("identity", candidate_id)
            candidate = {
                "candidate_id": candidate_id,
                "identity": identity,
                "profile": variant.get("profile", variant.get("profile_id")),
                "order": order,
                "search_mode": variant.get("search_mode"),
                "limiter_L": variant.get("limiter_L"),
                "pretest_period": _plain(variant.get("pretest_period")),
                "metrics": _plain(variant.get("metrics", {})),
                "strategy_payloads": _plain(variant.get("strategy_payloads", ())),
            }
            try:
                source_members = variant.get("members")
                if isinstance(source_members, (str, bytes)) or not isinstance(source_members, Sequence):
                    continue
                members = []
                for member in source_members:
                    if not isinstance(member, Mapping):
                        raise ValueError("candidate member is invalid")
                    x = Decimal(str(member.get("x_usdt")))
                    if not x.is_finite():
                        raise ValueError("candidate member size is invalid")
                    if x > 0:
                        members.append({key: member.get(key) for key in ("symbol", "side", "strategy_id", "result_id")})
                candidate["members"] = members
                cls._validate_stage1_executable_candidate(candidate)
                candidate["candidate_digest"] = _digest(_json(candidate))
            except (ArithmeticError, TypeError, ValueError):
                continue
            candidates.append(candidate)
        if not candidates:
            raise PortfolioPanelError("PORTFOLIO_STAGE1_EXECUTABLES_UNAVAILABLE", "no off-only executable candidate is available", status=422)
        if (
            not isinstance(campaign.get("campaign_id"), str)
            or re.fullmatch(r"campaign-[0-9a-f]{32}", campaign["campaign_id"]) is None
            or any(not isinstance(campaign.get(key), str) or re.fullmatch(r"[0-9a-f]{64}", campaign[key]) is None for key in ("input_digest", "config_digest"))
        ):
            raise PortfolioPanelError("PORTFOLIO_STAGE1_EXECUTABLES_UNAVAILABLE", "stage 1 campaign bindings are invalid", status=422)
        body = {
            "schema_version": 1,
            "campaign_id": campaign.get("campaign_id"),
                "input_digest": campaign.get("input_digest"),
                "config_digest": campaign.get("config_digest"),
                "candidates": candidates,
        }
        body["payload_digest"] = _digest(_json(body))
        return body

    def _load_stage1_executables(
        self,
        campaign_id: str,
        input_digest: str,
        config_digest: str,
        expected_digest: str,
        *,
        require_committed: bool = True,
    ) -> dict[str, Any]:
        try:
            path = self._stage1_executables_path(campaign_id)
            if require_committed:
                saved, runtime = self._campaign_by_id(campaign_id)
                campaign = self._hydrate_campaign(saved, runtime, allow_terminal=True, expected_campaign_id=campaign_id)
                if (
                    self._project_state(saved) != "SUCCEEDED"
                    or not isinstance(campaign, Mapping)
                    or campaign.get("input_digest") != input_digest
                    or campaign.get("config_digest") != config_digest
                    or runtime.get("executables_path") != str(path)
                    or runtime.get("executables_digest") != expected_digest
                ):
                    raise ValueError("artifact is not tied to a committed Stage 1 job")
            if not path.is_absolute() or path.resolve(strict=True) != path or not path.is_file():
                raise ValueError("artifact path is invalid")
            raw = path.read_bytes()
            document = json.loads(raw.decode("utf-8"))
            if (
                not isinstance(document, dict)
                or set(document) != {"schema_version", "campaign_id", "input_digest", "config_digest", "candidates", "payload_digest"}
                or type(document["schema_version"]) is not int
                or document["schema_version"] != 1
                or document["campaign_id"] != campaign_id
                or document["input_digest"] != input_digest
                or document["config_digest"] != config_digest
                or document["payload_digest"] != expected_digest
                or raw != (_json(document) + "\n").encode("utf-8")
            ):
                raise ValueError("artifact binding is invalid")
            candidates = document["candidates"]
            if not isinstance(candidates, list) or not candidates:
                raise ValueError("artifact candidates are invalid")
            seen_ids: set[str] = set()
            seen_identities: set[str] = set()
            previous_order = -1
            for candidate in candidates:
                fields = {"candidate_id", "identity", "profile", "order", "search_mode", "limiter_L", "pretest_period", "metrics", "strategy_payloads", "members", "candidate_digest"}
                if not isinstance(candidate, dict) or set(candidate) != fields:
                    raise ValueError("artifact candidate shape is invalid")
                candidate_body = {key: value for key, value in candidate.items() if key != "candidate_digest"}
                if (
                    type(candidate.get("order")) is not int
                    or candidate["order"] <= previous_order
                    or candidate["candidate_id"] in seen_ids
                    or candidate["identity"] in seen_identities
                    or candidate["candidate_digest"] != _digest(_json(candidate_body))
                ):
                    raise ValueError("artifact candidate binding is invalid")
                self._validate_stage1_executable_candidate(candidate)
                seen_ids.add(candidate["candidate_id"])
                seen_identities.add(candidate["identity"])
                previous_order = candidate["order"]
            body = {key: value for key, value in document.items() if key != "payload_digest"}
            if document["payload_digest"] != _digest(_json(body)):
                raise ValueError("artifact digest is invalid")
            return document
        except (ArithmeticError, OSError, UnicodeDecodeError, json.JSONDecodeError, KeyError, TypeError, ValueError) as error:
            raise PortfolioPanelError("PORTFOLIO_STAGE1_EXECUTABLES_UNAVAILABLE", "stage 1 executables are unavailable", status=409) from error

    def _prepare_stage2_baseline(self, campaign_id: str) -> dict[str, Any]:
        """Build one exact Stage 2 input package without running or mutating anything."""
        try:
            saved, runtime = self._campaign_by_id(campaign_id)
            campaign = self._hydrate_campaign(saved, runtime, allow_terminal=True, expected_campaign_id=campaign_id)
            input_digest = campaign.get("input_digest")
            config_digest = campaign.get("config_digest")
            artifact_digest = runtime.get("executables_digest")
            if (
                not isinstance(input_digest, str)
                or re.fullmatch(r"[0-9a-f]{64}", input_digest) is None
                or not isinstance(config_digest, str)
                or re.fullmatch(r"[0-9a-f]{64}", config_digest) is None
                or not isinstance(artifact_digest, str)
                or re.fullmatch(r"[0-9a-f]{64}", artifact_digest) is None
            ):
                raise ValueError("stage 1 bindings are invalid")
            artifact = self._load_stage1_executables(campaign_id, input_digest, config_digest, artifact_digest)
            candidate = artifact["candidates"][0]

            material = _stage2_material(candidate)
            return {
                "campaign_id": campaign_id,
                "input_digest": input_digest,
                "config_digest": config_digest,
                "artifact_digest": artifact_digest,
                "candidate_id": material["candidate_id"],
                "candidate_digest": candidate["candidate_digest"],
                "portfolio_name": material["candidate_id"],
                "expected_names": material["expected_names"],
                "pretest_period": material["pretest_period"],
                "tester_config_json": material["tester_config_json"],
                "strategy_jsons": material["strategy_jsons"],
                "receipt": material["receipt"],
            }
        except PortfolioPanelError as error:
            raise PortfolioPanelError("PORTFOLIO_STAGE2_INPUT_INVALID", "stage 2 baseline input is invalid", status=409) from error
        except (ArithmeticError, OSError, UnicodeDecodeError, json.JSONDecodeError, KeyError, TypeError, ValueError) as error:
            raise PortfolioPanelError("PORTFOLIO_STAGE2_INPUT_INVALID", "stage 2 baseline input is invalid", status=409) from error

    def _verify_stage2_prepared(self, prepared: Mapping[str, Any]) -> None:
        """Re-read the committed artifact and prove the Stage 2 receipt still matches."""
        try:
            artifact = self._load_stage1_executables(
                prepared["campaign_id"],
                prepared["input_digest"],
                prepared["config_digest"],
                prepared["artifact_digest"],
            )
            candidates = artifact.get("candidates")
            receipt = prepared.get("receipt")
            if not isinstance(candidates, list) or not candidates or not isinstance(receipt, Mapping):
                raise ValueError("stage 2 receipt is invalid")
            candidate = candidates[0]
            if not isinstance(candidate, Mapping):
                raise ValueError("stage 2 candidate is invalid")
            payloads = candidate.get("strategy_payloads")
            first = payloads[0] if isinstance(payloads, Sequence) and payloads else None
            facts = first.get("facts") if isinstance(first, Mapping) else None
            material = _stage2_material(candidate)
            if (
                candidate.get("candidate_id") != prepared.get("candidate_id")
                or candidate.get("candidate_digest") != prepared.get("candidate_digest")
                or material["expected_names"] != prepared.get("expected_names")
                or material["pretest_period"] != prepared.get("pretest_period")
                or material["tester_config_json"] != prepared.get("tester_config_json")
                or material["strategy_jsons"] != prepared.get("strategy_jsons")
                or material["receipt"] != receipt
            ):
                raise ValueError("stage 2 receipt no longer matches")
        except (
            ArithmeticError, InvalidOperation, KeyError, OSError, TypeError,
            UnicodeDecodeError, ValueError, PortfolioPanelError,
        ) as error:
            raise PortfolioPanelError(
                "PORTFOLIO_JOB_STAGE2_INPUT_CHANGED", "stage 2 receipt no longer matches", status=500
            ) from error

    @staticmethod
    def _stage2_result(
        prepared: Mapping[str, Any],
        result: WizardResult,
        *,
        report_evidence: Mapping[str, Any] | None = None,
    ) -> dict[str, Any]:
        names = tuple(result.strategy_names)
        expected = tuple(prepared["expected_names"])
        if len(names) != len(set(names)) or len(names) < 2 or set(names) != set(expected):
            raise PortfolioPanelError("PORTFOLIO_JOB_STAGE2_RESULT_INVALID", "tester result strategy names are invalid", status=500)
        metrics: dict[str, float] = {}
        for key in CORE_METRICS:
            value = result.stats.get(key)
            if isinstance(value, bool) or value is None:
                raise PortfolioPanelError("PORTFOLIO_JOB_STAGE2_RESULT_INVALID", "tester result metrics are invalid", status=500)
            try:
                number = float(Decimal(str(value)))
            except (ArithmeticError, TypeError, ValueError):
                raise PortfolioPanelError("PORTFOLIO_JOB_STAGE2_RESULT_INVALID", "tester result metrics are invalid", status=500) from None
            if not math.isfinite(number):
                raise PortfolioPanelError("PORTFOLIO_JOB_STAGE2_RESULT_INVALID", "tester result metrics are invalid", status=500)
            if key in {"MaxDrawdown", "MaxDrawdownPercent"} and number < 0:
                raise PortfolioPanelError("PORTFOLIO_JOB_STAGE2_RESULT_INVALID", "tester result drawdown is invalid", status=500)
            metrics[key] = number
        output = {
            "campaign_id": prepared["campaign_id"],
            "input_digest": prepared["input_digest"],
            "config_digest": prepared["config_digest"],
            "artifact_digest": prepared["artifact_digest"],
            "candidate_id": prepared["candidate_id"],
            "candidate_digest": prepared["candidate_digest"],
            "report_folder": prepared["candidate_id"],
            "name_comment": prepared["candidate_id"],
            "strategy_names": list(names),
            "period": result.period,
            "pretest_period": _plain(prepared["pretest_period"]),
            "run_id": result.run_id,
            "report_name": result.report_name,
            "report_link": result.chart_url or None,
            "metrics": metrics,
            "expected": {
                "sizing": _plain(prepared["receipt"].get("expected_sizing", [])),
                "max_balance": _plain(prepared["receipt"].get("expected_max_balance", [])),
            },
        }
        if report_evidence is not None:
            output["report_evidence"] = dict(report_evidence)
        return output

    def _stage2_cancel_finish(self, job_id: str) -> None:
        try:
            saved = self.registry.get(job_id)
            if saved["state"] in {"QUEUED", "RUNNING"}:
                saved = self.registry.cancel(job_id)
            runtime = self.registry.runtime(job_id)
            runtime.pop("stage2_result", None)
            runtime["finished_at"] = _now()
            if saved["state"] in {"QUEUED", "RUNNING", "CANCELLING"}:
                self.registry.sync(job_id, {"state": "CANCELLED", "phase": "CANCELLED"}, runtime=runtime)
            else:
                self.registry.sync(job_id, {"state": saved["state"], "phase": saved.get("phase")}, runtime=runtime)
        except PanelJobError:
            pass

    def _stage2_fail(self, job_id: str, error: BaseException) -> None:
        try:
            saved = self.registry.get(job_id)
            code = error.code if isinstance(error, PortfolioPanelError) else "PORTFOLIO_JOB_STAGE2_FAILED"
            message = str(error) if isinstance(error, PortfolioPanelError) else "portfolio tester submission failed"
            diagnostics = [{"severity": "ERROR", "code": code, "message": _redact_text(message)}]
            runtime = self.registry.runtime(job_id)
            runtime.pop("stage2_result", None)
            runtime["diagnostics"] = diagnostics
            runtime["finished_at"] = _now()
            if saved["state"] in _TERMINAL:
                state, phase = saved["state"], saved.get("phase")
            else:
                state, phase = "FAILED", "FAILED"
            self.registry.sync(job_id, {"state": state, "phase": phase, "error": {"code": code, "message": _redact_text(message)}, "result": {}}, runtime=runtime)
        except PanelJobError:
            pass

    def _run_stage2(self, job_id: str, prepared: Mapping[str, Any]) -> None:
        tester: Any = None
        filled = False
        failure: BaseException | None = None
        completed_result: dict[str, Any] | None = None
        try:
            if self._cancelled(job_id):
                raise asyncio.CancelledError
            self.registry.transition(job_id, "RUNNING", phase=STAGE2_STAGES[0])
            self._sync_runtime(job_id, started_at=_now(), stage_index=0, completed_stages=0)
            tester = self._local_testing_service_provider()
            if self._cancelled(job_id):
                raise asyncio.CancelledError
            self._verify_stage2_prepared(prepared)
            fill_readback = tester.fill_prebuilt(
                tester_config_json=prepared["tester_config_json"],
                strategy_jsons=prepared["strategy_jsons"],
                delete_old_reports=False,
            )
            filled = True
            receipt = prepared["receipt"]
            actual_names = fill_readback.get("strategy_names") if isinstance(fill_readback, Mapping) else None
            actual_manifest = fill_readback.get("strategy_file_manifest") if isinstance(fill_readback, Mapping) else None
            if isinstance(actual_manifest, list):
                actual_manifest = sorted(
                    actual_manifest,
                    key=lambda item: str(item.get("filename", "")) if isinstance(item, Mapping) else "",
                )
            if (
                not isinstance(fill_readback, Mapping)
                or fill_readback.get("tester_config_hash") != receipt["tester_config_sha256"]
                or not isinstance(actual_names, Sequence)
                or not all(isinstance(name, str) for name in actual_names)
                or len(actual_names) != len(set(actual_names))
                or sorted(actual_names) != sorted(prepared["expected_names"])
                or actual_manifest != receipt["strategy_manifest"]
            ):
                raise PortfolioPanelError("PORTFOLIO_JOB_STAGE2_FILL_READBACK_INVALID", "tester staging readback is invalid", status=500)
            self._sync_runtime(job_id, stage_index=1, completed_stages=1)
            if self._cancelled(job_id):
                raise asyncio.CancelledError
            baseline = file_fingerprint(tester.config.wizard_result)
            candidate_id = prepared.get("candidate_id")
            if not isinstance(candidate_id, str) or re.fullmatch(r"[0-9a-f]{64}", candidate_id) is None:
                raise PortfolioPanelError("PORTFOLIO_JOB_STAGE2_RESULT_INVALID", "tester report folder is invalid", status=500)
            report_root = Path(tester.config.report_dir).parent.resolve()
            report_folder = report_root / candidate_id
            if report_folder.is_symlink():
                raise PortfolioPanelError("PORTFOLIO_JOB_STAGE2_RESULT_INVALID", "tester report folder is invalid", status=500)
            report_folder = report_folder.resolve()
            try:
                report_folder.relative_to(report_root)
            except ValueError as error:
                raise PortfolioPanelError("PORTFOLIO_JOB_STAGE2_RESULT_INVALID", "tester report folder is invalid", status=500) from error
            folder_snapshot = _report_folder_snapshot(report_folder)
            if self._cancelled(job_id):
                raise asyncio.CancelledError
            self._verify_stage2_prepared(prepared)
            start_wall_ns = time.time_ns()
            tester.start()
            if self._cancelled(job_id):
                raise asyncio.CancelledError
            self._sync_runtime(job_id, stage_index=2, completed_stages=2)
            deadline = time.monotonic() + float(tester.config.stall_timeout_seconds)
            previous_report_fingerprint: tuple[int, int, str] | None = None
            stable_report_polls = 0
            while time.monotonic() < deadline:
                if self._cancelled(job_id):
                    raise asyncio.CancelledError
                try:
                    results = read_stable_file(
                        tester.config.wizard_result,
                        lambda path: load_wizard_results(path, expected_report_folder=prepared["candidate_id"]),
                        baseline=baseline,
                    )
                except ResultParseError as error:
                    raise PortfolioPanelError("PORTFOLIO_JOB_STAGE2_RESULT_INVALID", "tester result is invalid", status=500) from error
                if results is not None:
                    if len(results) != 1:
                        raise PortfolioPanelError("PORTFOLIO_JOB_STAGE2_RESULT_INVALID", "tester result count is invalid", status=500)
                    report_fingerprint = _fresh_report_fingerprint(
                        report_folder,
                        results[0].report_name,
                        folder_snapshot,
                        start_wall_ns,
                    )
                    if report_fingerprint is None:
                        interval = float(getattr(tester.config, "poll_interval_seconds", 1.0))
                        event = self._cancel_events.get(job_id)
                        if event is not None:
                            event.wait(interval)
                        continue
                    if report_fingerprint == previous_report_fingerprint:
                        stable_report_polls += 1
                    else:
                        previous_report_fingerprint = report_fingerprint
                        stable_report_polls = 1
                    if stable_report_polls < 2:
                        interval = float(getattr(tester.config, "poll_interval_seconds", 1.0))
                        event = self._cancel_events.get(job_id)
                        if event is not None:
                            event.wait(interval)
                        continue
                    self._verify_stage2_prepared(prepared)
                    completed_result = self._stage2_result(
                        prepared,
                        results[0],
                        report_evidence={
                            "fingerprint": {
                                "mtime_ns": report_fingerprint[0],
                                "size": report_fingerprint[1],
                                "sha256": report_fingerprint[2],
                            },
                            "start_wall_ns": start_wall_ns,
                            "accepted_wall_ns": time.time_ns(),
                            "stable_polls": stable_report_polls,
                        },
                    )
                    self._sync_runtime(job_id, stage_index=3, completed_stages=3, stage2_result=completed_result)
                    break
                interval = float(getattr(tester.config, "poll_interval_seconds", 1.0))
                event = self._cancel_events.get(job_id)
                if event is None or not event.wait(interval):
                    continue
            else:
                raise PortfolioPanelError("PORTFOLIO_JOB_STAGE2_TIMEOUT", "tester result did not become fresh", status=500)
            if self._cancelled(job_id):
                raise asyncio.CancelledError
        except asyncio.CancelledError as error:
            failure = error
        except BaseException as error:
            failure = error
        finally:
            if filled and tester is not None:
                try:
                    tester.stop()
                except BaseException as error:
                    if failure is None or isinstance(failure, asyncio.CancelledError):
                        failure = error
            if failure is None and completed_result is not None:
                if self._cancelled(job_id):
                    failure = asyncio.CancelledError()
                else:
                    try:
                        runtime = self.registry.runtime(job_id)
                        self.registry.sync(job_id, {"state": "COMMITTED", "phase": "COMMITTED", "result": completed_result}, runtime={**runtime, "finished_at": _now()})
                    except BaseException as error:
                        failure = asyncio.CancelledError() if self._cancelled(job_id) else error
            if isinstance(failure, asyncio.CancelledError):
                self._stage2_cancel_finish(job_id)
            elif failure is not None:
                self._stage2_fail(job_id, failure)
            self._cancel_events.pop(job_id, None)
            self._threads.pop(job_id, None)

    @staticmethod
    def _verify_workbook(path: Path) -> None:
        from openpyxl import load_workbook

        workbook = load_workbook(path, data_only=False)
        try:
            expected = ["Summary", "Finalists", "Portfolios", "Members", "Excluded", "Metadata"]
            weighted = ["Итог", "Варианты", "Состав", "Финалисты", "Исключено", "Metadata"]
            if workbook.sheetnames not in (expected, [*expected[:-1], "Metadata", "Profile Status"], weighted):
                raise ValueError("workbook sheets are invalid")
            if getattr(workbook, "_external_links", ()):
                raise ValueError("workbook external links are not allowed")
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    for cell in row:
                        if cell.data_type == "f":
                            raise ValueError("workbook formulas are not allowed")
                        if getattr(cell, "hyperlink", None) is not None:
                            raise ValueError("workbook links are not allowed")
        finally:
            workbook.close()

    def _rollback_result(self, path: Path, previous: bytes | None) -> None:
        try:
            if previous is None:
                path.unlink(missing_ok=True)
            else:
                self._replace_bytes(path, previous)
        except OSError:
            pass

    def _run(self, job_id: str) -> None:
        published: Path | None = None
        previous: bytes | None = None
        published_executables: Path | None = None
        previous_executables: bytes | None = None
        staged_executables: Path | None = None
        executable_document: dict[str, Any] | None = None
        heartbeat_stop: threading.Event | None = None
        heartbeat_thread: threading.Thread | None = None
        progress_generation: int | None = None
        saved = self.registry.get(job_id)
        if saved.get("state") in _TERMINAL:
            return
        try:
            runtime = self.registry.runtime(job_id)
            campaign = self._hydrate_campaign(
                saved,
                runtime,
                expected_campaign_id=(saved.get("campaign_id") if isinstance(saved, Mapping) else None),
            )
            self.registry.transition(job_id, "RUNNING", phase=STAGES[0])
            saved = self.registry.get(job_id)
            launch = campaign["launch"]
            try:
                frozen_raw = base64.b64decode(campaign["config_bytes"], validate=True)
                frozen_digest = campaign.get("frozen_config_digest", campaign["config_digest"])
                if _digest(frozen_raw) != frozen_digest or json.loads(frozen_raw.decode("utf-8")) != campaign["config_document"]:
                    raise ValueError("frozen config digest mismatch")
            except (ValueError, TypeError, UnicodeDecodeError):
                raise PortfolioPanelError("CONFIG_INVALID", "frozen portfolio settings are invalid", status=422) from None
            self._sync_runtime(job_id, started_at=_now())
            finalists = tuple(dict(row) for row in campaign.get("finalists", ()) if isinstance(row, Mapping))
            selected: tuple[dict[str, Any], ...] = ()
            variants: tuple[Any, ...] = ()
            excluded: tuple[dict[str, Any], ...] = ()
            optimizer_excluded: tuple[dict[str, Any], ...] = ()
            optimizer_blockers: list[str] = []
            optimizer_warnings: list[str] = []
            optimizer_status = "PASS"
            for index, stage in enumerate(STAGES):
                if self._cancelled(job_id):
                    self._cancel_finish(job_id, published, previous, published_executables, previous_executables, staged_executables)
                    published = None
                    published_executables = None
                    return
                self._append_journal(job_id, stage=stage, severity="INFO", code="STAGE_STARTED", text=stage)
                known_total = {
                    "LOAD_FINALISTS": len(finalists),
                    "SELECT_CANDIDATES": len(finalists),
                    "GENERATE_VARIANTS": len(launch["profiles"]),
                    "VALIDATE_VARIANTS": len(variants),
                }.get(stage)
                known_total = known_total if known_total and known_total > 0 else None
                self._set_progress(
                    job_id,
                    stage_index=index,
                    completed=index,
                    stage_total=known_total,
                    stage_percent=0 if known_total is not None else None,
                )
                if stage == "VALIDATE_SNAPSHOT":
                    continue
                if stage == "LOAD_FINALISTS":
                    self._sync_runtime(job_id, counters={"finalists_read": len(finalists)})
                elif stage == "SELECT_CANDIDATES":
                    selected_pairs = tuple((item[0], item[1]) for item in launch["selected_pairs"])
                    maximums = {tuple(key.rsplit("|", 1)): value for key, value in launch["maximums"].items()}
                    rows = self.cutoff_selector(finalists, selected_pairs=selected_pairs, maximums=maximums)
                    selected = tuple(dict(row) for row in (rows or ()) if row.get("selection_status") == "SELECTED")
                    excluded = tuple(dict(row) for row in (rows or ()) if row.get("selection_status") != "SELECTED")
                    if not selected:
                        raise PortfolioPanelError("PORTFOLIO_JOB_NO_ELIGIBLE_CANDIDATES", "no eligible candidates", status=422)
                elif stage == "GENERATE_VARIANTS":
                    generator = self.variant_generator or self._package_variant_generator
                    progress_generation = self._activate_progress(job_id)
                    self._progress_reporter.start(job_id)
                    heartbeat_stop = threading.Event()
                    heartbeat_thread = threading.Thread(
                        target=lambda: self._progress_heartbeat_loop(job_id, heartbeat_stop),
                        name="mrs3-portfolio-progress-heartbeat",
                        daemon=True,
                    )
                    heartbeat_thread.start()
                    generated = _invoke(generator, selected, campaign, launch["profiles"], self._progress_callback(job_id, progress_generation))
                    if isinstance(generated, Mapping) and "variants" in generated:
                        variants = tuple(generated.get("variants") or ())
                        optimizer_blockers = [str(item) for item in generated.get("blockers", ()) if isinstance(item, str)]
                        generated_excluded = generated.get("excluded", ())
                        optimizer_excluded = tuple(dict(item) for item in generated_excluded if isinstance(item, Mapping))
                        optimizer_warnings = [str(item) for item in generated.get("warnings", ()) if isinstance(item, str)]
                        optimizer_status = str(generated.get("status", "PASS"))
                    else:
                        variants = tuple(generated or ())
                    if not variants and not optimizer_blockers:
                        optimizer_blockers = ["PORTFOLIO_JOB_VARIANTS_NOT_READY"]
                    variants, capped, cap_blockers = self._cap_variants(variants, launch["profiles"])
                    optimizer_excluded = tuple((*optimizer_excluded, *capped))
                    optimizer_blockers.extend(cap_blockers)
                    if not variants:
                        details = [*optimizer_blockers]
                        details.extend(
                            f"{item.get('symbol')}:{item.get('selection_reason')}"
                            for item in optimizer_excluded
                            if item.get("symbol") and item.get("selection_reason") not in {"DIRECTION_TOP_N", "MAX_CANDIDATES"}
                        )
                        message = "; ".join(dict.fromkeys(details)) or "portfolio variant generation produced no variants"
                        raise PortfolioPanelError(
                            "PORTFOLIO_JOB_VARIANTS_NOT_READY",
                            message,
                            status=422,
                        )
                elif stage == "VALIDATE_VARIANTS":
                    if self.variant_validator is not None:
                        checked = _invoke(self.variant_validator, variants, campaign)
                        if isinstance(checked, Mapping) and "variants" in checked:
                            variants = tuple(checked.get("variants") or ())
                            optimizer_blockers.extend(str(item) for item in checked.get("blockers", ()) if isinstance(item, str))
                            optimizer_excluded = tuple((*optimizer_excluded, *(dict(item) for item in checked.get("excluded", ()) if isinstance(item, Mapping))))
                        else:
                            variants = tuple(checked or ())
                    variants, capped, cap_blockers = self._cap_variants(variants, launch["profiles"])
                    optimizer_excluded = tuple((*optimizer_excluded, *capped))
                    optimizer_blockers.extend(cap_blockers)
                    if not variants:
                        raise PortfolioPanelError(
                            "PORTFOLIO_JOB_VARIANTS_NOT_READY",
                            "portfolio variant validation produced no variants",
                            status=422,
                        )
                elif stage == "BUILD_WORKBOOK":
                    final_workbook = self._results_root / campaign["campaign_id"] / "stage1.xlsx"
                    workbook = self._staging_target(campaign["campaign_id"])
                    executable_document = self._build_stage1_executables(campaign, variants)
                    staged_executables = self._staging_executables_path(campaign["campaign_id"])
                    executable_bytes = (_json(executable_document) + "\n").encode("utf-8")
                    self._replace_bytes(staged_executables, executable_bytes)
                    if staged_executables.read_bytes() != executable_bytes:
                        raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "stage 1 executable artifact could not be verified", status=500)
                    if self.workbook_builder is not None:
                        built = _invoke(self.workbook_builder, workbook, campaign, finalists, selected, variants, excluded)
                        workbook = Path(built) if built is not None else workbook
                    else:
                        self._write_workbook(workbook, campaign, finalists, selected, variants, excluded, optimizer_excluded=optimizer_excluded, blockers=optimizer_blockers, warnings=optimizer_warnings, optimizer_status=optimizer_status)
                    self._staging_path(workbook, campaign["campaign_id"])
                    summary = self._summary(campaign, finalists, selected, variants, excluded, optimizer_excluded=optimizer_excluded, blockers=optimizer_blockers, warnings=optimizer_warnings, optimizer_status=optimizer_status)
                    runtime_values = {"staged_workbook_path": str(workbook), "final_workbook_path": str(final_workbook), "summary": _plain(summary), "counters": {"finalists_read": len(finalists), "candidates_selected": len(selected), "variants_created": len(variants)}}
                    self._sync_runtime(job_id, **runtime_values)
                elif stage == "PUBLISH_RESULTS":
                    pass
                self._set_progress(job_id, stage_index=index + 1 if index + 1 < len(STAGES) else index, completed=index + 1)
            if self._cancelled(job_id):
                self._cancel_finish(job_id, published, previous, published_executables, previous_executables, staged_executables)
                published = None
                published_executables = None
                return
            runtime = self.registry.runtime(job_id)
            staged = self._staging_path(Path(runtime["staged_workbook_path"]), campaign["campaign_id"])
            final = self._result_path(campaign["campaign_id"])
            final_executables = self._stage1_executables_path(campaign["campaign_id"], create=True)
            staged_executables = self._staging_executables_path(campaign["campaign_id"])
            if final.is_symlink():
                raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "result workbook is unsafe", status=500)
            if staged_executables.is_symlink() or staged_executables.resolve(strict=True).parent != staged.parent.resolve(strict=True):
                raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "staged executable artifact is unsafe", status=500)
            if final.is_file():
                previous = final.read_bytes()
            if final_executables.is_file():
                previous_executables = final_executables.read_bytes()
            os.replace(staged, final)
            published = final
            os.replace(staged_executables, final_executables)
            published_executables = final_executables
            self._verify_workbook(final)
            if executable_document is None:
                raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "stage 1 executable artifact is unavailable", status=500)
            verified_executables = self._load_stage1_executables(
                campaign["campaign_id"], campaign["input_digest"], campaign["config_digest"], executable_document["payload_digest"],
                require_committed=False,
            )
            if verified_executables["payload_digest"] != executable_document["payload_digest"]:
                raise PortfolioPanelError("PORTFOLIO_JOB_FAILED", "stage 1 executable artifact could not be verified", status=500)
            if self._cancelled(job_id):
                self._cancel_finish(job_id, published, previous, published_executables, previous_executables, staged_executables)
                published = None
                published_executables = None
                return
            self._append_journal(job_id, stage="PUBLISH_RESULTS", severity="INFO", code="COMPLETED", text="stage1 workbook published")
            runtime = self.registry.runtime(job_id)
            committed_runtime = {
                **runtime,
                "workbook_path": str(final),
                "executables_path": str(final_executables),
                "executables_digest": verified_executables["payload_digest"],
                "finished_at": _now(),
            }
            try:
                self.registry.sync(job_id, {"state": "COMMITTED", "phase": "COMMITTED", "result": runtime.get("summary", {})}, runtime=committed_runtime)
            except BaseException:
                self._rollback_result(final, previous)
                self._rollback_result(final_executables, previous_executables)
                published = None
                published_executables = None
                raise
            return
        except (KeyboardInterrupt, SystemExit, GeneratorExit, asyncio.CancelledError):
            raise
        except BaseException as error:
            if self._cancelled(job_id):
                try:
                    self._cancel_finish(job_id, published, previous, published_executables, previous_executables, staged_executables)
                except PanelJobError:
                    pass
                return
            if published is not None:
                self._rollback_result(published, previous)
            if published_executables is not None:
                self._rollback_result(published_executables, previous_executables)
            if staged_executables is not None:
                try:
                    staged_executables.unlink(missing_ok=True)
                except OSError:
                    pass
            try:
                failed_runtime_value = self.registry.runtime(job_id)
                failed_runtime = dict(failed_runtime_value) if isinstance(failed_runtime_value, Mapping) else {}
                candidate = failed_runtime.get("staged_workbook_path")
                if not isinstance(candidate, str):
                    campaign = failed_runtime.get("campaign")
                    campaign_id = campaign.get("campaign_id") if isinstance(campaign, Mapping) else None
                    if isinstance(campaign_id, str):
                        fallback = self._staging_root / campaign_id / "stage1.xlsx"
                        candidate = str(fallback) if fallback.exists() or fallback.is_symlink() else None
                if isinstance(candidate, str):
                    try:
                        self._staging_path(Path(candidate), failed_runtime["campaign"]["campaign_id"]).unlink(missing_ok=True)
                    except (KeyError, OSError, PortfolioPanelError):
                        pass
            except PanelJobError:
                pass
            try:
                current = self.registry.get(job_id)
                if current["state"] not in _TERMINAL:
                    self.registry.transition(job_id, "FAILED", phase="FAILED")
                code = error.code if isinstance(error, PortfolioPanelError) else "PORTFOLIO_JOB_FAILED"
                message = str(error) if isinstance(error, PortfolioPanelError) else "portfolio calculation failed"
                diagnostics = [{"severity": "ERROR", "code": code, "message": _redact_text(message)}]
                failed_runtime_value = self.registry.runtime(job_id)
                failed_runtime = dict(failed_runtime_value) if isinstance(failed_runtime_value, Mapping) else {}
                failed_runtime.pop("workbook_path", None)
                failed_runtime.pop("executables_path", None)
                failed_runtime.pop("executables_digest", None)
                entries = failed_runtime.get("journal") if isinstance(failed_runtime.get("journal"), list) else []
                stage_index = failed_runtime.get("stage_index")
                stage = STAGES[stage_index] if isinstance(stage_index, int) and 0 <= stage_index < len(STAGES) else "FAILED"
                journal_code = code if code.startswith(("PORTFOLIO_JOB_", "CAMPAIGN_")) else f"PORTFOLIO_JOB_{code}"
                entries.append({"timestamp_utc": _now(), "stage": stage, "severity": "ERROR", "code": journal_code, "text": _redact_text(message), "counters": _plain(failed_runtime.get("counters", {}))})
                failed_runtime["journal"] = entries[-200:]
                self.registry.sync(job_id, {"state": "FAILED", "phase": "FAILED", "error": {"code": code, "message": _redact_text(message)}, "result": {}}, runtime={**failed_runtime, "diagnostics": diagnostics, "finished_at": _now()})
            except PanelJobError:
                pass
        finally:
            self._deactivate_progress(job_id)
            self._progress_reporter.stop(job_id)
            if heartbeat_stop is not None:
                heartbeat_stop.set()
            if heartbeat_thread is not None and hasattr(heartbeat_thread, "join"):
                heartbeat_thread.join(timeout=1.0)
            self._cancel_events.pop(job_id, None)
            self._threads.pop(job_id, None)
            self._cleanup_terminal_snapshot(job_id)

    @staticmethod
    def _summary(campaign: Mapping[str, Any], finalists: Sequence[Mapping[str, Any]], selected: Sequence[Mapping[str, Any]], variants: Sequence[Any], excluded: Sequence[Mapping[str, Any]], *, optimizer_excluded: Sequence[Mapping[str, Any]] = (), blockers: Sequence[str] = (), warnings: Sequence[str] = (), optimizer_status: str | None = None) -> dict[str, Any]:
        blockers = list(dict.fromkeys(str(item) for item in blockers if item))
        warnings = list(dict.fromkeys(str(item) for item in warnings if item))
        finalist_reasons: dict[str, int] = {}
        for item in excluded:
            reason = str(item.get("selection_reason") or "UNKNOWN")
            finalist_reasons[reason] = finalist_reasons.get(reason, 0) + 1
        optimizer_reasons: dict[str, int] = {}
        for item in optimizer_excluded:
            reason = str(item.get("selection_reason") or item.get("reason") or "UNKNOWN")
            optimizer_reasons[reason] = optimizer_reasons.get(reason, 0) + 1
        variants_by_profile: dict[str, int] = {}
        for item in variants:
            profile = str(_get(item, "profile", "UNASSIGNED") or "UNASSIGNED")
            variants_by_profile[profile] = variants_by_profile.get(profile, 0) + 1
        prepared_by_profile = {
            str(profile["profile_id"]): variants_by_profile.get(str(profile["profile_id"]), 0)
            for profile in campaign.get("launch", {}).get("profiles", ())
            if isinstance(profile, Mapping)
        }
        status = optimizer_status or ("PASS" if variants and not blockers else ("PARTIAL" if variants else "FAIL"))
        summary = {
            "optimizer_status": status,
            "finalists_read": len(finalists),
            "candidates_selected": len(selected),
            "variants_created": len(variants),
            "variants_by_profile": dict(sorted(variants_by_profile.items())),
            "prepared": sum(prepared_by_profile.values()),
            "prepared_by_profile": dict(sorted(prepared_by_profile.items())),
            "excluded": len(excluded),
            "finalist_exclusions_by_reason": dict(sorted(finalist_reasons.items())),
            "optimizer_excluded": len(optimizer_excluded),
            "optimizer_exclusions_by_reason": dict(sorted(optimizer_reasons.items())),
            "blockers": blockers,
            "optimizer_blockers": blockers,
            "warnings": warnings,
            "optimizer_warnings": warnings,
            "campaign_id": campaign["campaign_id"],
        }
        weighted_variants = tuple(item for item in variants if _get(item, "search_mode") == CAMPAIGN_SEARCH_MODE)
        if weighted_variants:
            variant = weighted_variants[0]
            metrics = _get(variant, "metrics", {})
            if not isinstance(metrics, Mapping):
                metrics = {}
            profile_id = _get(variant, "profile", _get(variant, "profile_id"))
            config_profiles = _get(_get(campaign, "config_document", {}), "profiles", {})
            profile_config = _get(config_profiles, profile_id, {}) if isinstance(config_profiles, Mapping) else {}
            profile_dd_limit = _weighted_value(
                profile_config,
                "max_actual_equity_dd_pct", "maximum_actual_equity_dd_pct", "dd_limit_pct",
                default=_weighted_value(variant, "maximum_individual_dd_pct", "max_actual_equity_dd_pct", default="UNKNOWN"),
            )
            summary["Weighted candidate ID"] = _weighted_value(variant, "candidate_id", "strategy_id", "id")
            available = _weighted_value(metrics, "B_available_usdt", "bank_available_usdt", "available_bank_usdt")
            if available == "UNKNOWN":
                profiles = _get(_get(campaign, "launch", {}), "profiles", ())
                if isinstance(profiles, Sequence) and not isinstance(profiles, (str, bytes)):
                    profile = next((item for item in profiles if _get(item, "profile_id") == profile_id), None)
                    available = _weighted_value(profile, "bank_available_usdt", default="UNCAPPED") if profile is not None else "UNKNOWN"
            payload_pairs = _weighted_payload_pairs(variant)
            max_balances = tuple(
                _weighted_number(_get(_get(payload, "strategy", {}), "basic", {}), "max_balance")
                for _member, payload in payload_pairs
                if isinstance(payload, Mapping)
                and isinstance(_get(payload, "strategy", {}), Mapping)
                and isinstance(_get(_get(payload, "strategy", {}), "basic", {}), Mapping)
                and _weighted_number(_get(_get(payload, "strategy", {}), "basic", {}), "max_balance") != "UNKNOWN"
            )
            saturation = _weighted_value(metrics, "B_sat_settings_usdt", "B_saturation_usdt", "B_sat_usdt")
            if saturation == "UNKNOWN" and max_balances:
                saturation = max(max_balances, key=lambda value: Decimal(str(value)))
            before_reserve = _weighted_value(metrics, "reserve_before_limiter_pct", "reserve_full_load_pct", "reserve_before_pct")
            after_reserve = _weighted_value(metrics, "reserve_after_limiter_pct", "reserve_after_pct")
            release_status = _weighted_value(metrics, "limiter_release_status")
            reserve_reason = _weighted_value(metrics, "reserve_unknown_reason")
            required_bank = _weighted_value(metrics, "B_required_usdt", "required_bank_usdt", "B_required_margin_usdt")
            historical_bank = _weighted_value(metrics, "historical_bank_usdt", "bank_for_path_usdt")
            stress_bank = _weighted_value(metrics, "B_risk_usdt", "stress_p95_bank_usdt", "stress_bank_p95_usdt")
            margin_bank = _weighted_value(metrics, "B_margin_usdt")
            if stress_bank == "UNKNOWN":
                bootstrap_banks = _weighted_value(metrics, "bootstrap_p95_banks_usdt", default=())
                flattened = []
                def collect_bootstrap(value: Any) -> None:
                    if isinstance(value, (str, bytes)):
                        return
                    if isinstance(value, Sequence):
                        for item in value:
                            collect_bootstrap(item)
                        return
                    try:
                        number = Decimal(str(value))
                    except (InvalidOperation, TypeError, ValueError):
                        return
                    if number.is_finite() and number > 0:
                        flattened.append(number)
                collect_bootstrap(bootstrap_banks)
                if flattened:
                    stress_bank = max(flattened)
            source_evidence = _weighted_source_evidence(campaign)

            def percentage_of_saturation(value: Any) -> Any:
                if value == "UNKNOWN" or saturation == "UNKNOWN":
                    return "UNKNOWN"
                try:
                    amount = Decimal(str(value))
                    denominator = Decimal(str(saturation))
                except (InvalidOperation, TypeError, ValueError):
                    return "UNKNOWN"
                return amount / denominator * Decimal(100) if amount.is_finite() and denominator.is_finite() and denominator > 0 else "UNKNOWN"

            def percentage_of(value: Any, denominator: Any) -> Any:
                if value == "UNKNOWN" or denominator in (None, "UNKNOWN", "UNCAPPED"):
                    return "UNKNOWN"
                try:
                    amount = Decimal(str(value))
                    base = Decimal(str(denominator))
                except (InvalidOperation, TypeError, ValueError):
                    return "UNKNOWN"
                return amount / base * Decimal(100) if amount.is_finite() and base.is_finite() and base > 0 else "UNKNOWN"

            def percentage_of_target(value: Any) -> Any:
                return "—" if available == "UNCAPPED" else percentage_of(value, available)

            weighted_members = []
            total_notional = Decimal(0)
            for member, payload in payload_pairs:
                facts = _get(payload, "facts", {})
                basic = _get(_get(payload, "strategy", {}), "basic", {})
                raw_x = _get(member, "x_usdt")
                if raw_x is None:
                    raw_x = _get(facts, "x")
                try:
                    position = Decimal(str(raw_x))
                except (InvalidOperation, TypeError, ValueError):
                    continue
                if not position.is_finite() or position <= 0:
                    continue
                total_notional += position
                q = _weighted_number(facts, "q")
                try:
                    bank_share = Decimal(str(q)) * Decimal(100)
                    if not bank_share.is_finite() or bank_share <= 0:
                        bank_share = None
                except (InvalidOperation, TypeError, ValueError):
                    bank_share = None
                def known(value: Any) -> Any:
                    return None if value == "UNKNOWN" else value
                direction = known(_weighted_value(member, "side", "direction", default=_weighted_value(payload, "side")))
                balance_percentage = _weighted_value(basic, f"balance_percentage_{str(direction).lower()}", default=bank_share if bank_share is not None else "UNKNOWN")
                source = source_evidence.get(_weighted_identity(member), {})
                if not isinstance(source, Mapping) or source.get("__invalid_source_evidence"):
                    source = {}
                source_metrics = _get(source, "metrics", {})
                source_pnl = _weighted_value(source, "total_pnl", "source_pnl", "source_pnl_usdt", "pnl", default=_weighted_value(source_metrics, "total_pnl", "source_pnl", "pnl"))
                source_maxdd = _weighted_value(source, "max_drawdown", "max_drawdown_usdt", "source_max_drawdown", default=_weighted_value(source_metrics, "max_drawdown", "max_drawdown_usdt"))
                source_maxdd_pct = _weighted_value(source, "max_drawdown_pct", "source_max_drawdown_pct", default=_weighted_value(source_metrics, "max_drawdown_pct", "source_max_drawdown_pct"))
                source_initial = _weighted_value(source, "source_initial_balance", "initial_balance", "result_initial_balance", default=_weighted_value(source_metrics, "source_initial_balance", "initial_balance"))
                source_dd_number = _weighted_decimal_or_unknown(source_maxdd, nonnegative=True)
                source_initial_number = _weighted_decimal_or_unknown(source_initial)
                scaled_maxdd = (
                    source_dd_number * position / source_initial_number
                    if isinstance(source_dd_number, Decimal) and isinstance(source_initial_number, Decimal)
                    else "UNKNOWN"
                )
                capacity = known(_weighted_value(member, "capacity_usdt", "capacity", default=_weighted_value(facts, "C", "capacity_usdt")))
                liquidity_utilization = percentage_of(position, capacity)
                mrs3 = _get(_get(payload, "strategy", {}), "mrs3", {})
                order_key = "ma_long" if str(direction).upper() == "LONG" else "ma_short"
                payload_orders = _get(mrs3, order_key, ())
                order_count = len(payload_orders) if isinstance(payload_orders, Sequence) and not isinstance(payload_orders, (str, bytes)) and payload_orders else _weighted_value(source, "order_count", default=_weighted_value(source_metrics, "order_count"))
                weighted_members.append({
                    "pair": known(_weighted_value(member, "symbol", "pair", default=_weighted_value(basic, "symbol"))),
                    "direction": direction,
                    "strategy_id": known(_weighted_value(member, "strategy_id", "strategyId")),
                    "result_id": known(_weighted_value(member, "result_id", "resultId")),
                    "position_usdt": position,
                    "bank_share_pct": bank_share,
                    "balance_percentage": known(balance_percentage),
                    "pair_multiplier_pct": known(balance_percentage),
                    "max_balance": known(_weighted_number(basic, "max_balance")),
                    "entry_order_percentages": _weighted_entry_order_percentages(payload, direction),
                    "timeframe": known(_weighted_value(source, "timeframe", default=_weighted_value(source_metrics, "timeframe"))),
                    "user_rank": known(_weighted_value(source, "user_rank", default=_weighted_value(source_metrics, "user_rank"))),
                    "source_pnl": known(source_pnl),
                    "source_max_drawdown_usdt": known(source_maxdd),
                    "source_max_drawdown_pct": known(source_maxdd_pct),
                    "order_count": known(order_count),
                    "liquidity_utilization_pct": known(liquidity_utilization),
                    "scaled_max_drawdown_usdt": known(scaled_maxdd),
                    "leverage": known(_weighted_number(basic, "leverage", "planned_leverage", "max_leverage")),
                    "capacity_usdt": capacity,
                })
            summary.update({
                "weighted_result_schema": 1,
                "B required USDT": required_bank,
                "Required bank USDT": required_bank,
                "Historical bank USDT": historical_bank,
                "Stress bank P95 USDT": stress_bank,
                "Margin-only bank USDT": margin_bank,
                "B available USDT": available,
                "Target bank USDT": available,
                "Profile DD limit %": profile_dd_limit,
                "B saturation USDT": saturation,
                "B margin USDT": _weighted_value(metrics, "B_margin_usdt"),
                "P30 common USDT/30d": _weighted_value(metrics, "p30_common_usdt_30d"),
                "P30 limiter USDT/30d": _weighted_value(metrics, "p30_limiter_model_usdt_30d"),
                "MaxDD %": _weighted_value(metrics, "max_drawdown_pct"),
                "CDaR peak80 USDT": _weighted_value(metrics, "cdar_peak80_usdt"),
                "CDaR peak90 USDT": _weighted_value(metrics, "cdar_peak90_usdt"),
                "Reserve before limiter": before_reserve,
                "Reserve after limiter": after_reserve,
                "Reserve UNKNOWN reason": reserve_reason,
                "MM all USDT": _weighted_value(metrics, "M_all_usdt"),
                "Bottleneck": _weighted_value(metrics, "bottleneck", "limiter_bottleneck", "margin_bottleneck"),
                "Limiter L": _weighted_value(metrics, "limiter_L"),
                "Limiter P30 status": _weighted_value(metrics, "limiter_p30_status", "p30_status"),
                "Limiter release status": release_status,
                "Joint status": _weighted_value(metrics, "joint_status", "joint_metrics", default="NOT_TESTED"),
                "Search status": _weighted_value(variant, "status", default=_weighted_value(metrics, "status", default="NOT_TESTED")),
                "Budget status": _weighted_value(metrics, "budget_status", "budget_limited"),
                "IM all USDT": _weighted_value(metrics, "I_all_usdt"),
                "CDaR peak80 %": percentage_of_saturation(_weighted_value(metrics, "cdar_peak80_usdt")),
                "CDaR peak90 %": percentage_of_saturation(_weighted_value(metrics, "cdar_peak90_usdt")),
                "CDaR peak80 target %": percentage_of_target(_weighted_value(metrics, "cdar_peak80_usdt")),
                "CDaR peak90 target %": percentage_of_target(_weighted_value(metrics, "cdar_peak90_usdt")),
                "CDaR peak80 saturation %": percentage_of_saturation(_weighted_value(metrics, "cdar_peak80_usdt")),
                "CDaR peak90 saturation %": percentage_of_saturation(_weighted_value(metrics, "cdar_peak90_usdt")),
                "IM target %": percentage_of_target(_weighted_value(metrics, "I_all_usdt")),
                "IM saturation %": percentage_of_saturation(_weighted_value(metrics, "I_all_usdt")),
                "MM target %": percentage_of_target(_weighted_value(metrics, "M_all_usdt")),
                "MM saturation %": percentage_of_saturation(_weighted_value(metrics, "M_all_usdt")),
                "MaxDD SUM USDT": _weighted_source_maxdd_sum(campaign, weighted_members),
                "Total full notional USDT": total_notional,
                "Weighted members": weighted_members,
            })
        return summary

    @staticmethod
    def _cap_variants(variants: Sequence[Any], profiles: Sequence[Mapping[str, Any]]) -> tuple[tuple[Any, ...], tuple[dict[str, Any], ...], tuple[str, ...]]:
        limits = {
            str(profile["profile_id"]): int(profile["max_candidates"])
            for profile in profiles
            if isinstance(profile, Mapping) and isinstance(profile.get("profile_id"), str) and isinstance(profile.get("max_candidates"), int)
        }
        kept: list[Any] = []
        excluded: list[dict[str, Any]] = []
        blockers: list[str] = []
        for variant in variants:
            profile = _get(variant, "profile")
            profile_id = str(profile) if profile is not None else ""
            if profile_id not in limits:
                blockers.append("UNKNOWN_PROFILE")
                excluded.append(
                    {
                        "profile": profile_id,
                        "stage": "GENERATE_VARIANTS",
                        "selection_reason": "UNKNOWN_PROFILE",
                        "message": "UNKNOWN_PROFILE",
                        "candidate_id": next((_get(variant, key) for key in ("candidate_id", "strategy_id", "id") if _get(variant, key) is not None), ""),
                    }
                )
                continue
            kept.append(variant)
        return tuple(kept), tuple(excluded), tuple(dict.fromkeys(blockers))

    def _write_workbook(self, path: Path, campaign: Mapping[str, Any], finalists: Sequence[Mapping[str, Any]], selected: Sequence[Mapping[str, Any]], variants: Sequence[Any], excluded: Sequence[Mapping[str, Any]], *, optimizer_excluded: Sequence[Mapping[str, Any]] = (), blockers: Sequence[str] = (), warnings: Sequence[str] = (), optimizer_status: str | None = None) -> Path:
        def val(item: Any, *keys: str, default: Any = None) -> Any:
            for key in keys:
                found = _get(item, key, None)
                if found is not None:
                    return found
            return default

        finalist_rows = []
        for row in tuple(selected) + tuple(excluded):
            if not any(val(row, key) is not None for key in ("strategy_id", "strategyId", "result_id", "resultId", "symbol", "pair")):
                continue
            finalist_rows.append([campaign["campaign_id"], val(row, "strategy_id", "strategyId"), val(row, "result_id", "resultId"), val(row, "symbol", "pair"), val(row, "side", "direction"), val(row, "user_status", "status"), val(row, "user_rank"), val(row, "effective_maximum"), val(row, "selection_status"), val(row, "selection_reason")])
        portfolio_rows = []
        member_rows = []
        weighted_mode = any(val(item, "search_mode") == CAMPAIGN_SEARCH_MODE for item in variants)
        def display(value: Any) -> Any:
            if value is None or value == "UNKNOWN":
                return None
            try:
                number = Decimal(str(value))
            except (InvalidOperation, TypeError, ValueError):
                return value
            return format(number.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP), "f") if number.is_finite() else None

        def display_ratio(value: Any) -> str:
            if value in (None, "UNKNOWN", "NOT_TESTED"):
                return "—"
            try:
                number = Decimal(str(value))
            except (InvalidOperation, TypeError, ValueError):
                return str(value)
            return format(number.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP), "f") + "%" if number.is_finite() else "—"

        def display_amount_ratios(amount: Any, target: Any, saturation: Any) -> Any:
            shown = display(amount)
            if shown is None:
                return None
            return f"{shown} USDT · {display_ratio(target)} / {display_ratio(saturation)}"

        def display_integer(value: Any) -> Any:
            if value in (None, "UNKNOWN", "NOT_TESTED"):
                return None
            try:
                number = Decimal(str(value))
            except (InvalidOperation, TypeError, ValueError):
                return value
            return str(int(number)) if number.is_finite() and number == number.to_integral_value() else str(value)

        for index, variant in enumerate(variants):
            candidate_id = val(variant, "candidate_id", "strategy_id", "id", default=f"candidate-{index + 1}")
            profile = val(variant, "profile", default="")
            directions = val(variant, "directions")
            members = val(variant, "members")
            members_is_sequence = isinstance(members, Sequence) and not isinstance(members, (str, bytes))
            if not members_is_sequence:
                members = ()
            member_count = len(directions) if isinstance(directions, Mapping) else (len(members) if members_is_sequence else val(variant, "member_count"))
            pair_count = val(variant, "pair_count")
            if pair_count is None:
                symbols = {
                    str(symbol)
                    for symbol in (
                        [val(details, "symbol", "pair") for details in directions.values()]
                        if isinstance(directions, Mapping)
                        else [val(details, "symbol", "pair") for details in members] if isinstance(members, Sequence)
                        else []
                    )
                    if symbol is not None
                }
                if not symbols:
                    symbol = val(variant, "symbol", "pair", default=val(val(variant, "slot"), "symbol"))
                    if symbol is not None:
                        symbols.add(str(symbol))
                pair_count = len(symbols) if symbols else None
            metrics = val(variant, "metrics", default={})
            period = val(variant, "pretest_period", default={})
            if val(variant, "search_mode") == CAMPAIGN_SEARCH_MODE:
                variant_summary = self._summary(campaign, (), (), (variant,), ())
                weighted_members = variant_summary.get("Weighted members", ())
                member_names = "; ".join(
                    f"{item.get('pair')} {item.get('direction')}"
                    for item in weighted_members
                    if isinstance(item, Mapping) and item.get("pair") and item.get("direction")
                )
                portfolio_rows.append([
                    index + 1,
                    candidate_id,
                    profile,
                    len(weighted_members),
                    member_names,
                    display(variant_summary.get("B saturation USDT")),
                    display(variant_summary.get("Target bank USDT")),
                    display(variant_summary.get("Historical bank USDT")),
                    display(variant_summary.get("Stress bank P95 USDT")),
                    display(variant_summary.get("Margin-only bank USDT")),
                    display(_weighted_value(metrics, "p30_limiter_model_usdt_30d", "p30_common_usdt_30d")),
                    display(variant_summary.get("MaxDD SUM USDT")) or "—",
                    display(_weighted_value(metrics, "max_drawdown_pct")) or "—",
                    display_amount_ratios(variant_summary.get("CDaR peak80 USDT"), variant_summary.get("CDaR peak80 target %"), variant_summary.get("CDaR peak80 saturation %")) or "—",
                    display_amount_ratios(variant_summary.get("CDaR peak90 USDT"), variant_summary.get("CDaR peak90 target %"), variant_summary.get("CDaR peak90 saturation %")) or "—",
                    display_amount_ratios(variant_summary.get("IM all USDT"), variant_summary.get("IM target %"), variant_summary.get("IM saturation %")) or "—",
                    display_amount_ratios(variant_summary.get("MM all USDT"), variant_summary.get("MM target %"), variant_summary.get("MM saturation %")) or "—",
                    display(variant_summary.get("Total full notional USDT")),
                ])
                for member in weighted_members:
                    if not isinstance(member, Mapping):
                        continue
                    order_percentages = member.get("entry_order_percentages")
                    order_columns = [
                        display_ratio(order_percentages[ordinal])
                        if isinstance(order_percentages, Sequence) and not isinstance(order_percentages, (str, bytes)) and ordinal < len(order_percentages)
                        else None
                        for ordinal in range(4)
                    ]
                    member_rows.append([
                        index + 1,
                        candidate_id,
                        profile,
                        member.get("pair"),
                        member.get("direction"),
                        member.get("strategy_id"),
                        member.get("result_id"),
                        display(member.get("position_usdt")),
                        display(member.get("balance_percentage", member.get("bank_share_pct"))),
                         display(member.get("max_balance")),
                         *order_columns,
                         display(member.get("leverage")),
                         display(member.get("capacity_usdt")),
                         member.get("timeframe"),
                         display_integer(member.get("user_rank")),
                         display(member.get("source_pnl")),
                         display(member.get("source_max_drawdown_usdt")),
                         display_ratio(member.get("source_max_drawdown_pct")),
                         display_integer(member.get("order_count")),
                         display_ratio(member.get("liquidity_utilization_pct")),
                         display(member.get("scaled_max_drawdown_usdt")),
                     ])
                continue
            portfolio_rows.append([campaign["campaign_id"], candidate_id, profile, val(variant, "final_pretest_rank", default=index + 1), val(variant, "scheduling_key_id", default=val(val(variant, "scheduling_key", default={}), "identity")), val(variant, "scheduling_score", "score"), member_count, pair_count, val(variant, "limiter"), val(variant, "maximum_individual_dd_pct"), val(variant, "minimum_free_margin_reserve_pct"), val(variant, "maximum_account_mm_load_pct"), val(variant, "gate", "gate_result", default="UNKNOWN"), val(variant, "blocking_reasons", "reasons", default=""), "", None, "", val(variant, "search_mode", default=val(metrics, "metric_basis")), val(variant, "evaluations", default="UNKNOWN"), val(metrics, "proxy_pnl_usdt", default=val(metrics, "proxy_end_pnl_usdt")), val(metrics, "proxy_max_drawdown_usdt"), val(metrics, "proxy_max_drawdown_pct"), val(metrics, "proxy_recovery_factor"), val(metrics, "proxy_reserve_usdt"), val(metrics, "proxy_reserve_pct"), val(metrics, "k"), val(metrics, "tested_size_usdt"), val(metrics, "actual_size_usdt"), val(metrics, "tested_size_basis", default=val(metrics, "sizing_basis")), f"{val(period, 'start_utc')}..{val(period, 'end_utc')}" if val(period, 'start_utc') and val(period, 'end_utc') else "UNKNOWN", val(variant, "refinement", default=val(metrics, "refinement", default="DAILY")), val(metrics, "k1"), val(metrics, "corrective_reduction_applied", default=False), val(variant, "daily_pretest_rank"), val(variant, "final_pretest_rank")])
            if isinstance(directions, Mapping) and directions:
                for ordinal, (direction, details) in enumerate(directions.items(), 1):
                    member_rows.append([campaign["campaign_id"], candidate_id, profile, ordinal, val(details, "strategy_id", "strategyId"), val(details, "result_id", "resultId"), val(details, "symbol", "pair", default=val(variant, "symbol")), direction, val(details, "user_rank"), val(details, "scalar", "scalar_pct"), val(details, "quantity", "rounded_quantity"), val(details, "leverage"), val(details, "notional_usdt"), val(details, "estimated_individual_dd_usdt"), val(details, "estimated_individual_dd_pct"), val(details, "liquidity_scalar_ceiling_pct"), val(details, "calculated_initial_margin_usdt"), val(details, "gate", "gate_result", default="UNKNOWN"), val(details, "reasons", default=""), *([None] * 12)])
            elif isinstance(members, Sequence):
                weighted = val(variant, "search_mode") == CAMPAIGN_SEARCH_MODE
                payload_pairs = _weighted_payload_pairs(variant) if weighted else ()
                for ordinal, details in enumerate(members, 1):
                    if weighted:
                        payload = payload_pairs[ordinal - 1][1] if ordinal <= len(payload_pairs) else None
                        facts = _get(payload, "facts", {})
                        basic = _get(_get(payload, "strategy", {}), "basic", {})
                        strategy_id = val(details, "strategy_id", "strategyId")
                        result_id = val(details, "result_id", "resultId")
                        finalist = f"{strategy_id}/{result_id}" if strategy_id is not None and result_id is not None else "UNKNOWN"
                        warnings_value = val(details, "warnings", "warning", default="UNKNOWN")
                        member_rows.append([
                            campaign["campaign_id"], candidate_id, profile, ordinal, strategy_id, result_id,
                            val(details, "symbol", "pair", default="UNKNOWN"), val(details, "side", default="UNKNOWN"),
                            val(details, "user_rank", default="UNKNOWN"), "UNKNOWN", "UNKNOWN", "UNKNOWN", "UNKNOWN",
                            "UNKNOWN", "UNKNOWN", "UNKNOWN", "UNKNOWN", val(variant, "gate", "gate_result", default="UNKNOWN"),
                            warnings_value, "UNKNOWN", "UNKNOWN", "UNKNOWN", "UNKNOWN", "UNKNOWN", "UNKNOWN", "UNKNOWN",
                            "UNKNOWN", "UNKNOWN", "UNKNOWN", "UNKNOWN", "UNKNOWN",
                            finalist, val(details, "side", default="UNKNOWN"), val(details, "x_usdt", default="UNKNOWN"),
                            val(details, "capacity_usdt", default="UNKNOWN"), _weighted_number(facts, "q"),
                            val(details, "priority", default="UNKNOWN"), _weighted_number(basic, "max_balance"),
                            val(details, "source_scale", "source_scale_pct", "scale", "size_scale"),
                            val(details, "hold90", "hold90_hours", "hold_90", "p90_hold"), warnings_value,
                        ])
                        continue
                    calendar = val(details, "calendar_7d", default={})
                    weekday = val(details, "weekday_5d", default={})
                    member_rows.append([campaign["campaign_id"], candidate_id, profile, ordinal, val(details, "strategy_id", "strategyId"), val(details, "result_id", "resultId"), val(details, "symbol", "pair"), val(details, "side", "direction"), val(details, "user_rank"), None, val(details, "quantity", "maximum_closing_quantity"), val(details, "planned_leverage"), val(details, "actual_size_usdt", "position_size_usdt"), None, val(details, "max_drawdown_pct"), None, None, "PASS", val(details, "spread_diagnostics", default=""), val(details, "capacity_status"), val(calendar, "available_days"), val(calendar, "mean_minute_turnover"), val(calendar, "rounded_cap_usdt"), val(weekday, "available_days"), val(weekday, "mean_minute_turnover"), val(weekday, "rounded_cap_usdt"), val(details, "spread_status"), val(details, "spread_mean_bps"), val(details, "sizing_digest"), val(details, "capacity_digest"), val(details, "reference_digest"), *([None] * 10)])
        profile_status_rows = []
        pretest_mode = any(val(item, "search_mode") == "PRETEST_PROXY" or (isinstance(val(item, "metrics"), Mapping) and val(item, "metrics").get("metric_basis") == "PRETEST_PROXY") for item in variants)
        if pretest_mode:
            configured_budget = val(campaign.get("config_document", {}).get("search", {}), "max_enumerated_combinations", default="UNKNOWN")
            for profile in campaign.get("launch", {}).get("profiles", ()):
                if not isinstance(profile, Mapping):
                    continue
                profile_id = str(profile.get("profile_id", ""))
                profile_variants = tuple(item for item in variants if str(val(item, "profile", default="")) == profile_id)
                profile_blockers = tuple(str(item) for item in blockers if str(item).startswith(f"{profile_id}:"))
                status = "PASS" if profile_variants and not profile_blockers else "FAILED"
                first_metrics = val(profile_variants[0], "metrics", default={}) if profile_variants else {}
                first_period = val(profile_variants[0], "pretest_period", default={}) if profile_variants else {}
                coverage = "UNKNOWN"
                if isinstance(first_period, Mapping):
                    coverage_map = first_period.get("coverage_pct")
                    if isinstance(coverage_map, Mapping):
                        coverage = ", ".join(f"{key}={value}" for key, value in sorted(coverage_map.items(), key=lambda item: str(item[0])))
                profile_status_rows.append([
                    campaign["campaign_id"], profile_id, status, profile.get("max_candidates", "UNKNOWN"),
                    val(profile_variants[0], "evaluations", default=0) if profile_variants else 0,
                    configured_budget,
                    f"{first_period.get('start_utc')}..{first_period.get('end_utc')}" if isinstance(first_period, Mapping) and first_period.get("start_utc") and first_period.get("end_utc") else "UNKNOWN",
                    coverage, "; ".join(profile_blockers) or "", val(first_metrics, "metric_basis", default="PRETEST_PROXY"), val(first_metrics, "joint_metrics", default="NOT_TESTED"),
                ])
        excluded_rows = []
        for is_optimizer, rows in ((False, excluded), (True, optimizer_excluded)):
            excluded_rows.extend(
                [campaign["campaign_id"], "PORTFOLIO" if is_optimizer else "FINALIST", val(row, "strategy_id", "result_id", default=""), val(row, "symbol", "pair"), val(row, "side", "direction"), val(row, "profile"), val(row, "stage", default="GENERATE_VARIANTS" if is_optimizer else "SELECT_CANDIDATES"), "BLOCKED" if is_optimizer else "EXCLUDED", val(row, "selection_reason"), _redact_text(val(row, "message", default=val(row, "selection_reason", default="")))]
                for row in rows
            )
        summary = self._summary(campaign, finalists, selected, variants, excluded, optimizer_excluded=optimizer_excluded, blockers=blockers, warnings=warnings, optimizer_status=optimizer_status)
        metadata = {"schema_version": "portfolio_panel_stage1_v1", "campaign_id": campaign["campaign_id"], "created_at_utc": campaign.get("created_at_utc", "2000-01-01T00:00:00Z"), "input_digest": campaign["input_digest"], "config_digest": campaign["config_digest"], "policy_version": campaign["versions"]["policy_version"], "algorithm_versions": _json(campaign["versions"])}
        def frame(rows: Sequence[Sequence[Any]], headers: Sequence[str]) -> pd.DataFrame:
            return pd.DataFrame([[ _safe_cell(value, key=header) for value, header in zip(row, headers)] for row in rows], columns=headers)

        if weighted_mode:
            weighted_variant = next((item for item in variants if val(item, "search_mode") == CAMPAIGN_SEARCH_MODE), None)
            weighted_summary = self._summary(campaign, finalists, selected, variants, excluded, optimizer_excluded=optimizer_excluded, blockers=blockers, warnings=warnings, optimizer_status=optimizer_status)
            top_metrics = val(weighted_variant, "metrics", default={}) if weighted_variant is not None else {}
            top_members = weighted_summary.get("Weighted members", ())
            limit_text = display(weighted_summary.get("Profile DD limit %")) or "—"
            summary_rows = [
                ["ID варианта", weighted_summary.get("Weighted candidate ID")],
                ["Профиль", val(weighted_variant, "profile") if weighted_variant is not None else None],
                ["Целевой банк, USDT", display(weighted_summary.get("Target bank USDT")) or ("UNCAPPED" if weighted_summary.get("Target bank USDT") == "UNCAPPED" else None)],
                [f"Минимальный банк для DD ≤ {limit_text}% на истории", display(weighted_summary.get("Historical bank USDT"))],
                [f"Банк для DD ≤ {limit_text}% в 95% стресс-сценариев", display(weighted_summary.get("Stress bank P95 USDT"))],
                ["Банк для профильных лимитов, USDT", display(weighted_summary.get("Margin-only bank USDT"))],
                ["PnL, USDT", display(_weighted_value(top_metrics, "p30_limiter_model_usdt_30d", "p30_common_usdt_30d"))],
                ["MaxDD SUM, USDT", display(weighted_summary.get("MaxDD SUM USDT")) or "—"],
                ["Исторический рассчитанный DD, %", display(_weighted_value(top_metrics, "max_drawdown_pct")) or "—"],
                ["CDaR худшие 20%, USDT · target/saturation", display_amount_ratios(weighted_summary.get("CDaR peak80 USDT"), weighted_summary.get("CDaR peak80 target %"), weighted_summary.get("CDaR peak80 saturation %")) or "—"],
                ["CDaR худшие 10%, USDT · target/saturation", display_amount_ratios(weighted_summary.get("CDaR peak90 USDT"), weighted_summary.get("CDaR peak90 target %"), weighted_summary.get("CDaR peak90 saturation %")) or "—"],
                ["IM", display_amount_ratios(weighted_summary.get("IM all USDT"), weighted_summary.get("IM target %"), weighted_summary.get("IM saturation %")) or "—"],
                ["MM", display_amount_ratios(weighted_summary.get("MM all USDT"), weighted_summary.get("MM target %"), weighted_summary.get("MM saturation %")) or "—"],
                ["Полный номинал портфеля, USDT", display(weighted_summary.get("Total full notional USDT"))],
                ["Позиций", len(top_members) if isinstance(top_members, Sequence) else 0],
            ]
            tables = {
                "Итог": frame(summary_rows, WEIGHTED_SUMMARY_HEADERS),
                "Варианты": frame(portfolio_rows, WEIGHTED_VARIANT_HEADERS),
                "Состав": frame(member_rows, WEIGHTED_MEMBER_HEADERS),
                "Финалисты": frame(finalist_rows, WEIGHTED_FINALIST_HEADERS),
                "Исключено": frame(excluded_rows, WEIGHTED_EXCLUDED_HEADERS),
                "Metadata": frame([[key, value] for key, value in metadata.items()], METADATA_HEADERS),
            }
        else:
            tables = {
                "Summary": frame([[key, value] for key, value in summary.items()], SUMMARY_HEADERS),
                "Finalists": frame(finalist_rows, FINALIST_HEADERS),
                "Portfolios": frame(portfolio_rows, PORTFOLIO_HEADERS),
                "Members": frame(member_rows, MEMBER_HEADERS),
                "Excluded": frame(excluded_rows, EXCLUDED_HEADERS),
                "Metadata": frame([[key, value] for key, value in metadata.items()], METADATA_HEADERS),
            }
            if pretest_mode:
                tables["Profile Status"] = frame(profile_status_rows, PROFILE_STATUS_HEADERS)
        write_audit_workbook(tables, path)
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment
        workbook = load_workbook(path)
        temporary: Path | None = None
        try:
            workbook["Metadata"].sheet_state = "hidden"
            if weighted_mode:
                width_limits = {"Итог": 34, "Варианты": 20, "Состав": 18, "Финалисты": 18, "Исключено": 20}
                for sheet_name, maximum_width in width_limits.items():
                    worksheet = workbook[sheet_name]
                    worksheet.row_dimensions[1].height = 48
                    for cell in worksheet[1]:
                        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                        dimension = worksheet.column_dimensions[cell.column_letter]
                        dimension.width = min(maximum_width, max(12, float(dimension.width or 12)))
            with tempfile.NamedTemporaryFile(suffix=".xlsx", dir=path.parent, delete=False) as handle:
                temporary = Path(handle.name)
            workbook.save(temporary)
            workbook.close()
            workbook = None
            normalize_xlsx_workbook(temporary)
            os.replace(temporary, path)
            temporary = None
        finally:
            if workbook is not None:
                workbook.close()
            if temporary is not None:
                temporary.unlink(missing_ok=True)
        return path

    def cancel(self, job_id: str) -> dict[str, Any]:
        try:
            saved = self.registry.get(job_id)
        except PanelJobError as error:
            code = "PORTFOLIO_JOB_NOT_FOUND" if error.code == "NOT_FOUND" else "PORTFOLIO_JOB_RUNTIME_UNAVAILABLE"
            raise PortfolioPanelError(code, "job is not available" if code == "PORTFOLIO_JOB_NOT_FOUND" else "portfolio job runtime is unavailable", status=404 if code == "PORTFOLIO_JOB_NOT_FOUND" else 500) from error
        except (OSError, TypeError, ValueError) as error:
            raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500) from error
        if saved.get("kind") not in {"portfolio.stage1", "portfolio.stage2"}:
            raise PortfolioPanelError("PORTFOLIO_JOB_NOT_FOUND", "job is not available", status=404)
        if saved["state"] in _TERMINAL:
            raise PortfolioPanelError("PORTFOLIO_JOB_TERMINAL", "job is terminal", status=409)
        event = self._cancel_events.get(job_id)
        if event is not None:
            event.set()
        if saved["state"] == "QUEUED":
            try:
                saved = self.registry.cancel(job_id)
            except PanelJobError as error:
                if error.code == "NOT_FOUND":
                    raise PortfolioPanelError("PORTFOLIO_JOB_NOT_FOUND", "job is not available", status=404) from error
                raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500) from error
            except (OSError, TypeError, ValueError) as error:
                raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500) from error
        elif saved["state"] == "RUNNING":
            try:
                saved = self.registry.cancel(job_id)
            except PanelJobError as error:
                if error.code == "NOT_FOUND":
                    raise PortfolioPanelError("PORTFOLIO_JOB_NOT_FOUND", "job is not available", status=404) from error
                raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500) from error
            except (OSError, TypeError, ValueError) as error:
                raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500) from error
        else:
            try:
                saved = self.registry.get(job_id)
            except PanelJobError as error:
                code = "PORTFOLIO_JOB_NOT_FOUND" if error.code == "NOT_FOUND" else "PORTFOLIO_JOB_RUNTIME_UNAVAILABLE"
                raise PortfolioPanelError(code, "job is not available" if code == "PORTFOLIO_JOB_NOT_FOUND" else "portfolio job runtime is unavailable", status=404 if code == "PORTFOLIO_JOB_NOT_FOUND" else 500) from error
            except (OSError, TypeError, ValueError) as error:
                raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500) from error
        if saved.get("state") in {"CANCELLED", "FAILED"}:
            self._cleanup_terminal_snapshot(job_id)
        return {"job_id": job_id, "status": self._project_state(saved)}

    def active_or_job(self) -> dict[str, Any] | None:
        active = self.active_job()
        if active is not None:
            return active
        try:
            jobs = self.registry.list()
        except (PanelJobError, KeyError, OSError, TypeError, ValueError) as error:
            raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500) from error
        portfolio_jobs = [saved for saved in jobs if saved.get("kind") in {"portfolio.stage1", "portfolio.stage2"} and saved.get("state") in _TERMINAL]
        if portfolio_jobs:
            latest = max(portfolio_jobs, key=lambda saved: str(saved.get("created_at_utc") or ""))
            return self.job(latest["job_id"])
        return None

    def results(self, campaign_id: str) -> dict[str, Any]:
        saved, runtime = self._campaign_by_id(campaign_id)
        if self._project_state(saved) != "SUCCEEDED":
            raise PortfolioPanelError("PORTFOLIO_JOB_RESULTS_UNAVAILABLE", "results are not available", status=409)
        try:
            campaign = self._hydrate_campaign(saved, runtime, allow_terminal=True, expected_campaign_id=campaign_id)
        except PortfolioPanelError as error:
            if error.code.startswith("CAMPAIGN_"):
                raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500) from error
            raise
        if not isinstance(campaign.get("input_digest"), str) or not isinstance(campaign.get("config_digest"), str):
            raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500)
        summary = runtime.get("summary") if isinstance(runtime.get("summary"), Mapping) else {}
        try:
            workbook = Path(runtime.get("workbook_path", ""))
            expected = self._results_root / campaign_id / "stage1.xlsx"
            if workbook != expected or workbook.is_symlink() or workbook.parent.is_symlink() or self._results_root.is_symlink() or workbook.resolve(strict=True) != expected:
                raise OSError
        except (OSError, TypeError, ValueError, FileNotFoundError):
            raise PortfolioPanelError("PORTFOLIO_JOB_RESULTS_UNAVAILABLE", "results are not available", status=409) from None
        if not isinstance(summary, dict):
            summary = dict(summary)
        if summary.get("weighted_result_schema") != 1:
            artifact_digest = runtime.get("executables_digest")
            if isinstance(artifact_digest, str) and re.fullmatch(r"[0-9a-f]{64}", artifact_digest):
                try:
                    artifact = self._load_stage1_executables(
                        campaign_id,
                        campaign["input_digest"],
                        campaign["config_digest"],
                        artifact_digest,
                    )
                    candidates = artifact.get("candidates") if isinstance(artifact, Mapping) else None
                    candidate = candidates[0] if isinstance(candidates, list) and candidates else None
                    if isinstance(candidate, Mapping):
                        enriched = self._summary(campaign, (), (), (candidate,), ())
                        for key in (
                            "weighted_result_schema",
                            "Weighted candidate ID", "B required USDT", "Required bank USDT", "Historical bank USDT",
                            "Stress bank P95 USDT", "Margin-only bank USDT", "B saturation USDT", "B margin USDT",
                            "P30 common USDT/30d", "P30 limiter USDT/30d", "MaxDD %", "CDaR peak80 USDT",
                            "CDaR peak90 USDT", "IM all USDT", "MM all USDT", "Total full notional USDT", "Weighted members",
                            "CDaR peak80 %", "CDaR peak90 %",
                            "Target bank USDT", "Profile DD limit %", "MaxDD SUM USDT",
                            "CDaR peak80 target %", "CDaR peak90 target %", "CDaR peak80 saturation %", "CDaR peak90 saturation %",
                            "IM target %", "IM saturation %", "MM target %", "MM saturation %",
                        ):
                            if key in enriched:
                                summary[key] = _plain(enriched[key])
                except PortfolioPanelError:
                    pass
        return {"campaign_id": campaign_id, "input_digest": campaign["input_digest"], "config_digest": campaign["config_digest"], "summary": summary, "blockers": summary.get("blockers", []), "workbook_available": workbook.is_file()}

    def workbook(self, campaign_id: str) -> bytes:
        saved, runtime = self._campaign_by_id(campaign_id)
        if self._project_state(saved) != "SUCCEEDED":
            raise PortfolioPanelError("PORTFOLIO_JOB_WORKBOOK_UNAVAILABLE", "workbook is not available", status=409)
        raw = runtime.get("workbook_path")
        if not isinstance(raw, str):
            raise PortfolioPanelError("PORTFOLIO_JOB_WORKBOOK_UNAVAILABLE", "workbook is not available", status=409)
        path = Path(raw)
        try:
            expected = self._results_root / campaign_id / "stage1.xlsx"
            if path != expected or not path.is_absolute() or path.is_symlink() or path.parent.is_symlink() or self._results_root.is_symlink():
                raise OSError
            resolved = path.resolve(strict=True)
            if resolved != expected or not resolved.is_file():
                raise OSError
            return path.read_bytes()
        except (OSError, ValueError, FileNotFoundError):
            raise PortfolioPanelError("PORTFOLIO_JOB_WORKBOOK_UNAVAILABLE", "workbook is not available", status=409) from None

    def submit_tester_submission(self, campaign_id: str, payload: Mapping[str, Any]) -> dict[str, Any]:
        if self._local_testing_service_provider is None:
            raise PortfolioPanelError("PORTFOLIO_JOB_STAGE2_NOT_AUTHORIZED", "stage 2 is not authorized", status=409)
        if (
            not isinstance(payload, Mapping)
            or set(payload) != {"confirmed", "campaign_id"}
            or payload.get("confirmed") is not True
            or payload.get("campaign_id") != campaign_id
        ):
            raise PortfolioPanelError("PORTFOLIO_CAMPAIGN_INVALID", "campaign confirmation is invalid", status=422)
        with self._lock, self.registry.lock:
            for saved in self.registry.list():
                if saved.get("kind") != "portfolio.stage2":
                    continue
                try:
                    runtime = self.registry.runtime(saved["job_id"])
                except (PanelJobError, KeyError, OSError, TypeError, ValueError) as error:
                    raise PortfolioPanelError("PORTFOLIO_JOB_RUNTIME_UNAVAILABLE", "portfolio job runtime is unavailable", status=500) from error
                binding = runtime.get("stage2") if isinstance(runtime.get("stage2"), Mapping) else {}
                if binding.get("campaign_id") == campaign_id:
                    return {"campaign_id": campaign_id, "job_id": saved["job_id"], "status": self._project_state(saved)}
            saved_stage1, _runtime = self._campaign_by_id(campaign_id)
            if self._project_state(saved_stage1) != "SUCCEEDED":
                raise PortfolioPanelError("PORTFOLIO_STAGE2_CAMPAIGN_NOT_READY", "stage 1 campaign is not ready", status=409)
            prepared = self._prepare_stage2_baseline(campaign_id)
            binding = {
                "campaign_id": prepared["campaign_id"],
                "input_digest": prepared["input_digest"],
                "config_digest": prepared["config_digest"],
                "artifact_digest": prepared["artifact_digest"],
                "candidate_id": prepared["candidate_id"],
                "candidate_digest": prepared["candidate_digest"],
                "expected_names": list(prepared["expected_names"]),
                "pretest_period": _plain(prepared["pretest_period"]),
                "receipt": _plain(prepared["receipt"]),
            }
            submission_key = f"portfolio-stage2:{campaign_id}"
            created_job_id: str | None = None
            try:
                saved = self.registry.submit("portfolio.stage2", {"campaign_id": campaign_id, "candidate_id": prepared["candidate_id"]}, submission_key, ("portfolio_optimizer",))
                created_job_id = saved["job_id"]
                self.registry.reserve_runtime(saved["job_id"], "stage2", binding)
            except PanelJobError as error:
                if created_job_id is not None:
                    try:
                        self.registry.discard_queued(created_job_id)
                    except Exception:
                        pass
                code = "PORTFOLIO_JOB_BUSY" if error.code in {"RESOURCE_BUSY", "JOB_CAPACITY_EXHAUSTED"} else error.code
                raise PortfolioPanelError(code, "portfolio optimizer is busy" if code == "PORTFOLIO_JOB_BUSY" else code, status=409 if code == "PORTFOLIO_JOB_BUSY" else 400) from error
            except Exception as error:
                if created_job_id is not None:
                    try:
                        self.registry.discard_queued(created_job_id)
                    except Exception:
                        pass
                raise PortfolioPanelError("PORTFOLIO_JOB_START_FAILED", "portfolio job could not start", status=503) from error
            event = threading.Event()
            self._cancel_events[saved["job_id"]] = event
            try:
                worker = threading.Thread(target=self._run_stage2, args=(saved["job_id"], prepared), name="mrs3-portfolio-stage2", daemon=True)
                self._threads[saved["job_id"]] = worker
                worker.start()
            except BaseException as error:
                self._threads.pop(saved["job_id"], None)
                self._cancel_events.pop(saved["job_id"], None)
                self.registry.discard_queued(saved["job_id"])
                raise PortfolioPanelError("PORTFOLIO_JOB_START_FAILED", "portfolio job could not start", status=503) from error
        return {"campaign_id": campaign_id, "job_id": saved["job_id"], "status": "QUEUED"}


__all__ = ["PortfolioPanelError", "PortfolioPanelService", "STAGES"]
