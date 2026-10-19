"""The listing workbook generated from the pair registry (registry stage 4).

The PerformanceDB import, the pipeline, the finalist retest and plateau research read listing days through
`loader.load_listing_dates` from a workbook; its path and source are recorded in PerformanceDB, so that contract
stays. The registry writes the workbook: sheet «Пары» (Пара, Дата листинга на Bybit (UTC)) and «Пары_BingX» (Пара,
Листинг), one row per pair with a listing day, the UTC day at midnight as in the hand-kept workbook. Contract:
docs/specs/2026-10-18-pair-registry.md, «Stage 4 contract».
"""

from __future__ import annotations

from datetime import datetime, timezone
import os
from pathlib import Path
import tempfile

import duckdb

WORKBOOK_NAME = "pairs.xlsx"
SHEETS = (("bybit", "Пары", "Дата листинга на Bybit (UTC)"), ("bingx", "Пары_BingX", "Листинг"))


def write_listing_workbook(con: duckdb.DuckDBPyConnection, out: Path) -> dict[str, int]:
    """Write the workbook atomically (a temporary file in the same folder, then a replace); rows per exchange."""
    from openpyxl import Workbook

    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    counts: dict[str, int] = {}
    for exchange, title, column in SHEETS:
        sheet = book.create_sheet(title)
        sheet.append(["Пара", column])
        rows = con.execute("select symbol, listing_ms from pairs where exchange = ? and listing_ms is not null "
                           "order by symbol", [exchange]).fetchall()
        for symbol, listing_ms in rows:
            day = datetime.fromtimestamp(listing_ms / 1000, timezone.utc).date()
            sheet.append([symbol, datetime(day.year, day.month, day.day)])
            sheet.cell(sheet.max_row, 2).number_format = "yyyy-mm-dd"
        sheet.column_dimensions["A"].width = 18
        sheet.column_dimensions["B"].width = 30
        counts[exchange] = len(rows)
    handle, temporary = tempfile.mkstemp(prefix=f".{out.name}.", suffix=".tmp", dir=out.parent)
    os.close(handle)
    try:
        book.save(temporary)
        os.replace(temporary, out)
    except BaseException:
        Path(temporary).unlink(missing_ok=True)
        raise
    return counts
