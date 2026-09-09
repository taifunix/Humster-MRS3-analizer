"""Server-owned provenance and bulk control for Performance v2 finalist retests.

The module deliberately keeps cohort lineage in ordinary job/selection metadata.  It
does not add a second database schema or a second status vocabulary.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from typing import Mapping, Sequence

import duckdb
from openpyxl import Workbook, load_workbook

from .config import AlgorithmConfig
from .lots import LotMethod
from .performance_v2_store import PerformanceV2StoreError, require_performance_v2
from .strategy_json import generate_strategy


LISTING_WARMUP_HOURS = 120
MAX_COHORT_MEMBERS = 10_000


_RUNTIME_KEYS = {
    "created_at", "updated_at", "timestamp", "exported_at", "imported_at",
    "started_at", "finished_at", "completed_at",
    "at_utc",
}


def _runtime_key(key: object) -> bool:
    name = str(key).casefold()
    return (
        name in _RUNTIME_KEYS
        or name.endswith("_at_utc")
        or name.endswith("_timestamp")
        or name.endswith("_path")
        or name in {"path", "local_path", "runtime_path"}
    )


def _json_value(value: object) -> object:
    if isinstance(value, Decimal):
        if not value.is_finite():
            raise ValueError("canonical JSON does not support non-finite Decimal values")
        return str(value)
    if isinstance(value, datetime):
        current = value if value.tzinfo is not None else value.replace(tzinfo=timezone.utc)
        return current.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Path):
        return None
    if isinstance(value, Mapping):
        return {
            str(key): _json_value(child)
            for key, child in sorted(value.items(), key=lambda item: str(item[0]))
            if not _runtime_key(key)
        }
    if isinstance(value, (list, tuple)):
        return [_json_value(child) for child in value]
    if isinstance(value, (set, frozenset)):
        return sorted((_json_value(child) for child in value), key=lambda child: json.dumps(child, ensure_ascii=False, sort_keys=True, separators=(",", ":")))
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    raise TypeError(f"unsupported canonical JSON value: {type(value).__name__}")


def canonical_json(value: object) -> str:
    """Return stable UTF-8 JSON used by cohort, manifest and config digests."""
    return json.dumps(_json_value(value), ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def canonical_digest(value: object) -> str:
    return sha256(canonical_json(value).encode("utf-8")).hexdigest()


def cohort_digest(members: Sequence[Mapping[str, object]]) -> str:
    """Hash sorted members while ignoring runtime timestamps and local paths."""
    ordered = sorted(
        (dict(member) for member in members),
        key=lambda member: (int(member.get("strategy_id", 0)), canonical_json(member)),
    )
    return canonical_digest({"members": ordered})


def review_key(workbook_sha256: str, selection_run_id: str) -> str:
    """Derive an idempotent existing-table review key from actual workbook bytes."""
    if not isinstance(workbook_sha256, str) or len(workbook_sha256) != 64:
        raise ValueError("workbook_sha256 must be a SHA-256 hash")
    try:
        int(workbook_sha256, 16)
    except ValueError:
        raise ValueError("workbook_sha256 must be a SHA-256 hash") from None
    if not isinstance(selection_run_id, str) or not selection_run_id.strip():
        raise ValueError("selection_run_id must be non-empty")
    return sha256((workbook_sha256.lower() + "\0" + selection_run_id).encode("utf-8")).hexdigest()


# Names used by callers that describe the same contract more explicitly.
canonical_provenance_json = canonical_json
canonical_provenance_digest = canonical_digest
deterministic_review_key = review_key


class FinalistRetestError(PerformanceV2StoreError):
    """A typed failure in the server-owned bulk finalist flow."""

    def __init__(self, code: str, message: str | None = None, *, details: object = None) -> None:
        self.code = code
        self.details = details
        super().__init__(message or code)


@dataclass(frozen=True, slots=True)
class RetestExclusion:
    strategy_id: int | None
    strategy_name: str | None
    symbol: str | None
    reason: str
    side: str | None = None

    def as_dict(self) -> dict[str, object]:
        return {
            "strategy_id": self.strategy_id,
            "strategy_name": self.strategy_name,
            "symbol": self.symbol,
            "side": self.side,
            "reason": self.reason,
        }


@dataclass(frozen=True, slots=True)
class FinalistRetestCohort:
    scope: str
    requested_start: str
    requested_end: str
    warmup_hours: int
    members: tuple[Mapping[str, object], ...]
    exclusions: tuple[RetestExclusion, ...]
    cohort_sha256: str

    @property
    def member_count(self) -> int:
        return len(self.members)

    @property
    def excluded_count(self) -> int:
        return len(self.exclusions)

    def as_manifest(self) -> dict[str, object]:
        return {
            "scope": self.scope,
            "requested_start": self.requested_start,
            "requested_end": self.requested_end,
            "warmup_hours": self.warmup_hours,
            "cohort_sha256": self.cohort_sha256,
            "members": [_json_value(dict(member)) for member in self.members],
            "exclusions": [item.as_dict() for item in self.exclusions],
        }


@dataclass(frozen=True, slots=True)
class FinalistRetestBatch:
    job_id: str
    strategies_path: Path
    manifest_path: Path
    cohort: FinalistRetestCohort
    config_sha256: str
    strategy_count: int

    @property
    def run_id(self) -> str:
        return self.job_id

    @property
    def output_dir(self) -> Path:
        return self.manifest_path.parent


@dataclass(frozen=True, slots=True)
class FinalistRetestImportResult:
    job_id: str
    status: str
    successful_replacements: tuple[Mapping[str, int], ...] = ()
    failures: tuple[Mapping[str, object], ...] = ()

    @property
    def success_count(self) -> int:
        return len(self.successful_replacements)

    @property
    def failure_count(self) -> int:
        return len(self.failures)

    def terminal_metadata(
        self,
        *,
        scope: str,
        cohort_sha256: str,
        manifest_sha256: str,
        config_sha256: str,
    ) -> dict[str, object]:
        return {
            "job_id": self.job_id,
            "scope": scope,
            "cohort_sha256": cohort_sha256,
            "manifest_sha256": manifest_sha256,
            "config_sha256": config_sha256,
            "successful_replacements": [dict(item) for item in self.successful_replacements],
            "failures": [dict(item) for item in self.failures],
            "status": self.status,
        }


def _outcome_for(outcomes: Mapping[object, object], member: Mapping[str, object]) -> Mapping[str, object] | None:
    value = outcomes.get(member["strategy_id"])
    if value is None:
        value = outcomes.get(str(member["strategy_id"]))
    if value is None:
        value = outcomes.get(member["strategy_name"])
    return value if isinstance(value, Mapping) else None


def apply_finalist_retest_outcomes(
    connection: duckdb.DuckDBPyConnection,
    cohort: FinalistRetestCohort,
    outcomes: Mapping[object, object],
    *,
    job_id: str,
) -> FinalistRetestImportResult:
    """Apply already-imported result IDs independently for each frozen member.

    The native HTML importer remains responsible for parsing and writing result,
    action, equity and metrics facts.  This narrow coordinator only advances a
    strategy pointer after checking the frozen result ID, which makes fake/native
    importer integration equally easy to verify.
    """
    require_performance_v2(connection)
    if not isinstance(outcomes, Mapping) or not isinstance(job_id, str) or not job_id.strip():
        raise FinalistRetestError("INVALID_REQUEST", "bulk import outcome request is invalid")
    successful: list[Mapping[str, int]] = []
    failures: list[Mapping[str, object]] = []
    for member in cohort.members:
        strategy_id = int(member["strategy_id"])
        name = str(member["strategy_name"])
        outcome = _outcome_for(outcomes, member)
        if outcome is None:
            failures.append({"strategy_id": strategy_id, "strategy_name": name, "reason": "UNVISITED"})
            continue
        outcome_status = str(outcome.get("status", "FAILED")).upper()
        if outcome_status not in {"SUCCESS", "IMPORTED", "REPLACED", "COMMITTED"}:
            failures.append({
                "strategy_id": strategy_id, "strategy_name": name,
                "reason": str(outcome.get("reason") or outcome.get("error_code") or "IMPORT_FAILED"),
            })
            continue
        raw_new = outcome.get("new_result_id", outcome.get("result_id"))
        if isinstance(raw_new, bool) or not isinstance(raw_new, int) or raw_new <= 0:
            failures.append({"strategy_id": strategy_id, "strategy_name": name, "reason": "MISSING_NEW_RESULT"})
            continue
        try:
            connection.execute("begin transaction")
            current_row = connection.execute(
                "select current_result_id from strategies where strategy_id = ? and lifecycle_status = 'ACTIVE'",
                [strategy_id],
            ).fetchone()
            if current_row is None or current_row[0] != member["result_id"]:
                connection.execute("rollback")
                failures.append({"strategy_id": strategy_id, "strategy_name": name, "reason": "STALE_RESULT"})
                continue
            result_row = connection.execute(
                "select strategy_id from strategy_results where result_id = ?",
                [raw_new],
            ).fetchone()
            if result_row is None or int(result_row[0]) != strategy_id:
                connection.execute("rollback")
                failures.append({"strategy_id": strategy_id, "strategy_name": name, "reason": "RESULT_ID_MISMATCH"})
                continue
            connection.execute(
                "update strategies set current_result_id = ?, updated_at_utc = now() where strategy_id = ?",
                [raw_new, strategy_id],
            )
            connection.execute("commit")
        except Exception:
            try:
                connection.execute("rollback")
            except Exception:
                pass
            failures.append({"strategy_id": strategy_id, "strategy_name": name, "reason": "REPLACE_FAILED"})
            continue
        successful.append({"strategy_id": strategy_id, "old_result_id": int(member["result_id"]), "new_result_id": raw_new})
    status = "COMMITTED" if successful else "FAILED"
    return FinalistRetestImportResult(job_id, status, tuple(successful), tuple(failures))


execute_finalist_retest_import = apply_finalist_retest_outcomes
import_finalist_retest = apply_finalist_retest_outcomes


CONTROL_WORKBOOK_SHEETS = ("Candidates", "Groups", "Retest Failures", "_MRS_SELECTION_META")
CONTROL_CANDIDATE_HEADERS = (
    "Pair", "Direction", "Strategy ID", "Result ID", "User Status", "User Rank", "RETEST", "Comment",
    "Auto Status", "Auto Rank", "Auto Analog Of ID", "Analog Of ID", "Auto Reason", "Effective Start", "Effective End", "Score",
)
CONTROL_GROUP_HEADERS = (
    "Pair", "Direction", "Frozen Count", "Success Count", "Failure Count", "Auto Status Count",
)
CONTROL_FAILURE_HEADERS = ("Pair", "Direction", "Strategy ID", "Strategy", "Result ID", "Reason")


def _control_row_value(row: Mapping[str, object], *keys: str) -> object:
    for key in keys:
        if key in row:
            value = row[key]
            if isinstance(value, (datetime, date)):
                return _json_value(value)
            return value
    return None


def _control_records(rows: Sequence[Mapping[str, object]] | None) -> list[Mapping[str, object]]:
    return [row for row in (rows or ()) if isinstance(row, Mapping)]


def _control_metadata_value(value: object) -> object:
    """Store structured provenance as one deterministic XLSX string cell."""
    value = _json_value(value)
    if isinstance(value, (Mapping, list, tuple)):
        return canonical_json(value)
    return value


def _metadata_json(metadata: Mapping[str, object], key: str, default: object = None) -> object:
    value = metadata.get(key, default)
    if not isinstance(value, str):
        return value
    try:
        return json.loads(value)
    except (TypeError, ValueError):
        return value


def write_combined_control_workbook(
    candidates: Sequence[Mapping[str, object]],
    groups: Sequence[Mapping[str, object]] = (),
    failures: Sequence[Mapping[str, object]] = (),
    metadata: Mapping[str, object] | None = None,
    path: Path | None = None,
) -> Path:
    """Write the fixed, reviewable bulk control workbook atomically."""
    candidate_rows = _control_records(candidates)
    group_rows = _control_records(groups)
    failure_rows = _control_records(failures)
    workbook = Workbook()
    first = workbook.active
    first.title = "Candidates"
    sheets = {name: (first if name == "Candidates" else workbook.create_sheet(name)) for name in CONTROL_WORKBOOK_SHEETS}
    for name, headers in (
        ("Candidates", CONTROL_CANDIDATE_HEADERS),
        ("Groups", CONTROL_GROUP_HEADERS),
        ("Retest Failures", CONTROL_FAILURE_HEADERS),
    ):
        sheet = workbook[name]
        sheet.append(list(headers))
    candidate_rows.sort(key=lambda row: (
        str(_control_row_value(row, "Pair", "symbol") or ""),
        str(_control_row_value(row, "Direction", "side") or ""),
        int(_control_row_value(row, "Strategy ID", "strategy_id") or 0),
    ))
    seen: set[tuple[object, object]] = set()
    for row in candidate_rows:
        strategy_id = _control_row_value(row, "Strategy ID", "strategy_id")
        result_id = _control_row_value(row, "Result ID", "result_id")
        key = (strategy_id, result_id)
        if key in seen:
            raise FinalistRetestError("CONTROL_ROWSET_DUPLICATE", "combined workbook contains duplicate candidates")
        seen.add(key)
        values = [
            _control_row_value(row, "Pair", "symbol"), _control_row_value(row, "Direction", "side"),
            strategy_id, result_id,
            _control_row_value(row, "User Status", "user_status", "effective_status"),
            _control_row_value(row, "User Rank", "user_rank"), _control_row_value(row, "RETEST", "retest"),
            _control_row_value(row, "Comment", "comment"),
            _control_row_value(row, "Auto Status", "auto_status"), _control_row_value(row, "Auto Rank", "auto_rank"),
            _control_row_value(row, "Auto Analog Of ID", "auto_analog_of_strategy_id"), _control_row_value(row, "Analog Of ID", "analog_of_strategy_id"),
            _control_row_value(row, "Auto Reason", "auto_reason", "elimination_reason"),
            _control_row_value(row, "Effective Start", "effective_start"), _control_row_value(row, "Effective End", "effective_end"),
            _control_row_value(row, "Score", "score", "final_score"),
        ]
        sheets["Candidates"].append(values)
    for row in sorted(group_rows, key=lambda item: (str(_control_row_value(item, "Pair", "symbol") or ""), str(_control_row_value(item, "Direction", "side") or ""))):
        sheets["Groups"].append([
            _control_row_value(row, "Pair", "symbol"), _control_row_value(row, "Direction", "side"),
            _control_row_value(row, "Frozen Count", "frozen_count", "count"),
            _control_row_value(row, "Success Count", "success_count"), _control_row_value(row, "Failure Count", "failure_count"),
            _control_row_value(row, "Auto Status Count", "auto_status_count"),
        ])
    for row in failure_rows:
        sheets["Retest Failures"].append([
            _control_row_value(row, "Pair", "symbol"), _control_row_value(row, "Direction", "side"),
            _control_row_value(row, "Strategy ID", "strategy_id"), _control_row_value(row, "Strategy", "strategy_name"),
            _control_row_value(row, "Result ID", "result_id"), _control_row_value(row, "Reason", "reason"),
        ])
    meta = sheets["_MRS_SELECTION_META"]
    for key, value in sorted((metadata or {}).items(), key=lambda item: str(item[0])):
        meta.append([str(key), _control_metadata_value(value)])
    meta.sheet_state = "veryHidden"
    for sheet in (sheets["Candidates"], sheets["Groups"], sheets["Retest Failures"]):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    destination = Path(path) if path is not None else Path.cwd() / ".combined-control.xlsx"
    temporary = destination.with_name(f".{destination.name}.tmp")
    destination.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(temporary)
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)
    return destination


def combined_control_workbook_bytes(
    candidates: Sequence[Mapping[str, object]],
    groups: Sequence[Mapping[str, object]] = (),
    failures: Sequence[Mapping[str, object]] = (),
    metadata: Mapping[str, object] | None = None,
) -> bytes:
    """Build control bytes before any database write."""
    import tempfile

    with tempfile.TemporaryDirectory() as directory:
        path = write_combined_control_workbook(candidates, groups, failures, metadata, Path(directory) / "control.xlsx")
        return path.read_bytes()


def validate_combined_control_workbook(data: bytes) -> tuple[dict[str, object], list[dict[str, object]], list[dict[str, object]], list[dict[str, object]]]:
    """Validate fixed sheets, editable vocabulary and per-group rank uniqueness."""
    if not isinstance(data, bytes) or not data:
        raise FinalistRetestError("CONTROL_INVALID_FILE")
    try:
        from .performance_v2_selection_review import _bounded_xlsx

        _bounded_xlsx(data)
    except Exception as error:
        raise FinalistRetestError("CONTROL_INVALID_FILE") from error
    try:
        workbook = load_workbook(BytesIO(data), data_only=False, read_only=False)
    except Exception as error:
        raise FinalistRetestError("CONTROL_INVALID_FILE") from error
    try:
        if tuple(workbook.sheetnames) != CONTROL_WORKBOOK_SHEETS or workbook["_MRS_SELECTION_META"].sheet_state != "veryHidden":
            raise FinalistRetestError("CONTROL_SCHEMA_MISMATCH")
        if any(cell.data_type == "f" for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row):
            raise FinalistRetestError("CONTROL_FORMULA_FORBIDDEN")
        expected = {
            "Candidates": CONTROL_CANDIDATE_HEADERS, "Groups": CONTROL_GROUP_HEADERS, "Retest Failures": CONTROL_FAILURE_HEADERS,
        }
        parsed: dict[str, list[dict[str, object]]] = {}
        for name, headers in expected.items():
            sheet = workbook[name]
            actual = tuple(cell.value for cell in sheet[1])
            if actual != headers:
                raise FinalistRetestError("CONTROL_SCHEMA_MISMATCH")
            rows: list[dict[str, object]] = []
            for values in sheet.iter_rows(min_row=2, values_only=True):
                if all(value is None for value in values):
                    continue
                if any(isinstance(value, str) and value.startswith("=") for value in values):
                    raise FinalistRetestError("CONTROL_FORMULA_FORBIDDEN")
                rows.append(dict(zip(headers, values, strict=True)))
            parsed[name] = rows
        metadata = {str(row[0]): row[1] for row in workbook["_MRS_SELECTION_META"].iter_rows(values_only=True) if row[0] is not None}
        if len(metadata) != sum(1 for row in workbook["_MRS_SELECTION_META"].iter_rows(values_only=True) if row[0] is not None):
            raise FinalistRetestError("CONTROL_METADATA_INVALID")
        seen: set[tuple[object, object]] = set()
        ranks: dict[tuple[object, object], set[int]] = {}
        analogs: dict[tuple[object, object], list[tuple[int, str, int | None]]] = {}
        allowed = {"FINALIST", "RESERVE", "ANALOG", "FILTERED", "REJECTED"}
        for row in parsed["Candidates"]:
            key = (row["Strategy ID"], row["Result ID"])
            if key in seen:
                raise FinalistRetestError("CONTROL_ROWSET_DUPLICATE")
            seen.add(key)
            status = str(row["User Status"] or "").strip().upper()
            if status not in allowed:
                raise FinalistRetestError("CONTROL_INVALID_STATUS")
            raw_rank = row["User Rank"]
            if raw_rank not in (None, ""):
                if isinstance(raw_rank, bool) or int(raw_rank) != raw_rank or int(raw_rank) <= 0 or status not in {"FINALIST", "RESERVE"}:
                    raise FinalistRetestError("CONTROL_INVALID_RANK")
                group = (row["Pair"], row["Direction"])
                rank = int(raw_rank)
                if rank in ranks.setdefault(group, set()):
                    raise FinalistRetestError("CONTROL_DUPLICATE_RANK")
                ranks[group].add(rank)
            if row["RETEST"] not in (None, "", "RETEST"):
                raise FinalistRetestError("CONTROL_INVALID_RETEST")
            if row["Comment"] is not None and len(str(row["Comment"])) > 1000:
                raise FinalistRetestError("CONTROL_COMMENT_TOO_LONG")
            raw_analog = row["Analog Of ID"]
            analog = None
            if raw_analog not in (None, ""):
                try:
                    analog = int(raw_analog)
                except (TypeError, ValueError, OverflowError):
                    raise FinalistRetestError("CONTROL_INVALID_ANALOG") from None
                if isinstance(raw_analog, bool) or analog != raw_analog or analog <= 0:
                    raise FinalistRetestError("CONTROL_INVALID_ANALOG")
            strategy_id = int(row["Strategy ID"])
            if (status == "ANALOG") != (analog is not None) or analog == strategy_id:
                raise FinalistRetestError("CONTROL_INVALID_ANALOG")
            analogs.setdefault((row["Pair"], row["Direction"]), []).append((strategy_id, status, analog))
        for entries in analogs.values():
            statuses = {strategy_id: status for strategy_id, status, _analog in entries}
            if any(analog is not None and statuses.get(analog) not in {"FINALIST", "RESERVE"} for _strategy_id, _status, analog in entries):
                raise FinalistRetestError("CONTROL_INVALID_ANALOG")
        # The server includes digests for the immutable sheets.  Checking them
        # at upload time catches row deletion and edits before any DB lookup or
        # transaction begins.  Older ordinary fixtures may omit these keys.
        def rows_digest(items: Sequence[Mapping[str, object]]) -> str:
            return canonical_digest([dict(sorted(item.items(), key=lambda pair: pair[0])) for item in items])
        if metadata.get("groups_sha256") is not None and str(metadata["groups_sha256"]) != rows_digest(parsed["Groups"]):
            raise FinalistRetestError("CONTROL_IMMUTABLE_FIELDS_CHANGED")
        if metadata.get("failures_sha256") is not None and str(metadata["failures_sha256"]) != rows_digest(parsed["Retest Failures"]):
            raise FinalistRetestError("CONTROL_IMMUTABLE_FIELDS_CHANGED")
        immutable = []
        for item in parsed["Candidates"]:
            immutable.append({key: item.get(key) for key in (
                "Pair", "Direction", "Strategy ID", "Result ID", "Auto Status", "Auto Rank",
                "Auto Analog Of ID", "Auto Reason", "Effective Start", "Effective End", "Score",
            )})
        if metadata.get("immutable_content_sha256") is not None and str(metadata["immutable_content_sha256"]) != canonical_digest(immutable):
            raise FinalistRetestError("CONTROL_IMMUTABLE_FIELDS_CHANGED")
        return metadata, parsed["Candidates"], parsed["Groups"], parsed["Retest Failures"]
    except FinalistRetestError:
        raise
    except Exception as error:
        raise FinalistRetestError("CONTROL_INVALID_FILE") from error
    finally:
        workbook.close()


read_combined_control_workbook = validate_combined_control_workbook


def _legacy_import_combined_control_workbook(
    connection: duckdb.DuckDBPyConnection,
    data: bytes,
) -> dict[str, object]:
    """Validate and apply all group reviews in one DuckDB transaction.

    Existing selection review tables remain the ledger.  Each group gets a
    deterministic review id derived from the actual uploaded bytes and its
    selection run, while a failed group or row aborts the complete operation.
    """
    metadata, rows, _groups, _failures = validate_combined_control_workbook(data)
    require_performance_v2(connection)
    if metadata.get("ranking_scope") not in {None, "RETEST_COHORT"}:
        raise FinalistRetestError("CONTROL_SCOPE_MISMATCH")
    expected_instance = metadata.get("database_instance_id")
    instance = connection.execute("select value from schema_info where key = 'database_instance_id'").fetchone()
    if expected_instance is not None and str(expected_instance) != str(instance[0] if instance else ""):
        raise FinalistRetestError("CONTROL_DATABASE_MISMATCH")
    groups: dict[tuple[str, str], list[dict[str, object]]] = {}
    for row in rows:
        groups.setdefault((str(row["Pair"]), str(row["Direction"])), []).append(row)
    if not groups:
        raise FinalistRetestError("CONTROL_ROWSET_MISMATCH")
    actual_hash = sha256(data).hexdigest()
    prepared: list[tuple[str, str, list[dict[str, object]], str]] = []
    for (symbol, side), group_rows in sorted(groups.items()):
        run_row = connection.execute(
            "select selection_run_id from selection_runs where symbol = ? and side = ? order by created_at_utc desc, selection_run_id desc limit 1",
            [symbol, side],
        ).fetchone()
        if not run_row:
            raise FinalistRetestError("CONTROL_SELECTION_RUN_MISSING")
        run_id = str(run_row[0])
        snapshot_rows = connection.execute(
            "select strategy_id, result_id_at_selection, auto_status, auto_rank, auto_analog_of_strategy_id from selection_results where selection_run_id = ?",
            [run_id],
        ).fetchall()
        snapshot = {int(row[0]): row[1:] for row in snapshot_rows}
        submitted_ids = {int(row["Strategy ID"]) for row in group_rows}
        if submitted_ids != set(snapshot):
            raise FinalistRetestError("CONTROL_ROWSET_MISMATCH")
        names = {int(strategy_id): str(name) for strategy_id, name in connection.execute(
            "select strategy_id, strategy_name from strategies where strategy_id in (select unnest(?::bigint[]))", [list(snapshot)]
        ).fetchall()}
        current = {int(strategy_id): result_id for strategy_id, result_id in connection.execute(
            "select strategy_id, current_result_id from strategies where strategy_id in (select unnest(?::bigint[]))", [list(snapshot)]
        ).fetchall()}
        ranks: set[int] = set()
        statuses: dict[int, str] = {}
        for row in group_rows:
            strategy_id = int(row["Strategy ID"])
            if int(row["Result ID"]) != int(snapshot[strategy_id][0]) or current.get(strategy_id) != snapshot[strategy_id][0]:
                raise FinalistRetestError("CONTROL_STALE_RESULTS")
            submitted_auto_rank = None if row["Auto Rank"] in (None, "") else int(row["Auto Rank"])
            submitted_auto_analog = None if row["Auto Analog Of ID"] in (None, "") else int(row["Auto Analog Of ID"])
            if (
                str(row["Auto Status"]) != str(snapshot[strategy_id][1])
                or submitted_auto_rank != snapshot[strategy_id][2]
                or submitted_auto_analog != snapshot[strategy_id][3]
            ):
                raise FinalistRetestError("CONTROL_IMMUTABLE_FIELDS_CHANGED")
            status = str(row["User Status"]).strip().upper()
            rank = row["User Rank"]
            if rank not in (None, ""):
                rank = int(rank)
                if rank in ranks:
                    raise FinalistRetestError("CONTROL_DUPLICATE_RANK")
                ranks.add(rank)
            statuses[strategy_id] = status
        for row in group_rows:
            analog = row["Analog Of ID"]
            if statuses[int(row["Strategy ID"])] == "ANALOG":
                if analog in (None, "") or int(analog) == int(row["Strategy ID"]) or int(analog) not in statuses or statuses[int(analog)] not in {"FINALIST", "RESERVE"}:
                    raise FinalistRetestError("CONTROL_INVALID_ANALOG")
            elif analog not in (None, ""):
                raise FinalistRetestError("CONTROL_INVALID_ANALOG")
        prepared.append((symbol, side, group_rows, run_id))
    now = datetime.now(timezone.utc)
    review_ids: list[str] = []
    connection.execute("begin transaction")
    try:
        for symbol, side, group_rows, run_id in prepared:
            review_id = review_key(actual_hash, run_id)
            if connection.execute("select 1 from selection_review_imports where workbook_sha256 = ?", [review_id]).fetchone():
                review_ids.append(review_id)
                continue
            connection.execute(
                "insert into selection_review_imports (review_import_id, selection_run_id, workbook_sha256, imported_at_utc, row_count) values (?, ?, ?, ?, ?)",
                [review_id, run_id, review_id, now, len(group_rows)],
            )
            decision_rows = []
            retest_ids = []
            for row in group_rows:
                strategy_id = int(row["Strategy ID"])
                rank = None if row["User Rank"] in (None, "") else int(row["User Rank"])
                analog = None if row["Analog Of ID"] in (None, "") else int(row["Analog Of ID"])
                decision_rows.append([review_id, strategy_id, str(row["User Status"]).strip().upper(), rank, analog, "" if row["Comment"] is None else str(row["Comment"])])
                if row["RETEST"] == "RETEST":
                    retest_ids.append(strategy_id)
            connection.executemany(
                "insert into selection_review_rows (review_import_id, strategy_id, user_status, user_rank, user_analog_of_strategy_id, comment) values (?, ?, ?, ?, ?, ?)",
                decision_rows,
            )
            ids = [int(row["Strategy ID"]) for row in group_rows]
            connection.execute("delete from strategy_tags where tag = 'REJECTED' and strategy_id in (select unnest(?::bigint[]))", [ids])
            rejected = [[strategy_id, "REJECTED", "SELECTION_REVIEW", review_id, now] for _review, strategy_id, status, *_ in decision_rows if status == "REJECTED"]
            if rejected:
                connection.executemany("insert into strategy_tags (strategy_id, tag, source, source_ref, updated_at_utc) values (?, ?, ?, ?, ?)", rejected)
            connection.execute("delete from strategy_tags where tag = 'RETEST' and source = 'SELECTION_REVIEW' and strategy_id in (select unnest(?::bigint[]))", [ids])
            retest = [[strategy_id, "RETEST", "SELECTION_REVIEW", review_id, now] for strategy_id in retest_ids]
            if retest:
                connection.executemany("insert into strategy_tags (strategy_id, tag, source, source_ref, updated_at_utc) values (?, ?, ?, ?, ?) on conflict (strategy_id, tag) do update set source = excluded.source, source_ref = excluded.source_ref, updated_at_utc = excluded.updated_at_utc", retest)
            review_ids.append(review_id)
        connection.execute("commit")
    except Exception:
        try:
            connection.execute("rollback")
        except Exception:
            pass
        raise
    return {"review_import_ids": review_ids, "group_count": len(prepared), "row_count": len(rows), "workbook_sha256": actual_hash}


def import_combined_control_workbook(
    connection: duckdb.DuckDBPyConnection,
    data: bytes,
) -> dict[str, object]:
    """Apply a server-issued combined workbook after exact provenance checks.

    Workbooks produced by the finalist retest exporter carry exact selection
    run IDs and immutable rowset digests.  The old hand-authored workbook shape
    remains accepted for the ordinary compatibility tests and is delegated to
    the historical implementation above.
    """
    metadata, rows, groups, failures = validate_combined_control_workbook(data)
    if "group_run_ids_json" not in metadata:
        return _legacy_import_combined_control_workbook(connection, data)
    require_performance_v2(connection)
    if metadata.get("ranking_scope") != "RETEST_COHORT":
        raise FinalistRetestError("CONTROL_SCOPE_MISMATCH")
    expected_instance = metadata.get("database_instance_id")
    instance = connection.execute("select value from schema_info where key = 'database_instance_id'").fetchone()
    if str(expected_instance or "") != str(instance[0] if instance else ""):
        raise FinalistRetestError("CONTROL_DATABASE_MISMATCH")
    actual_hash = sha256(data).hexdigest()
    try:
        run_ids = _metadata_json(metadata, "group_run_ids_json", {})
        run_hashes = _metadata_json(metadata, "group_workbook_sha256_json", {})
        issued_rowsets = _metadata_json(metadata, "exact_rowsets_json", {})
        if not isinstance(run_ids, Mapping) or not run_ids:
            raise FinalistRetestError("CONTROL_METADATA_INVALID")
        if run_hashes is None:
            run_hashes = {}
        if not isinstance(run_hashes, Mapping):
            raise FinalistRetestError("CONTROL_METADATA_INVALID")
        if not isinstance(issued_rowsets, Mapping):
            raise FinalistRetestError("CONTROL_METADATA_INVALID")
    except FinalistRetestError:
        raise
    except Exception as error:
        raise FinalistRetestError("CONTROL_METADATA_INVALID") from error

    def group_key(symbol: object, side: object) -> str:
        return f"{symbol}|{side}"

    candidate_groups: dict[str, list[dict[str, object]]] = {}
    for row in rows:
        candidate_groups.setdefault(group_key(row["Pair"], row["Direction"]), []).append(row)
    success_keys = set(candidate_groups)
    run_keys = {str(key) for key in run_ids}
    if success_keys != run_keys:
        raise FinalistRetestError("CONTROL_ROWSET_MISMATCH")
    group_rows_by_key = {group_key(row["Pair"], row["Direction"]): row for row in groups}
    if len(group_rows_by_key) != len(groups) or not success_keys.issubset(set(group_rows_by_key)):
        raise FinalistRetestError("CONTROL_ROWSET_MISMATCH")
    failure_groups: dict[str, int] = {}
    for row in failures:
        failure_groups[group_key(row["Pair"], row["Direction"])] = failure_groups.get(group_key(row["Pair"], row["Direction"]), 0) + 1
    if not set(failure_groups).issubset(set(group_rows_by_key)):
        raise FinalistRetestError("CONTROL_ROWSET_MISMATCH")
    for key, summary in group_rows_by_key.items():
        success_count = len(candidate_groups.get(key, ()))
        failure_count = failure_groups.get(key, 0)
        frozen_count = int(summary.get("Frozen Count") or 0)
        if frozen_count != success_count + failure_count or int(summary.get("Success Count") or 0) != success_count:
            raise FinalistRetestError("CONTROL_ROWSET_MISMATCH")
        if int(summary.get("Failure Count") or 0) != failure_count:
            raise FinalistRetestError("CONTROL_ROWSET_MISMATCH")

    immutable = []
    exact_rowsets: dict[str, list[dict[str, object]]] = {}
    prepared: list[tuple[str, str, list[dict[str, object]], list[tuple[object, ...]], str]] = []
    shared_issued_hash: str | None = None
    expected_job = str(metadata.get("bulk_retest_job_id") or "")
    expected_cohort = str(metadata.get("cohort_sha256") or "")
    expected_manifest = str(metadata.get("manifest_sha256") or "")
    expected_config = str(metadata.get("config_sha256") or "")
    for key, group in sorted(candidate_groups.items()):
        symbol, side = key.split("|", 1)
        run_id = str(run_ids.get(key) or "")
        if not run_id:
            raise FinalistRetestError("CONTROL_METADATA_INVALID")
        run = connection.execute(
            """select selection_run_id, database_instance_id, symbol, side, request_json,
                      config_sha256, workbook_sha256
                 from selection_runs where selection_run_id = ?""", [run_id],
        ).fetchone()
        if not run:
            raise FinalistRetestError("CONTROL_SELECTION_RUN_MISSING")
        if str(run[2]) != symbol or str(run[3]) != side:
            raise FinalistRetestError("CONTROL_SCOPE_MISMATCH")
        latest = connection.execute(
            """select selection_run_id from selection_runs where symbol = ? and side = ?
               order by created_at_utc desc, selection_run_id desc limit 1""", [symbol, side],
        ).fetchone()
        if not latest or str(latest[0]) != run_id:
            raise FinalistRetestError("CONTROL_SELECTION_RUN_NOT_LATEST")
        stored_run_hash = str(run[6] or "")
        if not stored_run_hash:
            raise FinalistRetestError("CONTROL_WORKBOOK_CHANGED")
        if shared_issued_hash is None:
            shared_issued_hash = stored_run_hash
        elif shared_issued_hash != stored_run_hash:
            raise FinalistRetestError("CONTROL_WORKBOOK_CHANGED")
        expected_run_hash = str(run_hashes.get(key) or stored_run_hash)
        if expected_run_hash != stored_run_hash:
            raise FinalistRetestError("CONTROL_WORKBOOK_CHANGED")
        expected_selection_config = metadata.get("selection_config_sha256")
        if expected_selection_config is not None and str(run[5]) != str(expected_selection_config):
            raise FinalistRetestError("CONTROL_COHORT_MISMATCH")
        try:
            request_json = json.loads(str(run[4]))
        except (TypeError, ValueError) as error:
            raise FinalistRetestError("CONTROL_METADATA_INVALID") from error
        if not isinstance(request_json, Mapping) or request_json.get("ranking_scope") != "RETEST_COHORT":
            raise FinalistRetestError("CONTROL_SCOPE_MISMATCH")
        if str(request_json.get("bulk_retest_job_id") or "") != expected_job:
            raise FinalistRetestError("CONTROL_COHORT_MISMATCH")
        if str(request_json.get("cohort_sha256") or "") != expected_cohort or str(request_json.get("manifest_sha256") or "") != expected_manifest or str(request_json.get("config_sha256") or "") != expected_config:
            raise FinalistRetestError("CONTROL_COHORT_MISMATCH")
        if expected_selection_config is not None and str(request_json.get("selection_config_sha256") or "") != str(expected_selection_config):
            raise FinalistRetestError("CONTROL_COHORT_MISMATCH")
        snapshot = connection.execute(
            """select strategy_id, result_id_at_selection, auto_status, auto_score,
                      auto_rank, auto_reason, auto_analog_of_strategy_id
                 from selection_results where selection_run_id = ? order by strategy_id""", [run_id]
        ).fetchall()
        raw_members = request_json.get("cohort_members")
        if not isinstance(raw_members, list):
            raise FinalistRetestError("CONTROL_COHORT_MISMATCH")
        try:
            request_members = tuple(sorted((int(pair[0]), int(pair[1])) for pair in raw_members if isinstance(pair, (list, tuple)) and len(pair) == 2))
        except (TypeError, ValueError, IndexError):
            raise FinalistRetestError("CONTROL_COHORT_MISMATCH") from None
        if request_members != tuple(sorted((int(item[0]), int(item[1])) for item in snapshot)):
            raise FinalistRetestError("CONTROL_COHORT_MISMATCH")
        submitted = {int(row["Strategy ID"]): row for row in group}
        if set(submitted) != {int(item[0]) for item in snapshot}:
            raise FinalistRetestError("CONTROL_ROWSET_MISMATCH")
        ids = sorted(submitted)
        current = dict(connection.execute(
            "select strategy_id, current_result_id from strategies where strategy_id in (select unnest(?::bigint[]))", [ids]
        ).fetchall())
        snapshot_by_id = {int(item[0]): item for item in snapshot}
        exact_rows: list[dict[str, object]] = []
        for strategy_id in ids:
            row = submitted[strategy_id]
            snap = snapshot_by_id[strategy_id]
            if current.get(strategy_id) != snap[1] or int(row["Result ID"]) != int(snap[1]):
                raise FinalistRetestError("CONTROL_STALE_RESULTS")
            def normalized(value: object) -> str | None:
                return None if value in (None, "") else str(value)
            try:
                score_matches = (
                    row["Score"] in (None, "") and snap[3] is None
                    or Decimal(str(row["Score"])) == Decimal(str(snap[3]))
                )
            except (ArithmeticError, ValueError):
                score_matches = False
            if normalized(row["Auto Status"]) != normalized(snap[2]) or normalized(row["Auto Rank"]) != normalized(snap[4]) or normalized(row["Auto Reason"]) != normalized(snap[5]) or normalized(row["Auto Analog Of ID"]) != normalized(snap[6]) or not score_matches:
                raise FinalistRetestError("CONTROL_IMMUTABLE_FIELDS_CHANGED")
            exact_rows.append({key_name: row.get(key_name) for key_name in (
                "Pair", "Direction", "Strategy ID", "Result ID", "Auto Status", "Auto Rank", "Auto Analog Of ID", "Auto Reason", "Effective Start", "Effective End", "Score",
            )})
        exact_rowsets[key] = exact_rows
        prepared.append((symbol, side, group, snapshot, run_id))
        immutable.extend(exact_rows)
    expected_immutable = metadata.get("immutable_content_sha256")
    if expected_immutable is not None and str(expected_immutable) != canonical_digest(immutable):
        raise FinalistRetestError("CONTROL_IMMUTABLE_FIELDS_CHANGED")
    expected_rowsets = metadata.get("exact_rowsets_sha256")
    if expected_rowsets is not None and str(expected_rowsets) != canonical_digest(exact_rowsets):
        raise FinalistRetestError("CONTROL_ROWSET_MISMATCH")
    if canonical_json(issued_rowsets) != canonical_json(exact_rowsets):
        raise FinalistRetestError("CONTROL_ROWSET_MISMATCH")
    issued_hash_metadata = metadata.get("issued_workbook_sha256")
    if issued_hash_metadata is not None and str(issued_hash_metadata) != str(shared_issued_hash or ""):
        raise FinalistRetestError("CONTROL_WORKBOOK_CHANGED")

    now = datetime.now(timezone.utc)
    review_ids: list[str] = []
    connection.execute("begin transaction")
    try:
        for _symbol, _side, group, _snapshot, run_id in prepared:
            review_id = review_key(actual_hash, run_id)
            existing = connection.execute("select review_import_id from selection_review_imports where review_import_id = ?", [review_id]).fetchone()
            if existing:
                review_ids.append(review_id)
                continue
            connection.execute(
                "insert into selection_review_imports (review_import_id, selection_run_id, workbook_sha256, imported_at_utc, row_count) values (?, ?, ?, ?, ?)",
                [review_id, run_id, review_id, now, len(group)],
            )
            decisions = []
            retest_ids = []
            for row in group:
                strategy_id = int(row["Strategy ID"])
                rank = None if row["User Rank"] in (None, "") else int(row["User Rank"])
                analog = None if row["Analog Of ID"] in (None, "") else int(row["Analog Of ID"])
                status = str(row["User Status"]).strip().upper()
                decisions.append([review_id, strategy_id, status, rank, analog, "" if row["Comment"] is None else str(row["Comment"])])
                if row["RETEST"] == "RETEST":
                    retest_ids.append(strategy_id)
            connection.executemany("insert into selection_review_rows (review_import_id, strategy_id, user_status, user_rank, user_analog_of_strategy_id, comment) values (?, ?, ?, ?, ?, ?)", decisions)
            ids = [int(row["Strategy ID"]) for row in group]
            connection.execute("delete from strategy_tags where tag = 'REJECTED' and strategy_id in (select unnest(?::bigint[]))", [ids])
            rejected = [[item[1], "REJECTED", "SELECTION_REVIEW", review_id, now] for item in decisions if item[2] == "REJECTED"]
            if rejected:
                connection.executemany("insert into strategy_tags (strategy_id, tag, source, source_ref, updated_at_utc) values (?, ?, ?, ?, ?)", rejected)
            connection.execute("delete from strategy_tags where tag = 'RETEST' and source = 'SELECTION_REVIEW' and strategy_id in (select unnest(?::bigint[]))", [ids])
            retest = [[strategy_id, "RETEST", "SELECTION_REVIEW", review_id, now] for strategy_id in retest_ids]
            if retest:
                connection.executemany("insert into strategy_tags (strategy_id, tag, source, source_ref, updated_at_utc) values (?, ?, ?, ?, ?) on conflict (strategy_id, tag) do update set source=excluded.source, source_ref=excluded.source_ref, updated_at_utc=excluded.updated_at_utc", retest)
            review_ids.append(review_id)
        connection.execute("commit")
    except Exception:
        _rollback_quietly(connection)
        raise
    return {"review_import_ids": review_ids, "group_count": len(prepared), "row_count": len(rows), "workbook_sha256": actual_hash}


def _config_digest(config: object) -> str:
    if isinstance(config, Path):
        try:
            return sha256(config.read_bytes()).hexdigest()
        except OSError as error:
            raise FinalistRetestError("CONFIG_INVALID", "tester config is unavailable") from error
    return canonical_digest(config if config is not None else {
        "InitialBalance": 1000,
        "balance_percentage_long": 100,
        "balance_percentage_short": 100,
        "risk_long": 1,
        "risk_short": 1,
    })


def finalist_retest_config_digest(templates: Mapping[object, object], tester_config: object = None) -> str:
    """Hash the inputs that can change generated retest strategy JSON."""
    if not isinstance(templates, Mapping):
        raise FinalistRetestError("CONFIG_INVALID", "retest templates are invalid")
    return canonical_digest({
        "tester_config_sha256": _config_digest(tester_config),
        "template_sha256": {
            str(side): _config_digest(Path(template) if isinstance(template, str) else template)
            for side, template in sorted(templates.items(), key=lambda item: str(item[0]))
        },
    })


def build_finalist_retest_manifest(
    connection: duckdb.DuckDBPyConnection,
    templates: Mapping[object, object],
    output_dir: Path,
    *,
    test_start: object,
    test_end: object,
    include_reserve: bool = False,
    listing_dates: Mapping[str, object] | Path | None = None,
    tester_config: object = None,
    job_id: str | None = None,
) -> FinalistRetestBatch:
    """Freeze and publish one native SINGLE_MODE manifest for the cohort."""
    cohort = freeze_finalist_cohort(
        connection,
        test_start=test_start,
        test_end=test_end,
        include_reserve=include_reserve,
        listing_dates=listing_dates,
    )
    if not cohort.members:
        raise FinalistRetestError("COHORT_EMPTY", "bulk retest cohort has no runnable members")
    if not isinstance(templates, Mapping):
        raise FinalistRetestError("CONFIG_INVALID", "retest templates are invalid")
    run_id = job_id or f"finalist-retest-{sha256(cohort.cohort_sha256.encode('ascii')).hexdigest()[:24]}"
    generated: list[dict[str, object]] = []
    strategy_runs: dict[str, str] = {}
    names_by_candidate: dict[str, list[str]] = {}
    for member in cohort.members:
        try:
            template = templates[member["side"]]
        except (KeyError, TypeError) as error:
            raise FinalistRetestError("CONFIG_INVALID", f"template is missing for {member['side']}") from error
        if isinstance(template, Mapping):
            source = dict(template)
        else:
            path = Path(template)
            try:
                source = json.loads(path.read_text(encoding="utf-8"))
            except (OSError, UnicodeDecodeError, json.JSONDecodeError) as error:
                raise FinalistRetestError("CONFIG_INVALID", "retest template is invalid") from error
            if not isinstance(source, Mapping):
                raise FinalistRetestError("CONFIG_INVALID", "retest template is invalid")
        orders = [
            {"open_ma": item["open_ma"], "shift_bp": item["shift_bp"]}
            for item in member["orders"]  # type: ignore[index]
        ]
        strategy = generate_strategy(
            source,
            {
                "symbol": member["symbol"], "side": member["side"], "timeframe": member["timeframe"],
                "common_close_ma": member["close_ma_len"], "order_count": member["order_count"],
                "structure_id": f"RETEST_{member['strategy_id']}", "orders": orders,
            },
            tuple(Decimal(str(item["lot_x"])) for item in member["orders"]),  # type: ignore[index]
            LotMethod.EQUAL,
            AlgorithmConfig.defaults(),
        )
        strategy["name"] = str(member["strategy_name"])
        basic = strategy.get("basic")
        if isinstance(basic, dict):
            basic.update({
                "my_fix_balance": 1000.0,
                "balance_percentage_long": 100.0,
                "balance_percentage_short": 100.0,
                "risk_long": 1.0,
                "risk_short": 1.0,
            })
        filename = f"{member['strategy_name']}.json"
        strategy_runs[filename] = str(member["analysis_run_id"])
        names_by_candidate.setdefault(str(member["candidate_identity"]), []).append(str(member["strategy_name"]))
        generated.append(strategy)
    hashes = {
        f"{member['strategy_name']}.json": canonical_digest(strategy)
        for member, strategy in zip(cohort.members, generated, strict=True)
    }
    config_sha256 = finalist_retest_config_digest(templates, tester_config)
    unsigned: dict[str, object] = {
        "format_version": 1,
        "analysis_run_id": run_id,
        "event_mode": "real_independent_events",
        "strategy_count": len(generated),
        "strategy_json_sha256": hashes,
        "strategy_analysis_run_ids": strategy_runs,
        "candidate_identity_to_strategy_names": {
            key: sorted(names) for key, names in sorted(names_by_candidate.items())
        },
        "finalist_retest": {
            **cohort.as_manifest(),
            "job_id": run_id,
            "tester_config_sha256": config_sha256,
        },
        "scope": cohort.scope,
        "test_start": cohort.requested_start,
        "test_end": cohort.requested_end,
        "cohort_sha256": cohort.cohort_sha256,
        "tester_config_sha256": config_sha256,
    }
    manifest = {**unsigned, "generation_manifest_sha256": canonical_digest(unsigned)}
    try:
        from .performance_v2_retest import _publish_retest

        strategies_path, manifest_path = _publish_retest(Path(output_dir).resolve(), generated, manifest)
    except FinalistRetestError:
        raise
    except Exception as error:
        raise FinalistRetestError("PUBLICATION_FAILED", "bulk retest manifest publication failed") from error
    return FinalistRetestBatch(run_id, strategies_path, manifest_path, cohort, config_sha256, len(generated))


def _boundary(value: object, field: str) -> tuple[datetime, str]:
    if isinstance(value, datetime):
        parsed = value if value.tzinfo is not None else value.replace(tzinfo=timezone.utc)
        parsed = parsed.astimezone(timezone.utc)
        return parsed, parsed.date().isoformat()
    if isinstance(value, date):
        return datetime.combine(value, time.min, tzinfo=timezone.utc), value.isoformat()
    if isinstance(value, str):
        try:
            parsed_date = date.fromisoformat(value)
        except ValueError:
            raise FinalistRetestError("INVALID_TEST_RANGE", f"{field} must be an ISO date") from None
        if parsed_date.isoformat() != value:
            raise FinalistRetestError("INVALID_TEST_RANGE", f"{field} must be an ISO date")
        return datetime.combine(parsed_date, time.min, tzinfo=timezone.utc), value
    raise FinalistRetestError("INVALID_TEST_RANGE", f"{field} must be an ISO date")


def _timestamp(value: object) -> datetime | None:
    if value is None:
        return None
    if hasattr(value, "to_pydatetime"):
        value = value.to_pydatetime()
    if isinstance(value, datetime):
        return (value if value.tzinfo is not None else value.replace(tzinfo=timezone.utc)).astimezone(timezone.utc)
    if isinstance(value, date):
        return datetime.combine(value, time.min, tzinfo=timezone.utc)
    if isinstance(value, str):
        try:
            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            return None
        return (parsed if parsed.tzinfo is not None else parsed.replace(tzinfo=timezone.utc)).astimezone(timezone.utc)
    return None


def _listing_map(value: Mapping[str, object] | Path | None) -> Mapping[str, object]:
    if value is None:
        return {}
    if isinstance(value, Mapping):
        return value
    if isinstance(value, Path):
        try:
            from .loader import load_listing_dates

            loaded = load_listing_dates(value)
        except Exception as error:
            raise FinalistRetestError("LISTING_DATE_INVALID", "listing dates are invalid") from error
        if not isinstance(loaded, Mapping):
            raise FinalistRetestError("LISTING_DATE_INVALID", "listing dates are invalid")
        return loaded
    raise FinalistRetestError("LISTING_DATE_INVALID", "listing dates are invalid")


def _latest_selection_runs(connection: duckdb.DuckDBPyConnection) -> dict[tuple[str, str], str]:
    rows = connection.execute(
        """
        select symbol, side, selection_run_id
          from (
            select symbol, side, selection_run_id,
                   row_number() over (
                     partition by symbol, side order by created_at_utc desc, selection_run_id desc
                   ) as rn
              from selection_runs
          )
         where rn = 1
        """
    ).fetchall()
    return {(str(symbol), str(side)): str(run_id) for symbol, side, run_id in rows}


def _effective_selection_decisions(connection: duckdb.DuckDBPyConnection) -> dict[int, tuple[str, int | None, str | None]]:
    runs = _latest_selection_runs(connection)
    if not runs:
        return {}
    output: dict[int, tuple[str, int | None, str | None]] = {}
    for run_id in runs.values():
        reviewed = connection.execute(
            """
            select review_import_id from selection_review_imports
             where selection_run_id = ? order by imported_at_utc desc, review_import_id desc limit 1
            """,
            [run_id],
        ).fetchone()
        review_id = str(reviewed[0]) if reviewed else None
        rows = connection.execute(
            """
            select r.strategy_id, r.auto_status, r.auto_rank, r.prior_rejected,
                   v.user_status, v.user_rank
              from selection_results r
              left join selection_review_rows v
                on v.strategy_id = r.strategy_id and v.review_import_id = ?
             where r.selection_run_id = ?
            """,
            [review_id, run_id],
        ).fetchall()
        for strategy_id, auto_status, auto_rank, prior_rejected, user_status, user_rank in rows:
            status = str(user_status or ("REJECTED" if prior_rejected else auto_status))
            rank = user_rank if user_rank is not None else auto_rank
            output[int(strategy_id)] = (status, None if rank is None else int(rank), run_id)
    return output


def _source_orders(connection: duckdb.DuckDBPyConnection, strategy_id: int, order_count: object) -> list[dict[str, object]] | None:
    if isinstance(order_count, bool) or not isinstance(order_count, int) or not 1 <= order_count <= 4:
        return None
    rows = connection.execute(
        """
        select o.order_id, o.open_ma_len, o.shift_bp, o.lot_x, o.analysis_run_id, o.plateau_id,
               p.plateau_point_count, o.base_point_trades, p.plateau_total_trades
          from strategy_orders o
          left join analysis_plateaus p
            on p.analysis_run_id = o.analysis_run_id and p.plateau_id = o.plateau_id
         where o.strategy_id = ? order by o.order_id
        """,
        [strategy_id],
    ).fetchall()
    if len(rows) != order_count or [row[0] for row in rows] != list(range(1, order_count + 1)):
        return None
    result: list[dict[str, object]] = []
    for row in rows:
        if (
            not isinstance(row[1], int) or row[1] <= 0
            or not isinstance(row[2], int) or row[2] < 0
            or not isinstance(row[4], str) or not row[4].strip()
            or not isinstance(row[5], str) or not row[5].strip()
            or row[3] is None or row[6] is None or row[8] is None
        ):
            return None
        result.append({
            "order_id": int(row[0]), "open_ma": int(row[1]), "shift_bp": int(row[2]),
            "lot_x": str(row[3]), "analysis_run_id": str(row[4]), "plateau_id": str(row[5]),
            "plateau_point_count": None if row[6] is None else int(row[6]),
            "base_point_trades": None if row[7] is None else int(row[7]),
            "plateau_total_trades": None if row[8] is None else int(row[8]),
        })
    return result


def freeze_finalist_cohort(
    connection: duckdb.DuckDBPyConnection,
    *,
    test_start: object,
    test_end: object,
    include_reserve: bool = False,
    listing_dates: Mapping[str, object] | Path | None = None,
    warmup_hours: int = LISTING_WARMUP_HOURS,
) -> FinalistRetestCohort:
    """Read and freeze the current effective finalist/reserve population."""
    require_performance_v2(connection)
    if type(include_reserve) is not bool or isinstance(warmup_hours, bool) or not isinstance(warmup_hours, int) or warmup_hours < 0:
        raise FinalistRetestError("INVALID_REQUEST", "bulk retest options are invalid")
    requested_start_dt, requested_start = _boundary(test_start, "test_start")
    requested_end_dt, requested_end = _boundary(test_end, "test_end")
    if requested_start_dt >= requested_end_dt:
        raise FinalistRetestError("INVALID_TEST_RANGE", "test_start must be before test_end")
    scope = "FINALIST_RESERVE" if include_reserve else "FINALIST"
    allowed = {"FINALIST", "RESERVE"} if include_reserve else {"FINALIST"}
    decisions = _effective_selection_decisions(connection)
    listings = _listing_map(listing_dates)
    rows = connection.execute(
        """
        select s.strategy_id, s.strategy_name, s.symbol, s.side, s.timeframe, s.close_ma_len,
               s.order_count, s.analysis_run_id, s.candidate_identity, s.current_result_id,
               r.report_start_utc, r.report_end_utc
          from strategies s
          join strategy_results r on r.result_id = s.current_result_id and r.strategy_id = s.strategy_id
         where s.lifecycle_status = 'ACTIVE' and s.current_result_id is not null
         order by s.symbol, s.side, s.strategy_id
        """
    ).fetchall()
    members: list[Mapping[str, object]] = []
    exclusions: list[RetestExclusion] = []
    for row in rows:
        strategy_id, name, symbol, side, timeframe, close_ma, order_count, analysis_run, candidate, result_id, report_start, report_end = row
        decision = decisions.get(int(strategy_id))
        if decision is None or decision[0] not in allowed:
            continue
        if not isinstance(name, str) or not name.strip() or not isinstance(symbol, str) or not symbol.strip() or side not in {"LONG", "SHORT"}:
            exclusions.append(RetestExclusion(int(strategy_id), str(name) if name is not None else None, str(symbol) if symbol is not None else None, "INVALID_STRATEGY_SOURCE", str(side) if side in {"LONG", "SHORT"} else None))
            continue
        orders = _source_orders(connection, int(strategy_id), order_count)
        if orders is None or not isinstance(analysis_run, str) or not analysis_run.strip() or not isinstance(candidate, str) or not candidate.strip():
            exclusions.append(RetestExclusion(int(strategy_id), name, symbol, "INVALID_STRATEGY_SOURCE", str(side) if side in {"LONG", "SHORT"} else None))
            continue
        listing = _timestamp(listings.get(symbol))
        if listing is None:
            reason = "LISTING_DATE_MISSING" if symbol not in listings else "LISTING_DATE_INVALID"
            exclusions.append(RetestExclusion(int(strategy_id), name, symbol, reason, str(side) if side in {"LONG", "SHORT"} else None))
            continue
        effective_start = max(requested_start_dt, listing + timedelta(hours=warmup_hours))
        if effective_start >= requested_end_dt:
            exclusions.append(RetestExclusion(int(strategy_id), name, symbol, "EMPTY_EFFECTIVE_PERIOD", str(side) if side in {"LONG", "SHORT"} else None))
            continue
        members.append({
            "strategy_id": int(strategy_id), "strategy_name": name, "symbol": symbol, "side": str(side),
            "timeframe": str(timeframe), "close_ma_len": int(close_ma), "order_count": int(order_count),
            "analysis_run_id": analysis_run, "candidate_identity": candidate,
            "result_id": int(result_id), "effective_status": decision[0], "effective_rank": decision[1],
            "report_start_utc": _timestamp(report_start), "report_end_utc": _timestamp(report_end),
            "listing_date_utc": listing, "requested_start": requested_start, "requested_end": requested_end,
            "effective_start": effective_start, "effective_end": requested_end_dt,
            "orders": orders,
        })
    if len(members) > MAX_COHORT_MEMBERS:
        raise FinalistRetestError("COHORT_TOO_LARGE", "bulk retest cohort exceeds 10000 members")
    return FinalistRetestCohort(
        scope, requested_start, requested_end, warmup_hours, tuple(members), tuple(exclusions), cohort_digest(members)
    )


__all__ = [
    "canonical_json", "canonical_digest", "cohort_digest", "review_key",
    "canonical_provenance_json", "canonical_provenance_digest", "deterministic_review_key",
    "LISTING_WARMUP_HOURS", "MAX_COHORT_MEMBERS", "FinalistRetestError", "RetestExclusion",
    "FinalistRetestCohort", "FinalistRetestBatch", "FinalistRetestImportResult",
    "freeze_finalist_cohort", "build_finalist_retest_manifest", "apply_finalist_retest_outcomes",
    "finalist_retest_config_digest",
    "execute_finalist_retest_import", "import_finalist_retest", "CONTROL_WORKBOOK_SHEETS",
    "write_combined_control_workbook", "combined_control_workbook_bytes",
    "validate_combined_control_workbook", "read_combined_control_workbook",
    "import_combined_control_workbook",
]
