"""MRS2 history in the trade store: P2 segment plan per pair, simulation of a segment, the history run.

Contract: docs/specs/2026-10-17-mrs2-history-store-and-tester-sampling.md («Stage 1 contract»).
"""

from __future__ import annotations

from concurrent.futures import FIRST_COMPLETED, Executor, ProcessPoolExecutor, ThreadPoolExecutor, wait
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
import json
from pathlib import Path
from typing import Callable, Mapping, Sequence
import uuid

import duckdb
import openpyxl

from .minute_data import ensure_days
from .mrs2_simulator import MinuteBars, RunSettings, load_minutes, simulate_mrs2, simulate_period
from .period_csv import GENERATOR_VERSION, PERIOD_DAYS, Point, first_period_day, has_warmup_data, journal_append, ms_of, split_periods
from .trade_store import (
    SegmentKey, SegmentTables, TradeStore, TradeStoreError, build_tables, hourly_rows, metrics_columns, point_id, price_rows,
    trade_columns, write_segment_files,
)

ENGINE = GENERATOR_VERSION
SOURCE = "simulator"
LATEST_DAYS = PERIOD_DAYS
VERDICT_RANK = {"GO": 0, "CHECK": 1, None: 2, "": 2, "STOP": 3}
EXCHANGE_FOLDERS = {"bybit": "Bybit", "bingx": "BingX"}


@dataclass(frozen=True)
class BingXMinutes:
    """What the plan step of a BingX pair needs to ensure and verify its minute days (Panel stage 5b): the data
    folder, (api symbol, listing ms) per pair, the verification records known from the registry, the download
    mode. `client` is for tests (None: a real BingX client in the worker)."""

    data_root: str
    pairs: Mapping[str, tuple[str, int]]
    known: Mapping[str, Mapping[date, Mapping]] = field(default_factory=dict)
    mode: str = "incremental"
    client: object | None = None

    def symbol_dir(self, symbol: str) -> str:
        return str(Path(self.data_root) / self.pairs[symbol][0])


def exchange_run_settings(repo_root: Path, exchange: str) -> RunSettings:
    """The simulator settings of an exchange: Bybit keeps the reference `RunSettings()`; another exchange takes
    its fees, slippage and funding from `templates/exchanges/<Exchange>/mrs3-tester.json`."""
    if exchange == "bybit":
        return RunSettings()
    tester = json.loads((Path(repo_root) / "templates" / "exchanges" / EXCHANGE_FOLDERS[exchange]
                         / "mrs3-tester.json").read_text(encoding="utf-8"))
    return RunSettings(maker_fee=float(tester["MakerFee"]), taker_fee=float(tester["TakerFee"]),
                       slippage=float(tester["SlippagePercent"]), funding_rate=float(tester["FundingRate"]),
                       funding_interval_hours=int(tester["FundingIntervalHours"]))


def pair_order(registry: Path) -> list[tuple[str, date]]:
    """Pairs with their listing day: screening verdict GO, CHECK, none, STOP; inside a group by 7-day turnover."""
    workbook = openpyxl.load_workbook(registry, read_only=True, data_only=True)
    try:
        pairs = workbook["Пары"]
        head = list(next(pairs.iter_rows(min_row=1, max_row=1, values_only=True)))
        listing = {}
        for row in pairs.iter_rows(min_row=2, values_only=True):
            if row[0]:
                value = row[head.index("Дата листинга на Bybit (UTC)")]
                listing[str(row[0]).strip().upper()] = value.date() if isinstance(value, datetime) else None
        screening = {}
        if "Скрининг" in workbook.sheetnames:
            sheet = workbook["Скрининг"]
            shead = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if row[0] and (row[shead.index("Сторона")] in (None, "LONG")):
                    volume = row[shead.index("Объём 7д, USDT")] if "Объём 7д, USDT" in shead else None
                    screening[str(row[0]).strip().upper()] = (row[shead.index("Вердикт")], float(volume or 0))
    finally:
        workbook.close()

    def key(symbol: str):
        verdict, volume = screening.get(symbol, (None, 0.0))
        return (VERDICT_RANK.get(verdict, 2), -volume, symbol)
    return [(symbol, listing[symbol]) for symbol in sorted(listing, key=key) if listing[symbol] is not None]


def pair_segments(listing: date, end_day: date, warmup_days: int = 4) -> list[tuple[date, date]]:
    """P2 history periods over [listing + warm-up, end − 20 days), then the latest window [end − 20, end)."""
    first = first_period_day(listing, warmup_days)
    latest = end_day - timedelta(days=LATEST_DAYS)
    if first > latest:
        return []
    return [*split_periods(first, latest), (latest, end_day)]


def simulate_segment(minutes: MinuteBars, start: date, end: date, *, tick: float, qty_step: float,
                     points: Sequence[Point], settings: RunSettings = RunSettings()) -> SegmentTables:
    """Fresh simulator run of [start, end) for every point: trades, hourly series (with the price series), metrics."""
    start_ms, end_ms = ms_of(start), ms_of(end)
    trades, hourly, metrics = [], [price_rows(minutes, start_ms, end_ms)], []
    for point in sorted(points, key=lambda p: point_id(p.timeframe, p.shift_bp, p.open_len, p.close_len)):
        pid = point_id(point.timeframe, point.shift_bp, point.open_len, point.close_len)
        cycles = simulate_mrs2(minutes, start_ms, end_ms, timeframe=point.timeframe, open_len=point.open_len,
                               open_mult=(10_000 - point.shift_bp) / 10_000, close_len=point.close_len, tick=tick,
                               settings=settings)
        result = simulate_period(cycles, minutes, start_ms, end_ms, qty_step=qty_step, settings=settings)
        trades.append(trade_columns(pid, result))
        hourly.append(hourly_rows(pid, result, start_ms, end_ms))
        metrics.append(metrics_columns(pid, result))
    return build_tables(trades, hourly, metrics)


# -- worker tasks ---------------------------------------------------------------------


def prepare_pair(symbol: str, listing: date, data_root: str, end_day: date, download: bool, warmup_days: int,
                 download_threads: int, bingx: BingXMinutes | None = None, rate_share: int = 1
                 ) -> tuple[list[tuple[date, date]], dict, frozenset[date] | None, list[dict]]:
    """Plan the pair's segments and ensure the candle days of the plan.

    Returns (plan, journal record, accepted days, verification records). Bybit: accepted days None (a day counts
    when its file exists). BingX: the days `verify_bingx_days` accepted (without download: the days the registry
    already accepts), and its new records for the main process to write."""
    plan = pair_segments(listing, end_day, warmup_days)
    record = {"pair": symbol, "listing_date": listing.isoformat(), "end_day": end_day.isoformat()}
    if not plan:
        return [], {**record, "status": "NO_FULL_PERIOD"}, None, []
    first, last = plan[0][0] - timedelta(days=warmup_days), end_day - timedelta(days=1)
    if bingx is not None:
        from .pair_registry.minute_days import ACCEPTED, verify_bingx_days

        known = bingx.known.get(symbol, {})
        if not download:
            # no request: a recorded day counts only while its file is the one verified (or still absent)
            accepted = frozenset(day for day, row in known.items() if row["status"] in ACCEPTED
                                 and _file_sha(Path(bingx.symbol_dir(symbol)) / f"{symbol}{day.isoformat()}_1m.parquet")
                                 == row["file_sha256"])
            return plan, {**record, "status": "PLANNED", "segments": len(plan)}, accepted, []
        api_symbol, launch_ms = bingx.pairs[symbol]
        # the days of the pair in parallel (docs/specs/2026-10-20-panel-sync-by-exchange.md, part 2); the pairs
        # prepared at the same time share the request rate, so each worker's client takes its share of it
        days = verify_bingx_days(Path(bingx.data_root), symbol, api_symbol, launch_ms, first, last, known=known,
                                 client=bingx.client, mode=bingx.mode, threads=download_threads,
                                 rate_share=rate_share)
        return plan, {**record, "status": "PLANNED", "segments": len(plan), "days_verified": len(days.records),
                      "days_refused": {day.isoformat(): reason for day, reason in days.refused.items()}}, \
            frozenset(days.accepted), days.records
    if not download:
        return plan, {**record, "status": "PLANNED", "segments": len(plan)}, None, []
    days = ensure_days(Path(data_root), symbol, first, last, workers=download_threads)
    return plan, {**record, "status": "PLANNED", "segments": len(plan), "days_created": len(days["created"]),
                  "days_failed": [day.isoformat() for day in days["failed"]]}, None, []


def _file_sha(path: Path) -> str | None:
    import hashlib

    return hashlib.sha256(path.read_bytes()).hexdigest() if path.is_file() else None


def missing_days(data_root: Path, symbol: str, first: date, end: date) -> list[date]:
    """Days of [first, end) without a candle file in the tester's data folder."""
    folder = Path(data_root) / symbol
    return [first + timedelta(days=offset) for offset in range((end - first).days)
            if not (folder / f"{symbol}{(first + timedelta(days=offset)).isoformat()}_1m.parquet").is_file()]


def simulate_and_write(store_root: str, key: SegmentKey, data_root: str, tick: float, qty_step: float,
                       points: Sequence[Point], settings: RunSettings, replace_unregistered: bool = False,
                       symbol_dir: str | None = None, accepted: frozenset[date] | None = None) -> dict:
    """Simulate one segment and write its files; returns the catalog record or a skip status.

    Every candle day of the warm-up and the span must exist (Bybit) or be accepted by the minute verification
    (BingX, `accepted`; a verified day without a file had no trades): a day lost to a failed download would
    otherwise be stored as a quiet market under the segment's permanent id."""
    first = key.start - timedelta(days=settings.warmup_days)
    folder = Path(symbol_dir) if symbol_dir else Path(data_root) / key.symbol
    if accepted is None:
        absent = missing_days(Path(data_root), key.symbol, first, key.end)
    else:
        absent = [first + timedelta(days=offset) for offset in range((key.end - first).days)
                  if first + timedelta(days=offset) not in accepted]
    if absent:
        return {"status": "MISSING_DAYS", "days": [day.isoformat() for day in absent]}
    minutes = load_minutes(folder, key.symbol, ms_of(first), ms_of(key.end))
    if not has_warmup_data(minutes, key.start, settings.warmup_days):
        return {"status": "NO_WARMUP_DATA"}
    tables = simulate_segment(minutes, key.start, key.end, tick=tick, qty_step=qty_step, points=points,
                              settings=settings)
    record = write_segment_files(Path(store_root), key, tables, settings=settings, tick_size=tick, qty_step=qty_step,
                                 replace_unregistered=replace_unregistered)
    return {"status": "OK", "record": record}


# -- history run ----------------------------------------------------------------------------

PREPARE_IN_FLIGHT = 2


def _reactivate(store: TradeStore, segment: SegmentKey, journal: Path, base: Mapping) -> bool:
    """Reactivate a superseded segment planned again; False (journaled) when its files are missing or damaged."""
    try:
        store.reactivate(segment.segment_id)
    except (TradeStoreError, OSError, duckdb.Error) as error:
        journal_append(journal, {**base, "segment_id": segment.segment_id, "status": "REACTIVATE_FAILED",
                                 "error": f"{type(error).__name__}: {error}"})
        return False
    journal_append(journal, {**base, "segment_id": segment.segment_id, "status": "REACTIVATED"})
    return True


def run_history(store: TradeStore, *, data_root: Path, pairs: Sequence[tuple[str, date]], end_day: date,
                instruments: Mapping[str, tuple[float, float]], points: Sequence[Point], workers: int,
                download: bool = True, settings: RunSettings = RunSettings(), exchange: str = "bybit",
                engine: str = ENGINE, run_args: Mapping | None = None,
                progress: Callable[[dict], None] | None = None, bingx: BingXMinutes | None = None,
                on_minute_records: Callable[[list[dict]], None] | None = None) -> dict:
    """Bring every pair's planned segments into the store, pairs in the given priority order.

    Workers simulate and write segment files; only this process writes the catalog and the journal. Pairs are
    prepared lazily (at most PREPARE_IN_FLIGHT at a time) so the first pairs' segments start before later pairs
    are even planned; download threads share the `workers` limit. BingX (`bingx`): the plan step verifies the
    minute days; their records go to `on_minute_records` in this process.
    """
    if (exchange == "bingx") != (bingx is not None):
        raise ValueError("a BingX history run needs its minute description, and only BingX takes one")
    run_id = f"{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S')}-{uuid.uuid4().hex[:8]}"
    workers = max(1, workers)
    queue = list(dict.fromkeys(pairs))[::-1]          # duplicates dropped, popped from the end in priority order
    store.register_points(point_id(p.timeframe, p.shift_bp, p.open_len, p.close_len) for p in points)
    store.start_run(run_id, engine=engine, settings=settings, end_day=end_day, args=dict(run_args or {}))
    journal = store.root / "journal.jsonl"
    counts = {"written": 0, "reactivated": 0, "skipped": 0, "failed": 0}
    units = {"done": 0, "total": len(queue)}           # one unit per pair plan, plus one per segment to simulate

    def report(symbol: str, step: str, error: str | None = None) -> None:
        if progress is None:
            return
        if error is not None:
            progress({"error": error, "pair": symbol, "step": step})
        progress({"progress": {"done": units["done"], "total": units["total"], "step": step, "pair": symbol}})

    download_threads = max(1, workers // PREPARE_IN_FLIGHT)
    rate_share = min(PREPARE_IN_FLIGHT, workers)        # the BingX prepares running at once share the request rate
    pool: Executor = ThreadPoolExecutor(max_workers=1) if workers == 1 else ProcessPoolExecutor(max_workers=workers)
    pending: dict = {}

    def in_flight_prepares() -> int:
        return sum(1 for kind, _, _ in pending.values() if kind == "prepare")

    def fill() -> None:
        while queue and len(pending) < workers and in_flight_prepares() < PREPARE_IN_FLIGHT:
            symbol, listing = queue.pop()
            pending[pool.submit(prepare_pair, symbol, listing, str(data_root), end_day, download, settings.warmup_days,
                                download_threads, bingx, rate_share)] = ("prepare", symbol, None)

    try:
        with pool:
            fill()
            while pending:
                done, _ = wait(pending, return_when=FIRST_COMPLETED)
                for future in done:
                    kind, symbol, key = pending.pop(future)
                    base = {"run_id": run_id, "pair": symbol}
                    try:
                        result = future.result()
                    except Exception as error:  # one pair or segment never stops the others
                        counts["failed"] += 1
                        journal_append(journal, {**base, "stage": kind, "status": "ERROR",
                                                 "segment_id": key.segment_id if key else None,
                                                 "error": f"{type(error).__name__}: {error}"})
                        units["done"] += 1
                        report(symbol, kind, f"{type(error).__name__}: {error}")
                        continue
                    if kind == "prepare":
                        plan, record, accepted, records = result
                        if records and on_minute_records is not None:
                            on_minute_records(records)
                        journal_append(journal, {**base, **record})
                        symbol_dir = bingx.symbol_dir(symbol) if bingx is not None else None
                        known = store.statuses(exchange, symbol)
                        tick, qty_step = instruments[symbol]
                        for start, end in plan:
                            segment = SegmentKey(exchange, symbol, SOURCE, start, end, engine)
                            status = known.get(segment.segment_id)
                            if status == "active":
                                counts["skipped"] += 1
                            elif status == "superseded" and _reactivate(store, segment, journal, base):
                                counts["reactivated"] += 1
                            else:
                                # absent, or superseded with damaged files: simulate again and own the files
                                pending[pool.submit(simulate_and_write, str(store.root), segment, str(data_root),
                                                    tick, qty_step, points, settings, True, symbol_dir,
                                                    accepted)] = ("segment", symbol, segment)
                                units["total"] += 1
                        units["done"] += 1
                        report(symbol, "plan")
                        fill()
                        continue
                    if result["status"] == "OK":
                        try:
                            store.register(result["record"], replace_superseded=True)
                        except Exception as error:   # a segment already registered (duplicate plan, race): count,
                            counts["failed"] += 1      # keep the files, move on
                            journal_append(journal, {**base, "segment_id": key.segment_id, "status": "ERROR",
                                                     "error": f"{type(error).__name__}: {error}"})
                            continue
                        counts["written"] += 1
                        journal_append(journal, {**base, "segment_id": key.segment_id, "status": "OK",
                                                 "trades": result["record"]["trades"]})
                    else:
                        counts["skipped"] += 1
                        journal_append(journal, {**base, "segment_id": key.segment_id,
                                                 **{name: value for name, value in result.items()}})
                    units["done"] += 1
                    report(symbol, "segment", None if result.get("status") == "OK" else
                           f"{key.segment_id}: {result.get('status')} {result.get('missing_days', '')}".strip())
                fill()
    finally:
        store.finish_run(run_id, written=counts["written"] + counts["reactivated"], skipped=counts["skipped"],
                         failed=counts["failed"])
    return {"run_id": run_id, **counts}


__all__ = ["ENGINE", "SOURCE", "BingXMinutes", "exchange_run_settings", "missing_days", "pair_order",
           "pair_segments", "run_history", "simulate_segment"]
