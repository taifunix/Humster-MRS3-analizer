"""The «Пары и стратегии» main screen (Panel rework stage 3): actions with plan dialogs, views, pair card.

Contract: docs/specs/2026-10-18-panel-simulation-workflow.md, «Stage 3 contract». The registry is read read-only
(a writer — a running «Синхронизировать» — makes the views answer «реестр занят»); decisions and thresholds are
written through `pair_registry.records`. Every long action runs as one task of `panel_tasks.TaskRunner`.
"""

from __future__ import annotations

from dataclasses import asdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
import json
import math
from pathlib import Path
import re
import sys
import tempfile
from typing import Callable, Mapping, Sequence

import duckdb

from .pair_registry.history import sufficiency
from .pair_registry.liquidity import read_liquidity
from .pair_registry.records import FILTERS, effective_ready_rule, record_decision, record_setting
from .pair_registry.store import DATABASE, RegistryError, connect, load_registry_settings
from .error_sanitization import redact_local_paths
from .panel_strategies import (
    ALL_COLUMNS as STRATEGY_COLUMNS, DEFAULT_COLUMNS as STRATEGY_DEFAULT_COLUMNS, MARK_LABELS, StrategiesError,
    hide_cut, parse_rank, parse_status, strategy_rows,
)
from .panel_tasks import TaskRunner, TaskRunnerError
from .ready_store import ReadyRule, inputs_sha256, pair_periods

ACTIONS = ("sync", "history", "ready", "tester", "tick", "retest")
EXCHANGES = ("bybit", "bingx")
BINGX_PENDING = "BingX — этап 5: минутные данные BingX ещё не подключены"


class MainScreenError(ValueError):
    """A main-screen request is refused (message safe to show)."""


# -- views: columns, filter and sort shared by the table and the Excel export ----------------------------

VIEWS: dict[str, tuple[tuple[str, str], ...]] = {
    "pairs": (("symbol", "Пара"), ("stage", "Стадия"), ("history_status", "История"), ("periods", "Периодов"),
              ("last_period", "Последний период"), ("ready_status", "READY"),
              ("ready_source", "Источник последнего периода"), ("admitted", "Структур"),
              ("best_category", "Категория"), ("best_score", "Score"), ("best_dd", "DD"),
              ("tester_window", "Тестер"), ("tick_tests", "Тик-тестов"), ("finalist_strategies", "Финалистов"),
              ("decision", "Решение"), ("track_cuts", "Срезов"), ("track_admitted_cuts", "Допущено"),
              ("track_loss_rate", "Доля убыточных"), ("base_lot", "Базовый лот"), ("liquidity_day", "Ликвидность на")),
    "reference": (("symbol", "Пара"), ("display_name", "Название"), ("status", "Статус"),
                  ("symbol_type", "Тип"), ("listing_day", "Листинг"), ("tick_size", "Тик"), ("qty_step", "Шаг лота"),
                  ("min_notional", "Мин. ордер USDT"), ("max_order_qty", "Макс. ордер"),
                  ("max_leverage", "Макс. плечо"), ("risk_tiers", "Тиров риска"), ("funding_interval_min", "Фандинг, мин"),
                  ("maker_fee", "Maker"), ("taker_fee", "Taker"), ("api_symbol", "API символ"),
                  ("api_open", "API открытие")),
    "liquidity": (("symbol", "Пара"), ("as_of_day", "На день"), ("status", "Статус"),
                  ("window_weekdays", "Будней"), ("v25", "V25"), ("a15", "A15"), ("l", "L"), ("base_lot", "Базовый лот"),
                  ("missing_days", "Нет дней"), ("reason", "Причина")),
    "changes": (("changed", "Когда (UTC)"), ("symbol", "Пара"), ("kind", "Событие"), ("field", "Поле"),
                ("old_value", "Было"), ("new_value", "Стало")),
    "strategies": STRATEGY_COLUMNS,             # Panel stage 4a (panel_strategies)
}


_DECIMAL = re.compile(r"[+-]?(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?")


def _number(value: object) -> float | None:
    """A number for sorting: int/float values and decimal strings only (the same rule as app.js)."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(value) else None
    if isinstance(value, str) and _DECIMAL.fullmatch(value.strip()):
        return float(value)
    return None


def _text(value: object) -> str:
    """The value as JavaScript's String() shows it, so the filter matches the table in app.js."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and math.isfinite(value):
        if value.is_integer() and abs(value) < 1e21:
            return str(int(value))
        text = repr(value)
        if "e" in text and 1e-6 <= abs(value) < 1e21:      # JS writes these without an exponent
            text = format(Decimal(text), "f")
        return text
    return str(value)


def filter_rows(rows: Sequence[Mapping], columns: Sequence[str], text: str) -> list[Mapping]:
    """Case-insensitive substring over the shown columns (the same rule as the table in app.js)."""
    needle = (text or "").strip().casefold()
    if not needle:
        return list(rows)
    return [row for row in rows if any(needle in _text(row.get(key)).casefold() for key in columns)]


def sort_rows(rows: Sequence[Mapping], key: str | None, descending: bool = False) -> list[Mapping]:
    """Numbers numerically, text alphabetically, empty values last (the same rule as the table in app.js)."""
    if not key:
        return list(rows)
    filled = [row for row in rows if row.get(key) not in (None, "")]
    empty = [row for row in rows if row.get(key) in (None, "")]
    numeric = all(_number(row.get(key)) is not None for row in filled)
    ordered = sorted(filled, key=(lambda row: _number(row[key])) if numeric else (lambda row: _text(row[key]).casefold()),
                     reverse=descending)
    return ordered + empty


def _periods(count: int, new: bool = True) -> str:
    """«1 новый период», «2 новых периода», «5 новых периодов» (without «новый» when `new` is false)."""
    tail, tens = count % 10, count % 100
    if tail == 1 and tens != 11:
        adjective, noun = "новый ", "период"
    elif tail in (2, 3, 4) and tens not in (12, 13, 14):
        adjective, noun = "новых ", "периода"
    else:
        adjective, noun = "новых ", "периодов"
    return f"{count} {adjective if new else ''}{noun}"


def _iso(ms: object) -> str | None:
    return datetime.fromtimestamp(ms / 1000, timezone.utc).date().isoformat() if isinstance(ms, int) else None


class MainScreenService:
    def __init__(self, repo_root: Path, config_path: Path, runner: TaskRunner, *, python: str = sys.executable,
                 today: Callable[[], date] = lambda: datetime.now(timezone.utc).date(),
                 controller=None, chain_root: Path | None = None) -> None:
        self.repo_root = Path(repo_root)
        self.config_path = Path(config_path)
        self.runner = runner
        self._python = python
        self._today = today
        self._controller = controller
        self._chain_root = chain_root                # folder of the steps files; default: the system temp folder

    # -- configuration and storage --------------------------------------------------------------------------

    def _config(self) -> dict:
        try:
            return json.loads(self.config_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            raise MainScreenError("конфиг не читается") from None

    def _settings(self):
        try:
            return load_registry_settings(self.config_path)
        except RegistryError as error:
            raise MainScreenError(redact_local_paths(str(error))) from None

    def _read(self) -> duckdb.DuckDBPyConnection:
        database = self._settings().root / DATABASE
        if not database.is_file():
            raise MainScreenError("реестр ещё не создан: нажмите «Синхронизировать»")
        try:
            return duckdb.connect(str(database), read_only=True)
        except duckdb.Error:
            raise MainScreenError("реестр занят задачей: повторите, когда она закончится") from None

    def _write(self) -> duckdb.DuckDBPyConnection:
        try:
            return connect(self._settings().root)
        except duckdb.Error:
            raise MainScreenError("реестр занят задачей: повторите, когда она закончится") from None

    def _store_root(self) -> Path:
        root = (self._config().get("trade_store") or {}).get("root")
        if not isinstance(root, str) or not root.strip():
            raise MainScreenError("в config не задан trade_store.root")
        return Path(root)

    def _reader(self):
        """The trade store for reading only (no writable catalog connection in the Panel process), or None while
        the store has no catalog yet."""
        from .trade_store import TradeStore, TradeStoreError

        try:
            return TradeStore.reader(self._store_root())
        except TradeStoreError:
            return None

    @staticmethod
    def _exchange(value: object) -> str:
        if value not in EXCHANGES:
            raise MainScreenError("биржа: bybit или bingx")
        return str(value)

    # -- views ----------------------------------------------------------------------------------------------

    def view(self, name: str, exchange: str, options: Mapping | None = None) -> dict:
        if not isinstance(name, str) or name not in VIEWS:
            raise MainScreenError(f"неизвестный вид {name!r}")
        exchange = self._exchange(exchange)
        if name == "strategies":
            return self._view_strategies(exchange, dict(options or {}))
        con = self._read()
        try:
            rows = getattr(self, f"_view_{name}")(con, exchange)
        finally:
            con.close()
        return {"name": name, "exchange": exchange, "columns": [{"key": key, "label": label} for key, label in VIEWS[name]],
                "rows": rows}

    def _view_pairs(self, con, exchange: str) -> list[dict]:
        rows = con.execute(
            "select * from (select *, row_number() over (partition by symbol order by as_of_day desc) as latest "
            "from pipeline where exchange = ? and side = 'LONG') where latest = 1 order by symbol", [exchange]).fetchall()
        names = [column[0] for column in con.description]
        sources = self._ready_sources(exchange)
        out = []
        for row in rows:
            item = dict(zip(names, row))
            item["ready_source"] = sources.get(item["symbol"])
            liquidity = read_liquidity(con, exchange, item["symbol"])
            item["base_lot"] = liquidity["base_lot"] if liquidity else None
            item["liquidity_day"] = liquidity["as_of_day"] if liquidity else None
            out.append({key: (value.isoformat() if isinstance(value, date) else value) for key, value in item.items()
                        if key in dict(VIEWS["pairs"])})
        return out

    def _ready_sources(self, exchange: str) -> dict[str, str]:
        """Source of the newest period of each pair's latest READY row («simulator» / «tester»); empty when the
        store is not configured or busy."""
        store = self._reader()
        if store is None:
            return {}
        try:
            with store.connect(read_only=True) as catalog:
                rows = catalog.execute(
                    "select symbol, payload_json from (select *, row_number() over (partition by symbol order by "
                    "created_utc desc, ready_id desc) as latest from ready where exchange = ?) where latest = 1",
                    [exchange]).fetchall()
        except (MainScreenError, duckdb.Error, OSError):
            return {}
        out = {}
        for symbol, payload in rows:
            entry = json.loads(payload)
            out[symbol] = entry.get("newest_source") or entry.get("latest_source")
        return out

    # -- «Стратегии» (Panel stage 4a): READY structures + tick tests, marks in the registry -----------------------

    def _performance_db(self) -> Path | None:
        try:
            from .performance_v2_store import load_performance_v2_config, performance_v2_database_path

            return performance_v2_database_path(load_performance_v2_config(
                self.config_path.with_name("config.performance.json")))
        except Exception:                        # no PerformanceDB configured: READY rows only
            return None

    def _selection_config(self):
        from .performance_v2_selection import SelectionConfig, load_selection_config

        path = self.config_path.with_name("config.performance.json")
        try:
            return load_selection_config(path) if path.is_file() else SelectionConfig()
        except Exception:
            return SelectionConfig()

    @staticmethod
    def _flag(options: Mapping, key: str, default: bool) -> bool:
        value = options.get(key, default)
        return value in (True, "1", "true", 1) if not isinstance(value, bool) else value

    def _view_strategies(self, exchange: str, options: Mapping, *, metrics: bool = True) -> dict:
        hard, top5 = self._flag(options, "hard", True), self._flag(options, "top5", True)
        hide = self._flag(options, "hide", False)
        try:
            store_root = self._store_root()
        except MainScreenError:
            store_root = None
        from .pair_registry.records import current_finalists

        database = self._settings().root / DATABASE
        marks: list[dict] = []
        if database.is_file():                   # a short connection: the slow PerformanceDB read holds no registry lock
            try:
                with duckdb.connect(str(database), read_only=True) as con:
                    marks = current_finalists(con, exchange)
            except duckdb.Error:
                raise MainScreenError("реестр занят задачей: повторите, когда она закончится") from None
        config = self._selection_config()
        try:
            rows, notes = strategy_rows(exchange=exchange, store_root=store_root, performance_db=self._performance_db(),
                                        marks=marks, config=config, hard=hard, top5=top5, metrics=metrics)
        except duckdb.Error:
            raise MainScreenError("хранилище или PerformanceDB заняты: повторите") from None
        if hide:
            rows = hide_cut(rows)
        return {"name": "strategies", "exchange": exchange,
                "columns": [{"key": key, "label": label} for key, label in VIEWS["strategies"]],
                "default_columns": list(STRATEGY_DEFAULT_COLUMNS), "rows": rows, "notes": notes,
                "options": {"hard": hard, "top5": top5, "hide": hide},
                "top5_share_pct": float(config.top5_share_pct)}

    def finalist(self, payload: Mapping) -> dict:
        """A mark set in the table: written to the registry `finalists` (ADR-0080), source `panel`."""
        from .pair_registry.records import record_finalist

        from .pair_registry.records import current_finalists

        exchange = self._exchange(payload.get("exchange"))
        strategy = payload.get("strategy")
        if not isinstance(strategy, str) or not strategy.strip():
            raise MainScreenError("стратегия: строка")
        strategy = strategy.strip()
        try:
            status, rank = parse_status(payload.get("status")), parse_rank(payload.get("user_rank"))
        except StrategiesError as error:
            raise MainScreenError(str(error)) from None
        row = next((item for item in self._view_strategies(exchange, {}, metrics=False)["rows"]
                    if item["strategy"] == strategy), None)
        if row is None:                          # pair and side come from the view, never from the browser
            raise MainScreenError(f"стратегии {strategy} нет в виде «Стратегии»")
        rank = rank if status != "NONE" else None
        con = self._write()
        try:
            in_force = {item["strategy"]: item for item in current_finalists(con, exchange)}
            if status == "NONE" and strategy not in in_force:
                changed = False                  # no mark to clear: nothing is written
            else:
                changed = record_finalist(con, exchange, str(row["symbol"]), str(row["side"] or "LONG"), strategy,
                                          status, "panel", user_rank=rank)
        except RegistryError as error:
            raise MainScreenError(str(error)) from None
        finally:
            con.close()
        return {"strategy": strategy, "status": None if status == "NONE" else status, "user_rank": rank,
                "changed": changed}

    def import_marks(self, exchange: str, data: bytes) -> dict:
        """«Загрузить отметки»: rows matched by «Стратегия»; «Отметка» and «Ранг» read; only changes written
        (source `excel`); nothing is written when a row is invalid."""
        from openpyxl import load_workbook

        from .pair_registry.records import current_finalists, record_finalist

        exchange = self._exchange(exchange)
        try:
            sheet = load_workbook(BytesIO(data), read_only=True, data_only=True).active
            rows = list(sheet.iter_rows(values_only=True))
        except Exception:
            raise MainScreenError("файл не читается как книга Excel") from None
        if not rows:
            raise MainScreenError("книга пуста")
        header = [str(value).strip() if value is not None else "" for value in rows[0]]
        index = {MARK_LABELS[label]: position for position, label in enumerate(header) if label in MARK_LABELS}
        if not {"strategy", "status"} <= set(index):
            raise MainScreenError("нужны столбцы «Стратегия» и «Отметка»")
        known = {row["strategy"]: row for row in self._view_strategies(exchange, {}, metrics=False)["rows"]}
        planned, unknown, errors, seen = [], [], [], set()
        for number, values in enumerate(rows[1:], start=2):
            name = values[index["strategy"]] if index["strategy"] < len(values) else None
            if name is None or not str(name).strip():
                continue
            name = str(name).strip()
            if name in seen:
                errors.append({"row": number, "strategy": name, "error": "стратегия повторяется в книге"})
                continue
            seen.add(name)
            if name not in known:
                unknown.append(name)
                continue
            try:
                status = parse_status(values[index["status"]] if index["status"] < len(values) else None)
                rank = parse_rank(values[index["user_rank"]]) if "user_rank" in index and index["user_rank"] < len(
                    values) else None
            except StrategiesError as error:
                errors.append({"row": number, "strategy": name, "error": str(error)})
                continue
            planned.append((known[name], status, rank if status != "NONE" else None))
        if errors:
            return {"changed": [], "unchanged": 0, "unknown": unknown, "errors": errors, "written": False}
        con = self._write()
        try:
            in_force = {row["strategy"]: (row["status"], row["user_rank"]) for row in current_finalists(con, exchange)}
            changed, unchanged = [], 0
            con.begin()
            for row, status, rank in planned:
                before = in_force.get(row["strategy"], ("NONE", None))
                if before == (status, rank) or (status == "NONE" and row["strategy"] not in in_force):
                    unchanged += 1
                    continue
                record_finalist(con, exchange, str(row["symbol"]), str(row["side"] or "LONG"), row["strategy"],
                                status, "excel", user_rank=rank)
                changed.append({"strategy": row["strategy"], "from": before[0], "to": status, "user_rank": rank})
            con.commit()
        except BaseException:
            try:
                con.rollback()
            except duckdb.Error:
                pass
            raise
        finally:
            con.close()
        return {"changed": changed, "unchanged": unchanged, "unknown": unknown, "errors": [], "written": True}

    def _view_reference(self, con, exchange: str) -> list[dict]:
        tiers = dict(con.execute(
            "select symbol, count(*) from risk_tiers where exchange = ? and as_of_day = "
            "(select max(as_of_day) from risk_tiers where exchange = ?) group by symbol", [exchange, exchange]).fetchall())
        rows = con.execute("select * from pairs where exchange = ? order by symbol", [exchange]).fetchall()
        names = [column[0] for column in con.description]
        out = []
        for row in rows:
            item = dict(zip(names, row))
            item["listing_day"] = _iso(item.get("listing_ms"))
            item["risk_tiers"] = tiers.get(item["symbol"], 0)
            out.append({key: item.get(key) for key, _label in VIEWS["reference"]})
        return out

    def _view_liquidity(self, con, exchange: str) -> list[dict]:
        rows = con.execute(
            "select * from (select *, row_number() over (partition by symbol order by as_of_day desc) as latest "
            "from liquidity where exchange = ?) where latest = 1 order by symbol", [exchange]).fetchall()
        names = [column[0] for column in con.description]
        out = []
        for row in rows:
            item = dict(zip(names, row))
            item["as_of_day"] = item["as_of_day"].isoformat()
            item["missing_days"] = ", ".join(json.loads(item["missing_days"] or "[]"))
            out.append({key: item.get(key) for key, _label in VIEWS["liquidity"]})
        return out

    def _view_changes(self, con, exchange: str) -> list[dict]:
        rows = con.execute("select changed_ms, symbol, kind, field, old_value, new_value from pair_changes "
                           "where exchange = ? order by seq desc limit 2000", [exchange]).fetchall()
        return [{"changed": datetime.fromtimestamp(row[0] / 1000, timezone.utc).strftime("%Y-%m-%d %H:%M"),
                 "symbol": row[1], "kind": row[2], "field": row[3], "old_value": row[4], "new_value": row[5]}
                for row in rows]

    def view_xlsx(self, name: str, exchange: str, columns: Sequence[str] | None, text: str = "",
                  sort: str | None = None, descending: bool = False,
                  options: Mapping | None = None) -> tuple[str, bytes]:
        """Exactly the rows and columns shown: same filter and sort as the table; no data is changed."""
        from openpyxl import Workbook

        document = self.view(name, exchange, options)
        known = [column["key"] for column in document["columns"]]
        shown = [key for key in (columns or document.get("default_columns") or known) if key in known] or known
        rows = sort_rows(filter_rows(document["rows"], shown, text), sort if sort in shown else None, descending)
        labels = {column["key"]: column["label"] for column in document["columns"]}
        book = Workbook()
        sheet = book.active
        sheet.title = name
        sheet.append([labels[key] for key in shown])
        for row in rows:
            sheet.append([row.get(key) for key in shown])
            for cell in sheet[sheet.max_row]:          # registry text is never written as a formula
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
        buffer = BytesIO()
        book.save(buffer)
        return f"{name}_{exchange}_{self._today().isoformat()}.xlsx", buffer.getvalue()

    def _view_settings_path(self) -> Path:
        return self._settings().root / "view_settings.json"

    def view_settings(self) -> dict:
        path = self._view_settings_path()
        try:
            document = json.loads(path.read_text(encoding="utf-8")) if path.is_file() else {}
        except (OSError, json.JSONDecodeError):
            document = {}
        return {name: document.get(name, {}) for name in VIEWS}

    def save_view_settings(self, payload: Mapping) -> dict:
        name, columns = payload.get("view"), payload.get("columns")
        if not isinstance(name, str) or name not in VIEWS:
            raise MainScreenError(f"неизвестный вид {name!r}")
        known = [key for key, _label in VIEWS[name]]
        if not isinstance(columns, list) or not columns or any(key not in known for key in columns):
            raise MainScreenError("столбцы: непустой список столбцов этого вида")
        document = self.view_settings()
        document[name] = {"columns": list(dict.fromkeys(columns))}
        path = self._view_settings_path()
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_suffix(".tmp")
        temporary.write_text(json.dumps(document, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(path)
        return document

    # -- pair card -------------------------------------------------------------------------------------------

    def pair(self, exchange: str, symbol: str) -> dict:
        exchange = self._exchange(exchange)
        symbol = str(symbol or "").strip().upper()
        con = self._read()
        try:
            row = con.execute("select * from pairs where exchange = ? and symbol = ?", [exchange, symbol]).fetchone()
            if row is None:
                raise MainScreenError(f"пары {symbol} нет в реестре")
            reference = dict(zip([column[0] for column in con.description], row))
            tiers = con.execute("select tier_id, risk_limit_value, initial_margin, maintenance_margin, max_leverage "
                                "from risk_tiers where exchange = ? and symbol = ? and as_of_day = (select max(as_of_day) "
                                "from risk_tiers where exchange = ? and symbol = ?) order by tier_id",
                                [exchange, symbol, exchange, symbol]).fetchall()
            pipeline = con.execute("select * from pipeline where exchange = ? and symbol = ? and side = 'LONG' "
                                   "order by as_of_day desc limit 1", [exchange, symbol]).fetchone()
            pipeline_names = [column[0] for column in con.description]
            tests = con.execute("select kind, structure, window_start, window_end, pnl_pct, maxdd_pct, trades, source "
                                "from pair_tests where exchange = ? and symbol = ? order by window_end desc nulls last, "
                                "source limit 500", [exchange, symbol]).fetchall()
            decisions = con.execute("select side, decision, reason, author, created_ms from decisions where exchange = ? "
                                    "and symbol = ? order by seq desc", [exchange, symbol]).fetchall()
            settings_rows = con.execute("select side, filter, parameter, value, reason, author, created_ms from "
                                        "pair_settings where exchange = ? and symbol = ? order by seq desc",
                                        [exchange, symbol]).fetchall()
            rule, overridden = effective_ready_rule(con, exchange, symbol)
            liquidity = read_liquidity(con, exchange, symbol)
        finally:
            con.close()
        periods, ready = [], None
        try:
            store = self._reader()
            if store is not None:
                periods = [{"start": _iso(segment["start_ms"]), "end": _iso(segment["end_ms"]),
                            "source": segment.get("source")} for segment in store.segments(symbol, exchange)]
                with store.connect(read_only=True) as catalog:
                    found = catalog.execute("select payload_json from ready where exchange = ? and symbol = ? "
                                            "order by created_utc desc, ready_id desc limit 1",
                                            [exchange, symbol]).fetchone()
                ready = json.loads(found[0]) if found else None
        except (MainScreenError, duckdb.Error, OSError):
            pass
        defaults = asdict(ReadyRule())
        listing = reference.get("listing_ms")
        return {
            "exchange": exchange, "symbol": symbol,
            "reference": {key: (value.isoformat() if isinstance(value, date) else value) for key, value in reference.items()},
            "history": sufficiency(listing, self._today()) if isinstance(listing, int) else None,
            "risk_tiers": [dict(zip(("tier", "risk_limit", "initial_margin", "maintenance_margin", "max_leverage"), tier))
                           for tier in tiers],
            "pipeline": ({key: (value.isoformat() if isinstance(value, date) else value)
                          for key, value in zip(pipeline_names, pipeline)} if pipeline else None),
            "periods": periods,
            "ready": ({"status": ready.get("status"), "timeframes": ready.get("timeframes", []),
                       "structures": [{key: item.get(key) for key in ("timeframe", "score", "dd", "category", "pair_rank",
                                                                      "strategy")} for item in ready.get("structures", [])]}
                      if ready else None),
            "liquidity": liquidity,
            "tests": [dict(zip(("kind", "structure", "window_start", "window_end", "pnl_pct", "maxdd_pct", "trades",
                                "source"), test)) for test in tests],
            "decisions": [{"side": side, "decision": decision, "reason": reason, "author": author,
                           "when": datetime.fromtimestamp(ms / 1000, timezone.utc).strftime("%Y-%m-%d %H:%M")}
                          for side, decision, reason, author, ms in decisions],
            "thresholds": [{"filter": "ready_v4", "parameter": parameter, "default": defaults[parameter],
                            "value": getattr(rule, parameter), "overridden": parameter in overridden,
                            "type": "int" if kind is int else "number", "min": low, "max": high}
                           for parameter, (kind, low, high) in FILTERS["ready_v4"].items()],
            "settings_history": [{"side": side, "filter": filt, "parameter": parameter, "value": value, "reason": reason,
                                  "author": author,
                                  "when": datetime.fromtimestamp(ms / 1000, timezone.utc).strftime("%Y-%m-%d %H:%M")}
                                 for side, filt, parameter, value, reason, author, ms in settings_rows],
        }

    def decision(self, payload: Mapping) -> dict:
        if not all(isinstance(payload.get(key, ""), (str, type(None))) for key in
                   ("exchange", "symbol", "side", "decision", "reason", "author")):
            raise MainScreenError("поля решения: строки")
        con = self._write()
        try:
            return record_decision(con, self._exchange(payload.get("exchange")), payload.get("symbol"),
                                   payload.get("side", "ALL"), payload.get("decision"), payload.get("reason"),
                                   payload.get("author"))
        except RegistryError as error:
            raise MainScreenError(str(error)) from None
        finally:
            con.close()

    def setting(self, payload: Mapping) -> dict:
        if not all(isinstance(payload.get(key, ""), (str, int, float, type(None))) for key in
                   ("exchange", "symbol", "side", "filter", "parameter", "value", "reason", "author")):
            raise MainScreenError("поля порога: строки или числа")
        con = self._write()
        try:
            return record_setting(con, self._exchange(payload.get("exchange")), payload.get("symbol"),
                                  payload.get("side", "ALL"), payload.get("filter", "ready_v4"), payload.get("parameter"),
                                  payload.get("value"), payload.get("reason"), payload.get("author"))
        except RegistryError as error:
            raise MainScreenError(str(error)) from None
        finally:
            con.close()

    # -- plans -----------------------------------------------------------------------------------------------

    def plan(self, action: str, exchange: str) -> dict:
        if action not in ACTIONS:
            raise MainScreenError(f"неизвестное действие {action!r}")
        exchange = self._exchange(exchange)
        if exchange == "bingx" and action in ("history", "ready", "tester", "tick"):
            return {"action": action, "exchange": exchange, "items": [], "disabled": BINGX_PENDING}
        builder = getattr(self, f"_plan_{action}")
        return {"action": action, "exchange": exchange, **builder(exchange)}

    def _trading_pairs(self, exchange: str) -> list[tuple[str, int | None]]:
        con = self._read()
        try:
            return con.execute("select symbol, listing_ms from pairs where exchange = ? and status = 'TRADING' "
                               "order by symbol", [exchange]).fetchall()
        finally:
            con.close()

    def _plan_sync(self, exchange: str) -> dict:
        return {"items": [], "note": "Обновить справочник обеих бирж, ликвидность, конвейер и книгу листингов."}

    def _history_end_day(self) -> date:
        from .minute_data import last_archived_day

        return last_archived_day(datetime.now(timezone.utc)) + timedelta(days=1)

    def _plan_history(self, exchange: str) -> dict:
        from .trade_history import pair_segments

        store = self._reader()
        end_day = self._history_end_day()
        items = []
        for symbol, listing_ms in self._trading_pairs(exchange):
            if listing_ms is None:
                continue
            listing = datetime.fromtimestamp(listing_ms / 1000, timezone.utc).date()
            planned = pair_segments(listing, end_day)
            stored = store.segments(symbol, exchange, source="simulator") if store is not None else []
            have = {(_iso(s["start_ms"]), _iso(s["end_ms"])) for s in stored}
            have_end = max((_iso(s["end_ms"]) for s in stored), default=None)
            missing = [segment for segment in planned if (segment[0].isoformat(), segment[1].isoformat()) not in have]
            if not planned:
                label, checked = "мало истории", False
            elif not have and missing:
                label, checked = f"новая пара: {_periods(len(missing), new=False)}", True
            elif missing == planned[-1:]:
                label, checked = "продлить последнее окно", True
            elif missing and any(segment[0].isoformat() < have_end for segment in missing):
                # the P2 grid realigns the last windows when the end day moves
                label, checked = f"пересчитать последние окна: {_periods(len(missing), new=False)}", True
            elif missing:
                label, checked = _periods(len(missing)), True
            else:
                label, checked = "без изменений", False
            items.append({"key": symbol, "label": symbol, "detail": label, "checked": checked})
        return {"items": items, "note": f"Период до {(end_day - timedelta(days=1)).isoformat()} включительно."}

    def _plan_ready(self, exchange: str) -> dict:
        store = self._reader()
        con = self._read()
        try:
            rules = {symbol: effective_ready_rule(con, exchange, symbol)[0] for symbol, _ in
                     con.execute("select symbol, listing_ms from pairs where exchange = ? and status = 'TRADING' "
                                 "order by symbol", [exchange]).fetchall()}
        finally:
            con.close()
        if store is None:
            return {"items": [{"key": symbol, "label": symbol, "detail": "мало периодов (0)", "checked": False}
                              for symbol in rules]}
        with store.connect(read_only=True) as catalog:
            latest = dict(catalog.execute(
                "select symbol, inputs_sha256 from (select *, row_number() over (partition by symbol order by "
                "created_utc desc, ready_id desc) as latest from ready where exchange = ?) where latest = 1",
                [exchange]).fetchall())
        items = []
        for symbol, rule in rules.items():
            segments, newest_source = pair_periods(store, symbol, rule, exchange, "auto")
            if len(segments) < rule.min_periods:
                items.append({"key": symbol, "label": symbol, "detail": f"мало периодов ({len(segments)})",
                              "checked": False})
                continue
            current = inputs_sha256({"periods": [{"metrics_sha256": s["metrics_sha256"]} for s in segments]}, rule)
            changed = latest.get(symbol) != current
            detail = ("READY ещё не собирался" if symbol not in latest else
                      "появились новые периоды" if changed else "без изменений")
            if newest_source == "tester":
                detail += " · последний период из тестера"
            items.append({"key": symbol, "label": symbol, "detail": detail, "checked": changed})
        return {"items": items}

    def _plan_tester(self, exchange: str) -> dict:
        store = self._reader()
        if store is None:
            return {"items": [], "note": "В хранилище ещё нет истории."}
        items = []
        for symbol, _listing in self._trading_pairs(exchange):
            segments = store.segments(symbol, exchange, source="simulator")
            if not segments:
                continue
            newest = segments[-1]
            tested = any((s["start_ms"], s["end_ms"]) == (newest["start_ms"], newest["end_ms"])
                         for s in store.segments(symbol, exchange, source="tester"))
            window = f"{_iso(newest['start_ms'])}..{_iso(newest['end_ms'])}"
            items.append({"key": symbol, "label": symbol,
                          "detail": f"окно {window}: {'уже проверено тестером' if tested else 'не проверено'}",
                          "checked": False})
        return {"items": items, "note": "Выборка точек select-v1 (лучшие 10 % и соседи); план, запуск и импорт."}

    def _plan_tick(self, exchange: str) -> dict:
        from .store_ready_tick import list_ready_folders, resolve_ready_folder

        root = self._store_root()
        folders = [folder for folder in list_ready_folders(root) if folder.get("testable")]
        if not folders:
            return {"items": [], "note": "Нет папок READY: сначала «Пересобрать READY»."}
        manifests = {folder["name"]: resolve_ready_folder(root, folder["name"]) / "ready_manifest.json" for folder in folders}
        name = max(manifests, key=lambda key: (manifests[key].stat().st_mtime, key))   # newest build, whatever its name
        document = json.loads(manifests[name].read_text(encoding="utf-8"))
        items = [{"key": item["strategy"], "label": item["strategy"],
                  "detail": f"{pair['pair']} {item['timeframe']} · {item.get('category')} · score {item.get('score', 0):.1f}",
                  "checked": True}
                 for pair in document["pairs"] if pair.get("status") == "READY" for item in pair.get("structures", [])]
        return {"items": items, "ready": name, "note": f"Папка READY {name}."}

    def _plan_retest(self, exchange: str) -> dict:
        return {"items": [], "options": {"include_reserve": False},
                "note": "Ретест текущих финалистов на новом периоде; опция «включая резерв»."}

    # -- runs ------------------------------------------------------------------------------------------------

    def run(self, action: str, exchange: str, items: Sequence[str], options: Mapping | None = None) -> dict:
        if action not in ACTIONS:
            raise MainScreenError(f"неизвестное действие {action!r}")
        exchange = self._exchange(exchange)
        if exchange == "bingx" and action != "sync":
            raise MainScreenError(BINGX_PENDING)
        if not isinstance(items, list) or any(not isinstance(item, str) or not item.strip() for item in items):
            raise MainScreenError("отметьте строки плана")
        options = dict(options or {})
        if action in ("history", "ready", "tester", "tick") and not items:
            raise MainScreenError("не отмечено ни одной строки")
        if action == "tick":
            return self._run_tick(items, options)
        if action == "retest":
            return self._run_retest(options)
        python, config = self._python, str(self.config_path)
        scripts = self.repo_root / "scripts"
        pairs = ",".join(item.strip().upper() for item in items)
        store_root = str(self._store_root()) if action != "sync" else None
        if action == "sync":
            steps = [("Синхронизировать", [python, "-m", "mrs3.pair_registry", "sync", "--config", config])]
            label, stop_file = "Синхронизировать", None
        elif action == "history":
            steps = [("История", [python, str(scripts / "simulate_history.py"), "--store", store_root, "--config", config,
                                  "--pairs", pairs]),
                     ("Конвейер", [python, "-m", "mrs3.pair_registry", "pipeline", "--config", config])]
            label, stop_file = "Пересчитать историю", None
        elif action == "ready":
            out = Path(store_root) / "ready" / f"{self._today().isoformat()}_ready_{datetime.now(timezone.utc).strftime('%H%M%S')}"
            steps = [("READY", [python, str(scripts / "ready_from_store.py"), "--store", store_root, "--config", config,
                                "--pairs", pairs, "--out", str(out), "--latest-source", "auto"]),
                     ("Конвейер", [python, "-m", "mrs3.pair_registry", "pipeline", "--config", config])]
            label, stop_file = "Пересобрать READY", None
        else:
            tester = str(scripts / "tester_jobs.py")
            steps = [(step, [python, tester, step, "--store", store_root, "--config", config, "--pairs", pairs])
                     for step in ("plan", "run", "import")]
            steps.append(("Конвейер", [python, "-m", "mrs3.pair_registry", "pipeline", "--config", config]))
            label, stop_file = "Проверить тестером", Path(store_root) / "jobs" / ".stop"
        return self._start_chain(action, label, steps, stop_file)

    def _start_chain(self, action: str, label: str, steps, stop_file: Path | None) -> dict:
        folder = Path(tempfile.mkdtemp(prefix="mrs3-chain-", dir=self._chain_root))
        steps_file = folder / "steps.json"
        steps_file.write_text(json.dumps({"steps": [{"name": name, "argv": argv} for name, argv in steps]},
                                         ensure_ascii=False), encoding="utf-8")
        argv = [self._python, str(self.repo_root / "scripts" / "panel_chain.py"), str(steps_file), "--cleanup"]
        if stop_file is not None:
            argv += ["--stop-file", str(stop_file)]
        try:                                        # a started chain deletes the file itself once it has read it
            return self.runner.start(action, label, argv, total=len(steps), stop_file=stop_file)
        except BaseException as error:
            steps_file.unlink(missing_ok=True)
            try:
                folder.rmdir()
            except OSError:
                pass
            if isinstance(error, TaskRunnerError):
                raise MainScreenError(str(error)) from None
            raise

    def _run_tick(self, strategies: Sequence[str], options: Mapping) -> dict:
        """Prepare a batch of the checked strategies of the newest READY folder and start the SINGLE_MODE tick test
        over the chosen window (default: the last 30 days)."""
        if self._controller is None:
            raise MainScreenError("тик-тест недоступен")
        try:
            end = options.get("end_date") or (self._today() - timedelta(days=1)).isoformat()
            start = options.get("start_date") or (date.fromisoformat(end) - timedelta(days=30)).isoformat()
            if date.fromisoformat(start) >= date.fromisoformat(end) or date.fromisoformat(end) >= self._today():
                raise ValueError
        except (TypeError, ValueError):
            raise MainScreenError("окно тик-теста: начало раньше конца, конец не позже вчерашнего дня") from None
        plan = self._plan_tick("bybit")
        if not plan.get("ready"):
            raise MainScreenError(plan.get("note") or "нет папки READY")
        prepared = self._controller.store_ready_prepare({"ready": plan["ready"]}, strategies=list(strategies))
        job = self._controller.panel_job_submit({"kind": "strategies.tester.start", "request": {
            "analysis_run_id": prepared["analysis_run_id"], "start_date": start, "end_date": end,
            "collect_reports": False}})
        return {"tick": {**prepared, "job": job}}

    def _run_retest(self, options: Mapping) -> dict:
        if self._controller is None:
            raise MainScreenError("ретест недоступен")
        include_reserve = options.get("include_reserve", False)
        if type(include_reserve) is not bool:
            raise MainScreenError("«включая резерв»: да или нет")
        job = self._controller.panel_job_submit({"kind": "strategies.performance.v2.finalist-retest.start",
                                                 "request": {"include_reserve": include_reserve}})
        return {"retest": {"job": job}}

    def stop(self) -> dict:
        return self.runner.stop()

    def task(self) -> dict:
        return self.runner.status()
