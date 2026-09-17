from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal
from http.client import HTTPConnection
import json
from pathlib import Path
import subprocess
import threading
from types import SimpleNamespace

import duckdb
import pandas as pd
from openpyxl import Workbook, load_workbook
from io import BytesIO
import pytest

import mrs3.panel as panel_module
from mrs3.panel import PerformanceV2ApiError, PanelController, _control_score, create_panel_server
from mrs3.performance_v2_store import initialize_performance_v2
from mrs3.performance_v2_retest import RetestBatch
from mrs3.performance_v2_finalist_retest import FinalistRetestError, validate_combined_control_workbook


def _controller(tmp_path: Path, *, seed: bool = True) -> PanelController:
    bot = tmp_path / "bot"
    dates = tmp_path / "input" / "dates.xlsx"
    dates.parent.mkdir(parents=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["BTCUSDT", datetime(2026, 2, 1)])
    workbook.save(dates)
    template = tmp_path / "templates" / "base.json"
    template.parent.mkdir()
    template.write_text("{}", encoding="utf-8")
    config = {
        "tester_runner": {
            "bot_root": str(bot), "executable": "hb_c.exe", "base_url": "http://127.0.0.1:80",
            "port": 80, "strategy_dir": "settings_strategy", "report_dir": "tester/report/my_test",
            "wizard_result": "tester/wizard_result.json", "wizard_progress": "tester/wizard_progress.json",
            "tester_config": "tester/tester_config.json", "inbox_root": str(tmp_path / "inbox"),
        },
        "panel_paths": {"performance_db_root": "legacy"},
        "panel_workflow": {
            "listing_dates_path": "input/dates.xlsx",
            "strategy_templates": {"LONG": "templates/base.json", "SHORT": "templates/base.json"},
        },
    }
    config_path = tmp_path / "config.local.json"
    config_path.write_text(json.dumps(config), encoding="utf-8")
    (tmp_path / "config.performance.json").write_text(
        json.dumps({"unified_performance_v2": {"database_root": "performance-v2", "workers": 1}}),
        encoding="utf-8",
    )
    if seed:
        database = tmp_path / "performance-v2" / "strategy_performance.duckdb"
        database.parent.mkdir()
        with duckdb.connect(str(database)) as connection:
            initialize_performance_v2(connection)
            strategy_id = connection.execute(
                """
                insert into strategies (
                    strategy_name, symbol, side, timeframe, close_ma_len, order_count,
                    analysis_run_id, candidate_identity, lifecycle_status, current_result_id,
                    created_at_utc, updated_at_utc
                ) values ('alpha', 'BTCUSDT', 'LONG', '1h', 20, 1, 'run', 'candidate', 'ACTIVE', null, now(), now())
                returning strategy_id
                """
            ).fetchone()[0]
            result_id = connection.execute(
                """
                insert into strategy_results (
                    strategy_id, report_start_utc, report_end_utc, exchange,
                    commission_rate, initial_balance, final_balance, imported_at_utc
                ) values (?, '2026-01-01 00:00:00+00', '2026-01-09 00:00:00+00', 'Bybit', .0004, 100, 101, now())
                returning result_id
                """,
                [strategy_id],
            ).fetchone()[0]
            connection.execute("update strategies set current_result_id = ? where strategy_id = ?", [result_id, strategy_id])
            connection.execute(
                "insert into strategy_tags values (?, 'RETEST', 'TEST', 'fixture', now())", [strategy_id]
            )
    return PanelController(tmp_path, config_path)


def test_control_score_drops_missing_and_non_finite_values() -> None:
    assert [_control_score(value) for value in (
        None, pd.NA, float("nan"), float("inf"), float("-inf"),
        Decimal("NaN"), Decimal("Infinity"), Decimal("-Infinity"),
    )] == [None] * 8
    assert _control_score(Decimal("5.5")) == 5.5


def test_finalist_control_candidate_workbook_normalizes_rank_values(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller = _controller(tmp_path, seed=False)
    captured: dict[str, object] = {}

    def capture_writer(result, path, request, review_metadata, user_review_rows):
        captured["result"] = result.copy()
        captured["reviews"] = user_review_rows
        path.write_bytes(b"candidate-workbook")
        return path

    monkeypatch.setattr(panel_module, "write_selection_workbook", capture_writer)
    controller._finalist_control_candidate_workbook(
        [{
            "strategy_id": 1, "result_id": 11, "strategy_name": "alpha", "symbol": "BTCUSDT", "side": "LONG",
            "timeframe": "1h", "order_count": 1, "close_ma_len": 20, "auto_status": "FINALIST", "finalist": True,
            "user_status": "FINALIST", "user_rank": pd.NA, "auto_rank": float("inf"), "final_rank": float("-inf"),
        }],
        panel_module.SelectionRequest("BTCUSDT", "LONG", ()),
    )

    result = captured["result"].iloc[0]
    assert result["user_rank"] is None
    assert result["auto_rank"] is None
    assert result["final_rank"] is None
    assert captured["reviews"][1]["user_rank"] is None


def test_current_control_export_is_read_only_until_review_and_preserves_outside_scope(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    controller = _controller(tmp_path)
    database = tmp_path / "performance-v2" / "strategy_performance.duckdb"
    now = datetime(2026, 1, 1, tzinfo=timezone.utc)
    with duckdb.connect(str(database)) as connection:
        beta_id = connection.execute(
            """insert into strategies (strategy_name, symbol, side, timeframe, close_ma_len, order_count,
               analysis_run_id, candidate_identity, lifecycle_status, created_at_utc, updated_at_utc)
               values ('beta', 'BTCUSDT', 'LONG', '1h', 20, 1, 'run', 'beta', 'ACTIVE', ?, ?)
               returning strategy_id""", [now, now],
        ).fetchone()[0]
        beta_result_id = connection.execute(
            """insert into strategy_results (strategy_id, report_start_utc, report_end_utc, exchange,
               commission_rate, initial_balance, final_balance, imported_at_utc)
               values (?, '2026-01-01', '2026-01-09', 'Bybit', .0004, 100, 101, ?)
               returning result_id""", [beta_id, now],
        ).fetchone()[0]
        previous_alpha_result_id = connection.execute("select current_result_id from strategies where strategy_name = 'alpha'").fetchone()[0]
        connection.execute(
            """insert into strategies (strategy_name, symbol, side, timeframe, close_ma_len, order_count,
               analysis_run_id, candidate_identity, lifecycle_status, created_at_utc, updated_at_utc)
               values ('old-alpha', 'BTCUSDT', 'LONG', '1h', 20, 1, 'old-run', 'old-alpha', 'DISCARDED', ?, ?)""",
            [now, now],
        )
        old_holder_id = connection.execute("select strategy_id from strategies where strategy_name = 'old-alpha'").fetchone()[0]
        connection.execute("update strategies set current_result_id = null where strategy_id = 1")
        connection.execute("update strategy_results set strategy_id = ? where result_id = ?", [old_holder_id, previous_alpha_result_id])
        alpha_result_id = connection.execute(
            """insert into strategy_results (strategy_id, report_start_utc, report_end_utc, exchange,
               commission_rate, initial_balance, final_balance, imported_at_utc)
               values (1, '2026-01-02', '2026-01-10', 'Bybit', .0004, 100, 102, ?)
               returning result_id""", [now],
        ).fetchone()[0]
        connection.execute("update strategies set current_result_id = ? where strategy_id = 1", [alpha_result_id])
        connection.execute("update strategies set current_result_id = ? where strategy_id = ?", [beta_result_id, beta_id])
        instance = connection.execute("select value from schema_info where key = 'database_instance_id'").fetchone()[0]
        connection.execute(
            """insert into selection_runs (selection_run_id, database_instance_id, symbol, side, selection_contract_version,
               request_json, request_sha256, config_json, config_sha256, candidate_count, representative_count,
               auto_finalist_count, top_n, workbook_sha256, created_at_utc)
               values ('ordinary', ?, 'BTCUSDT', 'LONG', 'v1', '{}', ?, '{}', ?, 2, 2, 1, 20, ?, ?)""",
            [instance, "r" * 64, "s" * 64, "w" * 64, now],
        )
        connection.executemany(
            """insert into selection_results (selection_run_id, strategy_id, result_id_at_selection, auto_status,
               auto_score, auto_rank, prior_rejected, stage_trace_json) values ('ordinary', ?, ?, ?, ?, ?, false, '{}')""",
            # Strategy IDs deliberately run opposite to rank/score order.
            [(1, alpha_result_id, "FINALIST", 80, 2), (beta_id, beta_result_id, "RESERVE", 90, 1)],
        )

    def effective() -> dict[int, tuple[object, object, object]]:
        with duckdb.connect(str(database), read_only=True) as connection:
            from mrs3.performance_v2_selection_review import effective_selection_decisions
            return {
                int(strategy_id): decision
                for strategy_id, decision in effective_selection_decisions(connection).items()
            }

    before = effective()
    assert before == {}
    with pytest.raises(FinalistRetestError, match="there are no current effective finalists"):
        controller.strategies_performance_v2_finalist_retest_export({"include_reserve": True})
    with pytest.raises(FinalistRetestError, match="there are no current effective finalists"):
        controller.strategies_performance_v2_finalist_retest_export({"include_reserve": False})
    with duckdb.connect(str(database), read_only=True) as connection:
        assert connection.execute("select count(*) from selection_review_imports").fetchone() == (0,)
        assert connection.execute("select count(*) from selection_review_rows").fetchone() == (0,)
        assert connection.execute("select count(*) from selection_runs").fetchone() == (1,)
    with duckdb.connect(str(database)) as connection:
        connection.execute(
            """insert into selection_review_imports
               (review_import_id, selection_run_id, workbook_sha256, imported_at_utc, row_count)
               values ('reviewed', 'ordinary', ?, ?, 2)""", ["a" * 64, now],
        )
        connection.executemany(
            """insert into selection_review_rows
               (review_import_id, strategy_id, user_status, user_rank, user_analog_of_strategy_id, comment)
               values ('reviewed', ?, ?, ?, null, null)""",
            [(1, "FINALIST", 1), (beta_id, "RESERVE", None)],
        )
    before = effective()
    assert before[1][:2] == ("FINALIST", 1)
    assert before[beta_id][:2] == ("RESERVE", None)
    with duckdb.connect(str(database), read_only=True) as connection:
        assert connection.execute("select count(*) from window_metrics").fetchone() == (0,)
    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        http = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        http.request("GET", "/api/v2/strategies/performance-v2/finalist-retest/export?include_reserve=true")
        response = http.getresponse()
        body = response.read()
        content_type = response.getheader("Content-Type", "")
        http.close()
        assert response.status == 409
        assert content_type.startswith("application/json")
        assert not body.startswith(b"PK")
        document = json.loads(body.decode("utf-8"))
        assert document["error"]["code"] == "SELECTION_CACHE_INCOMPLETE"
        assert any(token in document["error"]["message"].casefold() for token in ("cache", "prepare", "recalculate"))
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)
    with pytest.raises(FinalistRetestError, match="prepare.*cache|cache.*prepare"):
        controller.strategies_performance_v2_finalist_retest_export({"include_reserve": True})
    assert controller.strategies_performance_v2_recalculate({"symbol": "BTCUSDT", "side": "LONG"}) == {"status": "READY"}
    reserve_filename, reserve_issued = controller.strategies_performance_v2_finalist_retest_export({"include_reserve": True})
    assert reserve_filename == "performance-v2-current-finalists-with-reserve.xlsx"
    reserve_sheet = load_workbook(BytesIO(reserve_issued))["Candidates"]
    assert reserve_sheet.max_row == 3
    reserve_headers = {cell.value: cell.column for cell in reserve_sheet[1]}
    assert {reserve_sheet.cell(row, reserve_headers["ID"]).value for row in range(2, reserve_sheet.max_row + 1)} == {1, beta_id}
    reserve_rows = {
        reserve_sheet.cell(row, reserve_headers["ID"]).value: row
        for row in range(2, reserve_sheet.max_row + 1)
    }
    assert reserve_sheet.cell(reserve_rows[1], reserve_headers["Auto Rank"]).value == 2
    assert reserve_sheet.cell(reserve_rows[beta_id], reserve_headers["Auto Rank"]).value == 1
    assert reserve_sheet.cell(reserve_rows[1], reserve_headers["Final score (Pair+Side)"]).value == 80
    assert reserve_sheet.cell(reserve_rows[beta_id], reserve_headers["Final score (Pair+Side)"]).value == 90
    unedited = controller.strategies_performance_v2_finalist_retest_control_import(reserve_issued)
    assert unedited["group_count"] == 1 and unedited["row_count"] == 2
    fresh_filename, fresh_issued = controller.strategies_performance_v2_finalist_retest_export({"include_reserve": True})
    assert fresh_filename == "performance-v2-current-finalists-with-reserve.xlsx"
    fresh_workbook = load_workbook(BytesIO(fresh_issued))
    fresh_sheet = fresh_workbook["Candidates"]
    fresh_headers = {cell.value: cell.column for cell in fresh_sheet[1]}
    fresh_beta_row = next(
        row for row in range(2, fresh_sheet.max_row + 1)
        if fresh_sheet.cell(row, fresh_headers["ID"]).value == beta_id
    )
    fresh_sheet.cell(fresh_beta_row, fresh_headers["Comment"]).value = "fresh allowed edit"
    edited_io = BytesIO()
    fresh_workbook.save(edited_io)
    edited = controller.strategies_performance_v2_finalist_retest_control_import(edited_io.getvalue())
    assert edited["group_count"] == 1 and edited["row_count"] == 2
    filename, issued = controller.strategies_performance_v2_finalist_retest_export({"include_reserve": False})
    assert filename == "performance-v2-current-finalists.xlsx"
    assert {strategy_id: decision[:2] for strategy_id, decision in effective().items()} == {
        strategy_id: decision[:2] for strategy_id, decision in before.items()
    }
    with duckdb.connect(str(database), read_only=True) as connection:
        assert connection.execute("select count(*) from selection_runs").fetchone() == (4,)

    workbook = load_workbook(BytesIO(issued))
    sheet = workbook["Candidates"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.max_row == 2 and sheet.cell(2, headers["ID"]).value == 1
    assert sheet.cell(2, headers["Result ID"]).value == alpha_result_id
    assert {sheet.cell(row, headers["ID"]).value for row in range(2, sheet.max_row + 1)} == {1}
    assert {"PnL/30", "DD", "Lots", "Close", "User Status", "User Rank", "RETEST", "Analog Of ID", "Comment"}.issubset(headers)
    assert "Pair" not in headers
    assert any(validation.formula1 == '"FINALIST,RESERVE,ANALOG,FILTERED,REJECTED"' for validation in sheet.data_validations.dataValidation)
    forbidden_refs = ("Finalists", "All candidates", "_MRS_SELECTION_META")
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                assert cell.data_type != "f"
                if cell.data_type == "f":
                    assert not any(reference in str(cell.value) for reference in forbidden_refs)
        for validation in worksheet.data_validations.dataValidation:
            assert not any(reference in str(validation.formula1) for reference in forbidden_refs)
            assert not any(reference in str(validation.sqref) for reference in forbidden_refs)
        for sqref in worksheet.conditional_formatting:
            assert not any(reference in str(sqref) for reference in forbidden_refs)
            for rule in worksheet.conditional_formatting[sqref]:
                for formula in rule.formula or ():
                    assert not any(forbidden_reference in str(formula) for forbidden_reference in forbidden_refs)
    assert not workbook.defined_names
    with duckdb.connect(str(database), read_only=True) as connection:
        config = panel_module.load_selection_config(controller.default_config.with_name("config.performance.json"))
        request = panel_module.SelectionRequest("BTCUSDT", "LONG", ())
        ordinary = panel_module.load_selection_candidates(connection, request, config, cache_only=True)
    ordinary = ordinary.copy()
    ordinary = ordinary.loc[ordinary["strategy_id"] == 1].copy()
    ordinary["auto_status"] = "FINALIST"
    ordinary["finalist"] = True
    ordinary["final_rank"] = 1
    ordinary["final_score"] = 90
    ordinary_path = panel_module.write_selection_workbook(
        ordinary, tmp_path / "ordinary-candidates.xlsx", request, {"workbook_schema_version": "2"},
        {1: {"user_status": "FINALIST", "user_rank": 1, "user_analog_of_strategy_id": None, "comment": None}},
    )
    ordinary_sheet = load_workbook(ordinary_path)["All candidates"]
    ordinary_headers = [cell.value for cell in ordinary_sheet[1]]
    assert all(isinstance(header, str) and header.strip() for header in ordinary_headers)
    assert len(ordinary_headers) == len(set(ordinary_headers))
    assert list(headers) == ordinary_headers
    for name in ("Close", "Lots", "User Status"):
        current_cell = sheet.cell(1, headers[name])
        ordinary_cell = ordinary_sheet.cell(1, ordinary_headers.index(name) + 1)
        assert current_cell.style_id == ordinary_cell.style_id
        assert sheet.column_dimensions[current_cell.column_letter].hidden == ordinary_sheet.column_dimensions[ordinary_cell.column_letter].hidden
        assert sheet.cell(2, headers[name]).number_format == ordinary_sheet.cell(2, ordinary_headers.index(name) + 1).number_format
    assert {(validation.formula1, str(validation.sqref)) for validation in sheet.data_validations.dataValidation} == {
        (validation.formula1, str(validation.sqref)) for validation in ordinary_sheet.data_validations.dataValidation
    }
    original_loader = panel_module.load_selection_candidates

    def mismatched_loader(*args, **kwargs):
        loaded = original_loader(*args, **kwargs)
        loaded = loaded.copy()
        loaded.loc[:, "result_id"] = 999999
        return loaded

    with monkeypatch.context() as context:
        context.setattr(panel_module, "load_selection_candidates", mismatched_loader)
        with pytest.raises(FinalistRetestError) as raised:
            controller.strategies_performance_v2_finalist_retest_export({"include_reserve": False})
    assert raised.value.code == "CONTROL_CANONICAL_ROW_MISMATCH"
    assert "BTCUSDT" in str(raised.value) and "LONG" in str(raised.value)
    assert "recalculate" in str(raised.value)
    def missing_loader(*args, **kwargs):
        loaded = original_loader(*args, **kwargs)
        return loaded.iloc[0:0].copy()

    with monkeypatch.context() as context:
        context.setattr(panel_module, "load_selection_candidates", missing_loader)
        with pytest.raises(FinalistRetestError) as missing:
            controller.strategies_performance_v2_finalist_retest_export({"include_reserve": False})
    assert missing.value.code == "CONTROL_CANONICAL_ROW_MISMATCH"
    assert "BTCUSDT" in str(missing.value) and "LONG" in str(missing.value)
    assert "recalculate" in str(missing.value)
    sheet.cell(2, headers["User Status"]).value = "REJECTED"
    sheet.cell(2, headers["User Rank"]).value = None
    sheet.cell(2, headers["RETEST"]).value = "RETEST"
    edited_io = BytesIO()
    workbook.save(edited_io)
    imported = controller.strategies_performance_v2_finalist_retest_control_import(edited_io.getvalue())
    assert imported["group_count"] == imported["row_count"] == 1
    after = effective()
    assert after[1][0:2] == ("REJECTED", None)
    assert after[beta_id][0:2] == before[beta_id][0:2]
    replay = controller.strategies_performance_v2_finalist_retest_control_import(edited_io.getvalue())
    assert replay["review_import_ids"] == imported["review_import_ids"]
    with duckdb.connect(str(database)) as connection:
        current_run_id = connection.execute(
            "select selection_run_id from selection_runs where request_json like '%CURRENT_EFFECTIVE%' order by created_at_utc desc limit 1"
        ).fetchone()[0]
        connection.execute("update selection_runs set request_json = '{}' where selection_run_id = ?", [current_run_id])
        with pytest.raises(FinalistRetestError, match="CONTROL_SCOPE_MISMATCH"):
            controller.strategies_performance_v2_finalist_retest_control_import(edited_io.getvalue())


def test_current_control_export_import_keeps_multiple_pair_side_groups_local(tmp_path: Path) -> None:
    controller = _controller(tmp_path)
    database = tmp_path / "performance-v2" / "strategy_performance.duckdb"
    now = datetime(2026, 1, 2, tzinfo=timezone.utc)
    with duckdb.connect(str(database)) as connection:
        eth_id = connection.execute(
            """insert into strategies (strategy_name, symbol, side, timeframe, close_ma_len, order_count,
               analysis_run_id, candidate_identity, lifecycle_status, current_result_id, created_at_utc, updated_at_utc)
               values ('eth', 'ETHUSDT', 'SHORT', '1h', 7, 1, 'run-eth', 'eth', 'ACTIVE', null, ?, ?)
               returning strategy_id""", [now, now],
        ).fetchone()[0]
        eth_result_id = connection.execute(
            """insert into strategy_results (strategy_id, report_start_utc, report_end_utc, exchange,
               commission_rate, initial_balance, final_balance, imported_at_utc)
               values (?, '2026-01-01', '2026-01-09', 'Bybit', .0004, 100, 103, ?)
               returning result_id""", [eth_id, now],
        ).fetchone()[0]
        connection.execute("update strategies set current_result_id = ? where strategy_id = ?", [eth_result_id, eth_id])
        instance = connection.execute("select value from schema_info where key = 'database_instance_id'").fetchone()[0]
        alpha_result_id = connection.execute("select current_result_id from strategies where strategy_id = 1").fetchone()[0]
        for run_id, strategy_id, result_id, symbol, side, auto_score, auto_rank, user_rank, auto_reason in (
                ("ordinary-btc", 1, alpha_result_id, "BTCUSDT", "LONG", Decimal("90"), 1, 1, None),
                ("ordinary-eth", eth_id, eth_result_id, "ETHUSDT", "SHORT", Decimal("80.523456789123456789"), 2, 2, "LOT_VARIANT_REDUNDANT"),
        ):
            connection.execute(
                """insert into selection_runs (selection_run_id, database_instance_id, symbol, side, selection_contract_version,
                   request_json, request_sha256, config_json, config_sha256, candidate_count, representative_count,
                   auto_finalist_count, top_n, workbook_sha256, created_at_utc)
                   values (?, ?, ?, ?, 'v1', '{}', ?, '{}', ?, 1, 1, 1, 20, ?, ?)""",
                [run_id, instance, symbol, side, "r" * 64, "s" * 64, "w" * 64, now],
            )
            connection.execute(
                """insert into selection_results (selection_run_id, strategy_id, result_id_at_selection, auto_status,
                   auto_score, auto_rank, auto_reason, prior_rejected, stage_trace_json)
                   values (?, ?, ?, 'FINALIST', ?, ?, ?, false, '{}')""",
                [run_id, strategy_id, result_id, auto_score, auto_rank, auto_reason],
            )
            connection.execute(
                """insert into selection_review_imports
                   (review_import_id, selection_run_id, workbook_sha256, imported_at_utc, row_count)
                   values (?, ?, ?, ?, 1)""", [f"review-{run_id}", run_id, ("a" if strategy_id == 1 else "b") * 64, now],
            )
            connection.execute(
                """insert into selection_review_rows
                   (review_import_id, strategy_id, user_status, user_rank, user_analog_of_strategy_id, comment)
                   values (?, ?, 'FINALIST', ?, null, null)""", [f"review-{run_id}", strategy_id, user_rank],
            )

    controller.strategies_performance_v2_recalculate({"symbol": "BTCUSDT", "side": "LONG"})
    controller.strategies_performance_v2_recalculate({"symbol": "ETHUSDT", "side": "SHORT"})
    filename, issued = controller.strategies_performance_v2_finalist_retest_export({"include_reserve": False})
    assert filename == "performance-v2-current-finalists.xlsx"
    metadata, candidates, groups, failures = validate_combined_control_workbook(issued)
    assert len(candidates) == 2 and len(groups) == 2 and not failures
    assert {(row["Pair"], row["Direction"]) for row in candidates} == {("BTCUSDT", "LONG"), ("ETHUSDT", "SHORT")}
    candidates_by_group = {(row["Pair"], row["Direction"]): row for row in candidates}
    assert candidates_by_group[("BTCUSDT", "LONG")]["Score"] == 90
    assert candidates_by_group[("BTCUSDT", "LONG")]["Auto Reason"] is None
    assert candidates_by_group[("ETHUSDT", "SHORT")]["Score"] == pytest.approx(80.523456789123456789)
    assert candidates_by_group[("ETHUSDT", "SHORT")]["Auto Rank"] == 2
    assert candidates_by_group[("ETHUSDT", "SHORT")]["User Rank"] == 2
    assert candidates_by_group[("ETHUSDT", "SHORT")]["Auto Reason"] == "LOT_VARIANT_REDUNDANT"
    assert set(json.loads(str(metadata["exact_rowsets_json"]))) == {"BTCUSDT|LONG", "ETHUSDT|SHORT"}
    unedited = controller.strategies_performance_v2_finalist_retest_control_import(issued)
    assert unedited["group_count"] == unedited["row_count"] == 2

    workbook = load_workbook(BytesIO(issued))
    sheet = workbook["Candidates"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    rows_by_group = {
        (sheet.cell(row, headers["Пара"]).value, sheet.cell(row, headers["Side"]).value): row
        for row in range(2, sheet.max_row + 1)
    }
    assert set(rows_by_group) == {("BTCUSDT", "LONG"), ("ETHUSDT", "SHORT")}
    sheet.cell(rows_by_group[("BTCUSDT", "LONG")], headers["User Status"]).value = "REJECTED"
    sheet.cell(rows_by_group[("BTCUSDT", "LONG")], headers["User Rank"]).value = None
    sheet.cell(rows_by_group[("ETHUSDT", "SHORT")], headers["User Status"]).value = "RESERVE"
    sheet.cell(rows_by_group[("ETHUSDT", "SHORT")], headers["User Rank"]).value = 1
    edited_io = BytesIO()
    workbook.save(edited_io)
    imported = controller.strategies_performance_v2_finalist_retest_control_import(edited_io.getvalue())
    assert imported["group_count"] == imported["row_count"] == 2
    with duckdb.connect(str(database), read_only=True) as connection:
        from mrs3.performance_v2_selection_review import effective_selection_decisions

        decisions = effective_selection_decisions(connection)
    assert decisions[1][:2] == ("REJECTED", None)
    assert decisions[eth_id][:2] == ("RESERVE", 1)


def test_retest_status_is_db_authoritative_and_defaults_from_current_result(tmp_path: Path) -> None:
    controller = _controller(tmp_path)

    status = controller.strategies_performance_v2_retest_status()

    assert status["count"] == status["retest_count"] == status["active_count"] == 1
    assert status["default_start"] == "2026-01-01"
    assert status["default_end"] == "2026-01-09"


def test_retest_status_is_safe_when_database_is_missing_or_empty(tmp_path: Path) -> None:
    missing = _controller(tmp_path / "missing", seed=False)
    assert missing.strategies_performance_v2_retest_status() == {
        "count": 0, "retest_count": 0, "active_count": 0, "phase": "IDLE"
    }

    empty_root = tmp_path / "empty"
    empty = _controller(empty_root, seed=False)
    database = empty_root / "performance-v2" / "strategy_performance.duckdb"
    database.parent.mkdir()
    with duckdb.connect(str(database)) as connection:
        initialize_performance_v2(connection)
    status = empty.strategies_performance_v2_retest_status()
    assert status["count"] == 0 and status["phase"] == "IDLE"
    assert status["default_start"] is None and status["default_end"] is None


def test_metadata_retest_inbox_resolves_relative_artifacts_from_runner_dirs(tmp_path: Path) -> None:
    controller = _controller(tmp_path)
    bot_root = tmp_path / "bot"
    report_dir = bot_root / "tester" / "report" / "my_test"
    strategy_dir = bot_root / "settings_strategy"
    report_dir.mkdir(parents=True)
    strategy_dir.mkdir(parents=True)
    (report_dir / "alpha.html").write_text("<html></html>", encoding="utf-8")
    (strategy_dir / "alpha.json").write_text("{}", encoding="utf-8")
    inbox = tmp_path / "inbox" / "retest-relative"
    inbox.mkdir(parents=True)
    (inbox / "inbox_manifest.json").write_text(json.dumps({
        "schema_version": 1,
        "run_mode": "SINGLE_MODE",
        "source_mode": "metadata_only",
        "inbox_ready": True,
        "expected_strategy_names": ["alpha"],
        "entries": [{"strategy_name": "alpha", "strategy_path": "alpha.json", "report_path": "alpha.html"}],
    }), encoding="utf-8")

    controller._validate_metadata_inbox(inbox)


def test_metadata_retest_inbox_accepts_project_strategy_source(tmp_path: Path) -> None:
    controller = _controller(tmp_path)
    report_dir = tmp_path / "bot" / "tester" / "report" / "my_test"
    report_dir.mkdir(parents=True)
    (report_dir / "alpha.html").write_text("<html></html>", encoding="utf-8")
    project_strategy = tmp_path / "Output" / "strategies" / "alpha.json"
    project_strategy.parent.mkdir(parents=True)
    project_strategy.write_text("{}", encoding="utf-8")
    inbox = tmp_path / "inbox" / "retest-project-source"
    inbox.mkdir(parents=True)
    (inbox / "inbox_manifest.json").write_text(json.dumps({
        "schema_version": 1,
        "run_mode": "SINGLE_MODE",
        "source_mode": "metadata_only",
        "inbox_ready": True,
        "expected_strategy_names": ["alpha"],
        "entries": [{
            "strategy_name": "alpha",
            "strategy_path": "alpha.json",
            "report_path": "alpha.html",
        }],
    }), encoding="utf-8")

    controller._validate_metadata_inbox(inbox)


def test_metadata_retest_inbox_resolves_relative_strategy_from_performance_root(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    controller = _controller(tmp_path)
    report_dir = tmp_path / "bot" / "tester" / "report" / "my_test"
    report_dir.mkdir(parents=True)
    (report_dir / "alpha.html").write_text("<html></html>", encoding="utf-8")
    project_strategy = tmp_path / "Output" / "strategies" / "alpha.json"
    project_strategy.parent.mkdir(parents=True)
    project_strategy.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(
        controller, "_performance_v2_config", lambda: SimpleNamespace(strategy_root=Path("Output/strategies"))
    )
    monkeypatch.chdir(tmp_path.parent)
    inbox = tmp_path / "inbox" / "retest-relative-project-source"
    inbox.mkdir(parents=True)
    (inbox / "inbox_manifest.json").write_text(json.dumps({
        "schema_version": 1,
        "run_mode": "SINGLE_MODE",
        "source_mode": "metadata_only",
        "inbox_ready": True,
        "expected_strategy_names": ["alpha"],
        "entries": [{"strategy_name": "alpha", "strategy_path": "alpha.json", "report_path": "alpha.html"}],
    }), encoding="utf-8")

    controller._validate_metadata_inbox(inbox)


@pytest.mark.parametrize("performance_config", [None, "{invalid"])
def test_metadata_retest_inbox_is_safe_without_valid_performance_config(
    tmp_path: Path, performance_config: str | None
) -> None:
    controller = _controller(tmp_path)
    report_dir = tmp_path / "bot" / "tester" / "report" / "my_test"
    strategy_dir = tmp_path / "bot" / "settings_strategy"
    report_dir.mkdir(parents=True)
    strategy_dir.mkdir(parents=True)
    (report_dir / "alpha.html").write_text("<html></html>", encoding="utf-8")
    (strategy_dir / "alpha.json").write_text("{}", encoding="utf-8")
    inbox = tmp_path / "inbox" / "retest-without-performance-config"
    inbox.mkdir(parents=True)
    (inbox / "inbox_manifest.json").write_text(json.dumps({
        "schema_version": 1,
        "run_mode": "SINGLE_MODE",
        "source_mode": "metadata_only",
        "inbox_ready": True,
        "expected_strategy_names": ["alpha"],
        "entries": [{"strategy_name": "alpha", "strategy_path": "alpha.json", "report_path": "alpha.html"}],
    }), encoding="utf-8")
    performance_path = tmp_path / "config.performance.json"
    if performance_config is None:
        performance_path.unlink()
    else:
        performance_path.write_text(performance_config, encoding="utf-8")

    controller._validate_metadata_inbox(inbox)


def test_metadata_retest_inbox_rejects_strategy_outside_both_configured_roots(tmp_path: Path) -> None:
    controller = _controller(tmp_path)
    report_dir = tmp_path / "bot" / "tester" / "report" / "my_test"
    report_dir.mkdir(parents=True)
    (report_dir / "alpha.html").write_text("<html></html>", encoding="utf-8")
    outside_strategy = tmp_path / "not-a-configured-strategy-root" / "alpha.json"
    outside_strategy.parent.mkdir(parents=True)
    outside_strategy.write_text("{}", encoding="utf-8")
    inbox = tmp_path / "inbox" / "retest-outside-strategy"
    inbox.mkdir(parents=True)
    (inbox / "inbox_manifest.json").write_text(json.dumps({
        "schema_version": 1,
        "run_mode": "SINGLE_MODE",
        "source_mode": "metadata_only",
        "inbox_ready": True,
        "expected_strategy_names": ["alpha"],
        "entries": [{
            "strategy_name": "alpha",
            "strategy_path": str(outside_strategy),
            "report_path": "alpha.html",
        }],
    }), encoding="utf-8")

    with pytest.raises(ValueError, match="outside configured directory"):
        controller._validate_metadata_inbox(inbox)


def test_metadata_retest_inbox_rejects_strategy_symlink_to_outside_root(tmp_path: Path) -> None:
    controller = _controller(tmp_path)
    report_dir = tmp_path / "bot" / "tester" / "report" / "my_test"
    strategy_dir = tmp_path / "bot" / "settings_strategy"
    report_dir.mkdir(parents=True)
    strategy_dir.mkdir(parents=True)
    (report_dir / "alpha.html").write_text("<html></html>", encoding="utf-8")
    outside_strategy = tmp_path / "outside-alpha.json"
    outside_strategy.write_text("{}", encoding="utf-8")
    try:
        (strategy_dir / "alpha.json").symlink_to(outside_strategy)
    except (OSError, NotImplementedError):
        pytest.skip("symlink creation is unavailable")
    inbox = tmp_path / "inbox" / "retest-symlink-strategy"
    inbox.mkdir(parents=True)
    (inbox / "inbox_manifest.json").write_text(json.dumps({
        "schema_version": 1,
        "run_mode": "SINGLE_MODE",
        "source_mode": "metadata_only",
        "inbox_ready": True,
        "expected_strategy_names": ["alpha"],
        "entries": [{"strategy_name": "alpha", "strategy_path": "alpha.json", "report_path": "alpha.html"}],
    }), encoding="utf-8")

    with pytest.raises(ValueError, match="strategy_path is missing"):
        controller._validate_metadata_inbox(inbox)


@pytest.mark.parametrize(
    "payload",
    [
        {"test_start": "2026-01-01", "test_end": "2026-01-01"},
        {"test_start": "2026-02-30", "test_end": "2026-03-01"},
        {"test_start": "2026-01-01"},
        {"test_start": "2026-01-01", "test_end": "2026-01-02", "start_date": "2026-01-01", "end_date": "2026-01-02"},
    ],
)
def test_retest_start_rejects_invalid_or_mixed_date_contract(tmp_path: Path, payload: dict[str, str]) -> None:
    controller = _controller(tmp_path)
    with duckdb.connect(":memory:") as connection:
        initialize_performance_v2(connection)
        with pytest.raises(ValueError):
            controller._retest_range(payload, connection)


def test_retest_start_uses_native_single_mode_and_keeps_database_read_only(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    controller = _controller(tmp_path)
    database = tmp_path / "performance-v2" / "strategy_performance.duckdb"
    before = database.read_bytes()
    outside_inbox = tmp_path / "outside-inbox"
    outside_inbox.mkdir()
    controller._panel_jobs.submit(
        "strategies.tester.native.start", {"retest": True}, "panel:outside-inbox",
        ("strategies.tester",), job_id="outside-inbox",
    )
    controller._panel_jobs.transition("outside-inbox", "RUNNING")
    controller._panel_jobs.sync(
        "outside-inbox",
        {"state": "COMMITTED", "phase": "COMMITTED", "inbox_ready": True},
        runtime={"retest": True, "inbox_path": str(outside_inbox)},
    )
    batch = RetestBatch("batch-1", tmp_path / "output" / "strategies", tmp_path / "output" / "strategy_manifest.json", 1)
    seen: dict[str, object] = {}

    monkeypatch.setattr(panel_module, "build_retest_manifest", lambda *args: batch)

    class FakeTester:
        def start(self, manifest_path, *, analysis_run_id, start_date, end_date, job_id):
            seen.update(locals())
            return {"job_id": job_id, "state": "RUNNING", "phase": "RUNNING"}

    monkeypatch.setattr(controller, "_single_mode_strategy_test", lambda: FakeTester())
    result = controller.strategies_performance_v2_retest_start({})

    assert result["state"] == "RUNNING"
    assert result["job_id"] != "outside-inbox"
    assert seen["analysis_run_id"] == "batch-1"
    assert seen["start_date"] == "2026-01-01"
    assert seen["end_date"] == "2026-01-09"
    assert controller._panel_jobs.runtime(result["job_id"])["mode"] == "SINGLE_MODE"
    runtime = controller._panel_jobs.runtime(result["job_id"])
    assert runtime["retest"] is True and runtime["manifest_path"].endswith("strategy_manifest.json")
    assert database.read_bytes() == before


@pytest.mark.parametrize("payload", [{"unknown": True}, {"test_start": "2026-01-01"}])
def test_retest_start_validates_payload_before_reusing_committed_inbox(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, payload: dict[str, object]
) -> None:
    controller = _controller(tmp_path)
    _committed_retest_job(controller, tmp_path, job_id="payload-guard")
    monkeypatch.setattr(controller, "_validate_metadata_inbox", lambda _inbox: None)

    with pytest.raises(ValueError):
        controller.strategies_performance_v2_retest_start(payload)


def test_retest_start_checks_database_before_reusing_committed_inbox(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    controller = _controller(tmp_path, seed=False)
    _committed_retest_job(controller, tmp_path, job_id="missing-database")
    monkeypatch.setattr(controller, "_validate_metadata_inbox", lambda _inbox: None)

    with pytest.raises(PerformanceV2ApiError) as error:
        controller.strategies_performance_v2_retest_start({})

    assert error.value.status == 404


def test_retest_start_reuses_oldest_valid_inbox_when_newest_is_missing_reports(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    controller = _controller(tmp_path)
    report_dir = tmp_path / "bot" / "tester" / "report" / "my_test"
    strategy_dir = tmp_path / "bot" / "settings_strategy"
    report_dir.mkdir(parents=True)
    strategy_dir.mkdir(parents=True)
    (report_dir / "alpha.html").write_text("<html></html>", encoding="utf-8")
    (strategy_dir / "alpha.json").write_text("{}", encoding="utf-8")
    inbox = tmp_path / "inbox" / "persisted-retest"
    inbox.mkdir(parents=True)
    (inbox / "inbox_manifest.json").write_text(json.dumps({
        "schema_version": 1,
        "run_mode": "SINGLE_MODE",
        "source_mode": "metadata_only",
        "inbox_ready": True,
        "expected_strategy_names": ["alpha"],
        "entries": [{"strategy_name": "alpha", "strategy_path": "alpha.json", "report_path": "alpha.html"}],
    }), encoding="utf-8")
    controller._panel_jobs.submit(
        "strategies.tester.native.start", {"retest": True}, "panel:persisted-retest",
        ("strategies.tester",), job_id="persisted-retest",
    )
    controller._panel_jobs.transition("persisted-retest", "RUNNING")
    controller._panel_jobs.sync(
        "persisted-retest",
        {"state": "COMMITTED", "phase": "COMMITTED", "inbox_ready": True},
        runtime={
            "retest": True,
            "inbox_path": str(inbox),
            "test_start": "2026-01-01",
            "test_end": "2026-01-09",
            "listing_dates_path": "input/dates.xlsx",
        },
    )
    newer_inbox = tmp_path / "inbox" / "newer-invalid"
    newer_inbox.mkdir(parents=True)
    (newer_inbox / "inbox_manifest.json").write_text(json.dumps({
        "schema_version": 1,
        "run_mode": "SINGLE_MODE",
        "source_mode": "metadata_only",
        "inbox_ready": True,
        "expected_strategy_names": ["alpha"],
        "entries": [{"strategy_name": "alpha", "strategy_path": "alpha.json", "report_path": "missing.html"}],
    }), encoding="utf-8")
    controller._panel_jobs.submit(
        "strategies.tester.native.start", {"retest": True}, "panel:newer-invalid",
        ("strategies.tester",), job_id="newer-invalid",
    )
    controller._panel_jobs.transition("newer-invalid", "RUNNING")
    controller._panel_jobs.sync(
        "newer-invalid",
        {"state": "COMMITTED", "phase": "COMMITTED", "inbox_ready": True},
        runtime={"retest": True, "inbox_path": str(newer_inbox)},
    )

    # The journal sorts object keys when persisting, so recovery must use the
    # explicit creation timestamp rather than dictionary/reverse-list order.
    controller = PanelController(tmp_path, tmp_path / "config.local.json")

    class FakeTester:
        def status(self, job_id: str) -> dict[str, object]:
            return {"job_id": job_id, "state": "COMMITTED", "phase": "COMMITTED", "inbox_ready": True}

        def capture_inbox(self, *args: object, **kwargs: object) -> Path:
            pytest.fail("a committed persisted inbox must not be recaptured")

    monkeypatch.setattr(controller, "_single_mode_strategy_test", lambda: FakeTester())
    monkeypatch.setattr(panel_module, "build_retest_manifest", lambda *args: pytest.fail("a new RETEST must not be built"))
    monkeypatch.setattr(controller, "_retest_range", lambda *args: pytest.fail("reusable inbox must skip date validation"))

    result = controller.strategies_performance_v2_retest_start({})

    assert result["job_id"] == "persisted-retest"
    assert result["state"] == "COMMITTED"
    assert result["phase"] == "COMMITTED"
    assert isinstance(result["progress"], dict)
    assert result["inbox_ready"] is True
    assert result["inbox_path"] == str(inbox)
    mapping, runtime = controller._retest_mapping(result["job_id"])
    assert mapping == {"alpha": 1}
    assert runtime["test_start"] == "2026-01-01"
    assert runtime["listing_dates_path"] == "input/dates.xlsx"

    captured: dict[str, object] = {}

    def fake_import(payload: dict[str, object], *, _internal: bool = False) -> dict[str, object]:
        assert _internal is True
        captured.update(payload)
        controller._panel_jobs.submit(
            "strategies.performance.v2.import", {"retest": True}, "panel:import-after-reuse",
            ("performance-v2-db",), job_id="import-after-reuse",
        )
        controller._panel_jobs.transition("import-after-reuse", "RUNNING")
        return controller._panel_jobs.get("import-after-reuse")

    monkeypatch.setattr(controller, "strategies_performance_v2_import", fake_import)
    imported = controller.strategies_performance_v2_retest_import({"tester_job_id": result["job_id"]})

    assert imported["job_id"] == "import-after-reuse"
    assert captured["mode"] == "REPLACE"
    assert captured["replacement_strategy_ids"] == {"alpha": 1}
    assert captured["clear_retest_on_success"] is True


def test_committed_native_retest_verify_reuses_persisted_inbox_without_capture(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    controller = _controller(tmp_path)
    tester_job_id = _committed_retest_job(controller, tmp_path, job_id="verify-persisted-retest")
    report_dir = tmp_path / "bot" / "tester" / "report" / "my_test"
    strategy_dir = tmp_path / "bot" / "settings_strategy"
    report_dir.mkdir(parents=True)
    strategy_dir.mkdir(parents=True)
    (report_dir / "alpha.html").write_text("<html></html>", encoding="utf-8")
    (strategy_dir / "alpha.json").write_text("{}", encoding="utf-8")
    inbox = Path(controller._panel_jobs.runtime(tester_job_id)["inbox_path"])
    (inbox / "inbox_manifest.json").write_text(json.dumps({
        "schema_version": 1,
        "run_mode": "SINGLE_MODE",
        "source_mode": "metadata_only",
        "inbox_ready": True,
        "expected_strategy_names": ["alpha"],
        "entries": [{"strategy_name": "alpha", "strategy_path": "alpha.json", "report_path": "alpha.html"}],
    }), encoding="utf-8")
    monkeypatch.setattr(
        controller, "_single_mode_strategy_test", lambda: pytest.fail("committed RETEST must not load current tester state")
    )
    validated: list[Path] = []
    monkeypatch.setattr(controller, "_validate_metadata_inbox", validated.append)

    result = controller.strategies_tester_verify_inbox(tester_job_id)

    assert result["job_id"] == tester_job_id
    assert result["state"] == "COMMITTED"
    assert result["phase"] == "COMMITTED"
    assert result["inbox_ready"] is True
    assert result["inbox_path"] == str(inbox.resolve())
    assert controller._panel_jobs.runtime(tester_job_id)["retest"] is True
    assert validated == [inbox.resolve()]


def test_committed_native_retest_verify_reports_missing_source_artifacts_for_fresh_run(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    controller = _controller(tmp_path)
    tester_job_id = _committed_retest_job(controller, tmp_path, job_id="verify-missing-sources")
    inbox = Path(controller._panel_jobs.runtime(tester_job_id)["inbox_path"])
    (inbox / "inbox_manifest.json").write_text(json.dumps({
        "schema_version": 1,
        "run_mode": "SINGLE_MODE",
        "source_mode": "metadata_only",
        "inbox_ready": True,
        "expected_strategy_names": ["alpha"],
        "entries": [{"strategy_name": "alpha", "strategy_path": "alpha.json", "report_path": "alpha.html"}],
    }), encoding="utf-8")
    monkeypatch.setattr(
        controller, "_single_mode_strategy_test", lambda: pytest.fail("missing source artifacts must not recapture")
    )

    with pytest.raises(PerformanceV2ApiError) as error:
        controller.strategies_tester_verify_inbox(tester_job_id)

    assert error.value.code == "RETEST_SOURCE_ARTIFACTS_UNAVAILABLE"
    assert "source artifacts" in str(error.value)


def test_committed_native_retest_verify_rejects_same_size_foreign_manifest_without_capture(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    controller = _controller(tmp_path)
    tester_job_id = _committed_retest_job(controller, tmp_path, job_id="verify-foreign-retest")
    inbox = Path(controller._panel_jobs.runtime(tester_job_id)["inbox_path"])
    (inbox / "inbox_manifest.json").write_text(json.dumps({
        "schema_version": 1,
        "batch_id": tester_job_id,
        "run_mode": "SINGLE_MODE",
        "source_mode": "metadata_only",
        "inbox_ready": True,
        "expected_strategy_names": ["alpha"],
        "entries": [{"strategy_name": "foreign", "strategy_path": "foreign.json", "report_path": "foreign.html"}],
    }), encoding="utf-8")
    monkeypatch.setattr(
        controller, "_single_mode_strategy_test", lambda: pytest.fail("broken committed RETEST must not recapture")
    )

    with pytest.raises(ValueError, match="manifest"):
        controller.strategies_tester_verify_inbox(tester_job_id)


def test_committed_native_retest_verify_survives_restart_with_stale_tester_manifest(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    controller = _controller(tmp_path)
    tester_job_id = _committed_retest_job(controller, tmp_path, job_id="verify-restart-retest")
    inbox = Path(controller._panel_jobs.runtime(tester_job_id)["inbox_path"])
    (inbox / "inbox_manifest.json").write_text(json.dumps({
        "schema_version": 1,
        "batch_id": tester_job_id,
        "run_mode": "SINGLE_MODE",
        "source_mode": "metadata_only",
        "inbox_ready": True,
        "expected_strategy_names": ["alpha"],
        "entries": [{"strategy_name": "alpha", "strategy_path": "alpha.json", "report_path": "alpha.html"}],
    }), encoding="utf-8")
    report_dir = tmp_path / "bot" / "tester" / "report" / "my_test"
    report_dir.mkdir(parents=True)
    strategy_dir = tmp_path / "bot" / "settings_strategy"
    strategy_dir.mkdir(parents=True)
    (report_dir / "alpha.html").write_text("<html></html>", encoding="utf-8")
    (strategy_dir / "alpha.json").write_text("{}", encoding="utf-8")
    (report_dir / "tester_manifest.json").write_text(json.dumps({"job_id": "foreign-job", "phase": "COMMITTED"}), encoding="utf-8")
    restored = PanelController(tmp_path, tmp_path / "config.local.json")
    monkeypatch.setattr(
        restored, "_single_mode_strategy_test", lambda: pytest.fail("stale tester manifest must not be loaded")
    )

    result = restored.strategies_tester_verify_inbox(tester_job_id)

    assert result["job_id"] == tester_job_id
    assert result["inbox_ready"] is True


@pytest.mark.parametrize(
    ("import_state", "reusable"),
    [("COMMITTED", False), ("FAILED", True), ("CANCELLED", True)],
)
def test_reusable_retest_job_checks_persisted_import_state(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, import_state: str, reusable: bool
) -> None:
    controller = _controller(tmp_path)
    tester_job_id = _committed_retest_job(controller, tmp_path, job_id=f"tester-import-{import_state.lower()}")
    import_job_id = f"import-{import_state.lower()}"
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {"retest": True}, f"panel:{import_job_id}",
        ("performance-v2-db",), job_id=import_job_id,
    )
    controller._panel_jobs.transition(import_job_id, "RUNNING")
    if import_state == "COMMITTED":
        controller._panel_jobs.sync(import_job_id, {"state": "COMMITTED", "phase": "COMMITTED"})
    elif import_state == "FAILED":
        controller._panel_jobs.sync(import_job_id, {"state": "FAILED", "phase": "FAILED"})
    else:
        controller._panel_jobs.transition(import_job_id, "CANCELLING")
        controller._panel_jobs.transition(import_job_id, "CANCELLED")
    runtime = controller._panel_jobs.runtime(tester_job_id)
    runtime["retest_import_job_id"] = import_job_id
    controller._panel_jobs.sync(tester_job_id, {"state": "COMMITTED"}, runtime=runtime)
    monkeypatch.setattr(controller, "_validate_metadata_inbox", lambda _inbox: None)

    result = controller._reusable_retest_job()

    assert (result is not None) is reusable
    if reusable:
        assert result["job_id"] == tester_job_id


def test_reusable_retest_job_rejects_inbox_root_itself(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    controller = _controller(tmp_path)
    inbox_root = tmp_path / "inbox"
    inbox_root.mkdir()
    tester_job_id = _committed_retest_job(controller, tmp_path, job_id="tester-inbox-root")
    runtime = controller._panel_jobs.runtime(tester_job_id)
    runtime["inbox_path"] = str(inbox_root)
    controller._panel_jobs.sync(tester_job_id, {"state": "COMMITTED"}, runtime=runtime)
    monkeypatch.setattr(controller, "_validate_metadata_inbox", lambda _inbox: None)

    assert controller._reusable_retest_job() is None


def test_reusable_retest_job_uses_creation_order_after_registry_restore(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    controller = _controller(tmp_path)
    old_job_id = _committed_retest_job(controller, tmp_path, job_id="z-old-retest")
    new_job_id = _committed_retest_job(controller, tmp_path, job_id="a-new-retest")
    controller._panel_jobs.jobs[old_job_id].pop("created_at_utc", None)
    controller._panel_jobs.jobs[new_job_id]["created_at_utc"] = "2026-01-01T00:00:01+00:00"
    controller._panel_jobs._save()
    restored = PanelController(tmp_path, tmp_path / "config.local.json")
    monkeypatch.setattr(restored, "_validate_metadata_inbox", lambda _inbox: None)

    result = restored._reusable_retest_job()

    assert "created_at_utc" not in restored._panel_jobs.get(old_job_id)
    assert result["job_id"] == new_job_id


@pytest.mark.parametrize("link_kind", ["resource", "tester_job", "inbox"])
def test_reusable_retest_job_blocks_reverse_committed_import_link(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, link_kind: str
) -> None:
    controller = _controller(tmp_path)
    tester_job_id = _committed_retest_job(controller, tmp_path, job_id=f"reverse-link-{link_kind}")
    inbox = Path(controller._panel_jobs.runtime(tester_job_id)["inbox_path"])
    runtime: dict[str, object] = {}
    resource_keys: tuple[str, ...] = ()
    if link_kind == "resource":
        resource_keys = (f"tester:{tester_job_id}",)
    elif link_kind == "tester_job":
        runtime["tester_job_id"] = tester_job_id
    else:
        runtime["inbox_path"] = str(inbox)
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {"retest": True}, f"panel:reverse-{link_kind}",
        resource_keys, job_id=f"reverse-import-{link_kind}",
    )
    controller._panel_jobs.transition(f"reverse-import-{link_kind}", "RUNNING")
    controller._panel_jobs.sync(
        f"reverse-import-{link_kind}",
        {"state": "COMMITTED", "phase": "COMMITTED"},
        runtime=runtime or None,
    )
    monkeypatch.setattr(controller, "_validate_metadata_inbox", lambda _inbox: None)

    assert controller._reusable_retest_job() is None


@pytest.mark.parametrize("kind", ["tester", "import"])
@pytest.mark.parametrize("state", ["QUEUED", "RUNNING", "CANCELLING"])
def test_reusable_retest_job_blocks_active_retest_jobs(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, kind: str, state: str
) -> None:
    controller = _controller(tmp_path)
    _committed_retest_job(controller, tmp_path, job_id=f"active-candidate-{kind}-{state}")
    job_id = f"active-{kind}-{state}"
    panel_kind = "strategies.tester.native.start" if kind == "tester" else "strategies.performance.v2.import"
    resource_keys = ("strategies.tester",) if kind == "tester" else ("performance-v2-db",)
    controller._panel_jobs.submit(panel_kind, {"retest": True}, f"panel:{job_id}", resource_keys, job_id=job_id)
    if state != "QUEUED":
        controller._panel_jobs.transition(job_id, "RUNNING")
    if state == "CANCELLING":
        controller._panel_jobs.transition(job_id, "CANCELLING")
    monkeypatch.setattr(controller, "_validate_metadata_inbox", lambda _inbox: None)

    assert controller._reusable_retest_job() is None


def _committed_retest_job(controller: PanelController, tmp_path: Path, *, state: str = "COMMITTED", job_id: str = "retest-job") -> str:
    inbox = tmp_path / "inbox" / job_id
    inbox.mkdir(parents=True)
    (inbox / "inbox_manifest.json").write_text(
        json.dumps({"expected_strategy_names": ["alpha"], "run_mode": "SINGLE_MODE"}), encoding="utf-8"
    )
    controller._panel_jobs.submit(
        "strategies.tester.native.start", {"retest": True}, f"panel:{job_id}", ("strategies.tester",), job_id=job_id
    )
    controller._panel_jobs.transition(job_id, "RUNNING")
    controller._panel_jobs.sync(
        job_id,
        {"state": state, "phase": state, "inbox_ready": state == "COMMITTED"},
        runtime={
            "retest": True, "inbox_path": str(inbox), "test_start": "2026-01-01",
            "test_end": "2026-01-09", "listing_dates_path": "input/dates.xlsx",
        },
    )
    return job_id


def test_retest_import_requires_committed_inbox_and_builds_mapping_on_server(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    controller = _controller(tmp_path)
    job_id = _committed_retest_job(controller, tmp_path)
    assert controller._panel_jobs.get(job_id)["retest"] is True
    captured: dict[str, object] = {}
    def fake_import(payload: dict[str, object], **_kwargs: object) -> dict[str, object]:
        captured.update(payload)
        controller._panel_jobs.submit(
            "strategies.performance.v2.import", {"retest": True}, "panel:import",
            ("performance-v2-db",), job_id="import",
        )
        controller._panel_jobs.transition("import", "RUNNING")
        return controller._panel_jobs.get("import")

    monkeypatch.setattr(controller, "strategies_performance_v2_import", fake_import)

    with pytest.raises(ValueError, match="server-built"):
        controller.strategies_performance_v2_retest_import({"tester_job_id": job_id, "replacement_strategy_ids": {"alpha": 1}})
    result = controller.strategies_performance_v2_retest_import({"tester_job_id": job_id})

    assert result["job_id"] == "import"
    assert captured["mode"] == "REPLACE"
    assert captured["replacement_strategy_ids"] == {"alpha": 1}
    assert captured["clear_retest_on_success"] is True
    assert captured["test_start"] == "2026-01-01"
    assert captured["listing_dates_path"] == "input/dates.xlsx"
    assert controller._panel_jobs.runtime(job_id)["retest_import_job_id"] == "import"
    with pytest.raises(ValueError, match="already started"):
        controller.strategies_performance_v2_retest_import({"tester_job_id": job_id})
    restarted = PanelController(tmp_path, tmp_path / "config.local.json")
    with pytest.raises(ValueError, match="already started"):
        restarted.strategies_performance_v2_retest_import({"tester_job_id": job_id})

    _committed_retest_job(controller, tmp_path / "not-ready", state="RUNNING", job_id="not-ready")
    with pytest.raises(ValueError, match="not committed"):
        controller.strategies_performance_v2_retest_import({"tester_job_id": "not-ready"})


def test_retest_import_allows_retry_after_failed_inner_job(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    controller = _controller(tmp_path)
    tester_job_id = _committed_retest_job(controller, tmp_path, job_id="tester-retry-failed-import")
    failed_import_id = "failed-retest-import"
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {"retest": True}, "panel:failed-retest-import",
        ("performance-v2-db",), job_id=failed_import_id,
    )
    controller._panel_jobs.transition(failed_import_id, "RUNNING")
    controller._panel_jobs.sync(
        failed_import_id,
        {"state": "FAILED", "phase": "FAILED", "error": {"code": "PERFORMANCE_V2_IMPORT_FAILED"}},
    )
    runtime = controller._panel_jobs.runtime(tester_job_id)
    runtime["retest_import_job_id"] = failed_import_id
    controller._panel_jobs.sync(tester_job_id, {"state": "COMMITTED"}, runtime=runtime)

    def fake_import(_payload: dict[str, object], *, _internal: bool = False) -> dict[str, object]:
        assert _internal is True
        retry_id = "retry-retest-import"
        controller._panel_jobs.submit(
            "strategies.performance.v2.import", {"retest": True}, f"panel:{retry_id}",
            ("performance-v2-db",), job_id=retry_id,
        )
        controller._panel_jobs.transition(retry_id, "RUNNING")
        return controller._panel_jobs.get(retry_id)

    monkeypatch.setattr(controller, "strategies_performance_v2_import", fake_import)

    result = controller.strategies_performance_v2_retest_import({"tester_job_id": tester_job_id})

    assert result["job_id"] == "retry-retest-import"
    assert result["job_id"] != failed_import_id
    assert controller._panel_jobs.runtime(tester_job_id)["retest_import_job_id"] == "retry-retest-import"
    with pytest.raises(ValueError, match="already started"):
        controller.strategies_performance_v2_retest_import({"tester_job_id": tester_job_id})


def test_retest_import_allows_retry_after_cancelled_inner_job(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    controller = _controller(tmp_path)
    tester_job_id = _committed_retest_job(controller, tmp_path, job_id="tester-retry-cancelled-import")
    cancelled_import_id = "cancelled-retest-import"
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {"retest": True}, "panel:cancelled-retest-import",
        ("performance-v2-db",), job_id=cancelled_import_id,
    )
    controller._panel_jobs.transition(cancelled_import_id, "RUNNING")
    controller._panel_jobs.transition(cancelled_import_id, "CANCELLING")
    controller._panel_jobs.transition(cancelled_import_id, "CANCELLED")
    runtime = controller._panel_jobs.runtime(tester_job_id)
    runtime["retest_import_job_id"] = cancelled_import_id
    controller._panel_jobs.sync(tester_job_id, {"state": "COMMITTED"}, runtime=runtime)

    def fake_import(_payload: dict[str, object], *, _internal: bool = False) -> dict[str, object]:
        assert _internal is True
        retry_id = "retry-cancelled-retest-import"
        controller._panel_jobs.submit(
            "strategies.performance.v2.import", {"retest": True}, f"panel:{retry_id}",
            ("performance-v2-db",), job_id=retry_id,
        )
        controller._panel_jobs.transition(retry_id, "RUNNING")
        return controller._panel_jobs.get(retry_id)

    monkeypatch.setattr(controller, "strategies_performance_v2_import", fake_import)
    result = controller.strategies_performance_v2_retest_import({"tester_job_id": tester_job_id})

    assert result["job_id"] == "retry-cancelled-retest-import"
    assert controller._panel_jobs.runtime(tester_job_id)["retest_import_job_id"] == "retry-cancelled-retest-import"
    with pytest.raises(ValueError, match="already started"):
        controller.strategies_performance_v2_retest_import({"tester_job_id": tester_job_id})


def test_retest_import_blocks_interrupted_marker_with_extra_error_fields(tmp_path: Path) -> None:
    controller = _controller(tmp_path)
    tester_job_id = _committed_retest_job(controller, tmp_path, job_id="tester-interrupted-import")
    interrupted_import_id = "interrupted-retest-import"
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {"retest": True}, "panel:interrupted-retest-import",
        ("performance-v2-db",), job_id=interrupted_import_id,
    )
    controller._panel_jobs.transition(interrupted_import_id, "RUNNING")
    controller._panel_jobs.sync(
        interrupted_import_id,
        {"state": "FAILED", "phase": "FAILED", "error": {"code": "INTERRUPTED", "message": "restart"}},
    )
    runtime = controller._panel_jobs.runtime(tester_job_id)
    runtime["retest_import_job_id"] = interrupted_import_id
    controller._panel_jobs.sync(tester_job_id, {"state": "COMMITTED"}, runtime=runtime)

    with pytest.raises(ValueError, match="already started"):
        controller.strategies_performance_v2_retest_import({"tester_job_id": tester_job_id})


def test_retest_import_blocks_failed_marker_with_unknown_error_code(tmp_path: Path) -> None:
    controller = _controller(tmp_path)
    tester_job_id = _committed_retest_job(controller, tmp_path, job_id="tester-unknown-import-error")
    unknown_import_id = "unknown-retest-import"
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {"retest": True}, "panel:unknown-retest-import",
        ("performance-v2-db",), job_id=unknown_import_id,
    )
    controller._panel_jobs.transition(unknown_import_id, "RUNNING")
    controller._panel_jobs.sync(
        unknown_import_id,
        {"state": "FAILED", "phase": "FAILED", "error": {"code": "UNKNOWN_FAILURE"}},
    )
    runtime = controller._panel_jobs.runtime(tester_job_id)
    runtime["retest_import_job_id"] = unknown_import_id
    controller._panel_jobs.sync(tester_job_id, {"state": "COMMITTED"}, runtime=runtime)

    with pytest.raises(ValueError, match="already started"):
        controller.strategies_performance_v2_retest_import({"tester_job_id": tester_job_id})


def test_retest_recovery_selector_prioritizes_ready_and_newest_jobs() -> None:
    utility = Path(__file__).parents[1] / "src" / "mrs3" / "panel_web" / "retest_recovery.js"
    tester = lambda job_id, state, inbox_ready=False: {
        "job_id": job_id, "kind": "strategies.tester.native.start", "retest": True,
        "state": state, "inbox_ready": inbox_ready,
    }
    fixtures = [
        [tester("ready-old", "COMMITTED", True), tester("failed-new", "FAILED")],
        [tester("ready-old", "COMMITTED", True), tester("running-new", "RUNNING")],
        [tester("ready-old", "COMMITTED", True), tester("ready-new", "COMMITTED", True)],
        [tester("failed-old", "FAILED"), tester("failed-new", "FAILED")],
        [tester("failed-old", "FAILED"), tester("committed-no-ready", "COMMITTED")],
    ]
    script = (
        f"const {{selectRetestTester}} = require({json.dumps(str(utility))});"
        f"const fixtures = {json.dumps(fixtures)};"
        "process.stdout.write(JSON.stringify(fixtures.map((jobs) => selectRetestTester(jobs)?.job_id ?? null)));"
    )
    result = subprocess.run(("node", "-e", script), capture_output=True, text=True, check=True)

    assert json.loads(result.stdout) == [
        "ready-old", "ready-old", "ready-new", "failed-new", "committed-no-ready",
    ]


def test_retest_check_selector_uses_newest_committed_native_job() -> None:
    utility = Path(__file__).parents[1] / "src" / "mrs3" / "panel_web" / "retest_recovery.js"
    tester = lambda job_id, state, inbox_ready=False: {
        "job_id": job_id, "kind": "strategies.tester.native.start", "retest": True,
        "state": state, "inbox_ready": inbox_ready,
    }
    jobs = [
        tester("ready-old", "COMMITTED", True),
        tester("committed-new", "COMMITTED"),
        tester("running-newest", "RUNNING"),
    ]
    script = (
        f"const {{selectCommittedRetestTester}} = require({json.dumps(str(utility))});"
        f"const jobs = {json.dumps(jobs)};"
        "process.stdout.write(selectCommittedRetestTester(jobs)?.job_id ?? '');"
    )
    result = subprocess.run(("node", "-e", script), capture_output=True, text=True, check=True)

    assert result.stdout == "committed-new"


def test_retest_import_reserves_before_inner_launch(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    controller = _controller(tmp_path)
    tester_job_id = _committed_retest_job(controller, tmp_path, job_id="tester-reservation")
    entered = threading.Event()
    release = threading.Event()

    def fake_import(_payload: dict[str, object], *, _internal: bool = False) -> dict[str, object]:
        assert _internal is True
        entered.set()
        assert release.wait(2)
        controller._panel_jobs.submit(
            "strategies.performance.v2.import", {"retest": True}, "panel:reserved-import",
            ("performance-v2-db",), job_id="reserved-import",
        )
        controller._panel_jobs.transition("reserved-import", "RUNNING")
        return controller._panel_jobs.get("reserved-import")

    monkeypatch.setattr(controller, "strategies_performance_v2_import", fake_import)
    errors: list[BaseException] = []

    def launch() -> None:
        try:
            controller.strategies_performance_v2_retest_import({"tester_job_id": tester_job_id})
        except BaseException as error:
            errors.append(error)

    worker = threading.Thread(target=launch)
    worker.start()
    assert entered.wait(2)
    with pytest.raises(ValueError, match="already started"):
        controller.strategies_performance_v2_retest_import({"tester_job_id": tester_job_id})
    release.set()
    worker.join(2)
    assert not worker.is_alive()
    assert not errors
    assert controller._panel_jobs.runtime(tester_job_id)["retest_import_job_id"] == "reserved-import"


@pytest.mark.parametrize("malformed", [None, {}, {"job_id": ""}, {"job_id": "unregistered"}])
def test_retest_import_clears_reservation_when_inner_job_is_malformed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, malformed: object
) -> None:
    controller = _controller(tmp_path)
    tester_job_id = _committed_retest_job(controller, tmp_path, job_id="tester-retry")
    calls = 0

    def fake_import(_payload: dict[str, object], *, _internal: bool = False) -> object:
        nonlocal calls
        assert _internal is True
        calls += 1
        if calls == 1:
            return malformed
        controller._panel_jobs.submit(
            "strategies.performance.v2.import", {"retest": True}, "panel:retry-import",
            ("performance-v2-db",), job_id="retry-import",
        )
        controller._panel_jobs.transition("retry-import", "RUNNING")
        return controller._panel_jobs.get("retry-import")

    monkeypatch.setattr(controller, "strategies_performance_v2_import", fake_import)
    with pytest.raises(ValueError, match="invalid|not registered"):
        controller.strategies_performance_v2_retest_import({"tester_job_id": tester_job_id})
    assert "retest_import_job_id" not in controller._panel_jobs.runtime(tester_job_id)

    result = controller.strategies_performance_v2_retest_import({"tester_job_id": tester_job_id})
    assert result["job_id"] == "retry-import"
    assert controller._panel_jobs.runtime(tester_job_id)["retest_import_job_id"] == "retry-import"


def test_retest_import_route_returns_inner_job_and_serves_its_failure_report(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    controller = _controller(tmp_path)
    tester_job_id = _committed_retest_job(controller, tmp_path, job_id="tester-route")
    report = tmp_path / "performance-v2" / "retest-failures.csv"
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text("reason\nTEST\n", encoding="utf-8")

    def fake_import(payload: dict[str, object], *, _internal: bool = False) -> dict[str, object]:
        assert _internal is True and payload["mode"] == "REPLACE"
        inner_id = "retest-import-route"
        controller._panel_jobs.submit(
            "strategies.performance.v2.import", {"retest": True}, f"panel:{inner_id}",
            ("performance-v2-db",), job_id=inner_id,
        )
        controller._panel_jobs.transition(inner_id, "RUNNING")
        document = {"job_id": inner_id, "state": "COMMITTED", "result": {
            "failure_report_path": str(report),
            "database_path": str(tmp_path / "private.duckdb"),
            "audit_path": str(tmp_path / "private.audit.json"),
        }}
        controller._panel_jobs.sync(inner_id, document)
        controller._record_special_job(document)
        return controller._panel_jobs.get(inner_id)

    monkeypatch.setattr(controller, "strategies_performance_v2_import", fake_import)
    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        try:
            connection.request(
                "POST", "/api/v2/strategies/performance-v2/retest/import",
                body=json.dumps({"tester_job_id": tester_job_id}),
                headers={"Content-Type": "application/json"},
            )
            response = connection.getresponse()
            body = json.loads(response.read())
        finally:
            connection.close()
        assert response.status == 202
        inner_id = body["job"]["job_id"]
        assert inner_id == "retest-import-route" and body["job"]["retest"] is True
        assert body["job"]["result"]["failure_report_available"] is True
        assert "failure_report_path" not in body["job"]["result"]
        assert "database_path" not in body["job"]["result"]
        assert "audit_path" not in body["job"]["result"]

        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        try:
            connection.request("GET", f"/api/artifact?name=performance-v2-failure-report:{inner_id}")
            artifact_response = connection.getresponse()
            artifact = artifact_response.read()
        finally:
            connection.close()
        assert artifact_response.status == 200
        assert artifact == report.read_bytes()
        restarted = PanelController(tmp_path, tmp_path / "config.local.json")
        assert restarted.artifact(f"performance-v2-failure-report:{inner_id}") == report.resolve()
        restarted_server = create_panel_server("127.0.0.1", 0, restarted)
        restarted_thread = threading.Thread(target=restarted_server.serve_forever, daemon=True)
        restarted_thread.start()
        try:
            connection = HTTPConnection("127.0.0.1", restarted_server.server_port, timeout=2)
            try:
                connection.request("GET", f"/api/artifact?name=performance-v2-failure-report:{inner_id}")
                restarted_response = connection.getresponse()
                restarted_artifact = restarted_response.read()
            finally:
                connection.close()
            assert restarted_response.status == 200
            assert restarted_artifact == report.read_bytes()
        finally:
            restarted_server.shutdown()
            restarted_server.server_close()
    finally:
        server.shutdown()
        server.server_close()


def test_public_performance_v2_dispatcher_rejects_replacement_controls(tmp_path: Path) -> None:
    controller = _controller(tmp_path)
    job_id = _committed_retest_job(controller, tmp_path, job_id="tester-public")

    with pytest.raises(ValueError, match="internal only"):
        controller.panel_job_submit({
            "kind": "strategies.performance.v2.import",
            "request": {"tester_job_id": job_id, "replacement_strategy_ids": {"alpha": 1}},
        })
    with pytest.raises(ValueError, match="internal only"):
        controller.panel_job_submit({
            "kind": "strategies.performance.v2.import",
            "request": {"tester_job_id": job_id, "mode": "REPLACE"},
        })
    with pytest.raises(ValueError, match="internal only"):
        controller.panel_job_submit({
            "kind": "strategies.performance.v2.import",
            "request": {"tester_job_id": job_id, "clear_retest_on_success": True},
        })
    with pytest.raises(ValueError, match="internal only"):
        controller.panel_job_submit({
            "kind": "strategies.performance.v2.import",
            "request": {"tester_job_id": job_id, "_retest": True},
        })


def test_retest_http_start_and_import_return_job_envelopes(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    controller = _controller(tmp_path, seed=False)
    calls: list[dict[str, object]] = []

    def fake_submit(document: dict[str, object]) -> dict[str, str]:
        calls.append(document)
        return {"job_id": f"job-{len(calls)}"}

    monkeypatch.setattr(controller, "panel_job_submit", fake_submit)
    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    requests = (
        ("/api/v2/strategies/performance-v2/retest/start", {"test_start": "2026-01-01", "test_end": "2026-01-09"}),
        ("/api/v2/strategies/performance-v2/retest/import", {"tester_job_id": "tester-1"}),
    )
    try:
        for endpoint, payload in requests:
            connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
            try:
                connection.request(
                    "POST", endpoint, body=json.dumps(payload),
                    headers={"Content-Type": "application/json"},
                )
                response = connection.getresponse()
                body = json.loads(response.read())
            finally:
                connection.close()
            assert response.status == 202
            assert body == {"job": {"job_id": f"job-{len(calls)}"}}
    finally:
        server.shutdown()
        server.server_close()

    assert calls == [
        {"kind": "strategies.performance.v2.retest.start", "request": requests[0][1]},
        {"kind": "strategies.performance.v2.retest.import", "request": requests[1][1]},
    ]


def test_retest_status_http_is_safe_without_database(tmp_path: Path) -> None:
    controller = _controller(tmp_path, seed=False)
    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        try:
            connection.request("GET", "/api/v2/strategies/performance-v2/retest/status")
            response = connection.getresponse()
            body = json.loads(response.read())
        finally:
            connection.close()
        assert response.status == 200
        assert body["count"] == 0 and body["phase"] == "IDLE"
    finally:
        server.shutdown()
        server.server_close()


def test_failure_report_is_served_only_for_committed_import_job(tmp_path: Path) -> None:
    controller = _controller(tmp_path)
    job_id = "import-job"
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {"tester_job_id": "tester"}, f"panel:{job_id}", ("performance-v2-db",), job_id=job_id
    )
    controller._panel_jobs.transition(job_id, "RUNNING")
    report = tmp_path / "performance-v2" / "performance_v2_failures_import.csv"
    report.write_text("reason\nTEST\n", encoding="utf-8")
    controller._panel_jobs.sync(job_id, {"state": "COMMITTED", "result": {"failure_report_path": str(report)}})
    controller._record_special_job({"job_id": job_id, "state": "COMMITTED", "result": {"failure_report_path": str(report)}})
    assert controller.artifact(f"performance-v2-failure-report:{job_id}") == report.resolve()
    with pytest.raises(ValueError):
        controller.artifact("performance-v2-failure-report:missing")
    restarted = PanelController(tmp_path, tmp_path / "config.local.json")
    assert restarted.artifact(f"performance-v2-failure-report:{job_id}") == report.resolve()

    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        try:
            connection.request("GET", f"/api/artifact?name=performance-v2-failure-report:{job_id}")
            response = connection.getresponse()
            body = response.read()
        finally:
            connection.close()
        assert response.status == 200
        assert body == report.read_bytes()
    finally:
        server.shutdown()
        server.server_close()


def test_failure_report_requires_committed_import_and_contained_regular_file(tmp_path: Path) -> None:
    controller = _controller(tmp_path)
    root = tmp_path / "performance-v2"
    report = root / "report.csv"
    report.write_text("ok\n", encoding="utf-8")

    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {"tester_job_id": "tester"}, "import-running", ("db",), job_id="import-running"
    )
    controller._panel_jobs.transition("import-running", "RUNNING")
    controller._panel_jobs.sync("import-running", {"state": "RUNNING", "result": {"failure_report_path": str(report)}})
    with pytest.raises(ValueError):
        controller.performance_v2_failure_report("import-running")

    controller._panel_jobs.submit(
        "strategies.tester.native.start", {}, "foreign", (), job_id="foreign"
    )
    controller._panel_jobs.transition("foreign", "RUNNING")
    controller._panel_jobs.sync("foreign", {"state": "COMMITTED", "result": {"failure_report_path": str(report)}})
    with pytest.raises(ValueError):
        controller.performance_v2_failure_report("foreign")

    nested = root / "nested" / "report.csv"
    nested.parent.mkdir()
    nested.write_text("nested\n", encoding="utf-8")
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {}, "nested", (), job_id="import-nested"
    )
    controller._panel_jobs.transition("import-nested", "RUNNING")
    controller._panel_jobs.sync("import-nested", {"state": "COMMITTED", "result": {"failure_report_path": str(nested)}})
    with pytest.raises(ValueError):
        controller.performance_v2_failure_report("import-nested")

    outside = tmp_path / "outside.csv"
    outside.write_text("outside\n", encoding="utf-8")
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {}, "outside", (), job_id="import-outside"
    )
    controller._panel_jobs.transition("import-outside", "RUNNING")
    controller._panel_jobs.sync("import-outside", {"state": "COMMITTED", "result": {"failure_report_path": str(outside)}})
    with pytest.raises(ValueError):
        controller.performance_v2_failure_report("import-outside")

    symlink = root / "report-link.csv"
    try:
        symlink.symlink_to(outside)
    except OSError:
        pytest.skip("Windows environment does not provide usable symlink creation")
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {}, "symlink", (), job_id="import-symlink"
    )
    controller._panel_jobs.transition("import-symlink", "RUNNING")
    controller._panel_jobs.sync("import-symlink", {"state": "COMMITTED", "result": {"failure_report_path": str(symlink)}})
    with pytest.raises(ValueError):
        controller.performance_v2_failure_report("import-symlink")


def test_later_listing_date_is_valid_warmup_input(tmp_path: Path) -> None:
    controller = _controller(tmp_path)
    database = tmp_path / "performance-v2" / "strategy_performance.duckdb"
    with duckdb.connect(str(database), read_only=True) as connection:
        controller._validate_retest_listing(connection, tmp_path / "input" / "dates.xlsx", "2026-01-01")
