from __future__ import annotations

from dataclasses import replace
from copy import deepcopy
from contextlib import contextmanager
from hashlib import sha256
from http.client import HTTPConnection
from decimal import Decimal
import json
from io import BytesIO
from pathlib import Path
import threading
import time
from types import SimpleNamespace
from datetime import datetime, timedelta, timezone

import duckdb
import pandas as pd
from openpyxl import Workbook, load_workbook
import pytest

from tests._protocol_v2 import use_protocol_v2_target

import mrs3.panel as panel_module
import mrs3.panel_jobs as panel_jobs_module
from mrs3.panel_jobs import PanelJobRegistry
from mrs3.performance_v2_store import (
    PerformanceV2Config,
    PerformanceV2StoreError,
    initialize_performance_v2,
    performance_v2_database_path,
)
from mrs3.panel import PanelController, create_panel_server
from mrs3.performance_v2_import import PerformanceV2ImportError, PerformanceV2ImportResult
from mrs3.performance_v2_finalist_retest import FinalistRetestError, combined_control_workbook_bytes
from mrs3.performance_v2_equity_quality import EquitySample, calculate_equity_quality_facts
from mrs3.performance_v2_equity_regime import (
    ALGORITHM_VERSION as EQUITY_REGIME_ALGORITHM_VERSION,
    EquityRegimeSample,
    classify_equity_regime,
)
from mrs3.performance_v2_equity_regime_cache import (
    encode_equity_regime_assessment,
    encode_equity_regime_facts,
)
from mrs3.performance_v2_selection import (
    PerformanceV2SelectionError,
    SelectionConfig,
    SelectionRequest,
    parse_selection_request,
    selection_cache_missing_strategy_ids,
    write_selection_workbook,
)
from mrs3.panel_performance_v2 import (
    PerformanceV2ApiError,
    PerformanceV2PanelRequest,
    PerformanceV2PanelResult,
    LocalPerformanceV2Jobs,
    LocalPerformanceV2Service,
)
from mrs3.performance_v2_windows import WindowMetrics


FIXTURE = Path(__file__).parent / "fixtures" / "performance" / "report_current_v2.html"
UTC = timezone.utc


def test_panel_performance_v2_uses_common_import_workers(tmp_path: Path) -> None:
    config = tmp_path / "config.local.json"
    config.write_text(json.dumps({"duckdb_import": {"workers": 7}}), encoding="utf-8")
    (tmp_path / "config.performance.json").write_text(
        json.dumps({"unified_performance_v2": {"database_root": "performance-v2", "workers": 1}}),
        encoding="utf-8",
    )
    controller = PanelController(tmp_path, config)

    assert controller._performance_v2_config().workers == 7
    config.write_text(json.dumps({"duckdb_import": {"workers": 11}}), encoding="utf-8")
    assert controller._performance_v2_config().workers == 11


def _metrics_for_normalization(
    start: datetime,
    end: datetime | None,
    *,
    growth: object = Decimal("1.1"),
    trade_count: object = 5,
) -> WindowMetrics:
    return WindowMetrics(
        1, start, end or start, "test", start, end, "AVAILABLE", None,
        growth, Decimal("1.2345"), Decimal(".01"), Decimal("1.02"), Decimal("3.4"),
        Decimal("2.3"), Decimal(".4"), Decimal("1.5"), trade_count, Decimal("50"),
    )


def _initialize_v11(connection: duckdb.DuckDBPyConnection) -> None:
    """Create a schema-v11 database: the current initializer stops before the v12 migration.

    DuckDB cannot drop columns from FK-referenced tables, so legacy-catalog
    fixtures start from v11 instead of downgrading a v12 database.
    """
    import mrs3.performance_v2_store as store

    migrate, require = store._migrate_schema_v11_to_v12, store.require_performance_v2
    store._migrate_schema_v11_to_v12 = lambda _connection: None
    store.require_performance_v2 = store._require_v11_catalog
    try:
        store.initialize_performance_v2(connection)
    finally:
        store._migrate_schema_v11_to_v12, store.require_performance_v2 = migrate, require


def _rebuild_as_v11(database: Path) -> None:
    """Copy a Panel-migrated v12 fixture into a fresh v11 database with the same markers and rows."""
    from mrs3.performance_v2_compact import TABLES

    rebuilt = database.with_name(f"{database.stem}.v11-rebuild.duckdb")
    with duckdb.connect(str(rebuilt)) as target:
        _initialize_v11(target)
        origin = database.as_posix().replace("'", "''")
        target.execute(f"attach '{origin}' as origin (read_only)")
        target.execute("delete from main.schema_info")
        target.execute("insert into main.schema_info select * from origin.main.schema_info")
        target.execute("update main.schema_info set value = '11' where key = 'schema_version'")
        for table in TABLES:
            columns = [row[0] for row in target.execute(
                "select column_name from information_schema.columns where table_catalog = current_database() "
                "and table_schema = 'main' and table_name = ? order by ordinal_position", [table],
            ).fetchall()]
            if columns:
                names = ", ".join(columns)
                target.execute(f"insert into main.{table} ({names}) select {names} from origin.main.{table}")
        target.execute("detach origin")
    database.unlink()
    rebuilt.rename(database)


def _db(tmp_path: Path, *, legacy_v11: bool = False) -> tuple[duckdb.DuckDBPyConnection, int]:
    tmp_path.mkdir(parents=True, exist_ok=True)
    target = (tmp_path / "strategy_performance.duckdb").resolve()
    known_production_targets = {
        (Path(__file__).resolve().parents[1] / "data" / "performanceDB" / "strategy_performance.duckdb").resolve(),
        (Path(__file__).resolve().parents[1] / "data" / "performance-v2" / "strategy_performance.duckdb").resolve(),
    }
    assert target.is_relative_to(tmp_path.resolve()), "writable database target must be beneath its test fixture root"
    assert target not in known_production_targets, "known production PerformanceDB paths are forbidden"
    assert not target.is_relative_to(Path(__file__).resolve().parents[1]), "repository databases are read-only"
    connection = duckdb.connect(str(target))
    if legacy_v11:
        _initialize_v11(connection)
    else:
        initialize_performance_v2(connection)
    now = datetime(2026, 1, 1, tzinfo=UTC)
    strategy_id = connection.execute(
        """insert into strategies (strategy_name, symbol, side, timeframe, close_ma_len,
           order_count, analysis_run_id, candidate_identity, lifecycle_status,
           created_at_utc, updated_at_utc) values ('alpha', 'BTCUSDT', 'LONG', '1h',
           3, 1, 'run', 'candidate', 'ACTIVE', ?, ?) returning strategy_id""",
        [now, now],
    ).fetchone()[0]
    result_id = connection.execute(
        """insert into strategy_results (strategy_id, report_start_utc, report_end_utc, exchange,
           commission_rate, initial_balance, final_balance, total_pnl, total_pnl_pct,
           max_drawdown, max_drawdown_pct, total_fees, total_trades, imported_at_utc)
           values (?, ?, ?, 'Bybit', .0004, 100, 110, 10, 10, 0, 0, 2, 2, ?) returning result_id""",
        [strategy_id, now, datetime(2026, 1, 5, tzinfo=UTC), now],
    ).fetchone()[0]
    connection.execute("update strategies set current_result_id = ? where strategy_id = ?", [result_id, strategy_id])
    connection.execute(
        """insert into analysis_plateaus (analysis_run_id, plateau_id, plateau_point_count, plateau_total_trades)
           values ('run', 'P1', 12, 34)"""
    )
    connection.execute(
        """insert into strategy_orders (strategy_id, order_id, open_ma_len, open_multiplier,
           shift_bp, lot_x, analysis_run_id, plateau_id, base_point_trades)
           values (?, 1, 7, 0.995, 125, 1, 'run', 'P1', 8)""",
        [strategy_id],
    )
    connection.executemany(
        "insert into strategy_actions (result_id, action_index, timestamp_utc, symbol, order_id, action, size, post_size, post_side, pnl, fee, balance, raw_action_json) values (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        [
            (result_id, 0, now, "BTCUSDT", 1, "opened", 1, 1, "long", 0, 1, 100, None),
            (result_id, 1, datetime(2026, 1, 2, tzinfo=UTC), "BTCUSDT", 1, "closed", 1, 0, "", 10, 1, 110, None),
        ],
    )
    connection.executemany(
        "insert into strategy_equity values (?, ?, ?, ?, ?)",
        [
            (result_id, 0, now, 100, 100),
            (result_id, 1, datetime(2026, 1, 2, tzinfo=UTC), 110, 110),
            (result_id, 2, datetime(2026, 1, 5, tzinfo=UTC), 110, 110),
        ],
    )
    return connection, int(result_id)


def _strategy(name: str, orders: int) -> dict[str, object]:
    return {
        "name": name,
        "exchange": {"name": "Bybit", "use_upnl": True},
        "basic": {
            "strategy": "mrs3", "symbol": "ONUSDT", "time_frame": "1h",
            "use_long": True, "use_short": False,
        },
        "mrs3": {
            "ma_long": [
                {"id": order_id, "len": 6 + order_id, "multiplier": 0.995, "lot_x": 1 / orders}
                for order_id in range(1, orders + 1)
            ],
            "ma_short": [],
            "ma_close_long": {"len": 3},
            "ma_close_short": {"len": 3},
        },
    }


def _make_inbox(tmp_path: Path) -> tuple[Path, Path, dict[Path, bytes]]:
    inbox = tmp_path / "inbox"
    strategies = inbox / "strategies"
    reports = tmp_path / "tester-reports"
    strategies.mkdir(parents=True)
    reports.mkdir()
    report = FIXTURE.read_bytes()
    entries: list[dict[str, object]] = []
    diagnostics: dict[str, object] = {}
    for index, (name, orders) in enumerate((("P1", 1), ("P2", 2)), start=1):
        strategy = _strategy(name, orders)
        strategy_bytes = json.dumps(strategy, separators=(",", ":")).encode()
        (strategies / f"{name}.json").write_bytes(strategy_bytes)
        (reports / f"{name}.html").write_bytes(report)
        candidate = f"candidate-{index}"
        diagnostics[candidate] = {
            "order_count": orders,
            "orders": [
                {
                    "order_id": order_id,
                    "plateau_id": "P1" if order_id == 1 else "P2",
                    "plateau_point_count": 4,
                    "base_point_trades": 20,
                    "plateau_total_trades": 80,
                }
                for order_id in range(1, orders + 1)
            ],
        }
        entries.append({
            "manifest_entry_id": f"{index:032x}",
            "strategy_name": name,
            "strategy_version_id": sha256(
                json.dumps(strategy, sort_keys=True, separators=(",", ":")).encode()
            ).hexdigest(),
            "strategy_path": str((strategies / f"{name}.json").resolve()),
            "report_path": str((reports / f"{name}.html").resolve()),
            "wizard_run_id": "run-1",
            "exchange_name": "Bybit",
            "source_strategy_sha256": sha256(strategy_bytes).hexdigest(),
            "source_report_sha256": sha256(report).hexdigest(),
        })
    commission_contract = {
        "MakerFee": "0.0002", "TakerFee": "0.0004", "SlippagePercent": "0.01",
        "FundingRate": "0.0001", "FundingIntervalHours": "8",
    }
    manifest = {
        "schema_version": 1,
        "batch_id": "panel-v2-test",
        "expected_strategy_names": ["P1", "P2"],
        "tester_config_sha256": "t" * 64,
        "commission_contract": commission_contract,
        "commission_contract_id": sha256(json.dumps(commission_contract, sort_keys=True, separators=(",", ":")).encode()).hexdigest(),
        "run_mode": "FAST",
        "test_start": "2026-01-01",
        "test_end": "2026-01-09",
        "entries": entries,
        "v6_provenance": {
            "analysis_run_id": "a" * 64,
            "generation_manifest_sha256": "g" * 64,
            "strategy_json_sha256": {f"{name}.json": entries[index]["strategy_version_id"] for index, name in enumerate(("P1", "P2"))},
            "candidate_identity_to_strategy_names": {
                "candidate-1": ["P1"], "candidate-2": ["P2"],
            },
            "candidate_diagnostics": diagnostics,
        },
    }
    manifest_path = inbox / "inbox_manifest.json"
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    return inbox, reports, {
        path.relative_to(inbox): path.read_bytes()
        for path in inbox.rglob("*") if path.is_file()
    }


def _request(tmp_path: Path) -> tuple[PerformanceV2PanelRequest, dict[Path, bytes]]:
    inbox, reports, snapshot = _make_inbox(tmp_path)
    config = PerformanceV2Config(tmp_path / "performance-v2", workers=2)
    target = performance_v2_database_path(config)
    target.parent.mkdir(parents=True)
    with duckdb.connect(str(target)) as connection:
        initialize_performance_v2(connection)
    dates = tmp_path / "Input" / "dates.xlsx"
    dates.parent.mkdir()
    workbook = Workbook()
    workbook.active.append(["ONUSDT", datetime(2025, 12, 25)])
    workbook.save(dates)
    return PerformanceV2PanelRequest(
        inbox=inbox,
        report_root=reports,
        config=config,
        listing_dates_path=Path("Input/dates.xlsx"),
    ), snapshot


def test_v2_panel_service_imports_committed_inbox_without_eager_window_calculation(tmp_path: Path) -> None:
    request, before = _request(tmp_path)
    result = LocalPerformanceV2Service().run(request)

    assert result.status == "COMMITTED"
    assert result.imported_count == 2
    assert result.skipped_count == result.rejected_count == 0
    assert result.strategy_count == 2
    assert result.order_count == 3
    assert result.plateau_count == 2
    assert result.result_count == 2
    assert result.audit_path is not None and result.audit_path.is_file()
    with duckdb.connect(str(request.config.database_root / "strategy_performance.duckdb"), read_only=True) as connection:
        assert connection.execute("select count(*) from window_metrics").fetchone() == (0,)
        assert connection.execute("select count(*) from strategies").fetchone() == (2,)
        assert connection.execute("select count(*) from strategy_orders").fetchone() == (3,)
        assert connection.execute("select count(*) from strategy_results").fetchone() == (2,)
    assert {
        path.relative_to(request.inbox): path.read_bytes()
        for path in request.inbox.rglob("*") if path.is_file()
    } == before


def test_v2_panel_service_carries_phase_evidence_into_readback_progress(tmp_path: Path) -> None:
    request, _ = _request(tmp_path)
    target = performance_v2_database_path(request.config)
    imported = PerformanceV2ImportResult(
        "import-1", "COMMITTED", 1, 0, 0, target, None,
        {"PUBLISH_ROWS": 0.125, "COMMIT": 0.25},
    )
    progress: list[object] = []
    service = LocalPerformanceV2Service(import_func=lambda _request, **_kwargs: imported)

    service.run(request, progress=progress.append)

    verified = next(item for item in progress if isinstance(item, dict) and item.get("stage") == "READBACK_VERIFIED")
    phases = verified["evidence"]["phase_seconds"]
    assert phases["PUBLISH_ROWS"] == 0.125
    assert phases["COMMIT"] == 0.25
    assert "PANEL_READBACK" in phases


@pytest.mark.parametrize(
    "phases",
    [
        {"UNKNOWN_PHASE": 0.1},
        {f"PUBLISH_ROWS_{index}": 0.1 for index in range(16)},
    ],
)
def test_v2_panel_service_omits_malformed_phase_evidence(tmp_path: Path, phases: dict[str, float]) -> None:
    request, _ = _request(tmp_path)
    target = performance_v2_database_path(request.config)
    imported = PerformanceV2ImportResult("import-1", "COMMITTED", 1, 0, 0, target, None, phases)
    progress: list[object] = []

    LocalPerformanceV2Service(import_func=lambda _request, **_kwargs: imported).run(
        request, progress=progress.append,
    )

    verified = next(item for item in progress if isinstance(item, dict) and item.get("stage") == "READBACK_VERIFIED")
    assert "evidence" not in verified


def test_v2_panel_service_keeps_readback_timing_without_importer_phases(tmp_path: Path) -> None:
    request, _ = _request(tmp_path)
    target = performance_v2_database_path(request.config)
    imported = PerformanceV2ImportResult("import-1", "COMMITTED", 1, 0, 0, target, None, None)
    progress: list[object] = []

    LocalPerformanceV2Service(import_func=lambda _request, **_kwargs: imported).run(
        request, progress=progress.append,
    )

    verified = next(item for item in progress if isinstance(item, dict) and item.get("stage") == "READBACK_VERIFIED")
    assert set(verified["evidence"]["phase_seconds"]) == {"PANEL_READBACK"}


def test_v2_panel_service_resolves_listing_dates_from_trusted_project_root(tmp_path: Path) -> None:
    request, _ = _request(tmp_path)
    project_root = tmp_path / "project-root"
    dates = project_root / "dates.xlsx"
    dates.parent.mkdir(parents=True)
    workbook = Workbook()
    workbook.active.append(["ONUSDT", datetime(2025, 12, 25)])
    workbook.save(dates)
    request = replace(
        request,
        listing_dates_path=Path("dates.xlsx"),
        listing_dates_root=project_root,
    )

    result = LocalPerformanceV2Service().run(request)

    assert result.status == "COMMITTED"
    assert result.imported_count == 2


def test_v2_panel_service_falls_back_to_inbox_parent_for_legacy_listing_path(tmp_path: Path) -> None:
    request, _ = _request(tmp_path)
    request = replace(request, listing_dates_root=tmp_path / "missing-project-root")

    result = LocalPerformanceV2Service().run(request)

    assert result.status == "COMMITTED"
    assert result.imported_count == 2


def test_v2_all_rejected_import_does_not_cleanup_tester_sources(tmp_path: Path) -> None:
    request, _ = _request(tmp_path)
    manifest_path = request.inbox / "inbox_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    reports = []
    for index, name in enumerate(("P1", "P2")):
        report = request.report_root / f"{name}.html"
        report.write_text("<html>invalid</html>", encoding="utf-8")
        manifest["entries"][index]["source_report_sha256"] = sha256(report.read_bytes()).hexdigest()
        reports.append(report)
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    request = replace(request, strategy_root=request.inbox / "strategies")

    result = LocalPerformanceV2Service().run(request)

    assert result.status == "FAILED"
    assert result.imported_count == 0
    assert result.rejected_count == 2
    assert result.failure_report_path is not None and result.failure_report_path.is_file()
    assert all(report.is_file() for report in reports)
    assert (request.strategy_root / "P1.json").is_file()  # type: ignore[union-attr]


def test_v2_failed_import_without_database_keeps_sources(tmp_path: Path) -> None:
    request, _ = _request(tmp_path)
    result = LocalPerformanceV2Service(
        import_func=lambda _request, **_kwargs: PerformanceV2ImportResult(
            "failed", "FAILED", 0, 0, 1, None, None
        )
    ).run(request)

    assert result.status == "FAILED"
    assert (request.report_root / "P1.html").is_file()


def test_v2_panel_service_passes_frozen_result_ids_directly(tmp_path: Path) -> None:
    request, _ = _request(tmp_path)
    request = replace(
        request, mode="REPLACE", replacement_strategy_ids={"P1": 1},
        expected_current_result_ids={"P1": 11},
    )
    captured = {}

    def import_result(import_request, **_kwargs):
        captured["expected"] = import_request.expected_current_result_ids
        return PerformanceV2ImportResult("failed", "FAILED", 0, 0, 1, None, None)

    LocalPerformanceV2Service(import_func=import_result).run(request)

    assert captured["expected"] == {"P1": 11}


def test_v2_panel_service_real_importer_isolates_frozen_result_mismatch(tmp_path: Path) -> None:
    request, _ = _request(tmp_path)
    assert LocalPerformanceV2Service().run(request).imported_count == 2
    target = performance_v2_database_path(request.config)
    with duckdb.connect(str(target), read_only=True) as connection:
        old = {
            str(name): (int(strategy_id), int(result_id), final_balance)
            for name, strategy_id, result_id, final_balance in connection.execute(
                "select s.strategy_name, s.strategy_id, s.current_result_id, r.final_balance "
                "from strategies s join strategy_results r using (strategy_id) order by s.strategy_name"
            ).fetchall()
        }

    changed_report = FIXTURE.read_bytes().replace(b"1009.9", b"1019.9")
    manifest_path = request.inbox / "inbox_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    for entry in manifest["entries"]:
        Path(entry["report_path"]).write_bytes(changed_report)
        entry["source_report_sha256"] = sha256(changed_report).hexdigest()
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

    result = LocalPerformanceV2Service().run(replace(
        request,
        mode="REPLACE",
        replacement_strategy_ids={name: identity[0] for name, identity in old.items()},
        expected_current_result_ids={"P1": old["P1"][1] + 1, "P2": old["P2"][1]},
    ))

    assert result.status == "COMMITTED"
    assert result.imported_count == 1
    assert result.rejected_count == 1
    assert result.failures == ({
        "strategy_id": old["P1"][0],
        "strategy_name": "P1",
        "symbol": "ONUSDT",
        "reason": "STALE_RESULT",
    },)
    assert result.failure_report_path is not None
    assert "STALE_RESULT" in result.failure_report_path.read_text(encoding="utf-8")
    with duckdb.connect(str(target), read_only=True) as connection:
        rows = connection.execute(
            "select s.strategy_name, s.current_result_id, r.final_balance "
            "from strategies s join strategy_results r using (strategy_id)"
        ).fetchall()
        result_counts = dict(connection.execute(
            "select s.strategy_name, count(*) from strategies s "
            "join strategy_results r using (strategy_id) group by s.strategy_name"
        ).fetchall())
    current = {str(name): (int(result_id), balance) for name, result_id, balance in rows}
    assert current["P1"] == (old["P1"][1], old["P1"][2])
    assert result_counts == {"P1": 1, "P2": 1}
    assert current["P2"][0] == old["P2"][1]
    assert current["P2"][1] != old["P2"][2]


@pytest.mark.parametrize(
    "private_controls",
    [
        {"_retest": True},
        {"_expected_current_result_ids": {"P1": 11}},
    ],
)
def test_v2_panel_controller_rejects_external_retest_import_controls(
    tmp_path: Path, private_controls: dict[str, object],
) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")

    with pytest.raises(ValueError, match="internal only"):
        controller.strategies_performance_v2_import({"tester_job_id": "unused", **private_controls})


def test_v2_panel_controller_rejects_external_import_job_id(tmp_path: Path) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")

    with pytest.raises(ValueError, match="job ID is internal only"):
        controller.strategies_performance_v2_import({"tester_job_id": "unused"}, job_id="caller-chosen")


def test_v2_panel_service_passes_collection_consumption_expectations(tmp_path: Path) -> None:
    request, _ = _request(tmp_path)
    request = replace(
        request,
        expected_inbox_manifest_sha256="a" * 64,
        expected_collection_id="collection-1",
    )
    captured = {}

    def import_result(import_request, **_kwargs):
        captured["digest"] = import_request.expected_inbox_manifest_sha256
        captured["collection_id"] = import_request.expected_collection_id
        return PerformanceV2ImportResult("failed", "FAILED", 0, 0, 1, None, None)

    LocalPerformanceV2Service(import_func=import_result).run(request)

    assert captured == {"digest": "a" * 64, "collection_id": "collection-1"}


@pytest.mark.parametrize(
    "listing_path",
    [Path("../dates.xlsx"), Path("C:/absolute/dates.xlsx")],
)
def test_v2_panel_service_rejects_unsafe_listing_dates_paths(tmp_path: Path, listing_path: Path) -> None:
    request, _ = _request(tmp_path)
    request = replace(request, listing_dates_path=listing_path, listing_dates_root=tmp_path)

    with pytest.raises(ValueError, match="relative"):
        LocalPerformanceV2Service().run(request)


def test_v2_panel_service_rejects_symlinked_listing_dates_path(tmp_path: Path) -> None:
    request, _ = _request(tmp_path)
    project_root = tmp_path / "project-root"
    project_root.mkdir()
    target = tmp_path / "external-dates.xlsx"
    target.write_bytes(b"not a workbook")
    link = project_root / "dates.xlsx"
    try:
        link.symlink_to(target)
    except OSError:
        pytest.skip("symlink creation is unavailable")
    request = replace(request, listing_dates_path=Path("dates.xlsx"), listing_dates_root=project_root)

    with pytest.raises(PerformanceV2ImportError, match="symlink"):
        LocalPerformanceV2Service().run(request)


def test_v2_panel_service_rejects_listing_dates_resolved_outside_trusted_roots(tmp_path: Path) -> None:
    request, _ = _request(tmp_path)
    project_root = tmp_path / "project-root"
    project_root.mkdir()
    external_dir = tmp_path / "external"
    external_dir.mkdir()
    (external_dir / "dates.xlsx").write_bytes(b"not a workbook")
    alias = project_root / "alias"
    try:
        alias.symlink_to(external_dir, target_is_directory=True)
    except OSError:
        pytest.skip("directory symlink creation is unavailable")
    request = replace(
        request,
        listing_dates_path=Path("alias/dates.xlsx"),
        listing_dates_root=project_root,
    )

    with pytest.raises(PerformanceV2ImportError, match="outside the trusted input root"):
        LocalPerformanceV2Service().run(request)


def test_visible_performance_card_targets_only_v2_import_job_and_status() -> None:
    panel_web = Path(__file__).parents[1] / "src" / "mrs3" / "panel_web"
    html = (panel_web / "index.html").read_text(encoding="utf-8")
    js = (panel_web / "app.js").read_text(encoding="utf-8")
    card = html.split("3. Test and Import to Performance DB", 1)[1].split("</details>", 1)[0]
    handler = js.split("importStartV2?.addEventListener", 1)[1].split("const recoverSplitJobs", 1)[0]
    recovery = js.split("const recoverSplitJobs = async", 1)[1].split("const settingsStatus", 1)[0]
    v2_slice = js.split("const importStartV2", 1)[1].split("const settingsStatus", 1)[0]

    assert 'id="performance-import-start"' in card
    assert "delete-tested-html-v2" not in card
    assert "performance-db-refresh" not in card
    assert "performance-audit-open" not in card
    assert 'id="strategy-dd5-card-v2"' not in html
    assert "strategies.performance.v2.import" in handler
    assert "/api/v2/strategies/performance-v2/import/status" in handler
    assert "strategies.performance.import" not in handler
    assert "strategies.dd5.start" not in handler
    assert "dd5-workbook" not in handler
    assert "strategies.performance.v2.import" in recovery
    assert "/api/v2/strategies/performance-v2/import/status" in recovery
    assert "refreshPerformanceCatalog();" not in v2_slice


def test_v2_panel_controller_initializes_missing_target_and_uses_committed_tester_job(
    tmp_path: Path, monkeypatch,
) -> None:
    import mrs3.panel as panel_module

    request, _ = _request(tmp_path)
    config_path = tmp_path / "config.local.json"
    config_path.write_text(json.dumps({
        "panel_paths": {"tester_report_dir": "wrong-reports"},
        "panel_workflow": {"listing_dates_path": "Input/dates.xlsx"},
    }), encoding="utf-8")
    (tmp_path / "config.performance.json").write_text(
        json.dumps({"unified_performance_v2": {"database_root": "performance-v2", "workers": 1}}),
        encoding="utf-8",
    )
    monkeypatch.setattr(panel_module.RunnerConfig, "from_json", lambda _path: type("Runner", (), {"report_dir": request.report_root})())
    controller = PanelController(tmp_path, config_path)
    monkeypatch.setattr(controller, "_validate_metadata_inbox", lambda _inbox: None)
    performance_v2_database_path(request.config).unlink()
    job = controller._panel_jobs.submit(
        "strategies.tester.start", {"mode": "SINGLE_MODE", "analysis_run_id": "a", "start_date": "2026-01-01", "end_date": "2026-01-09"},
        "tester-v2", ("strategies.tester",), job_id="tester-v2",
    )
    controller._panel_jobs.transition("tester-v2", "RUNNING")
    controller._panel_jobs.sync(
        "tester-v2", {"state": "COMMITTED", "phase": "COMMITTED", "inbox_ready": True},
        runtime={"inbox_path": str(request.inbox), "performance_v2_import_verified": True},
    )

    with pytest.raises(ValueError, match="unsupported fields"):
        controller.panel_job_submit({
            "kind": "strategies.performance.v2.import",
            "request": {
                "tester_job_id": "tester-v2",
                "listing_dates_path": "Input/dates.xlsx",
                "listing_dates_root": "/",
            },
        })

    started = controller.panel_job_submit({
        "kind": "strategies.performance.v2.import",
        "request": {
            "tester_job_id": "tester-v2",
            "window_a": ["2026-01-01T00:00:00Z", "2026-01-09T00:00:00Z"],
            "window_b": ["2026-01-01T00:00:00Z", "2026-01-03T12:00:00Z"],
        },
    })
    v2_job_id = started["job_id"]
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        status = controller.strategies_performance_v2_import_status(v2_job_id)
        if status["state"] in {"COMMITTED", "FAILED"}:
            break
        time.sleep(0.02)
    assert status["state"] == "COMMITTED"
    assert status["result"]["imported_count"] == 2
    assert "database_path" not in status["result"]
    assert "audit_path" not in status["result"]
    assert status["result"]["order_count"] == 3
    assert status["result"]["window_count"] == 0

    # The controller's new route is independently exposed by the HTTP server.
    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        try:
            connection.request("GET", f"/api/v2/strategies/performance-v2/import/status?job_id={v2_job_id}")
            response = connection.getresponse()
            document = json.loads(response.read())
        finally:
            connection.close()
        assert response.status == 200
        assert document["state"] == "COMMITTED"
        assert "database_path" not in document["result"]
        assert "failure_report_path" not in document["result"]

        restarted = PanelController(tmp_path, config_path)
        persisted = restarted.strategies_performance_v2_import_status(v2_job_id)
        assert persisted["state"] == "COMMITTED"
        assert persisted["result"]["imported_count"] == 2
        assert "audit_path" not in persisted["result"]
        assert "failure_report_path" not in persisted["result"]
        assert "windows" not in persisted["result"]
    finally:
        server.shutdown()
        server.server_close()


@pytest.mark.parametrize(("existing", "should_initialize"), [(b"", True), (b"foreign target", False)])
def test_v2_panel_first_import_bootstrap_initializes_empty_and_preserves_foreign_target(
    tmp_path: Path, existing: bytes, should_initialize: bool,
) -> None:
    config = PerformanceV2Config(tmp_path / "performance-v2")
    target = performance_v2_database_path(config)
    target.parent.mkdir(parents=True)
    target.write_bytes(existing)

    PanelController._initialize_missing_performance_v2_target(config)

    if should_initialize:
        with duckdb.connect(str(target), read_only=True) as connection:
            assert connection.execute(
                "select value from schema_info where key = 'schema_version'"
            ).fetchone() == ("12",)
    else:
        assert target.read_bytes() == existing


def test_v2_panel_first_import_bootstrap_preserves_unsupported_duckdb(tmp_path: Path) -> None:
    config = PerformanceV2Config(tmp_path / "performance-v2")
    target = performance_v2_database_path(config)
    target.parent.mkdir(parents=True)
    with duckdb.connect(str(target)) as connection:
        connection.execute("create table foreign_schema(value integer)")
    before = target.read_bytes()

    PanelController._initialize_missing_performance_v2_target(config)

    assert target.read_bytes() == before


def test_v2_panel_first_import_bootstrap_does_not_clobber_target_created_during_publish(
    tmp_path: Path, monkeypatch,
) -> None:
    config = PerformanceV2Config(tmp_path / "performance-v2")
    target = performance_v2_database_path(config)
    real_link = panel_module.os.link

    def competing_publish(source: Path, destination: Path) -> None:
        Path(destination).write_bytes(b"foreign target")
        real_link(source, destination)

    monkeypatch.setattr(panel_module.os, "link", competing_publish)

    PanelController._initialize_missing_performance_v2_target(config)

    assert target.read_bytes() == b"foreign target"


def test_v2_panel_first_import_bootstrap_rejects_dangling_target_symlink(tmp_path: Path) -> None:
    config = PerformanceV2Config(tmp_path / "performance-v2")
    lexical_target = config.database_root / "strategy_performance.duckdb"
    lexical_target.parent.mkdir(parents=True)
    try:
        lexical_target.symlink_to(config.database_root / "missing.duckdb")
    except OSError:
        pytest.skip("file symlink creation is unavailable")

    with pytest.raises(PerformanceV2StoreError, match="redirected"):
        PanelController._initialize_missing_performance_v2_target(config)

    assert lexical_target.is_symlink()
    assert not (config.database_root / "missing.duckdb").exists()


def test_v2_panel_first_import_bootstrap_rejects_redirected_canonical_entry(
    tmp_path: Path, monkeypatch,
) -> None:
    config = PerformanceV2Config(tmp_path / "performance-v2")
    lexical_target = config.database_root / "strategy_performance.duckdb"
    real_is_symlink = Path.is_symlink
    monkeypatch.setattr(
        Path,
        "is_symlink",
        lambda path: path == lexical_target or real_is_symlink(path),
    )

    with pytest.raises(PerformanceV2StoreError, match="redirected"):
        PanelController._initialize_missing_performance_v2_target(config)

    assert not lexical_target.exists()


def test_v2_panel_controller_injects_server_owned_listing_root(tmp_path: Path, monkeypatch) -> None:
    import mrs3.panel as panel_module

    request, _ = _request(tmp_path)
    config_path = tmp_path / "config.local.json"
    config_path.write_text(json.dumps({"panel_workflow": {"listing_dates_path": "input/dates.xlsx"}}), encoding="utf-8")
    (tmp_path / "config.performance.json").write_text(
        json.dumps({"unified_performance_v2": {"database_root": "performance-v2", "workers": 1}}),
        encoding="utf-8",
    )
    monkeypatch.setattr(panel_module.RunnerConfig, "from_json", lambda _path: type("Runner", (), {"report_dir": request.report_root})())
    controller = PanelController(tmp_path, config_path)
    monkeypatch.setattr(controller, "_validate_metadata_inbox", lambda _inbox: None)
    controller._panel_jobs.submit(
        "strategies.tester.start", {"mode": "SINGLE_MODE"}, "tester-root", (), job_id="tester-root"
    )
    controller._panel_jobs.transition("tester-root", "RUNNING")
    controller._panel_jobs.sync(
        "tester-root",
        {"state": "COMMITTED", "phase": "COMMITTED", "inbox_ready": True},
        runtime={"inbox_path": str(request.inbox), "performance_v2_import_verified": True},
    )
    captured: dict[str, object] = {}

    class StubJobs:
        def start(self, panel_request: PerformanceV2PanelRequest, *, job_id: str | None = None) -> dict[str, object]:
            captured["request"] = panel_request
            return {"job_id": job_id, "state": "COMMITTED"}

    controller._performance_v2_jobs = StubJobs()  # type: ignore[assignment]
    controller.strategies_performance_v2_import({"tester_job_id": "tester-root"})

    assert captured["request"].listing_dates_root == tmp_path.resolve()  # type: ignore[union-attr]
    assert captured["request"].listing_dates_path == Path("input/dates.xlsx")  # type: ignore[union-attr]
    assert captured["request"].strategy_root == (tmp_path / "Output").resolve()  # type: ignore[union-attr]
    assert controller._panel_jobs.runtime("tester-root")["performance_v2_import_verified"] is False
    with pytest.raises(ValueError, match="explicit inbox verification"):
        controller.strategies_performance_v2_import({"tester_job_id": "tester-root"})
    runtime = controller._panel_jobs.runtime("tester-root")
    runtime["performance_v2_import_verified"] = True
    controller._panel_jobs.sync("tester-root", {"state": "COMMITTED"}, runtime=runtime)
    with pytest.raises(ValueError, match="listing_dates_path is configured by the server"):
        controller.strategies_performance_v2_import({"tester_job_id": "tester-root", "listing_dates_path": "override.xlsx"})


def _write_fresh_metadata_inbox(
    tmp_path: Path,
    *,
    analysis_id: str,
    batch_name: str,
) -> tuple[Path, Path, Path]:
    batch = tmp_path / "Output" / "fresh-shortlist-v2" / analysis_id / batch_name
    strategies = batch / "strategies"
    strategies.mkdir(parents=True)
    strategy = {"name": "S1"}
    strategy_bytes = json.dumps(
        strategy, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
    ).encode("utf-8")
    strategy_path = strategies / "S1.json"
    strategy_path.write_bytes(strategy_bytes)
    generation = {
        "format_version": 1,
        "analysis_run_id": analysis_id,
        "event_mode": "real_independent_events",
        "strategy_count": 1,
        "strategy_json_sha256": {"S1.json": sha256(strategy_bytes).hexdigest()},
    }
    generation["generation_manifest_sha256"] = sha256(json.dumps(
        generation, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
    ).encode("utf-8")).hexdigest()
    (batch / "strategy_manifest.json").write_text(
        json.dumps(generation), encoding="utf-8",
    )

    report_root = tmp_path / "bot" / "tester" / "report" / "my_test"
    report_root.mkdir(parents=True)
    (report_root / "S1.html").write_text("report", encoding="utf-8")
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    (inbox / "inbox_manifest.json").write_text(json.dumps({
        "schema_version": 1,
        "run_mode": "SINGLE_MODE",
        "source_mode": "metadata_only",
        "inbox_ready": True,
        "expected_strategy_names": ["S1"],
        "entries": [{
            "strategy_name": "S1",
            "strategy_path": str(strategy_path),
            "report_path": "S1.html",
        }],
        "v6_provenance": {"analysis_run_id": analysis_id},
    }), encoding="utf-8")
    return inbox, report_root, strategy_path


def test_metadata_inbox_accepts_its_validated_published_fresh_batch(tmp_path: Path, monkeypatch) -> None:
    analysis_id = "a" * 64
    inbox, report_root, _strategy_path = _write_fresh_metadata_inbox(
        tmp_path,
        analysis_id=analysis_id,
        batch_name=f"{'b' * 64}-{'c' * 32}",
    )
    config = tmp_path / "config.local.json"
    config.write_text("{}", encoding="utf-8")
    controller = PanelController(tmp_path, config)
    monkeypatch.setattr(
        panel_module.RunnerConfig,
        "from_json",
        lambda _path: SimpleNamespace(
            strategy_dir=tmp_path / "bot" / "settings_strategy",
            report_dir=report_root,
        ),
    )
    monkeypatch.setattr(
        controller,
        "_performance_v2_config",
        lambda: SimpleNamespace(strategy_root=tmp_path / "Output" / "strategies"),
    )

    controller._validate_metadata_inbox(inbox)


def test_metadata_inbox_rejects_symlinked_artifact_without_creating_symlink(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    analysis_id = "a" * 64
    inbox, report_root, strategy_path = _write_fresh_metadata_inbox(
        tmp_path,
        analysis_id=analysis_id,
        batch_name=f"{'b' * 64}-{'c' * 32}",
    )
    config = tmp_path / "config.local.json"
    config.write_text("{}", encoding="utf-8")
    controller = PanelController(tmp_path, config)
    monkeypatch.setattr(
        panel_module.RunnerConfig,
        "from_json",
        lambda _path: SimpleNamespace(
            strategy_dir=tmp_path / "bot" / "settings_strategy",
            report_dir=report_root,
        ),
    )
    monkeypatch.setattr(
        controller,
        "_performance_v2_config",
        lambda: SimpleNamespace(strategy_root=tmp_path / "Output" / "strategies"),
    )
    real_is_symlink = Path.is_symlink
    inspected: list[Path] = []

    def report_symlink(path: Path) -> bool:
        inspected.append(path)
        return path == strategy_path or real_is_symlink(path)

    monkeypatch.setattr(Path, "is_symlink", report_symlink)

    with pytest.raises(ValueError, match="strategy_path is missing"):
        controller._validate_metadata_inbox(inbox)

    assert strategy_path in inspected


def test_metadata_inbox_accepts_strategy_anywhere_under_output(tmp_path: Path, monkeypatch) -> None:
    analysis_id = "a" * 64
    inbox, report_root, _strategy_path = _write_fresh_metadata_inbox(
        tmp_path,
        analysis_id=analysis_id,
        batch_name=".stage-unpublished",
    )
    config = tmp_path / "config.local.json"
    config.write_text("{}", encoding="utf-8")
    controller = PanelController(tmp_path, config)
    monkeypatch.setattr(
        panel_module.RunnerConfig,
        "from_json",
        lambda _path: SimpleNamespace(
            strategy_dir=tmp_path / "bot" / "settings_strategy",
            report_dir=report_root,
        ),
    )
    monkeypatch.setattr(
        controller,
        "_performance_v2_config",
        lambda: SimpleNamespace(strategy_root=tmp_path / "Output" / "strategies"),
    )

    controller._validate_metadata_inbox(inbox)


def test_metadata_inbox_rejects_redirected_output_root(tmp_path: Path, monkeypatch) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    output_root = tmp_path / "Output"
    real_is_junction = Path.is_junction
    monkeypatch.setattr(
        Path,
        "is_junction",
        lambda path: path == output_root or real_is_junction(path),
    )

    with pytest.raises(ValueError, match="Output strategy root is redirected"):
        controller._output_strategy_root()


def test_metadata_inbox_rejects_redirected_fresh_analysis_directory(tmp_path: Path, monkeypatch) -> None:
    analysis_id = "a" * 64
    inbox, report_root, _strategy_path = _write_fresh_metadata_inbox(
        tmp_path,
        analysis_id=analysis_id,
        batch_name=f"{'b' * 64}-{'c' * 32}",
    )
    analysis_root = tmp_path / "Output" / "fresh-shortlist-v2" / analysis_id
    external_root = tmp_path / "external-analysis"
    analysis_root.rename(external_root)
    try:
        analysis_root.symlink_to(external_root, target_is_directory=True)
    except OSError:
        pytest.skip("directory symlink creation is unavailable")

    config = tmp_path / "config.local.json"
    config.write_text("{}", encoding="utf-8")
    controller = PanelController(tmp_path, config)
    monkeypatch.setattr(
        panel_module.RunnerConfig,
        "from_json",
        lambda _path: SimpleNamespace(
            strategy_dir=tmp_path / "bot" / "settings_strategy",
            report_dir=report_root,
        ),
    )
    monkeypatch.setattr(
        controller,
        "_performance_v2_config",
        lambda: SimpleNamespace(strategy_root=tmp_path / "Output" / "strategies"),
    )

    with pytest.raises(ValueError, match="strategy_path is outside configured directory"):
        controller._validate_metadata_inbox(inbox)


def test_v2_panel_controller_rejects_committed_job_without_verified_inbox(tmp_path: Path) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    controller._panel_jobs.submit(
        "strategies.tester.start", {"mode": "SINGLE_MODE"}, "tester", (), job_id="tester"
    )
    controller._panel_jobs.transition("tester", "RUNNING")
    controller._panel_jobs.sync("tester", {"state": "COMMITTED", "phase": "COMMITTED"})

    with pytest.raises(ValueError, match="committed tester inbox"):
        controller.strategies_performance_v2_import({"tester_job_id": "tester"})


def test_v2_panel_controller_requires_explicit_verify_before_import(tmp_path: Path) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    controller._panel_jobs.submit(
        "strategies.tester.start", {"mode": "SINGLE_MODE"}, "tester", (), job_id="tester"
    )
    controller._panel_jobs.transition("tester", "RUNNING")
    controller._panel_jobs.sync(
        "tester", {"state": "COMMITTED", "phase": "COMMITTED", "inbox_ready": True},
        runtime={"inbox_path": str(tmp_path / "inbox" )},
    )

    with pytest.raises(ValueError, match="explicit inbox verification"):
        controller.strategies_performance_v2_import({"tester_job_id": "tester"})


def test_v2_failed_import_keeps_failure_report_available(tmp_path: Path) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    job_id = "failed-import"
    report = database.parent / "performance_v2_failures_failed.csv"
    report.write_text("reason\nINVALID_REPORT\n", encoding="utf-8")
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {}, f"panel:{job_id}",
        ("performance-v2-db",), job_id=job_id,
    )
    controller._panel_jobs.transition(job_id, "RUNNING")
    document = {
        "job_id": job_id,
        "state": "FAILED",
        "error": {"code": "PERFORMANCE_V2_IMPORT_FAILED", "message": "all reports rejected"},
        "result": {
            "status": "FAILED",
            "imported_count": 0,
            "rejected_count": 2,
            "failure_report_path": str(report),
        },
    }
    controller._record_special_job(document)

    assert controller.artifact(f"performance-v2-failure-report:{job_id}") == report.resolve()


def test_terminal_performance_callback_skips_three_identical_polls_and_saves_changed_payload(tmp_path, monkeypatch, caplog):
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    job_id = "terminal-performance"
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {}, f"panel:{job_id}", (), job_id=job_id,
    )
    controller._panel_jobs.transition(job_id, "RUNNING")
    caplog.set_level("INFO", logger="mrs3.panel")
    controller._record_special_job({"job_id": job_id, "state": "RUNNING", "phase": "PUBLISHING"})
    assert not [record for record in caplog.records if "PANEL_TERMINAL_SYNC" in record.message]
    document = {
        "job_id": job_id,
        "state": "COMMITTED",
        "phase": "COMMITTED",
        "inbox_path": str(tmp_path / "private" / "inbox"),
        "result": {
            "status": "COMMITTED", "imported_count": 1,
            "database_path": str(tmp_path / "private" / "db.duckdb"),
            "audit_path": str(tmp_path / "private" / "audit.json"),
            "failure_report_path": str(tmp_path / "private" / "failures.csv"),
        },
    }
    replacements = []
    temporary_sources = []
    real_replace = panel_jobs_module.os.replace
    journal = controller._panel_jobs.journal

    def counted_replace(source, destination):
        temporary_sources.append(source)
        if Path(destination) == journal:
            replacements.append(destination)
        return real_replace(source, destination)

    monkeypatch.setattr(panel_jobs_module.os, "replace", counted_replace)
    controller._record_special_job(document)
    assert replacements == [journal]
    sync_records = [record for record in caplog.records if "PANEL_TERMINAL_SYNC" in record.message]
    assert len(sync_records) == 1
    assert f"job_id={job_id}" in sync_records[0].message
    duration = float(sync_records[0].message.split("duration_seconds=", 1)[1])
    assert 0 <= duration < float("inf")
    assert controller._panel_jobs.get(job_id)["result"] == {
        "status": "COMMITTED", "imported_count": 1,
        "failure_report_available": True,
        "failure_report_token": f"performance-v2-failure-report:{job_id}",
    }
    assert controller._panel_jobs.runtime(job_id) == {
        "inbox_path": str(tmp_path / "private" / "inbox"),
        "failure_report_path": str(tmp_path / "private" / "failures.csv"),
    }

    for _ in range(3):
        controller._record_special_job(document)
    assert replacements == [journal]
    assert len([record for record in caplog.records if "PANEL_TERMINAL_SYNC" in record.message]) == 1

    changed = {**document, "result": {**document["result"], "imported_count": 2}}
    controller._record_special_job(changed)
    assert replacements == [journal, journal]
    assert len([record for record in caplog.records if "PANEL_TERMINAL_SYNC" in record.message]) == 1
    reloaded = PanelController(tmp_path, tmp_path / "config.local.json")._panel_jobs
    assert reloaded.get(job_id)["result"]["imported_count"] == 2
    assert "database_path" not in reloaded.get(job_id)["result"]
    assert reloaded.runtime(job_id)["failure_report_path"] == str(tmp_path / "private" / "failures.csv")
    assert all(not Path(source).exists() for source in temporary_sources)


def test_v2_worker_keeps_readback_evidence_private_until_terminal_save(tmp_path, monkeypatch) -> None:
    request, _ = _request(tmp_path)
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    job_id = "evidence-terminal"
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {}, f"panel:{job_id}", (), job_id=job_id,
    )
    controller._panel_jobs.transition(job_id, "RUNNING")
    entered = threading.Event()
    release = threading.Event()

    class BlockingService:
        def run(self, _request, *, progress=None):
            progress({
                "stage": "READBACK_VERIFIED", "completed": 1, "total": 1,
                "evidence": {"phase_seconds": {"PANEL_READBACK": 0.25}},
            })
            entered.set()
            assert release.wait(5)
            return PerformanceV2PanelResult(
                "import-1", "COMMITTED", 1, 0, 0,
                performance_v2_database_path(request.config), None,
                1, 0, 0, 1,
            )

    replacements = []
    journal = controller._panel_jobs.journal
    real_replace = panel_jobs_module.os.replace

    def counted_replace(source, destination):
        if Path(destination) == journal:
            replacements.append(destination)
        return real_replace(source, destination)

    monkeypatch.setattr(panel_jobs_module.os, "replace", counted_replace)
    jobs = LocalPerformanceV2Jobs(service=BlockingService(), on_update=controller._record_special_job)
    jobs.start(request, job_id=job_id)
    assert entered.wait(5)
    assert "evidence" not in jobs.status(job_id)
    assert replacements == []
    for _ in range(3):
        controller._record_special_job(jobs.status(job_id))
    assert replacements == []

    release.set()
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline and jobs.status(job_id)["state"] == "RUNNING":
        time.sleep(0.01)
    assert jobs.status(job_id)["state"] == "COMMITTED"
    assert jobs.status(job_id)["evidence"] == {"phase_seconds": {"PANEL_READBACK": 0.25}}
    while time.monotonic() < deadline and len(replacements) < 1:
        time.sleep(0.01)
    assert replacements == [journal]
    restored = PanelJobRegistry(journal, recover_on_load=False)
    assert restored.get(job_id)["evidence"] == {"phase_seconds": {"PANEL_READBACK": 0.25}}


@pytest.mark.parametrize(
    "kind,document",
    [
        ("strategies.performance.v2.import", {"state": "RUNNING", "phase": "PUBLISHING", "progress": {"current": 3, "total": 3}}),
        ("analysis.local", {"state": "COMMITTED", "phase": "COMMITTED"}),
    ],
)
def test_j5_import_skips_unchanged_polls_and_other_kind_still_saves_each_poll(
    tmp_path, monkeypatch, kind, document,
):
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    job_id = "unchanged-boundary"
    controller._panel_jobs.submit(kind, {}, job_id, (), job_id=job_id)
    controller._panel_jobs.transition(job_id, "RUNNING")
    journal = controller._panel_jobs.journal
    replacements = []
    real_replace = panel_jobs_module.os.replace

    def counted_replace(source, destination):
        if Path(destination) == journal:
            replacements.append(destination)
        return real_replace(source, destination)

    monkeypatch.setattr(panel_jobs_module.os, "replace", counted_replace)
    controller._record_special_job({"job_id": job_id, **document})
    first = 0 if kind == "strategies.performance.v2.import" else 1
    assert replacements == [journal] * first
    for _ in range(3 if kind == "strategies.performance.v2.import" else 1):
        controller._record_special_job({"job_id": job_id, **document})
    expected = 0 if kind == "strategies.performance.v2.import" else 2
    assert replacements == [journal] * expected


def test_running_performance_import_progress_never_serializes_the_journal(tmp_path, monkeypatch):
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    job_id = "volatile-performance-import"
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {}, job_id, (), job_id=job_id,
    )
    controller._panel_jobs.transition(job_id, "RUNNING", phase="PARSING")
    saves = []
    monkeypatch.setattr(controller._panel_jobs, "_save", lambda: saves.append(True))

    controller._record_special_job({
        "job_id": job_id,
        "state": "RUNNING",
        "phase": "PARSING",
        "progress": {"current": 136, "total": 2070, "unit": "reports"},
        "error": None,
    })
    controller._record_special_job({
        "job_id": job_id,
        "state": "RUNNING",
        "phase": "PARSING",
        "progress": {"current": 119, "total": 2070, "unit": "reports"},
        "error": None,
    })

    assert saves == []
    assert controller._panel_jobs.get(job_id)["progress"]["current"] == 136


def test_running_finalist_retest_progress_is_volatile_and_terminal_state_is_durable(tmp_path, monkeypatch):
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    job_id = "volatile-finalist-retest"
    controller._panel_jobs.submit(
        "strategies.performance.v2.finalist-retest", {}, job_id,
        ("strategies.tester", "performance-v2-finalist-retest"), job_id=job_id,
    )
    controller._panel_jobs.transition(job_id, "RUNNING", phase="BOT_RUN")
    frozen_runtime = {
        "bulk_retest": True,
        "cohort_members": [{"strategy_id": 1, "result_id": 11}],
    }
    controller._panel_jobs.sync(
        job_id, {"state": "RUNNING", "phase": "BOT_RUN"}, runtime=frozen_runtime,
    )
    saves = []
    save = controller._panel_jobs._save

    def count_save() -> None:
        saves.append(True)
        save()

    monkeypatch.setattr(controller._panel_jobs, "_save", count_save)

    for current in (71, 72, 73):
        controller._record_special_job({
            "job_id": job_id,
            "state": "RUNNING",
            "phase": "BOT_RUN",
            "progress": {"current": current, "total": 250, "unit": "reports"},
            "error": None,
            **({"runtime": frozen_runtime, "inbox_ready": False} if current == 71 else {}),
        })

    assert saves == []
    assert controller._panel_jobs.jobs[job_id]["progress"]["current"] == 73
    assert controller._panel_jobs._journal_dirty is True
    assert controller._panel_jobs.runtime(job_id) == frozen_runtime

    controller._record_special_job({
        "job_id": job_id,
        "state": "FAILED",
        "phase": "FAILED",
        "progress": {"current": 73, "total": 250, "unit": "reports"},
        "error": {"code": "SINGLE_MODE_TEST_FAILED", "message": "tester failed"},
    })

    assert saves == [True]
    assert controller._panel_jobs.get(job_id)["state"] == "FAILED"
    assert controller._panel_jobs._journal_dirty is False
    restored = PanelJobRegistry(controller._panel_jobs.journal, recover_on_load=False)
    assert restored.get(job_id)["state"] == "FAILED"
    assert restored.get(job_id)["progress"]["current"] == 73


def test_finalist_retest_commit_merges_runtime_and_only_committed_inbox_is_ready(tmp_path):
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    job_id = "runtime-finalist-retest"
    controller._panel_jobs.submit(
        "strategies.performance.v2.finalist-retest", {}, job_id,
        ("strategies.tester", "performance-v2-finalist-retest"), job_id=job_id,
    )
    controller._panel_jobs.transition(job_id, "RUNNING", phase="BOT_RUN")
    controller._panel_jobs.sync(
        job_id, {"state": "RUNNING", "phase": "BOT_RUN"},
        runtime={"bulk_retest": True, "cohort_members": [{"strategy_id": 7}]},
    )

    controller._record_special_job({
        "job_id": job_id, "state": "RUNNING", "phase": "BOT_RUN",
        "progress": {"current": 1, "total": 250, "unit": "reports"},
        "inbox_ready": True,
    })
    assert "inbox_ready" not in controller._panel_jobs.get(job_id)

    inbox_path = str(tmp_path / "inbox")
    controller._record_special_job({
        "job_id": job_id, "state": "COMMITTED", "phase": "COMMITTED",
        "mode": "SINGLE_MODE", "inbox_path": inbox_path,
        "runtime": {"import_job_id": "import-1"}, "inbox_ready": True,
    })

    stored = controller._panel_jobs.get(job_id)
    runtime = controller._panel_jobs.runtime(job_id)
    assert stored["inbox_ready"] is True
    assert runtime["inbox_path"] == inbox_path
    assert runtime["mode"] == "SINGLE_MODE"
    assert runtime["import_job_id"] == "import-1"
    assert runtime["cohort_members"] == [{"strategy_id": 7}]

    failed_id = "failed-finalist-inbox"
    controller._panel_jobs.submit(
        "strategies.performance.v2.finalist-retest", {}, failed_id,
        ("strategies.tester", "performance-v2-finalist-retest"), job_id=failed_id,
    )
    controller._panel_jobs.transition(failed_id, "RUNNING", phase="BOT_RUN")
    controller._record_special_job({
        "job_id": failed_id, "state": "FAILED", "phase": "FAILED",
        "progress": {"current": 2, "total": 250, "unit": "reports"},
        "inbox_ready": True,
    })
    assert "inbox_ready" not in controller._panel_jobs.get(failed_id)


def test_terminal_finalist_callback_does_not_keep_transient_progress_warning(tmp_path):
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    job_id = "terminal-clears-progress-warning"
    controller._panel_jobs.submit(
        "strategies.performance.v2.finalist-retest", {}, "terminal-warning",
        ("strategies.tester",), job_id=job_id,
    )
    controller._panel_jobs.transition(job_id, "RUNNING")
    controller._record_special_job({
        "job_id": job_id,
        "state": "RUNNING",
        "phase": "BOT_RUN",
        "progress": {"current": 1, "total": 2, "unit": "reports"},
        "progress_publication_error": "temporary journal failure",
    })
    assert controller._panel_jobs.get(job_id)["progress"]["publication_error"] == "temporary journal failure"

    controller._record_special_job({
        "job_id": job_id,
        "state": "COMMITTED",
        "phase": "COMMITTED",
        "progress": {"current": 2, "total": 2, "unit": "reports"},
        "progress_publication_error": "temporary journal failure",
    })

    assert "publication_error" not in controller._panel_jobs.get(job_id)["progress"]


def test_tester_reconciliation_runs_once_during_controller_startup(tmp_path, monkeypatch):
    calls = []
    original = PanelController._reconcile_interrupted_tester_jobs

    def counted(controller):
        calls.append(controller)
        original(controller)

    monkeypatch.setattr(PanelController, "_reconcile_interrupted_tester_jobs", counted)
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    assert calls == [controller]

    controller._panel_jobs.submit("strategies.performance.v2.finalist-retest", {}, "live-after-startup")
    assert calls == [controller]


def test_j5_nonterminal_import_callback_reloads_full_snapshot_and_skips_repeats(tmp_path, monkeypatch):
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    job_id = "running-import"
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {}, job_id, (), job_id=job_id,
    )
    controller._panel_jobs.transition(job_id, "RUNNING")
    document = {
        "job_id": job_id,
        "state": "RUNNING",
        "phase": "PUBLISHING",
        "progress": {"current": 3, "total": 5, "unit": "items"},
        "error": {"code": "TRANSIENT", "message": "retrying"},
        "evidence": {"source_revision": "r1"},
        "result": {"status": "RUNNING", "imported_count": 0},
        "inbox_path": str(tmp_path / "inbox"),
    }
    journal = controller._panel_jobs.journal
    replacements = []
    temporary_sources = []
    real_replace = panel_jobs_module.os.replace

    def counted_replace(source, destination):
        temporary_sources.append(source)
        if Path(destination) == journal:
            replacements.append(destination)
        return real_replace(source, destination)

    monkeypatch.setattr(panel_jobs_module.os, "replace", counted_replace)
    controller._record_special_job(document)
    for _ in range(3):
        controller._record_special_job(document)

    assert replacements == [journal]
    restored = PanelJobRegistry(journal, recover_on_load=False)
    assert restored.jobs == controller._panel_jobs.jobs
    assert restored.jobs[job_id]["phase"] == "PUBLISHING"
    assert all(not Path(source).exists() for source in temporary_sources)


def test_j5_state_only_import_transition_skips_identical_repeats(tmp_path, monkeypatch):
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    job_id = "state-only-import"
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {}, job_id, (), job_id=job_id,
    )
    controller._panel_jobs.transition(job_id, "RUNNING")
    document = {"job_id": job_id, "state": "COMMITTED", "phase": "COMMITTED"}
    journal = controller._panel_jobs.journal
    replacements = []
    temporary_sources = []
    real_replace = panel_jobs_module.os.replace

    def counted_replace(source, destination):
        temporary_sources.append(source)
        if Path(destination) == journal:
            replacements.append(destination)
        return real_replace(source, destination)

    monkeypatch.setattr(panel_jobs_module.os, "replace", counted_replace)
    controller._record_special_job(document)
    for _ in range(3):
        controller._record_special_job(document)

    assert replacements == [journal]
    saved = controller._panel_jobs.jobs[job_id]
    assert saved["state"] == "COMMITTED"
    assert saved["phase"] == "COMMITTED"
    assert "result" not in saved
    assert controller._panel_jobs.runtime(job_id) == {}
    assert PanelJobRegistry(journal, recover_on_load=False).jobs == controller._panel_jobs.jobs
    assert all(not Path(source).exists() for source in temporary_sources)


def test_j5_phase_progress_and_runtime_changes_each_save_once(tmp_path, monkeypatch):
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    job_id = "changed-import"
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {}, job_id, (), job_id=job_id,
    )
    controller._panel_jobs.transition(job_id, "RUNNING")
    journal = controller._panel_jobs.journal
    replacements = []
    temporary_sources = []
    real_replace = panel_jobs_module.os.replace

    def counted_replace(source, destination):
        temporary_sources.append(source)
        if Path(destination) == journal:
            replacements.append(destination)
        return real_replace(source, destination)

    monkeypatch.setattr(panel_jobs_module.os, "replace", counted_replace)
    statuses = (
        {"job_id": job_id, "state": "RUNNING", "phase": "PUBLISHING"},
        {"job_id": job_id, "state": "RUNNING", "phase": "PUBLISHING", "progress": {"current": 1, "total": 2}},
        {"job_id": job_id, "state": "RUNNING", "phase": "PUBLISHING", "progress": {"current": 1, "total": 2}, "inbox_path": str(tmp_path / "inbox")},
    )
    for document in statuses:
        controller._record_special_job(document)
        count = len(replacements)
        controller._record_special_job(document)
        assert len(replacements) == count

    assert len(replacements) == 1
    assert controller._panel_jobs.runtime(job_id) == {"inbox_path": str(tmp_path / "inbox")}
    assert PanelJobRegistry(journal, recover_on_load=False).jobs == controller._panel_jobs.jobs
    assert all(not Path(source).exists() for source in temporary_sources)


def test_j5_missing_error_and_evidence_save_once_and_retain_runtime(tmp_path, monkeypatch):
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    job_id = "normalized-import"
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {}, job_id, (), job_id=job_id,
    )
    controller._panel_jobs.transition(job_id, "RUNNING")
    controller._panel_jobs.sync(
        job_id,
        {"state": "RUNNING", "phase": "PUBLISHING", "inbox_ready": True},
        runtime={"existing": "value"},
    )
    initial = {
        "job_id": job_id,
        "state": "RUNNING",
        "phase": "PUBLISHING",
        "progress": {"current": 2, "total": 3},
        "error": {"code": "RETRY"},
        "evidence": {"source_revision": "r1"},
        "inbox_ready": True,
        "inbox_path": str(tmp_path / "inbox"),
    }
    sparse = {"job_id": job_id, "state": "RUNNING", "phase": "PUBLISHING"}
    journal = controller._panel_jobs.journal
    replacements = []
    temporary_sources = []
    real_replace = panel_jobs_module.os.replace

    def counted_replace(source, destination):
        temporary_sources.append(source)
        if Path(destination) == journal:
            replacements.append(destination)
        return real_replace(source, destination)

    monkeypatch.setattr(panel_jobs_module.os, "replace", counted_replace)
    controller._record_special_job(initial)
    assert len(replacements) == 1
    controller._record_special_job(sparse)
    assert len(replacements) == 1
    for _ in range(3):
        controller._record_special_job(sparse)

    saved = controller._panel_jobs.jobs[job_id]
    assert len(replacements) == 1
    assert saved["error"] == initial["error"]
    assert saved["evidence"] == initial["evidence"]
    assert "inbox_ready" not in saved
    assert controller._panel_jobs.runtime(job_id) == {
        "existing": "value", "inbox_path": str(tmp_path / "inbox"),
    }
    restored = PanelJobRegistry(journal, recover_on_load=False)
    assert restored.jobs == controller._panel_jobs.jobs
    assert all(not Path(source).exists() for source in temporary_sources)


def test_j5_dirty_volatile_job_is_persisted_by_identical_import_callback(tmp_path, monkeypatch):
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    import_id = "dirty-import"
    other_id = "dirty-other"
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {}, import_id, (), job_id=import_id,
    )
    controller._panel_jobs.submit("analysis.local", {}, other_id, (), job_id=other_id)
    controller._panel_jobs.transition(import_id, "RUNNING")
    controller._panel_jobs.transition(other_id, "RUNNING")
    document = {
        "job_id": import_id,
        "state": "RUNNING",
        "phase": "PUBLISHING",
        "progress": {"current": 1, "total": 2},
    }
    controller._record_special_job(document)
    journal = controller._panel_jobs.journal
    replacements = []
    temporary_sources = []
    real_replace = panel_jobs_module.os.replace

    def counted_replace(source, destination):
        temporary_sources.append(source)
        if Path(destination) == journal:
            replacements.append(destination)
        return real_replace(source, destination)

    monkeypatch.setattr(panel_jobs_module.os, "replace", counted_replace)
    controller._panel_jobs.volatile_sync(
        other_id,
        {"state": "RUNNING", "phase": "RUNNING", "progress": {"current": 1, "total": 2}},
    )
    controller._record_special_job(document)

    assert replacements == []
    assert all(not Path(source).exists() for source in temporary_sources)


def test_j4_characterization_normal_producer_resource_key_and_false_tester_marker_are_unchanged(tmp_path, monkeypatch):
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    tester_id = "characterization-tester"
    import_id = "characterization-import"
    controller._panel_jobs.submit("strategies.tester.start", {}, "characterization-tester", (), job_id=tester_id)
    controller._panel_jobs.transition(tester_id, "RUNNING")
    controller._panel_jobs.sync(
        tester_id,
        {"state": "COMMITTED", "phase": "COMMITTED", "inbox_ready": True},
        runtime={"performance_v2_import_verified": False},
    )
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {}, "characterization-import",
        (f"tester:{tester_id}",), job_id=import_id,
    )
    assert f"tester:{tester_id}" in controller._panel_jobs.jobs[import_id]["resource_keys"]
    controller._panel_jobs.transition(import_id, "RUNNING")
    assert "request" not in controller._panel_jobs.jobs[import_id]
    tester_before = deepcopy(controller._panel_jobs.jobs[tester_id])
    tester_runtime_before = controller._panel_jobs.runtime(tester_id)
    sync_calls = []
    real_sync = controller._panel_jobs.sync

    def record_sync(job_id, status, *args, **kwargs):
        sync_calls.append((job_id, status, kwargs))
        return real_sync(job_id, status, *args, **kwargs)

    monkeypatch.setattr(controller._panel_jobs, "sync", record_sync)
    replacements = []
    real_replace = panel_jobs_module.os.replace
    monkeypatch.setattr(panel_jobs_module.os, "replace", lambda source, destination: (replacements.append(destination) if Path(destination) == controller._panel_jobs.journal else None, real_replace(source, destination))[1])

    controller._record_special_job({"job_id": import_id, "state": "COMMITTED", "phase": "COMMITTED", "result": {"status": "COMMITTED"}})

    assert controller._panel_jobs.jobs[tester_id] == tester_before
    assert controller._panel_jobs.runtime(tester_id) == tester_runtime_before
    assert [job_id for job_id, _status, _kwargs in sync_calls] == [import_id]
    assert len(replacements) == 1
    restored = PanelJobRegistry(controller._panel_jobs.journal, recover_on_load=False)
    assert restored.jobs[tester_id] == tester_before
    assert restored.runtime(tester_id) == tester_runtime_before


def test_j4_legacy_tester_marker_resets_once_on_terminal_import(tmp_path, monkeypatch):
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    tester_id = "legacy-tester"
    import_id = "legacy-import"
    controller._panel_jobs.submit("strategies.tester.start", {}, "legacy-tester", (), job_id=tester_id)
    controller._panel_jobs.transition(tester_id, "RUNNING")
    controller._panel_jobs.sync(
        tester_id,
        {"state": "COMMITTED", "phase": "COMMITTED", "inbox_ready": True},
        runtime={"performance_v2_import_verified": True},
    )
    controller._panel_jobs.submit(
        "strategies.performance.v2.import", {}, "legacy-import", (), job_id=import_id,
    )
    controller._panel_jobs.jobs[import_id]["request"] = {"tester_job_id": tester_id}
    controller._panel_jobs.transition(import_id, "RUNNING")
    tester_before = deepcopy(controller._panel_jobs.jobs[tester_id])
    tester_runtime_before = controller._panel_jobs.runtime(tester_id)
    sync_calls = []
    real_sync = controller._panel_jobs.sync

    def record_sync(job_id, status, *args, **kwargs):
        sync_calls.append((job_id, status, kwargs))
        return real_sync(job_id, status, *args, **kwargs)

    monkeypatch.setattr(controller._panel_jobs, "sync", record_sync)
    replacements = []
    real_replace = panel_jobs_module.os.replace
    monkeypatch.setattr(panel_jobs_module.os, "replace", lambda source, destination: (replacements.append(destination) if Path(destination) == controller._panel_jobs.journal else None, real_replace(source, destination))[1])

    controller._record_special_job({"job_id": import_id, "state": "COMMITTED", "phase": "COMMITTED", "result": {"status": "COMMITTED"}})

    assert controller._panel_jobs.get(tester_id) == {key: value for key, value in tester_before.items() if key != "runtime"}
    assert tester_runtime_before["performance_v2_import_verified"] is True
    assert controller._panel_jobs.runtime(tester_id)["performance_v2_import_verified"] is False
    assert [job_id for job_id, _status, _kwargs in sync_calls] == [tester_id, import_id]
    assert len(replacements) == 2
    controller._record_special_job({"job_id": import_id, "state": "COMMITTED", "phase": "COMMITTED", "result": {"status": "COMMITTED"}})
    assert [job_id for job_id, _status, _kwargs in sync_calls] == [tester_id, import_id, import_id]
    assert len(replacements) == 2
    restored = PanelJobRegistry(controller._panel_jobs.journal, recover_on_load=False)
    assert restored.get(tester_id) == {key: value for key, value in tester_before.items() if key != "runtime"}
    assert restored.runtime(tester_id)["performance_v2_import_verified"] is False


def test_j4_legacy_marker_two_save_failures_retry_with_dirty_journal(tmp_path, monkeypatch):
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    registry = controller._panel_jobs
    tester_id, import_id = "legacy-tester", "legacy-import"
    registry.submit("strategies.tester.start", {}, tester_id, (), job_id=tester_id)
    registry.transition(tester_id, "RUNNING")
    registry.sync(tester_id, {"state": "COMMITTED"}, runtime={"performance_v2_import_verified": True})
    registry.submit("strategies.performance.v2.import", {}, import_id, (), job_id=import_id)
    registry.jobs[import_id]["request"] = {"tester_job_id": tester_id}
    registry.transition(import_id, "RUNNING")
    real_save = registry._save
    attempts = 0

    def fail_twice():
        nonlocal attempts
        attempts += 1
        if attempts <= 2:
            raise OSError("injected journal write failure")
        real_save()

    monkeypatch.setattr(registry, "_save", fail_twice)
    terminal = {"job_id": import_id, "state": "COMMITTED", "phase": "COMMITTED", "result": {"status": "COMMITTED"}}
    with pytest.raises(OSError, match="injected journal write failure"):
        controller._record_special_job(terminal)
    assert PanelJobRegistry(registry.journal, recover_on_load=False).runtime(tester_id)["performance_v2_import_verified"] is True
    assert registry._journal_dirty is True
    controller._record_special_job(terminal)
    restored = PanelJobRegistry(registry.journal, recover_on_load=False)
    assert restored.runtime(tester_id)["performance_v2_import_verified"] is False
    assert restored.get(import_id)["state"] == "COMMITTED"


def _terminal_collection_controller(tmp_path: Path) -> tuple[PanelController, str, str]:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    collection_id, import_id = "collection-test", "collection-import"
    registry = controller._panel_jobs
    registry.submit("strategies.tester.collection", {}, collection_id, (), job_id=collection_id)
    registry.transition(collection_id, "RUNNING")
    registry.sync(
        collection_id,
        {"state": "COMMITTED", "phase": "COMMITTED", "inbox_ready": True},
        runtime={"import_in_progress": import_id},
    )
    registry.submit("strategies.performance.v2.import", {}, import_id, (), job_id=import_id)
    registry.jobs[import_id]["request"] = {"tester_job_id": collection_id}
    registry.transition(import_id, "RUNNING")
    controller._collection_import_claim_ids.add(import_id)
    return controller, collection_id, import_id


def test_j4_collection_finish_and_terminal_import_save_once(tmp_path, monkeypatch):
    controller, collection_id, import_id = _terminal_collection_controller(tmp_path)
    registry = controller._panel_jobs
    calls = []

    def finish(_collection_id, _import_id, *, committed):
        calls.append((_collection_id, _import_id, committed))
        runtime = registry.runtime(collection_id)
        runtime.pop("import_in_progress")
        registry.sync(collection_id, registry.get(collection_id), runtime=runtime)

    controller._report_collection_service = SimpleNamespace(finish_import=finish)
    replacements = []
    real_replace = panel_jobs_module.os.replace
    monkeypatch.setattr(panel_jobs_module.os, "replace", lambda source, destination: (replacements.append(destination) if Path(destination) == registry.journal else None, real_replace(source, destination))[1])
    terminal = {"job_id": import_id, "state": "COMMITTED", "phase": "COMMITTED", "result": {"status": "COMMITTED"}}
    controller._record_special_job(terminal)
    assert calls == [(collection_id, import_id, True)]
    assert import_id not in controller._collection_import_claim_ids
    assert registry.get(import_id)["state"] == "COMMITTED"
    assert len(replacements) == 2
    controller._record_special_job(terminal)
    assert calls == [(collection_id, import_id, True)]
    assert len(replacements) == 2


@pytest.mark.parametrize("error_type", [panel_jobs_module.PanelJobError, OSError])
def test_j4_collection_finish_error_still_publishes_terminal_import(tmp_path, error_type):
    controller, collection_id, import_id = _terminal_collection_controller(tmp_path)
    calls = []

    def fail_finish(_collection_id, _import_id, *, committed):
        calls.append((_collection_id, _import_id, committed))
        raise error_type("injected collection finish failure")

    controller._report_collection_service = SimpleNamespace(finish_import=fail_finish)
    controller._record_special_job(
        {"job_id": import_id, "state": "FAILED", "phase": "FAILED", "result": {"status": "FAILED"}}
    )
    assert calls == [(collection_id, import_id, False)]
    assert import_id not in controller._collection_import_claim_ids
    assert controller._panel_jobs.get(import_id)["state"] == "FAILED"


def test_j4_missing_tester_still_publishes_terminal_import(tmp_path):
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    import_id = "missing-tester-import"
    registry = controller._panel_jobs
    registry.submit("strategies.performance.v2.import", {}, import_id, (), job_id=import_id)
    registry.jobs[import_id]["request"] = {"tester_job_id": "missing-tester"}
    registry.transition(import_id, "RUNNING")
    controller._collection_import_claim_ids.add(import_id)
    controller._record_special_job(
        {"job_id": import_id, "state": "FAILED", "phase": "FAILED", "result": {"status": "FAILED"}}
    )
    assert import_id not in controller._collection_import_claim_ids
    assert registry.get(import_id)["state"] == "FAILED"


def _controller_for_windows(tmp_path: Path, *, legacy_v11: bool = False) -> tuple[PanelController, Path, int]:
    connection, result_id = _db(tmp_path / "data", legacy_v11=legacy_v11)
    connection.close()
    config = tmp_path / "config.local.json"
    config.write_text(json.dumps({"panel_paths": {"performance_db_root": "v1"}}), encoding="utf-8")
    (tmp_path / "config.performance.json").write_text(
        json.dumps({"unified_performance_v2": {"database_root": "data"}}), encoding="utf-8"
    )
    return PanelController(tmp_path, config), tmp_path / "data" / "strategy_performance.duckdb", result_id


def _remove_v11_catalog_additions(connection: duckdb.DuckDBPyConnection) -> None:
    connection.execute("drop index selection_review_imports_aggregate_run_uq")
    connection.execute("drop index selection_review_imports_run_imported_idx")
    connection.execute("drop table selection_aggregate_imports")
    connection.execute("drop table selection_publication_runs")
    connection.execute("drop table selection_publications")
    connection.execute(
        "create temp table __selection_review_rows_before_v11 as "
        "select * from selection_review_rows"
    )
    connection.execute("drop table selection_review_rows")
    connection.execute("alter table selection_review_imports drop column aggregate_import_id")
    connection.execute(
        """create table selection_review_rows (
            review_import_id VARCHAR NOT NULL REFERENCES selection_review_imports(review_import_id),
            strategy_id BIGINT NOT NULL,
            user_status VARCHAR NOT NULL CHECK (
                user_status IN ('FINALIST', 'RESERVE', 'ANALOG', 'FILTERED', 'REJECTED')
            ),
            user_rank INTEGER CHECK (user_rank IS NULL OR user_rank > 0),
            user_analog_of_strategy_id BIGINT,
            comment VARCHAR,
            PRIMARY KEY (review_import_id, strategy_id)
        )"""
    )
    connection.execute(
        """insert into selection_review_rows (
            review_import_id, strategy_id, user_status, user_rank,
            user_analog_of_strategy_id, comment
        ) select review_import_id, strategy_id, user_status, user_rank,
            user_analog_of_strategy_id, comment
        from __selection_review_rows_before_v11"""
    )
    connection.execute("drop table __selection_review_rows_before_v11")
    connection.execute(
        "create index selection_review_imports_run_imported_idx "
        "on selection_review_imports(selection_run_id, imported_at_utc)"
    )


def _make_v5_catalog(connection: duckdb.DuckDBPyConnection) -> None:
    from mrs3.performance_v2_store import _require_v5_catalog

    _remove_v11_catalog_additions(connection)
    connection.execute("drop table strategy_rejection_sources")
    connection.execute("alter table selection_results drop column equity_regime_json")
    connection.execute("drop table equity_quality_metrics")
    connection.execute("update schema_info set value = '5' where key = 'schema_version'")
    _require_v5_catalog(connection)


def _make_v8_catalog(connection: duckdb.DuckDBPyConnection) -> None:
    from mrs3.performance_v2_store import _require_v8_catalog

    _remove_v11_catalog_additions(connection)
    connection.execute("drop table strategy_rejection_sources")
    connection.execute("alter table selection_results drop column equity_regime_json")
    connection.execute("update schema_info set value = '8' where key = 'schema_version'")
    _require_v8_catalog(connection)


def _http_server(controller: PanelController):
    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    return server, thread


def _http_json(connection: HTTPConnection, method: str, path: str, payload: object | None = None) -> tuple[int, dict]:
    body = None if payload is None else json.dumps(payload).encode("utf-8")
    connection.request(method, path, body=body, headers={"Content-Type": "application/json"} if body is not None else {})
    response = connection.getresponse()
    return response.status, json.loads(response.read().decode("utf-8"))


def test_performance_v2_maintenance_http_preview_apply_and_reused_token(tmp_path: Path) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    server, thread = _http_server(controller)
    try:
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        status, catalog = _http_json(connection, "GET", "/api/v2/strategies/performance-v2/maintenance/catalog")
        assert status == 200 and catalog["symbols"] == ["BTCUSDT"]

        with duckdb.connect(str(database), read_only=True) as db:
            before = db.execute("select count(*) from strategies").fetchone()[0]
        status, bad = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/preview", {
            "operation": "full", "symbols": ["BTCUSDT"], "strategy_ids": [1],
        })
        assert status == 400 and bad["error"]["code"] == "INVALID_REQUEST"

        status, preview = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/preview", {
            "operation": "full", "symbols": ["BTCUSDT"],
        })
        assert status == 200
        assert preview["pairs"][0]["symbol"] == "BTCUSDT"
        assert "token" in preview and "fingerprint" not in preview
        assert preview["global_counts"]["import_files"] >= 0
        with duckdb.connect(str(database), read_only=True) as db:
            assert db.execute("select count(*) from strategies").fetchone()[0] == before

        status, accepted = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/apply", {
            "token": preview["token"],
        })
        assert status == 202 and accepted["job_id"]
        job_id = accepted["job_id"]
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            status, job = _http_json(connection, "GET", f"/api/v2/strategies/performance-v2/maintenance/status?job_id={job_id}")
            assert status == 200
            if job["status"] in {"COMMITTED", "FAILED"}:
                break
            time.sleep(.02)
        assert job["status"] == "COMMITTED"
        assert job["strategy_total"] == preview["pairs"][0]["strategy_count"]
        assert job["pair_scoped_deleted"] == job["pair_scoped_total"]

        status, reused = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/apply", {
            "token": preview["token"],
        })
        assert status == 409 and reused["error"]["code"] == "PREVIEW_TOKEN_INVALID"
        connection.close()
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_performance_v2_maintenance_preview_tokens_expire_and_evict_oldest(tmp_path: Path) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    payload = {"operation": "full", "symbols": ["BTCUSDT"]}
    preview = controller.strategies_performance_v2_maintenance_preview(payload)
    token = str(preview["token"])
    created_at, saved_preview = controller._performance_v2_maintenance_previews[token]
    controller._performance_v2_maintenance_previews[token] = (
        created_at - panel_module._PERFORMANCE_V2_MAINTENANCE_PREVIEW_TTL_SECONDS - 1,
        saved_preview,
    )
    with pytest.raises(PerformanceV2ApiError) as expired:
        controller.strategies_performance_v2_maintenance_apply({"token": token})
    assert expired.value.code == "PREVIEW_TOKEN_INVALID"

    now = panel_module.perf_counter()
    limit = panel_module._PERFORMANCE_V2_MAINTENANCE_MAX_PREVIEWS
    controller._performance_v2_maintenance_previews = {
        f"old-{index}": (now, {}) for index in range(limit)
    }
    newest = controller.strategies_performance_v2_maintenance_preview(payload)
    assert len(controller._performance_v2_maintenance_previews) == limit
    assert "old-0" not in controller._performance_v2_maintenance_previews
    assert f"old-{limit - 1}" in controller._performance_v2_maintenance_previews
    with pytest.raises(PerformanceV2ApiError) as evicted:
        controller.strategies_performance_v2_maintenance_apply({"token": "old-0"})
    assert evicted.value.code == "PREVIEW_TOKEN_INVALID"
    assert database.is_file()


def test_performance_v2_maintenance_plateau_recovery_expires_after_fifteen_minutes(tmp_path: Path) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    preview = controller.strategies_performance_v2_maintenance_preview({
        "operation": "full", "symbols": ["BTCUSDT"],
    })
    saved_preview = controller._performance_v2_maintenance_previews[preview["token"]][1]
    with duckdb.connect(str(database)) as connection:
        connection.execute("delete from strategy_orders")
    controller._performance_v2_maintenance_recovery_preview = (
        panel_module.perf_counter() - panel_module._PERFORMANCE_V2_MAINTENANCE_RECOVERY_TTL_SECONDS - 1,
        saved_preview,
    )

    after_expiry = controller.strategies_performance_v2_maintenance_preview({
        "operation": "full", "symbols": ["BTCUSDT"],
    })

    assert after_expiry["table_counts"]["analysis_plateaus"] == 0
    assert controller._performance_v2_maintenance_recovery_preview is None


def test_performance_v2_maintenance_apply_admission_is_atomic(tmp_path: Path, monkeypatch) -> None:
    controller, _, _ = _controller_for_windows(tmp_path)
    payload = {"operation": "full", "symbols": ["BTCUSDT"]}
    tokens = [
        controller.strategies_performance_v2_maintenance_preview(payload)["token"]
        for _ in range(2)
    ]
    worker_entered = threading.Event()
    release_worker = threading.Event()
    worker_ids: list[str] = []

    def blocked_worker(job_id, _preview, _started_at):
        worker_ids.append(job_id)
        worker_entered.set()
        assert release_worker.wait(5)
        with controller._performance_v2_maintenance_state_lock:
            job = controller._performance_v2_maintenance_job
            if job and job.get("job_id") == job_id:
                job.update({"status": "COMMITTED", "phase": "committed", "elapsed_seconds": 0.0})
                controller._performance_v2_maintenance_active_job = None

    monkeypatch.setattr(controller, "_run_performance_v2_maintenance", blocked_worker)
    start = threading.Barrier(3)
    outcomes: list[tuple[int, str]] = []

    def submit(token: str) -> None:
        start.wait()
        try:
            result = controller.strategies_performance_v2_maintenance_apply({"token": token})
            outcomes.append((202, str(result["job_id"])))
        except PerformanceV2ApiError as error:
            outcomes.append((error.status, error.code))

    requests = [threading.Thread(target=submit, args=(token,)) for token in tokens]
    try:
        for request in requests:
            request.start()
        start.wait()
        for request in requests:
            request.join(timeout=3)
        assert all(not request.is_alive() for request in requests)
        assert worker_entered.wait(2)
        assert sorted(status for status, _ in outcomes) == [202, 409]
        assert sum(status == 409 and value == "PERFORMANCE_V2_MAINTENANCE_BUSY" for status, value in outcomes) == 1
        assert len(worker_ids) == 1
    finally:
        release_worker.set()
        for request in requests:
            request.join(timeout=2)


@pytest.mark.parametrize(
    ("operation", "message", "expected_code", "expected_status"),
    [
        ("catalog", "Could not set lock on file", "PERFORMANCE_V2_LOCKED", 409),
        ("catalog", "database page checksum mismatch", "PERFORMANCE_V2_DATABASE_ERROR", 500),
        ("preview", "Could not set lock on file", "PERFORMANCE_V2_LOCKED", 409),
        ("preview", "database page checksum mismatch", "PERFORMANCE_V2_DATABASE_ERROR", 500),
    ],
)
def test_performance_v2_maintenance_reports_database_errors_by_cause(
    tmp_path: Path, monkeypatch, operation: str, message: str,
    expected_code: str, expected_status: int,
) -> None:
    controller, _, _ = _controller_for_windows(tmp_path)

    def fail_connect(*_args, **_kwargs):
        raise duckdb.IOException(message)

    monkeypatch.setattr(panel_module.duckdb, "connect", fail_connect)
    with pytest.raises(PerformanceV2ApiError) as raised:
        if operation == "catalog":
            controller.strategies_performance_v2_maintenance_catalog()
        else:
            controller.strategies_performance_v2_maintenance_preview(
                {"operation": "full", "symbols": ["BTCUSDT"]}
            )
    assert raised.value.code == expected_code
    assert raised.value.status == expected_status
    assert message in str(raised.value)


def test_performance_v2_maintenance_schema_errors_are_typed_for_catalog_and_preview(tmp_path: Path) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    with duckdb.connect(str(database)) as db:
        db.execute("create table unclassified_maintenance_table (id integer)")
    server, thread = _http_server(controller)
    try:
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        status, catalog = _http_json(connection, "GET", "/api/v2/strategies/performance-v2/maintenance/catalog")
        assert status == 409
        assert catalog["error"]["code"] == "PERFORMANCE_V2_SCHEMA_INVALID"
        status, preview = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/preview", {
            "operation": "full", "symbols": ["BTCUSDT"],
        })
        assert status == 409
        assert preview["error"]["code"] == "PERFORMANCE_V2_SCHEMA_INVALID"
        assert "unclassified_maintenance_table" in preview["error"]["message"]
        connection.close()
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_performance_v2_maintenance_catalog_and_preview_return_promptly_when_panel_writer_lock_is_held(tmp_path: Path) -> None:
    controller, _, _ = _controller_for_windows(tmp_path)
    lock_entered = threading.Event()
    release_lock = threading.Event()

    def hold_writer_lock() -> None:
        with controller._performance_v2_writer_lock:
            lock_entered.set()
            assert release_lock.wait(5)

    holder = threading.Thread(target=hold_writer_lock, daemon=True)
    holder.start()
    assert lock_entered.wait(2)
    server, thread = _http_server(controller)
    try:
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        status, catalog = _http_json(connection, "GET", "/api/v2/strategies/performance-v2/maintenance/catalog")
        assert status == 409 and catalog["error"]["code"] == "PERFORMANCE_V2_LOCKED"
        status, preview = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/preview", {
            "operation": "full", "symbols": ["BTCUSDT"],
        })
        assert status == 409 and preview["error"]["code"] == "PERFORMANCE_V2_LOCKED"
        connection.close()
    finally:
        release_lock.set()
        holder.join(timeout=2)
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_performance_v2_maintenance_progress_callback_accepts_unpreviewed_symbol_and_table(
    tmp_path: Path, monkeypatch,
) -> None:
    controller, _, _ = _controller_for_windows(tmp_path)
    preview = controller.strategies_performance_v2_maintenance_preview({
        "operation": "full", "symbols": ["BTCUSDT"],
    })

    def report_committed_row(_connection, _preview, *, on_phase=None, on_commit=None):
        if on_phase:
            on_phase("unexpected_table")
        if on_commit:
            on_commit({
                "table": "unexpected_table", "rows": 1,
                "rows_by_symbol": {"UNPREVIEWED": 1}, "shared_plateau_rows": 0, "global": False,
            })
        return {
            "pair_scoped_deleted": 1, "global_deleted": {},
            "table_counts": {"unexpected_table": 1}, "shared_plateau_deleted": 0,
        }

    monkeypatch.setattr(panel_module, "apply_performance_v2_maintenance_preview", report_committed_row)
    accepted = controller.strategies_performance_v2_maintenance_apply({"token": preview["token"]})
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        job = controller.strategies_performance_v2_maintenance_status(accepted["job_id"])
        if job["status"] in {"COMMITTED", "FAILED"}:
            break
        time.sleep(.01)

    assert job["status"] == "COMMITTED"
    assert job["pair_counts"]["UNPREVIEWED"] == 1
    assert job["pair_table_counts"]["UNPREVIEWED"]["unexpected_table"] == 1


def test_performance_v2_maintenance_returns_conflicts_while_apply_runs(tmp_path: Path, monkeypatch) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    existing_writer_entered = threading.Event()
    release_existing_writer = threading.Event()
    entered = threading.Event()
    release = threading.Event()
    real_apply = panel_module.apply_performance_v2_maintenance_preview

    def existing_writer():
        with controller._performance_v2_writer_guard(database):
            existing_writer_entered.set()
            assert release_existing_writer.wait(5)

    def blocked_apply(connection, preview, *, on_phase=None, on_commit=None):
        entered.set()
        assert release.wait(5)
        return real_apply(connection, preview, on_phase=on_phase, on_commit=on_commit)

    monkeypatch.setattr(panel_module, "apply_performance_v2_maintenance_preview", blocked_apply)
    server, thread = _http_server(controller)
    try:
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        status, preview = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/preview", {
            "operation": "rejected", "symbols": ["BTCUSDT"],
        })
        assert status == 200
        existing_writer_thread = threading.Thread(target=existing_writer, daemon=True)
        existing_writer_thread.start()
        assert existing_writer_entered.wait(2)
        status, accepted = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/apply", {"token": preview["token"]})
        assert status == 202
        status, waiting_job = _http_json(connection, "GET", f"/api/v2/strategies/performance-v2/maintenance/status?job_id={accepted['job_id']}")
        assert status == 200 and waiting_job["status"] == "RUNNING"
        assert not entered.is_set(), "apply must wait for an existing Panel writer"
        release_existing_writer.set()
        assert entered.wait(2)
        def database_must_not_be_reopened():
            raise AssertionError("status polling must use memory only")
        monkeypatch.setattr(controller, "_performance_v2_maintenance_target", database_must_not_be_reopened)

        status, catalog = _http_json(connection, "GET", "/api/v2/strategies/performance-v2/maintenance/catalog")
        assert status == 409 and catalog["error"]["code"] == "PERFORMANCE_V2_MAINTENANCE_BUSY"
        status, preview_busy = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/preview", {
            "operation": "full", "symbols": ["BTCUSDT"],
        })
        assert status == 409 and preview_busy["error"]["code"] == "PERFORMANCE_V2_MAINTENANCE_BUSY"
        status, job = _http_json(connection, "GET", f"/api/v2/strategies/performance-v2/maintenance/status?job_id={accepted['job_id']}")
        assert status == 200 and job["status"] == "RUNNING"

        new_writer_entered = threading.Event()
        # Use a complete guard scope so this simulates a new writer path.
        def new_writer():
            with controller._performance_v2_writer_guard(database):
                new_writer_entered.set()
        new_writer_thread = threading.Thread(target=new_writer, daemon=True)
        new_writer_thread.start()
        assert not new_writer_entered.wait(.1), "new Panel writer must wait behind apply"

        release.set()
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            status, job = _http_json(connection, "GET", f"/api/v2/strategies/performance-v2/maintenance/status?job_id={accepted['job_id']}")
            if job["status"] in {"COMMITTED", "FAILED"}:
                break
            time.sleep(.02)
        assert job["status"] == "COMMITTED"
        assert new_writer_entered.wait(2)
        existing_writer_thread.join(timeout=2)
        new_writer_thread.join(timeout=2)
        connection.close()
    finally:
        release.set()
        release_existing_writer.set()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_performance_v2_maintenance_keeps_actual_counts_after_late_failure(tmp_path: Path, monkeypatch) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    now = datetime(2026, 1, 1, tzinfo=timezone.utc)
    with duckdb.connect(str(database)) as db:
        eth_strategy_id = db.execute(
            """insert into strategies (strategy_name, symbol, side, timeframe, close_ma_len,
                   order_count, analysis_run_id, candidate_identity, lifecycle_status,
                   created_at_utc, updated_at_utc)
               values ('beta', 'ETHUSDT', 'LONG', '1h', 3, 1, 'run', 'beta', 'ACTIVE', ?, ?)
               returning strategy_id""",
            [now, now],
        ).fetchone()[0]
        db.execute(
            """insert into strategy_orders (strategy_id, order_id, open_ma_len, open_multiplier,
                   shift_bp, lot_x, analysis_run_id, plateau_id, base_point_trades)
               values (?, 1, 7, 0.995, 125, 1, 'run', 'P1', 8)""",
            [eth_strategy_id],
        )
        import_run_id = db.execute(
            """insert into import_runs (source_inbox_sha256, expected_report_count,
                   imported_count, skipped_count, rejected_count, status, started_at_utc)
               values ('maintenance-test-run', 1, 1, 0, 0, 'IMPORTED', ?)
               returning import_run_id""",
            [now],
        ).fetchone()[0]
        db.execute(
            """insert into import_files (import_run_id, source_filename, source_html_sha256,
                   source_size_bytes, status) values (?, 'report.html', 'maintenance-test-file', 1, 'IMPORTED')""",
            [import_run_id],
        )
    real_apply = panel_module.apply_performance_v2_maintenance_preview

    def fail_after_actions(connection, preview, *, on_phase=None, on_commit=None):
        def phase(name):
            if on_phase is not None:
                on_phase(name)
            if name == "import_runs":
                raise RuntimeError("injected failure before import_runs")
        return real_apply(connection, preview, on_phase=phase, on_commit=on_commit)

    monkeypatch.setattr(panel_module, "apply_performance_v2_maintenance_preview", fail_after_actions)
    server, thread = _http_server(controller)
    try:
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        status, preview = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/preview", {
            "operation": "full", "symbols": ["BTCUSDT", "ETHUSDT"],
        })
        assert status == 200
        assert preview["shared_plateau_rows"] == 1
        status, accepted = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/apply", {"token": preview["token"]})
        assert status == 202
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            status, job = _http_json(connection, "GET", f"/api/v2/strategies/performance-v2/maintenance/status?job_id={accepted['job_id']}")
            if job["status"] in {"COMMITTED", "FAILED"}:
                break
            time.sleep(.02)
        assert status == 200 and job["status"] == "FAILED"
        assert job["current_table"] == "import_runs"
        assert "injected failure" in job["error"]
        assert job["error"] == "injected failure before import_runs"
        assert not any(word in job["error"].casefold() for word in ("rollback", "restored", "restore"))
        assert job["table_counts"]["strategy_actions"] == 2
        assert job["pair_table_counts"]["BTCUSDT"]["strategy_actions"] == 2
        assert job["pair_counts"]["ETHUSDT"] >= 1
        assert job["shared_plateau_deleted"] == 1
        assert job["global_journal_counts"] == {"import_files": 1, "import_runs": 0}
        assert job["pair_scoped_deleted"] == sum(job["pair_counts"].values())
        assert job["pair_scoped_deleted"] + job["shared_plateau_deleted"] == sum(job["table_counts"].values())
        assert job["elapsed_seconds"] > 0
        assert 0 < job["pair_scoped_deleted"] < job["pair_scoped_total"]
        with duckdb.connect(str(database), read_only=True) as db:
            assert db.execute("select count(*) from strategy_actions").fetchone()[0] == 0
            assert db.execute("select count(*) from strategy_equity").fetchone()[0] == 0
            assert db.execute("select count(*) from import_files").fetchone()[0] == 0
            assert db.execute("select count(*) from import_runs").fetchone()[0] == 1
        connection.close()
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_performance_v2_maintenance_retry_reuses_only_failed_job_plateau_keys(
    tmp_path: Path, monkeypatch,
) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    with duckdb.connect(str(database)) as db:
        db.execute(
            "insert into analysis_plateaus values ('run', 'UNUSED', 5, 7)"
        )
    real_apply = panel_module.apply_performance_v2_maintenance_preview
    failed_once = False

    class FailPlateauOnce:
        def __init__(self, connection):
            self.connection = connection

        def execute(self, sql, parameters=None):
            nonlocal failed_once
            if not failed_once and sql.lstrip().casefold().startswith("delete from analysis_plateaus"):
                failed_once = True
                raise duckdb.IOException("injected failure at analysis_plateaus")
            if parameters is None:
                return self.connection.execute(sql)
            return self.connection.execute(sql, parameters)

    def fail_first_plateau_delete(connection, preview, *, on_phase=None, on_commit=None):
        if not failed_once:
            return real_apply(FailPlateauOnce(connection), preview, on_phase=on_phase, on_commit=on_commit)
        return real_apply(connection, preview, on_phase=on_phase, on_commit=on_commit)

    monkeypatch.setattr(panel_module, "apply_performance_v2_maintenance_preview", fail_first_plateau_delete)
    server, thread = _http_server(controller)
    try:
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        payload = {"operation": "full", "symbols": ["BTCUSDT"]}
        status, first_preview = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/preview", payload)
        assert status == 200 and first_preview["table_counts"]["analysis_plateaus"] == 1
        status, accepted = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/apply", {"token": first_preview["token"]})
        assert status == 202
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            status, first_job = _http_json(connection, "GET", f"/api/v2/strategies/performance-v2/maintenance/status?job_id={accepted['job_id']}")
            if first_job["status"] in {"COMMITTED", "FAILED"}:
                break
            time.sleep(.02)
        assert status == 200 and first_job["status"] == "FAILED"
        assert "15 minutes" in first_job["recovery_warning"]
        forbidden_sidecars = {".wal", ".bak", ".backup", ".snapshot"}
        assert not [path.name for path in database.parent.iterdir() if path.suffix.casefold() in forbidden_sidecars]
        status, reused = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/apply", {"token": first_preview["token"]})
        assert status == 409 and reused["error"]["code"] == "PREVIEW_TOKEN_INVALID"
        with duckdb.connect(str(database), read_only=True) as db:
            assert db.execute("select count(*) from strategy_orders").fetchone() == (0,)

        status, intervening_preview = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/preview", {
            "operation": "rejected", "symbols": ["BTCUSDT"],
        })
        assert status == 200 and intervening_preview["pair_scoped_total"] == 0
        status, intervening_job = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/apply", {
            "token": intervening_preview["token"],
        })
        assert status == 202
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            status, next_job = _http_json(connection, "GET", f"/api/v2/strategies/performance-v2/maintenance/status?job_id={intervening_job['job_id']}")
            if next_job["status"] in {"COMMITTED", "FAILED"}:
                break
            time.sleep(.02)
        assert status == 200 and next_job["status"] == "COMMITTED"
        status, forgotten_job = _http_json(connection, "GET", f"/api/v2/strategies/performance-v2/maintenance/status?job_id={accepted['job_id']}")
        assert status == 404 and forgotten_job["error"]["code"] == "MAINTENANCE_JOB_NOT_FOUND"

        status, retry_preview = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/preview", payload)
        assert status == 200
        assert retry_preview["table_counts"]["analysis_plateaus"] == 1
        status, retry_job = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/apply", {"token": retry_preview["token"]})
        assert status == 202
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            status, final_job = _http_json(connection, "GET", f"/api/v2/strategies/performance-v2/maintenance/status?job_id={retry_job['job_id']}")
            if final_job["status"] in {"COMMITTED", "FAILED"}:
                break
            time.sleep(.02)
        assert status == 200 and final_job["status"] == "COMMITTED"
        assert controller._performance_v2_maintenance_recovery_preview is None
        with duckdb.connect(str(database)) as db:
            assert db.execute("select plateau_id from analysis_plateaus order by plateau_id").fetchall() == [("UNUSED",)]
            db.execute("insert into analysis_plateaus values ('run', 'P1', 5, 7)")
            db.execute(
                """insert into strategies (strategy_name, symbol, side, timeframe, close_ma_len,
                       order_count, analysis_run_id, candidate_identity, lifecycle_status,
                       created_at_utc, updated_at_utc)
                   values ('gamma', 'BTCUSDT', 'LONG', '1h', 3, 1, 'run', 'gamma', 'ACTIVE', current_timestamp, current_timestamp)"""
            )
        status, next_preview = _http_json(connection, "POST", "/api/v2/strategies/performance-v2/maintenance/preview", payload)
        assert status == 200
        assert next_preview["table_counts"]["analysis_plateaus"] == 0
        connection.close()
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_finalist_retest_preview_http_uses_server_owned_scope(tmp_path: Path, monkeypatch) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    seen: list[bool] = []

    def preview(include_reserve: bool = False) -> dict[str, object]:
        seen.append(include_reserve)
        return {"scope": "FINALIST_RESERVE", "test_start": "2025-01-01", "test_end": "2026-09-07", "cohort_count": 7, "excluded_count": 1}

    monkeypatch.setattr(controller, "strategies_performance_v2_finalist_retest_preview", preview)
    server, thread = _http_server(controller)
    try:
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        status, document = _http_json(connection, "GET", "/api/v2/strategies/performance-v2/finalist-retest/preview?include_reserve=true")
        connection.close()
        assert status == 200 and document["cohort_count"] == 7
        assert seen == [True]

        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        status, document = _http_json(connection, "GET", "/api/v2/strategies/performance-v2/finalist-retest/preview?include_reserve=1")
        connection.close()
        assert status == 400 and document["error"]["code"] == "INVALID_REQUEST"
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_finalist_retest_preview_http_takes_an_optional_exchange(tmp_path: Path, monkeypatch) -> None:
    # Panel stage 5e-3b: the route passes a known exchange through, refuses another, and answers errors with a code
    from mrs3.evaluation_protocol import EvaluationProtocolError

    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    seen: list[dict[str, object]] = []

    def preview(include_reserve: bool = False, **kwargs) -> dict[str, object]:
        seen.append(kwargs)
        if kwargs.get("exchange") == "bingx" and len(seen) > 2:
            raise EvaluationProtocolError("EXCHANGE_TEMPLATES_MISSING", "C:\\secret\\templates is missing")
        return {"scope": "FINALIST", "test_start": "2025-01-01", "test_end": "2026-09-07", "cohort_count": 1,
                "excluded_count": 0}

    monkeypatch.setattr(controller, "strategies_performance_v2_finalist_retest_preview", preview)
    server, thread = _http_server(controller)
    base = "/api/v2/strategies/performance-v2/finalist-retest/preview?include_reserve=false"
    try:
        answers = []
        for suffix in ("", "&exchange=bingx", "&exchange=kraken", "&exchange=bingx"):
            connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
            answers.append(_http_json(connection, "GET", base + suffix))
            connection.close()
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)
    assert [status for status, _document in answers] == [200, 200, 400, 409]
    assert seen == [{}, {"exchange": "bingx"}, {"exchange": "bingx"}]           # kraken never reaches the controller
    assert answers[3][1]["error"]["code"] == "EXCHANGE_TEMPLATES_MISSING" and "secret" not in json.dumps(answers[3][1])


def test_current_finalist_control_export_http_accepts_reserve_scope_without_job(tmp_path: Path, monkeypatch) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    seen: list[dict[str, object]] = []
    request = parse_selection_request({"symbol": "BTCUSDT", "side": "LONG", "stages": []})
    ordinary_path = write_selection_workbook(
        pd.DataFrame([{
            "strategy_id": 1, "result_id": 11, "strategy_name": "alpha", "symbol": "BTCUSDT", "side": "LONG",
            "timeframe": "1h", "order_count": 1, "close_ma_len": 7, "pnl_30d_pct": Decimal("12"),
            "dd5_proxy": Decimal("4"), "max_drawdown_pct": Decimal("8"), "order_1_open_ma_len": 7,
            "order_1_lot_x": Decimal("0.25"), "order_1_plateau_point_count": 4, "auto_status": "FINALIST",
            "finalist": True, "final_rank": 1, "final_score": Decimal("5.5"), "elimination_reason": None,
        }]),
        tmp_path / "ordinary-candidates.xlsx", request, {"workbook_schema_version": "2"},
        {1: {"user_status": "FINALIST", "user_rank": 1, "user_analog_of_strategy_id": None, "comment": None}},
    )
    control_data = combined_control_workbook_bytes(
        [{"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 1, "result_id": 11,
          "user_status": "FINALIST", "user_rank": 1, "auto_status": "FINALIST", "auto_rank": 1,
          "score": Decimal("5.5")}],
        metadata={}, candidate_workbook=ordinary_path.read_bytes(),
    )

    def export(payload: dict[str, object]) -> tuple[str, bytes]:
        seen.append(payload)
        return "performance-v2-current-finalists-with-reserve.xlsx", control_data

    monkeypatch.setattr(controller, "strategies_performance_v2_finalist_retest_export", export)
    server, thread = _http_server(controller)
    try:
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        connection.request("GET", "/api/v2/strategies/performance-v2/finalist-retest/export?include_reserve=true")
        response = connection.getresponse()
        assert response.status == 200
        assert response.getheader("Content-Disposition") == 'attachment; filename="performance-v2-current-finalists-with-reserve.xlsx"'
        document = response.read()
        issued_sheet = load_workbook(BytesIO(document))["Candidates"]
        ordinary_sheet = load_workbook(ordinary_path)["All candidates"]
        issued_headers = [cell.value for cell in issued_sheet[1]]
        ordinary_headers = [cell.value for cell in ordinary_sheet[1]]
        assert issued_headers == ordinary_headers
        assert all(isinstance(header, str) and header.strip() for header in ordinary_headers)
        assert len(ordinary_headers) == len(set(ordinary_headers))
        headers = {header: index + 1 for index, header in enumerate(issued_headers)}
        assert issued_sheet.cell(2, headers["PnL/30"]).value == 12
        assert issued_sheet.cell(2, headers["DD"]).value == 8
        assert issued_sheet.cell(2, headers["Lots"]).value == "25"
        assert issued_sheet.cell(2, headers["Points"]).value == "4"
        assert issued_sheet.cell(2, headers["MA"]).value == "7"
        assert issued_sheet.cell(2, headers["Close"]).value == 7
        assert seen == [{"include_reserve": True}]
        connection.close()

        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        connection.request("GET", "/api/v2/strategies/performance-v2/finalist-retest/export?strategy_ids=1")
        response = connection.getresponse()
        assert response.status == 409
        connection.close()
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_current_finalist_control_export_http_empty_scope_is_unavailable(tmp_path: Path) -> None:
    (tmp_path / "config.local.json").write_text("{}", encoding="utf-8")
    (tmp_path / "config.performance.json").write_text(
        json.dumps({"unified_performance_v2": {"database_root": "performance-v2", "workers": 1}}),
        encoding="utf-8",
    )
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    database = tmp_path / "performance-v2" / "strategy_performance.duckdb"
    database.parent.mkdir()
    with duckdb.connect(str(database)) as connection:
        initialize_performance_v2(connection)
    server, thread = _http_server(controller)
    try:
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        connection.request("GET", "/api/v2/strategies/performance-v2/finalist-retest/export?include_reserve=false")
        response = connection.getresponse()
        status = response.status
        content_type = response.getheader("Content-Type", "")
        body = response.read()
        connection.close()
        assert status == 409
        assert "spreadsheet" not in content_type.casefold()
        assert not body.startswith(b"PK")
        document = json.loads(body.decode("utf-8"))
        assert document["error"]["code"] == "CONTROL_EXPORT_UNAVAILABLE"
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_current_finalist_control_export_http_hides_untrusted_error_details(tmp_path: Path, monkeypatch) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    leaked_path = tmp_path / "private" / "secret.xlsx"

    def export(_payload: dict[str, object]) -> tuple[str, bytes]:
        raise FinalistRetestError("UNSAFE_INTERNAL", f"failed to open {leaked_path}")

    monkeypatch.setattr(controller, "strategies_performance_v2_finalist_retest_export", export)
    server, thread = _http_server(controller)
    try:
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        connection.request("GET", "/api/v2/strategies/performance-v2/finalist-retest/export?include_reserve=false")
        response = connection.getresponse()
        body = response.read()
        connection.close()
        assert response.status == 409
        document = json.loads(body.decode("utf-8"))
        assert document["error"] == {
            "code": "CONTROL_EXPORT_UNAVAILABLE",
            "message": "bulk control workbook is unavailable",
        }
        assert str(leaked_path) not in body.decode("utf-8")
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_job_id_finalist_export_uses_stage_free_rich_control_schema(tmp_path: Path, monkeypatch) -> None:
    controller, database, result_id = _controller_for_windows(tmp_path)
    with duckdb.connect(str(database)) as connection:
        connection.executemany(
            "insert into strategy_actions (result_id, action_index, timestamp_utc, symbol, order_id, action, size, post_size, post_side, pnl, fee, balance, raw_action_json) values (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            [
                (result_id, 2, datetime(2026, 1, 2, 12, tzinfo=UTC), "BTCUSDT", 1, "opened", 1, 1, "long", 0, 1, 110, None),
                (result_id, 3, datetime(2026, 1, 3, 12, tzinfo=UTC), "BTCUSDT", 1, "closed", 1, 0, "", 5, 1, 115, None),
                (result_id, 4, datetime(2026, 1, 3, 18, tzinfo=UTC), "BTCUSDT", 1, "opened", 1, 1, "long", 0, 1, 115, None),
                (result_id, 5, datetime(2026, 1, 4, 12, tzinfo=UTC), "BTCUSDT", 1, "closed", 1, 0, "", 5, 1, 120, None),
            ],
        )
        connection.execute(
            "insert into strategy_equity values (?, ?, ?, ?, ?)",
            [result_id, 3, datetime(2026, 1, 3, 12, tzinfo=UTC), 115, 115],
        )
        connection.execute(
            "insert into strategy_equity values (?, ?, ?, ?, ?)",
            [result_id, 4, datetime(2026, 1, 4, 12, tzinfo=UTC), 120, 120],
        )
        connection.execute("update strategy_equity set wallet = 120, equity = 120 where result_id = ? and sample_index = 2", [result_id])
    controller.strategies_performance_v2_recalculate_all()
    stageful_payload = {
        "symbol": "BTCUSDT", "side": "LONG", "stages": [
            {"id": "filter_min_shift", "enabled": True, "scope": "pair_side", "min_shift_pct": "0.3"},
        ],
    }
    stageful_request = parse_selection_request(stageful_payload)
    runtime = {
        "bulk_retest": True, "scope": "FINALIST", "cohort_sha256": "c" * 64,
        "manifest_sha256": "m" * 64, "config_sha256": "g" * 64,
        "successful_replacements": [{"strategy_id": 1, "new_result_id": result_id}],
        "cohort_members": [{
            "strategy_id": 1, "strategy_name": "alpha", "symbol": "BTCUSDT", "side": "LONG",
            "result_id": result_id, "effective_status": "FINALIST", "effective_rank": 1,
            "effective_start": "2026-01-01", "effective_end": "2026-01-05",
        }],
    }
    monkeypatch.setattr(controller, "_bulk_retest_status_document", lambda _job_id: {"state": "COMMITTED"})
    monkeypatch.setattr(controller._panel_jobs, "runtime", lambda _job_id: runtime)
    monkeypatch.setattr(panel_module, "parse_selection_request", lambda _payload: stageful_request)
    filename, issued = controller.strategies_performance_v2_finalist_retest_export({"job_id": "bulk-export"})
    assert filename == "performance-v2-finalist-retest-bulk-export.xlsx"
    workbook = load_workbook(BytesIO(issued))
    sheet = workbook["Candidates"]
    headers = [cell.value for cell in sheet[1]]
    assert "eliminated_by_filter_min_shift" not in headers         # stage-free: no per-filter columns
    columns = {header: index + 1 for index, header in enumerate(headers)}
    assert sheet.cell(2, columns["PnL/30"]).value is not None
    assert sheet.cell(2, columns["DD"]).value == 0
    assert sheet.cell(2, columns["Lots"]).value == "100"
    assert sheet.cell(2, columns["Points"]).value == "12"
    assert sheet.cell(2, columns["MA"]).value == "7"
    assert sheet.cell(2, columns["Close"]).value == 3
    assert not any(cell.data_type == "f" for worksheet in workbook.worksheets for row in worksheet.iter_rows() for cell in row)
    imported = controller.strategies_performance_v2_finalist_retest_control_import(issued)
    assert imported["group_count"] == imported["row_count"] == 1


def test_finalist_retest_start_rejects_client_member_ids_before_io(tmp_path: Path) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")

    with pytest.raises(FinalistRetestError, match="unsupported fields"):
        controller.strategies_performance_v2_finalist_retest_start({"strategy_ids": [1]})


def test_finalist_retest_start_http_preserves_typed_validation_error(tmp_path: Path) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    server, thread = _http_server(controller)
    try:
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        status, document = _http_json(
            connection, "POST", "/api/v2/strategies/performance-v2/finalist-retest/start",
            {"strategy_ids": [1]},
        )
        connection.close()
        assert status == 400
        assert document["error"]["code"] == "INVALID_REQUEST"
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_finalist_retest_default_end_is_two_complete_utc_days_back(monkeypatch) -> None:
    import mrs3.panel as panel_module

    class FixedDateTime:
        @classmethod
        def now(cls, _timezone):
            return datetime(2026, 9, 9, 12, tzinfo=timezone.utc)

    monkeypatch.setattr(panel_module, "datetime", FixedDateTime)
    monkeypatch.setattr(panel_module, "freeze_finalist_cohort", lambda *_args, **_kwargs: SimpleNamespace(
        members=({"listing_date_utc": datetime(2020, 1, 1, tzinfo=timezone.utc)},)
    ))

    assert PanelController._bulk_retest_range({}, {}, object(), False) == ("2020-01-01", "2026-09-07")


def test_finalist_retest_replays_only_exact_cohort_and_config(tmp_path: Path, monkeypatch) -> None:
    import mrs3.panel as panel_module

    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    target = use_protocol_v2_target(monkeypatch)
    database = tmp_path / "performance.duckdb"
    database.touch()
    listing = tmp_path / "dates.xlsx"
    listing.touch()
    template = tmp_path / "base.json"
    template.write_text("{}", encoding="utf-8")
    manifest = tmp_path / "manifest.json"
    manifest.write_text("{}", encoding="utf-8")
    cohort = SimpleNamespace(
        scope="FINALIST", cohort_sha256="new-cohort",
        members=({"strategy_id": 1, "effective_start": datetime(2026, 1, 1, tzinfo=timezone.utc)},), exclusions=(),
    )

    class Connection:
        def __enter__(self): return self
        def __exit__(self, *_args): return None

    class Jobs:
        def __init__(self, old_cohort: str): self.old_cohort = old_cohort
        def list(self): return [{"kind": "strategies.performance.v2.finalist-retest", "job_id": "old", "state": "COMMITTED"}]
        def runtime(self, _job_id):
            return {"scope": "FINALIST", "test_start": "2025-01-01", "test_end": "2026-09-07",
                    "cohort_sha256": self.old_cohort, "config_sha256": panel_module.finalist_retest_config_digest({"LONG": str(template)}, evaluation=target),
                    "outcomes_finalized": True, "successful_replacements": [{"strategy_id": 1}],
                    **controller._protocol_v2_runtime(target)}

    monkeypatch.setattr(controller, "_retest_listing_context", lambda: (listing, Path("dates.xlsx"), None))
    monkeypatch.setattr(controller, "_performance_v2_config", lambda: SimpleNamespace(strategy_root=tmp_path / "Output" / "strategies"))
    monkeypatch.setattr(controller, "_workflow_defaults", lambda: {"strategy_templates": {"LONG": str(template)}})
    monkeypatch.setattr(controller, "_bulk_retest_range", lambda *_args: ("2025-01-01", "2026-09-07"))
    monkeypatch.setattr(controller, "_bulk_retest_status_document", lambda job_id: {"job_id": job_id, "replayed": True})
    captured: dict[str, object] = {}
    monkeypatch.setattr(
        controller, "_start_tracked_panel_job",
        lambda *_args, **kwargs: captured.update(runtime=kwargs["runtime"], request=_args[1], submit=_args[3]) or {"job_id": "new"},
    )
    monkeypatch.setattr(panel_module, "performance_v2_database_path", lambda _config: database)
    monkeypatch.setattr(panel_module, "load_listing_dates", lambda _path: {})
    monkeypatch.setattr(panel_module.duckdb, "connect", lambda *_args, **_kwargs: Connection())
    monkeypatch.setattr(panel_module, "freeze_finalist_cohort", lambda *_args, **_kwargs: cohort)
    monkeypatch.setattr(panel_module, "build_finalist_retest_manifest", lambda *_args, **_kwargs: SimpleNamespace(
        cohort=cohort, config_sha256=panel_module.finalist_retest_config_digest({"LONG": str(template)}, evaluation=target),
        manifest_path=manifest, run_id="new-run",
    ))

    controller._panel_jobs = Jobs("old-cohort")  # type: ignore[assignment]
    started: dict[str, object] = {}
    monkeypatch.setattr(controller, "_single_mode_strategy_test", lambda: SimpleNamespace(start=lambda *_args, **kwargs: started.update(kwargs)))
    assert controller.strategies_performance_v2_finalist_retest_start({"clear_reports": True})["job_id"] == "new"
    assert captured["runtime"]["cohort_members"] == [{"strategy_id": 1, "effective_start": "2026-01-01T00:00:00Z"}]
    snapshot = captured["runtime"]["equity_filter_snapshot"]
    assert snapshot["schema_version"] == 1
    assert snapshot["request"]["stages"] == [{"id": "filter_equity_regime", "enabled": True, "scope": "pair_side"}]
    assert 1 <= snapshot["cache_workers"] <= 16
    assert snapshot["selection_config"]
    assert snapshot["regime_policy"]
    assert snapshot["selection_config_sha256"]
    assert snapshot["regime_policy_sha256"]
    captured["submit"]("new")
    assert captured["request"]["clear_reports"] is True
    assert started["clear_reports"] is True
    # Protocol v2 §4.3: the retest runs on the base of its template, no start balance.
    assert "initial_balance" not in started
    controller._panel_jobs = Jobs("new-cohort")  # type: ignore[assignment]
    assert controller.strategies_performance_v2_finalist_retest_start({}) == {"job_id": "old", "replayed": True}


def test_finalist_retest_status_exposes_started_import_job(tmp_path: Path, monkeypatch) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    controller._panel_jobs.submit(
        "strategies.performance.v2.finalist-retest", {}, "test", job_id="bulk-job",
    )
    controller._panel_jobs.transition("bulk-job", "RUNNING")
    controller._panel_jobs.transition("bulk-job", "COMMITTED")
    controller._panel_jobs.sync(
        "bulk-job", {"state": "COMMITTED", "phase": "COMMITTED", "inbox_ready": True},
        runtime={
            "bulk_retest": True, "scope": "FINALIST", "cohort_sha256": "cohort",
            "cohort_members": [], "successful_replacements": [], "failures": [],
            "bulk_import_job_id": "import-job",
        },
    )
    monkeypatch.setattr(controller, "_single_mode_strategy_test", lambda: SimpleNamespace(status=lambda _job_id: (_ for _ in ()).throw(KeyError())))

    status = controller.strategies_performance_v2_finalist_retest_status("bulk-job")
    assert status["import_job_id"] == "import-job"
    assert status["equity_filter"]["state"] == "IDLE"
    assert status["equity_filter"]["eligible"] is False
    assert status["equity_filter"]["ineligible_reason"] == "IMPORT_NOT_FINALIZED"
    assert {"application_key", "attempt", "progress", "result", "error"}.issubset(status["equity_filter"])


def test_finalist_retest_equity_filter_status_allows_retry_after_failed_child(tmp_path: Path) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    registry = controller._panel_jobs
    registry.submit(
        "strategies.performance.v2.finalist-retest", {}, "equity-status-parent", job_id="equity-status-parent",
    )
    registry.transition("equity-status-parent", "RUNNING")
    registry.transition("equity-status-parent", "COMMITTED")
    registry.submit(
        "strategies.performance.v2.finalist-retest.equity-filter", {}, "equity-status-child",
        ("performance-v2-equity-filter",), job_id="equity-status-child",
    )
    registry.transition("equity-status-child", "FAILED")
    registry.sync("equity-status-parent", {"state": "COMMITTED"}, runtime={
        "bulk_retest": True,
        "outcomes_finalized": True,
        "successful_replacements": [{"strategy_id": 1, "new_result_id": 2}],
        "equity_filter_job_id": "equity-status-child",
        "equity_filter_published": False,
    })

    status = controller._equity_filter_child_status("equity-status-parent")

    assert status["state"] == "FAILED"
    assert status["eligible"] is True


def test_finalist_retest_equity_filter_retries_failed_child_with_new_registry_key(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    registry = controller._panel_jobs
    parent_id = "equity-retry-parent"
    registry.submit(
        "strategies.performance.v2.finalist-retest", {}, "equity-retry-parent", job_id=parent_id,
    )
    registry.transition(parent_id, "RUNNING")
    registry.transition(parent_id, "COMMITTED")
    registry.sync(parent_id, {"state": "COMMITTED"}, runtime={
        "bulk_retest": True,
        "scope": "FINALIST",
        "outcomes_finalized": True,
        "successful_replacements": [{"strategy_id": 1, "new_result_id": 2}],
        "cohort_members": [{"strategy_id": 1, "result_id": 1, "symbol": "BTCUSDT", "side": "LONG"}],
        "equity_filter_snapshot": {"schema_version": 1},
    })

    class Connection:
        def __enter__(self): return self
        def __exit__(self, *_args): return None

    monkeypatch.setattr(controller, "_performance_v2_config", lambda: SimpleNamespace())
    monkeypatch.setattr(panel_module, "performance_v2_database_path", lambda _config: tmp_path / "performance.duckdb")
    monkeypatch.setattr(panel_module.duckdb, "connect", lambda *_args, **_kwargs: Connection())
    monkeypatch.setattr(panel_module, "require_performance_v2", lambda _connection: None)
    monkeypatch.setattr(panel_module, "_current_equity_revisions", lambda _connection, _ids: {1: (2, "source")})
    monkeypatch.setattr(
        controller, "_run_bulk_retest_equity_filter_job",
        lambda bulk_job_id, child_id: None,
    )

    first = controller.strategies_performance_v2_finalist_retest_equity_filter({"job_id": parent_id})
    first_id = first["job_id"]
    registry.transition(first_id, "FAILED")

    second = controller.strategies_performance_v2_finalist_retest_equity_filter({"job_id": parent_id})

    assert second["job_id"] != first_id
    assert registry.get(second["job_id"])["idempotency_key"].endswith(":attempt:2")


def test_finalist_retest_equity_filter_bounds_retry_child_allocation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    registry = controller._panel_jobs
    parent_id = "equity-retry-exhausted-parent"
    registry.submit("strategies.performance.v2.finalist-retest", {}, parent_id, job_id=parent_id)
    registry.transition(parent_id, "RUNNING")
    registry.transition(parent_id, "COMMITTED")
    registry.sync(parent_id, {"state": "COMMITTED"}, runtime={
        "bulk_retest": True,
        "scope": "FINALIST",
        "outcomes_finalized": True,
        "successful_replacements": [{"strategy_id": 1, "new_result_id": 2}],
        "cohort_members": [{"strategy_id": 1, "result_id": 1, "symbol": "BTCUSDT", "side": "LONG"}],
        "equity_filter_snapshot": {"schema_version": 1},
    })

    class Connection:
        def __enter__(self): return self
        def __exit__(self, *_args): return None

    monkeypatch.setattr(controller, "_performance_v2_config", lambda: SimpleNamespace())
    monkeypatch.setattr(panel_module, "performance_v2_database_path", lambda _config: tmp_path / "performance.duckdb")
    monkeypatch.setattr(panel_module.duckdb, "connect", lambda *_args, **_kwargs: Connection())
    monkeypatch.setattr(panel_module, "require_performance_v2", lambda _connection: None)
    monkeypatch.setattr(panel_module, "_current_equity_revisions", lambda _connection, _ids: {1: (2, "source")})
    calls = 0

    def return_terminal_child(*_args, **_kwargs):
        nonlocal calls
        calls += 1
        return {"job_id": f"terminal-{calls}", "state": "FAILED"}

    monkeypatch.setattr(registry, "submit", return_terminal_child)

    with pytest.raises(PerformanceV2ApiError) as error:
        controller.strategies_performance_v2_finalist_retest_equity_filter({"job_id": parent_id})

    assert error.value.code == "RETEST_EQUITY_ACTION_RETRY_EXHAUSTED"
    assert calls == 8


@pytest.mark.parametrize(
    ("parent_state", "scope", "outcomes_finalized", "successes", "expected"),
    [
        ("RUNNING", "FINALIST", True, [{"strategy_id": 1, "new_result_id": 2}], "RETEST_COHORT_JOB_NOT_COMMITTED"),
        ("COMMITTED", "FINALIST", False, [{"strategy_id": 1, "new_result_id": 2}], "RETEST_IMPORT_NOT_FINALIZED"),
        ("COMMITTED", "FINALIST", True, [], "RETEST_COHORT_NO_SUCCESSFUL_MEMBERS"),
        ("COMMITTED", "OTHER", True, [{"strategy_id": 1, "new_result_id": 2}], "RETEST_COHORT_INVALID"),
    ],
)
def test_finalist_retest_equity_filter_rejects_invalid_frozen_cohort_metadata(
    tmp_path: Path, parent_state: str, scope: str, outcomes_finalized: bool,
    successes: list[dict[str, int]], expected: str,
) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    parent_id = f"equity-invalid-cohort-{expected.lower()}"
    registry = controller._panel_jobs
    registry.submit("strategies.performance.v2.finalist-retest", {}, parent_id, job_id=parent_id)
    registry.transition(parent_id, "RUNNING")
    if parent_state == "COMMITTED":
        registry.transition(parent_id, "COMMITTED")
    registry.sync(parent_id, {"state": parent_state}, runtime={
        "bulk_retest": True,
        "scope": scope,
        "outcomes_finalized": outcomes_finalized,
        "successful_replacements": successes,
    })

    with pytest.raises(PerformanceV2ApiError) as error:
        controller.strategies_performance_v2_finalist_retest_equity_filter({"job_id": parent_id})

    assert error.value.code == expected


def test_finalist_retest_equity_filter_rejects_changed_request_while_child_is_active(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    registry = controller._panel_jobs
    parent_id = "equity-active-conflict-parent"
    child_id = "equity-active-conflict-child"
    registry.submit("strategies.performance.v2.finalist-retest", {}, parent_id, job_id=parent_id)
    registry.transition(parent_id, "RUNNING")
    registry.transition(parent_id, "COMMITTED")
    registry.submit(
        "strategies.performance.v2.finalist-retest.equity-filter", {}, child_id,
        ("performance-v2-equity-filter",), job_id=child_id,
    )
    registry.transition(child_id, "RUNNING")
    registry.sync(parent_id, {"state": "COMMITTED"}, runtime={
        "bulk_retest": True,
        "scope": "FINALIST",
        "outcomes_finalized": True,
        "successful_replacements": [{"strategy_id": 1, "new_result_id": 2}],
        "cohort_members": [{"strategy_id": 1, "result_id": 1, "symbol": "BTCUSDT", "side": "LONG"}],
        "equity_filter_snapshot": {"schema_version": 1},
        "equity_filter_job_id": child_id,
        "equity_filter_application_key": "previous-application",
        "equity_filter_lock_expires_at": (datetime.now(UTC) + timedelta(minutes=5)).isoformat(),
    })

    class Connection:
        def __enter__(self): return self
        def __exit__(self, *_args): return None

    monkeypatch.setattr(controller, "_performance_v2_config", lambda: SimpleNamespace())
    monkeypatch.setattr(panel_module, "performance_v2_database_path", lambda _config: tmp_path / "performance.duckdb")
    monkeypatch.setattr(panel_module.duckdb, "connect", lambda *_args, **_kwargs: Connection())
    monkeypatch.setattr(panel_module, "require_performance_v2", lambda _connection: None)
    monkeypatch.setattr(panel_module, "_current_equity_revisions", lambda _connection, _ids: {1: (2, "source")})

    with pytest.raises(PerformanceV2ApiError) as error:
        controller.strategies_performance_v2_finalist_retest_equity_filter({"job_id": parent_id})

    assert error.value.code == "RETEST_EQUITY_ACTION_ACTIVE"
    assert "equity_filter_attempt" not in registry.runtime(parent_id)


def test_finalist_retest_status_exposes_pending_import_handoff(tmp_path: Path, monkeypatch) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    controller._panel_jobs.submit(
        "strategies.performance.v2.finalist-retest", {}, "pending-test", job_id="pending-bulk-job",
    )
    controller._panel_jobs.transition("pending-bulk-job", "RUNNING")
    controller._panel_jobs.transition("pending-bulk-job", "COMMITTED")
    controller._panel_jobs.sync(
        "pending-bulk-job", {"state": "COMMITTED", "phase": "COMMITTED", "inbox_ready": True},
        runtime={
            "bulk_retest": True, "scope": "FINALIST", "cohort_sha256": "cohort",
            "cohort_members": [], "successful_replacements": [], "failures": [],
            "bulk_import_job_id": "pending:child-link-not-saved",
        },
    )
    monkeypatch.setattr(
        controller, "_single_mode_strategy_test",
        lambda: SimpleNamespace(status=lambda _job_id: (_ for _ in ()).throw(KeyError())),
    )

    status = controller.strategies_performance_v2_finalist_retest_status("pending-bulk-job")

    assert status["import_job_id"] is None
    assert status["import_pending"] is True


def test_finalist_retest_import_captures_committed_inbox_before_import(tmp_path: Path, monkeypatch) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    job_id = "bulk-inbox-before-import"
    registry = controller._panel_jobs
    registry.submit(
        "strategies.performance.v2.finalist-retest", {}, "bulk-inbox-before-import",
        ("strategies.tester", "performance-v2-finalist-retest"), job_id=job_id,
    )
    registry.transition(job_id, "RUNNING")
    runtime = {
        "bulk_retest": True,
        "scope": "FINALIST",
        "test_start": "2026-01-01",
        "test_end": "2026-10-01",
        "cohort_members": [{"strategy_id": 7, "strategy_name": "alpha", "result_id": 17}],
    }
    registry.sync(job_id, {"state": "COMMITTED", "phase": "COMMITTED"}, runtime=runtime)
    inbox_root = tmp_path / "inbox"
    inbox = inbox_root / job_id
    captured: list[tuple[str, dict[str, object]]] = []

    class TesterService:
        def capture_inbox(self, captured_job_id: str, *, force_single_mode: bool = False) -> Path:
            captured.append((captured_job_id, {"force_single_mode": force_single_mode}))
            inbox.mkdir(parents=True)
            (inbox / "inbox_manifest.json").write_text(json.dumps({
                "schema_version": 1,
                "run_mode": "SINGLE_MODE",
                "source_mode": "metadata_only",
                "inbox_ready": True,
                "batch_id": job_id,
                "expected_strategy_names": ["alpha"],
                "entries": [{
                    "strategy_name": "alpha",
                    "strategy_path": "alpha.json",
                    "report_path": "alpha.html",
                }],
            }), encoding="utf-8")
            return inbox

        def mark_inbox_ready(self, captured_job_id: str, path: Path) -> None:
            assert captured_job_id == job_id
            assert path == inbox

    service = TesterService()
    monkeypatch.setattr(panel_module.RunnerConfig, "from_json", staticmethod(lambda _path: SimpleNamespace(inbox_root=inbox_root)))
    monkeypatch.setattr(controller, "_single_mode_strategy_test", lambda: service)
    monkeypatch.setattr(controller, "_validate_metadata_inbox", lambda _path: None)
    import_calls: list[tuple[dict[str, object], bool]] = []

    child_ids: list[str] = []

    def start_import(
        payload: dict[str, object], *, _internal: bool = False, job_id: str | None = None,
    ) -> dict[str, object]:
        import_calls.append((payload, _internal))
        assert job_id is not None
        child_ids.append(job_id)
        assert registry.get("bulk-inbox-before-import")["inbox_ready"] is True
        child = registry.submit(
            "strategies.performance.v2.import", {}, f"bulk-inbox-import-child:{job_id}",
            ("performance-v2-db",), job_id=job_id,
        )
        registry.transition(child["job_id"], "RUNNING")
        return {"job_id": child["job_id"]}

    monkeypatch.setattr(controller, "strategies_performance_v2_import", start_import)

    result = controller.strategies_performance_v2_finalist_retest_import({"tester_job_id": job_id})

    assert result["job_id"] == child_ids[0]
    assert captured == [(job_id, {"force_single_mode": True})]
    assert len(import_calls) == 1
    import_payload, internal = import_calls[0]
    assert internal is True
    assert import_payload["replacement_strategy_ids"] == {"alpha": 7}
    assert import_payload["test_start"] == "2026-01-01"
    assert import_payload["test_end"] == "2026-10-01"
    assert registry.runtime(job_id)["inbox_path"] == str(inbox.resolve())
    assert registry.get(job_id)["inbox_ready"] is True
    repeated = controller.strategies_performance_v2_finalist_retest_import({"tester_job_id": job_id})
    assert repeated["job_id"] == result["job_id"]
    assert len(import_calls) == 1
    assert sum(job["kind"] == "strategies.performance.v2.import" for job in registry.list()) == 1


def test_finalist_retest_import_serializes_concurrent_handoffs(tmp_path: Path, monkeypatch) -> None:
    controller, registry = _bulk_import_controller(tmp_path, "concurrent-handoff")
    import_entered = threading.Event()
    release_import = threading.Event()
    second_waiting_for_lock = threading.Event()
    second_thread_id: list[int] = []
    starts: list[str] = []
    results: list[dict[str, object]] = []
    errors: list[BaseException] = []

    class ObservedLock:
        def __init__(self) -> None:
            self.lock = threading.RLock()

        def acquire(self):
            if second_thread_id and threading.get_ident() == second_thread_id[0]:
                second_waiting_for_lock.set()
            return self.lock.acquire()

        def release(self) -> None:
            self.lock.release()

        def __enter__(self):
            self.acquire()
            return self

        def __exit__(self, *_args) -> None:
            self.release()

    controller._bulk_retest_import_lock = ObservedLock()  # type: ignore[attr-defined]

    def start_import(
        _payload: dict[str, object], *, _internal: bool = False, job_id: str | None = None,
    ) -> dict[str, object]:
        assert _internal and job_id
        starts.append(job_id)
        import_entered.set()
        assert release_import.wait(5)
        child = registry.submit(
            "strategies.performance.v2.import", {}, f"concurrent-child:{job_id}",
            ("performance-v2-db",), job_id=job_id,
        )
        registry.transition(job_id, "RUNNING")
        return {"job_id": child["job_id"]}

    monkeypatch.setattr(controller, "strategies_performance_v2_import", start_import)

    def call_import(*, second: bool = False) -> None:
        if second:
            second_thread_id.append(threading.get_ident())
        try:
            results.append(controller.strategies_performance_v2_finalist_retest_import({
                "tester_job_id": "concurrent-handoff",
            }))
        except BaseException as error:
            errors.append(error)

    first = threading.Thread(target=call_import)
    second = threading.Thread(target=call_import, kwargs={"second": True})
    first.start()
    try:
        assert import_entered.wait(5)
        second.start()
        assert second_waiting_for_lock.wait(5)
    finally:
        release_import.set()
        first.join(timeout=5)
        if second.ident is not None:
            second.join(timeout=5)

    assert not first.is_alive() and not second.is_alive()
    assert not errors
    assert len(results) == 2
    assert results[0]["job_id"] == results[1]["job_id"]
    assert starts == [results[0]["job_id"]]
    assert len([job for job in registry.list() if job["kind"] == "strategies.performance.v2.import"]) == 1


def _bulk_import_controller(tmp_path: Path, job_id: str) -> tuple[PanelController, PanelJobRegistry]:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    registry = controller._panel_jobs
    registry.submit(
        "strategies.performance.v2.finalist-retest", {}, f"bulk-import-{job_id}",
        ("strategies.tester", "performance-v2-finalist-retest"), job_id=job_id,
    )
    registry.transition(job_id, "RUNNING")
    registry.sync(job_id, {"state": "COMMITTED", "phase": "COMMITTED", "inbox_ready": True}, runtime={
        "bulk_retest": True,
        "scope": "FINALIST",
        "test_start": "2026-01-01",
        "test_end": "2026-10-01",
        "cohort_members": [{"strategy_id": 7, "strategy_name": "alpha", "result_id": 17}],
    })
    return controller, registry


@pytest.mark.parametrize("child_state", ["COMMITTED", "FAILED", "CANCELLED"])
def test_finalist_retest_import_recovers_preallocated_child_after_parent_link_failure(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, child_state: str,
) -> None:
    controller, registry = _bulk_import_controller(tmp_path, "link-save-failure")
    starts: list[str] = []

    def start_import(
        _payload: dict[str, object], *, _internal: bool = False, job_id: str | None = None,
    ) -> dict[str, object]:
        assert _internal is True
        assert job_id is not None
        starts.append(job_id)
        child_id = job_id
        active_registry = controller._panel_jobs
        child = active_registry.submit(
            "strategies.performance.v2.import", {}, f"link-save-child:{job_id}",
            ("performance-v2-db",), job_id=job_id,
        )
        active_registry.transition(child["job_id"], "RUNNING")
        if len(starts) == 1:
            if child_state == "COMMITTED":
                active_registry.sync(child_id, {"state": "COMMITTED"}, runtime={
                    "successful_replacements": [{"strategy_id": 7, "old_result_id": 17, "new_result_id": 18}],
                    "failures": [],
                })
            elif child_state == "CANCELLED":
                active_registry.transition(child_id, "CANCELLING")
                active_registry.transition(child_id, "CANCELLED")
            else:
                active_registry.transition(child_id, "FAILED")
        return {"job_id": job_id}

    monkeypatch.setattr(controller, "strategies_performance_v2_import", start_import)
    real_sync = registry.sync

    fail_once = [True]

    def fail_parent_link(job_id: str, status: dict[str, object], *, runtime: dict[str, object] | None = None, **kwargs: object) -> dict[str, object]:
        if fail_once[0] and job_id == "link-save-failure" and runtime and runtime.get("bulk_import_job_id") in starts:
            fail_once[0] = False
            raise OSError("parent link persistence failed")
        return real_sync(job_id, status, runtime=runtime, **kwargs)  # type: ignore[arg-type]

    monkeypatch.setattr(registry, "sync", fail_parent_link)

    with pytest.raises(OSError, match="parent link persistence failed"):
        controller.strategies_performance_v2_finalist_retest_import({"tester_job_id": "link-save-failure"})

    child_id = starts[0]
    marker = registry.runtime("link-save-failure").get("bulk_import_job_id")
    assert marker == f"pending:{child_id}"
    assert registry.get(child_id)["state"] == child_state
    assert len([job for job in registry.list() if job["kind"] == "strategies.performance.v2.import"]) == 1
    recovered_registry = PanelJobRegistry(registry.journal)
    controller._panel_jobs = recovered_registry
    monkeypatch.setattr(
        controller, "_performance_v2_config",
        lambda: PerformanceV2Config(tmp_path / "performance-v2"),
    )
    monkeypatch.setattr(
        controller, "_single_mode_strategy_test",
        lambda: SimpleNamespace(status=lambda _job_id: (_ for _ in ()).throw(KeyError())),
    )
    status = controller.strategies_performance_v2_finalist_retest_status("link-save-failure")
    assert status["import_pending"] is False
    assert status["import_job_id"] == child_id
    assert status["import_job_state"] == child_state
    assert recovered_registry.runtime("link-save-failure")["bulk_import_job_id"] == child_id
    assert recovered_registry.get(child_id)["state"] == child_state
    if child_state == "COMMITTED":
        assert status["successful_replacements"] == [{
            "strategy_id": 7, "old_result_id": 17, "new_result_id": 18,
        }]
        repeated = controller.strategies_performance_v2_finalist_retest_import({"tester_job_id": "link-save-failure"})
        assert repeated["job_id"] == child_id
        assert starts == [child_id]
    else:
        retry = controller.strategies_performance_v2_finalist_retest_import({"tester_job_id": "link-save-failure"})
        assert retry["job_id"] != child_id
        assert starts == [child_id, retry["job_id"]]
    assert len([job for job in recovered_registry.list() if job["kind"] == "strategies.performance.v2.import"]) == len(starts)


def test_finalist_retest_parent_status_preserves_member_outcomes_and_import_totals(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    registry = controller._panel_jobs
    parent_id, child_id = "bulk-outcome-parent", "bulk-outcome-child"
    registry.submit("strategies.performance.v2.finalist-retest", {}, "bulk-outcome-parent", job_id=parent_id)
    registry.transition(parent_id, "RUNNING")
    registry.transition(parent_id, "COMMITTED")
    members = [
        {"strategy_id": 7, "strategy_name": "alpha", "result_id": 17},
        {"strategy_id": 8, "strategy_name": "beta", "result_id": 18},
    ]
    registry.sync(parent_id, {"state": "COMMITTED"}, runtime={
        "bulk_retest": True, "scope": "FINALIST", "cohort_members": members,
        "bulk_import_job_id": child_id, "failures": [], "successful_replacements": [],
    })
    registry.submit("strategies.performance.v2.import", {}, "bulk-outcome-child", job_id=child_id)
    registry.transition(child_id, "RUNNING")
    registry.sync(child_id, {
        "state": "COMMITTED",
        "result": {"imported_count": 1, "skipped_count": 0, "rejected_count": 1},
    }, runtime={
        "successful_replacements": [{"strategy_id": 8, "old_result_id": 18, "new_result_id": 28}],
        "failures": [{
            "strategy_id": 7, "strategy_name": "alpha", "symbol": "BTCUSDT",
            "reason": "MISSING_CURRENT_RESULT",
        }],
    })
    monkeypatch.setattr(
        controller, "_single_mode_strategy_test",
        lambda: SimpleNamespace(status=lambda _job_id: (_ for _ in ()).throw(KeyError())),
    )

    status = controller.strategies_performance_v2_finalist_retest_status(parent_id)

    assert status["successful_replacements"] == [{
        "strategy_id": 8, "old_result_id": 18, "new_result_id": 28,
    }]
    assert status["failures"] == [{
        "strategy_id": 7, "strategy_name": "alpha", "symbol": "BTCUSDT",
        "reason": "MISSING_CURRENT_RESULT",
    }]
    assert (status["imported_count"], status["rejected_count"], status["expected_count"]) == (1, 1, 2)
    assert (status["success_count"], status["failure_count"]) == (1, 1)
    assert status["import_job_state"] == "COMMITTED"


@pytest.mark.parametrize("terminal_state", ["FAILED", "CANCELLED"])
def test_finalist_retest_import_retries_after_terminal_failed_child(
    tmp_path: Path, terminal_state: str, monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, registry = _bulk_import_controller(tmp_path, f"retry-{terminal_state.lower()}")
    starts: list[str] = []

    def start_import(
        _payload: dict[str, object], *, _internal: bool = False, job_id: str | None = None,
    ) -> dict[str, object]:
        assert job_id is not None
        child_id = job_id
        starts.append(child_id)
        child = registry.submit(
            "strategies.performance.v2.import", {}, child_id,
            ("performance-v2-db",), job_id=child_id,
        )
        registry.transition(child["job_id"], "RUNNING")
        return {"job_id": child_id}

    monkeypatch.setattr(controller, "strategies_performance_v2_import", start_import)
    job_id = f"retry-{terminal_state.lower()}"
    first = controller.strategies_performance_v2_finalist_retest_import({"tester_job_id": job_id})
    assert first["job_id"] == starts[0]
    if terminal_state == "CANCELLED":
        registry.transition(starts[0], "CANCELLING")
        registry.transition(starts[0], "CANCELLED")
    else:
        registry.transition(starts[0], "FAILED")
    runtime = registry.runtime(job_id)
    runtime["outcomes_finalized"] = True
    runtime["failures"] = [{"strategy_id": 7, "strategy_name": "alpha", "reason": "IMPORT_FAILED"}]
    registry.sync(job_id, {"state": "COMMITTED"}, runtime=runtime)

    second = controller.strategies_performance_v2_finalist_retest_import({"tester_job_id": job_id})

    assert second["job_id"] == starts[1]
    assert registry.runtime(job_id).get("outcomes_finalized") is False
    assert len(starts) == 2 and starts[0] != starts[1]


@pytest.mark.parametrize("payload", [{}, {"job_id": ""}, {"job_id": 7}, {"job_id": "bulk", "extra": True}])
def test_finalist_retest_equity_filter_requires_exact_parent_job_payload(
    tmp_path: Path, payload: dict[str, object],
) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    with pytest.raises(PerformanceV2ApiError) as error:
        controller.strategies_performance_v2_finalist_retest_equity_filter(payload)
    assert error.value.code == "INVALID_REQUEST"


def test_finalist_retest_equity_filter_identity_includes_frozen_snapshot() -> None:
    base = {
        "scope": "FINALIST",
        "cohort_sha256": "cohort",
        "cohort_members": [{"strategy_id": 1, "result_id": 2}],
        "successful_replacements": [{"strategy_id": 1, "new_result_id": 3}],
        "equity_filter_source_revision_digest": "source",
        "equity_filter_snapshot": {"selection_config_sha256": "config-a", "regime_policy_sha256": "policy"},
    }
    changed = deepcopy(base)
    changed["equity_filter_snapshot"]["selection_config_sha256"] = "config-b"

    assert PanelController._equity_filter_identity_seed(base) != PanelController._equity_filter_identity_seed(changed)


def test_finalist_retest_equity_filter_status_reads_live_child_progress(tmp_path: Path) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    registry = controller._panel_jobs
    parent_id = "equity-live-status-parent"
    child_id = "equity-live-status-child"
    registry.submit("strategies.performance.v2.finalist-retest", {}, parent_id, job_id=parent_id)
    registry.transition(parent_id, "RUNNING")
    registry.transition(parent_id, "COMMITTED")
    registry.submit(
        "strategies.performance.v2.finalist-retest.equity-filter", {}, child_id,
        ("performance-v2-equity-filter",), job_id=child_id,
    )
    registry.transition(child_id, "RUNNING")
    registry.sync(child_id, {"state": "RUNNING", "phase": "WARMING_CACHE", "progress": {"current": 2, "total": 5, "unit": "pairs"}})
    registry.sync(parent_id, {"state": "COMMITTED"}, runtime={
        "bulk_retest": True,
        "outcomes_finalized": True,
        "successful_replacements": [{"strategy_id": 1, "new_result_id": 2}],
        "equity_filter_snapshot": {"schema_version": 1},
        "equity_filter_job_id": child_id,
    })

    status = controller._equity_filter_child_status(parent_id)

    assert status["state"] == "RUNNING"
    assert status["phase"] == "WARMING_CACHE"
    assert status["progress"]["current"] == 2
    assert status["progress"]["total"] == 5


def test_finalist_retest_equity_filter_clears_lease_on_base_exception(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    registry = controller._panel_jobs
    parent_id = "equity-base-exception-parent"
    child_id = "equity-base-exception-child"
    registry.submit("strategies.performance.v2.finalist-retest", {}, parent_id, job_id=parent_id)
    registry.transition(parent_id, "RUNNING")
    registry.transition(parent_id, "COMMITTED")
    registry.submit(
        "strategies.performance.v2.finalist-retest.equity-filter", {}, child_id,
        ("performance-v2-equity-filter",), job_id=child_id,
    )
    registry.sync(parent_id, {"state": "COMMITTED"}, runtime={
        "equity_filter_lock_holder": child_id,
        "equity_filter_lock_expires_at": (datetime.now(UTC) + timedelta(minutes=5)).isoformat(),
    })

    def interrupt(*_args, **_kwargs):
        raise KeyboardInterrupt()

    monkeypatch.setattr(controller, "_run_bulk_retest_equity_filter", interrupt)

    with pytest.raises(KeyboardInterrupt):
        controller._run_bulk_retest_equity_filter_job(parent_id, child_id)

    assert registry.get(child_id)["state"] == "FAILED"
    assert registry.runtime(parent_id)["equity_filter_lock_holder"] is None


def test_recalculate_requires_offline_migration_for_existing_v8_database(tmp_path: Path) -> None:
    controller, database, _ = _controller_for_windows(tmp_path, legacy_v11=True)
    with duckdb.connect(str(database)) as connection:
        _make_v8_catalog(connection)
    before_file = (sha256(database.read_bytes()).hexdigest(), database.stat().st_mtime_ns)

    with pytest.raises(PerformanceV2ApiError) as raised:
        controller.strategies_performance_v2_recalculate_all()

    assert (raised.value.code, raised.value.status) == ("PERFORMANCE_V2_MIGRATION_REQUIRED", 409)
    assert (sha256(database.read_bytes()).hexdigest(), database.stat().st_mtime_ns) == before_file
    with duckdb.connect(str(database), read_only=True) as connection:
        assert connection.execute("select value from schema_info where key = 'schema_version'").fetchone() == ("8",)
        assert connection.execute(
            "select count(*) from information_schema.tables where table_name = 'strategy_rejection_sources'"
        ).fetchone() == (0,)
        assert connection.execute(
            "select count(*) from information_schema.columns where table_name = 'selection_results' "
            "and column_name = 'equity_regime_json'"
        ).fetchone() == (0,)


def test_panel_schema_preflight_initializes_an_empty_new_database(tmp_path: Path) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    database.unlink()
    database.write_bytes(b"")

    controller._ensure_performance_v2_schema(database)

    with duckdb.connect(str(database), read_only=True) as connection:
        assert connection.execute(
            "select value from schema_info where key = 'schema_version'"
        ).fetchone() == ("12",)
        assert connection.execute(
            "select count(*) from information_schema.tables where table_name = 'strategy_rejection_sources'"
        ).fetchone() == (1,)


def test_panel_schema_preflight_migrates_v9_to_v11(tmp_path: Path) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    from tests.test_performance_v2_store import _initialize_v9_fixture

    database.unlink()
    with duckdb.connect(str(database)) as connection:
        _initialize_v9_fixture(connection)

    controller._ensure_performance_v2_schema(database)

    with duckdb.connect(str(database), read_only=True) as connection:
        assert connection.execute("select value from schema_info where key = 'schema_version'").fetchone() == ("12",)
        assert connection.execute(
            "select is_nullable from information_schema.columns "
            "where table_name = 'selection_review_rows' and column_name = 'user_status'"
        ).fetchone() == ("YES",)


def test_panel_schema_preflight_migrates_exact_v10_catalog_to_v11(tmp_path: Path) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    from tests.test_performance_v2_store import _prepare_v10_migration_fixture

    database.unlink()
    with duckdb.connect(str(database)) as connection:
        _prepare_v10_migration_fixture(connection)

    controller._ensure_performance_v2_schema(database)

    with duckdb.connect(str(database), read_only=True) as connection:
        assert connection.execute(
            "select value from schema_info where key = 'schema_version'"
        ).fetchone() == ("12",)


def test_panel_schema_preflight_maps_corrupt_database_to_schema_invalid(tmp_path: Path) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    database.write_bytes(b"not a DuckDB database")
    before = (sha256(database.read_bytes()).hexdigest(), database.stat().st_mtime_ns)

    with pytest.raises(PerformanceV2ApiError) as raised:
        controller._ensure_performance_v2_schema(database)

    assert (raised.value.code, raised.value.status) == ("PERFORMANCE_V2_SCHEMA_INVALID", 500)
    assert (sha256(database.read_bytes()).hexdigest(), database.stat().st_mtime_ns) == before


@pytest.mark.parametrize("enabled", [False, True])
def test_selection_workbook_handles_equity_stage_only_when_enabled_column_is_present(
    tmp_path: Path, enabled: bool,
) -> None:
    request = parse_selection_request({"symbol": "BTCUSDT", "side": "LONG", "stages": [
        {"id": "filter_equity_regime", "enabled": enabled, "scope": "pair_side"},
    ]})
    row = {
        "strategy_id": 1, "strategy_name": "candidate", "timeframe": "1h", "order_count": 1,
        "robust_pnl_30d_pct": Decimal("10"), "worst_drawdown_pct": Decimal("5"),
        "worst_holding_p95_minutes": Decimal("30"), "ab_stability_ratio": Decimal(".9"),
        "first_shift_bp": 100, "minimum_plateau_point_count": 10, "close_ma_len": 3,
        "dd5_proxy": Decimal("10"), "capital_proxy": Decimal("5"),
        "holding_p95_minutes": Decimal("30"),
    }
    if enabled:
        start = datetime(2026, 1, 1, tzinfo=UTC)
        end = start + timedelta(days=28)
        facts = calculate_equity_quality_facts(1, start, end, (
            EquitySample(1, 0, start, Decimal("100")),
            EquitySample(1, 1, end, Decimal("110")),
        ))
        facts_sha256 = sha256(json.dumps(
            facts.to_canonical_dict(), sort_keys=True, separators=(",", ":"), ensure_ascii=False,
        ).encode("utf-8")).hexdigest()
        row.update({
            "result_id": 1,
            "_equity_cache": {
                "status": "FRESH", "facts": facts,
                "source_revision": "0" * 64, "facts_sha256": facts_sha256,
            },
        })
        regime_end = start + timedelta(days=42)
        regime_samples = tuple(
            EquityRegimeSample(1, index, start + timedelta(days=index), Decimal("100") * Decimal("1.01") ** index)
            for index in range(43)
        )
        assessment = classify_equity_regime(1, start, regime_end, regime_samples)
        regime_facts = assessment.facts
        regime_facts_json = encode_equity_regime_facts(regime_facts)
        row["_equity_regime_cache"] = {
            "status": "FRESH",
            "assessment": assessment,
            "facts": regime_facts,
            "source_revision": "0" * 64,
            "facts_sha256": sha256(regime_facts_json.encode("utf-8")).hexdigest(),
            "classifier_algo_version": EQUITY_REGIME_ALGORITHM_VERSION,
            "equity_regime_json": encode_equity_regime_assessment(assessment),
        }
    result = panel_module.run_selection(pd.DataFrame([row]), request)
    path = write_selection_workbook(result, tmp_path / f"equity-{enabled}.xlsx", request)
    headers = [cell.value for cell in load_workbook(path, data_only=True)["All candidates"][1]]

    if enabled:
        assert result.loc[0, "equity_regime_decision"] == "PASS"
    assert "eliminated_by_filter_equity_regime" not in headers
    assert ("Regime rank" in headers) is enabled
    assert "Equity state" not in headers


def test_valid_v11_schema_check_does_not_open_writer_with_reader_present(tmp_path: Path, monkeypatch) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    original_connect = duckdb.connect

    def reject_writer(*args, **kwargs):
        if not kwargs.get("read_only", False):
            raise AssertionError("v11 schema preflight attempted a writer open")
        return original_connect(*args, **kwargs)

    with duckdb.connect(str(database), read_only=True) as reader:
        assert reader.execute("select value from schema_info where key = 'schema_version'").fetchone() == ("12",)
        monkeypatch.setattr(panel_module.duckdb, "connect", reject_writer)
        controller._ensure_performance_v2_schema(database)


def test_cached_v11_schema_check_skips_reopen_with_reader_present(tmp_path: Path, monkeypatch) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    controller._ensure_performance_v2_schema(database)

    with duckdb.connect(str(database), read_only=True) as reader:
        assert reader.execute("select value from schema_info where key = 'schema_version'").fetchone() == ("12",)

        def unexpected_reopen(*_args, **_kwargs):
            raise AssertionError("schema cache hit reopened the database")

        monkeypatch.setattr(panel_module.duckdb, "connect", unexpected_reopen)
        controller._ensure_performance_v2_schema(database)


def test_v11_schema_check_rejects_malformed_catalog(tmp_path: Path) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    with duckdb.connect(str(database)) as connection:
        connection.execute("drop table strategy_rejection_sources")

    with pytest.raises(PerformanceV2ApiError) as raised:
        controller._ensure_performance_v2_schema(database)

    assert raised.value.code == "PERFORMANCE_V2_SCHEMA_INVALID"
    assert raised.value.status == 500


def test_v11_schema_check_repairs_missing_window_column(tmp_path: Path) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    with duckdb.connect(str(database)) as connection:
        connection.execute("alter table window_metrics drop column holding_seconds")

    controller._ensure_performance_v2_schema(database)

    with duckdb.connect(str(database), read_only=True) as connection:
        repaired = connection.execute(
            """select data_type from information_schema.columns
                where table_schema = 'main' and table_name = 'window_metrics'
                  and column_name = 'holding_seconds'"""
        ).fetchone()
    assert repaired == ("DECIMAL(38,12)",)


def test_selection_recalculate_all_does_not_prepare_optimizer_inputs(tmp_path: Path, monkeypatch) -> None:
    controller, _, _ = _controller_for_windows(tmp_path)
    import mrs3.panel as panel_module
    assert not hasattr(panel_module, "prepare_current_optimizer_inputs")

    assert controller.strategies_performance_v2_recalculate_all()["status"] == "READY"


def test_selection_recalculate_all_skips_pairs_with_ready_facts(tmp_path: Path) -> None:
    controller, database, result_id = _controller_for_windows(tmp_path)

    first = controller.strategies_performance_v2_recalculate_all()
    with duckdb.connect(str(database), read_only=True) as check:
        assert check.execute(
            "select count(*) from equity_quality_metrics where algo_version = ? and result_id = ?",
            [EQUITY_REGIME_ALGORITHM_VERSION, result_id],
        ).fetchone() == (1,)
        assert check.execute(
            "select count(*) from equity_quality_metrics where algo_version <> ? and result_id = ?",
            [EQUITY_REGIME_ALGORITHM_VERSION, result_id],
        ).fetchone() == (0,)
    second = controller.strategies_performance_v2_recalculate_all()

    assert first == {"status": "READY", "total_pairs": 1, "recalculated_pairs": 1, "ready_pairs": 0}
    assert second == {"status": "READY", "total_pairs": 1, "recalculated_pairs": 0, "ready_pairs": 1}


def test_recalculate_all_progress_is_visible_and_rejects_duplicate_run(tmp_path: Path, monkeypatch) -> None:
    controller, _, _ = _controller_for_windows(tmp_path)
    entered = threading.Event()
    release = threading.Event()
    responses: list[dict[str, object]] = []

    def paused_prepare(*_args, on_batch_complete, **_kwargs) -> None:
        on_batch_complete(1)
        entered.set()
        assert release.wait(5)

    monkeypatch.setattr(panel_module, "prepare_selection_window_cache", paused_prepare)
    worker = threading.Thread(target=lambda: responses.append(controller.strategies_performance_v2_recalculate_all()))
    worker.start()
    try:
        assert entered.wait(5)
        assert controller.strategies_performance_v2_recalculate_all_progress() == {
            "status": "RUNNING", "total_pairs": 1, "ready_pairs": 0,
            "planned_pairs": 1, "completed_pairs": 0,
            "total_strategies": 1, "completed_strategies": 1,
            "current_pair": "BTCUSDT/LONG", "error": None,
        }
        with pytest.raises(PerformanceV2ApiError) as raised:
            controller.strategies_performance_v2_recalculate_all()
        assert raised.value.status == 409
        assert raised.value.code == "PERFORMANCE_V2_RECALCULATE_IN_PROGRESS"
    finally:
        release.set()
        worker.join(timeout=5)
    assert not worker.is_alive()
    assert responses == [{"status": "READY", "total_pairs": 1, "recalculated_pairs": 1, "ready_pairs": 0}]
    assert controller.strategies_performance_v2_recalculate_all_progress()["status"] == "READY"
    assert controller.strategies_performance_v2_recalculate_all_progress()["completed_pairs"] == 1


def test_recalculate_all_progress_retains_committed_count_on_error(tmp_path: Path, monkeypatch) -> None:
    controller, _, _ = _controller_for_windows(tmp_path)

    def failing_prepare(*_args, on_batch_complete, **_kwargs) -> None:
        on_batch_complete(1)
        raise OSError("later batch failed")

    monkeypatch.setattr(panel_module, "prepare_selection_window_cache", failing_prepare)
    with pytest.raises(PerformanceV2ApiError, match="later batch failed"):
        controller.strategies_performance_v2_recalculate_all()
    status = controller.strategies_performance_v2_recalculate_all_progress()
    assert status["status"] == "FAILED"
    assert status["completed_strategies"] == 1
    assert status["completed_pairs"] == 0
    assert status["error"] == "later batch failed"


def test_recalculate_all_progress_http_reads_memory_only(tmp_path: Path, monkeypatch) -> None:
    controller, _, _ = _controller_for_windows(tmp_path)
    server, thread = _http_server(controller)
    monkeypatch.setattr(panel_module.duckdb, "connect", lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("DB opened")))
    try:
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        status, body = _http_json(connection, "GET", "/api/v2/strategies/performance-v2/recalculate-all/progress")
        connection.close()
        assert status == 200
        assert body["status"] == "IDLE"
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_panel_writer_guard_preserves_store_errors_raised_by_body(tmp_path: Path) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    from mrs3.performance_v2_equity_cache import EquitySourceChangedError
    from mrs3.performance_v2_store import PerformanceV2StoreError

    assert not issubclass(EquitySourceChangedError, PerformanceV2StoreError)
    error = PerformanceV2StoreError("body failure code")
    with pytest.raises(PerformanceV2StoreError, match="body failure code") as raised:
        with controller._performance_v2_writer_guard(database):
            raise error
    assert raised.value is error


def test_panel_creates_bound_database_and_refuses_a_different_exchange(tmp_path: Path) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    database.unlink()
    database.write_bytes(b"")
    (tmp_path / "config.performance.json").write_text(json.dumps({"unified_performance_v2": {
        "database_root": "data", "evaluation_protocol": "v2", "exchange": "Bybit"}}), encoding="utf-8")

    controller._ensure_performance_v2_schema(database)

    with duckdb.connect(str(database), read_only=True) as connection:
        assert dict(connection.execute(
            "select key, value from schema_info where key in ('schema_version', 'evaluation_protocol', 'exchange')"
        ).fetchall()) == {"schema_version": "12", "evaluation_protocol": "v2", "exchange": "Bybit"}

    (tmp_path / "config.performance.json").write_text(json.dumps({"unified_performance_v2": {
        "database_root": "data", "evaluation_protocol": "v2", "exchange": "OKX"}}), encoding="utf-8")
    restarted = PanelController(tmp_path, tmp_path / "config.local.json")
    restarted._ensure_performance_v2_schema(database)          # ADR-0083: the exchange no longer binds the database
    with duckdb.connect(str(database), read_only=True) as connection:
        assert connection.execute("select value from schema_info where key = 'exchange'").fetchone() == ("Bybit",)


def test_panel_rechecks_identity_when_only_the_configuration_changes(tmp_path: Path) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    database.unlink()
    database.write_bytes(b"")
    config_path = tmp_path / "config.performance.json"
    config_path.write_text(json.dumps({"unified_performance_v2": {
        "database_root": "data", "evaluation_protocol": "v2", "exchange": "Bybit"}}), encoding="utf-8")
    controller._ensure_performance_v2_schema(database)
    controller._ensure_performance_v2_schema(database)  # cached

    config_path.write_text(json.dumps({"unified_performance_v2": {"database_root": "data"}}), encoding="utf-8")
    with pytest.raises(PerformanceV2ApiError) as raised:
        controller._ensure_performance_v2_schema(database)
    assert (raised.value.code, raised.value.status) == ("EVALUATION_PROTOCOL_MISMATCH", 409)


def test_panel_refuses_mismatching_v11_database_without_migrating_it(tmp_path: Path) -> None:
    controller, database, _ = _controller_for_windows(tmp_path, legacy_v11=True)
    (tmp_path / "config.performance.json").write_text(json.dumps({"unified_performance_v2": {
        "database_root": "data", "evaluation_protocol": "v2", "exchange": "Bybit"}}), encoding="utf-8")

    with pytest.raises(PerformanceV2ApiError) as raised:
        controller._ensure_performance_v2_schema(database)

    assert raised.value.code == "EVALUATION_PROTOCOL_MISMATCH"
    with duckdb.connect(str(database), read_only=True) as connection:
        assert connection.execute(
            "select value from schema_info where key = 'schema_version'"
        ).fetchone() == ("11",)


# -- review of Panel stage 4c: the single-pair recalculate left with the Pareto card; its lock, worker, error and
# -- incremental-cache checks now run against recalculate-all («Посчитать метрики» of «Стратегии»).

def test_recalculate_all_prepares_the_cache_under_the_writer_lock(tmp_path: Path, monkeypatch) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    controller._ensure_performance_v2_schema(database)

    class TrackingRLock:
        def __init__(self) -> None:
            self._lock = threading.RLock()
            self.owner: int | None = None
            self.depth = 0

        def __enter__(self):
            self._lock.acquire()
            self.owner, self.depth = threading.get_ident(), self.depth + 1
            return self

        def __exit__(self, *_args) -> None:
            self.depth -= 1
            if self.depth == 0:
                self.owner = None
            self._lock.release()

    lock = TrackingRLock()
    controller._performance_v2_writer_lock = lock
    held_during_prepare: list[bool] = []

    def prepare(path, *_args, on_batch_complete=None, **_kwargs):
        held_during_prepare.append(lock.owner == threading.get_ident() and lock.depth > 0)
        if on_batch_complete is not None:
            on_batch_complete(1)

    monkeypatch.setattr(panel_module, "prepare_selection_window_cache", prepare)
    result = controller.strategies_performance_v2_recalculate_all()
    assert result["status"] == "READY" and result["total_pairs"] > 0 and result["recalculated_pairs"] > 0
    assert held_during_prepare and all(held_during_prepare)
    assert lock.depth == 0 and controller.strategies_performance_v2_recalculate_all_progress()["status"] == "READY"


def test_recalculate_all_runs_real_window_workers_under_controller_lock(tmp_path: Path, monkeypatch) -> None:
    controller, database, first_result_id = _controller_for_windows(tmp_path)
    (tmp_path / "config.performance.json").write_text(
        json.dumps({"unified_performance_v2": {"database_root": "data", "workers": 2}}), encoding="utf-8")
    with duckdb.connect(str(database)) as connection:
        now = datetime(2026, 1, 1, tzinfo=UTC)
        strategy_id = int(connection.execute(
            """insert into strategies (strategy_name, symbol, side, timeframe, close_ma_len,
               order_count, analysis_run_id, candidate_identity, lifecycle_status,
               created_at_utc, updated_at_utc) values ('beta', 'BTCUSDT', 'LONG', '1h',
               3, 1, 'run', 'candidate-beta', 'ACTIVE', ?, ?) returning strategy_id""", [now, now]).fetchone()[0])
        result_id = int(connection.execute(
            """insert into strategy_results (strategy_id, report_start_utc, report_end_utc, exchange,
               commission_rate, initial_balance, final_balance, total_pnl, total_pnl_pct,
               max_drawdown, max_drawdown_pct, total_fees, total_trades, imported_at_utc)
               select ?, report_start_utc, report_end_utc, exchange, commission_rate, initial_balance,
                      final_balance, total_pnl, total_pnl_pct, max_drawdown, max_drawdown_pct,
                      total_fees, total_trades, imported_at_utc
                 from strategy_results where result_id = ? returning result_id""",
            [strategy_id, first_result_id]).fetchone()[0])
        connection.execute("update strategies set current_result_id = ? where strategy_id = ?", [result_id, strategy_id])
        connection.execute(
            """insert into strategy_actions (result_id, action_index, timestamp_utc, symbol, order_id,
                   action, size, post_size, post_side, pnl, fee, balance, price, cost, raw_action_json)
               select ?, action_index, timestamp_utc, symbol, order_id, action, size, post_size,
                      post_side, pnl, fee, balance, price, cost, raw_action_json
                 from strategy_actions where result_id = ?""", [result_id, first_result_id])
        connection.execute(
            "insert into strategy_equity select ?, sample_index, timestamp_utc, wallet, equity from strategy_equity where result_id = ?",
            [result_id, first_result_id])

    workers_entered, release_workers = threading.Event(), threading.Event()
    worker_ids: list[int] = []
    results: list[dict[str, object]] = []
    errors: list[Exception] = []

    class TrackingRLock:
        def __init__(self) -> None:
            self._lock = threading.RLock()
            self.depth = 0
            self.owner_thread_id: int | None = None

        def __enter__(self):
            self._lock.acquire()
            if self.depth == 0:
                self.owner_thread_id = threading.get_ident()
            self.depth += 1
            return self

        def __exit__(self, *_args) -> None:
            self.depth -= 1
            if self.depth == 0:
                self.owner_thread_id = None
            self._lock.release()

    tracking_lock = TrackingRLock()
    controller._performance_v2_writer_lock = tracking_lock
    import mrs3.performance_v2_selection as selection_module
    real_worker = selection_module._selection_window_job_from_args
    orchestrator: list[int] = []

    def gated_worker(args):
        assert orchestrator and tracking_lock.owner_thread_id == orchestrator[0] and tracking_lock.depth > 0
        worker_ids.append(int(args[1]))
        if len(worker_ids) == 2:
            workers_entered.set()
        assert release_workers.wait(5)
        return real_worker(args)

    monkeypatch.setattr(selection_module, "_selection_window_job_from_args", gated_worker)

    def recalculate() -> None:
        orchestrator.append(threading.get_ident())
        try:
            results.append(controller.strategies_performance_v2_recalculate_all())
        except Exception as error:
            errors.append(error)

    writer = threading.Thread(target=recalculate)
    writer.start()
    try:
        assert workers_entered.wait(10)
        assert len(set(worker_ids)) == 2
    finally:
        release_workers.set()
        writer.join(timeout=10)
    assert not writer.is_alive() and not errors
    assert results[0]["status"] == "READY"


@pytest.mark.parametrize(
    ("error_name", "code", "status"),
    [("EquitySchemaUpgradeRequiredError", "EQUITY_SCHEMA_UPGRADE_REQUIRED", 409),
     ("EquityCacheSchemaInvalidError", "PERFORMANCE_V2_SCHEMA_INVALID", 500)],
)
def test_recalculate_all_maps_equity_cache_errors_to_typed_api_errors(
    tmp_path: Path, monkeypatch, error_name: str, code: str, status: int,
) -> None:
    # the equity cache errors arise while recalculate-all looks up the missing strategies of each pair
    import mrs3.performance_v2_selection as selection_module

    controller, _database, _ = _controller_for_windows(tmp_path)
    error_class = getattr(selection_module, error_name)

    def failing(*_args, **_kwargs):
        raise error_class("equity cache")

    monkeypatch.setattr(panel_module, "selection_cache_missing_strategy_ids", failing)
    with pytest.raises(PerformanceV2ApiError) as raised:
        controller.strategies_performance_v2_recalculate_all()

    assert (raised.value.code, raised.value.status) == (code, status)
    assert controller.strategies_performance_v2_recalculate_all_progress()["status"] == "FAILED"


def test_recalculate_all_tracks_only_new_current_results_for_add_replace_and_repeat(
    tmp_path: Path, monkeypatch,
) -> None:
    controller, database, _ = _controller_for_windows(tmp_path)
    request = parse_selection_request({"symbol": "BTCUSDT", "side": "LONG", "stages": []})
    controller.strategies_performance_v2_recalculate_all()

    def add_current(connection: duckdb.DuckDBPyConnection, name: str, owner_id: int | None = None) -> tuple[int, int]:
        now = datetime(2026, 1, 1, tzinfo=UTC)
        if owner_id is None:
            strategy_id = int(connection.execute(
                """insert into strategies (strategy_name, symbol, side, timeframe, close_ma_len,
                   order_count, analysis_run_id, candidate_identity, lifecycle_status,
                   created_at_utc, updated_at_utc) values (?, 'BTCUSDT', 'LONG', '1h',
                   3, 1, 'run', ?, 'ACTIVE', ?, ?) returning strategy_id""", [name, name, now, now]).fetchone()[0])
        else:
            old_result_id = int(connection.execute(
                "select current_result_id from strategies where strategy_id = ?", [owner_id]).fetchone()[0])
            old_holder_id = int(connection.execute(
                """insert into strategies (strategy_name, symbol, side, timeframe, close_ma_len,
                   order_count, analysis_run_id, candidate_identity, lifecycle_status,
                   created_at_utc, updated_at_utc) values (?, 'BTCUSDT', 'LONG', '1h',
                   3, 1, 'run', ?, 'DISCARDED', ?, ?) returning strategy_id""",
                [f"{name}-old-holder", f"{name}-old-holder", now, now]).fetchone()[0])
            connection.execute("update strategies set current_result_id = null where strategy_id = ?", [owner_id])
            for table in ("strategy_actions", "strategy_equity", "window_metrics", "optimizer_prepared_inputs"):
                connection.execute(f"delete from {table} where result_id = ?", [old_result_id])
            connection.execute("update strategy_results set strategy_id = ? where result_id = ?", [old_holder_id, old_result_id])
            strategy_id = owner_id
        result_id = int(connection.execute(
            """insert into strategy_results (strategy_id, report_start_utc, report_end_utc, exchange,
               commission_rate, initial_balance, final_balance, total_pnl, total_pnl_pct,
               max_drawdown, max_drawdown_pct, total_fees, total_trades, imported_at_utc)
               values (?, ?, ?, 'Bybit', .0004, 100, 110, 10, 10, 0, 0, 2, 2, ?) returning result_id""",
            [strategy_id, now, datetime(2026, 1, 5, tzinfo=UTC), now]).fetchone()[0])
        connection.execute("update strategies set current_result_id = ? where strategy_id = ?", [result_id, strategy_id])
        connection.executemany(
            "insert into strategy_actions (result_id, action_index, timestamp_utc, symbol, order_id, action, size, post_size, post_side, pnl, fee, balance, raw_action_json) values (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            [(result_id, 0, now, "BTCUSDT", 1, "opened", 1, 1, "long", 0, 1, 100, None),
             (result_id, 1, datetime(2026, 1, 2, tzinfo=UTC), "BTCUSDT", 1, "closed", 1, 0, "", 10, 1, 110, None)])
        connection.executemany(
            "insert into strategy_equity values (?, ?, ?, ?, ?)",
            [(result_id, 0, now, 100, 100), (result_id, 1, datetime(2026, 1, 2, tzinfo=UTC), 110, 110),
             (result_id, 2, datetime(2026, 1, 5, tzinfo=UTC), 110, 110)])
        return strategy_id, result_id

    with duckdb.connect(str(database)) as connection:
        beta_id, _ = add_current(connection, "beta")
        assert selection_cache_missing_strategy_ids(connection, request, SelectionConfig()) == (beta_id,)
    controller.strategies_performance_v2_recalculate_all()
    with duckdb.connect(str(database)) as connection:
        alpha_id = int(connection.execute("select strategy_id from strategies where strategy_name = 'alpha'").fetchone()[0])
        replacement_id, _ = add_current(connection, "alpha-replacement", alpha_id)
        assert replacement_id == alpha_id
        assert selection_cache_missing_strategy_ids(connection, request, SelectionConfig()) == (alpha_id,)
    controller.strategies_performance_v2_recalculate_all()
    with duckdb.connect(str(database)) as connection:
        assert selection_cache_missing_strategy_ids(connection, request, SelectionConfig()) == ()

    calls: list[tuple[object, ...]] = []
    monkeypatch.setattr(panel_module, "prepare_selection_window_cache", lambda *args, **kwargs: calls.append((args, kwargs)))
    assert controller.strategies_performance_v2_recalculate_all() == {
        "status": "READY", "total_pairs": 1, "recalculated_pairs": 0, "ready_pairs": 1,
    }
    assert calls == []                                          # nothing missing: nothing recalculated


def test_removed_screen_routes_are_disconnected(tmp_path: Path) -> None:
    # Panel stage 4c: removed — and disconnected, not just hidden
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    server, thread = _http_server(controller)
    base = "/api/v2/strategies/performance-v2"
    try:
        for method, path in [("GET", f"{base}/export"), ("GET", f"{base}/catalog"), ("GET", f"{base}/selection-settings"),
                             ("HEAD", f"{base}/export"), ("OPTIONS", f"{base}/export"),
                             *[("POST", f"{base}/{name}") for name in (
                                 "windows", "selection", "selection/preview", "selection/export",
                                 "selection/write-rejected", "selection-preview", "selection-cache-status",
                                 "recalculate", "selection-settings", "selection-review-import",
                                 "selection/aggregate-import", "selection-user-fields-import", "export")]]:
            connection = HTTPConnection("127.0.0.1", server.server_port, timeout=5)
            try:
                body = b"{}" if method == "POST" else None
                headers = {"Content-Type": "application/json", "Content-Length": "2"} if body else {}
                connection.request(method, path, body=body, headers=headers)
                response = connection.getresponse()
                response.read()
                assert response.status == 404, (method, path, response.status)
            finally:
                connection.close()
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)
