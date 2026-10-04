from datetime import UTC, datetime, timedelta
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import json

import duckdb
from openpyxl import load_workbook
import pandas as pd
import pytest

from mrs3.performance_v2_equity_regime import EquityRegimeSample, classify_equity_regime
from mrs3.performance_v2_equity_regime_cache import (
    ALGORITHM_VERSION,
    encode_equity_regime_assessment,
    encode_equity_regime_facts,
    equity_regime_source_revision,
    upsert_equity_regime_facts_checked,
)
from mrs3.performance_v2_equity_cache import current_equity_source_metadata
from mrs3.performance_v2_selection import (
    SelectionConfig,
    parse_selection_request,
    write_selection_workbook,
)
from mrs3.performance_v2_selection_review import (
    SelectionReviewError,
    apply_prior_rejected,
    effective_selection_decisions,
    new_run_metadata,
    persist_selection_snapshot,
    persist_selection_snapshots,
    import_selection_review,
)
from mrs3.performance_v2_store import initialize_performance_v2


END = datetime(2026, 9, 2, tzinfo=UTC)
START = END - timedelta(days=42)
NOW = END + timedelta(days=1)


def _database() -> duckdb.DuckDBPyConnection:
    connection = duckdb.connect(":memory:")
    initialize_performance_v2(connection)
    connection.execute(
        """insert into strategies values
           (1, 'strategy-1', 'BTCUSDT', 'LONG', '1h', 5, 1, 'run', 'candidate',
            'ACTIVE', null, ?, ?)""",
        [NOW, NOW],
    )
    connection.execute(
        """insert into strategy_results (
               result_id, strategy_id, report_start_utc, report_end_utc, exchange,
               commission_rate, initial_balance, final_balance, total_pnl, total_pnl_pct,
               max_drawdown, max_drawdown_pct, total_fees, total_trades, imported_at_utc
           ) values (101, 1, ?, ?, 'Bybit', .0004, 100, 90, -10, -10, 40, 40, 1, 10, ?)""",
        [START, END, NOW],
    )
    connection.execute("update strategies set current_result_id = 101 where strategy_id = 1")
    return connection


def _request(*, filter_enabled: bool = True, rank_enabled: bool = False):
    stages = []
    if filter_enabled:
        stages.append({"id": "filter_equity_regime", "enabled": True, "scope": "pair_side"})
    if rank_enabled:
        stages.append({
            "id": "rank_robust_top_n", "enabled": True, "scope": "pair_side",
            "top_n": 1, "method": "equity_quality_v1",
        })
    return parse_selection_request({"symbol": "BTCUSDT", "side": "LONG", "stages": stages})


def _samples():
    return tuple(
        EquityRegimeSample(101, index, END + timedelta(days=day), Decimal(f"{value}.000000000000"))
        for index, (day, value) in enumerate(((-42, 200), (-28, 100), (-14, 80), (-7, 60), (0, 90)))
    )


def _assessment(*, evaluated: bool = True):
    samples = _samples() if evaluated else ()
    return classify_equity_regime(101, START, END, samples)


def _result(connection, assessment, *, filter_enabled: bool, cache: bool = True):
    connection.execute("delete from strategy_equity where result_id = 101")
    if assessment.facts.raw_sample_count:
        connection.executemany(
            """insert into strategy_equity (result_id, sample_index, timestamp_utc, wallet, equity)
               values (101, ?, ?, 100, ?)""",
            [[sample.sample_index, sample.timestamp_utc, sample.equity] for sample in _samples()],
        )
    result = pd.DataFrame([{
        "strategy_id": 1,
        "result_id": 101,
        "strategy_name": "strategy-1",
        "symbol": "BTCUSDT",
        "side": "LONG",
        "auto_status": "FILTERED",
        "finalist": False,
        "elimination_reason": "EQUITY_REGIME_DROP",
        "final_score": None,
        "final_rank": None,
        "auto_analog_of_strategy_id": None,
        "prior_rejected": False,
    }])
    source = current_equity_source_metadata(connection, 101)
    source_revision = equity_regime_source_revision(source)
    now = NOW
    if cache:
        upsert_equity_regime_facts_checked(
            connection, source, assessment.facts, calculated_at_utc=now,
        )
    facts_payload = encode_equity_regime_facts(assessment.facts)
    result.attrs["equity_regime_evidence"] = {"1": {
        "strategy_id": 1,
        "result_id": 101,
        "source_revision": source_revision,
        "facts_sha256": sha256(facts_payload.encode()).hexdigest(),
        "classifier_algo_version": ALGORITHM_VERSION,
        "equity_regime_json": encode_equity_regime_assessment(assessment),
        "equity_filter_enabled": filter_enabled,
    }}
    return result


def test_filter_drop_publishes_canonical_assessment_and_each_hard_reason():
    connection = _database()
    request = _request()
    assessment = _assessment()
    assert assessment.state == assessment.decision == "DROP"
    assert set(assessment.reasons) == {"DD_14_7_GTE_23", "W28_DOWN"}
    result = _result(connection, assessment, filter_enabled=True)
    metadata = new_run_metadata(connection, request)
    payload = result.attrs["equity_regime_evidence"]["1"]["equity_regime_json"]

    persist_selection_snapshot(connection, request, SelectionConfig(), result, metadata, b"review")

    stored = connection.execute(
        "select equity_regime_json from selection_results where selection_run_id = ? and strategy_id = 1",
        [metadata["selection_run_id"]],
    ).fetchone()[0]
    assert stored == payload
    assert json.loads(stored)["state"] == "DROP"
    published_request = json.loads(connection.execute(
        "select request_json from selection_runs where selection_run_id = ?",
        [metadata["selection_run_id"]],
    ).fetchone()[0])
    assert published_request["equity_regime_snapshot"] == {
        "algorithm_version": ALGORITHM_VERSION,
        "sources": {"1": {
            "result_id": 101,
            "source_revision": result.attrs["equity_regime_evidence"]["1"]["source_revision"],
        }},
    }
    assert connection.execute(
        """select strategy_id, source_kind, reason_code, first_result_id,
                  first_selection_run_id, classifier_algo_version
             from strategy_rejection_sources order by reason_code"""
    ).fetchall() == [
        (1, "EQUITY_REGIME_FILTER", "DD_14_7_GTE_23", 101, metadata["selection_run_id"], ALGORITHM_VERSION),
        (1, "EQUITY_REGIME_FILTER", "W28_DOWN", 101, metadata["selection_run_id"], ALGORITHM_VERSION),
    ]


def test_filter_drop_from_preview_cache_miss_publishes_without_warming_cache():
    connection = _database()
    request = _request()
    result = _result(connection, _assessment(), filter_enabled=True, cache=False)
    metadata = new_run_metadata(connection, request)
    payload = result.attrs["equity_regime_evidence"]["1"]["equity_regime_json"]

    persist_selection_snapshot(connection, request, SelectionConfig(), result, metadata, b"preview-miss")

    assert connection.execute(
        "select count(*) from equity_quality_metrics where algo_version = ?",
        [ALGORITHM_VERSION],
    ).fetchone() == (0,)
    assert connection.execute(
        "select equity_regime_json from selection_results where selection_run_id = ?",
        [metadata["selection_run_id"]],
    ).fetchone() == (payload,)
    assert connection.execute(
        "select count(*) from strategy_rejection_sources where strategy_id = 1"
    ).fetchone() == (2,)


def test_stale_assessment_digest_aborts_every_publication_table():
    connection = _database()
    request = _request()
    result = _result(connection, _assessment(), filter_enabled=True)
    result.attrs["equity_regime_evidence"]["1"]["facts_sha256"] = "0" * 64
    metadata = new_run_metadata(connection, request)
    before = {
        table: connection.execute(f"select count(*) from {table}").fetchone()[0]
        for table in ("selection_runs", "selection_results", "strategy_rejection_sources", "strategy_tags")
    }

    with pytest.raises(SelectionReviewError, match="SELECTION_REVIEW_STALE_RESULTS"):
        persist_selection_snapshot(connection, request, SelectionConfig(), result, metadata, b"review")

    assert {
        table: connection.execute(f"select count(*) from {table}").fetchone()[0]
        for table in before
    } == before


def test_rank_only_drop_and_filter_not_evaluated_do_not_create_rejection_sources():
    connection = _database()
    request = _request(filter_enabled=False, rank_enabled=True)
    assessment = _assessment()
    result = _result(connection, assessment, filter_enabled=False)
    facts = {"result_id": 101}
    result.attrs["equity_quality_facts"] = {"1": {
        "result_id": 101,
        "source_revision": result.attrs["equity_regime_evidence"]["1"]["source_revision"],
        "facts_sha256": sha256(json.dumps(facts, sort_keys=True, separators=(",", ":")).encode()).hexdigest(),
        "facts": facts,
    }}
    metadata = new_run_metadata(connection, request)
    persist_selection_snapshot(connection, request, SelectionConfig(), result, metadata, b"rank-only")
    assert connection.execute("select count(*) from strategy_rejection_sources").fetchone() == (0,)

    connection = _database()
    request = _request()
    result = _result(connection, _assessment(evaluated=False), filter_enabled=True)
    metadata = new_run_metadata(connection, request)
    persist_selection_snapshot(connection, request, SelectionConfig(), result, metadata, b"not-evaluated")
    assert json.loads(connection.execute("select equity_regime_json from selection_results").fetchone()[0])["state"] == "NOT_EVALUATED"
    assert connection.execute("select count(*) from strategy_rejection_sources").fetchone() == (0,)


def test_both_equity_toggles_off_keeps_assessment_snapshot_null():
    connection = _database()
    request = _request(filter_enabled=False, rank_enabled=False)
    result = _result(connection, _assessment(), filter_enabled=False)
    metadata = new_run_metadata(connection, request)

    persist_selection_snapshot(connection, request, SelectionConfig(), result, metadata, b"both-off")

    assert connection.execute(
        "select equity_regime_json from selection_results where selection_run_id = ?",
        [metadata["selection_run_id"]],
    ).fetchone() == (None,)
    published_request = json.loads(connection.execute(
        "select request_json from selection_runs where selection_run_id = ?",
        [metadata["selection_run_id"]],
    ).fetchone()[0])
    assert "equity_regime_snapshot" not in published_request
    assert connection.execute("select count(*) from strategy_rejection_sources").fetchone() == (0,)


def test_duplicate_hard_reasons_across_group_runs_keep_first_evidence():
    connection = _database()
    request = _request()
    result = _result(connection, _assessment(), filter_enabled=True)
    first = new_run_metadata(connection, request)
    second = new_run_metadata(connection, request)

    persist_selection_snapshots(connection, tuple(
        {"request": request, "config": SelectionConfig(), "result": result, "metadata": metadata}
        for metadata in (first, second)
    ), workbook_bytes=b"combined")

    rows = connection.execute(
        "select reason_code, first_selection_run_id from strategy_rejection_sources order by reason_code"
    ).fetchall()
    assert rows == [
        ("DD_14_7_GTE_23", first["selection_run_id"]),
        ("W28_DOWN", first["selection_run_id"]),
    ]


def test_conflicting_assessment_for_same_strategy_aborts_combined_publication():
    connection = _database()
    request = _request()
    result = _result(connection, _assessment(), filter_enabled=True)
    conflicting = result.copy()
    conflicting.attrs["equity_regime_evidence"] = {
        "1": {
            **result.attrs["equity_regime_evidence"]["1"],
            "source_revision": "f" * 64,
        },
    }
    first = new_run_metadata(connection, request)
    second = new_run_metadata(connection, request)

    with pytest.raises(SelectionReviewError, match="SELECTION_REVIEW_STALE_RESULTS"):
        persist_selection_snapshots(connection, (
            {"request": request, "config": SelectionConfig(), "result": result, "metadata": first},
            {"request": request, "config": SelectionConfig(), "result": conflicting, "metadata": second},
        ), workbook_bytes=b"conflicting")

    assert connection.execute("select count(*) from selection_runs").fetchone() == (0,)
    assert connection.execute("select count(*) from selection_results").fetchone() == (0,)
    assert connection.execute("select count(*) from strategy_rejection_sources").fetchone() == (0,)


def test_manual_review_clearing_manual_tag_keeps_equity_rejection_effective(tmp_path):
    connection = _database()
    request = _request()
    result = _result(connection, _assessment(), filter_enabled=True)
    metadata = new_run_metadata(connection, request)
    path = tmp_path / "equity-review.xlsx"
    review = {1: {
        "user_status": "REJECTED", "user_rank": None,
        "user_analog_of_strategy_id": None, "comment": "manual",
    }}
    workbook_path = write_selection_workbook(result, path, request, metadata, review)
    workbook_bytes = workbook_path.read_bytes()
    persist_selection_snapshot(connection, request, SelectionConfig(), result, metadata, workbook_bytes)
    connection.execute(
        "insert into strategy_tags values (1, 'REJECTED', 'SELECTION_REVIEW', 'old-review', ?)",
        [NOW],
    )

    workbook = load_workbook(BytesIO(workbook_bytes))
    sheet = workbook["All candidates"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["User Status"], "FINALIST")
    sheet.cell(2, headers["User Rank"], 1)
    edited = BytesIO()
    workbook.save(edited)
    import_selection_review(connection, edited.getvalue())

    assert connection.execute("select count(*) from strategy_tags where tag = 'REJECTED'").fetchone() == (0,)
    assert connection.execute("select count(*) from strategy_rejection_sources where strategy_id = 1").fetchone()[0] == 2
    assert effective_selection_decisions(connection)[1][0] == "REJECTED"
    assert apply_prior_rejected(connection, pd.DataFrame([{"strategy_id": 1}])).loc[0, "prior_rejected"]
