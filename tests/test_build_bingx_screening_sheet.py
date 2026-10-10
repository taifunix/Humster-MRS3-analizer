from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from scripts.build_bingx_screening_sheet import SOURCE_SHEET, TARGET_SHEET, build_bingx_screening_sheet, main

HEADER = ["Пара", "Листинг", "Объём 7д, USDT", "Сторона", "Стадия", "Вердикт", "Точек", "BIG", "READY", "Plateu",
          "Финалист", "Стадия анализа", "Коммент", None, "date start", "date end"]
PANW_LAUNCH = int(datetime(2026, 6, 11, 10, 0, tzinfo=timezone.utc).timestamp() * 1000)
CONTRACTS = [
    {"symbol": "NCSKPANW2USD-USDT", "displayName": "PANW-USDT", "launchTime": PANW_LAUNCH},
    {"symbol": "ON-USDT", "displayName": "ON-USDT", "launchTime": 1},
    {"symbol": "NCSKON2USD-USDT", "displayName": "ON-USDT",
     "launchTime": int(datetime(2026, 7, 1, 23, 30, tzinfo=timezone.utc).timestamp() * 1000)},
    {"symbol": "NCSKAMD2USD-USDT", "displayName": "AMDUS-USDT",
     "launchTime": int(datetime(2026, 8, 2, tzinfo=timezone.utc).timestamp() * 1000)},
]


def _registry(path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "Пары"
    sheet = workbook.create_sheet(SOURCE_SHEET)
    sheet.append(HEADER)
    sheet.append(["PANWUSDT", datetime(2026, 7, 22), 1000, "LONG", "СКРИН", "GO", 5, 2, "ЕСТЬ", None, "+", "MRS3-READY",
                  "note", None, datetime(2026, 7, 1), datetime(2026, 9, 1)])
    sheet.append(["XAUUSDT", datetime(2026, 3, 9), 2000, "LONG", None, "STOP"])
    sheet.append(["ONUSDT", datetime(2026, 6, 1), 3000])
    sheet.append(["AMDSTOCKUSDT", datetime(2026, 6, 2), 4000])
    sheet["B2"].number_format = "DD.MM.YYYY"
    sheet["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    sheet.column_dimensions["A"].width = 21
    sheet.auto_filter.ref = "A1:P5"
    sheet.freeze_panes = "A2"
    sheet.sheet_view.pane.topLeftCell = "A165"  # bottom pane scrolled down, as in the real registry
    validation = DataValidation(type="list", formula1='"СКРИН,АНАЛИЗ"')
    validation.add("E2:E5")
    sheet.add_data_validation(validation)
    sheet.conditional_formatting.add("F2:F5", CellIsRule(operator="equal", formula=['"GO"'], fill=PatternFill("solid", fgColor="00FF00")))
    workbook.save(path)


def test_builds_bingx_sheet_with_same_form_and_bingx_tickers(tmp_path: Path):
    path = tmp_path / "registry.xlsx"
    _registry(path)
    report = build_bingx_screening_sheet(path, CONTRACTS)

    assert (report.added, report.existing, report.missing) == (3, 0, ["XAUUSDT"])
    assert report.backup is not None and report.backup.is_file()
    workbook = load_workbook(path)
    assert workbook.sheetnames.index(TARGET_SHEET) == workbook.sheetnames.index(SOURCE_SHEET) + 1
    sheet = workbook[TARGET_SHEET]
    assert [cell.value for cell in sheet[1]][: len(HEADER)] == HEADER
    rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2, max_col=len(HEADER))]
    filled = [row for row in rows if row[0] is not None]
    assert [row[:2] for row in filled] == [
        ["PANWUSDT", datetime(2026, 6, 11)], ["ONUSDT", datetime(2026, 7, 1)], ["AMDUSDT", datetime(2026, 8, 2)],
    ]
    assert all(value is None for row in rows for value in row[2:])  # Bybit results are not copied
    assert sheet["B2"].number_format == "DD.MM.YYYY"
    assert sheet["A1"].fill.fgColor.rgb.endswith("FFFF00")
    assert sheet.column_dimensions["A"].width == 21
    assert sheet.auto_filter.ref == "A1:P5" and sheet.freeze_panes == "A2"
    assert [str(item.sqref) for item in sheet.data_validations.dataValidation] == ["E2:E5"]
    assert [str(item.sqref) for item in sheet.conditional_formatting] == ["F2:F5"]
    assert workbook[SOURCE_SHEET]["A2"].value == "PANWUSDT" and workbook[SOURCE_SHEET]["F2"].value == "GO"


def test_rerun_keeps_manual_values_and_only_appends_new_pairs(tmp_path: Path):
    path = tmp_path / "registry.xlsx"
    _registry(path)
    build_bingx_screening_sheet(path, CONTRACTS[:1])
    workbook = load_workbook(path)
    workbook[TARGET_SHEET]["F2"] = "GO"
    workbook.save(path)

    report = build_bingx_screening_sheet(path, CONTRACTS)
    assert (report.added, report.existing) == (2, 1)
    sheet = load_workbook(path)[TARGET_SHEET]
    assert sheet["F2"].value == "GO"
    assert [sheet.cell(row=index, column=1).value for index in range(2, 5)] == ["PANWUSDT", "ONUSDT", "AMDUSDT"]

    again = build_bingx_screening_sheet(path, CONTRACTS)
    assert (again.added, again.existing, again.backup) == (0, 3, None)


def test_main_reports_summary(tmp_path: Path, capsys):
    path = tmp_path / "registry.xlsx"
    _registry(path)
    assert main(["--registry", str(path)], contracts=CONTRACTS) == 0
    output = capsys.readouterr().out
    assert "добавлено пар: 3" in output and "XAUUSDT" in output


def test_first_run_without_bingx_pairs_creates_empty_sheet_and_saves(tmp_path: Path):
    path = tmp_path / "registry.xlsx"
    _registry(path)
    report = build_bingx_screening_sheet(path, [])
    assert report.added == 0 and report.backup is not None
    sheet = load_workbook(path)[TARGET_SHEET]
    assert all(cell.value is None for row in sheet.iter_rows(min_row=2) for cell in row)


def test_unknown_launch_time_leaves_listing_empty_and_rows_below_operator_notes(tmp_path: Path):
    path = tmp_path / "registry.xlsx"
    _registry(path)
    build_bingx_screening_sheet(path, CONTRACTS[:1])
    workbook = load_workbook(path)
    workbook[TARGET_SHEET]["M3"] = "operator note without pair"
    workbook.save(path)
    no_launch = [{"symbol": "NCSKON2USD-USDT", "displayName": "ON-USDT"}]
    build_bingx_screening_sheet(path, CONTRACTS[:1] + no_launch)
    sheet = load_workbook(path)[TARGET_SHEET]
    assert sheet["M3"].value == "operator note without pair" and sheet["A3"].value is None
    assert sheet["A4"].value == "ONUSDT" and sheet["B4"].value is None
