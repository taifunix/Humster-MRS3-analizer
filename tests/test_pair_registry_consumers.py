"""Registry stage 4: the consumers read the pair registry.

Contract: docs/specs/2026-10-18-pair-registry.md, «Stage 4 contract»; decision ADR-0079.
"""

from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

import pytest

from mrs3.pair_registry.liquidity import read_liquidity, refresh_liquidity, registry_capacity
from mrs3.pair_registry.store import RegistrySettings
from mrs3.portfolio.minute_capacity import MinuteCapacityError
from tests.test_pair_registry_liquidity import _registry_with_pair, _two_weeks

ANCHOR = datetime(2026, 10, 19, 9, 30, tzinfo=timezone.utc)          # Monday: the window ends on Friday 16


def _ready_registry(tmp_path: Path):
    root = tmp_path / "data"
    _two_weeks(root)
    con = _registry_with_pair(tmp_path)
    refresh_liquidity(con, "bybit", date(2026, 10, 19), RegistrySettings(liquidity_data_root=root))
    return con


# -- 1. optimizer: LOT_MODEL_V2 capacity from the registry ---------------------------------------------------

def test_registry_capacity_is_lot_model_v2_from_the_stored_window(tmp_path: Path) -> None:
    con = _ready_registry(tmp_path)
    row = read_liquidity(con, "bybit", "XOMUSDT")
    capacity = registry_capacity(con, "bybit", "xomusdt", anchor_created_at=ANCHOR)
    assert (capacity.status, capacity.symbol, capacity.selected_basis) == ("READY", "XOMUSDT", "LOT_MODEL_V2")
    assert capacity.window_start == date(2026, 10, 8) and capacity.window_end == date(2026, 10, 16)
    features = capacity.lot_features
    assert features.v25 == Decimal(row["v25"]) and features.a15 == Decimal(row["a15"])
    assert features.active_15m_bins == row["active_bins"]
    assert [day.isoformat() for day in features.window_dates] == row["window_days"]
    assert features.file_hashes == tuple((day, digest) for day, digest in row["window_files"])
    assert features.publication_lag_hours is None and features.anchor_created_at_utc == ANCHOR.isoformat()
    assert capacity.round_down_usdt == Decimal(row["step"]) and len(capacity.content_digest) == 64
    # the same window on Saturday and Sunday: the last weekday before both days is Friday 16
    saturday = registry_capacity(con, "bybit", "XOMUSDT", anchor_created_at=datetime(2026, 10, 17, tzinfo=timezone.utc))
    assert saturday.content_digest == capacity.content_digest


def test_registry_capacity_fails_closed(tmp_path: Path) -> None:
    con = _ready_registry(tmp_path)
    later = datetime(2026, 10, 21, 9, tzinfo=timezone.utc)                # wants the window ending on Tuesday 20
    with pytest.raises(MinuteCapacityError, match="LIQUIDITY_MODEL_WINDOW_UNAVAILABLE"):
        registry_capacity(con, "bybit", "XOMUSDT", anchor_created_at=later)
    with pytest.raises(MinuteCapacityError, match="LIQUIDITY_MODEL_WINDOW_UNAVAILABLE"):
        registry_capacity(con, "bybit", "AAPLUSDT", anchor_created_at=ANCHOR)          # no row
    with pytest.raises(MinuteCapacityError, match="LIQUIDITY_MODEL_ANCHOR_UNAVAILABLE"):
        registry_capacity(con, "bybit", "XOMUSDT", anchor_created_at=datetime(2026, 10, 19))      # naive
    con.execute("update liquidity set inputs_sha256 = null")                # a stage 2 draft row
    with pytest.raises(MinuteCapacityError, match="LIQUIDITY_MODEL_WINDOW_UNAVAILABLE"):
        registry_capacity(con, "bybit", "XOMUSDT", anchor_created_at=ANCHOR)


def lot_model_sizing(capacity):
    """One XOMUSDT finalist sized in LOT_MODEL mode (position_sizing.enrich_finalist_rows)."""
    from mrs3.portfolio.liquidity import ReferenceReader
    from mrs3.portfolio.position_sizing import enrich_finalist_rows

    reference = ReferenceReader.from_records(
        instruments=[{"symbol": "XOMUSDT", "status": "Trading", "contract_type": "LinearPerpetual", "tick_size": "0.01",
                      "qty_step": "0.1", "min_qty": "0.1", "max_qty": "100000", "leverage_step": "0.1",
                      "max_leverage": "50", "min_notional": "5"}],
        risk_tiers=[{"symbol": "XOMUSDT", "risk_limit_value": "1000000", "max_leverage": "25"}],
        captured_at_ms=1_000,
    )
    row = {"user_status": "FINALIST", "symbol": "XOMUSDT", "side": "LONG", "strategy_id": 7, "result_id": 70,
           "strategy_orders": ({"order_id": 1, "lot_x": Decimal("1"), "shift_bp": 100},
                               {"order_id": 2, "lot_x": Decimal("1"), "shift_bp": 300})}
    return enrich_finalist_rows([row], {"XOMUSDT": capacity}, reference, {"XOMUSDT": Decimal("100")}, now_ms=1_000,
                                maximum_age_hours=2, lot_model_settings={"lot_model_base_coefficient": Decimal("9"),
                                                                         "lot_model_max_shift_bonus": Decimal("1.1")})


def test_position_sizing_accepts_a_v2_window_and_records_the_model(tmp_path: Path) -> None:
    con = _ready_registry(tmp_path)
    capacity = registry_capacity(con, "bybit", "XOMUSDT", anchor_created_at=ANCHOR)
    result = lot_model_sizing(capacity)
    assert result.status == "PASS", result
    row = result.rows[0]
    assert row["liquidity_lot_model"] == "LOT_MODEL_V2"
    assert row["liquidity_v25_usdt"] == capacity.lot_features.v25
    assert row["liquidity_window_dates"] == capacity.lot_features.window_dates


def test_position_sizing_refuses_an_inconsistent_v2_window(tmp_path: Path) -> None:
    from dataclasses import replace

    con = _ready_registry(tmp_path)
    capacity = registry_capacity(con, "bybit", "XOMUSDT", anchor_created_at=ANCHOR)
    features = capacity.lot_features
    weekend = replace(features, window_dates=(*features.window_dates[:-1], date(2026, 10, 18)))
    stale = replace(features, anchor_created_at_utc=datetime(2026, 10, 22, tzinfo=timezone.utc).isoformat())
    short_hashes = replace(features, file_hashes=features.file_hashes[:-1])
    too_many_bins = replace(features, active_15m_bins=len(features.window_dates) * 96 + 1)
    for bad in (weekend, stale, short_hashes, too_many_bins):
        result = lot_model_sizing(replace(capacity, lot_features=bad))
        assert result.status == "FAIL" and result.exclusions[0].reason == "LIQUIDITY_MODEL_WINDOW_UNAVAILABLE", bad


def test_the_panel_optimizer_reads_the_registry_when_it_exists(monkeypatch, tmp_path: Path) -> None:
    import json

    from mrs3.panel_portfolio import PortfolioPanelService

    con = _ready_registry(tmp_path)
    con.close()
    (tmp_path / "config.local.json").write_text(json.dumps({"pair_registry": {"root": "registry"}}), encoding="utf-8")
    observed = {}

    def fake_adapter(*_args, **kwargs):
        observed.update(kwargs)
        from types import SimpleNamespace
        return SimpleNamespace(variants=(), blockers=(), excluded=(), warnings=())

    monkeypatch.setattr("mrs3.portfolio.adapter.run_portfolio_adapter", fake_adapter)
    PortfolioPanelService(tmp_path)._package_variant_generator((), {"config_document": {}}, ())
    capacity = observed["liquidity_reader"]("XOMUSDT", ANCHOR, "10")
    assert capacity.selected_basis == "LOT_MODEL_V2" and capacity.round_down_usdt == Decimal("10")
    (tmp_path / "registry" / "registry.duckdb").unlink()                  # no registry: LOT_MODEL_V1 as before
    observed.clear()
    PortfolioPanelService(tmp_path)._package_variant_generator((), {"config_document": {}}, ())
    assert "liquidity_reader" not in observed


# -- 2. base-lot export -----------------------------------------------------------------------------------------

def test_base_lot_export_takes_each_cell_from_the_registry_window_of_its_day(tmp_path: Path) -> None:
    import json

    from openpyxl import load_workbook

    from mrs3.pair_registry.liquidity import liquidity_on
    from scripts.calculate_bybit_base_lots import BaseLotExportError, registry_lot_source, run_export

    con = _ready_registry(tmp_path)
    row = read_liquidity(con, "bybit", "XOMUSDT")
    assert liquidity_on(con, "bybit", "XOMUSDT", date(2026, 10, 16))["base_lot"] == row["base_lot"]
    assert liquidity_on(con, "bybit", "XOMUSDT", date(2026, 10, 15)) is None
    con.close()
    config = tmp_path / "config.local.json"
    config.write_text(json.dumps({"pair_registry": {"root": "registry"}}), encoding="utf-8")
    from tests.test_calculate_bybit_base_lots import _workbook

    _workbook(tmp_path / "lots.xlsx")                     # the legacy layout: BTCUSDT and ETHUSDT rows
    book = load_workbook(tmp_path / "lots.xlsx")
    book["Actual"]["A2"] = "XOMUSDT"
    book.save(tmp_path / "lots.xlsx")
    source = registry_lot_source(config, k="9", round_down_usdt="10")
    result = run_export(tmp_path / "lots.xlsx", tmp_path / "no-minute-files", anchor_created_at=datetime(
        2026, 10, 19, 8, tzinfo=timezone.utc), lot_source=source)
    written = {item.actualization_date: item.base_lot_usdt for item in result.rows if item.symbol == "XOMUSDT"}
    assert set(written) == {date(2026, 10, 17), date(2026, 10, 18), date(2026, 10, 19)}       # windows ending Fri 16
    assert all(value == Decimal(row["base_lot"]) for value in written.values())
    assert {error.actualization_date for error in result.errors if error.symbol == "XOMUSDT"} == {
        date(2026, 10, 10) + __import__("datetime").timedelta(days=offset) for offset in range(7)}
    assert "registry" in result.errors[0].message
    saved = load_workbook(tmp_path / "lots.xlsx")["Actual"]
    assert saved.cell(2, 13).value == int(Decimal(row["base_lot"]))
    with pytest.raises(BaseLotExportError):
        registry_lot_source(tmp_path / "missing.json", k="9", round_down_usdt="10")
    with pytest.raises(BaseLotExportError, match="K or step"):          # before any workbook cell is touched
        registry_lot_source(config, k="8", round_down_usdt="10")
    assert load_workbook(tmp_path / "lots.xlsx").defined_names["MRS3_BASE_LOT_MODEL"].attr_text == '"LOT_MODEL_V2"'


# -- 3. pairs and listing days of the history -------------------------------------------------------------------

def test_history_pairs_come_from_the_registry_by_liquidity(tmp_path: Path) -> None:
    import json

    from mrs3.pair_registry.history import history_pairs
    from mrs3.pair_registry.refresh import refresh_exchange

    config = tmp_path / "config.local.json"
    config.write_text(json.dumps({"pair_registry": {"root": "registry"}}), encoding="utf-8")
    assert history_pairs(config) is None                                    # no registry: the workbook is used
    con = _ready_registry(tmp_path)                                          # XOMUSDT with a READY window
    listing = int(datetime(2026, 6, 1, tzinfo=timezone.utc).timestamp() * 1000)
    refresh_exchange(con, "bybit", [
        {"symbol": "XOMUSDT", "status": "TRADING", "api_symbol": "XOMUSDT", "listing_ms": 1_700_000_000_000},
        {"symbol": "AAPLUSDT", "status": "TRADING", "api_symbol": "AAPLUSDT", "listing_ms": listing},
        {"symbol": "ZZZUSDT", "status": "TRADING", "api_symbol": "ZZZUSDT", "listing_ms": listing},
        {"symbol": "OLDUSDT", "status": "DELISTED", "api_symbol": "OLDUSDT", "listing_ms": listing},
        {"symbol": "NODATEUSDT", "status": "TRADING", "api_symbol": "NODATEUSDT", "listing_ms": None},
    ], [], now_ms=1_760_000_000_001)
    con.close()
    pairs = history_pairs(config)
    assert [symbol for symbol, _day in pairs] == ["XOMUSDT", "AAPLUSDT", "ZZZUSDT"]      # liquid first, then by name
    assert dict(pairs)["AAPLUSDT"] == date(2026, 6, 1)


def test_simulate_history_takes_the_registry_order(tmp_path: Path, monkeypatch) -> None:
    import json
    import sys

    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
    import simulate_history

    seen = {}
    monkeypatch.setattr(simulate_history, "history_pairs", lambda config, **_: [("XOMUSDT", date(2026, 6, 1))])
    monkeypatch.setattr(simulate_history, "pair_order", lambda path: pytest.fail("the workbook was read"))

    def stop(store, **kwargs):
        seen["order"] = list(kwargs["pairs"])
        return {"failed": 0}

    from types import SimpleNamespace

    monkeypatch.setattr(simulate_history, "run_history", stop)
    monkeypatch.setattr(simulate_history, "instrument", lambda symbol, cache: SimpleNamespace(tick_size=0.01, qty_step=0.01))
    config = tmp_path / "config.local.json"
    config.write_text(json.dumps({"tester_runner": {"bot_root": str(tmp_path / "bot")},
                                  "trade_store": {"root": str(tmp_path / "store")}}), encoding="utf-8")
    assert simulate_history.main(["--config", str(config), "--no-download", "--end-day", "2026-10-18"]) == 0
    assert seen["order"] == [("XOMUSDT", date(2026, 6, 1))]


# -- 4. the listing workbook for the loader-based consumers -------------------------------------------------------

def test_the_listing_workbook_is_read_by_the_loader_on_both_sheets(tmp_path: Path, capsys) -> None:
    import json

    import pandas as pd

    from mrs3.loader import load_listing_dates
    from mrs3.pair_registry import cli
    from mrs3.pair_registry.refresh import refresh_exchange
    from mrs3.pair_registry.workbook import write_listing_workbook

    con = _ready_registry(tmp_path)
    listing = int(datetime(2026, 6, 1, 8, 17, 11, tzinfo=timezone.utc).timestamp() * 1000)
    refresh_exchange(con, "bybit", [
        {"symbol": "XOMUSDT", "status": "TRADING", "api_symbol": "XOMUSDT", "listing_ms": 1_700_000_000_000},
        {"symbol": "OLDUSDT", "status": "DELISTED", "api_symbol": "OLDUSDT", "listing_ms": listing},
        {"symbol": "NODATEUSDT", "status": "TRADING", "api_symbol": "NODATEUSDT", "listing_ms": None},
    ], [], now_ms=1_760_000_000_001)
    refresh_exchange(con, "bingx", [{"symbol": "NVDA-USDT", "status": "TRADING", "api_symbol": "NCSKNVDA2USD-USDT",
                                     "listing_ms": listing}], [], now_ms=1_760_000_000_001)
    out = tmp_path / "out" / "pairs.xlsx"
    assert write_listing_workbook(con, out) == {"bybit": 2, "bingx": 1}
    con.close()
    bybit = load_listing_dates(out)
    assert bybit == {"OLDUSDT": pd.Timestamp("2026-06-01", tz="UTC"),                 # the UTC day at midnight,
                     "XOMUSDT": pd.Timestamp("2023-11-14", tz="UTC")}                 # as in the workbook today
    assert load_listing_dates(out, "Пары_BingX") == {"NVDA-USDT": pd.Timestamp("2026-06-01", tz="UTC")}
    assert not list(out.parent.glob("*.tmp*"))
    config = tmp_path / "config.local.json"
    config.write_text(json.dumps({"pair_registry": {"root": "registry"}}), encoding="utf-8")
    assert cli.main(["workbook", "--config", str(config)]) == 0
    assert json.loads(capsys.readouterr().out.splitlines()[-1])["workbook"] == {"bybit": 2, "bingx": 1}
    assert load_listing_dates(tmp_path / "registry" / "pairs.xlsx") == bybit


# -- review of registry stage 4 ----------------------------------------------------------------------------------

def test_registry_capacity_uses_the_anchor_window_after_a_later_sync_and_skips_non_trading(tmp_path: Path) -> None:
    con = _ready_registry(tmp_path)
    root = tmp_path / "data"
    refresh_liquidity(con, "bybit", date(2026, 10, 21), RegistrySettings(liquidity_data_root=root))   # files end 18
    assert read_liquidity(con, "bybit", "XOMUSDT") is None              # the newest window is UNAVAILABLE
    retried = registry_capacity(con, "bybit", "XOMUSDT", anchor_created_at=ANCHOR)        # a frozen Campaign retry
    assert retried.window_end == date(2026, 10, 16)
    con.execute("update pairs set status = 'DELISTED' where symbol = 'XOMUSDT'")
    with pytest.raises(MinuteCapacityError, match="LIQUIDITY_MODEL_WINDOW_UNAVAILABLE"):
        registry_capacity(con, "bybit", "XOMUSDT", anchor_created_at=ANCHOR)


def _held_registry(tmp_path: Path):
    """A writable connection that makes any read-only open of the same file fail in this process."""
    import duckdb

    return duckdb.connect(str(tmp_path / "registry" / "registry.duckdb"))


def test_busy_registry_paths(tmp_path: Path) -> None:
    import json

    from mrs3.pair_registry.history import history_pairs, listing_days
    from mrs3.panel_portfolio import _registry_liquidity_reader
    from scripts.calculate_bybit_base_lots import BaseLotExportError, registry_lot_source

    con = _ready_registry(tmp_path)
    con.close()
    config = tmp_path / "config.local.json"
    config.write_text(json.dumps({"pair_registry": {"root": "registry"}}), encoding="utf-8")
    reader = _registry_liquidity_reader(tmp_path)
    source = registry_lot_source(config, k="9", round_down_usdt="10")
    held = _held_registry(tmp_path)
    try:
        warnings = []
        assert history_pairs(config, on_busy=warnings.append) is None and "unavailable" in warnings[0]
        assert listing_days(config, on_busy=warnings.append) == {} and len(warnings) == 2
        with pytest.raises(MinuteCapacityError, match="PAIR_REGISTRY_BUSY"):
            reader("XOMUSDT", ANCHOR, "10")
        with pytest.raises(BaseLotExportError, match="busy"):          # an error cell, not an aborted export
            source("XOMUSDT", date(2026, 10, 19))
    finally:
        held.close()


def test_base_lot_export_never_mixes_models(tmp_path: Path) -> None:
    import json

    from openpyxl import load_workbook

    from scripts.calculate_bybit_base_lots import BaseLotExportError, recorded_model, registry_lot_source, run_export
    from tests.test_calculate_bybit_base_lots import _workbook

    con = _ready_registry(tmp_path)
    con.close()
    config = tmp_path / "config.local.json"
    config.write_text(json.dumps({"pair_registry": {"root": "registry"}}), encoding="utf-8")
    _workbook(tmp_path / "lots.xlsx")
    book = load_workbook(tmp_path / "lots.xlsx")
    book["Actual"]["A2"] = "XOMUSDT"
    book.save(tmp_path / "lots.xlsx")
    source = registry_lot_source(config, k="9", round_down_usdt="10")
    run_export(tmp_path / "lots.xlsx", tmp_path / "none", anchor_created_at=datetime(2026, 10, 17, 8, tzinfo=timezone.utc),
               lot_source=source)
    assert recorded_model(load_workbook(tmp_path / "lots.xlsx")) == "LOT_MODEL_V2"
    next_day = datetime(2026, 10, 18, 8, tzinfo=timezone.utc)
    with pytest.raises(BaseLotExportError, match="LOT_MODEL_V2"):       # a minute-file run must not add V1 values
        run_export(tmp_path / "lots.xlsx", tmp_path / "none", anchor_created_at=next_day)
    book = load_workbook(tmp_path / "lots.xlsx")
    del book.defined_names["MRS3_BASE_LOT_MODEL"]                       # an existing V1 workbook has no marker
    book.save(tmp_path / "lots.xlsx")
    before = (tmp_path / "lots.xlsx").read_bytes()
    with pytest.raises(BaseLotExportError, match="--rebuild-history"):
        run_export(tmp_path / "lots.xlsx", tmp_path / "none", anchor_created_at=next_day, lot_source=source)
    assert (tmp_path / "lots.xlsx").read_bytes() == before


def test_base_lot_cli_defaults_to_the_registry(tmp_path: Path, capsys) -> None:
    import json

    from scripts.calculate_bybit_base_lots import main

    config = tmp_path / "config.local.json"
    config.write_text(json.dumps({"pair_registry": {"root": "registry"}}), encoding="utf-8")
    assert main(["--input", str(tmp_path / "lots.xlsx"), "--config", str(config)]) == 1
    assert "pair registry not found" in capsys.readouterr().out


def test_generate_period_csv_takes_the_registry_order_and_any_status_for_explicit_pairs(tmp_path: Path, monkeypatch) -> None:
    import json
    import sys

    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
    import generate_period_csv

    class Stop(Exception):
        pass

    monkeypatch.setattr(generate_period_csv, "history_pairs", lambda config, **_: [("XOMUSDT", date(2026, 6, 1))])
    monkeypatch.setattr(generate_period_csv, "listing_days", lambda config, **_: {"OLDUSDT": date(2026, 5, 1)})
    monkeypatch.setattr(generate_period_csv, "pair_order", lambda path: pytest.fail("the workbook was read"))
    seen = []

    def instrument(symbol, cache):
        seen.append(symbol)
        raise Stop

    monkeypatch.setattr(generate_period_csv, "instrument", instrument)
    config = tmp_path / "config.local.json"
    config.write_text(json.dumps({"tester_runner": {"bot_root": str(tmp_path / "bot")}}), encoding="utf-8")
    with pytest.raises(Stop):
        generate_period_csv.main(["--out", str(tmp_path / "out"), "--config", str(config), "--pairs", "OLDUSDT"])
    assert seen == ["OLDUSDT"]


def test_registry_cli_workbook_and_liquidity_days(tmp_path: Path, capsys) -> None:
    import json

    import duckdb

    from mrs3.pair_registry import cli

    config = tmp_path / "config.local.json"
    config.write_text(json.dumps({"pair_registry": {"root": "registry"},
                                  "tester_runner": {"bot_root": str(tmp_path / "bot")}}), encoding="utf-8")
    assert cli.main(["workbook", "--config", str(config)]) == 1        # no registry yet: nothing is created
    assert "does not exist" in capsys.readouterr().out and not (tmp_path / "registry").exists()
    con = _ready_registry(tmp_path)
    con.close()
    root = tmp_path / "bot" / "tester" / "data" / "bybit"
    root.parent.mkdir(parents=True)
    (tmp_path / "data").rename(root)
    assert cli.main(["liquidity", "--config", str(config), "--today", "2026-10-20", "--days", "3"]) == 0
    runs = json.loads(capsys.readouterr().out.splitlines()[-1])["bybit"]["runs"]
    assert len(runs) == 3
    with duckdb.connect(str(tmp_path / "registry" / "registry.duckdb"), read_only=True) as check:
        days = [row[0].isoformat() for row in check.execute("select as_of_day from liquidity order by as_of_day").fetchall()]
    assert days == ["2026-10-16", "2026-10-19"]           # run days 18, 19, 20 -> windows ending Fri 16, Fri 16, Mon 19
    assert cli.main(["liquidity", "--config", str(config), "--days", "40"]) == 1


def test_a_locked_workbook_fails_only_its_sync_step(tmp_path: Path, monkeypatch, capsys) -> None:
    import json

    from mrs3.pair_registry import cli
    from tests.test_pair_registry import BINGX, BYBIT, RISK

    config = tmp_path / "config.local.json"
    config.write_text(json.dumps({"pair_registry": {"root": str(tmp_path / "registry")}}), encoding="utf-8")
    monkeypatch.setattr(cli, "fetch_bybit", lambda: (BYBIT, RISK))
    monkeypatch.setattr(cli, "fetch_bingx", lambda: BINGX)

    def locked(con, out):
        raise PermissionError(13, "denied", str(out))

    monkeypatch.setattr(cli, "write_listing_workbook", locked)
    cli.main(["sync", "--config", str(config), "--today", "2026-10-19", "--skip-track"])
    events = [json.loads(line) for line in capsys.readouterr().out.splitlines()]
    error = next(event for event in events if event.get("step") == "workbook")
    assert "listing workbook not written" in error["error"] and str(tmp_path) not in error["error"]
