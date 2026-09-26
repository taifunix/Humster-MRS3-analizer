from __future__ import annotations

from datetime import datetime, timedelta, timezone
from decimal import Decimal
from hashlib import sha256
from http.client import HTTPConnection
from io import BytesIO
import json
from pathlib import Path
import threading

import duckdb
from openpyxl import load_workbook
import pytest

from mrs3.panel import PanelController, create_panel_server
from mrs3.panel_performance_v2 import (
    PerformanceV2ApiError,
    PerformanceV2ExportSelection,
    export_performance_v2,
    parse_performance_v2_export_query,
)
from mrs3.performance_v2_store import initialize_performance_v2
from mrs3.performance_v2_equity_cache import current_equity_source_metadata, upsert_equity_quality_facts_checked
from mrs3.performance_v2_equity_quality import EquitySample, calculate_equity_quality_facts
import mrs3.performance_v2_selection as selection_module


UTC = timezone.utc


def _export_database(path: Path) -> tuple[Path, int]:
    now = datetime(2026, 1, 1, tzinfo=UTC)
    with duckdb.connect(str(path)) as connection:
        initialize_performance_v2(connection)
        strategy_id = connection.execute(
            """insert into strategies (
                   strategy_name, symbol, side, timeframe, close_ma_len, order_count,
                   analysis_run_id, candidate_identity, lifecycle_status, created_at_utc, updated_at_utc
               ) values ('alpha', 'BTCUSDT', 'LONG', '1h', 3, 1, 'run', 'candidate', 'ACTIVE', ?, ?)
               returning strategy_id""",
            [now, now],
        ).fetchone()[0]
        result_id = connection.execute(
            """insert into strategy_results (
                   strategy_id, report_start_utc, report_end_utc, exchange, commission_rate,
                   initial_balance, final_balance, total_pnl, total_pnl_pct, max_drawdown,
                   max_drawdown_pct, total_fees, total_trades, imported_at_utc
               ) values (?, ?, ?, 'Bybit', .0004, 100, 101, 1, 1, 0, 0, 0, 1, ?)
               returning result_id""",
            [strategy_id, now, datetime(2026, 1, 9, tzinfo=UTC), now],
        ).fetchone()[0]
        connection.execute(
            "update strategies set current_result_id = ? where strategy_id = ?", [result_id, strategy_id]
        )
        connection.execute(
            "insert into analysis_plateaus (analysis_run_id, plateau_id, plateau_point_count, plateau_total_trades) values ('run', 'P1', 1, 1)"
        )
        connection.execute(
            """insert into strategy_orders (
                   strategy_id, order_id, open_ma_len, open_multiplier, shift_bp, lot_x,
                   analysis_run_id, plateau_id, base_point_trades
               ) values (?, 1, 7, .995, 125, 1, 'run', 'P1', 1)""",
            [strategy_id],
        )
        connection.execute(
            "insert into strategy_tags values (?, 'RETEST', 'TEST', 'fixture', ?)", [strategy_id, now]
        )
        database_id = connection.execute(
            "select value from schema_info where key = 'database_instance_id'"
        ).fetchone()[0]
        connection.execute(
            """insert into selection_runs (
                   selection_run_id, database_instance_id, symbol, side, selection_contract_version,
                   request_json, request_sha256, config_json, config_sha256, candidate_count,
                   representative_count, auto_finalist_count, top_n, workbook_sha256, created_at_utc
               ) values ('run-1', ?, 'BTCUSDT', 'LONG', 'selection-v1', '{}', ?, '{}', ?, 1, 1, 1, 1, ?, ?)""",
            [database_id, "a" * 64, "b" * 64, "c" * 64, now],
        )
        connection.execute(
            """insert into selection_results (
                   selection_run_id, strategy_id, result_id_at_selection, auto_status, auto_score,
                   auto_rank, auto_reason, analog_group_key, auto_analog_of_strategy_id,
                   prior_rejected, stage_trace_json
               ) values ('run-1', ?, ?, 'FINALIST', 1, 1, null, null, null, false, '{}')""",
            [strategy_id, result_id],
        )
        connection.execute(
            "insert into selection_review_imports values ('review-1', 'run-1', ?, ?, 1)",
            ["d" * 64, now],
        )
        connection.execute(
            "insert into selection_review_rows values ('review-1', ?, 'REJECTED', null, null, 'operator decision')",
            [strategy_id],
        )
    return path, int(strategy_id)


def _catalog_identity(path: Path) -> tuple[object, ...]:
    with duckdb.connect(str(path), read_only=True) as connection:
        return (
            tuple(connection.execute("select key, value from schema_info order by key").fetchall()),
            tuple(connection.execute(
                "select table_name, column_name, data_type from information_schema.columns "
                "where table_schema = 'main' order by table_name, ordinal_position"
            ).fetchall()),
        )


def test_performance_v2_export_query_preserves_union_and_retest_intersection() -> None:
    selection = parse_performance_v2_export_query("status=FINALIST&status=RESERVE")
    assert selection.statuses == ("FINALIST", "RESERVE")
    assert parse_performance_v2_export_query("status=FINALIST&retest=1").retest is True
    assert parse_performance_v2_export_query("retest=1").statuses == ()
    assert parse_performance_v2_export_query("all_active=true").all_active is True
    with pytest.raises(PerformanceV2ApiError) as raised:
        parse_performance_v2_export_query("all_active=true&retest=1")
    assert (raised.value.code, raised.value.status) == ("all_active_mixed_with_status", 400)


def test_performance_v2_export_rejects_unknown_query_parameter(tmp_path: Path) -> None:
    controller = PanelController(tmp_path, tmp_path / "config.local.json")
    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        connection.request("GET", "/api/v2/strategies/performance-v2/export?wat=1")
        response = connection.getresponse()
        body = json.loads(response.read().decode("utf-8"))
        connection.close()
        assert response.status == 400
        assert body == {"error": "unknown_query_parameter", "message": "Unknown query parameter."}
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_read_only_export_accepts_v5_without_schema_or_review_changes(tmp_path: Path) -> None:
    database, strategy_id = _export_database(tmp_path / "strategy_performance.duckdb")
    selection = PerformanceV2ExportSelection(all_active=True)
    fixed_now = datetime(2026, 1, 2, tzinfo=UTC)

    _, v6_payload = export_performance_v2(database, selection, now=fixed_now)
    v6_sheet = load_workbook(BytesIO(v6_payload), data_only=True)["All candidates"]
    v6_headers = [cell.value for cell in v6_sheet[1]]

    with duckdb.connect(str(database)) as connection:
        connection.execute("drop table equity_quality_metrics")
        connection.execute("update schema_info set value = '5' where key = 'schema_version'")

    before_file = (sha256(database.read_bytes()).hexdigest(), database.stat().st_mtime_ns)
    before_catalog = _catalog_identity(database)
    _, payload = export_performance_v2(database, selection, now=fixed_now)
    after_file = (sha256(database.read_bytes()).hexdigest(), database.stat().st_mtime_ns)
    after_catalog = _catalog_identity(database)

    sheet = load_workbook(BytesIO(payload), data_only=True)["All candidates"]
    headers = [cell.value for cell in sheet[1]]
    assert headers == v6_headers
    assert not {"Equity state", "Equity basis", "Equity DD, %", "Equity smoothness"}.intersection(headers)
    assert sheet.cell(2, headers.index("ID") + 1).value == strategy_id
    assert sheet.cell(2, headers.index("Auto Status") + 1).value == "REJECTED"
    assert sheet.cell(2, headers.index("User Status") + 1).value == "REJECTED"
    assert sheet.cell(2, headers.index("RETEST") + 1).value == "RETEST"
    assert before_file == after_file
    assert before_catalog == after_catalog


def test_read_only_export_includes_only_fresh_cached_equity_facts_without_writes(tmp_path: Path, monkeypatch) -> None:
    database, strategy_id = _export_database(tmp_path / "strategy_performance.duckdb")
    with duckdb.connect(str(database)) as connection:
        result_id, report_start, report_end = connection.execute(
            "select s.current_result_id, r.report_start_utc, r.report_end_utc from strategies s "
            "join strategy_results r on r.result_id = s.current_result_id where s.strategy_id = ?", [strategy_id],
        ).fetchone()
        metadata = current_equity_source_metadata(connection, int(result_id))
        report_start = report_start.astimezone(UTC)
        report_end = report_end.astimezone(UTC)
        facts = calculate_equity_quality_facts(int(result_id), report_start, report_end, (
            EquitySample(int(result_id), 0, report_start, 100),
            EquitySample(int(result_id), 1, report_start + timedelta(days=1), 151),
            EquitySample(int(result_id), 2, report_start + timedelta(days=2), 120),
            EquitySample(int(result_id), 3, report_end, 140),
        ))
        upsert_equity_quality_facts_checked(connection, [(metadata, facts)], calculated_at_utc=report_end)

    monkeypatch.setattr(selection_module, "_load_source", lambda *args: (_ for _ in ()).throw(AssertionError("raw source read")))
    monkeypatch.setattr(selection_module, "_load_equity_samples_for_quality", lambda *args: (_ for _ in ()).throw(AssertionError("raw equity read")))
    monkeypatch.setattr(selection_module, "_persist", lambda *args: (_ for _ in ()).throw(AssertionError("cache write")))
    before_file = (sha256(database.read_bytes()).hexdigest(), database.stat().st_mtime_ns)
    before_catalog = _catalog_identity(database)
    _, payload = export_performance_v2(database, PerformanceV2ExportSelection(all_active=True))
    after_file = (sha256(database.read_bytes()).hexdigest(), database.stat().st_mtime_ns)
    after_catalog = _catalog_identity(database)

    sheet = load_workbook(BytesIO(payload), data_only=True)["All candidates"]
    headers = [cell.value for cell in sheet[1]]
    assert headers[-4:] == ["Equity state", "Equity basis", "Equity DD, %", "Equity smoothness"]
    assert sheet.cell(2, headers.index("Equity state") + 1).value == facts.state
    assert sheet.cell(2, headers.index("Equity basis") + 1).value == "7d / PROVISIONAL"
    drawdown_cell = sheet.cell(2, headers.index("Equity DD, %") + 1)
    smoothness_cell = sheet.cell(2, headers.index("Equity smoothness") + 1)
    assert facts.drawdown is not None and facts.drawdown * 100 > 0
    assert facts.drawdown * 100 != (facts.drawdown * 100).quantize(Decimal("0.01"))
    assert facts.windows[-1].er != facts.windows[-1].er.quantize(Decimal("0.01"))
    assert drawdown_cell.data_type == "n" and drawdown_cell.value == pytest.approx(float(facts.drawdown * 100))
    assert smoothness_cell.data_type == "n" and smoothness_cell.value == pytest.approx(float(facts.windows[-1].er))
    assert sheet.cell(2, headers.index("User Status") + 1).value == "REJECTED"
    assert sheet.cell(2, headers.index("RETEST") + 1).value == "RETEST"
    assert before_file == after_file
    assert before_catalog == after_catalog


def test_read_only_export_omits_invalid_cached_equity_facts_without_writes(tmp_path: Path, monkeypatch) -> None:
    database, strategy_id = _export_database(tmp_path / "strategy_performance.duckdb")
    with duckdb.connect(str(database)) as connection:
        result_id, report_start, report_end = connection.execute(
            "select s.current_result_id, r.report_start_utc, r.report_end_utc from strategies s "
            "join strategy_results r on r.result_id = s.current_result_id where s.strategy_id = ?", [strategy_id],
        ).fetchone()
        report_start = report_start.astimezone(UTC)
        report_end = report_end.astimezone(UTC)
        facts = calculate_equity_quality_facts(int(result_id), report_start, report_end, (
            EquitySample(int(result_id), 0, report_start, 100),
            EquitySample(int(result_id), 1, report_end, 110),
        ))
        metadata = current_equity_source_metadata(connection, int(result_id))
        upsert_equity_quality_facts_checked(connection, [(metadata, facts)], calculated_at_utc=report_end)
        connection.execute("update equity_quality_metrics set facts_json = '{}' where result_id = ?", [result_id])

    monkeypatch.setattr(selection_module, "_load_source", lambda *args: (_ for _ in ()).throw(AssertionError("raw source read")))
    monkeypatch.setattr(selection_module, "_load_equity_samples_for_quality", lambda *args: (_ for _ in ()).throw(AssertionError("raw equity read")))
    monkeypatch.setattr(selection_module, "_persist", lambda *args: (_ for _ in ()).throw(AssertionError("cache write")))
    before_file = (sha256(database.read_bytes()).hexdigest(), database.stat().st_mtime_ns)
    before_catalog = _catalog_identity(database)

    _, payload = export_performance_v2(database, PerformanceV2ExportSelection(all_active=True))

    after_file = (sha256(database.read_bytes()).hexdigest(), database.stat().st_mtime_ns)
    after_catalog = _catalog_identity(database)
    sheet = load_workbook(BytesIO(payload), data_only=True)["All candidates"]
    headers = [cell.value for cell in sheet[1]]
    assert sheet.cell(2, headers.index("ID") + 1).value == strategy_id
    assert not {"Equity state", "Equity basis", "Equity DD, %", "Equity smoothness"}.intersection(headers)
    assert before_file == after_file
    assert before_catalog == after_catalog


def test_read_only_export_rejects_v6_marker_with_missing_equity_table(tmp_path: Path) -> None:
    database, _ = _export_database(tmp_path / "strategy_performance.duckdb")
    with duckdb.connect(str(database)) as connection:
        connection.execute("drop table equity_quality_metrics")

    with pytest.raises(PerformanceV2ApiError) as raised:
        export_performance_v2(database, PerformanceV2ExportSelection(all_active=True))

    assert (raised.value.code, raised.value.status) == ("PERFORMANCE_V2_SCHEMA_INVALID", 500)


def test_read_only_export_maps_unsupported_schema_version_to_schema_invalid(tmp_path: Path) -> None:
    database, _ = _export_database(tmp_path / "strategy_performance.duckdb")
    with duckdb.connect(str(database)) as connection:
        connection.execute("update schema_info set value = '4' where key = 'schema_version'")

    with pytest.raises(PerformanceV2ApiError) as raised:
        export_performance_v2(database, PerformanceV2ExportSelection(all_active=True))

    assert (raised.value.code, raised.value.status) == ("PERFORMANCE_V2_SCHEMA_INVALID", 500)
