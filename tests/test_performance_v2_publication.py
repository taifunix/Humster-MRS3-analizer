from __future__ import annotations

from datetime import UTC, datetime
from dataclasses import asdict, replace
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path

import duckdb
import pytest
import mrs3.performance_v2_publication as publication_module
import mrs3.performance_v2_selection_review as selection_review_module
from openpyxl import load_workbook

from mrs3.performance_v2_publication import (
    PublicationConflict,
    PublicationNotFound,
    PublicationAggregate,
    PublicationPackage,
    PublicationPartition,
    PublicationReview,
    PublicationRun,
    PublicationValidationError,
    publish_publication,
    review_import_key,
)
from mrs3.performance_v2_selection import AllPairsPartition, build_combined_selection_workbook
from mrs3.performance_v2_selection_review import import_combined_selection_workbook
from mrs3.performance_v2_store import initialize_performance_v2


def _seed_database(tmp_path: Path) -> tuple[duckdb.DuckDBPyConnection, Path]:
    root = tmp_path / "performance"
    root.mkdir()
    connection = duckdb.connect(str(root / "strategy_performance.duckdb"))
    initialize_performance_v2(connection)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    now = datetime(2026, 10, 10, tzinfo=UTC)
    connection.execute(
        """insert into strategies (
            strategy_id, strategy_name, symbol, side, timeframe, close_ma_len,
            order_count, analysis_run_id, candidate_identity, lifecycle_status,
            created_at_utc, updated_at_utc
        ) values (1, 'one', 'BTCUSDT', 'LONG', '1h', 5, 1, 'analysis', 'candidate-1', 'ACTIVE', ?, ?)""",
        [now, now],
    )
    connection.execute(
        """insert into strategy_results (
            result_id, strategy_id, report_start_utc, report_end_utc, exchange,
            initial_balance, final_balance, imported_at_utc
        ) values (101, 1, ?, ?, 'BYBIT', 100, 101, ?)""",
        [now, now, now],
    )
    connection.execute("update strategies set current_result_id = 101 where strategy_id = 1")
    connection.execute(
        """insert into selection_runs values (
            'source-run', ?, 'BTCUSDT', 'LONG', 'selection-v1', '{}', 'request', '{}',
            'config', 1, 1, 1, 1, 'source-workbook', ?
        )""",
        [db_id, now],
    )
    connection.execute(
        """insert into selection_results (
            selection_run_id, strategy_id, result_id_at_selection, auto_status,
            auto_score, auto_rank, auto_reason, analog_group_key,
            auto_analog_of_strategy_id, prior_rejected, stage_trace_json
        ) values ('source-run', 1, 101, 'FINALIST', 1, 1, null, null, null, false, '{}')"""
    )
    connection.commit()
    return connection, root


@pytest.mark.parametrize("stored_instance_id", [None, "", "  "])
def test_database_instance_id_fails_closed_for_missing_identity(stored_instance_id: object) -> None:
    class FakeResult:
        def fetchone(self) -> tuple[object]:
            return (stored_instance_id,)

    class FakeConnection:
        def execute(self, query: str) -> FakeResult:
            return FakeResult()

    with pytest.raises(selection_review_module.SelectionReviewError) as error:
        selection_review_module.database_instance_id(FakeConnection())

    assert error.value.code == "SELECTION_REVIEW_DATABASE_MISMATCH"


def _authorize_publication_scope(
    connection: duckdb.DuckDBPyConnection,
    db_id: str,
    publication_id: str,
    decision_group_id: str,
) -> None:
    now = datetime(2026, 10, 10, tzinfo=UTC)
    connection.execute(
        """insert into selection_publications (
            publication_id, publication_kind, operation_key, operation_digest,
            manifest_contract_version, decision_group_id, database_instance_id,
            source_revision, controls_json, controls_sha256, render_model_json,
            render_model_sha256, evaluated_rowset_sha256, export_workbook_sha256,
            created_at_utc
        ) values (?, 'AUTO_REJECTION_OVERLAY', ?, 'digest', 'v11', ?, ?,
                  'source', '{}', 'controls', '{}', 'render', 'rowset', null, ?)""",
        [publication_id, f"scope-{publication_id}", decision_group_id, db_id, now],
    )
    connection.execute(
        """insert into selection_publication_runs
            (publication_id, pair, side, role, selection_run_id)
            values (?, 'BTCUSDT', 'LONG', 'SOURCE', 'source-run')""",
        [publication_id],
    )
    connection.commit()


def _add_second_candidate(connection: duckdb.DuckDBPyConnection) -> None:
    now = datetime(2026, 10, 10, tzinfo=UTC)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    connection.execute(
        """insert into strategies (
            strategy_id, strategy_name, symbol, side, timeframe, close_ma_len,
            order_count, analysis_run_id, candidate_identity, lifecycle_status,
            created_at_utc, updated_at_utc
        ) values (2, 'two', 'BTCUSDT', 'LONG', '1h', 5, 1, 'analysis', 'candidate-2', 'ACTIVE', ?, ?)""",
        [now, now],
    )
    connection.execute(
        """insert into strategy_results (
            result_id, strategy_id, report_start_utc, report_end_utc, exchange,
            initial_balance, final_balance, imported_at_utc
        ) values (102, 2, ?, ?, 'BYBIT', 100, 101, ?)""",
        [now, now, now],
    )
    connection.execute("update strategies set current_result_id = 102 where strategy_id = 2")
    connection.execute(
        """insert into selection_results (
            selection_run_id, strategy_id, result_id_at_selection, auto_status,
            auto_score, auto_rank, auto_reason, analog_group_key,
            auto_analog_of_strategy_id, prior_rejected, stage_trace_json
        ) values ('source-run', 2, 102, 'FILTERED', null, null, null, null, null, false, '{}')"""
    )
    connection.commit()


def _package(*, operation_key: str = "operation-1", operation_digest: str = "digest-1") -> PublicationPackage:
    return PublicationPackage(
        publication_id="publication-1",
        publication_kind="AUTO_REJECTION_OVERLAY",
        operation_key=operation_key,
        operation_digest=operation_digest,
        manifest_contract_version="v11",
        decision_group_id="decision-group-1",
        database_instance_id=None,
        source_revision="source-revision-1",
        controls_json="{}",
        controls_sha256=sha256(b"{}").hexdigest(),
        render_model_json="{}",
        render_model_sha256=sha256(b"{}").hexdigest(),
        evaluated_rowset_sha256="rowset-1",
        partitions=(
            PublicationPartition(
                pair="BTCUSDT",
                side="LONG",
                source_run_id="source-run",
                overlay_run=PublicationRun(
                    selection_run_id="overlay-run",
                    database_instance_id=None,
                    symbol="BTCUSDT",
                    side="LONG",
                    selection_contract_version="selection-v1",
                    request_json='{"ranking_scope":"AUTOMATIC_REJECTION_OVERLAY"}',
                    request_sha256="request-sha",
                    config_json="{}",
                    config_sha256="config-sha",
                    candidate_count=1,
                    representative_count=1,
                    auto_finalist_count=0,
                    top_n=1,
                    results=(
                        {
                            "strategy_id": 1,
                            "result_id_at_selection": 101,
                            "auto_status": "FILTERED",
                            "auto_score": None,
                            "auto_rank": None,
                            "auto_reason": "FILTER_HARD_CUTOFFS",
                            "analog_group_key": None,
                            "auto_analog_of_strategy_id": None,
                            "prior_rejected": False,
                            "stage_trace_json": "{}",
                            "equity_regime_json": None,
                        },
                    ),
                ),
                reviews=(
                    PublicationReview(
                        strategy_id=1,
                        user_status="REJECTED",
                        user_rank=None,
                        user_analog_of_strategy_id=None,
                        comment="Degraded Finalist",
                    ),
                ),
            ),
        ),
    )


def test_publication_uses_one_timestamp_for_header_children_and_tag(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)

    result = publish_publication(connection, root, _package())

    assert result.publication_id == "publication-1"
    assert result.decision_group_id == "decision-group-1"
    timestamps = connection.execute(
        """select created_at_utc from selection_publications where publication_id = 'publication-1'
           union all select imported_at_utc from selection_review_imports
           union all select updated_at_utc from strategy_tags"""
    ).fetchall()
    assert len(timestamps) == 3
    assert len({row[0] for row in timestamps}) == 1
    assert timestamps[0][0].microsecond >= 0
    assert connection.execute(
        "select aggregate_import_id from selection_review_imports"
    ).fetchone() == (None,)
    assert connection.execute(
        "select role from selection_publication_runs where publication_id = 'publication-1' order by role"
    ).fetchall() == [("OVERLAY",), ("SOURCE",)]


def test_aggregate_workbook_import_unchanged_row_is_a_noop(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-1")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST",
                "auto_rank": 1, "auto_reason": None,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-1", decision_group_id="decision-1",
        user_review_rows={1: {"retest": "RETEST"}},
    )

    result = import_combined_selection_workbook(
        connection, root, data, aggregate_import_id="aggregate-1",
        operation_key="aggregate-operation-1",
    )

    assert result["aggregate_import_id"] == "aggregate-1"
    assert result["changed_count"] == 0
    assert connection.execute("select count(*) from selection_review_imports").fetchone() == (0,)


def test_aggregate_import_requires_complete_partition_identity_set(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    _add_second_candidate(connection)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-complete")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-complete", decision_group_id="decision-complete",
    )

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, data, aggregate_import_id="aggregate-complete",
            operation_key="aggregate-operation-complete",
        )

    assert getattr(error.value, "code", None) == "SELECTION_REVIEW_CANDIDATE_MISMATCH"
    assert connection.execute("select count(*) from selection_aggregate_imports").fetchone() == (0,)


def test_aggregate_import_rejects_rank_collision_with_external_effective_holder(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    connection, root = _seed_database(tmp_path)
    _add_second_candidate(connection)
    connection.execute(
        "delete from selection_results where selection_run_id = 'source-run' and strategy_id = 2"
    )
    connection.commit()
    monkeypatch.setattr(
        selection_review_module,
        "effective_selection_decisions",
        lambda *args, **kwargs: {2: ("FINALIST", 1, "prior-run")},
    )
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-external-rank")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-external-rank", decision_group_id="decision-external-rank",
    )
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["All candidates"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["User Status"]).value = "FINALIST"
    sheet.cell(2, headers["User Rank"]).value = 1
    changed = BytesIO()
    workbook.save(changed)

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, changed.getvalue(), aggregate_import_id="aggregate-external-rank",
            operation_key="aggregate-operation-external-rank",
        )

    assert getattr(error.value, "code", None) == "SELECTION_REVIEW_INVALID_RANK"
    assert connection.execute("select count(*) from selection_aggregate_imports").fetchone() == (0,)


def test_aggregate_import_rejects_rank_collision_with_auto_only_holder(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    connection, root = _seed_database(tmp_path)
    _add_second_candidate(connection)
    connection.execute(
        "update selection_results set auto_status = 'FINALIST', auto_rank = 1 where selection_run_id = 'source-run' and strategy_id = 2"
    )
    connection.commit()
    monkeypatch.setattr(
        selection_review_module,
        "effective_selection_decisions",
        lambda *args, **kwargs: {2: ("FINALIST", 1, "source-run")},
    )
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-auto-rank")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([
                {"strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1},
                {"strategy_id": 2, "result_id": 102, "auto_status": "FINALIST", "auto_rank": 1},
            ]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-auto-rank", decision_group_id="decision-auto-rank",
    )
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["All candidates"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, headers["ID"]).value == 1:
            sheet.cell(row, headers["User Status"]).value = "FINALIST"
            sheet.cell(row, headers["User Rank"]).value = 1
    changed = BytesIO()
    workbook.save(changed)

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, changed.getvalue(), aggregate_import_id="aggregate-auto-rank",
            operation_key="aggregate-operation-auto-rank",
        )

    assert getattr(error.value, "code", None) == "SELECTION_REVIEW_INVALID_RANK"
    assert connection.execute("select count(*) from selection_aggregate_imports").fetchone() == (0,)


@pytest.mark.parametrize("effective_decisions", [
    {2: ("FINALIST", 1, "prior-run"), 3: ("RESERVE", 1, "prior-run")},
    {2: ("FINALIST", "malformed", "prior-run"), 3: ("RESERVE", 1, "prior-run")},
], ids=["duplicate", "malformed-unrelated"])
def test_aggregate_import_tolerates_unrelated_preexisting_effective_ranks(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, effective_decisions: dict[int, tuple[object, object, str]],
) -> None:
    connection, root = _seed_database(tmp_path)
    _add_second_candidate(connection)
    now = datetime(2026, 10, 10, tzinfo=UTC)
    connection.execute(
        """insert into strategies (
            strategy_id, strategy_name, symbol, side, timeframe, close_ma_len,
            order_count, analysis_run_id, candidate_identity, lifecycle_status,
            created_at_utc, updated_at_utc
        ) values (3, 'three', 'BTCUSDT', 'LONG', '1h', 5, 1, 'analysis', 'candidate-3', 'ACTIVE', ?, ?)""",
        [now, now],
    )
    connection.execute(
        """insert into strategy_results (
            result_id, strategy_id, report_start_utc, report_end_utc, exchange,
            initial_balance, final_balance, imported_at_utc
        ) values (103, 3, ?, ?, 'BYBIT', 100, 101, ?)""",
        [now, now, now],
    )
    connection.execute("update strategies set current_result_id = 103 where strategy_id = 3")
    connection.execute(
        "delete from selection_results where selection_run_id = 'source-run' and strategy_id = 2"
    )
    connection.commit()
    monkeypatch.setattr(
        selection_review_module,
        "effective_selection_decisions",
        lambda *args, **kwargs: effective_decisions,
    )
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-duplicate-effective")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-duplicate-effective",
        decision_group_id="decision-duplicate-effective",
    )

    workbook = load_workbook(BytesIO(data))
    sheet = workbook["All candidates"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["User Status"]).value = "FINALIST"
    sheet.cell(2, headers["User Rank"]).value = 2
    changed = BytesIO()
    workbook.save(changed)

    result = import_combined_selection_workbook(
        connection, root, changed.getvalue(), aggregate_import_id="aggregate-duplicate-effective",
        operation_key="aggregate-operation-duplicate-effective",
    )

    assert result["changed_count"] == 1
    assert connection.execute("select count(*) from selection_aggregate_imports").fetchone() == (1,)


def test_aggregate_import_rank_swap_projects_post_import_effective_map(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    _add_second_candidate(connection)
    now = datetime(2026, 10, 10, 1, tzinfo=UTC)
    connection.execute(
        "insert into selection_review_imports (review_import_id, selection_run_id, workbook_sha256, imported_at_utc, row_count) values ('swap-head', 'source-run', 'swap-workbook', ?, 2)",
        [now],
    )
    connection.executemany(
        "insert into selection_review_rows (review_import_id, strategy_id, user_status, user_rank, user_analog_of_strategy_id, comment) values ('swap-head', ?, 'FINALIST', ?, null, null)",
        [(1, 1), (2, 2)],
    )
    connection.commit()
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-swap")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([
                {"strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1},
                {"strategy_id": 2, "result_id": 102, "auto_status": "FILTERED"},
            ]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-swap", decision_group_id="decision-swap",
        user_review_rows={
            1: {"user_status": "FINALIST", "user_rank": 1},
            2: {"user_status": "FINALIST", "user_rank": 2},
        },
    )
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["All candidates"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    for row in range(2, sheet.max_row + 1):
        strategy_id = sheet.cell(row, headers["ID"]).value
        sheet.cell(row, headers["User Rank"]).value = 2 if strategy_id == 1 else 1
    changed = BytesIO()
    workbook.save(changed)

    result = import_combined_selection_workbook(
        connection, root, changed.getvalue(), aggregate_import_id="aggregate-swap",
        operation_key="aggregate-operation-swap",
    )

    assert result["changed_count"] == 2


def test_aggregate_import_comment_edit_does_not_recheck_old_duplicate_rank(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    connection, root = _seed_database(tmp_path)
    _add_second_candidate(connection)
    now = datetime(2026, 10, 10, 1, tzinfo=UTC)
    connection.execute(
        "insert into selection_review_imports (review_import_id, selection_run_id, workbook_sha256, imported_at_utc, row_count) values ('comment-head', 'source-run', 'comment-workbook', ?, 2)",
        [now],
    )
    connection.executemany(
        "insert into selection_review_rows (review_import_id, strategy_id, user_status, user_rank, user_analog_of_strategy_id, comment) values ('comment-head', ?, 'FINALIST', 1, null, null)",
        [(1,), (2,)],
    )
    connection.commit()
    monkeypatch.setattr(
        selection_review_module,
        "effective_selection_decisions",
        lambda *args, **kwargs: {
            1: ("FINALIST", 1, "source-run"), 2: ("FINALIST", 1, "source-run"),
        },
    )
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-comment-rank")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([
                {"strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1},
                {"strategy_id": 2, "result_id": 102, "auto_status": "FILTERED"},
            ]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-comment-rank", decision_group_id="decision-comment-rank",
        user_review_rows={
            1: {"user_status": "FINALIST", "user_rank": 1},
            2: {"user_status": "FINALIST", "user_rank": 1},
        },
    )
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["All candidates"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["Comment"]).value = "comment-only edit"
    changed = BytesIO()
    workbook.save(changed)

    result = import_combined_selection_workbook(
        connection, root, changed.getvalue(), aggregate_import_id="aggregate-comment-rank",
        operation_key="aggregate-operation-comment-rank",
    )

    assert result["changed_count"] == 1


def test_aggregate_import_stale_precedes_rank_collision(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    _add_second_candidate(connection)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-stale-rank")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([
                {"strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1},
                {"strategy_id": 2, "result_id": 102, "auto_status": "FILTERED"},
            ]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-stale-rank", decision_group_id="decision-stale-rank",
    )
    connection.execute(
        "insert into selection_review_imports (review_import_id, selection_run_id, workbook_sha256, imported_at_utc, row_count) values ('stale-rank-head', 'source-run', 'stale-rank-workbook', ?, 1)",
        [datetime(2026, 10, 10, 2, tzinfo=UTC)],
    )
    connection.execute(
        "insert into selection_review_rows (review_import_id, strategy_id, user_status, user_rank, user_analog_of_strategy_id, comment) values ('stale-rank-head', 1, 'RESERVE', null, null, 'newer')"
    )
    connection.commit()
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["All candidates"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, headers["User Status"]).value = "FINALIST"
        sheet.cell(row, headers["User Rank"]).value = 1
    changed = BytesIO()
    workbook.save(changed)

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, changed.getvalue(), aggregate_import_id="aggregate-stale-rank",
            operation_key="aggregate-operation-stale-rank",
        )

    assert getattr(error.value, "code", None) == "SELECTION_REVIEW_STALE"
    assert connection.execute("select count(*) from selection_aggregate_imports").fetchone() == (0,)


def test_aggregate_import_replay_cross_publication_is_operation_conflict(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-cross-1")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-cross", decision_group_id="decision-cross-1",
    )
    first = import_combined_selection_workbook(
        connection, root, data, aggregate_import_id="aggregate-cross",
        operation_key="aggregate-operation-cross",
    )
    assert first["code"] != "ALREADY_IMPORTED"
    _authorize_publication_scope(connection, db_id, "publication-other", "decision-cross-2")
    other = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-other",
        aggregate_import_id="aggregate-cross", decision_group_id="decision-cross-2",
    )

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, other, aggregate_import_id="aggregate-cross",
            operation_key="aggregate-operation-cross",
        )

    assert getattr(error.value, "code", None) == "OPERATION_KEY_CONFLICT"


def test_aggregate_import_missing_aggregate_header_is_typed_not_found(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-missing-aggregate")
    connection.execute(
        """insert into selection_publications (
            publication_id, publication_kind, operation_key, operation_digest,
            manifest_contract_version, decision_group_id, database_instance_id,
            source_revision, controls_json, controls_sha256, render_model_json,
            render_model_sha256, evaluated_rowset_sha256, export_workbook_sha256,
            created_at_utc
        ) values ('publication-missing-aggregate', 'AGGREGATE_REVIEW',
                  'aggregate-operation-missing', 'digest-missing', 'v11',
                  'decision-missing-aggregate', ?, 'source', '{}', 'controls',
                  '{}', 'render', 'rowset', null, ?)""",
        [db_id, datetime(2026, 10, 10, 1, tzinfo=UTC)],
    )
    connection.commit()
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-missing", decision_group_id="decision-missing-aggregate",
    )

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, data, aggregate_import_id="aggregate-missing",
            operation_key="aggregate-operation-missing",
        )

    assert getattr(error.value, "code", None) == "AGGREGATE_NOT_FOUND"


def test_aggregate_workbook_import_preserves_raw_sanitized_review_tuple_and_scope_case(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    _add_second_candidate(connection)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-raw")
    connection.execute("update selection_runs set database_instance_id = upper(database_instance_id)")
    now = datetime(2026, 10, 10, 1, tzinfo=UTC)
    connection.execute(
        "insert into selection_review_imports (review_import_id, selection_run_id, workbook_sha256, imported_at_utc, row_count) values ('raw-head', 'source-run', 'raw-head-workbook', ?, 2)",
        [now],
    )
    connection.executemany(
        "insert into selection_review_rows (review_import_id, strategy_id, user_status, user_rank, user_analog_of_strategy_id, comment) values ('raw-head', ?, ?, ?, ?, ?)",
        [
            (1, "FINALIST", 1, None, "=1+1"),
            (2, "ANALOG", None, 1, ""),
        ],
    )
    connection.commit()
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([
                {"strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1},
                {"strategy_id": 2, "result_id": 102, "auto_status": "FILTERED"},
            ]),
        )],
        database_instance_id=db_id,
        publication_id="PUBLICATION-SOURCE",
        aggregate_import_id="AGGREGATE-RAW",
        decision_group_id="DECISION-RAW",
        user_review_rows={
            1: {"user_status": "FINALIST", "user_rank": 1, "comment": "=1+1"},
            2: {"user_status": "ANALOG", "user_analog_of_strategy_id": 1, "comment": ""},
        },
    )

    result = import_combined_selection_workbook(
        connection, root, data,
        aggregate_import_id=" AgGrEgAtE-RaW ",
        operation_key=" AGGREGATE-RAW-OP ",
    )

    assert result["changed_count"] == 0
    assert result["unchanged_count"] == 2
    assert result["aggregate_import_id"] == "aggregate-raw"
    assert connection.execute("select count(*) from selection_review_imports").fetchone() == (1,)


def test_aggregate_status_edit_preserves_raw_frozen_comment(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    now = datetime(2026, 10, 10, 1, tzinfo=UTC)
    connection.execute(
        "insert into selection_review_imports (review_import_id, selection_run_id, workbook_sha256, imported_at_utc, row_count) values ('raw-status-head', 'source-run', 'raw-status-workbook', ?, 1)",
        [now],
    )
    connection.execute(
        "insert into selection_review_rows (review_import_id, strategy_id, user_status, user_rank, user_analog_of_strategy_id, comment) values ('raw-status-head', 1, 'FINALIST', 1, null, '=1+1')"
    )
    connection.commit()
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-status-raw")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-status-raw", decision_group_id="decision-status-raw",
        user_review_rows={1: {"user_status": "FINALIST", "user_rank": 1, "comment": "=1+1"}},
    )
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["All candidates"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["User Status"]).value = "REJECTED"
    sheet.cell(2, headers["User Rank"]).value = None
    changed = BytesIO()
    workbook.save(changed)

    result = import_combined_selection_workbook(
        connection, root, changed.getvalue(), aggregate_import_id="aggregate-status-raw",
        operation_key="aggregate-operation-status-raw",
    )

    assert result["changed_count"] == 1
    assert connection.execute(
        """select rows.user_status, rows.user_rank, rows.comment
             from selection_review_rows rows
             join selection_review_imports imports using (review_import_id)
            where imports.aggregate_import_id = 'aggregate-status-raw'"""
    ).fetchall()[-1] == ("REJECTED", None, "=1+1")


def test_aggregate_workbook_import_retest_change_is_typed_before_publication(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-retest")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-retest", decision_group_id="decision-retest",
        user_review_rows={1: {"retest": "RETEST"}},
    )
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["All candidates"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["RETEST"]).value = None
    changed = BytesIO()
    workbook.save(changed)

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, changed.getvalue(), aggregate_import_id="aggregate-retest",
            operation_key="aggregate-operation-retest",
        )

    assert getattr(error.value, "code", None) == "SELECTION_REVIEW_RETEST_READ_ONLY"
    assert connection.execute("select count(*) from selection_aggregate_imports").fetchone() == (0,)


def test_aggregate_workbook_import_changed_tuple_creates_manual_review_and_retries_idempotently(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-2")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST",
                "auto_rank": 1, "auto_reason": None,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-2", decision_group_id="decision-2",
    )
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["All candidates"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["User Status"], "FINALIST")
    sheet.cell(2, headers["User Rank"], 1)
    changed = BytesIO()
    workbook.save(changed)

    result = import_combined_selection_workbook(
        connection, root, changed.getvalue(), aggregate_import_id="aggregate-2",
        operation_key="aggregate-operation-2",
    )

    assert result["changed_count"] == 1
    assert connection.execute(
        "select user_status, user_rank from selection_review_rows"
    ).fetchone() == ("FINALIST", 1)
    retry = import_combined_selection_workbook(
        connection, root, changed.getvalue(), aggregate_import_id=" AGGREGATE-2 ",
        operation_key=" AGGREGATE-OPERATION-2 ",
    )
    assert retry["code"] == "ALREADY_IMPORTED"
    assert connection.execute("select count(*) from selection_review_imports").fetchone() == (1,)


def test_aggregate_import_replay_uses_case_insensitive_stored_keys(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-case")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-case", decision_group_id="decision-case",
    )
    workbook_sha256 = sha256(data).hexdigest()
    now = datetime(2026, 10, 10, tzinfo=UTC)
    connection.execute(
        """insert into selection_publications (
            publication_id, publication_kind, operation_key, operation_digest,
            manifest_contract_version, decision_group_id, database_instance_id,
            source_revision, controls_json, controls_sha256, render_model_json,
            render_model_sha256, evaluated_rowset_sha256, export_workbook_sha256,
            created_at_utc
        ) values ('publication-mixed-case', 'AGGREGATE_REVIEW',
                  'AgGrEgAtE-OpErAtIoN-CaSe', 'digest-case', 'v11',
                  'decision-case', ?, 'source', ?, 'controls', '{}',
                  'render', 'rowset', null, ?)""",
        [f" {db_id.upper()} ", json.dumps({"publication_id": "PUBLICATION-SOURCE"}), now],
    )
    connection.execute(
        """insert into selection_aggregate_imports (
            aggregate_import_id, publication_id, operation_key, operation_digest,
            manifest_contract_version, source_revision, partition_rowsets_json,
            candidate_identities_json, uploaded_workbook_sha256, lifecycle_status,
            imported_at_utc
        ) values ('AgGrEgAtE-CaSe', 'publication-mixed-case',
                  'AgGrEgAtE-OpErAtIoN-CaSe', 'digest-case', 'v11', 'source',
                  '[]', '[]', ?, 'ACTIVE', ?)""",
        [workbook_sha256, now],
    )
    connection.commit()

    retry = import_combined_selection_workbook(
        connection, root, data, aggregate_import_id="aggregate-case",
        operation_key="aggregate-operation-case",
    )

    assert retry["code"] == "ALREADY_IMPORTED"
    assert retry["aggregate_import_id"] == "aggregate-case"
    assert retry["publication_id"] == "publication-mixed-case"


def test_aggregate_import_replay_checks_manifest_database_before_idempotency(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-db-replay")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-db-replay", decision_group_id="decision-db-replay",
    )
    import_combined_selection_workbook(
        connection, root, data, aggregate_import_id="aggregate-db-replay",
        operation_key="aggregate-operation-db-replay",
    )
    workbook = load_workbook(BytesIO(data))
    metadata = workbook["_MRS_SELECTION_MANIFEST"]
    manifest = json.loads(metadata.cell(3, 2).value)
    manifest["database_instance_id"] = "different-database"
    manifest_json = json.dumps(manifest, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    metadata.cell(2, 2).value = sha256(manifest_json.encode("utf-8")).hexdigest()
    metadata.cell(3, 2).value = manifest_json
    tampered = BytesIO()
    workbook.save(tampered)

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, tampered.getvalue(), aggregate_import_id="aggregate-db-replay",
            operation_key="aggregate-operation-db-replay",
        )

    assert getattr(error.value, "code", None) == "SELECTION_REVIEW_DATABASE_MISMATCH"


def test_aggregate_import_replay_of_retired_publication_is_typed_not_found(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-retired-replay")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-retired-replay", decision_group_id="decision-retired-replay",
    )
    first = import_combined_selection_workbook(
        connection, root, data, aggregate_import_id="aggregate-retired-replay",
        operation_key="aggregate-operation-retired-replay",
    )
    connection.execute(
        "update selection_publications set retired_at_utc = ? where publication_id = ?",
        [datetime(2026, 10, 10, 3, tzinfo=UTC), first["publication_id"]],
    )
    connection.commit()

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, data, aggregate_import_id="aggregate-retired-replay",
            operation_key="aggregate-operation-retired-replay",
        )

    assert getattr(error.value, "code", None) == "AGGREGATE_NOT_FOUND"


def test_aggregate_workbook_import_rejects_newer_review_head_atomically(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-3")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST",
                "auto_rank": 1, "auto_reason": None,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-3", decision_group_id="decision-3",
    )
    now = datetime(2026, 10, 10, 2, tzinfo=UTC)
    connection.execute(
        "insert into selection_review_imports (review_import_id, selection_run_id, workbook_sha256, imported_at_utc, row_count) values ('newer-head', 'source-run', 'head-workbook', ?, 1)",
        [now],
    )
    connection.execute(
        "insert into selection_review_rows (review_import_id, strategy_id, user_status, user_rank, user_analog_of_strategy_id, comment) values ('newer-head', 1, 'RESERVE', NULL, NULL, 'newer')"
    )

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, data, aggregate_import_id="aggregate-3",
            operation_key="aggregate-operation-3",
        )

    assert getattr(error.value, "code", None) == "SELECTION_REVIEW_STALE"
    assert connection.execute("select count(*) from selection_aggregate_imports").fetchone() == (0,)


def test_aggregate_import_requires_plain_manifest_scope_and_authorized_publication(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1,
            }]),
        )],
        database_instance_id=db_id, aggregate_import_id="aggregate-scope",
    )

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, data, aggregate_import_id="aggregate-scope",
            operation_key="aggregate-operation-scope",
        )

    assert getattr(error.value, "code", None) == "INVALID_ARGUMENT"


def test_aggregate_import_rejects_blank_decision_group(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-blank-decision", decision_group_id=" ",
    )

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, data, aggregate_import_id="aggregate-blank-decision",
            operation_key="aggregate-operation-blank-decision",
        )

    assert getattr(error.value, "code", None) == "INVALID_ARGUMENT"


def test_aggregate_import_rejects_unknown_publication_scope(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1,
            }]),
        )],
        database_instance_id=db_id, publication_id="missing-publication",
        aggregate_import_id="aggregate-scope-missing", decision_group_id="decision-scope",
    )

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, data, aggregate_import_id="aggregate-scope-missing",
            operation_key="aggregate-operation-scope-missing",
        )

    assert getattr(error.value, "code", None) == "NOT_FOUND"


def test_aggregate_import_validates_empty_partitions_before_publish(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-empty")
    columns = ["strategy_id", "result_id", "auto_status", "auto_rank"]
    data = build_combined_selection_workbook(
        [
            AllPairsPartition(
                "BTCUSDT", "LONG", "source-run",
                __import__("pandas").DataFrame([{
                    "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1,
                }]),
            ),
            AllPairsPartition("ETHUSDT", "SHORT", "missing-empty-run", __import__("pandas").DataFrame(columns=columns)),
        ],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-empty", decision_group_id="decision-empty",
    )

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, data, aggregate_import_id="aggregate-empty",
            operation_key="aggregate-operation-empty",
        )

    assert getattr(error.value, "code", None) == "SOURCE_RUN_NOT_FOUND"
    assert connection.execute("select count(*) from selection_aggregate_imports").fetchone() == (0,)


def test_aggregate_workbook_import_reject_status_has_no_rank_leak(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-rejected")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST",
                "auto_rank": 1, "auto_reason": None,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-rejected", decision_group_id="decision-rejected",
    )
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["All candidates"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["User Status"], "REJECTED")
    changed = BytesIO()
    workbook.save(changed)

    result = import_combined_selection_workbook(
        connection, root, changed.getvalue(), aggregate_import_id="aggregate-rejected",
        operation_key="aggregate-operation-rejected",
    )

    assert result["changed_count"] == 1
    assert connection.execute(
        "select user_status, user_rank from selection_review_rows"
    ).fetchone() == ("REJECTED", None)


def test_aggregate_import_rejects_missing_scope_before_workbook_parse(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, b"not-an-xlsx", aggregate_import_id=" ", operation_key=" ",
        )

    assert getattr(error.value, "code", None) == "INVALID_ARGUMENT"


def test_aggregate_import_considers_unchanged_frozen_ranks(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    _add_second_candidate(connection)
    now = datetime(2026, 10, 10, 1, tzinfo=UTC)
    connection.execute(
        "insert into selection_review_imports (review_import_id, selection_run_id, workbook_sha256, imported_at_utc, row_count) values ('rank-head', 'source-run', 'rank-head-workbook', ?, 1)",
        [now],
    )
    connection.execute(
        "insert into selection_review_rows (review_import_id, strategy_id, user_status, user_rank, user_analog_of_strategy_id, comment) values ('rank-head', 1, 'FINALIST', 1, null, null)"
    )
    connection.commit()
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-rank")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([
                {"strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1},
                {"strategy_id": 2, "result_id": 102, "auto_status": "FILTERED"},
            ]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-rank", decision_group_id="decision-rank",
        user_review_rows={1: {"user_status": "FINALIST", "user_rank": 1}},
    )
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["All candidates"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(3, headers["User Status"], "FINALIST")
    sheet.cell(3, headers["User Rank"], 1)
    changed = BytesIO()
    workbook.save(changed)

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, changed.getvalue(), aggregate_import_id="aggregate-rank",
            operation_key="aggregate-operation-rank",
        )

    assert getattr(error.value, "code", None) == "SELECTION_REVIEW_INVALID_RANK"
    assert connection.execute("select count(*) from selection_aggregate_imports").fetchone() == (0,)


def test_aggregate_import_operation_key_collision_with_nonaggregate_is_typed(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    publish_publication(connection, root, _package(operation_key="shared-operation", operation_digest="shared-digest"))
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-collision")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-collision", decision_group_id="decision-collision",
    )

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, data, aggregate_import_id="aggregate-collision",
            operation_key=" SHARED-OPERATION ",
        )

    assert getattr(error.value, "code", None) == "OPERATION_KEY_CONFLICT"
    assert connection.execute("select count(*) from selection_aggregate_imports").fetchone() == (0,)


def test_aggregate_import_analog_target_must_be_in_same_partition(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-analog")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-analog", decision_group_id="decision-analog",
    )
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["All candidates"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["User Status"], "ANALOG")
    sheet.cell(2, headers["Analog Of ID"], 999)
    changed = BytesIO()
    workbook.save(changed)

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, changed.getvalue(), aggregate_import_id="aggregate-analog",
            operation_key="aggregate-operation-analog",
        )

    assert getattr(error.value, "code", None) == "SELECTION_REVIEW_INVALID_ANALOG"


def test_aggregate_import_automatic_integrity_precedes_rank_collision(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    db_id = connection.execute(
        "select value from schema_info where key = 'database_instance_id'"
    ).fetchone()[0]
    _authorize_publication_scope(connection, db_id, "publication-source", "decision-automatic-precedence")
    data = build_combined_selection_workbook(
        [AllPairsPartition(
            "BTCUSDT", "LONG", "source-run",
            __import__("pandas").DataFrame([{
                "strategy_id": 1, "result_id": 101, "auto_status": "FINALIST", "auto_rank": 1,
            }]),
        )],
        database_instance_id=db_id, publication_id="publication-source",
        aggregate_import_id="aggregate-automatic-precedence", decision_group_id="decision-automatic-precedence",
    )
    connection.execute(
        "update selection_results set auto_status = 'FILTERED' where selection_run_id = 'source-run' and strategy_id = 1"
    )
    connection.commit()
    workbook = load_workbook(BytesIO(data))
    sheet = workbook["All candidates"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["User Status"], "FINALIST")
    sheet.cell(2, headers["User Rank"], 2)
    changed = BytesIO()
    workbook.save(changed)

    with pytest.raises(Exception) as error:
        import_combined_selection_workbook(
            connection, root, changed.getvalue(), aggregate_import_id="aggregate-automatic-precedence",
            operation_key="aggregate-operation-automatic-precedence",
        )

    assert getattr(error.value, "code", None) == "SELECTION_REVIEW_AUTOMATIC_FIELDS_CHANGED"


def test_publication_preserves_long_review_comment_verbatim(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    comment = "x" * 1001

    publish_publication(
        connection,
        root,
        replace(_package(), partitions=(replace(
            _package().partitions[0],
            reviews=(PublicationReview(1, "REJECTED", None, None, comment),),
        ),)),
    )

    assert connection.execute(
        "select comment from selection_review_rows"
    ).fetchone() == (comment,)


def test_auto_rejection_source_projects_overlay_lineage_and_timestamp(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    base = _package()
    source = {
        "strategy_id": 1,
        "source_kind": "EQUITY_REGIME_FILTER",
        "reason_code": "W28_DOWN",
        "first_result_id": 101,
        "classifier_algo_version": "M3",
        "source_revision": "auto-revision",
        "facts_sha256": "auto-facts",
    }
    package = replace(
        base,
        partitions=(replace(base.partitions[0], rejection_sources=(source,)),),
    )

    result = publish_publication(connection, root, package)

    assert connection.execute(
        """select first_selection_run_id, source_revision, facts_sha256,
                  classifier_algo_version, created_at_utc
             from strategy_rejection_sources where strategy_id = 1"""
    ).fetchone() == (
        "overlay-run", "auto-revision", "auto-facts", "M3", result.normalized_latest,
    )


def test_auto_non_rejected_review_does_not_project_rejection_source(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    base = _package()
    source = {
        "strategy_id": 1,
        "source_kind": "EQUITY_REGIME_FILTER",
        "reason_code": "W28_DOWN",
        "first_result_id": 101,
        "classifier_algo_version": "M3",
        "source_revision": "ignored-revision",
        "facts_sha256": "ignored-facts",
    }
    package = replace(
        base,
        partitions=(replace(
            base.partitions[0],
            reviews=(PublicationReview(1, "FINALIST", 1, None, "not rejected"),),
            rejection_sources=(source,),
        ),),
    )

    publish_publication(connection, root, package)

    assert connection.execute("select count(*) from strategy_rejection_sources").fetchone() == (0,)


def test_publication_key_same_digest_is_zero_write_and_different_digest_conflicts(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    package = _package()

    first = publish_publication(connection, root, package)
    counts_before = connection.execute(
        "select (select count(*) from selection_publications), "
        "(select count(*) from selection_review_imports), "
        "(select count(*) from strategy_tags)"
    ).fetchone()
    retry = publish_publication(connection, root, package)

    assert first.publication_id == retry.publication_id
    assert retry.code == "ALREADY_COMMITTED_REEXPORT_REQUIRED"
    assert connection.execute(
        "select (select count(*) from selection_publications), "
        "(select count(*) from selection_review_imports), "
        "(select count(*) from strategy_tags)"
    ).fetchone() == counts_before

    with pytest.raises(PublicationConflict) as error:
        publish_publication(connection, root, _package(operation_digest="digest-2"))
    assert error.value.code == "OPERATION_KEY_CONFLICT"


def test_aggregate_publication_uses_same_timestamp_and_duplicate_upload_is_idempotent(
    tmp_path: Path,
) -> None:
    connection, root = _seed_database(tmp_path)
    aggregate = PublicationAggregate(
        aggregate_import_id="aggregate-1",
        uploaded_workbook_sha256="uploaded-1",
        partition_rowsets_json="{}",
        candidate_identities_json="{}",
    )
    package = replace(_package(), publication_kind="AGGREGATE_REVIEW", aggregate=aggregate)
    first = publish_publication(connection, root, package)

    timestamps = connection.execute(
        """select created_at_utc from selection_publications
           union all select imported_at_utc from selection_aggregate_imports
           union all select imported_at_utc from selection_review_imports
           union all select updated_at_utc from strategy_tags"""
    ).fetchall()
    assert len(timestamps) == 4
    assert len({row[0] for row in timestamps}) == 1
    assert connection.execute(
        "select aggregate_import_id from selection_review_imports"
    ).fetchone() == ("aggregate-1",)

    retry = publish_publication(connection, root, package)
    assert first.aggregate_import_id == retry.aggregate_import_id == "aggregate-1"
    assert retry.code == "ALREADY_IMPORTED"
    conflicting_upload = replace(
        package,
        aggregate=replace(aggregate, uploaded_workbook_sha256="uploaded-2"),
    )
    with pytest.raises(PublicationConflict) as error:
        publish_publication(connection, root, conflicting_upload)
    assert error.value.code == "UPLOAD_DIGEST_CONFLICT"


@pytest.mark.parametrize("publication_kind, aggregate", [
    ("UNKNOWN", None),
    ("AGGREGATE_REVIEW", None),
    ("AUTO_REJECTION_OVERLAY", PublicationAggregate("mixed", "upload", "{}", "{}")),
])
def test_publication_kind_invariants_reject_unknown_and_mixed_without_writes(
    tmp_path: Path, publication_kind: str, aggregate: PublicationAggregate | None,
) -> None:
    connection, root = _seed_database(tmp_path)
    package = replace(_package(), publication_kind=publication_kind, aggregate=aggregate)

    with pytest.raises(PublicationValidationError) as error:
        publish_publication(connection, root, package)

    assert error.value.code == "INVALID_ARGUMENT"
    assert connection.execute(
        "select (select count(*) from selection_publications), "
        "(select count(*) from selection_aggregate_imports), "
        "(select count(*) from selection_runs)"
    ).fetchone() == (0, 0, 1)


def test_rejection_sources_without_reviews_are_typed_and_zero_write(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    base = _package()
    source = {
        "strategy_id": 1,
        "source_kind": "EQUITY_REGIME_FILTER",
        "reason_code": "W28_DOWN",
        "first_result_id": 101,
        "classifier_algo_version": "M3",
        "source_revision": "orphan-revision",
        "facts_sha256": "orphan-facts",
    }
    package = replace(
        base,
        partitions=(replace(base.partitions[0], reviews=(), rejection_sources=(source,)),),
    )

    with pytest.raises(PublicationValidationError) as error:
        publish_publication(connection, root, package)

    assert error.value.code == "INVALID_REJECTION_SOURCE_SCOPE"
    assert connection.execute(
        "select (select count(*) from selection_publications), "
        "(select count(*) from strategy_rejection_sources)"
    ).fetchone() == (0, 0)


@pytest.mark.parametrize("replacement_source", [False, True])
def test_explicit_rejection_preserves_prior_provenance_and_replaces_only_tag(
    tmp_path: Path, replacement_source: bool,
) -> None:
    connection, root = _seed_database(tmp_path)
    connection.execute(
        """insert into strategy_tags (
            strategy_id, tag, source, source_ref, updated_at_utc
        ) values (1, 'REJECTED', 'EQUITY_REGIME_FILTER', 'equity-old', current_timestamp)"""
    )
    connection.execute(
        """insert into strategy_rejection_sources (
            strategy_id, source_kind, reason_code, first_result_id,
            first_selection_run_id, classifier_algo_version, source_revision,
            facts_sha256, created_at_utc
        ) values (1, 'EQUITY_REGIME_FILTER', 'W28_DOWN', 101, 'source-run', 'M3',
                  'equity-old', 'facts-old', current_timestamp)"""
    )
    connection.commit()
    base = _package()
    replacement = {
        "strategy_id": 1,
        "source_kind": "EQUITY_REGIME_FILTER",
        "reason_code": "W28_DOWN",
        "first_result_id": 101,
        "classifier_algo_version": "M3",
        "source_revision": "equity-new",
        "facts_sha256": "facts-new",
    }
    package = replace(
        base,
        publication_kind="AGGREGATE_REVIEW",
        aggregate=PublicationAggregate("aggregate-provenance", "upload-provenance", "{}", "{}"),
        partitions=(replace(
            base.partitions[0], overlay_run=None,
            reviews=(PublicationReview(1, "REJECTED", None, None, "explicit"),),
            rejection_sources=(replacement,) if replacement_source else (),
        ),),
    )

    publish_publication(connection, root, package)

    tag = connection.execute(
        "select tag, source, source_ref from strategy_tags where strategy_id = 1"
    ).fetchone()
    assert tag is not None
    assert tag[:2] == ("REJECTED", "SELECTION_REVIEW")
    assert tag[2]
    sources = connection.execute(
        """select reason_code, source_revision, facts_sha256
             from strategy_rejection_sources where strategy_id = 1
             order by source_revision"""
    ).fetchall()
    expected = [("W28_DOWN", "equity-old", "facts-old")]
    assert sources == expected


@pytest.mark.parametrize("current_result_id", [None, 999])
def test_overlay_stale_or_null_current_result_is_typed_and_zero_write(
    tmp_path: Path, current_result_id: int | None,
) -> None:
    connection, root = _seed_database(tmp_path)
    connection.execute("update strategies set current_result_id = ? where strategy_id = 1", [current_result_id])
    connection.commit()

    with pytest.raises(PublicationValidationError) as error:
        publish_publication(connection, root, _package())

    assert error.value.code == "STALE_RESULTS"
    assert connection.execute(
        "select (select count(*) from selection_publications), "
        "(select count(*) from selection_runs), "
        "(select count(*) from selection_review_imports)"
    ).fetchone() == (0, 1, 0)


def test_source_review_row_outside_frozen_run_is_typed_and_zero_write(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    base = _package()
    package = replace(
        base,
        publication_kind="AGGREGATE_REVIEW",
        aggregate=PublicationAggregate("aggregate-outside-source", "upload-outside-source", "{}", "{}"),
        partitions=(replace(
            base.partitions[0], overlay_run=None,
            reviews=(PublicationReview(999, "REJECTED", None, None, "outside"),),
        ),),
    )

    with pytest.raises(PublicationValidationError) as error:
        publish_publication(connection, root, package)

    assert error.value.code == "REVIEW_ROW_OUTSIDE_RUN"
    assert connection.execute("select count(*) from selection_publications").fetchone() == (0,)


def test_overlay_review_row_outside_frozen_overlay_is_typed_and_zero_write(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    base = _package()
    package = replace(
        base,
        partitions=(replace(
            base.partitions[0],
            reviews=(PublicationReview(999, "REJECTED", None, None, "outside"),),
        ),),
    )

    with pytest.raises(PublicationValidationError) as error:
        publish_publication(connection, root, package)

    assert error.value.code == "REVIEW_ROW_OUTSIDE_OVERLAY"
    assert connection.execute("select count(*) from selection_publications").fetchone() == (0,)


def test_duplicate_strategy_across_partitions_is_rejected_before_writes(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    base = _package()
    package = replace(
        base,
        publication_kind="AGGREGATE_REVIEW",
        aggregate=PublicationAggregate("aggregate-duplicate-strategy", "upload-duplicate-strategy", "{}", "{}"),
        partitions=base.partitions + (
            PublicationPartition(
                pair="ETHUSDT", side="LONG", source_run_id=None, overlay_run=None,
                reviews=(PublicationReview(1, "REJECTED", None, None, "duplicate"),),
            ),
        ),
    )

    with pytest.raises(PublicationValidationError) as error:
        publish_publication(connection, root, package)

    assert error.value.code == "DUPLICATE_STRATEGY"
    assert connection.execute("select count(*) from selection_publications").fetchone() == (0,)


def test_normalized_latest_includes_selection_run_creation(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    high = datetime(2050, 1, 1, 12, 0, 0, 223344, tzinfo=UTC)
    connection.execute(
        """insert into selection_runs values (
            'source-high', (select value from schema_info where key = 'database_instance_id'),
            'ETHUSDT', 'LONG', 'selection-v1', '{}', 'request-high', '{}', 'config-high',
            0, 0, 0, 1, 'source-workbook-high', ?
        )""",
        [high],
    )
    connection.commit()

    result = publish_publication(connection, root, _package())

    assert result.normalized_latest == high.replace(microsecond=223345)
    assert connection.execute(
        "select created_at_utc from selection_runs where selection_run_id = 'overlay-run'"
    ).fetchone() == (result.normalized_latest,)


@pytest.mark.parametrize("status", [None, "FINALIST"])
def test_aggregate_non_rejection_review_preserves_equity_rejection_projection(
    tmp_path: Path, status: str | None,
) -> None:
    connection, root = _seed_database(tmp_path)
    connection.execute(
        """insert into strategy_tags (
            strategy_id, tag, source, source_ref, updated_at_utc
        ) values (1, 'REJECTED', 'EQUITY_REGIME_FILTER', 'equity-source', current_timestamp)"""
    )
    connection.execute(
        """insert into strategy_rejection_sources (
            strategy_id, source_kind, reason_code, first_result_id,
            first_selection_run_id, classifier_algo_version, source_revision,
            facts_sha256, created_at_utc
        ) values (1, 'EQUITY_REGIME_FILTER', 'W28_DOWN', 101, 'source-run', 'M3',
                  'equity-revision', 'facts', current_timestamp)"""
    )
    connection.commit()
    base = _package()
    aggregate = PublicationAggregate("aggregate-preserve", "upload-preserve", "{}", "{}")
    package = replace(
        base,
        publication_kind="AGGREGATE_REVIEW",
        aggregate=aggregate,
        partitions=(replace(
            base.partitions[0], overlay_run=None,
            reviews=(PublicationReview(1, status, 1 if status == "FINALIST" else None, None, "keep"),),
            rejection_sources=(),
        ),),
    )

    publish_publication(connection, root, package)

    assert connection.execute(
        "select source, source_ref from strategy_tags where strategy_id = 1 and tag = 'REJECTED'"
    ).fetchone() == ("EQUITY_REGIME_FILTER", "equity-source")
    assert connection.execute(
        "select source_kind, reason_code from strategy_rejection_sources where strategy_id = 1"
    ).fetchone() == ("EQUITY_REGIME_FILTER", "W28_DOWN")


def test_aggregate_existing_operation_compares_id_and_lifecycle(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    aggregate = PublicationAggregate("aggregate-existing", "upload-existing", "{}", "{}")
    package = replace(_package(), publication_kind="AGGREGATE_REVIEW", aggregate=aggregate)
    publish_publication(connection, root, package)

    with pytest.raises(PublicationConflict) as aggregate_id:
        publish_publication(
            connection,
            root,
            replace(package, aggregate=replace(aggregate, aggregate_import_id="aggregate-other")),
        )
    assert aggregate_id.value.code == "AGGREGATE_ID_CONFLICT"

    connection.execute(
        "update selection_aggregate_imports set lifecycle_status = 'RETIRED' where aggregate_import_id = 'aggregate-existing'"
    )
    connection.commit()
    with pytest.raises(PublicationNotFound) as expired:
        publish_publication(connection, root, package)
    assert expired.value.code == "AGGREGATE_NOT_FOUND"


def test_normalized_latest_bumps_global_aggregate_timestamp(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    aggregate = PublicationAggregate("aggregate-clock", "upload-clock", "{}", "{}")
    first = replace(_package(), publication_kind="AGGREGATE_REVIEW", aggregate=aggregate)
    publish_publication(connection, root, first)
    high = datetime(2050, 1, 1, 12, 0, 0, 123456, tzinfo=UTC)
    connection.execute(
        "update selection_aggregate_imports set imported_at_utc = ? where aggregate_import_id = 'aggregate-clock'",
        [high],
    )
    connection.commit()

    second_base = _package(operation_key="operation-clock-2", operation_digest="digest-clock-2")
    second = replace(
        second_base,
        publication_id="publication-clock-2",
        publication_kind="AGGREGATE_REVIEW",
        aggregate=PublicationAggregate("aggregate-clock-2", "upload-clock-2", "{}", "{}"),
        partitions=(replace(
            second_base.partitions[0],
            overlay_run=replace(second_base.partitions[0].overlay_run, selection_run_id="overlay-clock-2"),
        ),),
    )
    result = publish_publication(connection, root, second)

    assert result.normalized_latest == high.replace(microsecond=123457)
    assert connection.execute(
        "select created_at_utc from selection_publications where publication_id = 'publication-clock-2'"
    ).fetchone() == (result.normalized_latest,)
    assert connection.execute(
        "select imported_at_utc from selection_review_imports where review_import_id = ?",
        [result.review_import_ids[0]],
    ).fetchone() == (result.normalized_latest,)


def test_lowercase_partition_identity_is_rejected(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    partition = replace(_package().partitions[0], pair="btcusdt", side="long")
    with pytest.raises(PublicationValidationError) as error:
        publish_publication(connection, root, replace(_package(), partitions=(partition,)))
    assert error.value.code == "INVALID_ARGUMENT"


def test_aggregate_overlay_requires_frozen_source_run(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    base = _package()
    package = replace(
        base,
        publication_kind="AGGREGATE_REVIEW",
        aggregate=PublicationAggregate("aggregate-overlay", "upload-overlay", "{}", "{}"),
        partitions=(replace(base.partitions[0], source_run_id=None, reviews=()),),
    )
    with pytest.raises(PublicationValidationError) as error:
        publish_publication(connection, root, package)
    assert error.value.code == "OVERLAY_SOURCE_REQUIRED"


def test_mapping_parser_rejects_malformed_integer_without_coercion() -> None:
    raw = asdict(_package())
    raw["partitions"][0]["overlay_run"]["candidate_count"] = "1"
    with pytest.raises(PublicationValidationError) as error:
        PublicationPackage.from_mapping(raw)
    assert error.value.code == "INVALID_ARGUMENT"


def test_lookup_requires_v11_guard(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    connection, root = _seed_database(tmp_path)
    calls = 0
    original = publication_module.require_performance_v2

    def guarded(db: duckdb.DuckDBPyConnection) -> None:
        nonlocal calls
        calls += 1
        original(db)

    monkeypatch.setattr(publication_module, "require_performance_v2", guarded)
    with pytest.raises(PublicationNotFound):
        publication_module.lookup_publication(connection, "missing")
    assert calls == 1


def test_publication_uses_one_writer_lock_for_all_partitions(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    connection, root = _seed_database(tmp_path)
    events: list[object] = []

    class FakeLock:
        def __init__(self, lock_root: Path) -> None:
            events.append(("init", lock_root))

        def __enter__(self):
            events.append("enter")
            return self

        def __exit__(self, *args: object) -> None:
            events.append("exit")

    monkeypatch.setattr(publication_module, "PerformanceV2WriterLock", FakeLock)
    publish_publication(connection, root, _package())

    assert events == [("init", root), "enter", "exit"]


def test_unique_operation_race_rechecks_committed_row(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    connection, root = _seed_database(tmp_path)
    package = _package()
    publish_publication(connection, root, package)
    original_lookup = publication_module._lookup_operation
    hidden = True

    def hide_first_lookup(db: duckdb.DuckDBPyConnection, operation_key: str):
        nonlocal hidden
        if hidden:
            hidden = False
            return None
        return original_lookup(db, operation_key)

    def simulate_race(*args, **kwargs):
        raise duckdb.ConstraintException("simulated unique operation race")

    monkeypatch.setattr(publication_module, "_lookup_operation", hide_first_lookup)
    monkeypatch.setattr(publication_module, "_revalidate_source_runs", lambda *args, **kwargs: None)
    monkeypatch.setattr(publication_module, "_insert_publication_header", simulate_race)
    result = publish_publication(connection, root, package)
    assert result.code == "ALREADY_COMMITTED_REEXPORT_REQUIRED"


def test_automatic_effective_rejection_is_a_projection_noop(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    connection.execute(
        """insert into strategy_tags (
            strategy_id, tag, source, source_ref, updated_at_utc
        ) values (1, 'REJECTED', 'EQUITY_REGIME_FILTER', 'existing', current_timestamp)"""
    )
    connection.commit()

    result = publish_publication(connection, root, _package())

    assert result.review_import_ids == ()
    assert connection.execute("select count(*) from selection_review_imports").fetchone() == (0,)
    assert connection.execute(
        "select source, source_ref from strategy_tags where strategy_id = 1 and tag = 'REJECTED'"
    ).fetchone() == ("EQUITY_REGIME_FILTER", "existing")


def test_unknown_and_retired_publications_are_typed_not_found(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    with pytest.raises(PublicationNotFound) as unknown:
        publication_module.lookup_publication(connection, "missing")
    assert unknown.value.code == "PUBLICATION_KEY_NOT_FOUND"

    publish_publication(connection, root, _package())
    connection.execute(
        "update selection_publications set retired_at_utc = current_timestamp where publication_id = 'publication-1'"
    )
    connection.commit()
    with pytest.raises(PublicationNotFound) as retired:
        publication_module.lookup_publication(connection, "operation-1")
    assert retired.value.code == "PUBLICATION_KEY_NOT_FOUND"
    with pytest.raises(PublicationNotFound) as retry:
        publish_publication(connection, root, _package())
    assert retry.value.code == "PUBLICATION_KEY_NOT_FOUND"


def test_publication_rolls_back_all_partitions_on_late_validation_failure(tmp_path: Path) -> None:
    connection, root = _seed_database(tmp_path)
    package = _package()
    bad_overlay = replace(
        package.partitions[0].overlay_run,
        selection_run_id="overlay-eth",
        symbol="ETHUSDT",
        results=(),
    )
    bad_partition = PublicationPartition(
        pair="ETHUSDT",
        side="LONG",
        source_run_id="missing-source-run",
        overlay_run=bad_overlay,
        reviews=(),
    )
    invalid = replace(package, partitions=package.partitions + (bad_partition,))

    with pytest.raises(PublicationNotFound) as error:
        publish_publication(connection, root, invalid)
    assert error.value.code == "SOURCE_RUN_NOT_FOUND"
    assert connection.execute("select count(*) from selection_publications").fetchone() == (0,)
    assert connection.execute("select count(*) from selection_review_imports").fetchone() == (0,)
    assert connection.execute("select count(*) from strategy_tags").fetchone() == (0,)


def test_publication_rolls_back_header_and_first_partition_when_late_projection_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    connection, root = _seed_database(tmp_path)
    now = datetime(2026, 10, 10, tzinfo=UTC)
    connection.execute(
        """insert into selection_runs values (
            'source-run-2', (select value from schema_info where key = 'database_instance_id'),
            'ETHUSDT', 'LONG', 'selection-v1', '{}', 'request', '{}', 'config',
            0, 0, 0, 1, 'source-workbook-2', ?
        )""",
        [now],
    )
    connection.execute(
        """insert into selection_results (
            selection_run_id, strategy_id, result_id_at_selection, auto_status,
            auto_score, auto_rank, auto_reason, analog_group_key,
            auto_analog_of_strategy_id, prior_rejected, stage_trace_json
        ) values ('source-run-2', 1, 101, 'FINALIST', 1, 1, null, null, null, false, '{}')"""
    )
    package = _package()
    failing = replace(
        package,
        publication_kind="AGGREGATE_REVIEW",
        aggregate=PublicationAggregate(
            aggregate_import_id="aggregate-late",
            uploaded_workbook_sha256="uploaded-late",
            partition_rowsets_json="{}",
            candidate_identities_json="{}",
        ),
        partitions=package.partitions + (
            PublicationPartition(
                pair="ETHUSDT",
                side="LONG",
                source_run_id="source-run-2",
                overlay_run=None,
                reviews=(),
                rejection_sources=(),
            ),
        ),
    )

    original = publication_module._write_review_projection
    calls = 0

    def fail_after_second_projection(*args, **kwargs):
        nonlocal calls
        calls += 1
        result = original(*args, **kwargs)
        if calls == 2:
            raise RuntimeError("forced projection failure")
        return result

    monkeypatch.setattr(publication_module, "_write_review_projection", fail_after_second_projection)
    with pytest.raises(RuntimeError, match="forced projection failure"):
        publish_publication(connection, root, failing)
    assert connection.execute("select count(*) from selection_publications").fetchone() == (0,)
    assert connection.execute("select count(*) from selection_publication_runs").fetchone() == (0,)
    assert connection.execute("select count(*) from selection_review_imports").fetchone() == (0,)
    assert connection.execute("select count(*) from strategy_tags").fetchone() == (0,)


def test_review_import_key_is_domain_separated_and_stable() -> None:
    expected = sha256(
        b"performance_v2_selection_review_import_v11\0publication_id\0pub\0run"
    ).hexdigest()
    assert review_import_key("publication_id", "PUB", "RUN") == expected
    assert review_import_key("aggregate_import_id", "pub", "run") != expected
