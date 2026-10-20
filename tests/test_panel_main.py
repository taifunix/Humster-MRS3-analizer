"""The «Пары и стратегии» main screen service (Panel rework stage 3).

Contract: docs/specs/2026-10-18-panel-simulation-workflow.md, «Stage 3 contract».
"""

from __future__ import annotations

from datetime import date, datetime, timezone
from http.client import HTTPConnection
from io import BytesIO
import json
from pathlib import Path
import subprocess

from openpyxl import load_workbook
import pytest

from mrs3.pair_registry.pipeline import compute_pipeline
from mrs3.pair_registry.refresh import refresh_exchange
from mrs3.pair_registry.store import connect
from mrs3.panel_main import MainScreenError, MainScreenService, filter_rows, sort_rows
from mrs3.ready_store import build_ready
from tests.test_panel_store_ready import served  # noqa: F401  (route-test fixture)
from tests.test_ready_store import REPO, _store_with_periods

NOW = 1_760_000_000_000
TEMPLATE = REPO / "templates" / "exchanges" / "Bybit" / "mrs3-strategy.json"


class FakeRunner:
    def __init__(self) -> None:
        self.started: list[tuple] = []

    def start(self, kind, label, argv, *, total=None, stop_file=None):
        self.started.append((kind, label, list(argv), total, stop_file))
        return {"task": {"state": "running", "kind": kind}}

    def stop(self):
        return {"task": {"state": "stopped"}}

    def status(self):
        return {"task": {"state": "idle"}}


def _world(tmp_path: Path):
    store = _store_with_periods(tmp_path / "store", "GOODUSDT", 3, 35.0)
    _store_with_periods(tmp_path / "store", "YOUNGUSDT", 1, 30.0)
    template = json.loads(TEMPLATE.read_text(encoding="utf-8"))
    build_ready(store, ["GOODUSDT", "YOUNGUSDT"], store.root / "ready" / "2026-10-18_simulation_120000", template)
    listing = int(datetime(2026, 6, 1, tzinfo=timezone.utc).timestamp() * 1000)
    con = connect(tmp_path / "registry")
    refresh_exchange(con, "bybit", [{"symbol": s, "status": "TRADING", "api_symbol": s, "listing_ms": listing,
                                     "tick_size": "0.01", "qty_step": "0.01", "display_name": s}
                                    for s in ("GOODUSDT", "YOUNGUSDT", "NEWUSDT")],
                     [{"symbol": "GOODUSDT", "tier_id": 1, "risk_limit_value": "10000", "initial_margin": "0.02",
                       "maintenance_margin": "0.01", "max_leverage": "50", "mm_deduction": "0", "is_lowest": True}],
                     now_ms=NOW)
    compute_pipeline(con, store, "bybit", date(2026, 10, 18), "pc1", now_ms=NOW, track=False)
    con.close()
    config = tmp_path / "config.local.json"
    config.write_text(json.dumps({"pair_registry": {"root": str(tmp_path / "registry")},
                                  "trade_store": {"root": str(store.root)}}), encoding="utf-8")
    runner = FakeRunner()
    service = MainScreenService(REPO, config, runner, python="py", today=lambda: date(2026, 10, 19),
                                chain_root=tmp_path)
    return service, runner, store


def test_views_list_the_registry_and_the_pipeline(tmp_path: Path) -> None:
    service, _runner, _store = _world(tmp_path)
    pairs = service.view("pairs", "bybit")
    assert [row["symbol"] for row in pairs["rows"]] == ["GOODUSDT", "NEWUSDT", "YOUNGUSDT"]
    good = pairs["rows"][0]
    assert (good["stage"], good["ready_status"], good["periods"], good["history_status"]) == (
        "READY", "READY", 3, "READY возможен")
    assert good["ready_source"] == "simulator" and pairs["rows"][1]["ready_source"] is None
    reference = service.view("reference", "bybit")["rows"][0]
    assert (reference["symbol"], reference["risk_tiers"], reference["listing_day"]) == ("GOODUSDT", 1, "2026-06-01")
    assert service.view("changes", "bybit")["rows"][0]["kind"] == "listed"
    assert service.view("liquidity", "bybit")["rows"] == []
    assert service.view("pairs", "bingx")["rows"] == []
    with pytest.raises(MainScreenError):
        service.view("unknown", "bybit")
    with pytest.raises(MainScreenError):
        service.view("pairs", "binance")


def test_filter_and_sort_match_the_table_rules() -> None:
    rows = [{"a": "XOM", "n": 10}, {"a": "aapl", "n": 2}, {"a": "Nvda", "n": None}]
    assert [row["a"] for row in filter_rows(rows, ["a"], "a")] == ["aapl", "Nvda"]
    assert [row["a"] for row in sort_rows(rows, "n")] == ["aapl", "XOM", "Nvda"]                # empty last
    assert [row["a"] for row in sort_rows(rows, "n", descending=True)] == ["XOM", "aapl", "Nvda"]
    assert [row["a"] for row in sort_rows(rows, "a")] == ["aapl", "Nvda", "XOM"]


def test_excel_export_has_exactly_the_shown_rows_and_columns(tmp_path: Path) -> None:
    service, _runner, _store = _world(tmp_path)
    name, data = service.view_xlsx("pairs", "bybit", ["symbol", "stage"], text="good", sort="symbol")
    assert name == "pairs_bybit_2026-10-19.xlsx"
    sheet = load_workbook(BytesIO(data)).active
    assert [[cell.value for cell in row] for row in sheet.iter_rows()] == [["Пара", "Стадия"], ["GOODUSDT", "READY"]]


def test_column_choice_is_saved_on_the_server(tmp_path: Path) -> None:
    service, _runner, _store = _world(tmp_path)
    assert service.view_settings()["pairs"] == {}
    service.save_view_settings({"view": "pairs", "columns": ["symbol", "stage", "symbol"]})
    assert service.view_settings()["pairs"] == {"columns": ["symbol", "stage"]}
    for bad in ({"view": "pairs", "columns": []}, {"view": "pairs", "columns": ["nope"]}, {"view": "x", "columns": ["a"]}):
        with pytest.raises(MainScreenError):
            service.save_view_settings(bad)


def test_pair_card_and_entries(tmp_path: Path) -> None:
    service, _runner, _store = _world(tmp_path)
    card = service.pair("bybit", "goodusdt")
    assert card["symbol"] == "GOODUSDT" and card["pipeline"]["stage"] == "READY"
    assert len(card["periods"]) == 3 and card["ready"]["status"] == "READY"
    assert card["risk_tiers"][0]["max_leverage"] == "50"
    admit = next(item for item in card["thresholds"] if item["parameter"] == "admit_score")
    assert (admit["default"], admit["value"], admit["overridden"]) == (30.0, 30.0, False)
    service.decision({"exchange": "bybit", "symbol": "GOODUSDT", "side": "ALL", "decision": "WATCH",
                      "reason": "проверить ликвидность", "author": "user"})
    service.setting({"exchange": "bybit", "symbol": "GOODUSDT", "side": "ALL", "parameter": "admit_score",
                     "value": "25", "reason": "волатильная", "author": "user"})
    card = service.pair("bybit", "GOODUSDT")
    assert card["decisions"][0]["decision"] == "WATCH"
    assert next(item for item in card["thresholds"] if item["parameter"] == "admit_score")["value"] == 25.0
    with pytest.raises(MainScreenError):
        service.setting({"exchange": "bybit", "symbol": "GOODUSDT", "parameter": "pool", "value": 99,
                         "reason": "r", "author": "u"})
    with pytest.raises(MainScreenError):
        service.pair("bybit", "MISSINGUSDT")


def test_a_panel_write_exports_the_registry_to_the_sync_repository(tmp_path: Path) -> None:
    from mrs3.pair_registry.sync import init_repo

    service, _runner, _store = _world(tmp_path)
    entry = {"exchange": "bybit", "symbol": "GOODUSDT", "side": "ALL", "decision": "WATCH", "reason": "r",
             "author": "user"}
    assert service.decision(entry)["sync"] == {"enabled": False}            # no sync repository: sync is off
    config = json.loads(service.config_path.read_text(encoding="utf-8"))
    config["pair_registry"]["machine"] = "pc1"
    service.config_path.write_text(json.dumps(config), encoding="utf-8")
    repo = tmp_path / "registry" / "sync"
    init_repo(repo, "pc1")
    result = service.setting({"exchange": "bybit", "symbol": "GOODUSDT", "side": "ALL", "parameter": "admit_score",
                              "value": "25", "reason": "r", "author": "user"})
    assert result["sync"] == {"enabled": True, "queued": True}
    service._push_thread.join(timeout=60)                                  # files and commit run in the background
    decisions = (repo / "decisions" / "pc1.jsonl").read_text(encoding="utf-8").strip().splitlines()
    assert len(decisions) == 2                                             # the earlier decision is exported too
    rows = (repo / "pair_settings" / "pc1.jsonl").read_text(encoding="utf-8").strip().splitlines()
    assert len(rows) == 2 and json.loads(rows[1])[1] == "pc1"
    log = subprocess.run(["git", "-C", str(repo), "log", "--format=%s"], capture_output=True, text=True, check=True)
    assert log.stdout.startswith("panel setting (pc1)")


def test_plans_for_history_ready_tester_and_tick(tmp_path: Path, monkeypatch) -> None:
    service, _runner, _store = _world(tmp_path)
    monkeypatch.setattr(service, "_history_end_day", lambda exchange="bybit": date(2026, 9, 10))
    history = {item["key"]: item for item in service.plan("history", "bybit")["items"]}
    assert history["NEWUSDT"]["detail"].startswith("новая пара") and history["NEWUSDT"]["checked"]
    ready = {item["key"]: item for item in service.plan("ready", "bybit")["items"]}
    assert (ready["GOODUSDT"]["detail"], ready["GOODUSDT"]["checked"]) == ("без изменений", False)
    assert ready["YOUNGUSDT"]["detail"].startswith("мало периодов")
    assert ready["NEWUSDT"]["detail"].startswith("мало периодов")
    from mrs3.panel_main import _periods
    assert [_periods(n) for n in (1, 2, 5, 11, 22)] == ["1 новый период", "2 новых периода", "5 новых периодов",
                                                        "11 новых периодов", "22 новых периода"]
    assert _periods(2, new=False) == "2 периода" and _periods(1, new=False) == "1 период"
    tester = service.plan("tester", "bybit")["items"]
    assert {item["key"] for item in tester} == {"GOODUSDT", "YOUNGUSDT"} and not any(item["checked"] for item in tester)
    tick = service.plan("tick", "bybit")
    assert tick["ready"] == "2026-10-18_simulation_120000" and len(tick["items"]) == 1 and tick["items"][0]["checked"]
    monkeypatch.setattr(service, "_history_end_day", lambda exchange="bybit": date(2026, 9, 10) if exchange == "bybit"
                        else (_ for _ in ()).throw(AssertionError("BingX end day expected")))
    with pytest.raises(AssertionError, match="BingX end day"):                       # BingX takes its own end day
        service.plan("history", "bingx")
    assert "disabled" not in service.plan("tester", "bingx")                        # 5c: tester sampling
    assert "disabled" not in service.plan("tick", "bingx")                          # 5e-3a: the BingX tick test
    assert "disabled" not in service.plan("retest", "bingx")                        # 5e-3b: the BingX retest
    assert "disabled" not in service.plan("ready", "bingx")                         # 5b: history and READY
    assert service.plan("retest", "bybit")["options"] == {"include_reserve": False}


def test_runs_start_one_chain_task(tmp_path: Path) -> None:
    service, runner, store = _world(tmp_path)
    service.run("sync", "bybit", [])
    kind, label, argv, total, stop_file = runner.started[-1]
    assert (kind, label, total, stop_file) == ("sync", "Синхронизировать", 1, None)
    steps = json.loads(Path(argv[2]).read_text(encoding="utf-8"))["steps"]
    assert steps[0]["argv"][1:4] == ["-m", "mrs3.pair_registry", "sync"]
    service.run("ready", "bybit", ["goodusdt"])
    steps = json.loads(Path(runner.started[-1][2][2]).read_text(encoding="utf-8"))["steps"]
    assert [step["name"] for step in steps] == ["READY", "Конвейер"]
    assert steps[0]["argv"][steps[0]["argv"].index("--pairs") + 1] == "GOODUSDT"
    assert steps[0]["argv"][steps[0]["argv"].index("--latest-source") + 1] == "auto"      # tester results when present
    service.run("tester", "bybit", ["GOODUSDT"])
    kind, _label, argv, total, stop_file = runner.started[-1]
    assert total == 4 and stop_file == store.root / "jobs" / ".stop" and argv[-2:] == ["--stop-file", str(stop_file)]
    assert argv[3] == "--cleanup" and Path(argv[2]).parent.parent == tmp_path
    with pytest.raises(MainScreenError):
        service.run("history", "bybit", [])
    service.run("tester", "bingx", ["AMDUSDT"])                                     # Panel stage 5c
    steps = json.loads(Path(runner.started[-1][2][2]).read_text(encoding="utf-8"))["steps"]
    assert [step["name"] for step in steps] == ["plan", "run", "import", "Конвейер"]
    assert all(step["argv"][-2:] == ["--exchange", "bingx"] for step in steps)
    service.run("history", "bingx", ["amdusdt"])                                    # Panel stage 5b
    steps = json.loads(Path(runner.started[-1][2][2]).read_text(encoding="utf-8"))["steps"]
    assert [step["name"] for step in steps] == ["История", "Конвейер"]
    assert all(step["argv"][-2:] == ["--exchange", "bingx"] for step in steps)
    service.run("ready", "bingx", ["AMDUSDT"])
    steps = json.loads(Path(runner.started[-1][2][2]).read_text(encoding="utf-8"))["steps"]
    assert all(step["argv"][-2:] == ["--exchange", "bingx"] for step in steps)
    with pytest.raises(MainScreenError):
        service.run("explode", "bybit", [])


def test_tick_and_retest_go_through_the_controller(tmp_path: Path) -> None:
    service, _runner, _store = _world(tmp_path)
    calls = []

    class Controller:
        def store_ready_prepare(self, payload, *, strategies=None):
            calls.append(("prepare", payload, strategies))
            return {"ready": payload["ready"], "analysis_run_id": "a" * 64, "strategy_count": len(strategies)}

        def panel_job_submit(self, payload):
            calls.append(("job", payload))
            return {"job_id": "j1"}

    service._controller = Controller()
    strategy = service.plan("tick", "bybit")["items"][0]["key"]
    result = service.run("tick", "bybit", [strategy], {"start_date": "2026-09-01", "end_date": "2026-10-01"})
    assert result["tick"]["job"] == {"job_id": "j1"} and calls[0][2] == [strategy]
    assert calls[1][1]["request"]["start_date"] == "2026-09-01"
    with pytest.raises(MainScreenError):
        service.run("tick", "bybit", [strategy], {"start_date": "2026-10-05", "end_date": "2026-10-01"})
    service.run("retest", "bybit", [], {"include_reserve": True})
    assert calls[-1][1] == {"kind": "strategies.performance.v2.finalist-retest.start", "request": {"include_reserve": True}}
    service.run("retest", "bingx", [], {"include_reserve": False})                 # Panel stage 5e-3b
    assert calls[-1][1]["request"] == {"include_reserve": False, "exchange": "bingx"}


def test_views_report_a_missing_registry_clearly(tmp_path: Path) -> None:
    config = tmp_path / "config.local.json"
    config.write_text(json.dumps({"pair_registry": {"root": str(tmp_path / "none")}}), encoding="utf-8")
    service = MainScreenService(REPO, config, FakeRunner())
    with pytest.raises(MainScreenError, match="Синхронизировать"):
        service.view("pairs", "bybit")


# -- routes (Panel rework stage 3) ------------------------------------------------------------------------------

def _served_world(tmp_path: Path, served):
    from mrs3.panel import PanelController

    service, runner, store = _world(tmp_path)
    controller = PanelController(tmp_path, service.config_path)
    service._controller = controller
    controller._main_screen_service = service
    return served(controller), controller, service, runner


def test_routes_serve_views_excel_settings_and_the_pair_card(tmp_path: Path, served) -> None:
    from tests.test_panel_store_ready import _call

    server, _controller, _service, _runner = _served_world(tmp_path, served)
    status, view = _call(server, "GET", "/api/v2/main/view?name=pairs&exchange=bybit")
    assert status == 200 and view["rows"][0]["symbol"] == "GOODUSDT"
    status, answer = _call(server, "GET", "/api/v2/main/view?name=nope&exchange=bybit")
    assert status == 400 and "nope" in answer["error"]
    status, settings = _call(server, "POST", "/api/v2/main/view-settings", {"view": "pairs", "columns": ["symbol"]})
    assert status == 200 and settings["pairs"] == {"columns": ["symbol"]}
    assert _call(server, "GET", "/api/v2/main/view-settings")[1]["pairs"] == {"columns": ["symbol"]}
    status, card = _call(server, "GET", "/api/v2/main/pair?exchange=bybit&symbol=GOODUSDT")
    assert status == 200 and card["pipeline"]["stage"] == "READY"
    status, _ = _call(server, "POST", "/api/v2/main/decision", {"exchange": "bybit", "symbol": "GOODUSDT",
                                                               "decision": "WATCH", "reason": "r", "author": "u"})
    assert status == 200
    status, answer = _call(server, "POST", "/api/v2/main/setting", {"exchange": "bybit", "symbol": "GOODUSDT",
                                                                   "parameter": "nope", "value": 1, "reason": "r",
                                                                   "author": "u"})
    assert status == 400 and isinstance(answer["error"], str)
    connection = HTTPConnection("127.0.0.1", server.server_port, timeout=10)
    try:
        connection.request("GET", "/api/v2/main/view.xlsx?name=pairs&exchange=bybit&columns=symbol,stage&text=good")
        response = connection.getresponse()
        data = response.read()
        assert response.status == 200 and "pairs_bybit_" in response.getheader("Content-Disposition")
    finally:
        connection.close()
    sheet = load_workbook(BytesIO(data)).active
    assert [[cell.value for cell in row] for row in sheet.iter_rows()] == [["Пара", "Стадия"], ["GOODUSDT", "READY"]]


def test_routes_plan_run_stop_and_task(tmp_path: Path, served) -> None:
    from tests.test_panel_store_ready import _call

    server, controller, _service, runner = _served_world(tmp_path, served)
    status, plan = _call(server, "POST", "/api/v2/main/plan", {"action": "ready", "exchange": "bybit"})
    assert status == 200 and {item["key"] for item in plan["items"]} == {"GOODUSDT", "NEWUSDT", "YOUNGUSDT"}
    status, started = _call(server, "POST", "/api/v2/main/run", {"action": "ready", "exchange": "bybit",
                                                                "items": ["GOODUSDT"], "options": {}})
    assert status == 200 and started["task"]["kind"] == "ready" and runner.started[-1][0] == "ready"
    status, answer = _call(server, "POST", "/api/v2/main/run", {"action": "ready", "exchange": "bybit", "items": []})
    assert status == 400 and isinstance(answer["error"], str)
    status, answer = _call(server, "POST", "/api/v2/main/run", {"action": "sync", "exchange": "bybit", "items": [],
                                                               "options": []})
    assert status == 400
    assert _call(server, "POST", "/api/v2/main/stop", {})[1]["task"]["state"] == "stopped"
    assert _call(server, "GET", "/api/v2/main/task")[1]["task"]["state"] == "idle"
    assert controller.has_active_panel_jobs() is False
    runner.status = lambda: {"task": {"state": "running"}}            # a running task blocks the panel restart
    assert controller.has_active_panel_jobs() is True


def test_routes_answer_json_for_every_failure(tmp_path: Path, served, monkeypatch) -> None:
    # review of stage 3: busy DuckDB files, bad field types and unexpected errors never drop the connection
    import duckdb
    from tests.test_panel_store_ready import _call

    server, _controller, service, _runner = _served_world(tmp_path, served)
    status, answer = _call(server, "POST", "/api/v2/main/view-settings", {"view": ["pairs"], "columns": ["symbol"]})
    assert status == 400 and isinstance(answer["error"], str)
    status, answer = _call(server, "POST", "/api/v2/main/setting", {"exchange": "bybit", "symbol": "GOODUSDT",
                                                                   "filter": ["x"], "parameter": "admit_score",
                                                                   "value": 1, "reason": "r", "author": "u"})
    assert status == 400
    status, answer = _call(server, "POST", "/api/v2/main/run", {"action": "tick", "exchange": "bybit",
                                                               "items": ["S"], "options": {"end_date": "abc"}})
    assert status == 400 and "окно" in answer["error"]

    def busy(*_args, **_kwargs):
        raise duckdb.ConnectionException("Can't open a connection to same database file with a different configuration")

    monkeypatch.setattr(service, "_read", busy)
    status, answer = _call(server, "GET", "/api/v2/main/view?name=pairs&exchange=bybit")
    assert status == 409 and "заняты" in answer["error"]

    def broken(*_args, **_kwargs):
        raise KeyError("unexpected")

    monkeypatch.setattr(service, "pair", broken)
    status, answer = _call(server, "GET", "/api/v2/main/pair?exchange=bybit&symbol=GOODUSDT")
    assert status == 500 and "журнале" in answer["error"]


def test_store_reads_work_while_another_thread_holds_the_catalog(tmp_path: Path) -> None:
    # review of stage 3: the Panel never opens the catalog writable, so a held read-only connection does not
    # break plans, the pair card or the READY source column
    import duckdb

    service, runner, store = _world(tmp_path)
    held = duckdb.connect(str(store.catalog_path), read_only=True)
    try:
        assert service.view("pairs", "bybit")["rows"][0]["ready_source"] == "simulator"
        assert len(service.pair("bybit", "GOODUSDT")["periods"]) == 3
        assert service.plan("ready", "bybit")["items"]
        service.run("ready", "bybit", ["GOODUSDT"])
        assert runner.started[-1][0] == "ready"
    finally:
        held.close()


def test_plans_and_views_work_before_the_store_has_a_catalog(tmp_path: Path) -> None:
    service, _runner, store = _world(tmp_path)
    config = json.loads(service.config_path.read_text(encoding="utf-8"))
    config["trade_store"]["root"] = str(tmp_path / "empty-store")
    service.config_path.write_text(json.dumps(config), encoding="utf-8")
    assert service.view("pairs", "bybit")["rows"][0]["ready_source"] is None
    assert service.plan("tester", "bybit")["items"] == []
    assert {item["detail"] for item in service.plan("ready", "bybit")["items"]} == {"мало периодов (0)"}
    assert service.pair("bybit", "GOODUSDT")["periods"] == []
    assert not (tmp_path / "empty-store").exists()                # nothing is created by a read


def test_history_plan_names_realigned_last_windows(tmp_path: Path, monkeypatch) -> None:
    # the stored GOODUSDT windows (2026-07-01 .. 2026-08-30) are off the listing grid of 2026-06-01: the planned
    # windows that start inside the stored history are a recalculation, not new periods
    service, _runner, _store = _world(tmp_path)
    monkeypatch.setattr(service, "_history_end_day", lambda exchange="bybit": date(2026, 8, 31))
    detail = {item["key"]: item for item in service.plan("history", "bybit")["items"]}["GOODUSDT"]["detail"]
    assert detail == "пересчитать последние окна: 4 периода"


def test_excel_never_writes_registry_text_as_a_formula(tmp_path: Path, monkeypatch) -> None:
    service, _runner, _store = _world(tmp_path)
    real = service.view

    def view(name, exchange, options=None):
        document = real(name, exchange, options)
        document["rows"][0]["symbol"] = "=HYPERLINK(\"x\")"
        return document

    monkeypatch.setattr(service, "view", view)
    _name, data = service.view_xlsx("pairs", "bybit", ["symbol"])
    cell = load_workbook(BytesIO(data)).active["A2"]
    assert cell.value == "=HYPERLINK(\"x\")" and cell.data_type == "s"


def test_filter_and_sort_follow_javascript_text_and_numbers() -> None:
    from mrs3.panel_main import _number, _text

    assert [_text(v) for v in (25.0, 1e-05, 0.1, True, None, 3, "a")] == ["25", "0.00001", "0.1", "true", "", "3", "a"]
    assert [_number(v) for v in ("0x10", "1_000", "inf", "1e3", " 2.5 ", True, float("nan"))] == [
        None, None, None, 1000.0, 2.5, None, None]
    rows = [{"a": 25.0}, {"a": 1e-05}]
    assert filter_rows(rows, ["a"], "0.00001") == [{"a": 1e-05}] and filter_rows(rows, ["a"], "25.0") == []


def test_tick_plan_takes_the_newest_ready_folder_by_its_manifest_time(tmp_path: Path) -> None:
    # re-review of stage 3: «<day>_ready_*» sorts before «<day>_simulation_*» by name, so the newest folder is
    # chosen by the manifest time, not by the name
    import os

    service, _runner, store = _world(tmp_path)
    template = json.loads(TEMPLATE.read_text(encoding="utf-8"))
    newer = store.root / "ready" / "2026-10-18_ready_130000"
    build_ready(store, ["GOODUSDT"], newer, template)
    old = store.root / "ready" / "2026-10-18_simulation_120000" / "ready_manifest.json"
    stamp = (newer / "ready_manifest.json").stat().st_mtime
    os.utime(old, (stamp - 60, stamp - 60))
    assert service.plan("tick", "bybit")["ready"] == "2026-10-18_ready_130000"
    # review of Panel stage 5b: a newer BingX READY folder never stands in for the Bybit tick test
    import shutil

    from mrs3.store_ready_tick import list_ready_folders

    bingx = store.root / "ready" / "2026-10-18_ready_140000"
    shutil.copytree(newer, bingx)
    manifest = json.loads((bingx / "ready_manifest.json").read_text(encoding="utf-8"))
    (bingx / "ready_manifest.json").write_text(json.dumps({**manifest, "exchange": "bingx"}), encoding="utf-8")
    os.utime(bingx / "ready_manifest.json", (stamp + 60, stamp + 60))
    listed = {folder["name"]: folder for folder in list_ready_folders(store.root)}
    assert listed["2026-10-18_ready_140000"]["testable"] and listed["2026-10-18_ready_140000"]["exchange"] == "bingx"
    assert service.plan("tick", "bybit")["ready"] == "2026-10-18_ready_130000"
