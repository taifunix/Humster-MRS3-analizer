from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
import pandas as pd
import pytest

from mrs3.config import AlgorithmConfig
from mrs3.screener.config import ScreenerConfig
from mrs3.screener.errors import ScreenerEvaluationError
from mrs3.screener.evaluate import evaluate_and_record, evaluate_pairs
from mrs3.screener.registry import SCREENING_SHEET

_ALGO = AlgorithmConfig.defaults()
_SCREENER = ScreenerConfig(
    stop_best_pnl30=Decimal("8"),
    go_min_good=2,
    big_min_good=1,
    good_pnl30=Decimal("10"),
    big_shift_bp=110,
    expected_combos_per_pair=2,
)

_START = "2026-08-01 00:00:00"
_END = "2026-08-31 00:00:00"


def _row(
    symbol: str,
    tf: str,
    multiplier: str,
    close_len: str,
    pnl: float,
    dd: float,
    win_rate: float,
    trades: int,
    *,
    side: str = "LONG",
    start: str = _START,
    end: str = _END,
) -> dict:
    mult_col = "settings[*].mrs2.ma_long.multiplier" if side == "LONG" else "settings[*].mrs2.ma_short.multiplier"
    close_col = "settings[*].mrs2.ma_close_long.len" if side == "LONG" else "settings[*].mrs2.ma_close_short.len"
    return {
        "StartDate": start,
        "EndDate": end,
        "TotalPnLPercent": pnl,
        "MaxDrawdownPercent": dd,
        "TotalTrades": trades,
        "WinRate": win_rate,
        "settings[*].basic.symbol": symbol,
        "settings[*].basic.time_frame": tf,
        close_col: close_len,
        mult_col: multiplier,
    }


def _write_csv(path: Path, rows: list[dict]) -> None:
    pd.DataFrame(rows).to_csv(path, index=False, encoding="utf-8-sig")


def _patch_listing(monkeypatch: pytest.MonkeyPatch, symbols: list[str]) -> None:
    fixed = {symbol: pd.Timestamp("2020-01-01", tz="UTC") for symbol in symbols}
    monkeypatch.setattr(
        "mrs3.screener.evaluate.resolve_listing_dates",
        lambda requested, **kwargs: {symbol: fixed[symbol] for symbol in requested},
    )


def test_evaluate_pairs_go_check_stop_incomplete(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    report_dir = tmp_path / "my_test"
    report_dir.mkdir()
    rows = [
        # GOUSDT: both rows economically pass with pnl30 well above good_pnl30.
        _row("GOUSDT", "1h", "0.99", "2", pnl=30, dd=5, win_rate=80, trades=20),
        _row("GOUSDT", "2h", "0.98", "6", pnl=25, dd=5, win_rate=80, trades=20),
        # CHECKUSDT: one good point, one failing WinRate -> not enough good points for GO.
        _row("CHECKUSDT", "1h", "0.99", "2", pnl=30, dd=5, win_rate=80, trades=20),
        _row("CHECKUSDT", "2h", "0.98", "6", pnl=30, dd=5, win_rate=50, trades=20),
        # STOPUSDT: both rows fail economic pass (WinRate too low).
        _row("STOPUSDT", "1h", "0.99", "2", pnl=30, dd=5, win_rate=10, trades=20),
        _row("STOPUSDT", "2h", "0.98", "6", pnl=30, dd=5, win_rate=10, trades=20),
        # INCOMPLETEUSDT: only one of the two expected combos present.
        _row("INCOMPLETEUSDT", "1h", "0.99", "2", pnl=30, dd=5, win_rate=80, trades=20),
    ]
    _write_csv(report_dir / "reports_history.csv", rows)
    _patch_listing(monkeypatch, ["GOUSDT", "CHECKUSDT", "STOPUSDT", "INCOMPLETEUSDT"])

    verdicts = evaluate_pairs(
        report_dir, algorithm_config=_ALGO, screener_config=_SCREENER, dates_path=None
    )
    by_symbol = {v.symbol: v for v in verdicts}

    assert by_symbol["GOUSDT"].verdict == "GO"
    assert by_symbol["GOUSDT"].n_good == 2
    assert by_symbol["CHECKUSDT"].verdict == "CHECK"
    assert by_symbol["CHECKUSDT"].n_good == 1
    assert by_symbol["STOPUSDT"].verdict == "STOP"
    assert by_symbol["STOPUSDT"].best_pnl30 is None
    assert by_symbol["INCOMPLETEUSDT"].verdict == "INCOMPLETE"
    assert by_symbol["INCOMPLETEUSDT"].n_unique_combos == 1


def test_evaluate_pairs_big_shift_flag(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    report_dir = tmp_path / "my_test"
    report_dir.mkdir()
    rows = [
        # shift_bp = round((1 - 0.98) * 10000) = 200 >= big_shift_bp(110).
        _row("BIGUSDT", "1h", "0.98", "2", pnl=30, dd=5, win_rate=80, trades=20),
        _row("BIGUSDT", "2h", "0.997", "6", pnl=30, dd=5, win_rate=80, trades=20),
    ]
    _write_csv(report_dir / "reports_history.csv", rows)
    _patch_listing(monkeypatch, ["BIGUSDT"])

    verdicts = evaluate_pairs(
        report_dir, algorithm_config=_ALGO, screener_config=_SCREENER, dates_path=None
    )
    assert verdicts[0].big_shift is True
    assert verdicts[0].n_good_big_shift == 1


def test_evaluate_pairs_short_side_shift_bp_and_comma_multiplier(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    report_dir = tmp_path / "my_test"
    report_dir.mkdir()
    rows = [
        # SHORT multiplier "1,020" -> shift_bp = round((1.02 - 1) * 10000) = 200.
        _row("SHORTUSDT", "1h", "1,020", "2", pnl=30, dd=5, win_rate=80, trades=20, side="SHORT"),
        _row("SHORTUSDT", "2h", "1,003", "6", pnl=30, dd=5, win_rate=80, trades=20, side="SHORT"),
    ]
    _write_csv(report_dir / "reports_history.csv", rows)
    _patch_listing(monkeypatch, ["SHORTUSDT"])

    verdicts = evaluate_pairs(
        report_dir, algorithm_config=_ALGO, screener_config=_SCREENER, dates_path=None
    )
    assert verdicts[0].side == "SHORT"
    # Both rows tie on pnl30 (identical pnl/dd/trades/win_rate/listing date);
    # max() is stable and keeps the first-seen point on ties.
    assert verdicts[0].best_shift_bp == 200


def test_evaluate_pairs_rejects_mixed_sides(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    report_dir = tmp_path / "my_test"
    report_dir.mkdir()
    long_row = _row("AUSDT", "1h", "0.99", "2", pnl=30, dd=5, win_rate=80, trades=20, side="LONG")
    short_row = _row("AUSDT", "2h", "1,003", "6", pnl=30, dd=5, win_rate=80, trades=20, side="SHORT")
    frame = pd.concat([pd.DataFrame([long_row]), pd.DataFrame([short_row])], ignore_index=True)
    frame.to_csv(report_dir / "reports_history.csv", index=False, encoding="utf-8-sig")
    _patch_listing(monkeypatch, ["AUSDT"])

    with pytest.raises(ScreenerEvaluationError, match="mix LONG and SHORT"):
        evaluate_pairs(report_dir, algorithm_config=_ALGO, screener_config=_SCREENER, dates_path=None)


def test_evaluate_pairs_rejects_mixed_windows(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    report_dir = tmp_path / "my_test"
    report_dir.mkdir()
    rows = [
        _row("AUSDT", "1h", "0.99", "2", pnl=30, dd=5, win_rate=80, trades=20),
        _row("AUSDT", "2h", "0.98", "6", pnl=30, dd=5, win_rate=80, trades=20, start="2026-07-01 00:00:00"),
    ]
    _write_csv(report_dir / "reports_history.csv", rows)
    _patch_listing(monkeypatch, ["AUSDT"])

    with pytest.raises(ScreenerEvaluationError, match="more than one"):
        evaluate_pairs(report_dir, algorithm_config=_ALGO, screener_config=_SCREENER, dates_path=None)


def test_evaluate_pairs_merges_partitioned_csv_files(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    report_dir = tmp_path / "my_test"
    report_dir.mkdir()
    _write_csv(
        report_dir / "reports_history.csv",
        [_row("AUSDT", "1h", "0.99", "2", pnl=30, dd=5, win_rate=80, trades=20)],
    )
    _write_csv(
        report_dir / "reports_history_p1.csv",
        [_row("AUSDT", "2h", "0.98", "6", pnl=30, dd=5, win_rate=80, trades=20)],
    )
    _patch_listing(monkeypatch, ["AUSDT"])

    verdicts = evaluate_pairs(
        report_dir, algorithm_config=_ALGO, screener_config=_SCREENER, dates_path=None
    )
    assert verdicts[0].n_reports == 2
    assert verdicts[0].verdict == "GO"


def test_evaluate_pairs_malformed_row_only_affects_its_own_symbol(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    report_dir = tmp_path / "my_test"
    report_dir.mkdir()
    good_row = _row("GOODUSDT", "1h", "0.99", "2", pnl=30, dd=5, win_rate=80, trades=20)
    good_row2 = _row("GOODUSDT", "2h", "0.98", "6", pnl=30, dd=5, win_rate=80, trades=20)
    bad_row = _row("BADUSDT", "1h", "0.99", "2", pnl=30, dd=5, win_rate=80, trades=20)
    bad_row["TotalPnLPercent"] = "not-a-number"
    bad_row2 = _row("BADUSDT", "2h", "0.98", "6", pnl=30, dd=5, win_rate=80, trades=20)
    _write_csv(report_dir / "reports_history.csv", [good_row, good_row2, bad_row, bad_row2])
    _patch_listing(monkeypatch, ["GOODUSDT", "BADUSDT"])

    verdicts = evaluate_pairs(
        report_dir, algorithm_config=_ALGO, screener_config=_SCREENER, dates_path=None
    )
    by_symbol = {v.symbol: v for v in verdicts}
    assert by_symbol["GOODUSDT"].verdict == "GO"
    assert by_symbol["BADUSDT"].verdict == "INCOMPLETE"
    assert by_symbol["BADUSDT"].n_unique_combos == 1


def test_evaluate_pairs_rejects_missing_required_column(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    report_dir = tmp_path / "my_test"
    report_dir.mkdir()
    rows = [_row("AUSDT", "1h", "0.99", "2", pnl=30, dd=5, win_rate=80, trades=20)]
    frame = pd.DataFrame(rows).drop(columns=["WinRate"])
    frame.to_csv(report_dir / "reports_history.csv", index=False, encoding="utf-8-sig")
    _patch_listing(monkeypatch, ["AUSDT"])

    with pytest.raises(ScreenerEvaluationError, match="missing required columns"):
        evaluate_pairs(report_dir, algorithm_config=_ALGO, screener_config=_SCREENER, dates_path=None)


def test_evaluate_pairs_rejects_blank_symbol(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    report_dir = tmp_path / "my_test"
    report_dir.mkdir()
    row = _row("AUSDT", "1h", "0.99", "2", pnl=30, dd=5, win_rate=80, trades=20)
    row["settings[*].basic.symbol"] = ""
    frame = pd.DataFrame([row])
    frame.loc[0, "settings[*].basic.symbol"] = None
    frame.to_csv(report_dir / "reports_history.csv", index=False, encoding="utf-8-sig")
    _patch_listing(monkeypatch, ["AUSDT"])

    with pytest.raises(ScreenerEvaluationError, match="blank basic.symbol"):
        evaluate_pairs(report_dir, algorithm_config=_ALGO, screener_config=_SCREENER, dates_path=None)


def test_evaluate_pairs_dd_non_positive_never_passes(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    report_dir = tmp_path / "my_test"
    report_dir.mkdir()
    rows = [
        _row("ZEROUSDT", "1h", "0.99", "2", pnl=30, dd=0, win_rate=80, trades=20),
        _row("ZEROUSDT", "2h", "0.98", "6", pnl=30, dd=0, win_rate=80, trades=20),
    ]
    _write_csv(report_dir / "reports_history.csv", rows)
    _patch_listing(monkeypatch, ["ZEROUSDT"])

    verdicts = evaluate_pairs(
        report_dir, algorithm_config=_ALGO, screener_config=_SCREENER, dates_path=None
    )
    assert verdicts[0].verdict == "STOP"
    assert verdicts[0].best_pnl30 is None


def test_evaluate_pairs_unreadable_duplicate_combo_still_marks_incomplete(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    # Two valid distinct combos (matches expected_combos_per_pair=2) plus a
    # third, unreadable row that duplicates one of those combos — combo
    # COUNT alone would look complete, but a report was still unreadable.
    report_dir = tmp_path / "my_test"
    report_dir.mkdir()
    good_row_1 = _row("DUPUSDT", "1h", "0.99", "2", pnl=30, dd=5, win_rate=80, trades=20)
    good_row_2 = _row("DUPUSDT", "2h", "0.98", "6", pnl=30, dd=5, win_rate=80, trades=20)
    duplicate_bad_row = _row("DUPUSDT", "1h", "0.99", "2", pnl=30, dd=5, win_rate=80, trades=20)
    duplicate_bad_row["TotalPnLPercent"] = "not-a-number"
    _write_csv(report_dir / "reports_history.csv", [good_row_1, good_row_2, duplicate_bad_row])
    _patch_listing(monkeypatch, ["DUPUSDT"])

    verdicts = evaluate_pairs(
        report_dir, algorithm_config=_ALGO, screener_config=_SCREENER, dates_path=None
    )
    assert verdicts[0].n_unique_combos == 2
    assert verdicts[0].verdict == "INCOMPLETE"


def test_evaluate_pairs_listing_date_after_window_end_is_incomplete_not_go(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    report_dir = tmp_path / "my_test"
    report_dir.mkdir()
    # A losing point (negative PnL); if effective_days were allowed to go
    # negative this would wrongly flip pnl30 positive instead of failing.
    rows = [
        _row("NEWUSDT", "1h", "0.99", "2", pnl=-10, dd=5, win_rate=80, trades=20),
        _row("NEWUSDT", "2h", "0.98", "6", pnl=-10, dd=5, win_rate=80, trades=20),
    ]
    _write_csv(report_dir / "reports_history.csv", rows)
    # Listing date after the report window's EndDate (2026-08-31).
    monkeypatch.setattr(
        "mrs3.screener.evaluate.resolve_listing_dates",
        lambda requested, **kwargs: {s: pd.Timestamp("2026-09-15", tz="UTC") for s in requested},
    )

    verdicts = evaluate_pairs(
        report_dir, algorithm_config=_ALGO, screener_config=_SCREENER, dates_path=None
    )
    assert verdicts[0].verdict == "INCOMPLETE"
    assert verdicts[0].n_unique_combos == 0


def test_evaluate_and_record_writes_verdicts_to_registry(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    from openpyxl import Workbook

    report_dir = tmp_path / "my_test"
    report_dir.mkdir()
    rows = [
        _row("GOUSDT", "1h", "0.99", "2", pnl=30, dd=5, win_rate=80, trades=20),
        _row("GOUSDT", "2h", "0.98", "6", pnl=30, dd=5, win_rate=80, trades=20),
    ]
    _write_csv(report_dir / "reports_history.csv", rows)
    _patch_listing(monkeypatch, ["GOUSDT"])

    registry_path = tmp_path / "registry.xlsx"
    workbook = Workbook()
    pairs = workbook.active
    pairs.title = "Пары"
    pairs.append(["Пара", "Дата листинга на Bybit (UTC)"])
    pairs.append(["GOUSDT", "2020-01-01"])
    workbook.save(registry_path)

    screener_config = replace(_SCREENER, liquidity_registry_path=registry_path)
    evaluate_and_record(
        report_dir, algorithm_config=_ALGO, screener_config=screener_config, dates_path=None
    )

    workbook = load_workbook(registry_path)
    sheet = workbook[SCREENING_SHEET]
    row = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
    assert row[0] == "GOUSDT"
    assert row[2] == "GO"
