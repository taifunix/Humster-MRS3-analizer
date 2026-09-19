from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest

from mrs3.screener.errors import ScreenerEvaluationError
from mrs3.screener.registry import (
    SCREENING_HEADER,
    SCREENING_SHEET,
    ScreeningRow,
    list_unscreened_symbols,
    read_registry_listing_dates,
    write_screening_results,
)


def _build_registry(tmp_path: Path, pairs_rows: list[tuple[str, str]]) -> Path:
    workbook = Workbook()
    pairs = workbook.active
    pairs.title = "Пары"
    pairs.append(["Пара", "Дата листинга на Bybit (UTC)", "Тип"])
    for symbol, listing_date in pairs_rows:
        pairs.append([symbol, listing_date, "stock"])
    daily = workbook.create_sheet("По дням")
    daily.append(["Пара", "Оборот 11.09 Fri, USDT"])
    daily.append(["SOXLUSDT", 12345])
    path = tmp_path / "registry.xlsx"
    workbook.save(path)
    return path


def test_read_registry_listing_dates_parses_pairs_sheet(tmp_path: Path) -> None:
    path = _build_registry(
        tmp_path,
        [("SOXLUSDT", "2026-05-19"), ("XAUUSDT", "2026-03-09")],
    )
    dates = read_registry_listing_dates(path, ("SOXLUSDT", "XAUUSDT"))
    assert dates["SOXLUSDT"] == pd.Timestamp("2026-05-19", tz="UTC")
    assert dates["XAUUSDT"] == pd.Timestamp("2026-03-09", tz="UTC")


def test_read_registry_listing_dates_rejects_duplicate_symbols(tmp_path: Path) -> None:
    path = _build_registry(
        tmp_path,
        [("SOXLUSDT", "2026-05-19"), ("SOXLUSDT", "2026-05-20")],
    )
    with pytest.raises(ScreenerEvaluationError, match="duplicate symbols"):
        read_registry_listing_dates(path, ("SOXLUSDT",))


def test_read_registry_listing_dates_ignores_unrelated_malformed_rows(
    tmp_path: Path,
) -> None:
    path = _build_registry(
        tmp_path,
        [("SOXLUSDT", "2026-05-19"), ("BADUSDT", "not-a-date")],
    )
    dates = read_registry_listing_dates(path, ("SOXLUSDT",))
    assert dates == {"SOXLUSDT": pd.Timestamp("2026-05-19", tz="UTC")}


def test_read_registry_listing_dates_ignores_unrelated_duplicate_rows(
    tmp_path: Path,
) -> None:
    path = _build_registry(
        tmp_path,
        [
            ("SOXLUSDT", "2026-05-19"),
            ("BADUSDT", "2026-01-01"),
            ("BADUSDT", "2026-01-02"),
        ],
    )
    dates = read_registry_listing_dates(path, ("SOXLUSDT",))
    assert dates == {"SOXLUSDT": pd.Timestamp("2026-05-19", tz="UTC")}


def test_read_registry_listing_dates_missing_file_raises_clear_error(tmp_path: Path) -> None:
    with pytest.raises(ScreenerEvaluationError, match="cannot read liquidity registry"):
        read_registry_listing_dates(tmp_path / "missing.xlsx", ("SOXLUSDT",))


def _row(symbol: str = "SOXLUSDT", side: str = "LONG") -> ScreeningRow:
    return ScreeningRow(
        symbol=symbol,
        side=side,
        verdict="GO",
        big_shift=True,
        n_good=9,
        n_good_big_shift=9,
        window_start="2026-08-01",
        window_end="2026-09-18",
    )


def test_write_screening_results_creates_sheet_with_header_when_missing(tmp_path: Path) -> None:
    path = _build_registry(tmp_path, [("SOXLUSDT", "2026-05-19")])

    write_screening_results(path, (_row(),))

    workbook = load_workbook(path)
    assert SCREENING_SHEET in workbook.sheetnames
    sheet = workbook[SCREENING_SHEET]
    header = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
    assert header == SCREENING_HEADER
    data_row = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
    assert data_row[:8] == ["SOXLUSDT", "LONG", "GO", "да", 9, 9, "2026-08-01", "2026-09-18"]
    assert data_row[9:] == [None, None, None]

    # other sheets survive untouched
    assert workbook["По дням"]["A2"].value == "SOXLUSDT"


def test_write_screening_results_updates_auto_columns_and_preserves_manual_columns(
    tmp_path: Path,
) -> None:
    path = _build_registry(tmp_path, [("SOXLUSDT", "2026-05-19")])
    write_screening_results(path, (_row(),))

    workbook = load_workbook(path)
    sheet = workbook[SCREENING_SHEET]
    sheet.cell(row=2, column=10, value="GOOD")
    sheet.cell(row=2, column=11, value="2026-10-01")
    sheet.cell(row=2, column=12, value="проверено вручную")
    workbook.save(path)

    updated = ScreeningRow(
        symbol="SOXLUSDT",
        side="LONG",
        verdict="CHECK",
        big_shift=False,
        n_good=3,
        n_good_big_shift=0,
        window_start="2026-08-01",
        window_end="2026-09-20",
    )
    write_screening_results(path, (updated,))

    workbook = load_workbook(path)
    sheet = workbook[SCREENING_SHEET]
    assert sheet.max_row == 2
    row_values = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
    assert row_values[:8] == ["SOXLUSDT", "LONG", "CHECK", "нет", 3, 0, "2026-08-01", "2026-09-20"]
    assert row_values[9:] == ["GOOD", "2026-10-01", "проверено вручную"]


def test_write_screening_results_appends_second_pair_without_touching_first(
    tmp_path: Path,
) -> None:
    path = _build_registry(tmp_path, [("SOXLUSDT", "2026-05-19"), ("XAUUSDT", "2026-03-09")])
    write_screening_results(path, (_row("SOXLUSDT"),))
    write_screening_results(path, (_row("XAUUSDT"),))

    workbook = load_workbook(path)
    sheet = workbook[SCREENING_SHEET]
    assert sheet.max_row == 3
    symbols = [sheet.cell(row=index, column=1).value for index in (2, 3)]
    assert symbols == ["SOXLUSDT", "XAUUSDT"]


def test_write_screening_results_rejects_unexpected_existing_header(tmp_path: Path) -> None:
    path = _build_registry(tmp_path, [("SOXLUSDT", "2026-05-19")])
    workbook = load_workbook(path)
    sheet = workbook.create_sheet(SCREENING_SHEET)
    sheet.append(["Symbol", "Side"])
    workbook.save(path)

    with pytest.raises(ScreenerEvaluationError, match="unexpected header"):
        write_screening_results(path, (_row(),))


def test_write_screening_results_missing_file_raises_clear_error(tmp_path: Path) -> None:
    with pytest.raises(ScreenerEvaluationError, match="cannot open liquidity registry"):
        write_screening_results(tmp_path / "missing.xlsx", (_row(),))


def test_list_unscreened_symbols_excludes_pairs_already_screened_for_side(
    tmp_path: Path,
) -> None:
    path = _build_registry(
        tmp_path,
        [("SOXLUSDT", "2026-05-19"), ("XAUUSDT", "2026-03-09"), ("MSTRUSDT", "2026-01-01")],
    )
    write_screening_results(path, (_row("SOXLUSDT", "LONG"),))

    assert list_unscreened_symbols(path, "LONG") == ("XAUUSDT", "MSTRUSDT")


def test_list_unscreened_symbols_treats_sides_independently(tmp_path: Path) -> None:
    path = _build_registry(tmp_path, [("SOXLUSDT", "2026-05-19")])
    write_screening_results(path, (_row("SOXLUSDT", "LONG"),))

    assert list_unscreened_symbols(path, "SHORT") == ("SOXLUSDT",)
    assert list_unscreened_symbols(path, "LONG") == ()


def test_list_unscreened_symbols_returns_all_pairs_when_screening_sheet_missing(
    tmp_path: Path,
) -> None:
    path = _build_registry(tmp_path, [("SOXLUSDT", "2026-05-19"), ("XAUUSDT", "2026-03-09")])

    assert list_unscreened_symbols(path, "LONG") == ("SOXLUSDT", "XAUUSDT")


def test_list_unscreened_symbols_normalizes_case_and_whitespace(tmp_path: Path) -> None:
    path = _build_registry(tmp_path, [(" soxlusdt ", "2026-05-19")])

    assert list_unscreened_symbols(path, "LONG") == ("SOXLUSDT",)


def test_list_unscreened_symbols_rejects_same_symbol_in_different_case(
    tmp_path: Path,
) -> None:
    # Normalization happens before duplicate detection, so a registry row
    # duplicated only by case/whitespace is caught the same way an exact
    # duplicate is (see test_list_unscreened_symbols_rejects_duplicate_symbols_in_pairs_sheet).
    path = _build_registry(tmp_path, [("soxlusdt", "2026-05-19"), ("SOXLUSDT ", "2026-05-19")])

    with pytest.raises(ScreenerEvaluationError, match="duplicate symbols"):
        list_unscreened_symbols(path, "LONG")


def test_list_unscreened_symbols_rejects_invalid_side(tmp_path: Path) -> None:
    path = _build_registry(tmp_path, [("SOXLUSDT", "2026-05-19")])
    with pytest.raises(ScreenerEvaluationError, match="invalid screener side"):
        list_unscreened_symbols(path, "BOTH")


def test_list_unscreened_symbols_missing_file_raises_clear_error(tmp_path: Path) -> None:
    with pytest.raises(ScreenerEvaluationError, match="cannot read liquidity registry"):
        list_unscreened_symbols(tmp_path / "missing.xlsx", "LONG")


def test_list_unscreened_symbols_rejects_duplicate_symbols_in_pairs_sheet(
    tmp_path: Path,
) -> None:
    path = _build_registry(
        tmp_path, [("SOXLUSDT", "2026-05-19"), ("SOXLUSDT", "2026-05-20")]
    )
    with pytest.raises(ScreenerEvaluationError, match="duplicate symbols"):
        list_unscreened_symbols(path, "LONG")


def test_list_unscreened_symbols_rejects_empty_pairs_sheet(tmp_path: Path) -> None:
    workbook = Workbook()
    pairs = workbook.active
    pairs.title = "Пары"
    path = tmp_path / "registry.xlsx"
    workbook.save(path)

    with pytest.raises(ScreenerEvaluationError, match="no header row"):
        list_unscreened_symbols(path, "LONG")


def test_list_unscreened_symbols_rejects_unexpected_screening_sheet_header(
    tmp_path: Path,
) -> None:
    path = _build_registry(tmp_path, [("SOXLUSDT", "2026-05-19")])
    workbook = load_workbook(path)
    sheet = workbook.create_sheet(SCREENING_SHEET)
    sheet.append(["Symbol", "Side"])
    workbook.save(path)

    with pytest.raises(ScreenerEvaluationError, match="unexpected header"):
        list_unscreened_symbols(path, "LONG")
