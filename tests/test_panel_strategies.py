"""The «Стратегии» view and the finalist marks (Panel rework stage 4a).

Contract: docs/specs/2026-10-18-panel-simulation-workflow.md, «Stage 4 contract»; ADR-0080.
"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO
import json
from pathlib import Path

import duckdb
from openpyxl import load_workbook
import pandas as pd
import pytest

import mrs3.panel_strategies as strategies_module
from mrs3.pair_registry.records import current_finalists, record_finalist
from mrs3.pair_registry.store import connect
from mrs3.panel_main import MainScreenError
from tests.test_panel_main import _world
from tests.test_panel_store_ready import served  # noqa: F401  (route-test fixture)

OLD = "GOODUSDT_45m_LONG_2ORD_CMA7_STR_old_EQUAL"


def _ready_names(store) -> list[str]:
    with store.connect(read_only=True) as catalog:
        payload = json.loads(catalog.execute("select payload_json from ready where symbol = 'GOODUSDT'").fetchone()[0])
    return [item["strategy"] for item in payload["structures"]]


def _performance_db(path: Path, names: list[str]) -> Path:
    from mrs3.performance_v2_store import initialize_performance_v2

    with duckdb.connect(str(path)) as con:
        initialize_performance_v2(con, evaluation_protocol="v2", exchange="Bybit")
        now = "2026-10-01 00:00:00+00"
        for strategy_id, name in enumerate(names, 1):
            con.execute("insert into strategies (strategy_id, strategy_name, symbol, side, timeframe, close_ma_len, "
                        "order_count, analysis_run_id, candidate_identity, lifecycle_status, created_at_utc, "
                        "updated_at_utc) values (?, ?, 'GOODUSDT', 'LONG', '1h', 2, 2, 'run', ?, 'ACTIVE', ?, ?)",
                        [strategy_id, name, name, now, now])
            con.execute("insert into strategy_results (result_id, strategy_id, report_start_utc, report_end_utc, "
                        "exchange, initial_balance, final_balance, total_pnl_pct, max_drawdown_pct, total_trades, "
                        "imported_at_utc) values (?, ?, '2026-09-01 00:00:00+00', '2026-10-01 00:00:00+00', 'Bybit', "
                        "10000, 10300, ?, 1.5, 40, ?)", [10 + strategy_id, strategy_id, 3 + strategy_id, now])
            con.execute("update strategies set current_result_id = ? where strategy_id = ?", [10 + strategy_id, strategy_id])
        tables = {row[0] for row in con.execute("select table_name from information_schema.tables").fetchall()}
        if "strategy_report_metrics" in tables:
            for strategy_id in range(1, len(names) + 1):
                con.execute("insert into strategy_report_metrics (result_id, ulcer_index_hourly_pct, time_under_water_pct) "
                            "values (?, ?, ?)", [10 + strategy_id, 5.0 - strategy_id, 40.0])
    return path


def _strategies_world(tmp_path: Path, monkeypatch):
    service, runner, store = _world(tmp_path)
    names = _ready_names(store)
    database = _performance_db(tmp_path / "perf.duckdb", [names[0], OLD])
    monkeypatch.setattr(service, "_performance_db", lambda: database)
    return service, store, names, database


def test_rows_join_ready_structures_with_tick_tests_and_rank_in_pair(tmp_path: Path, monkeypatch) -> None:
    service, _store, names, _db = _strategies_world(tmp_path, monkeypatch)
    view = service.view("strategies", "bybit")
    rows = {row["strategy"]: row for row in view["rows"]}
    assert set(rows) == {*names, OLD}                                    # the union of READY and tested strategies
    tested = rows[names[0]]
    assert (tested["tested"], tested["total_pnl_pct"], tested["window"]) == (True, 4.0, "2026-09-01..2026-10-01")
    assert tested["category"] in ("low", "moderate", "high")
    assert tested["orders"] and tested["score"] is not None and tested["ready_rank"] == 1
    assert tested["ulcer_index_hourly_pct"] == 4.0 and tested["time_under_water_pct"] == 40.0
    untested = [rows[name] for name in names[1:]]
    assert all(row["tested"] is False and row["total_pnl_pct"] is None for row in untested)
    old = rows[OLD]
    assert old["tested"] is True and old["category"] is None and old["rank_in_pair"] is None
    ranked = sorted((row for row in view["rows"] if row["rank_in_pair"]), key=lambda row: row["rank_in_pair"])
    assert [row["rank_in_pair"] for row in ranked] == list(range(1, len(names) + 1))
    assert view["default_columns"][:3] == ["strategy", "symbol", "side"]


def test_filter_toggles_explain_cut_rows_and_hide_them(tmp_path: Path, monkeypatch) -> None:
    service, _store, names, _db = _strategies_world(tmp_path, monkeypatch)

    def candidates(connection, request, config, *, cache_only):
        assert cache_only is True                                         # read only: no window cache writes
        return pd.DataFrame([
            {"strategy_id": 1, "pnl_30d_pct": Decimal("2"), "max_drawdown_pct": Decimal("1.5"), "history_days": 30},
            {"strategy_id": 2, "pnl_30d_pct": Decimal("30"), "max_drawdown_pct": Decimal("5"), "history_days": 90,
             "completed_profitable_cycle_count": 40, "completed_cycle_net_pnl": Decimal("100"),
             "top5_share_pct": Decimal("85"), "pnl_after_top5": Decimal("15"), "completed_cycles_reliable": True},
        ])

    monkeypatch.setattr(strategies_module, "load_selection_candidates", candidates)
    rows = {row["strategy"]: row for row in service.view("strategies", "bybit")["rows"]}
    assert rows[names[0]]["reason"].startswith("Жёсткие отсечки: PNL30_FLOOR")
    assert rows[names[0]]["pnl_30d_pct"] == 2.0
    assert rows[OLD]["reason"] == "Топ-5: доля 85.0 %"
    assert all(rows[name]["reason"] is None for name in names[1:])
    rows = {row["strategy"]: row for row in service.view("strategies", "bybit", {"hard": "0", "top5": "0"})["rows"]}
    assert rows[names[0]]["reason"] is None and rows[OLD]["reason"] is None
    hidden = {row["strategy"] for row in service.view("strategies", "bybit", {"hide": "1"})["rows"]}
    assert hidden == set(names[1:])


def test_marks_are_written_to_the_registry_and_shown(tmp_path: Path, monkeypatch) -> None:
    service, _store, names, _db = _strategies_world(tmp_path, monkeypatch)
    result = service.finalist({"exchange": "bybit", "strategy": names[0], "symbol": "goodusdt", "side": "LONG",
                               "status": "FINALIST", "user_rank": "2"})
    assert result == {"strategy": names[0], "status": "FINALIST", "user_rank": 2, "changed": True}
    row = next(row for row in service.view("strategies", "bybit")["rows"] if row["strategy"] == names[0])
    assert (row["status"], row["user_rank"]) == ("FINALIST", 2)
    service.finalist({"exchange": "bybit", "strategy": names[0], "symbol": "GOODUSDT", "status": "", "user_rank": 2})
    row = next(row for row in service.view("strategies", "bybit")["rows"] if row["strategy"] == names[0])
    assert (row["status"], row["user_rank"]) == (None, None)
    con = connect(tmp_path / "registry")
    assert [r[0] for r in con.execute("select source from finalists order by seq").fetchall()] == ["panel", "panel"]
    con.close()
    for bad in ({"status": "MAYBE"}, {"user_rank": 0}, {"user_rank": "1.5"}, {"strategy": ""},
                {"strategy": "NOT_IN_THE_VIEW"}):
        with pytest.raises(MainScreenError):
            service.finalist({"exchange": "bybit", "strategy": names[0], "symbol": "GOODUSDT", "side": "LONG",
                              "status": "RESERVE", **bad})
    # review of 4a: pair and side come from the view, and clearing a strategy without a mark writes nothing
    assert service.finalist({"exchange": "bybit", "strategy": OLD, "status": ""})["changed"] is False
    con = connect(tmp_path / "registry")
    assert con.execute("select count(*) from finalists where strategy = ?", [OLD]).fetchone()[0] == 0
    con.close()
    service.finalist({"exchange": "bybit", "strategy": OLD, "symbol": "WRONGUSDT", "side": "SHORT",
                      "status": "RESERVE"})
    con = connect(tmp_path / "registry")
    assert con.execute("select symbol, side from finalists where strategy = ?", [OLD]).fetchone() == ("GOODUSDT", "LONG")
    con.close()


def test_excel_round_trip_writes_only_changes_and_nothing_on_export(tmp_path: Path, monkeypatch) -> None:
    service, _store, names, database = _strategies_world(tmp_path, monkeypatch)
    before = database.read_bytes()
    _name, data = service.view_xlsx("strategies", "bybit", ["strategy", "symbol", "status", "user_rank"])
    assert database.read_bytes() == before                                # export has no PerformanceDB side effect
    book = load_workbook(BytesIO(data))
    sheet = book.active
    header = [cell.value for cell in sheet[1]]
    assert header == ["Стратегия", "Пара", "Отметка", "Ранг"]
    by_name = {sheet.cell(row, 1).value: row for row in range(2, sheet.max_row + 1)}
    sheet.cell(by_name[names[0]], 3).value = "finalist"
    sheet.cell(by_name[names[0]], 4).value = 1
    sheet.cell(by_name[OLD], 3).value = "REJECTED"
    sheet.append(["UNKNOWN_STRATEGY", "GOODUSDT", "FINALIST", None])
    buffer = BytesIO()
    book.save(buffer)
    result = service.import_marks("bybit", buffer.getvalue())
    assert result["written"] is True and result["unknown"] == ["UNKNOWN_STRATEGY"]
    assert {(item["strategy"], item["to"]) for item in result["changed"]} == {(names[0], "FINALIST"), (OLD, "REJECTED")}
    assert result["unchanged"] == len(names) - 1
    again = service.import_marks("bybit", buffer.getvalue())
    assert again["changed"] == [] and again["unchanged"] == len(names) + 1
    con = connect(tmp_path / "registry")
    marks = {row["strategy"]: (row["status"], row["user_rank"]) for row in current_finalists(con, "bybit")}
    assert marks == {names[0]: ("FINALIST", 1), OLD: ("REJECTED", None)}
    assert {r[0] for r in con.execute("select source from finalists").fetchall()} == {"excel"}
    con.close()
    sheet.cell(by_name[OLD], 3).value = "MAYBE"                            # one invalid row: nothing is written
    sheet.cell(by_name[names[0]], 3).value = "RESERVE"
    buffer = BytesIO()
    book.save(buffer)
    refused = service.import_marks("bybit", buffer.getvalue())
    assert refused["written"] is False and refused["errors"][0]["strategy"] == OLD
    con = connect(tmp_path / "registry")
    assert {row["strategy"]: row["status"] for row in current_finalists(con, "bybit")}[names[0]] == "FINALIST"
    con.close()
    with pytest.raises(MainScreenError):
        service.import_marks("bybit", b"not a workbook")
    sheet.cell(by_name[OLD], 3).value = "RESERVE"                          # review of 4a: a repeated strategy row
    sheet.append([names[0], "GOODUSDT", "REJECTED", None])
    buffer = BytesIO()
    book.save(buffer)
    repeated = service.import_marks("bybit", buffer.getvalue())
    assert repeated["written"] is False and repeated["errors"][0]["error"] == "стратегия повторяется в книге"


def test_the_pipeline_seeds_the_marks_from_the_performance_db_only_once(tmp_path: Path) -> None:
    from datetime import date

    from mrs3.pair_registry.pipeline import compute_pipeline
    from tests.test_pair_registry_records import NOW, _performance_db as registry_performance_db, _registry, _review
    from tests.test_ready_store import _store_with_periods

    con = _registry(tmp_path)
    database = registry_performance_db(tmp_path / "perf.duckdb")
    _review(database, "1", "2026-10-02 00:00:00+00", [(1, "FINALIST", 1)])
    store = _store_with_periods(tmp_path / "store", "GOODUSDT", 3, 35.0)
    first = compute_pipeline(con, store, "bybit", date(2026, 10, 19), "pc", performance_db=database, now_ms=NOW,
                             track=False)
    assert first["finalist_changes"] == 1                                 # the one-time migration
    record_finalist(con, "bybit", "GOODUSDT", "LONG", "GOODUSDT_15m_LONG_A", "RESERVE", "panel", now_ms=NOW + 1)
    _review(database, "2", "2026-10-03 00:00:00+00", [(1, "REJECTED", None)])
    second = compute_pipeline(con, store, "bybit", date(2026, 10, 19), "pc", performance_db=database, now_ms=NOW + 2,
                              track=False)
    assert second["finalist_changes"] == 0                                # the Panel mark is never overwritten
    assert current_finalists(con, "bybit")[0]["status"] == "RESERVE"


def test_a_mark_set_before_the_first_sync_does_not_block_the_seed(tmp_path: Path) -> None:
    # review of 4a: the seed still runs, and it never overrides the strategy marked in the Panel
    from datetime import date

    from mrs3.pair_registry.pipeline import compute_pipeline
    from tests.test_pair_registry_records import NOW, _performance_db as registry_performance_db, _registry, _review
    from tests.test_ready_store import _store_with_periods

    con = _registry(tmp_path)
    database = registry_performance_db(tmp_path / "perf.duckdb")
    _review(database, "1", "2026-10-02 00:00:00+00", [(1, "FINALIST", 1), (2, "REJECTED", None)])
    record_finalist(con, "bybit", "GOODUSDT", "LONG", "GOODUSDT_15m_LONG_A", "RESERVE", "panel", now_ms=NOW)
    store = _store_with_periods(tmp_path / "store", "GOODUSDT", 3, 35.0)
    result = compute_pipeline(con, store, "bybit", date(2026, 10, 19), "pc", performance_db=database, now_ms=NOW + 1,
                              track=False)
    assert result["finalist_changes"] == 1
    assert {row["strategy"]: row["status"] for row in current_finalists(con, "bybit")} == {
        "GOODUSDT_15m_LONG_A": "RESERVE", "GOODUSDT_1h_LONG_B": "REJECTED"}


def test_strategy_routes(tmp_path: Path, monkeypatch, served) -> None:
    from http.client import HTTPConnection

    from tests.test_panel_main import _served_world
    from tests.test_panel_store_ready import _call

    server, _controller, service, _runner = _served_world(tmp_path, served)
    status, view = _call(server, "GET", "/api/v2/main/view?name=strategies&exchange=bybit&hide=1")
    assert status == 200 and view["options"] == {"hard": True, "top5": True, "hide": True}
    strategy = view["rows"][0]["strategy"]
    status, answer = _call(server, "POST", "/api/v2/main/finalist", {"exchange": "bybit", "strategy": strategy,
                                                                    "symbol": "GOODUSDT", "status": "RESERVE"})
    assert status == 200 and answer["status"] == "RESERVE"
    status, answer = _call(server, "POST", "/api/v2/main/finalist", {"exchange": "bybit", "strategy": strategy,
                                                                    "symbol": "GOODUSDT", "status": "BAD"})
    assert status == 400 and isinstance(answer["error"], str)
    _name, data = service.view_xlsx("strategies", "bybit", ["strategy", "status"])
    for content_type, expected in (("text/plain", 415),
                                   ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 200)):
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=30)
        try:
            connection.request("POST", "/api/v2/main/finalists.xlsx?exchange=bybit", body=data,
                               headers={"Content-Type": content_type, "Content-Length": str(len(data))})
            response = connection.getresponse()
            body = json.loads(response.read())
            assert response.status == expected, body
        finally:
            connection.close()
    assert body["written"] is True and body["changed"] == []
