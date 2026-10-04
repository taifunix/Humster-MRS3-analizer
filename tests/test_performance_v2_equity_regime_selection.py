from __future__ import annotations

from dataclasses import replace
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from hashlib import sha256
import json

import duckdb
import pandas as pd
import pytest
from openpyxl import load_workbook

from mrs3.performance_v2_equity_quality import EquitySample, calculate_equity_quality_facts
from mrs3.performance_v2_equity_regime import (
    EquityRegimeSample,
    assess_equity_regime,
    calculate_equity_regime_facts,
)
from mrs3.performance_v2_equity_regime_cache import (
    ALGORITHM_VERSION as REGIME_ALGORITHM_VERSION,
    encode_equity_regime_assessment,
    encode_equity_regime_facts,
    equity_regime_source_revision,
)
import mrs3.performance_v2_selection as selection_module
from mrs3.performance_v2_selection_review import equity_quality_snapshot_metadata
from mrs3.performance_v2_selection import (
    SelectionConfig,
    parse_selection_request,
    run_selection,
    write_selection_workbook,
)


UTC = timezone.utc
END = datetime(2026, 10, 1, tzinfo=UTC)
START = END - timedelta(days=42)


def _regime_evidence(result_id: int, state: str) -> dict[str, object]:
    samples = tuple(
        EquityRegimeSample(
            result_id,
            index,
            START + timedelta(days=index),
            Decimal("100") * (Decimal("1.01") ** index),
        )
        for index in range(43)
    )
    facts = calculate_equity_regime_facts(result_id, START, END, samples)
    if state == "WEAKENING":
        windows = {
            key: replace(facts.windows[key], v=Decimal(value), p=Decimal(value), direction="UP")
            for key, value in (("28", "8"), ("14", "9"), ("7", "9"))
        }
        facts = replace(facts, windows_28=windows["28"], windows_14=windows["14"], windows_7=windows["7"])
    elif state == "RESUMED":
        flat = replace(facts.windows["28"], v=Decimal(0), p=Decimal(0), direction="FLAT")
        facts = replace(facts, windows_28=flat)
    elif state == "STALLED":
        facts = replace(facts, held_w7_breakout=False)
    elif state == "DROP":
        facts = replace(facts, dd14=Decimal("23"))
    elif state == "NOT_EVALUATED":
        facts = calculate_equity_regime_facts(result_id, START, END, ())
    assessment = assess_equity_regime(facts)
    assert assessment.state == state
    payload = encode_equity_regime_facts(facts)
    return {
        "status": "FRESH",
        "assessment": assessment,
        "facts": facts,
        "source_revision": "a" * 64,
        "facts_sha256": sha256(payload.encode("utf-8")).hexdigest(),
        "classifier_algo_version": REGIME_ALGORITHM_VERSION,
        "equity_regime_json": encode_equity_regime_assessment(assessment),
    }


def _quality_evidence(result_id: int, score: str) -> dict[str, object]:
    facts = calculate_equity_quality_facts(
        result_id,
        START,
        END,
        (
            EquitySample(result_id, 0, START, Decimal("100")),
            EquitySample(result_id, 1, END, Decimal("110")),
        ),
    )
    facts = replace(
        facts,
        state="GROWING",
        equity_class=0,
        score12=Decimal(score),
        drawdown=Decimal("0.1"),
        peak_gap=Decimal("0.05"),
        horizon_days=28,
    )
    payload = json.dumps(facts.to_canonical_dict(), sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return {
        "facts": facts,
        "source_revision": "b" * 64,
        "facts_sha256": sha256(payload.encode("utf-8")).hexdigest(),
    }


def _row(strategy_id: int, state: str, score: str = "1") -> dict[str, object]:
    result_id = strategy_id + 1000
    quality = _quality_evidence(result_id, score)
    return {
        "strategy_id": strategy_id,
        "strategy_name": f"s{strategy_id}",
        "symbol": "BTCUSDT",
        "side": "LONG",
        "timeframe": "1h",
        "close_ma_len": strategy_id + 2,
        "order_count": 1,
        "result_id": result_id,
        "_equity_quality": quality,
        "_equity_cache": {"status": "FRESH", **quality},
        "_equity_regime_cache": _regime_evidence(result_id, state),
    }


def _request(filter_enabled: bool, rank_enabled: bool, *, top_n: int = 2):
    return parse_selection_request({
        "symbol": "BTCUSDT",
        "side": "LONG",
        "stages": [
            {"id": "filter_equity_regime", "enabled": filter_enabled, "scope": "pair_side"},
            {"id": "rank_robust_top_n", "enabled": rank_enabled, "scope": "pair_side",
             "top_n": top_n, "method": "equity_quality_v1"},
        ],
    })


@pytest.mark.parametrize(
    ("filter_enabled", "rank_enabled", "drop_reason", "unscorable_reason"),
    [
        (True, False, "DD_14_7_GTE_23", "W28_UNAVAILABLE"),
        (False, True, "EQUITY_RANK_UNRANKABLE", "EQUITY_RANK_UNRANKABLE"),
        (True, True, "DD_14_7_GTE_23", "W28_UNAVAILABLE"),
    ],
)
def test_equity_regime_filter_and_rank_toggle_matrix(
    filter_enabled: bool, rank_enabled: bool, drop_reason: str, unscorable_reason: str,
) -> None:
    rows = [
        _row(1, "GROWING", "1"),
        _row(2, "WEAKENING", "100"),
        _row(3, "RESUMED", "1000"),
        _row(4, "STALLED", "10000"),
        _row(5, "DROP", "100000"),
        _row(6, "NOT_EVALUATED", "1000000"),
    ]
    if not rank_enabled:
        for row in rows:
            row.pop("_equity_cache")
            row.pop("_equity_quality")

    result = run_selection(
        pd.DataFrame(rows),
        _request(filter_enabled, rank_enabled),
        SelectionConfig(lot_variant_redundancy_enabled=False),
    ).set_index("strategy_name")

    assert result.loc["s1", "equity_regime_state"] == "GROWING"
    assert result.loc["s5", "auto_status"] == "FILTERED"
    assert result.loc["s5", "elimination_reason"] == drop_reason
    assert result.loc["s6", "auto_status"] == "FILTERED"
    assert result.loc["s6", "elimination_reason"] == unscorable_reason
    assert result.loc["s4", "auto_status"] == "RESERVE"
    assert not result.loc["s4", "finalist"]
    if rank_enabled:
        assert result.loc["s1", "auto_status"] == "FINALIST"
        assert result.loc["s2", "auto_status"] == "FINALIST"
        assert result.loc["s3", "auto_status"] == "RESERVE"
    else:
        assert all(result.loc[f"s{strategy_id}", "finalist"] for strategy_id in (1, 2, 3))
    if not filter_enabled:
        assert result.attrs["stage_counts"]["filter_equity_regime"]["eliminated"] == 0
    assert "equity_regime_evidence" in result.attrs
    expected_filter_rejected = 2 if filter_enabled else 0
    assert result.attrs["stage_counts"]["filter_equity_regime"]["eliminated"] == expected_filter_rejected


def test_equity_rank_sorts_regime_before_r7_quality_then_strategy_id() -> None:
    request = _request(False, True, top_n=4)
    rows = [
        _row(8, "GROWING", "1"),
        _row(2, "GROWING", "1"),
        _row(1, "WEAKENING", "100000"),
        _row(3, "RESUMED", "1000000"),
    ]

    result = run_selection(
        pd.DataFrame(rows), request, SelectionConfig(lot_variant_redundancy_enabled=False)
    ).sort_values("final_rank")

    assert result["strategy_id"].tolist() == [2, 8, 1, 3]


def test_equity_rank_filters_unrankable_before_top_n_and_reports_shortfall() -> None:
    rows = [_row(1, "DROP"), _row(2, "NOT_EVALUATED"), _row(3, "STALLED")]
    result = run_selection(
        pd.DataFrame(rows), _request(False, True, top_n=4),
        SelectionConfig(lot_variant_redundancy_enabled=False),
    ).set_index("strategy_id")

    assert result["finalist"].sum() == 0
    assert result.loc[1, "elimination_reason"] == "EQUITY_RANK_UNRANKABLE"
    assert result.loc[2, "elimination_reason"] == "EQUITY_RANK_UNRANKABLE"
    assert result.loc[3, "auto_status"] == "RESERVE"
    assert result.attrs["stage_counts"]["filter_equity_regime"]["eliminated"] == 0


def test_both_equity_consumers_off_keeps_legacy_result_without_regime_work() -> None:
    request = parse_selection_request({"symbol": "BTCUSDT", "side": "LONG", "stages": []})
    row = _row(1, "GROWING")
    row.pop("_equity_regime_cache")

    result = run_selection(
        pd.DataFrame([row]), request, SelectionConfig(lot_variant_redundancy_enabled=False)
    )

    assert not any(column.startswith("equity_regime_") for column in result.columns)
    assert "equity_regime_evidence" not in result.attrs
    assert result.loc[0, "auto_status"] == "FINALIST"


def test_equity_rank_backfills_and_uses_strategy_id_when_r7_facts_are_missing() -> None:
    rows = [_row(1, "DROP"), _row(5, "GROWING"), _row(2, "GROWING"), _row(3, "WEAKENING")]
    for row in rows:
        row.pop("_equity_cache")
        row.pop("_equity_quality")
    result = run_selection(
        pd.DataFrame(rows), _request(False, True, top_n=2),
        SelectionConfig(lot_variant_redundancy_enabled=False),
    )
    finalists = result.loc[result["finalist"], "strategy_id"].tolist()
    assert finalists == [2, 5]
    assert result.set_index("strategy_id").loc[1, "elimination_reason"] == "EQUITY_RANK_UNRANKABLE"


def test_rank_only_without_r73_cache_builds_publication_snapshot_in_memory() -> None:
    metadata = {
        "result_id": 7101,
        "imported_at_utc": END + timedelta(days=1),
        "report_start_utc": START,
        "report_end_utc": END,
        "effective_start_utc": START,
        "effective_end_utc": END,
        "optimizer_source_metadata_json": None,
    }
    source_row = (7, metadata["result_id"], START, END, metadata["imported_at_utc"], START, END, None)
    connection = duckdb.connect()
    connection.execute(
        "create table strategy_equity (result_id bigint, sample_index bigint, timestamp_utc timestamptz, equity decimal(38, 12))"
    )
    connection.executemany(
        "insert into strategy_equity values (?, ?, ?, ?)",
        [
            (metadata["result_id"], index, START + timedelta(days=index), Decimal("100") * (Decimal("1.01") ** index))
            for index in range(43)
        ],
    )
    try:
        facts_by_result = selection_module._selection_equity_quality_with_memory_fallback(
            connection, [source_row], {}
        )
    finally:
        connection.close()
    r7_evidence = facts_by_result[metadata["result_id"]]
    row = _row(7, "GROWING")
    row["result_id"] = metadata["result_id"]
    row.pop("_equity_quality")
    row["_equity_cache"] = r7_evidence
    row["_equity_regime_cache"] = _regime_evidence(metadata["result_id"], "GROWING")
    request = _request(False, True, top_n=1)
    config = SelectionConfig(lot_variant_redundancy_enabled=False)
    result = run_selection(pd.DataFrame([row]), request, config)
    snapshot = equity_quality_snapshot_metadata(request, config, result)
    assert snapshot is not None
    assert set(snapshot["sources"]) == {"7"}
    assert snapshot["sources"]["7"]["result_id"] == metadata["result_id"]


def test_equity_regime_cache_hit_and_preview_miss_are_read_only(monkeypatch) -> None:
    metadata = {
        "result_id": 7001,
        "imported_at_utc": END + timedelta(days=1),
        "report_start_utc": START,
        "report_end_utc": END,
        "effective_start_utc": START,
        "effective_end_utc": END,
        "optimizer_source_metadata_json": None,
    }
    row = (7, metadata["result_id"], START, END, metadata["imported_at_utc"], START, END, None)
    evidence = _regime_evidence(metadata["result_id"], "GROWING")
    facts_json = encode_equity_regime_facts(evidence["facts"])
    facts_digest = sha256(facts_json.encode("utf-8")).hexdigest()
    source_revision = equity_regime_source_revision(metadata)
    connection = duckdb.connect()
    connection.execute(
        "create table equity_quality_metrics (result_id bigint, source_revision varchar, algo_version varchar, facts_json varchar, facts_sha256 varchar)"
    )
    connection.execute(
        "create table strategy_equity (result_id bigint, sample_index bigint, timestamp_utc timestamptz, equity decimal(38, 12))"
    )
    samples = [
        (metadata["result_id"], index, START + timedelta(days=index), Decimal("100") * (Decimal("1.01") ** index))
        for index in range(43)
    ]
    connection.executemany("insert into strategy_equity values (?, ?, ?, ?)", samples)
    monkeypatch.setattr(selection_module, "_require_equity_read_schema", lambda _: 9)
    statements: list[str] = []

    class ReadCountingConnection:
        def execute(self, sql: str, parameters: object = None):
            statements.append(sql.lower().lstrip())
            return connection.execute(sql) if parameters is None else connection.execute(sql, parameters)

    connection.execute(
        "insert into equity_quality_metrics values (?, ?, ?, ?, ?)",
        [metadata["result_id"], source_revision, REGIME_ALGORITHM_VERSION, facts_json, facts_digest],
    )
    try:
        cached = selection_module._selection_equity_regime_by_result(ReadCountingConnection(), [row])
        assert cached[metadata["result_id"]]["assessment"].state == "GROWING"
        assert len(statements) == 1 and "from strategy_equity" not in statements[0]
        assert all(statement.startswith("select") for statement in statements)
        assert connection.execute("select count(*) from equity_quality_metrics").fetchone()[0] == 1

        connection.execute("delete from equity_quality_metrics")
        statements.clear()
        preview = selection_module._selection_equity_regime_by_result(ReadCountingConnection(), [row])
        assert preview[metadata["result_id"]]["assessment"].state == "GROWING", preview[metadata["result_id"]]["assessment"].reasons
        assert any("from strategy_equity" in statement for statement in statements)
        assert all(statement.startswith("select") for statement in statements)
        assert connection.execute("select count(*) from equity_quality_metrics").fetchone()[0] == 0
    finally:
        connection.close()


def test_workbook_shows_enabled_regime_facts_and_hides_them_when_both_off(tmp_path) -> None:
    row = _row(1, "GROWING")
    enabled_request = _request(True, False)
    enabled = run_selection(
        pd.DataFrame([row]), enabled_request,
        SelectionConfig(lot_variant_redundancy_enabled=False),
    )
    enabled_path = write_selection_workbook(enabled, tmp_path / "enabled.xlsx", enabled_request)
    workbook = load_workbook(enabled_path, read_only=True)
    headers = next(workbook["All candidates"].iter_rows(values_only=True))
    assert {"Regime state", "Regime decision", "W28 direction", "PRE28 v", "Regime DD14, %", "Previous ATH W7", "ATH W28"} <= set(headers)
    workbook.close()

    disabled_request = parse_selection_request({"symbol": "BTCUSDT", "side": "LONG", "stages": []})
    disabled = run_selection(
        pd.DataFrame([row]), disabled_request,
        SelectionConfig(lot_variant_redundancy_enabled=False),
    )
    disabled["equity_regime_json"] = [row["_equity_regime_cache"]["equity_regime_json"]]
    disabled_path = write_selection_workbook(disabled, tmp_path / "disabled.xlsx", disabled_request)
    workbook = load_workbook(disabled_path, read_only=True)
    headers = next(workbook["All candidates"].iter_rows(values_only=True))
    assert {"Regime state", "W28 direction", "PRE28 v", "Regime DD14, %", "ATH W28"} <= set(headers)
    workbook.close()

    disabled_null = disabled.drop(columns="equity_regime_json")
    disabled_null_path = write_selection_workbook(
        disabled_null, tmp_path / "disabled-null.xlsx", disabled_request
    )
    workbook = load_workbook(disabled_null_path, read_only=True)
    headers = next(workbook["All candidates"].iter_rows(values_only=True))
    assert "Regime state" not in headers and "W28 direction" not in headers
    workbook.close()
