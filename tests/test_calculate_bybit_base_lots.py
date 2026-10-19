from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.filters import FilterColumn, Filters
import pytest
from mrs3.portfolio.minute_capacity import resolve_liquidity_window

from scripts.calculate_bybit_base_lots import (
    BaseLotExportError,
    _validate_xlsx_package,
    _calculation_anchor,
    AVERAGE_HEADER,
    _clear_output_filters,
    _rotate_history,
    _atomic_save_workbook,
    calculate_base_lot,
    main,
    run_export,
)
import scripts.calculate_bybit_base_lots as base_lot_script
from tests.test_portfolio_minute_capacity import write_day


def _minute(day: date, index: int, value: str = "100") -> str:
    stamp = int(datetime(day.year, day.month, day.day, tzinfo=timezone.utc).timestamp() * 1000) + index * 60_000
    return f"{stamp},1,1,1,1,{value},{value},0,1\n"


def _write_daily_files(root: Path, symbol: str, anchor: datetime) -> None:
    days = tuple(anchor.date() - timedelta(days=offset) for offset in range(16, 0, -1))
    for day in days:
        write_day(root, symbol, day, [_minute(day, 0), _minute(day, 15)])


def _workbook(path: Path) -> None:
    workbook = Workbook()
    actual = workbook.active
    actual.title = "Actual"
    actual.append(("Pair", "Listing", "Base lot"))
    actual.append(("BTCUSDT", date(2026, 1, 1), "keep until success"))
    actual.append(("ETHUSDT", date(2026, 1, 2), 1234))
    actual["D1"] = "\u0414\u0430\u0442\u0430 \u0430\u043a\u0442\u0443\u0430\u043b\u0438\u0437\u0430\u0446\u0438\u0438"
    actual["N1"] = "Other data"
    actual["N2"] = "keep"
    actual["N3"] = 999
    actual.auto_filter.ref = "B1:N3"
    actual.auto_filter.filterColumn = [
        FilterColumn(colId=0, filters=Filters(filter=["listing filter"])),
        FilterColumn(colId=1, filters=Filters(filter=["legacy size filter"])),
        FilterColumn(colId=12, filters=Filters(filter=["keep"])),
    ]
    other = workbook.create_sheet("Other")
    other["A1"] = "preserve"
    workbook.save(path)


def _empty_workbook(path: Path) -> None:
    workbook = Workbook()
    actual = workbook.active
    actual.title = "Actual"
    actual.append(("Pair", "Listing", "Base lot"))
    workbook.save(path)


def test_calculate_base_lot_uses_decimal_formula_and_floors_without_b(tmp_path: Path) -> None:
    default_lot = calculate_base_lot("123.456", "0.75", k="9", round_down_usdt=base_lot_script.DEFAULT_ROUND_DOWN_USDT)
    assert default_lot == Decimal("830")
    assert default_lot % base_lot_script.DEFAULT_ROUND_DOWN_USDT == 0
    assert calculate_base_lot(Decimal("123.456"), Decimal("0.75"), k="9", round_down_usdt="1") == Decimal("833")


@pytest.mark.parametrize("lag", [0, 23, 24, 48])
def test_each_header_date_uses_the_seven_complete_preceding_days(lag: int) -> None:
    target = date(2026, 1, 2)
    assert resolve_liquidity_window(_calculation_anchor(target, lag), lag) == tuple(
        target - timedelta(days=offset) for offset in range(7, 0, -1)
    )


def test_date_window_handles_year_boundary() -> None:
    target = date(2026, 1, 2)
    assert resolve_liquidity_window(_calculation_anchor(target, 6), 6) == tuple(
        date(2025, 12, 26) + timedelta(days=offset) for offset in range(7)
    )


def test_early_utc_run_uses_utc_date_and_previous_seven_complete_days(tmp_path: Path) -> None:
    anchor = datetime(2026, 9, 8, 0, 30, tzinfo=timezone.utc)
    workbook_path = tmp_path / "early.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    _write_daily_files(data_root, "BTCUSDT", anchor)
    _write_daily_files(data_root, "ETHUSDT", anchor)

    result = run_export(workbook_path, data_root, anchor_created_at=anchor)

    saved = load_workbook(workbook_path, data_only=False)
    assert result.actualization_date == date(2026, 9, 8)
    assert saved["Actual"]["M1"].value.date() == date(2026, 9, 8)
    assert saved["Actual"]["M2"].value == 10
    assert result.errors == ()


def test_rotation_clears_stale_style_and_comment_from_target_cell() -> None:
    from openpyxl.comments import Comment
    from openpyxl.styles import PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet["D3"] = "ERROR: old"
    sheet["D3"].fill = PatternFill(fill_type="solid", fgColor="FFC7CE")
    sheet["D3"].comment = Comment("old error", "test")
    _rotate_history(sheet, date(2026, 10, 10))

    assert sheet["D3"].value is None
    assert sheet["D3"].fill.fill_type is None
    assert sheet["D3"].comment is None


def test_atomic_save_failure_keeps_target_and_removes_temporary_file(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    target = tmp_path / "liquidity.xlsx"
    target.write_bytes(b"original")
    workbook = Workbook()

    def fail_after_partial_write(path: Path) -> None:
        Path(path).write_bytes(b"partial")
        raise OSError("simulated save failure")

    monkeypatch.setattr(workbook, "save", fail_after_partial_write)
    with pytest.raises(BaseLotExportError, match="cannot atomically save"):
        _atomic_save_workbook(workbook, target)

    assert target.read_bytes() == b"original"
    assert list(tmp_path.iterdir()) == [target]


def test_run_export_creates_ten_date_size_columns_and_preserves_pair_listing(tmp_path: Path) -> None:
    anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "liquidity.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    _write_daily_files(data_root, "BTCUSDT", anchor)
    _write_daily_files(data_root, "ETHUSDT", anchor)

    result = run_export(
        workbook_path,
        data_root,
        anchor_created_at=anchor,
        fetch_day=lambda _symbol, _day: pytest.fail("existing official files should be reused"),
    )

    assert result.rows_written == 20
    assert result.errors == ()
    saved = load_workbook(workbook_path, data_only=False)
    assert saved.sheetnames == ["Actual", "Other"]
    assert [saved["Actual"].cell(1, col).value.date() for col in range(4, 14)] == [date(2026, 8, 30) + timedelta(days=offset) for offset in range(10)]
    assert all(saved["Actual"].cell(2, col).value == 10 for col in range(4, 14))
    assert all(saved["Actual"].cell(3, col).value == 10 for col in range(4, 14))
    assert all(isinstance(saved["Actual"].cell(row, col).value, (int, float, Decimal)) for row in (2, 3) for col in range(4, 14))
    assert saved["Actual"]["A2"].value == "BTCUSDT"
    assert saved["Actual"]["B2"].value.date() == date(2026, 1, 1)
    assert saved["Actual"]["A3"].value == "ETHUSDT"
    assert saved["Actual"]["B3"].value.date() == date(2026, 1, 2)
    assert saved["Actual"].max_column == 14
    assert saved["Actual"].auto_filter.ref == "A1:N3"
    assert [column_filter.colId for column_filter in saved["Actual"].auto_filter.filterColumn] == [1, 13]
    assert saved["Actual"].auto_filter.filterColumn[0].filters.filter == ["listing filter"]
    assert saved["Actual"].auto_filter.filterColumn[1].filters.filter == ["keep"]
    assert saved["Actual"]["N3"].value == 999
    assert all(saved["Actual"].cell(row, col).number_format == "#,##0" for row in (2, 3) for col in range(4, 14))
    assert saved["Other"]["A1"].value == "preserve"


def test_rebuild_history_recalculates_all_ten_existing_dates(tmp_path: Path) -> None:
    anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "partial-history.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, anchor)
    run_export(workbook_path, data_root, anchor_created_at=anchor)

    workbook = load_workbook(workbook_path)
    workbook["Actual"].auto_filter.ref = "A1:N3"
    workbook["Actual"].auto_filter.filterColumn = [
        FilterColumn(colId=0, filters=Filters(filter=["pair"])),
        FilterColumn(colId=1, filters=Filters(filter=["listing"])),
        FilterColumn(colId=2, filters=Filters(filter=["average criterion"])),
        FilterColumn(colId=3, filters=Filters(filter=["stale output"])),
        FilterColumn(colId=13, filters=Filters(filter=["other column"])),
    ]
    from openpyxl.comments import Comment
    from openpyxl.styles import PatternFill

    for row in (2, 3):
        for col in range(4, 14):
            cell = workbook["Actual"].cell(row, col)
            cell.value = f"ERROR: stale {row}/{col}" if (row + col) % 2 else 9999
            cell.fill = PatternFill(fill_type="solid", fgColor="FFC7CE")
            cell.comment = Comment("stale history", "test")
    workbook.save(workbook_path)
    workbook.close()

    result = run_export(workbook_path, data_root, anchor_created_at=anchor, rebuild_history=True)

    saved = load_workbook(workbook_path, data_only=False)
    assert result.rows_written == 20
    assert result.errors == ()
    assert all(saved["Actual"].cell(row, col).value == 10 for row in (2, 3) for col in range(4, 14))
    assert all(isinstance(saved["Actual"].cell(row, col).value, (int, float, Decimal)) for row in (2, 3) for col in range(4, 14))
    assert all(saved["Actual"].cell(row, col).number_format == "#,##0" for row in (2, 3) for col in range(4, 14))
    assert all(saved["Actual"].cell(row, col).fill.fill_type is None for row in (2, 3) for col in range(4, 14))
    assert all(saved["Actual"].cell(row, col).comment is None for row in (2, 3) for col in range(4, 14))
    assert saved["Actual"].auto_filter.ref == "A1:N3"
    assert [item.colId for item in saved["Actual"].auto_filter.filterColumn] == [0, 1, 2, 13]
    assert [item.filters.filter for item in saved["Actual"].auto_filter.filterColumn] == [["pair"], ["listing"], ["average criterion"], ["other column"]]
    saved.close()


def test_rebuild_history_aborts_without_replacing_workbook_on_calculation_error(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "failed-rebuild.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, anchor)
    run_export(workbook_path, data_root, anchor_created_at=anchor)
    before = workbook_path.read_bytes()

    original_calculate = base_lot_script.calculate_minute_capacity

    def fail_one_date(root, symbol, **kwargs):
        if symbol == "ETHUSDT" and kwargs["anchor_created_at"].date() == date(2026, 9, 4):
            raise OSError("archive unavailable")
        return original_calculate(root, symbol, **kwargs)

    monkeypatch.setattr(base_lot_script, "calculate_minute_capacity", fail_one_date)
    with pytest.raises(BaseLotExportError, match="workbook unchanged"):
        run_export(workbook_path, data_root, anchor_created_at=anchor, rebuild_history=True)

    assert workbook_path.read_bytes() == before


def test_rebuild_history_rejects_anchor_older_than_history(tmp_path: Path) -> None:
    anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "older-anchor.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, anchor)
    run_export(workbook_path, data_root, anchor_created_at=anchor)
    before = workbook_path.read_bytes()

    with pytest.raises(BaseLotExportError, match="older than the newest date"):
        run_export(
            workbook_path,
            data_root,
            anchor_created_at=datetime(2026, 9, 7, 8, tzinfo=timezone.utc),
            rebuild_history=True,
        )

    assert workbook_path.read_bytes() == before


def test_export_rejects_fractional_round_down_step(tmp_path: Path) -> None:
    workbook_path = tmp_path / "fractional-step.xlsx"
    _workbook(workbook_path)

    with pytest.raises(BaseLotExportError, match="round_down_usdt must be a whole number"):
        run_export(workbook_path, tmp_path / "bybit", round_down_usdt="0.5")


def test_legacy_blank_history_cells_get_numeric_size_format(tmp_path: Path) -> None:
    anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "blank-format.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    workbook = load_workbook(workbook_path)
    workbook["Actual"]["C4"].number_format = "yyyy-mm-dd"
    workbook.save(workbook_path)
    workbook.close()
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, anchor)

    run_export(workbook_path, data_root, anchor_created_at=anchor)

    saved = load_workbook(workbook_path, data_only=False)
    assert saved["Actual"]["D4"].value is None
    assert saved["Actual"]["D4"].number_format == "#,##0"


def test_end_to_end_run_preserves_filter_criteria_on_column_a(tmp_path: Path) -> None:
    anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "filter-a.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    workbook = load_workbook(workbook_path)
    sheet = workbook["Actual"]
    sheet.auto_filter.ref = "A1:N3"
    sheet.auto_filter.filterColumn = [
        FilterColumn(colId=0, filters=Filters(filter=["pair criterion"])),
        FilterColumn(colId=1, filters=Filters(filter=["listing criterion"])),
        FilterColumn(colId=2, filters=Filters(filter=["old size criterion"])),
        FilterColumn(colId=13, filters=Filters(filter=["right criterion"])),
    ]
    workbook.save(workbook_path)
    workbook.close()
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, anchor)

    run_export(workbook_path, data_root, anchor_created_at=anchor)

    saved = load_workbook(workbook_path, data_only=False)
    filters = saved["Actual"].auto_filter.filterColumn
    assert [column_filter.colId for column_filter in filters] == [0, 1, 13]
    assert [column_filter.filters.filter for column_filter in filters] == [
        ["pair criterion"], ["listing criterion"], ["right criterion"]
    ]


def test_gap_rebuild_retries_overlapping_error_on_next_run(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    first_anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "liquidity.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, first_anchor)
    run_export(workbook_path, data_root, anchor_created_at=first_anchor)

    workbook = load_workbook(workbook_path)
    error_column = next(col for col in range(4, 14) if workbook["Actual"].cell(1, col).value.date() == date(2026, 9, 4))
    workbook["Actual"].cell(3, error_column).value = "ERROR: ETHUSDT: earlier temporary error"
    workbook["Actual"].cell(3, error_column).fill = base_lot_script.ERROR_FILL
    workbook.save(workbook_path)
    workbook.close()

    fail_overlap = {True}
    original_calculate = base_lot_script.calculate_minute_capacity

    def fail_one_overlap(data_root, symbol, **kwargs):
        if fail_overlap and symbol == "ETHUSDT" and kwargs["anchor_created_at"].date() == date(2026, 9, 4):
            raise OSError("temporary read failure")
        return original_calculate(data_root, symbol, **kwargs)

    monkeypatch.setattr(base_lot_script, "calculate_minute_capacity", fail_one_overlap)
    gap_anchor = datetime(2026, 9, 11, 8, tzinfo=timezone.utc)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, gap_anchor)
    run_export(workbook_path, data_root, anchor_created_at=gap_anchor)

    saved = load_workbook(workbook_path, data_only=False)
    error_column = next(col for col in range(4, 14) if saved["Actual"].cell(1, col).value.date() == date(2026, 9, 4))
    assert str(saved["Actual"].cell(3, error_column).value).startswith("ERROR: ETHUSDT:")
    saved.close()

    fail_overlap.clear()
    next_anchor = datetime(2026, 9, 12, 8, tzinfo=timezone.utc)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, next_anchor)
    run_export(workbook_path, data_root, anchor_created_at=next_anchor)

    saved = load_workbook(workbook_path, data_only=False)
    error_column = next(col for col in range(4, 14) if saved["Actual"].cell(1, col).value.date() == date(2026, 9, 4))
    assert saved["Actual"].cell(3, error_column).value == 10
    assert saved["Actual"].cell(3, error_column).fill.fill_type is None


@pytest.mark.parametrize(("initial_ref", "expected_ref"), [("A1:C3", "A1:M3"), ("A1:L3", "A1:M3")])
def test_output_filter_always_covers_through_column_m(initial_ref: str, expected_ref: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Pair"
    sheet["B1"] = "Listing"
    sheet["A3"] = "row"
    sheet["B3"] = "date"
    sheet.auto_filter.ref = initial_ref
    sheet.auto_filter.filterColumn = [
        FilterColumn(colId=0, filters=Filters(filter=["pair"])),
        FilterColumn(colId=1, filters=Filters(filter=["listing"])),
        FilterColumn(colId=2, filters=Filters(filter=["average"])),
        FilterColumn(colId=3, filters=Filters(filter=["old-size"])),
    ]

    _clear_output_filters(sheet)

    assert sheet.auto_filter.ref == expected_ref
    assert all(column_filter.colId in {0, 1, 2} for column_filter in sheet.auto_filter.filterColumn)
    assert [column_filter.filters.filter for column_filter in sheet.auto_filter.filterColumn] == [["pair"], ["listing"], ["average"]]


def test_missing_autofilter_is_created_through_column_m() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Pair"
    sheet["A2"] = "BTCUSDT"

    _clear_output_filters(sheet)

    assert sheet.auto_filter.ref == "A1:M2"


def test_symbol_error_is_written_to_current_date_column_and_other_symbols_continue(tmp_path: Path) -> None:
    anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "liquidity.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    _write_daily_files(data_root, "BTCUSDT", anchor)

    result = run_export(
        workbook_path,
        data_root,
        anchor_created_at=anchor,
        fetch_day=lambda _symbol, _day: b"invalid archive",
    )

    assert result.rows_written == 10
    assert len(result.errors) == 10
    saved = load_workbook(workbook_path, data_only=False)
    assert saved["Actual"]["M2"].value == 10
    assert str(saved["Actual"]["M3"].value).startswith("ERROR: ETHUSDT:")
    assert saved["Actual"]["M3"].fill.fgColor.rgb in {"00FFC7CE", "FFC7CE"}


def test_next_date_rotates_history_and_same_date_is_a_noop(tmp_path: Path) -> None:
    first_anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "liquidity.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    _write_daily_files(data_root, "BTCUSDT", first_anchor)
    _write_daily_files(data_root, "ETHUSDT", first_anchor)
    run_export(workbook_path, data_root, anchor_created_at=first_anchor)

    next_anchor = datetime(2026, 9, 9, 8, tzinfo=timezone.utc)
    _write_daily_files(data_root, "BTCUSDT", next_anchor)
    _write_daily_files(data_root, "ETHUSDT", next_anchor)
    rotated = run_export(workbook_path, data_root, anchor_created_at=next_anchor)
    assert rotated.rows_written == 2
    saved = load_workbook(workbook_path, data_only=False)
    headers = [saved["Actual"].cell(1, col).value.date() for col in range(4, 14)]
    assert headers == [date(2026, 8, 31) + timedelta(days=offset) for offset in range(10)]
    assert saved["Actual"].max_column == 14
    assert saved["Actual"].auto_filter.ref == "A1:N3"
    assert saved["Actual"]["N3"].value == 999
    before = workbook_path.read_bytes()

    skipped = run_export(workbook_path, data_root, anchor_created_at=next_anchor)
    assert skipped.skipped_existing_date is True
    assert workbook_path.read_bytes() == before


def test_multi_day_gap_rebuilds_ten_contiguous_dates(tmp_path: Path) -> None:
    first_anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "liquidity.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    _write_daily_files(data_root, "BTCUSDT", first_anchor)
    _write_daily_files(data_root, "ETHUSDT", first_anchor)
    run_export(workbook_path, data_root, anchor_created_at=first_anchor)
    workbook = load_workbook(workbook_path)
    sheet = workbook["Actual"]
    for col in range(4, 14):
        distinct_value = sheet.cell(1, col).value.day * 100
        sheet.cell(2, col).value = distinct_value
        sheet.cell(3, col).value = distinct_value + 1
    workbook.save(workbook_path)
    workbook.close()

    later_anchor = datetime(2026, 9, 11, 8, tzinfo=timezone.utc)
    _write_daily_files(data_root, "BTCUSDT", later_anchor)
    _write_daily_files(data_root, "ETHUSDT", later_anchor)
    result = run_export(workbook_path, data_root, anchor_created_at=later_anchor)

    saved = load_workbook(workbook_path, data_only=False)
    assert result.rows_written == 6
    assert [saved["Actual"].cell(1, col).value.date() for col in range(4, 14)] == [
        date(2026, 9, 2) + timedelta(days=offset) for offset in range(10)
    ]
    for col in range(4, 14):
        header_date = saved["Actual"].cell(1, col).value.date()
        expected_btc = header_date.day * 100 if header_date <= date(2026, 9, 8) else 10
        expected_eth = expected_btc + 1 if header_date <= date(2026, 9, 8) else 10
        assert saved["Actual"].cell(2, col).value == expected_btc
        assert saved["Actual"].cell(3, col).value == expected_eth


def test_disjoint_gap_rebuilds_without_overlap(tmp_path: Path) -> None:
    first_anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "disjoint-gap.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, first_anchor)
    run_export(workbook_path, data_root, anchor_created_at=first_anchor)

    later_anchor = datetime(2026, 9, 20, 8, tzinfo=timezone.utc)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, later_anchor)
    result = run_export(workbook_path, data_root, anchor_created_at=later_anchor)

    saved = load_workbook(workbook_path, data_only=False)
    assert result.rows_written == 20
    assert [saved["Actual"].cell(1, col).value.date() for col in range(4, 14)] == [
        date(2026, 9, 11) + timedelta(days=offset) for offset in range(10)
    ]
    assert all(saved["Actual"].cell(row, col).value == 10 for row in (2, 3) for col in range(4, 14))
    assert saved["Actual"].auto_filter.ref == "A1:N3"
    assert saved["Actual"]["N3"].value == 999


def test_gap_rebuild_fills_every_listed_date_and_leaves_prelisting_dates_blank(tmp_path: Path) -> None:
    first_anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "listing-gap.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, first_anchor)
    run_export(workbook_path, data_root, anchor_created_at=first_anchor)

    workbook = load_workbook(workbook_path)
    workbook["Actual"].append(("SOLUSDT", date(2026, 1, 1)))
    workbook["Actual"].append(("LUNAUSDT", date(2026, 9, 10)))
    workbook.save(workbook_path)
    workbook.close()
    gap_anchor = datetime(2026, 9, 11, 8, tzinfo=timezone.utc)
    for symbol in ("BTCUSDT", "ETHUSDT", "SOLUSDT", "LUNAUSDT"):
        _write_daily_files(data_root, symbol, gap_anchor)

    run_export(workbook_path, data_root, anchor_created_at=gap_anchor)

    saved = load_workbook(workbook_path, data_only=False)
    assert all(saved["Actual"].cell(4, col).value == 10 for col in range(4, 14))
    future_listing_row = 5
    for col in range(4, 14):
        header_date = saved["Actual"].cell(1, col).value.date()
        if header_date < date(2026, 9, 10):
            assert saved["Actual"].cell(future_listing_row, col).value is None
            assert saved["Actual"].cell(future_listing_row, col).fill.fill_type is None
        else:
            assert saved["Actual"].cell(future_listing_row, col).value == 10


def test_new_pair_has_blank_earlier_history_and_current_value(tmp_path: Path) -> None:
    first_anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "liquidity.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, first_anchor)
    run_export(workbook_path, data_root, anchor_created_at=first_anchor)

    workbook = load_workbook(workbook_path)
    sheet = workbook["Actual"]
    sheet.append(("SOLUSDT", date(2026, 9, 9)))
    workbook.save(workbook_path)
    workbook.close()

    next_anchor = datetime(2026, 9, 9, 8, tzinfo=timezone.utc)
    _write_daily_files(data_root, "SOLUSDT", next_anchor)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, next_anchor)
    run_export(workbook_path, data_root, anchor_created_at=next_anchor)

    saved = load_workbook(workbook_path, data_only=False)
    assert all(saved["Actual"].cell(4, col).value is None for col in range(4, 13))
    assert all(saved["Actual"].cell(4, col).fill.fill_type is None for col in range(4, 13))
    assert all(saved["Actual"].cell(4, col).comment is None for col in range(4, 13))
    assert saved["Actual"]["M4"].value == 10


def _blank_history_cell(workbook_path: Path, row: int, column: int) -> None:
    workbook = load_workbook(workbook_path)
    workbook["Actual"].cell(row, column).value = None
    workbook.save(workbook_path)
    workbook.close()


def test_same_date_rerun_fills_blank_cells_in_every_retained_date(tmp_path: Path) -> None:
    anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "liquidity.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    for symbol in ("BTCUSDT", "ETHUSDT", "SOLUSDT"):
        _write_daily_files(data_root, symbol, anchor)
    run_export(workbook_path, data_root, anchor_created_at=anchor)
    workbook = load_workbook(workbook_path)
    workbook["Actual"].append(("SOLUSDT", date(2026, 1, 1)))
    workbook.save(workbook_path)
    workbook.close()
    _blank_history_cell(workbook_path, 2, 5)

    result = run_export(workbook_path, data_root, anchor_created_at=anchor)

    saved = load_workbook(workbook_path, data_only=False)
    assert not result.skipped_existing_date
    assert saved["Actual"]["E2"].value == 10
    assert all(saved["Actual"].cell(4, col).value == 10 for col in range(4, 14))


def test_next_date_rotation_fills_blank_cells_in_retained_dates(tmp_path: Path) -> None:
    first_anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "liquidity.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, first_anchor)
    run_export(workbook_path, data_root, anchor_created_at=first_anchor)
    _blank_history_cell(workbook_path, 3, 6)
    next_anchor = datetime(2026, 9, 9, 8, tzinfo=timezone.utc)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, next_anchor)

    run_export(workbook_path, data_root, anchor_created_at=next_anchor)

    saved = load_workbook(workbook_path, data_only=False)
    assert all(saved["Actual"].cell(row, col).value == 10 for row in (2, 3) for col in range(4, 14))


def test_successful_rerun_clears_previous_error_format(tmp_path: Path) -> None:
    anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "liquidity.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    _write_daily_files(data_root, "BTCUSDT", anchor)

    failed = run_export(
        workbook_path,
        data_root,
        anchor_created_at=anchor,
        fetch_day=lambda _symbol, _day: b"invalid archive",
    )
    assert len(failed.errors) == 10
    failed_book = load_workbook(workbook_path, data_only=False)
    headers_before_retry = [failed_book["Actual"].cell(1, col).value for col in range(4, 14)]
    btc_value_before_retry = failed_book["Actual"]["M2"].value
    failed_book.close()

    _write_daily_files(data_root, "ETHUSDT", anchor)
    recovered = run_export(
        workbook_path,
        data_root,
        anchor_created_at=anchor,
        fetch_day=lambda _symbol, _day: pytest.fail("all files should be present on rerun"),
    )
    assert recovered.errors == ()
    saved = load_workbook(workbook_path, data_only=False)
    assert [saved["Actual"].cell(1, col).value for col in range(4, 14)] == headers_before_retry
    assert saved["Actual"]["M2"].value == btc_value_before_retry
    assert saved["Actual"]["M3"].value == 10
    assert saved["Actual"]["M3"].number_format == "#,##0"
    assert saved["Actual"]["M3"].fill.fill_type is None


def test_no_symbols_fails_without_replacing_workbook(tmp_path: Path) -> None:
    workbook_path = tmp_path / "empty.xlsx"
    _empty_workbook(workbook_path)
    before = workbook_path.read_bytes()

    with pytest.raises(BaseLotExportError, match="contains no symbols"):
        run_export(workbook_path, tmp_path / "bybit")

    assert workbook_path.read_bytes() == before


def test_unrecognized_legacy_layout_fails_without_replacing_workbook(tmp_path: Path) -> None:
    workbook_path = tmp_path / "unrecognized.xlsx"
    _workbook(workbook_path)
    workbook = load_workbook(workbook_path)
    workbook["Actual"]["D1"] = "unrelated"
    workbook.save(workbook_path)
    before = workbook_path.read_bytes()

    with pytest.raises(BaseLotExportError, match="recognized legacy"):
        run_export(workbook_path, tmp_path / "bybit")

    assert workbook_path.read_bytes() == before


def test_legacy_layout_with_data_in_history_columns_is_rejected(tmp_path: Path) -> None:
    workbook_path = tmp_path / "occupied-history.xlsx"
    _workbook(workbook_path)
    workbook = load_workbook(workbook_path)
    workbook["Actual"]["E4"] = "preserve"
    workbook.save(workbook_path)
    before = workbook_path.read_bytes()

    with pytest.raises(BaseLotExportError, match="columns E:M contain existing data"):
        run_export(workbook_path, tmp_path / "bybit")

    assert workbook_path.read_bytes() == before


def test_partial_date_history_is_rejected_without_replacement(tmp_path: Path) -> None:
    workbook_path = tmp_path / "partial-history.xlsx"
    _workbook(workbook_path)
    workbook = load_workbook(workbook_path)
    workbook["Actual"]["C1"] = date(2026, 9, 30)
    workbook.save(workbook_path)
    before = workbook_path.read_bytes()

    with pytest.raises(BaseLotExportError, match="incomplete or unordered"):
        run_export(workbook_path, tmp_path / "bybit")

    assert workbook_path.read_bytes() == before


def test_nonconsecutive_date_history_is_rejected_without_replacement(tmp_path: Path) -> None:
    workbook_path = tmp_path / "gapped-history.xlsx"
    _workbook(workbook_path)
    workbook = load_workbook(workbook_path)
    sheet = workbook["Actual"]
    for offset, col in enumerate(range(4, 14)):
        gap = 1 if offset > 4 else 0
        sheet.cell(1, col).value = date(2026, 9, 30) + timedelta(days=offset + gap)
    workbook.save(workbook_path)
    before = workbook_path.read_bytes()

    with pytest.raises(BaseLotExportError, match="ten consecutive dates"):
        run_export(workbook_path, tmp_path / "bybit", anchor_created_at=datetime(2026, 10, 10, tzinfo=timezone.utc))

    assert workbook_path.read_bytes() == before


def test_duplicate_date_history_is_rejected_without_replacement(tmp_path: Path) -> None:
    workbook_path = tmp_path / "duplicate-history.xlsx"
    _workbook(workbook_path)
    workbook = load_workbook(workbook_path)
    sheet = workbook["Actual"]
    for offset, col in enumerate(range(4, 14)):
        sheet.cell(1, col).value = date(2026, 9, 30) + timedelta(days=offset)
    sheet["H1"] = sheet["G1"].value
    workbook.save(workbook_path)
    before = workbook_path.read_bytes()

    with pytest.raises(BaseLotExportError, match="incomplete or unordered"):
        run_export(workbook_path, tmp_path / "bybit", anchor_created_at=datetime(2026, 10, 10, tzinfo=timezone.utc))

    assert workbook_path.read_bytes() == before


def test_older_anchor_is_rejected_without_replacement(tmp_path: Path) -> None:
    workbook_path = tmp_path / "older-anchor.xlsx"
    _workbook(workbook_path)
    workbook = load_workbook(workbook_path)
    sheet = workbook["Actual"]
    sheet["C1"] = AVERAGE_HEADER
    for offset, col in enumerate(range(4, 14)):
        sheet.cell(1, col).value = date(2026, 10, 1) + timedelta(days=offset)
        sheet.cell(2, col).value = 10
        sheet.cell(3, col).value = 20
    workbook.save(workbook_path)
    before = workbook_path.read_bytes()

    with pytest.raises(BaseLotExportError, match="older than the newest date"):
        run_export(workbook_path, tmp_path / "bybit", anchor_created_at=datetime(2026, 9, 30, tzinfo=timezone.utc))

    assert workbook_path.read_bytes() == before


@pytest.mark.parametrize("feature", ["chart", "table"])
def test_unsupported_workbook_features_fail_without_replacement(tmp_path: Path, feature: str) -> None:
    workbook_path = tmp_path / f"unsupported-{feature}.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Actual"
    sheet.append(("Pair", "Listing", "Base lot"))
    sheet.append(("BTCUSDT", date(2026, 1, 1), None))
    if feature == "chart":
        sheet.append(("x", 1, 2))
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
        sheet.add_chart(chart, "E2")
    else:
        table = Table(displayName="Pairs", ref="A1:C2")
        sheet.add_table(table)
    workbook.save(workbook_path)
    before = workbook_path.read_bytes()

    with pytest.raises(BaseLotExportError, match="ordinary xlsx"):
        run_export(workbook_path, tmp_path / "bybit")

    assert workbook_path.read_bytes() == before


def test_legacy_vml_comment_package_is_supported(tmp_path: Path) -> None:
    from openpyxl.comments import Comment

    workbook_path = tmp_path / "comment.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Actual"
    sheet.append(("Pair", "Listing", "Base lot"))
    sheet.append(("BTCUSDT", date(2026, 1, 1), None))
    sheet["A1"].comment = Comment("legacy note", "test")
    workbook.save(workbook_path)

    _validate_xlsx_package(workbook_path)


@pytest.mark.parametrize("part", ["xl/media/image1.png", "xl/pivotCache/pivotCacheDefinition1.xml", "xl/vbaProject.bin"])
def test_unsupported_package_part_fails_without_replacement(tmp_path: Path, part: str) -> None:
    workbook_path = tmp_path / "unsupported-part.xlsx"
    _empty_workbook(workbook_path)
    with ZipFile(workbook_path, "a") as package:
        package.writestr(part, b"unsupported")
    before = workbook_path.read_bytes()

    with pytest.raises(BaseLotExportError, match="ordinary xlsx"):
        run_export(workbook_path, tmp_path / "bybit")

    assert workbook_path.read_bytes() == before


def test_cli_prints_errors_to_stdout_and_returns_nonzero(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["--input", str(tmp_path / "missing.xlsx")]) == 1
    assert "ERROR:" in capsys.readouterr().out


def test_cli_forwards_rebuild_history_flag(monkeypatch: pytest.MonkeyPatch) -> None:
    captured: dict[str, object] = {}
    monkeypatch.setattr(
        base_lot_script,
        "run_export",
        lambda **kwargs: captured.update(kwargs) or SimpleNamespace(skipped_existing_date=True, actualization_date=date(2026, 10, 10)),
    )

    assert main(["--rebuild-history"]) == 0
    assert captured["rebuild_history"] is True


def test_cli_reports_partial_failure_after_saving_workbook(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "liquidity.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    _write_daily_files(data_root, "BTCUSDT", anchor)
    invalid_dir = data_root / "ETHUSDT"
    invalid_dir.mkdir(parents=True, exist_ok=True)
    for day in (anchor.date() - timedelta(days=offset) for offset in range(16, 0, -1)):
        (invalid_dir / f"ETHUSDT{day.isoformat()}_1m.parquet").write_bytes(b"invalid")

    assert main([
        "--input", str(workbook_path),
        "--data-root", str(data_root),
        "--anchor-created-at", anchor.isoformat(), "--source", "minute-files",
    ]) == 1
    output = capsys.readouterr().out
    assert "ERROR: ETHUSDT " in output
    assert "FAILED: 10 symbol/date cell(s) failed" in output
    saved = load_workbook(workbook_path, data_only=False)
    assert str(saved["Actual"]["M3"].value).startswith("ERROR:")


def test_cli_success_path_has_no_report_option_or_output(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "liquidity.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    _write_daily_files(data_root, "BTCUSDT", anchor)
    _write_daily_files(data_root, "ETHUSDT", anchor)

    assert main([
        "--input", str(workbook_path),
        "--data-root", str(data_root),
        "--anchor-created-at", anchor.isoformat(), "--source", "minute-files",
    ]) == 0
    output = capsys.readouterr().out
    assert "OK: wrote 20 dated base lots" in output
    assert "REPORT:" not in output


def test_iso_text_date_headers_are_recognized_and_rewritten_as_dates(tmp_path: Path) -> None:
    first_anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "text-headers.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, first_anchor)
    run_export(workbook_path, data_root, anchor_created_at=first_anchor)
    workbook = load_workbook(workbook_path)
    for col in range(4, 14):
        cell = workbook["Actual"].cell(1, col)
        cell.value = cell.value.date().isoformat()
    workbook.save(workbook_path)
    workbook.close()
    next_anchor = datetime(2026, 9, 9, 8, tzinfo=timezone.utc)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, next_anchor)

    run_export(workbook_path, data_root, anchor_created_at=next_anchor)

    saved = load_workbook(workbook_path, data_only=False)["Actual"]
    headers = [saved.cell(1, col).value for col in range(4, 14)]
    assert all(isinstance(value, datetime) for value in headers)
    assert headers[-1].date() == date(2026, 9, 9)
    assert all(saved.cell(row, col).value == 10 for row in (2, 3) for col in range(4, 14))


@pytest.mark.parametrize("text", ["2026-9-1", "01.09.2026", "2026-09-01 00:00", "x"])
def test_non_iso_text_headers_are_not_dates(text: str) -> None:
    assert base_lot_script._date_header(text) is None


def _set_row_values(workbook_path: Path, row: int, values: list[object]) -> None:
    workbook = load_workbook(workbook_path)
    for offset, value in enumerate(values):
        workbook["Actual"].cell(row, 4 + offset).value = value
    workbook.save(workbook_path)
    workbook.close()


def test_average_column_holds_the_floored_mean_of_the_ten_dates(tmp_path: Path) -> None:
    anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "liquidity.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, anchor)
    run_export(workbook_path, data_root, anchor_created_at=anchor)
    saved = load_workbook(workbook_path, data_only=False)["Actual"]
    assert saved["C1"].value == AVERAGE_HEADER
    assert saved["C2"].value == 10 and saved["C3"].value == 10
    assert saved["N3"].value == 999

    _set_row_values(workbook_path, 2, [100, 110, 120, 130, 140, 150, 160, 170, 180, 190])
    result = run_export(workbook_path, data_root, anchor_created_at=anchor)

    saved = load_workbook(workbook_path, data_only=False)["Actual"]
    assert not result.skipped_existing_date
    assert saved["C2"].value == 140
    assert saved["C2"].number_format == "#,##0"


def test_average_ignores_error_cells_and_is_recomputed_after_rotation(tmp_path: Path) -> None:
    first_anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "liquidity.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, first_anchor)
    run_export(workbook_path, data_root, anchor_created_at=first_anchor)
    _set_row_values(workbook_path, 2, [1000, 1000, 1000, 1000, 1000, 1000, 1000, 1000, 1000, 1000])
    next_anchor = datetime(2026, 9, 9, 8, tzinfo=timezone.utc)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, next_anchor)

    run_export(workbook_path, data_root, anchor_created_at=next_anchor)

    saved = load_workbook(workbook_path, data_only=False)["Actual"]
    assert saved["C2"].value == 900
    assert base_lot_script._average_lot([100, "ERROR: x", None, 205, True], Decimal("10")) == 150
    assert base_lot_script._average_lot([None, "ERROR: x"], Decimal("10")) is None


def test_foreign_data_in_the_average_column_fails_without_replacement(tmp_path: Path) -> None:
    anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "liquidity.xlsx"
    data_root = tmp_path / "bybit"
    _workbook(workbook_path)
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, anchor)
    run_export(workbook_path, data_root, anchor_created_at=anchor)
    workbook = load_workbook(workbook_path)
    workbook["Actual"]["C1"] = "My notes"
    workbook.save(workbook_path)
    workbook.close()
    original = workbook_path.read_bytes()

    with pytest.raises(BaseLotExportError, match="column C is reserved"):
        run_export(workbook_path, data_root, anchor_created_at=datetime(2026, 9, 9, 8, tzinfo=timezone.utc))
    assert workbook_path.read_bytes() == original


def test_previous_layout_with_dates_in_c_to_l_moves_right_and_gets_the_average(tmp_path: Path) -> None:
    anchor = datetime(2026, 9, 8, 8, tzinfo=timezone.utc)
    workbook_path = tmp_path / "previous.xlsx"
    data_root = tmp_path / "bybit"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Actual"
    sheet.append(("Pair", "Listing", *[(date(2026, 8, 30) + timedelta(days=k)).isoformat() for k in range(10)]))
    sheet.append(("BTCUSDT", date(2026, 1, 1), *[100 + 10 * k for k in range(10)]))
    sheet.append(("ETHUSDT", date(2026, 1, 2), *([None] * 10)))
    sheet["N1"] = "Other data"
    sheet["N3"] = 999
    from openpyxl.worksheet.dimensions import ColumnDimension

    group = ColumnDimension(sheet, index="C", width=21, customWidth=True)
    group.min, group.max = 3, 12
    sheet.column_dimensions["C"] = group
    sheet.auto_filter.ref = "A1:L3"
    sheet.auto_filter.filterColumn = [FilterColumn(colId=1, filters=Filters(filter=["listing"]))]
    workbook.save(workbook_path)
    workbook.close()
    for symbol in ("BTCUSDT", "ETHUSDT"):
        _write_daily_files(data_root, symbol, anchor)

    result = run_export(workbook_path, data_root, anchor_created_at=anchor)

    saved = load_workbook(workbook_path, data_only=False)["Actual"]
    assert not result.skipped_existing_date
    assert saved["C1"].value == AVERAGE_HEADER
    assert [saved.cell(1, col).value.date() for col in range(4, 14)] == [date(2026, 8, 30) + timedelta(days=k) for k in range(10)]
    assert [saved.cell(2, col).value for col in range(4, 14)] == [100 + 10 * k for k in range(10)]
    assert all(saved.cell(3, col).value == 10 for col in range(4, 14))
    assert saved["C2"].value == 140 and saved["C3"].value == 10
    assert saved["N3"].value == 999
    assert saved.auto_filter.ref == "A1:N3"
    assert [f.colId for f in saved.auto_filter.filterColumn] == [1]
    import xml.etree.ElementTree as ET

    with ZipFile(workbook_path) as package:
        root = ET.fromstring(package.read("xl/worksheets/sheet1.xml"))
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    spans = [(int(c.get("min")), int(c.get("max")), float(c.get("width"))) for c in root.findall("m:cols/m:col", ns)]
    covered = [col for lo, hi, _ in spans for col in range(lo, hi + 1)]
    assert len(covered) == len(set(covered)), "overlapping <col> ranges"
    widths = {col: w for lo, hi, w in spans for col in range(lo, hi + 1)}
    assert all(widths[col] == 21 for col in range(3, 14))
