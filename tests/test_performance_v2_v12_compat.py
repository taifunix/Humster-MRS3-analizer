"""Old-database compatibility of the additive v11 -> v12 migration (spec 2026-10-12 new report format §7)."""

from __future__ import annotations

from pathlib import Path

import duckdb
from openpyxl import load_workbook

from mrs3.performance_v2_selection import (
    SelectionConfig,
    load_selection_candidates,
    parse_selection_request,
    run_selection,
    write_selection_workbook,
)
from mrs3.performance_v2_selection_review import apply_prior_rejected
from mrs3.performance_v2_store import (
    database_identity,
    initialize_performance_v2,
    require_performance_v2,
    require_performance_v2_readable,
)
from tests.test_panel_performance_v2 import _db


_REPORT_TABLES = ("strategy_report_metrics", "strategy_report_regimes", "strategy_report_monthly")


def _selection_outputs(connection: duckdb.DuckDBPyConnection, path: Path) -> tuple[list[tuple], list[tuple]]:
    request = parse_selection_request({"symbol": "BTCUSDT", "side": "LONG", "stages": []})
    candidates = load_selection_candidates(connection, request, SelectionConfig())
    selected = run_selection(apply_prior_rejected(connection, candidates), request)
    decisions = sorted(
        (int(row.strategy_id), str(row.auto_status), None if row.elimination_reason is None else str(row.elimination_reason))
        for row in selected.itertuples()
    )
    workbook = load_workbook(write_selection_workbook(selected, path, request), read_only=True)
    cells = [tuple(row) for row in workbook["All candidates"].iter_rows(values_only=True)]
    workbook.close()
    return decisions, cells


def test_v11_database_keeps_rows_and_outputs_after_v12_migration(tmp_path: Path) -> None:
    connection, _result_id = _db(tmp_path / "db", legacy_v11=True)
    try:
        assert require_performance_v2_readable(connection) == 11
        tables = [row[0] for row in connection.execute(
            "select table_name from information_schema.tables where table_schema = 'main' order by 1"
        ).fetchall()]
        outputs_before = _selection_outputs(connection, tmp_path / "before.xlsx")  # also warms window_metrics
        rows_before = {table: connection.execute(f"select * from {table} order by all").fetchall() for table in tables}

        initialize_performance_v2(connection, evaluation_protocol="v2", exchange="Bybit")

        require_performance_v2(connection)
        assert database_identity(connection) == ("v1", None)  # identity applies only at creation
        for table, before in rows_before.items():
            after = connection.execute(f"select * from {table} order by all").fetchall()
            if table == "schema_info":
                assert dict(after) == {**dict(before), "schema_version": "12"}
                continue
            width = len(before[0]) if before else 0
            assert [row[:width] for row in after] == before, table
            assert all(value is None for row in after for value in row[width:]), table
        for table in _REPORT_TABLES:
            assert connection.execute(f"select count(*) from {table}").fetchone() == (0,)
        assert _selection_outputs(connection, tmp_path / "after.xlsx") == outputs_before
    finally:
        connection.close()


def test_portfolio_source_records_ignore_unset_v12_columns(tmp_path: Path) -> None:
    from decimal import Decimal

    from mrs3.portfolio.input import _canonical_value, _records_for_ids

    connection, result_id = _db(tmp_path / "db", legacy_v11=True)
    try:
        strategy_id = connection.execute("select strategy_id from strategy_results where result_id = ?", [result_id]).fetchone()[0]
        lookups = (("strategies", "strategy_id", strategy_id), ("strategy_orders", "strategy_id", strategy_id),
                   ("strategy_results", "result_id", result_id))
        before = {table: _records_for_ids(connection, table, column, [value]) for table, column, value in lookups}

        initialize_performance_v2(connection)

        assert {table: _records_for_ids(connection, table, column, [value]) for table, column, value in lookups} == before
        connection.execute(
            "update strategy_results set sizing_my_fix_balance = 10000, sizing_balance_percentage_short = 100, "
            "sizing_risk_short = 1, tester_version = '2026/10/10', report_format = '2026-10'"
        )
        connection.execute("update strategy_orders set open_ma_type = 'SMA', open_ma_source = 'ohlc4'")
        result = _records_for_ids(connection, "strategy_results", "result_id", [result_id])[0]
        order = _records_for_ids(connection, "strategy_orders", "strategy_id", [strategy_id])[0]
        assert result["sizing_my_fix_balance"] == Decimal("10000")
        assert (order["open_ma_type"], order["open_ma_source"]) == ("SMA", "ohlc4")
        _canonical_value(result)
        _canonical_value(order)
    finally:
        connection.close()
