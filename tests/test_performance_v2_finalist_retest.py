from __future__ import annotations

from copy import deepcopy
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path
from hashlib import sha256
from io import BytesIO
import json
import math
from tempfile import TemporaryDirectory
import pytest
import pandas as pd
from openpyxl import Workbook, load_workbook

from mrs3.performance_v2_finalist_retest import (
    canonical_json,
    canonical_digest,
    cohort_digest,
    build_finalist_retest_manifest,
    FinalistRetestCohort,
    apply_finalist_retest_outcomes,
    freeze_finalist_cohort,
    review_key,
    combined_control_workbook_bytes,
    validate_combined_control_workbook,
    FinalistRetestError,
    import_combined_control_workbook,
    finalist_retest_config_digest,
    _control_reason,
)
from mrs3.performance_v2_store import initialize_performance_v2
from mrs3.performance_v2_selection import (
    SELECTION_REASON_ALIASES, SelectionRequest, parse_selection_request, write_selection_workbook,
)
import duckdb
from mrs3.panel_testing import render_strategy
from mrs3.panel_fast_strategy_test import LocalSingleModeStrategyTestService
from mrs3.panel_strategy_batch import validate_strategy_manifest


PORTFOLIO_ARITHMETIC_FIXTURE = Path(__file__).parent / "fixtures" / "portfolio" / "source_sizing_arithmetic.json"


def test_selection_reason_aliases_are_shared_and_reversible() -> None:
    assert SELECTION_REASON_ALIASES == {"PARETO_PLATEAU_POINTS_PER_ORDER": "PARETO_PL_PTS_PER_ORDER"}
    for reason, alias in SELECTION_REASON_ALIASES.items():
        assert _control_reason(alias) == reason
        assert _control_reason(f"7. {alias}") == reason


def _candidate_workbook(rows, request: SelectionRequest | None = None) -> bytes:
    request = request or parse_selection_request({"symbol": "BTCUSDT", "side": "LONG", "stages": []})
    normalized = []
    reviews = {}
    for source in rows:
        row = dict(source)
        row.update({
            "strategy_id": source.get("strategy_id"), "result_id": source.get("result_id"),
            "strategy_name": source.get("strategy_name", f"strategy-{source.get('strategy_id')}"),
            "symbol": source.get("symbol", request.symbol), "side": source.get("side", request.side),
            "timeframe": source.get("timeframe", "1h"), "order_count": source.get("order_count", 1),
            "close_ma_len": source.get("close_ma_len", 5),
            "auto_status": source.get("auto_status", source.get("user_status", "FINALIST")),
            "finalist": source.get("finalist", source.get("user_status", "FINALIST") in {"FINALIST", "RESERVE"}),
            "final_rank": source.get("final_rank", source.get("auto_rank")),
            "final_score": source.get("final_score", source.get("score")),
            "elimination_reason": source.get("elimination_reason", source.get("auto_reason")),
            "auto_analog_of_strategy_id": source.get("auto_analog_of_strategy_id"),
            "prior_rejected": False,
        })
        normalized.append(row)
        reviews[int(row["strategy_id"])] = {
            "user_status": source.get("user_status", "FINALIST"), "user_rank": source.get("user_rank"),
            "user_analog_of_strategy_id": source.get("user_analog_of_strategy_id"), "comment": source.get("comment"),
        }
    with TemporaryDirectory() as directory:
        path = write_selection_workbook(
            pd.DataFrame(normalized), Path(directory) / "candidates.xlsx", request,
            {"workbook_schema_version": "2"}, reviews,
        )
        return path.read_bytes()


def test_weighted_template_only_changes_name_and_mrs_priority() -> None:
    root = Path(__file__).resolve().parents[1]
    original_text = (root / "templates/strategies/retest-mrs3/base.json").read_text(encoding="utf-8")
    weighted_text = (root / "templates/strategies/portfolio-weighted-mrs/base.json").read_text(encoding="utf-8")
    original = json.loads(original_text)
    weighted = json.loads(weighted_text)

    assert original["mrs"] is None
    assert weighted["mrs"] == {"position_priority": 3}
    assert type(weighted["mrs"]["position_priority"]) is int
    assert weighted["exchange"]["use_upnl"] is True
    assert weighted["exchange"]["use_frozen_balance"] is True
    assert weighted["basic"]["use_fix"] is False
    assert weighted["basic"]["use_long"] is True
    assert weighted["basic"]["use_short"] is False
    assert isinstance(weighted["name"], str) and weighted["name"]
    assert weighted["name"] != original["name"]

    def keys(value: object):
        if isinstance(value, dict):
            for key, nested in value.items():
                yield key
                yield from keys(nested)
        elif isinstance(value, list):
            for nested in value:
                yield from keys(nested)

    assert not set(keys(weighted)).intersection({"open_positions_limiter", "k", "size_composition_vector"})

    normalized = deepcopy(original)
    normalized["name"] = weighted["name"]
    normalized["mrs"] = weighted["mrs"]
    assert weighted == normalized

    original_filename, rendered_original = render_strategy(original_text, "ONUSDT", "LONG")
    weighted_filename, rendered_weighted = render_strategy(weighted_text, "ONUSDT", "LONG")
    assert original_filename == f'{original["name"]}.json'
    assert weighted_filename == f'{weighted["name"]}.json'
    assert rendered_original["basic"]["symbol"] == "ONUSDT"
    assert rendered_weighted["basic"]["symbol"] == "ONUSDT"
    assert rendered_original["basic"]["use_long"] is True
    assert rendered_original["basic"]["use_short"] is False
    assert rendered_weighted["basic"]["use_long"] is True
    assert rendered_weighted["basic"]["use_short"] is False
    assert rendered_weighted["mrs"]["position_priority"] == 3


def test_canonical_json_encodes_typed_values_and_omits_runtime_provenance() -> None:
    value = {
        "b": Decimal("1.20"),
        "a": date(2026, 1, 2),
        "at_utc": datetime(2026, 1, 2, 3, 4, tzinfo=UTC),
        "local_path": "C:/private/inbox",
        "n": 7,
    }

    assert canonical_json(value) == '{"a":"2026-01-02","b":"1.20","n":7}'


def test_dynamic_source_sizing_arithmetic_uses_per_cycle_basis_and_reconstructs() -> None:
    values = json.loads(PORTFOLIO_ARITHMETIC_FIXTURE.read_text(encoding="utf-8"))
    assert values["use_fix"] is False
    assert values["rows"][0]["S"] != values["rows"][1]["S"]
    normalized = [row["delta"] / row["S"] for row in values["rows"]]
    increments = [value * values["x"] for value in normalized]
    assert normalized == [0.1, 0.1]
    assert increments == [5.0, 5.0]
    for row, increment in zip(values["rows"], increments, strict=True):
        assert math.isclose(increment * row["S"] / values["x"], row["delta"], rel_tol=1e-8, abs_tol=1e-8)


def test_cohort_digest_is_member_order_and_timestamp_independent() -> None:
    first = [
        {"strategy_id": 2, "result_id": 12, "status": "RESERVE", "rank": 1, "created_at_utc": "a"},
        {"strategy_id": 1, "result_id": 11, "status": "FINALIST", "rank": 2, "created_at_utc": "b"},
    ]
    second = [
        {"strategy_id": 1, "result_id": 11, "status": "FINALIST", "rank": 2, "created_at_utc": "changed"},
        {"strategy_id": 2, "result_id": 12, "status": "RESERVE", "rank": 1, "created_at_utc": "changed"},
    ]

    assert cohort_digest(first) == cohort_digest(second)


def test_review_key_is_deterministic_from_workbook_hash_and_selection_run() -> None:
    assert review_key("a" * 64, "selection-1") == review_key("a" * 64, "selection-1")
    assert review_key("a" * 64, "selection-1") != review_key("b" * 64, "selection-1")


def test_finalist_retest_config_digest_changes_with_template_content(tmp_path: Path) -> None:
    template = tmp_path / "base.json"
    template.write_text('{"basic":{"risk_long":1}}', encoding="utf-8")
    first = finalist_retest_config_digest({"LONG": str(template)})
    template.write_text('{"basic":{"risk_long":2}}', encoding="utf-8")

    assert finalist_retest_config_digest({"LONG": str(template)}) != first


def test_freeze_uses_effective_user_status_and_listing_warmup_without_writes() -> None:
    with duckdb.connect(":memory:") as connection:
        initialize_performance_v2(connection)
        now = datetime(2026, 1, 1, tzinfo=UTC)
        connection.execute(
            """insert into strategies values
            (1, 'final', 'BTCUSDT', 'LONG', '1h', 5, 1, 'run', 'candidate-1', 'ACTIVE', 11, ?, ?),
            (2, 'reserve', 'ETHUSDT', 'SHORT', '4h', 7, 1, 'run', 'candidate-2', 'ACTIVE', 12, ?, ?),
            (3, 'filtered', 'SOLUSDT', 'LONG', '1h', 5, 1, 'run', 'candidate-3', 'ACTIVE', 13, ?, ?)""",
            [now, now, now, now, now, now],
        )
        connection.execute(
            """insert into strategy_results
            (result_id, strategy_id, report_start_utc, report_end_utc, exchange, commission_rate,
             initial_balance, final_balance, imported_at_utc)
            values (11, 1, '2026-01-01', '2026-02-01', 'Bybit', .0004, 100, 101, ?),
                   (12, 2, '2026-01-01', '2026-02-01', 'Bybit', .0004, 100, 101, ?),
                   (13, 3, '2026-01-01', '2026-02-01', 'Bybit', .0004, 100, 101, ?)""",
            [now, now, now],
        )
        connection.execute("insert into analysis_plateaus values ('run', 'P1', 4, 80)")
        connection.execute(
            """insert into strategy_orders
            (strategy_id, order_id, open_ma_len, open_multiplier, shift_bp, lot_x, analysis_run_id, plateau_id, base_point_trades)
            values (1, 1, 7, .995, 50, 1, 'run', 'P1', 20),
                   (2, 1, 9, 1.005, 50, 1, 'run', 'P1', 20),
                   (3, 1, 7, .995, 50, 1, 'run', 'P1', 20)"""
        )
        for symbol, side, run_id in (("BTCUSDT", "LONG", "run-1"), ("ETHUSDT", "SHORT", "run-2"), ("SOLUSDT", "LONG", "run-3")):
            connection.execute(
                """insert into selection_runs
                (selection_run_id, database_instance_id, symbol, side, selection_contract_version,
                 request_json, request_sha256, config_json, config_sha256, candidate_count,
                 representative_count, auto_finalist_count, top_n, workbook_sha256, created_at_utc)
                values (?, (select value from schema_info where key='database_instance_id'), ?, ?, 'v1', '{}', ?, '{}', ?, 1, 1, 1, 20, ?, ?)""",
                [run_id, symbol, side, "a" * 64, "b" * 64, "c" * 64, now],
            )
        connection.execute(
            """insert into selection_results
            (selection_run_id, strategy_id, result_id_at_selection, auto_status, auto_score, auto_rank,
             auto_reason, analog_group_key, auto_analog_of_strategy_id, prior_rejected, stage_trace_json)
            values ('run-1', 1, 11, 'FINALIST', 90, 1, null, null, null, false, '{}'),
                   ('run-2', 2, 12, 'RESERVE', 80, 1, null, null, null, false, '{}'),
                   ('run-3', 3, 13, 'FILTERED', 70, null, null, null, null, false, '{}')"""
        )
        connection.execute(
            """insert into selection_review_imports
            values ('review-2', 'run-2', ?, ?, 1)""",
            ["d" * 64, now],
        )
        connection.execute(
            """insert into selection_review_rows values ('review-2', 2, 'RESERVE', 4, null, 'keep')"""
        )
        connection.execute(
            """insert into selection_review_imports
               values ('review-1', 'run-1', ?, ?, 1)""",
            ["e" * 64, now],
        )
        connection.execute(
            """insert into selection_review_rows values ('review-1', 1, 'FINALIST', 1, null, 'keep')"""
        )
        before = connection.execute("select count(*) from selection_review_imports").fetchone()

        finalist = freeze_finalist_cohort(
            connection,
            test_start="2026-01-01",
            test_end="2026-06-01",
            listing_dates={"BTCUSDT": date(2025, 12, 1), "ETHUSDT": date(2026, 2, 1), "SOLUSDT": date(2025, 12, 1)},
        )
        assert [item["strategy_id"] for item in finalist.members] == [1]
        assert finalist.members[0]["effective_start"].date() == date(2026, 1, 1)
        assert finalist.excluded_count == 0

        with_reserve = freeze_finalist_cohort(
            connection,
            test_start="2026-01-01",
            test_end="2026-06-01",
            include_reserve=True,
            listing_dates={"BTCUSDT": date(2025, 12, 1), "ETHUSDT": date(2026, 2, 1), "SOLUSDT": date(2025, 12, 1)},
        )
        assert [item["strategy_id"] for item in with_reserve.members] == [1, 2]
        assert with_reserve.members[1]["effective_rank"] == 4
        assert connection.execute("select count(*) from selection_review_imports").fetchone() == before


def test_manifest_has_one_common_period_and_native_strategy_provenance(tmp_path) -> None:
    with duckdb.connect(":memory:") as connection:
        initialize_performance_v2(connection)
        now = datetime(2026, 1, 1, tzinfo=UTC)
        connection.execute(
            "insert into strategies values (1, 'final', 'BTCUSDT', 'LONG', '1h', 5, 1, 'run', 'candidate', 'ACTIVE', 11, ?, ?)",
            [now, now],
        )
        connection.execute(
            "insert into strategy_results (result_id, strategy_id, report_start_utc, report_end_utc, exchange, commission_rate, initial_balance, final_balance, imported_at_utc) values (11, 1, '2026-01-01', '2026-06-01', 'Bybit', .0004, 100, 101, ?)",
            [now],
        )
        connection.execute("insert into analysis_plateaus values ('run', 'P1', 4, 80)")
        connection.execute(
            "insert into strategy_orders (strategy_id, order_id, open_ma_len, open_multiplier, shift_bp, lot_x, analysis_run_id, plateau_id, base_point_trades) values (1, 1, 7, .995, 50, 1, 'run', 'P1', 20)"
        )
        connection.execute(
            """insert into selection_runs
            (selection_run_id, database_instance_id, symbol, side, selection_contract_version,
             request_json, request_sha256, config_json, config_sha256, candidate_count,
             representative_count, auto_finalist_count, top_n, workbook_sha256, created_at_utc)
            values ('selection-1', (select value from schema_info where key='database_instance_id'), 'BTCUSDT', 'LONG', 'v1', '{}', ?, '{}', ?, 1, 1, 1, 20, ?, ?)""",
            ["a" * 64, "b" * 64, "c" * 64, now],
        )
        connection.execute(
            "insert into selection_results (selection_run_id, strategy_id, result_id_at_selection, auto_status, auto_rank, prior_rejected, stage_trace_json) values ('selection-1', 1, 11, 'FINALIST', 1, false, '{}')"
        )
        connection.execute(
            """insert into selection_review_imports
               values ('review-1', 'selection-1', ?, ?, 1)""",
            ["d" * 64, now],
        )
        connection.execute(
            """insert into selection_review_rows values ('review-1', 1, 'FINALIST', 1, null, 'keep')"""
        )
        batch = build_finalist_retest_manifest(
            connection,
            {"LONG": Path("templates/strategies/retest-mrs3/base.json")},
            tmp_path / "output",
            test_start="2026-01-01",
            test_end="2026-06-01",
            listing_dates={"BTCUSDT": date(2025, 12, 1)},
            job_id="bulk-1",
        )

    manifest = json.loads(batch.manifest_path.read_text(encoding="utf-8"))
    assert batch.strategy_count == 1
    assert manifest["scope"] == "FINALIST"
    assert manifest["finalist_retest"]["job_id"] == "bulk-1"
    assert manifest["finalist_retest"]["members"][0]["effective_end"] == "2026-06-01T00:00:00Z"
    validated = validate_strategy_manifest(batch.manifest_path)
    assert validated.analysis_run_id == "bulk-1"
    assert manifest["candidate_diagnostics"] == {
        "candidate": {
            "order_count": 1,
            "orders": [{"order_id": 1, "plateau_id": "P1", "plateau_point_count": 4,
                        "base_point_trades": 20, "plateau_total_trades": 80}],
        },
    }
    assert LocalSingleModeStrategyTestService._require_diagnostics(validated, ("final",)) == manifest["candidate_diagnostics"]
    strategy = json.loads((batch.strategies_path / "final.json").read_text(encoding="utf-8"))
    assert strategy["name"] == "final"
    assert strategy["basic"]["use_fix"] is False
    assert strategy["exchange"]["use_upnl"] is True
    assert strategy["exchange"]["use_frozen_balance"] is True
    assert strategy["basic"]["my_fix_balance"] == 1000.0


def test_import_outcomes_replace_successful_member_and_isolate_failure() -> None:
    with duckdb.connect(":memory:") as connection:
        initialize_performance_v2(connection)
        now = datetime(2026, 1, 1, tzinfo=UTC)
        connection.execute(
            """insert into strategies values
            (1, 'one', 'BTCUSDT', 'LONG', '1h', 5, 1, 'run', 'one', 'ACTIVE', 11, ?, ?),
            (2, 'two', 'BTCUSDT', 'LONG', '1h', 5, 1, 'run', 'two', 'ACTIVE', 12, ?, ?)""",
            [now, now, now, now],
        )
        connection.execute(
            """insert into strategy_results
            (result_id, strategy_id, report_start_utc, report_end_utc, exchange, commission_rate,
             initial_balance, final_balance, imported_at_utc)
            values (11, 1, '2026-01-01', '2026-02-01', 'Bybit', .0004, 100, 101, ?),
                   (12, 2, '2026-01-01', '2026-02-01', 'Bybit', .0004, 100, 102, ?)""",
            [now, now],
        )
        cohort = FinalistRetestCohort(
            "FINALIST", "2026-01-01", "2026-06-01", 120,
            (
                {"strategy_id": 1, "strategy_name": "one", "result_id": 11},
                {"strategy_id": 2, "strategy_name": "two", "result_id": 12},
            ), (), "c" * 64,
        )
        # v4 keeps one result row per strategy during REPLACE; a successful
        # importer therefore reports the same durable Result ID after facts
        # have been updated in place.
        connection.execute("update strategy_results set final_balance = 120 where result_id = 11")
        result = apply_finalist_retest_outcomes(
            connection, cohort,
            {1: {"status": "SUCCESS", "new_result_id": 11}, 2: {"status": "FAILED", "reason": "REPORT_MISSING"}},
            job_id="bulk-1",
        )

        assert result.status == "COMMITTED"
        assert result.success_count == 1
        assert result.failures == ({"strategy_id": 2, "strategy_name": "two", "reason": "REPORT_MISSING"},)
        assert connection.execute("select strategy_id, current_result_id from strategies order by strategy_id").fetchall() == [(1, 11), (2, 12)]


def test_combined_control_workbook_has_fixed_sheets_and_group_local_ranks() -> None:
    data = combined_control_workbook_bytes(
        [{"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 1, "result_id": 11,
          "user_status": "FINALIST", "user_rank": 1, "auto_status": "FINALIST"},
         {"symbol": "ETHUSDT", "side": "LONG", "strategy_id": 2, "result_id": 12,
          "user_status": "FINALIST", "user_rank": 1, "auto_status": "FINALIST"}],
        [{"symbol": "BTCUSDT", "side": "LONG", "frozen_count": 1}],
        [{"symbol": "DOGEUSDT", "side": "SHORT", "strategy_id": 3, "reason": "REPORT_MISSING"}],
        {"ranking_scope": "RETEST_COHORT", "bulk_retest_job_id": "bulk-1"},
        candidate_workbook=_candidate_workbook([
            {"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 1, "result_id": 11, "user_status": "FINALIST", "user_rank": 1},
            {"symbol": "ETHUSDT", "side": "LONG", "strategy_id": 2, "result_id": 12, "user_status": "FINALIST", "user_rank": 1},
        ]),
    )
    metadata, candidates, groups, failures = validate_combined_control_workbook(data)
    assert metadata["ranking_scope"] == "RETEST_COHORT"
    assert len(candidates) == 2 and len(groups) == 1 and len(failures) == 1

    duplicate = combined_control_workbook_bytes(
        [{"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 1, "result_id": 11,
          "user_status": "FINALIST", "user_rank": 1},
         {"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 2, "result_id": 12,
          "user_status": "FINALIST", "user_rank": 1}], metadata={},
        candidate_workbook=_candidate_workbook([
            {"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 1, "result_id": 11, "user_status": "FINALIST", "user_rank": 1},
            {"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 2, "result_id": 12, "user_status": "FINALIST", "user_rank": 1},
        ]),
    )
    with pytest.raises(FinalistRetestError, match="CONTROL_DUPLICATE_RANK"):
        validate_combined_control_workbook(duplicate)


def test_old_narrow_and_duplicate_alias_candidate_headers_are_rejected() -> None:
    workbook = Workbook()
    workbook.active.title = "Candidates"
    workbook["Candidates"].append([
        "Pair", "Direction", "Strategy ID", "Result ID", "User Status", "User Rank", "RETEST", "Comment",
        "Auto Status", "Auto Rank", "Auto Analog Of ID", "Analog Of ID", "Auto Reason", "Score",
    ])
    workbook.create_sheet("Groups").append(list(("Pair", "Direction", "Frozen Count", "Success Count", "Failure Count", "Auto Status Count")))
    workbook.create_sheet("Retest Failures").append(list(("Pair", "Direction", "Strategy ID", "Strategy", "Result ID", "Reason")))
    metadata = workbook.create_sheet("_MRS_SELECTION_META")
    metadata.sheet_state = "veryHidden"
    narrow_io = BytesIO()
    workbook.save(narrow_io)
    with pytest.raises(FinalistRetestError, match="CONTROL_SCHEMA_MISMATCH"):
        validate_combined_control_workbook(narrow_io.getvalue())

    rich_bytes = combined_control_workbook_bytes(
        [{"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 1, "result_id": 11,
          "user_status": "FINALIST", "user_rank": 1}],
        metadata={},
        candidate_workbook=_candidate_workbook([
            {"strategy_id": 1, "result_id": 11, "user_status": "FINALIST", "user_rank": 1},
        ]),
    )
    rich = load_workbook(BytesIO(rich_bytes))
    candidate_sheet = rich["Candidates"]
    candidate_sheet.cell(1, candidate_sheet.max_column + 1).value = "Pair"
    duplicate_io = BytesIO()
    rich.save(duplicate_io)
    with pytest.raises(FinalistRetestError, match="CONTROL_SCHEMA_MISMATCH"):
        validate_combined_control_workbook(duplicate_io.getvalue())


def test_combined_control_import_is_atomic_and_idempotent() -> None:
    with duckdb.connect(":memory:") as connection:
        initialize_performance_v2(connection)
        now = datetime(2026, 1, 1, tzinfo=UTC)
        connection.execute(
            "insert into strategies values (1, 'one', 'BTCUSDT', 'LONG', '1h', 5, 1, 'run', 'one', 'ACTIVE', 11, ?, ?)",
            [now, now],
        )
        connection.execute(
            "insert into strategy_results (result_id, strategy_id, report_start_utc, report_end_utc, exchange, commission_rate, initial_balance, final_balance, imported_at_utc) values (11, 1, '2026-01-01', '2026-02-01', 'Bybit', .0004, 100, 101, ?)",
            [now],
        )
        connection.execute("insert into selection_runs values ('sel', (select value from schema_info where key='database_instance_id'), 'BTCUSDT', 'LONG', 'v1', '{}', ?, '{}', ?, 1, 1, 1, 20, ?, ?)", ["a" * 64, "b" * 64, "c" * 64, now])
        connection.execute("insert into selection_results (selection_run_id, strategy_id, result_id_at_selection, auto_status, auto_rank, prior_rejected, stage_trace_json) values ('sel', 1, 11, 'FINALIST', 1, false, '{}')")
        instance = connection.execute("select value from schema_info where key='database_instance_id'").fetchone()[0]
        data = combined_control_workbook_bytes(
            [{"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 1, "result_id": 11, "user_status": "REJECTED", "user_rank": None, "auto_status": "FINALIST", "auto_rank": 1}],
            metadata={"database_instance_id": instance},
            candidate_workbook=_candidate_workbook([
                {"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 1, "result_id": 11, "user_status": "REJECTED", "user_rank": None, "auto_status": "FINALIST", "auto_rank": 1},
            ]),
        )
        result = import_combined_control_workbook(connection, data)
        assert result["group_count"] == 1
        assert connection.execute("select tag from strategy_tags where strategy_id=1").fetchall() == [("REJECTED",)]
        replay = import_combined_control_workbook(connection, data)
        assert replay["review_import_ids"] == result["review_import_ids"]


@pytest.mark.parametrize("reason", [
    "LOT_VARIANT_REDUNDANT", "ANALOG", "PRIOR_USER_REJECTED",
    "RANK_NOT_EVALUATED_INSUFFICIENT_DATA", "AB_NOT_EVALUATED_INSUFFICIENT_DATA",
    "PARETO_PLATEAU_POINTS_PER_ORDER",
])
def test_server_issued_control_accepts_user_edits_and_failure_only_group(reason: str) -> None:
    with duckdb.connect(":memory:") as connection:
        initialize_performance_v2(connection)
        now = datetime(2026, 1, 1, tzinfo=UTC)
        connection.execute(
            "insert into strategies values (1, 'one', 'BTCUSDT', 'LONG', '1h', 5, 1, 'run', 'one', 'ACTIVE', 11, ?, ?)",
            [now, now],
        )
        connection.execute(
            "insert into strategy_results (result_id, strategy_id, report_start_utc, report_end_utc, exchange, commission_rate, initial_balance, final_balance, imported_at_utc) values (11, 1, '2026-01-01', '2026-02-01', 'Bybit', .0004, 100, 101, ?)",
            [now],
        )
        instance = connection.execute("select value from schema_info where key='database_instance_id'").fetchone()[0]
        exact = {
            "Pair": "BTCUSDT", "Direction": "LONG", "Strategy ID": 1, "Result ID": 11,
            "Auto Status": "FINALIST", "Auto Rank": 1, "Auto Analog Of ID": None,
            "Auto Reason": reason, "Score": 5.5,
        }
        exact_rowsets = {"BTCUSDT|LONG": [exact]}
        groups = [
            {"symbol": "BTCUSDT", "side": "LONG", "frozen_count": 1, "success_count": 1, "failure_count": 0, "auto_status_count": 1},
            {"symbol": "ETHUSDT", "side": "SHORT", "frozen_count": 1, "success_count": 0, "failure_count": 1, "auto_status_count": 0},
        ]
        failures = [{"symbol": "ETHUSDT", "side": "SHORT", "strategy_id": 2, "strategy_name": "two", "result_id": 12, "reason": "REPORT_MISSING"}]
        group_digest_rows = [
            {"Pair": row["symbol"], "Direction": row["side"], "Frozen Count": row["frozen_count"], "Success Count": row["success_count"], "Failure Count": row["failure_count"], "Auto Status Count": row["auto_status_count"]}
            for row in groups
        ]
        failure_digest_rows = [{"Pair": "ETHUSDT", "Direction": "SHORT", "Strategy ID": 2, "Strategy": "two", "Result ID": 12, "Reason": "REPORT_MISSING"}]
        metadata = {
            "database_instance_id": instance, "ranking_scope": "RETEST_COHORT",
            "bulk_retest_job_id": "bulk-1", "cohort_sha256": "c" * 64,
            "manifest_sha256": "m" * 64, "config_sha256": "t" * 64,
            "selection_config_sha256": "s" * 64,
            "group_run_ids_json": {"BTCUSDT|LONG": "sel"},
            "exact_rowsets_json": exact_rowsets,
            "exact_rowsets_sha256": canonical_digest(exact_rowsets),
            "immutable_content_sha256": canonical_digest([exact]),
            "groups_sha256": canonical_digest(group_digest_rows),
            "failures_sha256": canonical_digest(failure_digest_rows),
        }
        issued = combined_control_workbook_bytes(
            [{"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 1, "result_id": 11,
              "user_status": "FINALIST", "user_rank": None, "auto_status": "FINALIST", "auto_rank": 1,
              "auto_reason": reason, "score": Decimal("5.5")}],
            groups, failures, metadata,
            candidate_workbook=_candidate_workbook([
                {"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 1, "result_id": 11,
                 "user_status": "FINALIST", "user_rank": None, "auto_status": "FINALIST", "auto_rank": 1,
                 "auto_reason": reason, "score": Decimal("5.5")},
            ], request=parse_selection_request({
                "symbol": "BTCUSDT", "side": "LONG", "stages": [{
                    "id": "pareto_plateau_points_per_order", "enabled": True, "scope": "pair_side",
                }],
            }) if reason == "PARETO_PLATEAU_POINTS_PER_ORDER" else None),
        )
        request_json = canonical_json({
            "ranking_scope": "RETEST_COHORT", "bulk_retest_job_id": "bulk-1",
            "cohort_sha256": "c" * 64, "manifest_sha256": "m" * 64,
            "config_sha256": "t" * 64, "selection_config_sha256": "s" * 64,
            "cohort_members": [[1, 11]],
        })
        connection.execute(
            """insert into selection_runs
               (selection_run_id, database_instance_id, symbol, side, selection_contract_version,
                request_json, request_sha256, config_json, config_sha256, candidate_count,
                representative_count, auto_finalist_count, top_n, workbook_sha256, created_at_utc)
               values ('sel', ?, 'BTCUSDT', 'LONG', 'v1', ?, ?, '{}', ?, 1, 1, 1, 20, ?, ?)""",
            [instance, request_json, "r" * 64, "s" * 64, sha256(issued).hexdigest(), now],
        )
        connection.execute(
            """insert into selection_results
               (selection_run_id, strategy_id, result_id_at_selection, auto_status, auto_score, auto_rank,
                auto_reason, analog_group_key, auto_analog_of_strategy_id, prior_rejected, stage_trace_json)
                values ('sel', 1, 11, 'FINALIST', 5.5, 1, ?, null, null, false, '{}')""", [reason]
        )

        workbook = load_workbook(BytesIO(issued))
        headers = {cell.value: cell.column for cell in workbook["Candidates"][1]}
        rendered_reason = workbook["Candidates"].cell(2, headers["Причина"]).value
        expected_reason = (
            "1. PARETO_PL_PTS_PER_ORDER" if reason == "PARETO_PLATEAU_POINTS_PER_ORDER"
            else reason
        )
        assert rendered_reason == expected_reason
        unedited = import_combined_control_workbook(connection, issued)
        assert unedited["group_count"] == unedited["row_count"] == 1
        workbook["Candidates"].cell(2, headers["User Status"]).value = "RESERVE"
        workbook["Candidates"].cell(2, headers["User Rank"]).value = 1
        edited_io = BytesIO()
        workbook.save(edited_io)
        edited = edited_io.getvalue()
        imported = import_combined_control_workbook(connection, edited)
        replay = import_combined_control_workbook(connection, edited)

        assert imported["group_count"] == 1 and imported["row_count"] == 1
        assert replay["review_import_ids"] == imported["review_import_ids"]
        assert connection.execute(
            "select user_status, user_rank from selection_review_rows order by rowid desc limit 1"
        ).fetchone() == ("RESERVE", 1)

        tampered = load_workbook(BytesIO(issued))
        tampered_headers = {cell.value: cell.column for cell in tampered["Candidates"][1]}
        tampered["Candidates"].cell(2, tampered_headers["Auto Status"]).value = "RESERVE"
        tampered_io = BytesIO()
        tampered.save(tampered_io)
        with pytest.raises(FinalistRetestError, match="CONTROL_IMMUTABLE_FIELDS_CHANGED"):
            import_combined_control_workbook(connection, tampered_io.getvalue())

        bad_analog = load_workbook(BytesIO(issued))
        bad_analog_headers = {cell.value: cell.column for cell in bad_analog["Candidates"][1]}
        bad_analog["Candidates"].cell(2, bad_analog_headers["User Status"]).value = "ANALOG"
        bad_analog_io = BytesIO()
        bad_analog.save(bad_analog_io)
        with pytest.raises(FinalistRetestError, match="CONTROL_INVALID_ANALOG"):
            import_combined_control_workbook(connection, bad_analog_io.getvalue())
