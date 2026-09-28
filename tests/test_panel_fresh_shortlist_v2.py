from __future__ import annotations

import json
from http.client import HTTPConnection
from pathlib import Path
import threading

import pytest

from mrs3.config import AlgorithmConfig
from mrs3.panel import PanelController, create_panel_server
from mrs3.fresh_shortlist import ShortlistBusyError
from tests.test_fresh_analysis_strategies import _make_analysis


def _controller(tmp_path: Path, *, ready: bool = True) -> tuple[PanelController, Path, str]:
    analysis_path = tmp_path / "run.analysis-v6.duckdb"
    analysis_id, _ = _make_analysis(analysis_path, ready=ready)
    config = tmp_path / "config.local.json"
    config.write_text(json.dumps({"duckdb_import": {"workers": 1}}), encoding="utf-8")
    controller = PanelController(
        tmp_path, config, analysis_config_loader=lambda _: AlgorithmConfig.defaults()
    )
    controller._fresh_analysis_paths[analysis_id] = analysis_path
    return controller, analysis_path, analysis_id


def _selection_request(analysis_id: str, **extra: object) -> dict[str, object]:
    return {
        "analysis_run_id": analysis_id,
        "filter_version": "shortlist-v2",
        "pretest_ab_enabled": False,
        "ladder_enabled": False,
        "pareto_enabled": False,
        **extra,
    }


def test_fresh_shortlist_returns_applied_snapshot_identity(tmp_path: Path) -> None:
    controller, _analysis_path, analysis_id = _controller(tmp_path)

    result = controller.strategies_fresh_shortlist({
        "analysis_run_id": analysis_id,
        "filter_version": "shortlist-v2",
        "pretest_ab_enabled": False,
        "ladder_enabled": False,
        "pareto_enabled": False,
    })

    assert result["filter_version"] == "shortlist-v2"
    assert result["analysis_run_id"] == analysis_id
    assert result["applied_options"] == {
        "pretest_ab_enabled": False,
        "ladder_enabled": False,
        "pareto_enabled": False,
    }
    assert len(result["selection_token"]) == 64
    assert result["groups"][0]["candidate_ids"] == ["STR-READY"]


@pytest.mark.parametrize("token", [None, "0" * 64], ids=["missing", "stale"])
@pytest.mark.parametrize(
    ("consumer", "extra"),
    [
        ("list_audit", {"audit": True}),
        ("list_audit", {"audit": True, "filters": {"source_pnl": False}, "close_support": False}),
        ("dedicated_audit", {}),
        ("generate", {"candidate_ids": ["STR-READY"], "selected_scopes": [["BTCUSDT", "LONG", "1h"]]}),
        ("runs", {"selected_scopes": [["BTCUSDT", "LONG", "1h"]], "start_date": "2026-08-01", "end_date": "2026-08-18"}),
    ],
)
def test_fresh_output_consumers_require_applied_selection_token(
    tmp_path: Path, consumer: str, extra: dict[str, object], token: str | None,
) -> None:
    controller, _analysis_path, analysis_id = _controller(tmp_path)
    payload = _selection_request(analysis_id, **extra)
    if token is not None:
        payload["selection_token"] = token
    invoke = {
        "list_audit": controller.strategies_fresh_shortlist,
        "dedicated_audit": controller.strategies_fresh_filter_audit,
        "generate": controller.strategies_fresh_generate,
        "runs": controller.strategies_fresh_generate_runs,
    }[consumer]

    with pytest.raises(ValueError, match="STALE_SHORTLIST_SELECTION"):
        invoke(payload)

    assert controller._fresh_generation_job is None
    assert not (tmp_path / "Output").exists()
    assert not (tmp_path / "bot").exists()


def test_json_generation_uses_server_ready_ids_not_browser_candidate_ids(
    tmp_path: Path, monkeypatch,
) -> None:
    controller, _analysis_path, analysis_id = _controller(tmp_path)
    snapshot = controller.strategies_fresh_shortlist(_selection_request(analysis_id))
    captured: dict[str, object] = {}
    template = tmp_path / "template.json"
    template.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(controller, "_workflow_default", lambda *_args, **_kwargs: template)

    def generate(*args: object, **kwargs: object) -> object:
        captured["candidate_ids"] = args[2]
        captured["selection"] = kwargs["selection"]
        return type("Generated", (), {"run_id": analysis_id, "surface_id": "surface", "strategy_count": 1})()

    monkeypatch.setattr("mrs3.panel.generate_fresh_analysis_strategies", generate)
    result = controller._generate_fresh_strategies(_selection_request(
        analysis_id,
        candidate_ids=["BROWSER-FORGED-ID"],
        selected_scopes=[["BTCUSDT", "LONG", "1h"]],
        selection_token=snapshot["selection_token"],
    ))

    assert captured["candidate_ids"] == ("STR-READY",)
    assert captured["selection"].selection_token == snapshot["selection_token"]
    assert result["phase"] == "COMMITTED"
    assert (tmp_path / "Output" / "fresh-shortlist-v2" / analysis_id).is_dir()


def test_json_generation_freezes_all_applied_flags_in_thread_payload(
    tmp_path: Path, monkeypatch,
) -> None:
    controller, _analysis_path, analysis_id = _controller(tmp_path)
    request = _selection_request(analysis_id, ladder_enabled=True, pareto_enabled=True)
    snapshot = controller.strategies_fresh_shortlist(request)
    captured: dict[str, object] = {}

    class CapturedThread:
        def __init__(self, *, target, args, **_kwargs):
            captured["payload"] = args[1]

        def start(self) -> None:
            pass

    monkeypatch.setattr("mrs3.panel.threading.Thread", CapturedThread)
    controller.strategies_fresh_generate({
        **request,
        "selection_token": snapshot["selection_token"],
        "selected_scopes": [["BTCUSDT", "LONG", "1h"]],
    })

    assert captured["payload"]["pretest_ab_enabled"] is False
    assert captured["payload"]["ladder_enabled"] is True
    assert captured["payload"]["pareto_enabled"] is True


def test_final_fresh_manifest_is_recovered_after_panel_cache_is_cleared(tmp_path: Path, monkeypatch) -> None:
    from tests.test_fresh_analysis_strategies import _template

    controller, _analysis_path, analysis_id = _controller(tmp_path)
    template = tmp_path / "template.json"
    template.write_text(json.dumps(_template()), encoding="utf-8")
    monkeypatch.setattr(controller, "_workflow_default", lambda *_args, **_kwargs: template)
    snapshot = controller.strategies_fresh_shortlist(_selection_request(analysis_id))

    result = controller._generate_fresh_strategies(_selection_request(
        analysis_id,
        candidate_ids=["BROWSER-FORGED-ID"],
        selected_scopes=[["BTCUSDT", "LONG", "1h"]],
        selection_token=snapshot["selection_token"],
    ))
    manifest = Path(result["output_dir"]) / "strategy_manifest.json"
    assert manifest.is_file()
    assert (manifest.parent / "strategies").is_dir()

    controller._fresh_strategy_manifests.clear()
    assert controller._fresh_strategy_manifest(analysis_id) == manifest
    assert controller.strategies_fresh_batch()["strategy_count"] == 2


def test_fresh_json_generation_rejects_browser_output_directory(tmp_path: Path) -> None:
    controller, _analysis_path, analysis_id = _controller(tmp_path)
    snapshot = controller.strategies_fresh_shortlist(_selection_request(analysis_id))

    with pytest.raises(ValueError, match="output_dir is server-controlled"):
        controller.strategies_fresh_generate(_selection_request(
            analysis_id,
            output_dir=str(tmp_path / "browser-output"),
            candidate_ids=["STR-READY"],
            selected_scopes=[["BTCUSDT", "LONG", "1h"]],
            selection_token=snapshot["selection_token"],
        ))

    assert controller._fresh_generation_job is None
    assert not (tmp_path / "browser-output").exists()


def test_fresh_json_generation_rejects_client_private_keys(tmp_path: Path) -> None:
    controller, _analysis_path, analysis_id = _controller(tmp_path)
    snapshot = controller.strategies_fresh_shortlist(_selection_request(analysis_id))

    with pytest.raises(ValueError, match="private fields are not allowed"):
        controller.strategies_fresh_generate(_selection_request(
            analysis_id,
            _shortlist_prepared=object(),
            selected_scopes=[["BTCUSDT", "LONG", "1h"]],
            selection_token=snapshot["selection_token"],
        ))

    assert controller._fresh_generation_job is None
    assert not (tmp_path / "Output").exists()


def test_empty_ready_selection_fails_before_json_publication(tmp_path: Path) -> None:
    controller, _analysis_path, analysis_id = _controller(tmp_path, ready=False)
    snapshot = controller.strategies_fresh_shortlist(_selection_request(analysis_id))

    with pytest.raises(ValueError, match="EMPTY_READY_SELECTION"):
        controller.strategies_fresh_generate(_selection_request(
            analysis_id,
            candidate_ids=["BROWSER-FORGED-ID"],
            selected_scopes=[["BTCUSDT", "LONG", "1h"]],
            selection_token=snapshot["selection_token"],
        ))

    assert controller._fresh_generation_job is None
    assert not (tmp_path / "Output").exists()


def test_changed_analysis_digest_rejects_old_token_before_any_output(tmp_path: Path) -> None:
    import duckdb

    controller, analysis_path, analysis_id = _controller(tmp_path)
    previous = controller.strategies_fresh_shortlist(_selection_request(analysis_id))
    connection = duckdb.connect(str(analysis_path))
    try:
        raw = connection.execute("select payload_json from structures").fetchone()[0]
        structure = json.loads(raw)
        structure["orders"][0]["source_pnl_pct"] += 1
        connection.execute(
            "update structures set payload_json=?",
            [json.dumps(structure, sort_keys=True, separators=(",", ":"))],
        )
    finally:
        connection.close()

    updated = controller.strategies_fresh_shortlist(_selection_request(analysis_id))
    assert updated["artifact_sha256"] != previous["artifact_sha256"]
    assert updated["selection_token"] != previous["selection_token"]

    stale = _selection_request(
        analysis_id,
        audit=True,
        selection_token=previous["selection_token"],
    )
    with pytest.raises(ValueError, match="STALE_SHORTLIST_SELECTION"):
        controller.strategies_fresh_shortlist(stale)
    with pytest.raises(ValueError, match="STALE_SHORTLIST_SELECTION"):
        controller.strategies_fresh_generate(_selection_request(
            analysis_id,
            candidate_ids=["STR-READY"],
            selected_scopes=[["BTCUSDT", "LONG", "1h"]],
            selection_token=previous["selection_token"],
        ))

    assert controller._fresh_generation_job is None
    assert not (tmp_path / "Output").exists()


def test_empty_evaluation_still_has_token_and_header_only_ready_audit(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    controller, _analysis_path, analysis_id = _controller(tmp_path, ready=False)
    snapshot = controller.strategies_fresh_shortlist(_selection_request(analysis_id))
    assert len(snapshot["selection_token"]) == 64
    assert snapshot["groups"][0]["ready_after_filters"] == 0
    result = controller.strategies_fresh_filter_audit(_selection_request(
        analysis_id, selection_token=snapshot["selection_token"],
    ))

    workbook = load_workbook(tmp_path / "Output" / result["filename"], read_only=True, data_only=True)
    try:
        assert workbook["READY"].max_row == 1
        assert workbook["READY"].max_column >= 5
    finally:
        workbook.close()


def test_scope_without_any_candidate_remains_in_empty_shortlist_snapshot(tmp_path: Path) -> None:
    import duckdb

    controller, analysis_path, analysis_id = _controller(tmp_path)
    connection = duckdb.connect(str(analysis_path))
    try:
        connection.execute("delete from structures")
    finally:
        connection.close()

    snapshot = controller.strategies_fresh_shortlist(_selection_request(analysis_id))

    assert len(snapshot["selection_token"]) == 64
    assert snapshot["groups"] == [{
        "scope_key": "BTCUSDT|LONG|1h", "pair": "BTCUSDT", "side": "LONG", "timeframe": "1h",
        "counts": {"1ORD": 0, "2ORD": 0, "3ORD": 0, "4ORD": 0},
        "ready": 0, "ready_after_filters": 0, "deferred": 0, "total": 0,
        "candidate_ids": [], "plateau_count": 0, "period": None,
    }]


def test_audit_and_json_staging_are_discarded_when_digest_changes_before_publish(
    tmp_path: Path, monkeypatch,
) -> None:
    controller, _analysis_path, analysis_id = _controller(tmp_path)
    snapshot = controller.strategies_fresh_shortlist(_selection_request(analysis_id))
    monkeypatch.setattr(
        controller, "_assert_fresh_analysis_digest",
        lambda *_args: (_ for _ in ()).throw(ValueError("STALE_SHORTLIST_SELECTION")),
    )

    with pytest.raises(ValueError, match="STALE_SHORTLIST_SELECTION"):
        controller.strategies_fresh_shortlist(_selection_request(
            analysis_id, audit=True, selection_token=snapshot["selection_token"],
        ))

    assert not list((tmp_path / "Output").glob("*.xlsx"))
    assert not list((tmp_path / "Output").glob("*.pending.xlsx"))


@pytest.mark.parametrize("dedicated", [False, True], ids=["shortlist-audit", "dedicated-audit"])
def test_audit_publishes_tokenized_workbook_with_missing_output_parent(
    tmp_path: Path, monkeypatch, dedicated: bool,
) -> None:
    from openpyxl import load_workbook
    from mrs3.analysis_filter_export import export_fresh_shortlist_audit as export

    controller, _analysis_path, analysis_id = _controller(tmp_path)
    snapshot = controller.strategies_fresh_shortlist(_selection_request(analysis_id))
    assert not (tmp_path / "Output").exists()
    observed_parent = []

    def export_with_parent_check(prepared, evaluation, path):
        observed_parent.append(Path(path).parent.is_dir())
        return export(prepared, evaluation, path)

    monkeypatch.setattr("mrs3.panel.export_fresh_shortlist_audit", export_with_parent_check)
    audit_request = _selection_request(analysis_id, selection_token=snapshot["selection_token"])
    result = (
        controller.strategies_fresh_filter_audit(audit_request)
        if dedicated else controller.strategies_fresh_shortlist({**audit_request, "audit": True})
    )
    workbook_path = tmp_path / "Output" / result["filename"]

    assert observed_parent == [True]
    assert f"{snapshot['selection_token']}" in result["filename"]
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        summary = dict(workbook["Summary"].iter_rows(min_row=2, values_only=True))
        assert summary["selection_token"] == snapshot["selection_token"]
        assert workbook["READY"].max_row == 2
    finally:
        workbook.close()


def test_json_generation_staging_is_discarded_if_digest_changes_before_publish(
    tmp_path: Path, monkeypatch,
) -> None:
    controller, _analysis_path, analysis_id = _controller(tmp_path)
    snapshot = controller.strategies_fresh_shortlist(_selection_request(analysis_id))
    template = tmp_path / "template.json"
    template.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(controller, "_workflow_default", lambda *_args, **_kwargs: template)
    monkeypatch.setattr(
        "mrs3.panel.generate_fresh_analysis_strategies",
        lambda *_args, **_kwargs: type("Generated", (), {
            "run_id": analysis_id, "surface_id": "surface", "strategy_count": 1,
        })(),
    )
    monkeypatch.setattr(
        controller, "_assert_fresh_analysis_digest",
        lambda *_args: (_ for _ in ()).throw(ValueError("STALE_SHORTLIST_SELECTION")),
    )

    with pytest.raises(ValueError, match="STALE_SHORTLIST_SELECTION"):
        controller._generate_fresh_strategies(_selection_request(
            analysis_id,
            candidate_ids=["BROWSER-FORGED-ID"],
            selected_scopes=[["BTCUSDT", "LONG", "1h"]],
            selection_token=snapshot["selection_token"],
        ))

    staging_root = tmp_path / "Output" / "fresh-shortlist-v2" / analysis_id
    assert not list(staging_root.glob("*/strategy_manifest.json"))
    assert not list(staging_root.glob(".stage-*"))


@pytest.mark.parametrize(
    ("route", "method", "extra"),
    [
        ("shortlist", "strategies_fresh_shortlist", {}),
        ("shortlist", "strategies_fresh_shortlist", {"audit": True}),
        ("generate", "strategies_fresh_generate", {}),
        ("runs", "strategies_fresh_generate_runs", {}),
    ],
)
def test_shortlist_busy_is_a_retryable_http_conflict_for_all_consumers(
    tmp_path: Path, monkeypatch, route: str, method: str, extra: dict[str, object],
) -> None:
    controller = PanelController(
        tmp_path, tmp_path / "config.local.json",
        analysis_config_loader=lambda _: AlgorithmConfig.defaults(),
    )
    monkeypatch.setattr(
        controller, method, lambda _payload: (_ for _ in ()).throw(ShortlistBusyError()),
    )
    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
    try:
        connection.request(
            "POST", f"/api/v2/strategies/fresh/{route}",
            json.dumps({"analysis_run_id": "a" * 64, "filter_version": "shortlist-v2", **extra}).encode("utf-8"),
            {"Content-Type": "application/json"},
        )
        response = connection.getresponse()
        body = json.loads(response.read().decode("utf-8"))
        retry_after = response.getheader("Retry-After")
    finally:
        connection.close()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)

    assert response.status == 409
    assert retry_after == "1"
    assert body == {"error": {"code": "SHORTLIST_BUSY", "message": "SHORTLIST_BUSY"}}


def test_generation_scope_normalizes_side_case(tmp_path: Path) -> None:
    controller, analysis_path, analysis_id = _controller(tmp_path)
    prepared, evaluation = controller._evaluate_fresh_shortlist_snapshot(
        analysis_path, analysis_id, (False, False, False),
    )

    scopes, ready_ids = controller._selected_fresh_ready(
        prepared, evaluation, [[" BTCUSDT ", "long", "1h"]], require_one_side=True,
    )

    assert scopes == (("BTCUSDT", "LONG", "1h"),)
    assert ready_ids == ("STR-READY",)
