from __future__ import annotations

import asyncio
import copy
import hashlib
import json
from http.client import HTTPConnection
import os
from pathlib import Path
import subprocess
import threading
import time
from decimal import Decimal
from types import SimpleNamespace
from typing import Any
import zipfile
from openpyxl import load_workbook
from openpyxl import Workbook

import pytest
import duckdb

import mrs3.panel as panel_module
from mrs3.panel import PanelController, create_panel_server
from mrs3.panel_portfolio import (
    PORTFOLIO_SNAPSHOT_UNSERIALIZABLE,
    PortfolioPanelError,
    PortfolioPanelService,
    STAGES,
    _redact_text,
    _repair_legacy_geometry,
    _safe_cell,
    _snapshot_bytes,
    _PortfolioProgressReporter,
    _stage2_material,
    _weighted_entry_order_percentages,
    _weighted_payload_pairs,
    _weighted_source_maxdd_sum,
)
from mrs3.panel_jobs import PanelJobError, PanelJobRegistry
from mrs3.config import DuckDBImportSettings
from mrs3.performance_v2_store import PerformanceV2StoreError
from mrs3.portfolio.config import WEIGHTED_SEARCH_DEFAULTS, PortfolioConfigError, migrate_portfolio_config_document
from mrs3.portfolio.adapter import (
    CAMPAIGN_CONTRACT_VERSION,
    CAMPAIGN_SEARCH_MODE,
    CAMPAIGN_WEIGHTED_ALGO_VERSION,
    CampaignContractError,
    MINUTE_CAPACITY_UNAVAILABLE,
)


def _config() -> dict:
    money = lambda amount: {"amount": str(amount), "currency": "USDT"}
    profile = lambda: {
        "pnl": {"policy_id": "operator_supplied_pnl_v1", "parameters": {"operator_supplied": True}},
        "liquidity": {"policy_id": "operator_supplied_liquidity_v1", "parameters": {"operator_supplied": True}},
        "ranking": {"id": "operator_supplied_ranking_v1", "parameters": {"operator_supplied": True}, "top_n": 1},
    }
    return migrate_portfolio_config_document({
        "schema_version": 1,
        "policy_version": "portfolio_optimizer_research_risk_v1",
        "algorithm_versions": {"sizing": "portfolio_optimizer_sizing_v1", "ranking": "portfolio_optimizer_ranking_v1"},
        "inputs": {"performance_db": "performance.duckdb", "portfolio_db": "portfolio.duckdb", "collector_root": "collector", "approved_templates": ["template.json"]},
        "scenarios": {name: {"account": name, "deposit": money(10000), "collateral": money(10000), "max_balance": money(10000), "sizing": {"upper_bound": money(10000), "grid": [money(1000)]}} for name in ("AGGRESSIVE", "BALANCED", "CONSERVATIVE")},
        "search": {"universe": {"policy_id": "u", "parameters": {"x": 1}}, "composition": {"policy_id": "c", "parameters": {"x": 1}}, "sizing": {"policy_id": "s", "parameters": {"x": 1}}, "limiter": {"policy_id": "l", "parameters": {"x": 1}}, "priority": {"policy_id": "p", "parameters": {"x": 1}}, "seed": 1, "rounds": 1, "total_test_budget": 1},
        "research": {"development_window": {"value": 1, "unit": "days"}, "validation_window": {"value": 1, "unit": "days"}, "warmup": {"value": 0, "unit": "days"}, "boundary": "fixture", "evidence_minimum": {"value": 1, "unit": "count"}},
        "liquidity": {"policy_id": "liquidity", "parameters": {"x": 1}}, "margin": {"policy_id": "margin", "parameters": {"x": 1}},
        "profiles": {name: profile() for name in ("AGGRESSIVE", "BALANCED", "CONSERVATIVE")},
        "runner": {"target": "local", "root": "tester", "timeout": {"value": 1, "unit": "seconds"}, "retries": 0},
    })[0]


def _write_config(path: Path) -> str:
    raw = json.dumps(_config(), ensure_ascii=False, indent=2).encode()
    path.write_bytes(raw)
    return hashlib.sha256(raw).hexdigest()


def _profiled_variants(selected, *_):
    return tuple({
        **dict(item),
        "candidate_id": f"candidate-{index + 1}",
        "identity": f"candidate-{index + 1}",
        "profile": "BALANCED",
        "search_mode": CAMPAIGN_SEARCH_MODE,
        "limiter_L": 0,
        "metrics": {"limiter_L": 0, "required_bank_usdt": Decimal("10000")},
        "pretest_period": {"start_utc": "2026-01-01T00:00:00Z", "end_utc": "2026-01-15T00:00:00Z"},
        "members": (
            {"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 700, "result_id": 711, "x_usdt": Decimal("100")},
            {"symbol": "BTCUSDT", "side": "SHORT", "strategy_id": 701, "result_id": 712, "x_usdt": Decimal("100")},
        ),
        "strategy_payloads": (
            _executable_payload("BTCUSDT", "LONG", name="PORTFOLIO_BTCUSDT_700_711"),
            _executable_payload("BTCUSDT", "SHORT", name="PORTFOLIO_BTCUSDT_701_712"),
        ),
    } for index, item in enumerate(selected))


def _executable_payload(symbol: str, side: str = "LONG", *, name: str | None = None, x: str = "100") -> dict:
    return {
        "side": side,
        "strategy": {
            "name": name or f"PORTFOLIO_{symbol}_{side}",
            "basic": {"symbol": symbol},
            "mrs": {"position_priority": 1},
        },
        "account": {"open_positions_limiter": 0},
        "facts": {"B": "10000", "x": x},
    }


def _candidate_member(symbol: str, side: str, strategy_id: int, result_id: int, x: str = "100") -> dict:
    return {"symbol": symbol, "side": side, "strategy_id": strategy_id, "result_id": result_id, "x_usdt": Decimal(x)}


def _weighted_executable_candidate(payloads: tuple[dict, ...], **updates) -> dict:
    members = tuple(
        _candidate_member(payload["strategy"]["basic"]["symbol"], payload["side"], index + 1, index + 101, payload["facts"]["x"])
        for index, payload in enumerate(payloads)
    )
    return {
        "candidate_id": "candidate-executable",
        "profile": "BALANCED",
        "search_mode": CAMPAIGN_SEARCH_MODE,
        "limiter_L": 0,
        "metrics": {"required_bank_usdt": Decimal("10000")},
        "pretest_period": {"start_utc": "2026-01-01T00:00:00Z", "end_utc": "2026-01-15T00:00:00Z"},
        "members": members,
        "strategy_payloads": payloads,
        **updates,
    }


def _stage2_payload(symbol: str, side: str, *, name: str, bank: str = "10000") -> dict:
    payload = _executable_payload(symbol, side, name=name)
    payload["facts"].update({"B": bank, "C": "100", "q": "1"})
    payload["strategy"]["basic"]["max_balance"] = "100"
    return payload


def _stage2_service(tmp_path: Path, candidate: dict) -> tuple[PortfolioPanelService, dict]:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    service = PortfolioPanelService(
        tmp_path,
        path,
        finalists_reader=lambda *_: [_finalist()],
        variant_generator=lambda *_: (candidate,),
    )
    result = service.submit_campaign({
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
        "expected_config_digest": digest,
    })
    assert _wait_stage1(service, result)["status"] == "SUCCEEDED"
    return service, result


def _wait_stage1(service: PortfolioPanelService, result: dict) -> dict:
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and service.job(result["job_id"])["status"] in {"QUEUED", "RUNNING"}:
        time.sleep(0.01)
    return service.job(result["job_id"])


def _finalist(**updates) -> dict:
    return {
        "strategy_id": 7,
        "result_id": 11,
        "symbol": "BTCUSDT",
        "side": "LONG",
        "user_status": "FINALIST",
        "user_rank": 1,
        "actions": [],
        "equity": [],
        **updates,
    }


def test_panel_freeze_emits_weighted_contract_without_legacy_stage1(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    finalists = [_finalist()]

    class IdleThread:
        def __init__(self, *args, **kwargs):
            pass

        def start(self):
            pass

    monkeypatch.setattr("mrs3.panel_portfolio.threading.Thread", IdleThread)
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: finalists)
    result = service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})
    campaign = service.registry.runtime(result["job_id"])["campaign"]
    assert campaign["campaign_contract_version"] == CAMPAIGN_CONTRACT_VERSION
    assert campaign["search_mode"] == CAMPAIGN_SEARCH_MODE
    assert campaign["weighted_algo_version"] == CAMPAIGN_WEIGHTED_ALGO_VERSION
    assert "stage1_mode" not in campaign
    assert all(campaign["versions"][key] == campaign[key] for key in ("campaign_contract_version", "search_mode", "weighted_algo_version"))


def test_panel_freeze_carries_dedicated_weighted_template_and_digest(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    template_path = Path(__file__).parents[1] / "templates" / "strategies" / "portfolio-weighted-mrs" / "base.json"
    expected_template = json.loads(template_path.read_text(encoding="utf-8"))
    expected_digest = hashlib.sha256(template_path.read_bytes()).hexdigest()

    class IdleThread:
        def __init__(self, *args, **kwargs):
            pass

        def start(self):
            pass

    monkeypatch.setattr("mrs3.panel_portfolio.threading.Thread", IdleThread)
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: [_finalist()])
    result = service.submit_campaign({
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
        "expected_config_digest": digest,
    })

    campaign = service.registry.runtime(result["job_id"])["campaign"]
    assert campaign["strategy_template"] == expected_template
    assert campaign["strategy_template_digest"] == expected_digest
    assert campaign["strategy_template"]["mrs"]["position_priority"] == 3


def test_panel_freeze_keeps_template_snapshot_after_file_changes(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    source_path = Path(__file__).parents[1] / "templates" / "strategies" / "portfolio-weighted-mrs" / "base.json"
    template_path = tmp_path / "base.json"
    template_path.write_bytes(source_path.read_bytes())
    monkeypatch.setattr("mrs3.panel_portfolio._WEIGHTED_TEMPLATE_PATH", template_path)

    class IdleThread:
        def __init__(self, *args, **kwargs):
            pass

        def start(self):
            pass

    monkeypatch.setattr("mrs3.panel_portfolio.threading.Thread", IdleThread)
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: [_finalist()])
    result = service.submit_campaign({
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
        "expected_config_digest": digest,
    })
    campaign = service.registry.runtime(result["job_id"])["campaign"]
    frozen_template = json.loads(json.dumps(campaign["strategy_template"]))
    frozen_digest = campaign["strategy_template_digest"]

    template_path.write_text('{"mrs":{"position_priority":1}}', encoding="utf-8")

    assert campaign["strategy_template"] == frozen_template
    assert campaign["strategy_template_digest"] == frozen_digest


@pytest.mark.parametrize(
    "template_bytes",
    [
        b"not-json",
        b"[]",
        b"null",
        b"true",
        b"NaN",
        b"{\"value\": Infinity}",
        b"{\"nested\": {\"value\": NaN}}",
        b"{\"nested\": {\"value\": -Infinity}}",
        b"\xff\xfe{}",
    ],
)
def test_panel_freeze_rejects_malformed_or_non_object_weighted_template_before_reader(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, template_bytes: bytes,
) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    template_path = tmp_path / "base.json"
    template_path.write_bytes(template_bytes)
    monkeypatch.setattr("mrs3.panel_portfolio._WEIGHTED_TEMPLATE_PATH", template_path)
    calls: list[tuple[object, ...]] = []

    def reader(*args):
        calls.append(args)
        return [_finalist()]

    service = PortfolioPanelService(tmp_path, path, finalists_reader=reader)
    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign({
            "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
            "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
            "expected_config_digest": digest,
        })

    assert error.value.code == "PORTFOLIO_TEMPLATE_INVALID"
    assert calls == []
    assert service.registry.list() == []


def test_panel_freeze_rejects_missing_weighted_template_before_reader(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    monkeypatch.setattr("mrs3.panel_portfolio._WEIGHTED_TEMPLATE_PATH", tmp_path / "missing.json")
    calls: list[tuple[object, ...]] = []

    def reader(*args):
        calls.append(args)
        return [_finalist()]

    service = PortfolioPanelService(tmp_path, path, finalists_reader=reader)
    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign({
            "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
            "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
            "expected_config_digest": digest,
        })

    assert error.value.code == "PORTFOLIO_TEMPLATE_INVALID"
    assert calls == []
    assert service.registry.list() == []


def test_campaign_launch_accepts_directional_shared_symbol_limits(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    service = PortfolioPanelService(tmp_path, path)
    config, _raw, _document = service._config()

    launch = service._normalise_campaign(
        {
            "pairs": [{"pair": " btCusdt ", "max_finalist_long": 1, "max_finalist_short": 1}],
            "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
            "expected_config_digest": digest,
        },
        config,
        digest,
    )

    assert launch["selected_pairs"] == [["BTCUSDT", "LONG"], ["BTCUSDT", "SHORT"]]
    assert launch["maximums"] == {"BTCUSDT|LONG": 1, "BTCUSDT|SHORT": 1}


def test_campaign_launch_accepts_one_enabled_direction(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    service = PortfolioPanelService(tmp_path, path)
    config, _raw, _document = service._config()

    launch = service._normalise_campaign(
        {
            "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 0, "max_finalist_short": 1}],
            "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
            "expected_config_digest": digest,
        },
        config,
        digest,
    )

    assert launch["selected_pairs"] == [["BTCUSDT", "SHORT"]]
    assert launch["maximums"] == {"BTCUSDT|SHORT": 1}


def test_campaign_launch_keeps_distinct_pairs_in_stable_long_only_order(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    service = PortfolioPanelService(tmp_path, path)
    config, _raw, _document = service._config()

    launch = service._normalise_campaign(
        {
            "pairs": [
                {"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0},
                {"pair": "ETHUSDT", "max_finalist_long": 1, "max_finalist_short": 0},
            ],
            "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
            "expected_config_digest": digest,
        },
        config,
        digest,
    )

    assert launch["selected_pairs"] == [["BTCUSDT", "LONG"], ["ETHUSDT", "LONG"]]
    assert launch["maximums"] == {"BTCUSDT|LONG": 1, "ETHUSDT|LONG": 1}


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("max_finalist_long", -1),
        ("max_finalist_short", -1),
        ("max_finalist_long", True),
        ("max_finalist_short", False),
        ("max_finalist_long", 1.5),
        ("max_finalist_short", "1"),
    ],
)
def test_campaign_rejects_invalid_finalist_limits_before_reader(tmp_path: Path, field: str, value: int) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    calls: list[tuple[object, ...]] = []

    def reader(*args):
        calls.append(args)
        return ()

    service = PortfolioPanelService(tmp_path, path, finalists_reader=reader)
    pair = {"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 1}
    pair[field] = value

    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign(
            {
                "pairs": [pair],
                "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
                "expected_config_digest": digest,
            }
        )

    assert error.value.code == "PORTFOLIO_CAMPAIGN_INVALID"
    assert error.value.field_errors[0]["field"] == f"pairs[0].{field}"
    assert calls == []


def test_campaign_rejects_missing_finalist_limit_with_field_error_before_reader(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    calls: list[tuple[object, ...]] = []
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *args: calls.append(args) or ())
    pair = {"pair": "BTCUSDT", "max_finalist_short": 1}

    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign(
            {
                "pairs": [pair],
                "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
                "expected_config_digest": digest,
            }
        )

    assert error.value.status == 422
    assert error.value.field_errors[0]["field"] == "pairs[0]"
    assert error.value.field_errors[0]["code"] == "INVALID_PAIR"
    assert calls == []


def test_campaign_rejects_both_directions_disabled_before_reader(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    calls: list[tuple[object, ...]] = []
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *args: calls.append(args) or ())

    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign(
            {
                "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 0, "max_finalist_short": 0}],
                "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
                "expected_config_digest": digest,
            }
        )

    assert error.value.code == "PORTFOLIO_CAMPAIGN_INVALID"
    assert error.value.field_errors[0]["field"] == "pairs[0].max_finalist_long"
    assert calls == []


@pytest.mark.parametrize(
    "pairs",
    [
        [
            {"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0},
            {"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0},
        ],
        [{"pair": "BTCUSDT", "max_finalist_long": 1}],
        [{"pair": "BTCUSDT", "max_finalist_long": None, "max_finalist_short": 0}],
        [{"pair": "BTCUSDT", "max_finalist_long": "1", "max_finalist_short": 0}],
    ],
)
def test_campaign_rejects_duplicate_or_malformed_limits_before_reader(tmp_path: Path, pairs: list[dict[str, object]]) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    calls: list[tuple[object, ...]] = []

    def reader(*args):
        calls.append(args)
        return ()

    service = PortfolioPanelService(tmp_path, path, finalists_reader=reader)

    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign(
            {
                "pairs": pairs,
                "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
                "expected_config_digest": digest,
            }
        )

    assert error.value.code == "PORTFOLIO_CAMPAIGN_INVALID"
    assert calls == []


def test_campaign_rejects_empty_profiles_before_reader(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    calls: list[tuple[object, ...]] = []
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *args: calls.append(args) or ())

    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign({
            "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
            "profiles": [],
            "expected_config_digest": digest,
        })

    assert error.value.code == "PORTFOLIO_CAMPAIGN_INVALID"
    assert calls == []


def test_invalid_resume_contract_maps_before_decode_or_runtime_writes(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    saved = registry.submit("portfolio.stage1", {}, "resume-contract", ("portfolio_optimizer",), job_id="resume-contract")
    registry.reserve_runtime(saved["job_id"], "campaign", {"search_mode": "PRETEST_PROXY"})
    service = PortfolioPanelService(tmp_path, registry=registry)
    monkeypatch.setattr("mrs3.panel_portfolio.base64.b64decode", lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("decode called")))
    calls: list[str] = []
    service.cutoff_selector = lambda *_args, **_kwargs: calls.append("cutoff")
    service.variant_generator = lambda *_args, **_kwargs: calls.append("variants")

    service._run(saved["job_id"])

    persisted = registry.get(saved["job_id"])
    runtime = registry.runtime(saved["job_id"])
    assert persisted["state"] == "FAILED"
    assert persisted["error"]["code"] == "CAMPAIGN_CONTRACT_VERSION_UNSUPPORTED"
    assert runtime["diagnostics"] == [{
        "severity": "ERROR",
        "code": "CAMPAIGN_CONTRACT_VERSION_UNSUPPORTED",
        "message": "CAMPAIGN_CONTRACT_VERSION_UNSUPPORTED",
    }]
    assert runtime["journal"][-1]["code"] == "CAMPAIGN_CONTRACT_VERSION_UNSUPPORTED"
    assert calls == []


def test_invalid_resume_nonmapping_runtime_persists_typed_contract_failure(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    saved = registry.submit("portfolio.stage1", {}, "resume-runtime", ("portfolio_optimizer",), job_id="resume-runtime")
    service = PortfolioPanelService(tmp_path, registry=registry)
    monkeypatch.setattr(registry, "runtime", lambda *_args: [])

    service._run(saved["job_id"])

    persisted = registry.get(saved["job_id"])
    stored_runtime = registry.jobs[saved["job_id"]]["runtime"]
    assert persisted["state"] == "FAILED"
    assert persisted["error"]["code"] == "CAMPAIGN_CONTRACT_VERSION_UNSUPPORTED"
    assert stored_runtime["diagnostics"][0]["code"] == "CAMPAIGN_CONTRACT_VERSION_UNSUPPORTED"


def test_invalid_resume_nonmapping_campaign_persists_typed_contract_failure(tmp_path: Path) -> None:
    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    saved = registry.submit("portfolio.stage1", {}, "resume-campaign", ("portfolio_optimizer",), job_id="resume-campaign")
    registry.reserve_runtime(saved["job_id"], "campaign", "not-a-campaign")
    service = PortfolioPanelService(tmp_path, registry=registry)

    service._run(saved["job_id"])

    persisted = registry.get(saved["job_id"])
    runtime = registry.runtime(saved["job_id"])
    assert persisted["state"] == "FAILED"
    assert persisted["error"]["code"] == "CAMPAIGN_CONTRACT_VERSION_UNSUPPORTED"
    assert runtime["diagnostics"][0]["code"] == "CAMPAIGN_CONTRACT_VERSION_UNSUPPORTED"


@pytest.mark.parametrize("signal_type", (KeyboardInterrupt, SystemExit, GeneratorExit, asyncio.CancelledError))
def test_worker_reraises_process_control_signals_before_failure_persistence(monkeypatch: pytest.MonkeyPatch, tmp_path: Path, signal_type: type[BaseException]) -> None:
    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    saved = registry.submit("portfolio.stage1", {}, f"signal-{signal_type.__name__}", ("portfolio_optimizer",), job_id=f"signal-{signal_type.__name__}")
    campaign = {
        "campaign_contract_version": CAMPAIGN_CONTRACT_VERSION,
        "search_mode": CAMPAIGN_SEARCH_MODE,
        "weighted_algo_version": CAMPAIGN_WEIGHTED_ALGO_VERSION,
        "versions": {
            "campaign_contract_version": CAMPAIGN_CONTRACT_VERSION,
            "search_mode": CAMPAIGN_SEARCH_MODE,
            "weighted_algo_version": CAMPAIGN_WEIGHTED_ALGO_VERSION,
        },
    }
    registry.reserve_runtime(saved["job_id"], "campaign", campaign)
    service = PortfolioPanelService(tmp_path, registry=registry)
    original_transition = registry.transition
    calls = 0

    def fail_once(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise signal_type()
        return original_transition(*args, **kwargs)

    monkeypatch.setattr(registry, "transition", fail_once)
    with pytest.raises(signal_type):
        service._run(saved["job_id"])

    assert calls == 1
    assert registry.get(saved["job_id"])["state"] == "QUEUED"


def test_invalid_panel_create_contract_maps_before_registry_write(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: [])
    monkeypatch.setattr(
        "mrs3.panel_portfolio.validate_campaign_contract",
        lambda _campaign: (_ for _ in ()).throw(CampaignContractError("CAMPAIGN_SEARCH_MODE_REQUIRED")),
    )
    payload = {
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
        "expected_config_digest": digest,
    }
    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign(payload)
    assert (error.value.code, error.value.status, str(error.value)) == (
        "CAMPAIGN_SEARCH_MODE_REQUIRED", 422, "CAMPAIGN_SEARCH_MODE_REQUIRED"
    )
    assert service.registry.list() == []


def test_valid_frozen_campaign_reaches_adapter_fact_blocker(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    finalists = [_finalist()]

    class IdleThread:
        def __init__(self, *args, **kwargs):
            pass

        def start(self):
            pass

    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: finalists)
    monkeypatch.setattr("mrs3.panel_portfolio.threading.Thread", IdleThread)
    result = service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})
    campaign = service.registry.runtime(result["job_id"])["campaign"]
    generated = service._package_variant_generator(finalists, campaign, campaign["launch"]["profiles"])
    assert generated["variants"] == ()
    assert generated["blockers"] == [MINUTE_CAPACITY_UNAVAILABLE]


def test_settings_get_put_uses_exact_byte_digest_and_cas(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    service = PortfolioPanelService(tmp_path, path)
    assert service.settings_get()["digest"] == digest
    changed = _config()
    changed["search"]["seed"] = 2
    saved = service.settings_put({"expected_digest": digest, "document": changed})
    assert saved["state"] == "READY"
    assert path.read_bytes() != json.dumps(changed).encode()
    with pytest.raises(Exception) as error:
        service.settings_put({"expected_digest": digest, "document": changed})
    assert getattr(error.value, "code", None) == "CONFIG_CHANGED"


def test_old_v2_settings_read_defaults_without_rewriting_source(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    document = _config()
    for profile in document["profiles"].values():
        for field in (
            "max_actual_equity_dd_pct",
            "min_calculated_free_margin_reserve_pct",
            "max_calculated_account_mm_load_pct",
        ):
            profile.pop(field)
    weighted = document["search"]["weighted_search"]
    weighted["repair_attempts"] = 3
    weighted.pop("lp_solutions_per_profile")
    document["search"]["composition"]["parameters"].pop("minimum_common_days")
    source = json.dumps(document, ensure_ascii=False, indent=2).encode()
    path.write_bytes(source)
    service = PortfolioPanelService(tmp_path, path)

    settings = service.settings_get()

    assert settings["digest"] == hashlib.sha256(source).hexdigest()
    assert path.read_bytes() == source
    assert settings["document"]["profiles"]["AGGRESSIVE"]["max_actual_equity_dd_pct"] == "20"
    assert settings["document"]["profiles"]["BALANCED"]["min_calculated_free_margin_reserve_pct"] == "40"
    assert settings["document"]["profiles"]["CONSERVATIVE"]["max_calculated_account_mm_load_pct"] == "20"
    assert settings["document"]["search"]["weighted_search"] == dict(WEIGHTED_SEARCH_DEFAULTS)
    assert settings["document"]["search"]["composition"]["parameters"]["minimum_common_days"] == 14


def test_old_v2_campaign_freezes_effective_document_pair(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    document = _config()
    for profile in document["profiles"].values():
        profile.pop("max_actual_equity_dd_pct")
        profile.pop("min_calculated_free_margin_reserve_pct")
        profile.pop("max_calculated_account_mm_load_pct")
    source = json.dumps(document, ensure_ascii=False, indent=2).encode()
    path.write_bytes(source)

    class IdleThread:
        def __init__(self, *args, **kwargs):
            pass

        def start(self):
            pass

    monkeypatch.setattr("mrs3.panel_portfolio.threading.Thread", IdleThread)
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: [_finalist()])
    digest = hashlib.sha256(source).hexdigest()
    result = service.submit_campaign({
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
        "expected_config_digest": digest,
    })
    campaign = service.registry.runtime(result["job_id"])["campaign"]
    frozen = __import__("base64").b64decode(campaign["config_bytes"])

    assert hashlib.sha256(source).hexdigest() == campaign["config_digest"]
    assert hashlib.sha256(frozen).hexdigest() == campaign["frozen_config_digest"]
    assert json.loads(frozen.decode()) == campaign["config_document"]


def test_settings_save_migrates_legacy_v1_document_to_v2(tmp_path: Path) -> None:
    source = Path(__file__).parents[1] / "portfolio_optimizer.local.json.example"
    path = tmp_path / "portfolio_optimizer.local.json"
    path.write_bytes(source.read_bytes())
    service = PortfolioPanelService(tmp_path, path)

    loaded = service.settings_get()
    assert loaded["state"] == "READY"
    assert loaded["schema_version"] == 2
    saved = service.settings_put({"expected_digest": loaded["digest"], "document": loaded["document"]})

    document = json.loads(path.read_text(encoding="utf-8"))
    assert saved["schema_version"] == 2
    assert document["schema_version"] == 2
    assert all("grid" not in scenario["sizing"] for scenario in document["scenarios"].values())


def test_settings_save_materializes_absent_v2_spread_history_bypass(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    document = _config()
    document["liquidity"].pop("spread_history_bypass_pretest", None)
    source = json.dumps(document, ensure_ascii=False, indent=2).encode()
    path.write_bytes(source)
    service = PortfolioPanelService(tmp_path, path)

    loaded = service.settings_get()
    assert loaded["document"]["liquidity"]["spread_history_bypass_pretest"] is False
    assert path.read_bytes() == source

    service.settings_put({"expected_digest": loaded["digest"], "document": loaded["document"]})

    assert json.loads(path.read_text(encoding="utf-8"))["liquidity"]["spread_history_bypass_pretest"] is False


def test_v2_legacy_algorithm_versions_are_resolved_for_panel_and_save(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    document = _config()
    document["algorithm_versions"] = {
        "sizing": "portfolio_optimizer_sizing_v1",
        "ranking": "portfolio_optimizer_ranking_v1",
    }
    source = json.dumps(document, ensure_ascii=False, indent=2).encode()
    path.write_bytes(source)
    digest = hashlib.sha256(source).hexdigest()
    service = PortfolioPanelService(tmp_path, path)

    settings = service.settings_get()
    assert settings["document"]["algorithm_versions"] == {
        "sizing": "portfolio_optimizer_sizing_v2",
        "ranking": "portfolio_optimizer_ranking_v2",
    }
    assert settings["digest"] == digest
    _, config_raw, campaign_document = service._config()
    assert config_raw == source
    assert campaign_document["algorithm_versions"] == settings["document"]["algorithm_versions"]

    saved = service.settings_put({"expected_digest": digest, "document": settings["document"]})
    assert saved["document"]["algorithm_versions"] == settings["document"]["algorithm_versions"]
    assert json.loads(path.read_text(encoding="utf-8"))["algorithm_versions"] == settings["document"]["algorithm_versions"]


def test_settings_put_restores_previous_bytes_when_readback_fails(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    previous = path.read_bytes()
    changed = _config()
    changed["search"]["seed"] = 2
    original_read_bytes = Path.read_bytes
    reads = 0

    def mismatched_readback(candidate: Path) -> bytes:
        nonlocal reads
        data = original_read_bytes(candidate)
        if candidate == path:
            reads += 1
            if reads == 2:
                return data + b" "
        return data

    monkeypatch.setattr(Path, "read_bytes", mismatched_readback)
    service = PortfolioPanelService(tmp_path, path)

    with pytest.raises(PortfolioPanelError) as error:
        service.settings_put({"expected_digest": digest, "document": changed})

    assert error.value.code == "CONFIG_WRITE_FAILED"
    assert original_read_bytes(path) == previous


def test_settings_put_does_not_relock_plain_lock(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    controller = PanelController(tmp_path, tmp_path / "panel-config.json")
    service = controller._portfolio_service
    service._lock = threading.Lock()
    changed = _config()
    changed["search"]["seed"] = 2

    saved = controller.portfolio_settings_put({"expected_digest": digest, "document": changed})

    assert saved["state"] == "READY"
    assert saved["digest"] == hashlib.sha256(path.read_bytes()).hexdigest()


def test_unsupported_settings_are_read_only(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    document = _config()
    document["schema_version"] = 3
    path.write_text(json.dumps(document), encoding="utf-8")
    service = PortfolioPanelService(tmp_path, path)
    assert service.settings_get()["state"] == "UNSUPPORTED_SCHEMA"
    with pytest.raises(Exception) as error:
        service.settings_put({"expected_digest": service.settings_get()["digest"], "document": document})
    assert getattr(error.value, "code", None) == "CONFIG_UNSUPPORTED_SCHEMA"


def test_readiness_reports_exact_config_digest(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    service = PortfolioPanelService(tmp_path, path)

    readiness = service.readiness()

    assert readiness["config_digest"] == digest
    assert readiness["schema_version"] == 2
    assert readiness["policy_version"] == "portfolio_optimizer_research_risk_v1"
    assert "search" not in readiness


def test_readiness_resets_pairs_and_counts_when_finalist_read_fails(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    _write_config(path)
    with duckdb.connect(str(tmp_path / "performance.duckdb")) as connection:
        connection.execute("create table selection_runs(symbol varchar, side varchar)")
        connection.execute("insert into selection_runs values ('BTCUSDT', 'LONG')")

    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: (_ for _ in ()).throw(OSError("finalists unavailable")))

    readiness = service.readiness()

    assert readiness["available_pairs"] == []
    assert readiness["current_finalists"] == {}
    assert readiness["stage1"]["blockers"] == ["FINALISTS_UNAVAILABLE", "NO_FINALISTS"]


def test_readiness_requests_metadata_only_finalists(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    _write_config(path)
    with duckdb.connect(str(tmp_path / "performance.duckdb")) as connection:
        connection.execute("create table selection_runs(symbol varchar, side varchar)")
        connection.execute("insert into selection_runs values ('BTCUSDT', 'LONG')")
    calls: list[tuple[object, ...]] = []

    def reader(*args):
        calls.append(args)
        return ({"symbol": "BTCUSDT", "side": "LONG"},)

    service = PortfolioPanelService(tmp_path, path, finalists_reader=reader)
    readiness = service.readiness()

    assert readiness["available_pairs"] == ["BTCUSDT|LONG"]
    assert len(calls) == 1
    assert calls[0][-1] is False


def test_active_campaign_conflict_is_reported_before_source_snapshot(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    calls = []

    def reader(*_args):
        calls.append(True)
        raise RuntimeError("source must not be read")

    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    service = PortfolioPanelService(tmp_path, path, registry=registry, finalists_reader=reader)
    payload = {
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
        "expected_config_digest": digest,
    }
    config, _raw, _document = service._config()
    launch = service._normalise_campaign(payload, config, digest)
    input_digest = hashlib.sha256(json.dumps(launch, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()
    registry.submit("portfolio.stage1", {"campaign_id": "active", "input_digest": input_digest, "config_digest": digest}, "active", ("portfolio_optimizer",), job_id="active")
    registry.reserve_runtime("active", "campaign", {"campaign_id": "active", "input_digest": input_digest, "config_digest": digest})
    service._threads["active"] = object()

    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign(payload)

    assert error.value.code == "PORTFOLIO_JOB_ACTIVE_DUPLICATE"
    assert calls == []


def test_other_active_campaign_is_busy_before_source_snapshot(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    calls = []

    def reader(*_args):
        calls.append(True)
        raise RuntimeError("source must not be read")

    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    service = PortfolioPanelService(tmp_path, path, registry=registry, finalists_reader=reader)
    payload = {
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
        "expected_config_digest": digest,
    }
    registry.submit("portfolio.stage1", {"campaign_id": "other", "input_digest": "other", "config_digest": digest}, "other", ("portfolio_optimizer",), job_id="other")
    registry.reserve_runtime("other", "campaign", {"campaign_id": "other", "input_digest": "other", "config_digest": digest})
    service._threads["other"] = object()

    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign(payload)

    assert error.value.code == "PORTFOLIO_JOB_BUSY"
    assert calls == []


def test_campaign_profile_rejects_unknown_fields_exactly(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: [])
    payload = {
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1, "unexpected": True}],
        "expected_config_digest": digest,
    }

    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign(payload)

    assert error.value.code == "PORTFOLIO_CAMPAIGN_INVALID"
    assert service.registry.list() == []


def test_campaign_profile_budgets_are_independent_of_configured_total(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    service = PortfolioPanelService(tmp_path, path)
    payload = {
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [
            {"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 2},
            {"profile_id": "AGGRESSIVE", "bank_available_usdt": "10000", "max_candidates": 3},
        ],
        "expected_config_digest": digest,
    }

    config, _raw, _document = service._config()
    launch = service._normalise_campaign(payload, config, digest)

    assert [profile["max_candidates"] for profile in launch["profiles"]] == [2, 3]


@pytest.mark.parametrize("max_candidates", [0, 51])
def test_campaign_rejects_max_candidates_outside_adapter_contract(tmp_path: Path, max_candidates: int) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    service = PortfolioPanelService(tmp_path, path)
    config, _raw, _document = service._config()

    with pytest.raises(PortfolioPanelError) as error:
        service._normalise_campaign(
            {
                "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
                "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": max_candidates}],
                "expected_config_digest": digest,
            },
            config,
            digest,
        )

    assert error.value.code == "PORTFOLIO_CAMPAIGN_INVALID"
    assert error.value.field_errors == [{
        "field": "profiles[0].max_candidates",
        "code": "MAX_CANDIDATES_RANGE",
        "message": "max_candidates must be between 1 and 50",
    }]


def test_multi_profile_budget_and_summary_stay_consistent(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    document = _config()
    document["search"]["total_test_budget"] = 2
    raw = json.dumps(document, ensure_ascii=False, indent=2).encode()
    path.write_bytes(raw)
    digest = hashlib.sha256(raw).hexdigest()
    service = PortfolioPanelService(tmp_path, path)
    payload = {
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [
            {"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1},
            {"profile_id": "AGGRESSIVE", "bank_available_usdt": "10000", "max_candidates": 1},
        ],
        "expected_config_digest": digest,
    }
    config, _, _ = service._config()
    launch = service._normalise_campaign(payload, config, digest)
    summary = service._summary(
        {"campaign_id": "campaign-fixed", "launch": launch},
        (),
        (),
        ({"profile": "BALANCED"}, {"profile": "AGGRESSIVE"}),
        (),
    )

    assert summary["prepared_by_profile"] == {"AGGRESSIVE": 1, "BALANCED": 1}
    assert summary["prepared"] == 2
    assert summary["variants_created"] == 2


@pytest.mark.parametrize("value", ["999999999999999999999999999", "1.0000000000001"])
def test_campaign_money_must_fit_decimal_38_12(tmp_path: Path, value: str) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: [])

    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": value, "max_candidates": 1}], "expected_config_digest": digest})

    assert error.value.code == "PORTFOLIO_CAMPAIGN_INVALID"
    assert service.registry.list() == []


@pytest.mark.parametrize("value", ("0", "-1", "NaN", "Infinity", True, {}))
def test_campaign_bank_ceiling_must_be_finite_and_positive(tmp_path: Path, value: object) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    service = PortfolioPanelService(tmp_path, path)
    config, _raw, _document = service._config()

    with pytest.raises(PortfolioPanelError) as error:
        service._normalise_campaign({
            "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
            "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": value, "max_candidates": 1}],
            "expected_config_digest": digest,
        }, config, digest)

    assert error.value.status == 422
    assert error.value.field_errors[0]["field"] == "profiles[0].bank_available_usdt"


@pytest.mark.parametrize("bank", (pytest.param(None, id="null"), pytest.param("missing", id="missing"), pytest.param("5000", id="capped"), pytest.param(0.1, id="fractional-float")))
def test_campaign_profile_normalises_optional_bank_ceiling(tmp_path: Path, bank: object) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    service = PortfolioPanelService(tmp_path, path)
    profile = {"profile_id": "BALANCED", "max_candidates": 1}
    if bank != "missing":
        profile["bank_available_usdt"] = bank
    config, _raw, _document = service._config()

    launch = service._normalise_campaign({
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [profile],
        "expected_config_digest": digest,
    }, config, digest)

    expected = None if bank in {None, "missing"} else str(bank)
    assert launch["profiles"][0]["bank_available_usdt"] == expected


@pytest.mark.parametrize("legacy_field", ("equity_usdt", "max_balance_usdt"))
def test_campaign_profile_rejects_legacy_bank_fields(tmp_path: Path, legacy_field: str) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    service = PortfolioPanelService(tmp_path, path)
    config, _raw, _document = service._config()

    with pytest.raises(PortfolioPanelError) as error:
        service._normalise_campaign({
            "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
            "profiles": [{"profile_id": "BALANCED", legacy_field: "5000", "max_candidates": 1}],
            "expected_config_digest": digest,
        }, config, digest)

    assert error.value.field_errors[0]["code"] == "LEGACY_PROFILE_FIELD_UNSUPPORTED"


def test_failed_campaign_snapshot_creates_no_job(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)

    def broken_reader(*_args):
        raise RuntimeError("source unavailable")

    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    service = PortfolioPanelService(tmp_path, path, registry=registry, finalists_reader=broken_reader)
    payload = {
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
        "expected_config_digest": digest,
    }

    with pytest.raises(Exception):
        service.submit_campaign(payload)

    assert registry.list() == []


def test_runtime_reservation_failure_discards_queued_job_and_retry_succeeds(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    finalists = [_finalist()]
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: finalists, variant_generator=_profiled_variants)
    original = service.registry.reserve_runtime
    calls = 0

    def fail_once(job_id: str, key: str, value: object) -> None:
        nonlocal calls
        calls += 1
        if calls == 1:
            raise OSError("runtime journal is unavailable")
        original(job_id, key, value)

    monkeypatch.setattr(service.registry, "reserve_runtime", fail_once)
    payload = {
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
        "expected_config_digest": digest,
    }

    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign(payload)
    assert error.value.code == "PORTFOLIO_JOB_START_FAILED"
    assert service.registry.list() == []

    retry = service.submit_campaign(payload)
    assert retry["status"] == "QUEUED"


def test_worker_start_failure_removes_queued_job(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)

    class BrokenThread:
        def __init__(self, *args, **kwargs):
            pass

        def start(self):
            raise RuntimeError("thread start failed")

    monkeypatch.setattr("mrs3.panel_portfolio.threading.Thread", BrokenThread)
    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    service = PortfolioPanelService(tmp_path, path, registry=registry, finalists_reader=lambda *_: [])

    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})

    assert error.value.code == "PORTFOLIO_JOB_START_FAILED"
    assert registry.list() == []


def test_campaign_uses_frozen_finalist_snapshot_after_source_changes(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    calls = []
    source_rows = [[_finalist()]]

    def reader(*_args):
        calls.append(True)
        return source_rows[0]

    service = PortfolioPanelService(tmp_path, path, finalists_reader=reader, variant_generator=_profiled_variants)
    payload = {
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
        "expected_config_digest": digest,
    }
    result = service.submit_campaign(payload)
    source_rows[0] = [_finalist(strategy_id=99, result_id=100)]
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and service.job(result["job_id"])["status"] in {"QUEUED", "RUNNING"}:
        time.sleep(0.01)

    assert len(calls) == 1
    assert service.job(result["job_id"])["status"] == "SUCCEEDED"
    assert service.registry.runtime(result["job_id"])["campaign"]["finalists"][0]["strategy_id"] == 7


def test_campaign_freezes_private_weighted_rows_without_public_series_aliases(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    source_rows = [{
        "strategy_id": 7,
        "result_id": 11,
        "symbol": "BTCUSDT",
        "side": "LONG",
        "user_status": "FINALIST",
        "user_rank": 1,
        "actions": [{"timestamp_utc": "2026-01-01T00:00:00Z", "size": "1"}],
        "action_series": [{"timestamp_utc": "2026-01-01T00:00:00Z", "size": "alias"}],
        "minute_actions": [{"timestamp_utc": "2026-01-01T00:00:00Z", "size": "minute"}],
        "equity": [{"timestamp_utc": "2026-01-01T00:00:00Z", "equity": "100"}],
        "equity_series": [{"timestamp_utc": "2026-01-01T00:00:00Z", "equity": "alias"}],
    }]

    class IdleThread:
        def __init__(self, *args, **kwargs):
            pass

        def start(self):
            pass

    monkeypatch.setattr("mrs3.panel_portfolio.threading.Thread", IdleThread)
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: source_rows)
    result = service.submit_campaign({
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
        "expected_config_digest": digest,
    })

    source_rows[0]["strategy_id"] = 99
    source_rows[0]["actions"][0]["size"] = "mutated"
    source_rows[0]["equity"][0]["equity"] = "mutated"

    campaign = service.registry.runtime(result["job_id"])["campaign"]
    assert "weighted_input_rows" in campaign
    weighted = campaign["weighted_input_rows"]
    assert weighted == [{
        "strategy_id": 7,
        "result_id": 11,
        "symbol": "BTCUSDT",
        "side": "LONG",
        "actions": [{"timestamp_utc": "2026-01-01T00:00:00Z", "size": "1"}],
        "equity": [{"timestamp_utc": "2026-01-01T00:00:00Z", "equity": "100"}],
    }]
    aliases = {"actions", "action_series", "minute_actions", "equity", "equity_series"}
    assert not aliases.intersection(campaign["finalists"][0])
    public_job = service.job(result["job_id"])
    assert not aliases.intersection(public_job)
    private_prepared_fields = {
        "prepared_json", "source_digest", "preparation_version", "_prepared_cycles",
        "raw_action_series", "raw_equity_series",
    }
    assert not private_prepared_fields.intersection(campaign["finalists"][0])
    assert not private_prepared_fields.intersection(public_job)
    assert "weighted_input_rows" not in public_job


def test_campaign_prepares_exact_production_finalist_result_ids_before_strict_read(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path,
) -> None:
    import mrs3.panel_portfolio as panel_portfolio_module

    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    (tmp_path / "config.local.json").write_text(json.dumps({"duckdb_import": {"workers": 3}}), encoding="utf-8")
    metadata_rows = [
        {"result_id": 22},
        {"result_id": 11},
        {"result_id": 22},
    ]
    full_rows = [
        _finalist(strategy_id=8, result_id=22),
        _finalist(strategy_id=7, result_id=11),
    ]
    reader_calls: list[tuple[Path, tuple[tuple[str, str], ...], bool]] = []
    preparation_calls: list[tuple[Path, tuple[int, ...], int]] = []

    def production_reader(database, pairs, include_series=True):
        reader_calls.append((database, tuple(pairs), include_series))
        return metadata_rows if not include_series else full_rows

    def preparer(database, result_ids, *, workers):
        preparation_calls.append((database, tuple(result_ids), workers))

    monkeypatch.setattr(panel_portfolio_module, "read_current_finalists", production_reader)

    class IdleThread:
        def __init__(self, *args, **kwargs):
            pass

        def start(self):
            pass

    monkeypatch.setattr(panel_portfolio_module.threading, "Thread", IdleThread)
    service = PortfolioPanelService(
        tmp_path,
        path,
        optimizer_input_preparer=preparer,
    )
    result = service.submit_campaign({
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 1}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
        "expected_config_digest": digest,
    })

    assert result["status"] == "QUEUED"
    expected_pairs = (("BTCUSDT", "LONG"), ("BTCUSDT", "SHORT"))
    assert reader_calls == [
        (tmp_path / "performance.duckdb", expected_pairs, False),
        (tmp_path / "performance.duckdb", expected_pairs, True),
    ]
    assert preparation_calls == [(tmp_path / "performance.duckdb", (22, 11), 3)]


def test_campaign_rejects_invalid_production_metadata_result_id_before_preparation(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path,
) -> None:
    import mrs3.panel_portfolio as panel_portfolio_module

    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    preparation_calls: list[tuple[object, ...]] = []

    def production_reader(_database, _pairs, include_series=True):
        return [{"result_id": True}] if not include_series else []

    def preparer(*args, **kwargs):
        preparation_calls.append((*args, *kwargs.values()))

    monkeypatch.setattr(panel_portfolio_module, "read_current_finalists", production_reader)
    service = PortfolioPanelService(
        tmp_path,
        path,
        optimizer_input_preparer=preparer,
    )

    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign({
            "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
            "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
            "expected_config_digest": digest,
        })

    assert error.value.code == "PORTFOLIO_FINALISTS_UNAVAILABLE"
    assert preparation_calls == []


def test_production_preparation_lock_failure_is_a_typed_snapshot_error(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path,
) -> None:
    import mrs3.panel_portfolio as panel_portfolio_module

    path = tmp_path / "portfolio_optimizer.local.json"
    _write_config(path)
    monkeypatch.setattr(
        panel_portfolio_module,
        "read_current_finalists",
        lambda *_args, **_kwargs: [{"result_id": 11}],
    )

    def locked(*_args, **_kwargs):
        raise PerformanceV2StoreError("writer lock busy")

    service = PortfolioPanelService(tmp_path, path, optimizer_input_preparer=locked)
    with pytest.raises(PortfolioPanelError) as error:
        service._snapshot_finalists(
            _config(),
            {"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 1}]},
        )

    assert error.value.code == "PORTFOLIO_FINALISTS_UNAVAILABLE"
    assert error.value.status == 422


def test_finalist_reader_wiring_controls_production_preparation(
    tmp_path: Path,
) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    _write_config(path)

    default_service = PortfolioPanelService(tmp_path, path)
    custom_service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: ())

    assert default_service._uses_production_finalists_reader is True
    assert custom_service._uses_production_finalists_reader is False


def test_default_snapshot_prepares_only_current_finalist_artifacts(tmp_path: Path) -> None:
    from datetime import datetime, timezone
    from tests.test_performance_v2_selection import _candidate_db
    from tests.test_portfolio_input import _add_review

    path = tmp_path / "portfolio_optimizer.local.json"
    _write_config(path)
    connection = _candidate_db(tmp_path)
    database = tmp_path / "performance.duckdb"
    try:
        strategy_id, result_id = connection.execute(
            "select strategy_id, current_result_id from strategies"
        ).fetchone()
        connection.execute(
            """update strategy_results set reported_start_utc = report_start_utc,
               reported_end_utc = report_end_utc, effective_start_utc = report_start_utc,
               effective_end_utc = report_end_utc, sizing_use_upnl = true,
               sizing_use_frozen_balance = true, sizing_use_fix = false,
               sizing_balance_percentage_long = 100, sizing_risk_long = 1,
               sizing_max_balance = 0 where result_id = ?""",
            [result_id],
        )
        connection.execute("update strategy_actions set price = 10, cost = 10")
        _add_review(
            connection,
            run_id="run-default",
            review_id="review-default",
            symbol="BTCUSDT",
            side="LONG",
            strategy_id=strategy_id,
            result_id=result_id,
            user_status="FINALIST",
            selection_time=datetime(2026, 1, 1, tzinfo=timezone.utc),
        )
    finally:
        connection.close()
    (tmp_path / "strategy_performance.duckdb").replace(database)

    service = PortfolioPanelService(tmp_path, path)
    finalists, weighted_rows = service._snapshot_finalists(
        _config(),
        {"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}]},
    )

    assert [row["result_id"] for row in finalists] == [result_id]
    assert len(weighted_rows) == 1
    with duckdb.connect(str(database), read_only=True) as connection:
        assert connection.execute(
            "select result_id from optimizer_prepared_inputs order by result_id"
        ).fetchall() == [(result_id,)]


def test_campaign_resolves_ordered_series_aliases_into_private_weighted_rows(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    source_rows = [{
        "strategy_id": 7,
        "result_id": 11,
        "symbol": "BTCUSDT",
        "side": "LONG",
        "user_status": "FINALIST",
        "user_rank": 1,
        "action_series": [{"timestamp_utc": "2026-01-01T00:00:00Z", "size": "1"}],
        "equity_series": [{"timestamp_utc": "2026-01-01T00:00:00Z", "equity": "100"}],
    }]

    class IdleThread:
        def __init__(self, *args, **kwargs):
            pass

        def start(self):
            pass

    monkeypatch.setattr("mrs3.panel_portfolio.threading.Thread", IdleThread)
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: source_rows)
    result = service.submit_campaign({
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
        "expected_config_digest": digest,
    })

    campaign = service.registry.runtime(result["job_id"])["campaign"]
    assert campaign["weighted_input_rows"][0]["actions"] == source_rows[0]["action_series"]
    assert campaign["weighted_input_rows"][0]["equity"] == source_rows[0]["equity_series"]


def test_campaign_rejects_missing_weighted_series_at_snapshot(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    source_rows = [{"strategy_id": 7, "result_id": 11, "symbol": "BTCUSDT", "side": "LONG"}]

    class IdleThread:
        def __init__(self, *args, **kwargs):
            pass

        def start(self):
            pass

    monkeypatch.setattr("mrs3.panel_portfolio.threading.Thread", IdleThread)
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: source_rows)

    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign({
            "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
            "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
            "expected_config_digest": digest,
        })

    assert error.value.code == "PORTFOLIO_FINALISTS_UNAVAILABLE"


def test_campaign_rejects_non_json_weighted_series_at_snapshot(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    source_rows = [_finalist(equity=[{"timestamp_utc": "2026-01-01T00:00:00Z", "equity": float("nan")}])]

    class IdleThread:
        def __init__(self, *args, **kwargs):
            pass

        def start(self):
            pass

    monkeypatch.setattr("mrs3.panel_portfolio.threading.Thread", IdleThread)
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: source_rows)

    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign({
            "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
            "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
            "expected_config_digest": digest,
        })

    assert error.value.code == "PORTFOLIO_FINALISTS_UNAVAILABLE"
    assert error.value.status == 422


def test_campaign_rejects_non_mapping_finalist_at_snapshot(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)

    class IdleThread:
        def __init__(self, *args, **kwargs):
            pass

        def start(self):
            pass

    monkeypatch.setattr("mrs3.panel_portfolio.threading.Thread", IdleThread)
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: [None])

    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign({
            "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
            "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
            "expected_config_digest": digest,
        })

    assert error.value.code == "PORTFOLIO_FINALISTS_UNAVAILABLE"
    assert error.value.status == 422


@pytest.mark.parametrize("missing", ("symbol", "side", "strategy_id", "result_id"))
def test_campaign_rejects_missing_weighted_identity_at_snapshot(monkeypatch: pytest.MonkeyPatch, tmp_path: Path, missing: str) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    source_row = {
        "strategy_id": 7,
        "result_id": 11,
        "symbol": "BTCUSDT",
        "side": "LONG",
        "actions": [],
        "equity": [],
    }
    source_row.pop(missing)

    class IdleThread:
        def __init__(self, *args, **kwargs):
            pass

        def start(self):
            pass

    monkeypatch.setattr("mrs3.panel_portfolio.threading.Thread", IdleThread)
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: [source_row])

    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign({
            "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
            "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
            "expected_config_digest": digest,
        })

    assert error.value.code == "PORTFOLIO_FINALISTS_UNAVAILABLE"
    assert error.value.field_errors == [{
        "field": f"finalists[0].{missing}",
        "code": "MISSING_FINALIST_IDENTITY",
        "message": f"finalist {missing} identity is unavailable",
    }]


@pytest.mark.parametrize("missing", ("actions", "equity"))
def test_campaign_rejects_missing_weighted_series_with_field_error(monkeypatch: pytest.MonkeyPatch, tmp_path: Path, missing: str) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    source_row = {
        "strategy_id": 7,
        "result_id": 11,
        "symbol": "BTCUSDT",
        "side": "LONG",
        "actions": [],
        "equity": [],
    }
    source_row.pop(missing)

    class IdleThread:
        def __init__(self, *args, **kwargs):
            pass

        def start(self):
            pass

    monkeypatch.setattr("mrs3.panel_portfolio.threading.Thread", IdleThread)
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: [source_row])

    with pytest.raises(PortfolioPanelError) as error:
        service.submit_campaign({
            "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
            "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
            "expected_config_digest": digest,
        })

    assert error.value.code == "PORTFOLIO_FINALISTS_UNAVAILABLE"
    assert error.value.field_errors == [{
        "field": f"finalists[0].{missing}",
        "code": "MISSING_FINALIST_SERIES",
        "message": f"finalist {missing} series is unavailable",
    }]


def test_invalid_published_workbook_fails_without_download(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    finalists = [_finalist()]

    def broken_builder(workbook_path, *_args):
        workbook_path.write_bytes(b"not an xlsx")
        return workbook_path

    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: finalists, variant_generator=_profiled_variants, workbook_builder=broken_builder)
    result = service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and service.job(result["job_id"])["status"] in {"QUEUED", "RUNNING"}:
        time.sleep(0.01)

    assert service.job(result["job_id"])["status"] == "FAILED"
    assert not (tmp_path / ".portfolio-results" / result["campaign_id"] / "stage1.xlsx").exists()
    assert not (tmp_path / ".portfolio-results" / result["campaign_id"] / "stage1-executables.json").exists()
    with pytest.raises(Exception):
        service.workbook(result["campaign_id"])


def test_commit_failure_rolls_back_published_workbook(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    finalists = [_finalist()]
    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    original_sync = registry.sync

    def fail_commit(job_id, status, *, runtime=None):
        if status.get("state") == "COMMITTED":
            raise PanelJobError("COMMIT_FAILED")
        return original_sync(job_id, status, runtime=runtime)

    registry.sync = fail_commit
    service = PortfolioPanelService(tmp_path, path, registry=registry, finalists_reader=lambda *_: finalists, variant_generator=_profiled_variants)
    result = service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and service.job(result["job_id"])["status"] in {"QUEUED", "RUNNING"}:
        time.sleep(0.01)

    assert service.job(result["job_id"])["status"] == "FAILED"
    assert not (tmp_path / ".portfolio-results" / result["campaign_id"] / "stage1.xlsx").exists()
    assert not (tmp_path / ".portfolio-results" / result["campaign_id"] / "stage1-executables.json").exists()
    with pytest.raises(PortfolioPanelError) as error:
        service.workbook(result["campaign_id"])
    assert error.value.code == "PORTFOLIO_JOB_WORKBOOK_UNAVAILABLE"


def test_journal_and_indeterminate_progress_are_public(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    finalists = [_finalist()]
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: finalists, variant_generator=_profiled_variants)
    result = service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and service.job(result["job_id"])["status"] in {"QUEUED", "RUNNING"}:
        time.sleep(0.01)

    job = service.job(result["job_id"])
    assert job["status"] == "SUCCEEDED"
    assert job["journal"]
    assert all(entry["code"].startswith("PORTFOLIO_JOB_") for entry in job["journal"])
    assert all(entry["stage"] in STAGES for entry in job["journal"])
    assert any(entry["code"] == "PORTFOLIO_JOB_COMPLETED" for entry in job["journal"])
    assert job["stage"]["completed"] == 1
    assert job["stage"]["total"] == 1
    assert job["stage"]["percent"] == 100


def test_zero_variants_fail_without_results_or_download(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    finalists = [_finalist()]
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: finalists, variant_generator=lambda *_: [])
    result = service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and service.job(result["job_id"])["status"] in {"QUEUED", "RUNNING"}:
        time.sleep(0.01)

    job = service.job(result["job_id"])
    assert job["status"] == "FAILED"
    assert job["diagnostics"][0]["code"] == "PORTFOLIO_JOB_VARIANTS_NOT_READY"
    with pytest.raises(PortfolioPanelError) as error:
        service.workbook(result["campaign_id"])
    assert error.value.code == "PORTFOLIO_JOB_WORKBOOK_UNAVAILABLE"
    assert job["journal"]
    assert all(entry["code"].startswith("PORTFOLIO_JOB_") for entry in job["journal"])
    assert all(entry["stage"] in STAGES for entry in job["journal"])
    assert any(entry["code"] == "PORTFOLIO_JOB_VARIANTS_NOT_READY" for entry in job["journal"])


def test_zero_adapter_variants_report_actionable_gate_details(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    finalists = [_finalist()]
    generated = {
        "variants": (),
        "blockers": ("BALANCED:INSUFFICIENT_DIRECTIONAL_UNIVERSE",),
        "excluded": ({"symbol": "BTCUSDT", "selection_reason": "SIZE_BELOW_MINIMUM_QTY"},),
    }
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: finalists, variant_generator=lambda *_: generated)
    result = service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and service.job(result["job_id"])["status"] in {"QUEUED", "RUNNING"}:
        time.sleep(0.01)

    assert "BTCUSDT:SIZE_BELOW_MINIMUM_QTY" in service.job(result["job_id"])["diagnostics"][0]["message"]


def test_typed_postsearch_blocker_survives_snapshot_cleanup_and_restart(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    typed_blocker = "PROFILE:AGGRESSIVE:WEIGHTED_CANDIDATE_BANK_INVALID"
    generated = {"variants": (), "blockers": (typed_blocker,), "excluded": ()}
    service = PortfolioPanelService(
        tmp_path,
        path,
        finalists_reader=lambda *_: [_finalist()],
        variant_generator=lambda *_: generated,
    )
    result = service.submit_campaign({
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "AGGRESSIVE", "bank_available_usdt": "10000", "max_candidates": 1}],
        "expected_config_digest": digest,
    })

    job = _wait_stage1(service, result)
    assert job["status"] == "FAILED"
    assert job["diagnostics"][0]["code"] == "PORTFOLIO_JOB_VARIANTS_NOT_READY"
    assert typed_blocker in job["diagnostics"][0]["message"]
    assert any(typed_blocker in entry["text"] for entry in job["journal"])
    snapshot = tmp_path / ".portfolio-results" / result["campaign_id"] / "campaign-input.json.gz"
    cleanup_deadline = time.monotonic() + 1
    while time.monotonic() < cleanup_deadline and snapshot.exists():
        time.sleep(0.01)
    assert not snapshot.exists()

    persisted = None
    journal_path = tmp_path / ".panel-jobs.json"
    persisted_deadline = time.monotonic() + 1
    while time.monotonic() < persisted_deadline and persisted is None:
        try:
            persisted = json.loads(journal_path.read_text(encoding="utf-8"))
        except PermissionError:
            time.sleep(0.01)
    assert persisted is not None
    record = persisted[result["job_id"]]
    assert record["state"] == "FAILED"
    assert typed_blocker in record["runtime"]["diagnostics"][0]["message"]
    assert any(typed_blocker in entry["text"] for entry in record["runtime"]["journal"])

    restarted = PortfolioPanelService(
        tmp_path,
        path,
        registry=PanelJobRegistry(tmp_path / ".panel-jobs.json"),
    )
    recovered = restarted.job(result["job_id"])
    assert recovered["status"] == "FAILED"
    assert typed_blocker in recovered["diagnostics"][0]["message"]
    assert any(typed_blocker in entry["text"] for entry in recovered["journal"])


def test_job_messages_redact_paths_and_secret_terms(tmp_path: Path) -> None:
    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    saved = registry.submit("portfolio.stage1", {}, "campaign", ("portfolio_optimizer",))
    registry.reserve_runtime(saved["job_id"], "campaign", {"campaign_id": "campaign"})
    registry.reserve_runtime(saved["job_id"], "journal", [{"text": r"C:\private\report token", "stage": "VALIDATE_SNAPSHOT"}])
    service = PortfolioPanelService(tmp_path, registry=registry)

    text = service.job(saved["job_id"])["journal"][0]["text"]

    assert "C:\\private" not in text
    assert "token" not in text
    assert "<redacted-" in text


@pytest.mark.parametrize("value", ["strategy/v1", "https://example.test/a/b", "foo/bar"])
def test_redaction_preserves_non_path_slashes(value: str) -> None:
    assert _redact_text(value) == value
    assert _safe_cell(value) == value


@pytest.mark.parametrize("value", [r"C:\private\report", r"\\server\share\report", "/var/lib/mrs3", "/a", "../private/report", "foo/../private"])
def test_redaction_hides_absolute_and_traversal_paths(value: str) -> None:
    assert _redact_text(value) == "<redacted-path>"
    assert _safe_cell(value) == "<redacted-path>"


def test_known_stage_progress_reports_local_denominator(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    finalists = [_finalist()]
    entered = threading.Event()
    release = threading.Event()

    def generator(*_args):
        entered.set()
        assert release.wait(timeout=3)
        return []

    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: finalists, variant_generator=generator)
    result = service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})
    assert entered.wait(timeout=3)
    try:
        job = service.job(result["job_id"])
        assert job["stage"] == {"index": 3, "name": "GENERATE_VARIANTS", "status": "RUNNING", "completed": 0, "total": 1, "percent": 0}
        assert job["overall_percent"] == 42
    finally:
        release.set()
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and service.job(result["job_id"])["status"] in {"QUEUED", "RUNNING"}:
        time.sleep(0.01)


def test_worker_start_preserves_queued_cancellation(tmp_path: Path) -> None:
    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    saved = registry.submit("portfolio.stage1", {}, "campaign", ("portfolio_optimizer",))
    registry.cancel(saved["job_id"])
    service = PortfolioPanelService(tmp_path, registry=registry)

    service._run(saved["job_id"])

    assert registry.get(saved["job_id"])["state"] == "CANCELLED"


def test_job_reports_deleted_settings_as_changed_since_freeze(tmp_path: Path) -> None:
    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    saved = registry.submit("portfolio.stage1", {}, "campaign", ("portfolio_optimizer",))
    registry.reserve_runtime(
        saved["job_id"],
        "campaign",
        {"campaign_id": "campaign", "config_digest": "frozen"},
    )
    service = PortfolioPanelService(tmp_path, registry=registry)

    job = service.job(saved["job_id"])
    assert job["stage"]["status"] == "INTERRUPTED"
    assert job["settings_changed_since_freeze"] is True


def test_summary_groups_exclusions_and_profile_counts() -> None:
    campaign = {
        "campaign_id": "campaign-fixed",
        "launch": {"profiles": [{"profile_id": "BALANCED", "max_candidates": 1}]},
    }

    summary = PortfolioPanelService._summary(
        campaign,
        ({"strategy_id": 1}, {"strategy_id": 2}),
        ({"strategy_id": 1},),
        ({"profile": "BALANCED"}, {"profile": "BALANCED"}),
        ({"selection_reason": "USER_RANK_CUTOFF"},),
        optimizer_excluded=({"selection_reason": "OPEN_POLICY"},),
        blockers=("BALANCED:OPEN_POLICY:RANKING_POLICY_REQUIRED",),
    )

    assert summary["variants_by_profile"] == {"BALANCED": 2}
    assert summary["prepared_by_profile"] == {"BALANCED": 2}
    assert summary["prepared"] == 2
    assert summary["finalist_exclusions_by_reason"] == {"USER_RANK_CUTOFF": 1}
    assert summary["optimizer_exclusions_by_reason"] == {"OPEN_POLICY": 1}


def test_workbook_bytes_are_deterministic_after_metadata_hide(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    _write_config(path)
    service = PortfolioPanelService(tmp_path, path)
    campaign = {"campaign_id": "campaign-fixed", "input_digest": "i", "config_digest": "c", "versions": {"policy_version": "p"}, "created_at_utc": "2000-01-01T00:00:00Z"}
    first = tmp_path / "one.xlsx"
    second = tmp_path / "two.xlsx"
    service._write_workbook(first, campaign, (), (), (), ())
    service._write_workbook(second, campaign, (), (), (), ())

    assert first.read_bytes() == second.read_bytes()
    with zipfile.ZipFile(first) as archive:
        core = archive.read("docProps/core.xml")
    assert b"2000-01-01T00:00:00Z" in core


def test_workbook_keeps_raw_series_out_of_public_cells(tmp_path: Path) -> None:
    service = PortfolioPanelService(tmp_path)
    campaign = {
        "campaign_id": "campaign-fixed",
        "input_digest": "i",
        "config_digest": "c",
        "versions": {"policy_version": "p", "algorithm_versions": {}},
        "created_at_utc": "2000-01-01T00:00:00Z",
    }
    selected = ({
        "strategy_id": 1,
        "result_id": 2,
        "symbol": "BTCUSDT",
        "side": "LONG",
        "actions": "PRIVATE_ACTIONS",
        "equity": "PRIVATE_EQUITY",
        "raw": "PRIVATE_RAW",
        "prepared_json": "PRIVATE_PREPARED_JSON",
        "source_digest": "PRIVATE_SOURCE_DIGEST",
        "preparation_version": "PRIVATE_PREPARATION_VERSION",
        "_prepared_cycles": "PRIVATE_PREPARED_CYCLES",
        "raw_action_series": "PRIVATE_RAW_ACTION_SERIES",
        "raw_equity_series": "PRIVATE_RAW_EQUITY_SERIES",
    },)
    variants = ({
        "candidate_id": "candidate",
        "profile": "BALANCED",
        "members": ({"symbol": "BTCUSDT", "raw": "PRIVATE_MEMBER_RAW"},),
        "metrics": {"raw": "PRIVATE_METRICS_RAW"},
    },)
    workbook_path = tmp_path / "raw-boundary.xlsx"

    service._write_workbook(workbook_path, campaign, (), selected, variants, ())

    workbook = load_workbook(workbook_path, data_only=False)
    try:
        values = [cell.value for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row]
        headers = {
            cell.value
            for worksheet in workbook.worksheets
            for cell in worksheet[1]
        }
    finally:
        workbook.close()
    assert not {
        "PRIVATE_ACTIONS", "PRIVATE_EQUITY", "PRIVATE_RAW", "PRIVATE_MEMBER_RAW", "PRIVATE_METRICS_RAW",
        "PRIVATE_PREPARED_JSON", "PRIVATE_SOURCE_DIGEST", "PRIVATE_PREPARATION_VERSION",
        "PRIVATE_PREPARED_CYCLES", "PRIVATE_RAW_ACTION_SERIES", "PRIVATE_RAW_EQUITY_SERIES",
    }.intersection(values)
    assert not {
        "prepared_json", "source_digest", "preparation_version", "_prepared_cycles",
        "raw_action_series", "raw_equity_series",
    }.intersection(headers)


def test_workbook_preserves_decimal_cells_as_exact_text(tmp_path: Path) -> None:
    service = PortfolioPanelService(tmp_path)
    campaign = {"campaign_id": "campaign-fixed", "input_digest": "i", "config_digest": "c", "versions": {"policy_version": "p"}, "created_at_utc": "2000-01-01T00:00:00Z"}
    workbook_path = tmp_path / "decimal.xlsx"
    service._write_workbook(
        workbook_path,
        campaign,
        (),
        (),
        ({"candidate_id": "candidate", "profile": "BALANCED", "scheduling_score": Decimal("12345678901234567890.123456789012"), "directions": {}},),
        (),
    )
    workbook = load_workbook(workbook_path, data_only=False)
    try:
        cell = workbook["Portfolios"][2][5]
        assert cell.value == "12345678901234567890.123456789012"
        assert isinstance(cell.value, str)
    finally:
        workbook.close()


def test_workbook_exposes_adapter_capacity_spread_and_warning_diagnostics(tmp_path: Path) -> None:
    service = PortfolioPanelService(tmp_path)
    campaign = {"campaign_id": "campaign-fixed", "input_digest": "i", "config_digest": "c", "versions": {"policy_version": "p"}, "created_at_utc": "2000-01-01T00:00:00Z"}
    member = {
        "strategy_id": 1, "result_id": 2, "symbol": "BTCUSDT", "side": "LONG",
        "position_size_usdt": Decimal("600"), "maximum_closing_quantity": Decimal("6"),
        "planned_leverage": Decimal("50"), "max_drawdown_pct": Decimal("10"),
        "capacity_status": "PRELIMINARY",
        "calendar_7d": SimpleNamespace(available_days=3, mean_minute_turnover=Decimal("2000"), rounded_cap_usdt=Decimal("600")),
        "weekday_5d": SimpleNamespace(available_days=2, mean_minute_turnover=Decimal("1800"), rounded_cap_usdt=Decimal("500")),
        "spread_status": "CLEAR", "spread_mean_bps": Decimal("8.5"),
        "sizing_digest": "sizing", "capacity_digest": "capacity", "reference_digest": "reference",
    }
    workbook_path = tmp_path / "adapter.xlsx"
    service._write_workbook(
        workbook_path, campaign, (), (),
        ({"candidate_id": "candidate", "profile": "BALANCED", "members": (member,)},), (),
        warnings=("LIQUIDITY_CAPACITY_PRELIMINARY",),
    )

    workbook = load_workbook(workbook_path, data_only=False)
    try:
        headers = {cell.value: index for index, cell in enumerate(workbook["Members"][1])}
        row = workbook["Members"][2]
        assert row[headers["7d Position Cap USDT"]].value == "600"
        assert row[headers["5d Analytic Cap USDT"]].value == "500"
        assert row[headers["Spread Status"]].value == "CLEAR"
        summary_cells = {row[0].value: row[1] for row in workbook["Summary"].iter_rows(min_row=2)}
        summary = {key: cell.value for key, cell in summary_cells.items()}
        assert "LIQUIDITY_CAPACITY_PRELIMINARY" in summary["optimizer_warnings"]
    finally:
        workbook.close()


def test_weighted_summary_labels_missing_bank_ceiling_as_uncapped() -> None:
    summary = PortfolioPanelService._summary(
        {
            "campaign_id": "campaign-fixed",
            "launch": {"profiles": [{"profile_id": "BALANCED", "bank_available_usdt": None, "max_candidates": 1}]},
        },
        (), (),
        ({
            "candidate_id": "candidate-weighted",
            "profile": "BALANCED",
            "search_mode": "WEIGHTED_V1",
            "metrics": {"required_bank_usdt": Decimal("1200")},
        },),
        (),
    )

    assert summary["B required USDT"] == Decimal("1200")
    assert summary["B available USDT"] == "UNCAPPED"


def test_weighted_payload_pairing_keeps_same_symbol_long_and_short_distinct() -> None:
    members = (
        {"strategy_id": 1, "result_id": 11, "symbol": "BTCUSDT", "side": "LONG", "x_usdt": Decimal("100")},
        {"strategy_id": 2, "result_id": 22, "symbol": "BTCUSDT", "side": "SHORT", "x_usdt": Decimal("100")},
    )
    payloads = (
        {"side": "LONG", "strategy": {"basic": {"symbol": "BTCUSDT"}}},
        {"side": "SHORT", "strategy": {"basic": {"symbol": "BTCUSDT"}}},
    )
    pairs = _weighted_payload_pairs({"members": members, "strategy_payloads": payloads})
    assert [payload["side"] for _member, payload in pairs] == ["LONG", "SHORT"]


def test_weighted_payload_pairing_ignores_missing_payload_side_without_crashing() -> None:
    members = ({"strategy_id": 1, "result_id": 11, "symbol": "BTCUSDT", "side": "SHORT", "x_usdt": Decimal("100")},)
    payloads = ({"strategy": {"basic": {"symbol": "BTCUSDT"}}},)

    assert _weighted_payload_pairs({"members": members, "strategy_payloads": payloads}) == ((members[0], None),)


def test_workbook_uses_final_pretest_metrics_rank_and_composition_member_size(tmp_path: Path) -> None:
    service = PortfolioPanelService(tmp_path)
    campaign = {
        "campaign_id": "campaign-fixed", "input_digest": "i", "config_digest": "c",
        "versions": {"policy_version": "p"}, "created_at_utc": "2000-01-01T00:00:00Z",
        "launch": {"profiles": [{"profile_id": "BALANCED", "max_candidates": 2}]},
        "config_document": {"search": {"max_enumerated_combinations": 100}},
    }
    member = {
        "strategy_id": 1, "result_id": 2, "symbol": "BTCUSDT", "side": "LONG",
        "quantity": Decimal("4"), "maximum_closing_quantity": Decimal("99"),
        "actual_size_usdt": Decimal("123"), "position_size_usdt": Decimal("999"),
    }
    variant = {
        "candidate_id": "candidate", "profile": "BALANCED", "search_mode": "PRETEST_PROXY",
        "daily_pretest_rank": 1, "final_pretest_rank": 2, "refinement": "MINUTE",
        "members": (member,), "evaluations": 3,
        "pretest_period": {"start_utc": "2026-01-01", "end_utc": "2026-01-15", "coverage_pct": {}},
        "metrics": {
            "metric_basis": "PRETEST_PROXY", "proxy_pnl_usdt": Decimal("77"),
            "proxy_max_drawdown_pct": Decimal("2"), "proxy_recovery_factor": Decimal("3"),
            "proxy_reserve_usdt": Decimal("100"), "k": Decimal("0.5"), "k1": Decimal("0.6"),
            "tested_size_usdt": Decimal("100"), "actual_size_usdt": Decimal("123"),
            "tested_size_basis": "SOURCE_INITIAL_BALANCE_X_OPENING_LOT", "joint_metrics": "NOT_TESTED",
        },
    }
    path = tmp_path / "final-pretest.xlsx"
    service._write_workbook(path, campaign, (), (), (variant,), ())

    workbook = load_workbook(path, data_only=False)
    try:
        portfolio_headers = {cell.value: index for index, cell in enumerate(workbook["Portfolios"][1])}
        portfolio = workbook["Portfolios"][2]
        assert portfolio[portfolio_headers["Pretest PnL USDT"]].value == "77"
        assert portfolio[portfolio_headers["Actual Size USDT"]].value == "123"
        assert portfolio[portfolio_headers["Refinement"]].value == "MINUTE"
        assert portfolio[portfolio_headers["Final PRETEST Rank"]].value == 2
        member_headers = {cell.value: index for index, cell in enumerate(workbook["Members"][1])}
        members = workbook["Members"][2]
        assert members[member_headers["Quantity"]].value == "4"
        assert members[member_headers["Notional USDT"]].value == "123"
        for header in ("Finalist", "Side", "x USDT", "C USDT", "q", "Priority", "max_balance", "Source Scale", "Hold90", "Warnings"):
            assert members[member_headers[header]].value is None
        status_headers = {cell.value: index for index, cell in enumerate(workbook["Profile Status"][1])}
        status = workbook["Profile Status"][2]
        assert status[status_headers["Metric Basis"]].value == "PRETEST_PROXY"
        assert status[status_headers["Joint Metrics"]].value == "NOT_TESTED"
    finally:
        workbook.close()


def test_workbook_does_not_invent_gate_or_counts(tmp_path: Path) -> None:
    service = PortfolioPanelService(tmp_path)
    campaign = {"campaign_id": "campaign-fixed", "input_digest": "i", "config_digest": "c", "versions": {"policy_version": "p"}, "created_at_utc": "2000-01-01T00:00:00Z"}
    workbook_path = tmp_path / "missing-evidence.xlsx"
    service._write_workbook(workbook_path, campaign, (), (), ({"candidate_id": "candidate", "profile": "BALANCED"},), ())
    workbook = load_workbook(workbook_path, data_only=False)
    try:
        row = workbook["Portfolios"][2]
        assert row[6].value is None
        assert row[7].value is None
        assert row[12].value == "UNKNOWN"
    finally:
        workbook.close()


def test_variant_validation_keeps_full_pretest_universe_without_max_candidates(tmp_path: Path) -> None:
    campaign = {
        "campaign_id": "campaign-fixed",
        "input_digest": "i",
        "config_digest": "c",
        "versions": {"policy_version": "p"},
        "created_at_utc": "2000-01-01T00:00:00Z",
        "launch": {"profiles": [{"profile_id": "BALANCED", "max_candidates": 1}]},
    }
    variants = ({"candidate_id": "first", "profile": "BALANCED"}, {"candidate_id": "second", "profile": "BALANCED"})
    kept, excluded, blockers = PortfolioPanelService._cap_variants(variants, campaign["launch"]["profiles"])
    assert tuple(item["candidate_id"] for item in kept) == ("first", "second")
    assert excluded == ()
    assert blockers == ()
    summary = PortfolioPanelService._summary(campaign, (), (), kept, (), optimizer_excluded=excluded, blockers=blockers)
    assert summary["variants_created"] == 2
    assert summary["optimizer_excluded"] == 0
    assert summary["prepared_by_profile"] == {"BALANCED": 2}
    assert summary["prepared"] == 2
    workbook_path = tmp_path / "capped.xlsx"
    PortfolioPanelService(tmp_path)._write_workbook(workbook_path, campaign, (), (), kept, (), optimizer_excluded=excluded, blockers=blockers)
    workbook = load_workbook(workbook_path, data_only=False)
    try:
        assert workbook["Portfolios"].max_row == 3
        assert workbook["Excluded"].max_row == 1
    finally:
        workbook.close()


def test_variant_without_known_profile_is_excluded_and_reported(tmp_path: Path) -> None:
    campaign = {
        "campaign_id": "campaign-fixed",
        "input_digest": "i",
        "config_digest": "c",
        "versions": {"policy_version": "p"},
        "created_at_utc": "2000-01-01T00:00:00Z",
        "launch": {"profiles": [{"profile_id": "BALANCED", "max_candidates": 1}]},
    }
    variants = ({"candidate_id": "unknown"}, {"candidate_id": "known", "profile": "BALANCED"})

    kept, excluded, blockers = PortfolioPanelService._cap_variants(variants, campaign["launch"]["profiles"])

    assert tuple(item["candidate_id"] for item in kept) == ("known",)
    assert excluded[0]["selection_reason"] == "UNKNOWN_PROFILE"
    assert blockers == ("UNKNOWN_PROFILE",)
    summary = PortfolioPanelService._summary(campaign, (), (), kept, (), optimizer_excluded=excluded, blockers=blockers)
    assert summary["variants_created"] == 1
    assert summary["prepared_by_profile"] == {"BALANCED": 1}
    workbook_path = tmp_path / "unknown-profile.xlsx"
    PortfolioPanelService(tmp_path)._write_workbook(workbook_path, campaign, (), (), kept, (), optimizer_excluded=excluded, blockers=blockers)
    workbook = load_workbook(workbook_path, data_only=False)
    try:
        assert workbook["Portfolios"][2][2].value == "BALANCED"
        assert workbook["Excluded"][2][8].value == "UNKNOWN_PROFILE"
    finally:
        workbook.close()


def test_verify_workbook_rejects_hyperlinks_without_publishing(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    finalists = [_finalist()]

    def linked_builder(workbook_path, *_args):
        workbook = Workbook()
        workbook.active.title = "Summary"
        for name in ("Finalists", "Portfolios", "Members", "Excluded", "Metadata"):
            workbook.create_sheet(name)
        workbook["Summary"]["A1"].hyperlink = "https://example.test/"
        workbook.save(workbook_path)
        workbook.close()
        return workbook_path

    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: finalists, variant_generator=_profiled_variants, workbook_builder=linked_builder)
    result = service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and service.job(result["job_id"])["status"] in {"QUEUED", "RUNNING"}:
        time.sleep(0.01)

    assert service.job(result["job_id"])["status"] == "FAILED"
    assert not (tmp_path / ".portfolio-results" / result["campaign_id"] / "stage1.xlsx").exists()
    assert not (tmp_path / ".portfolio-results" / result["campaign_id"] / "stage1-executables.json").exists()
    with pytest.raises(PortfolioPanelError):
        service.workbook(result["campaign_id"])


def test_package_adapter_status_is_a_structured_blocker(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    monkeypatch.setattr("mrs3.portfolio.adapter.run_portfolio_adapter", lambda *_args, **_kwargs: SimpleNamespace(variants=(), blockers=("BALANCED:POSITION_SIZING_FAILED",), excluded=(), warnings=()))

    result = PortfolioPanelService(tmp_path)._package_variant_generator(({"strategy_id": 1},), {}, ({"profile_id": "BALANCED"},))

    assert result["blockers"] == ["BALANCED:POSITION_SIZING_FAILED"]


def test_package_adapter_uses_duckdb_import_workers_only(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    (tmp_path / "config.local.json").write_text(json.dumps({
        "duckdb_import": {"workers": 3},
        "direct_materialization": {"workers": 99},
    }), encoding="utf-8")
    observed = {}

    def fake_adapter(*_args, **kwargs):
        observed.update(kwargs)
        return SimpleNamespace(variants=(), blockers=(), excluded=(), warnings=())

    monkeypatch.setattr("mrs3.portfolio.adapter.run_portfolio_adapter", fake_adapter)
    campaign = {"config_document": {"search": {"weighted_search": {}}}}
    PortfolioPanelService(tmp_path)._package_variant_generator((), campaign, ())

    assert observed["workers"] == 3
    assert set(observed) == {"workspace_root", "workers"}
    (tmp_path / "config.local.json").write_text(json.dumps({"direct_materialization": {"workers": 99}}), encoding="utf-8")
    observed.clear()
    PortfolioPanelService(tmp_path)._package_variant_generator((), campaign, ())
    assert observed["workers"] == DuckDBImportSettings().workers


def test_package_adapter_real_contract_rejects_invalid_campaign(tmp_path: Path) -> None:
    result = PortfolioPanelService(tmp_path)._package_variant_generator(
        ({"strategy_id": 1, "result_id": 2, "symbol": "BTCUSDT", "side": "LONG"},),
        {},
        ({"profile_id": "BALANCED", "max_candidates": 1},),
    )
    assert result["variants"] == ()
    assert result["blockers"] == ["CAMPAIGN_CONTRACT_VERSION_UNSUPPORTED"]


def test_package_adapter_passes_variants_to_profile_cap(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    passing = ({"candidate_id": "first", "profile": "BALANCED"}, {"candidate_id": "second", "profile": "BALANCED"})
    monkeypatch.setattr(
        "mrs3.portfolio.adapter.run_portfolio_adapter",
        lambda *_args, **_kwargs: SimpleNamespace(variants=passing, blockers=(), excluded=(), warnings=()),
    )
    profile = {"profile_id": "BALANCED", "max_candidates": 1}
    result = PortfolioPanelService(tmp_path)._package_variant_generator(({"strategy_id": 1},), {}, (profile,))
    assert result["variants"][0]["profile"] == "BALANCED"
    kept, excluded, blockers = PortfolioPanelService._cap_variants(result["variants"], (profile,))
    summary = PortfolioPanelService._summary(
        {"campaign_id": "campaign", "launch": {"profiles": [profile]}},
        (),
        (),
        kept,
        (),
        optimizer_excluded=excluded,
        blockers=blockers,
    )
    assert len(kept) == 2
    assert summary["variants_created"] == 2
    assert summary["prepared_by_profile"] == {"BALANCED": 2}


def test_package_search_exception_is_typed_failure(monkeypatch: pytest.MonkeyPatch) -> None:
    def broken(*_args, **_kwargs):
        raise RuntimeError("search crashed")

    monkeypatch.setattr("mrs3.portfolio.adapter.run_portfolio_adapter", broken)

    with pytest.raises(PortfolioPanelError) as error:
        PortfolioPanelService(Path.cwd())._package_variant_generator(({"strategy_id": 1},), {}, ({"profile_id": "BALANCED"},))

    assert error.value.code == "PORTFOLIO_JOB_FAILED"


def test_cancel_during_publication_rolls_back_every_artifact(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    finalists = [_finalist()]
    entered = threading.Event()
    release = threading.Event()

    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: finalists, variant_generator=_profiled_variants)

    def blocked_verify(workbook_path: Path) -> None:
        entered.set()
        assert release.wait(timeout=3)

    monkeypatch.setattr(service, "_verify_workbook", blocked_verify)
    result = service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})
    assert entered.wait(timeout=3)
    service.cancel(result["job_id"])
    release.set()
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and service.job(result["job_id"])["status"] in {"QUEUED", "RUNNING", "CANCEL_REQUESTED"}:
        time.sleep(0.01)

    assert service.job(result["job_id"])["status"] == "CANCELLED"
    assert not (tmp_path / ".portfolio-staging" / result["campaign_id"] / "stage1.xlsx").exists()
    assert not (tmp_path / ".portfolio-results" / result["campaign_id"] / "stage1.xlsx").exists()
    assert not (tmp_path / ".portfolio-results" / result["campaign_id"] / "stage1-executables.json").exists()
    with pytest.raises(PortfolioPanelError):
        service.workbook(result["campaign_id"])


def test_workbook_sanitizes_formula_text_and_rejects_formula_cells(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    _write_config(path)
    service = PortfolioPanelService(tmp_path, path)
    campaign = {"campaign_id": "campaign-fixed", "input_digest": "i", "config_digest": "c", "versions": {"policy_version": "p"}, "created_at_utc": "2000-01-01T00:00:00Z"}
    workbook_path = tmp_path / "literal.xlsx"
    service._write_workbook(workbook_path, campaign, (), (), (), (), optimizer_excluded=({"profile": "BALANCED", "selection_reason": "OPEN_POLICY", "message": '=HYPERLINK("http://evil")'},))
    workbook = load_workbook(workbook_path, data_only=False)
    try:
        assert workbook["Excluded"][2][9].data_type == "s"
        assert workbook["Excluded"][2][9].value.startswith("'=")
    finally:
        workbook.close()

    formula_path = tmp_path / "formula.xlsx"
    workbook = load_workbook(workbook_path)
    workbook["Summary"]["B2"] = "=1+1"
    workbook.save(formula_path)
    workbook.close()
    with pytest.raises(ValueError, match="formulas"):
        service._verify_workbook(formula_path)


def test_staging_file_root_is_rejected_before_builder_runs(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    staging = tmp_path / ".portfolio-staging"
    staging.write_text("not a directory", encoding="utf-8")
    finalists = [_finalist()]
    called = []

    def builder(*_args):
        called.append(True)

    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: finalists, variant_generator=_profiled_variants, workbook_builder=builder)
    result = service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and service.job(result["job_id"])["status"] in {"QUEUED", "RUNNING"}:
        time.sleep(0.01)

    assert service.job(result["job_id"])["status"] == "FAILED"
    assert called == []


def test_workbook_rejects_path_outside_results_root(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    finalists = [_finalist()]
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: finalists, variant_generator=_profiled_variants)
    result = service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and service.job(result["job_id"])["status"] in {"QUEUED", "RUNNING"}:
        time.sleep(0.01)
    runtime = service.registry.runtime(result["job_id"])
    runtime["workbook_path"] = str(tmp_path / "outside.xlsx")
    service.registry.sync(result["job_id"], {"state": "COMMITTED"}, runtime=runtime)

    with pytest.raises(Exception):
        service.workbook(result["campaign_id"])


def test_job_projection_progress_cancel_and_restart(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    registry.submit("portfolio.stage1", {}, "x", ("portfolio_optimizer",), job_id="j")
    registry.transition("j", "RUNNING")
    service = PortfolioPanelService(tmp_path, path, registry=registry)
    service._threads["j"] = object()
    job = service.job("j")
    assert job["status"] == "RUNNING"
    assert job["overall_percent"] == 0
    assert service.cancel("j")["status"] == "CANCEL_REQUESTED"
    restarted = PortfolioPanelService(tmp_path, path, registry=PanelJobRegistry(tmp_path / ".panel-jobs.json"))
    assert restarted.job("j")["status"] == "INTERRUPTED"


def test_threadless_orphan_with_runtime_failure_is_projected_and_releases_submit(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    registry.submit("portfolio.stage1", {}, "orphan", ("portfolio_optimizer",), job_id="orphan")
    service = PortfolioPanelService(tmp_path, path, registry=registry, finalists_reader=lambda *_: [_finalist()], variant_generator=_profiled_variants)
    original_runtime = registry.runtime
    monkeypatch.setattr(registry, "runtime", lambda *_: (_ for _ in ()).throw(OSError("runtime unavailable")))

    assert service.active_job() is None
    assert registry.get("orphan")["state"] == "FAILED"
    assert registry.get("orphan")["error"]["code"] == "INTERRUPTED"

    monkeypatch.setattr(registry, "runtime", original_runtime)
    retry = service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})
    assert retry["status"] == "QUEUED"


def test_threadless_orphan_without_campaign_is_projected(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    _write_config(path)
    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    registry.submit("portfolio.stage1", {}, "orphan-no-runtime", ("portfolio_optimizer",), job_id="orphan-no-runtime")
    service = PortfolioPanelService(tmp_path, path, registry=registry)

    assert service.active_job() is None
    assert registry.get("orphan-no-runtime")["state"] == "FAILED"
    assert registry.get("orphan-no-runtime")["error"]["code"] == "INTERRUPTED"


def test_active_or_job_restores_latest_terminal_job(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    service = PortfolioPanelService(tmp_path)
    jobs = [
        {"job_id": "latest", "kind": "portfolio.stage1", "state": "FAILED", "created_at_utc": "2026-09-23T14:00:00+00:00"},
        {"job_id": "unrelated", "kind": "strategies.tester", "state": "FAILED", "created_at_utc": "2026-09-23T15:00:00+00:00"},
        {"job_id": "old", "kind": "portfolio.stage1", "state": "COMMITTED", "created_at_utc": "2026-09-22T14:00:00+00:00"},
        {"job_id": "stale", "kind": "portfolio.stage1", "state": "RUNNING", "created_at_utc": "2026-09-24T14:00:00+00:00"},
        {"job_id": "undated", "kind": "portfolio.stage1", "state": "FAILED", "created_at_utc": None},
    ]
    monkeypatch.setattr(service, "active_job", lambda: None)
    monkeypatch.setattr(service.registry, "list", lambda: jobs)
    monkeypatch.setattr(service, "job", lambda job_id: {"job_id": job_id, "status": "FAILED"})

    assert service.active_or_job() == {"job_id": "latest", "status": "FAILED"}


def test_stage2_route_is_permanently_blocked(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    _write_config(path)
    service = PortfolioPanelService(tmp_path, path)
    with pytest.raises(Exception) as error:
        service.submit_tester_submission("campaign", {"confirmed": True, "campaign_id": "campaign"})
    assert getattr(error.value, "code", None) == "PORTFOLIO_JOB_STAGE2_NOT_AUTHORIZED"


@pytest.mark.parametrize("payload", ({}, {"confirmed": True, "campaign_id": "other"}))
def test_stage2_route_blocks_before_confirmation_validation(tmp_path: Path, payload: dict) -> None:
    service = PortfolioPanelService(tmp_path)
    with pytest.raises(PortfolioPanelError) as error:
        service.submit_tester_submission("campaign", payload)
    assert error.value.code == "PORTFOLIO_JOB_STAGE2_NOT_AUTHORIZED"


@pytest.mark.parametrize("failure", (PanelJobError("RESOURCE_BUSY"), OSError("registry unavailable"), TypeError("registry unavailable"), ValueError("registry unavailable")))
def test_cancel_maps_non_not_found_failures_to_runtime_unavailable(monkeypatch: pytest.MonkeyPatch, tmp_path: Path, failure: Exception) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    _write_config(path)
    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    registry.submit("portfolio.stage1", {}, "cancel-runtime", ("portfolio_optimizer",), job_id="cancel-runtime")
    service = PortfolioPanelService(tmp_path, path, registry=registry)
    monkeypatch.setattr(registry, "cancel", lambda *_: (_ for _ in ()).throw(failure))

    with pytest.raises(PortfolioPanelError) as error:
        service.cancel("cancel-runtime")

    assert error.value.code == "PORTFOLIO_JOB_RUNTIME_UNAVAILABLE"
    assert error.value.status == 500


def test_failed_finalization_does_not_wedge_future_submission(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    finalists = [_finalist()]
    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    original_transition = registry.transition
    original_sync = registry.sync

    def fail_terminal_transition(job_id, state, *, phase=None):
        if state == "FAILED":
            raise PanelJobError("FINALIZATION_FAILED")
        return original_transition(job_id, state, phase=phase)

    def fail_terminal_sync(job_id, status, *, runtime=None):
        if status.get("state") == "FAILED":
            raise PanelJobError("FINALIZATION_FAILED")
        return original_sync(job_id, status, runtime=runtime)

    registry.transition = fail_terminal_transition
    registry.sync = fail_terminal_sync
    service = PortfolioPanelService(
        tmp_path,
        path,
        registry=registry,
        finalists_reader=lambda *_: finalists,
        variant_generator=lambda *_: (_ for _ in ()).throw(RuntimeError("generation failed")),
    )
    payload = {"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest}
    first = service.submit_campaign({**payload})
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and first["job_id"] in service._threads:
        time.sleep(0.01)
    assert service.job(first["job_id"])["status"] == "INTERRUPTED"

    service.variant_generator = lambda selected, *_: selected
    second = service.submit_campaign({**payload})
    assert second["job_id"] != first["job_id"]


def test_running_transition_failure_releases_queued_orphan(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    finalists = [_finalist()]
    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    original_transition = registry.transition
    original_sync = registry.sync

    def fail_terminal_transition(job_id, state, *, phase=None):
        if state in {"RUNNING", "FAILED"}:
            raise PanelJobError("FINALIZATION_FAILED")
        return original_transition(job_id, state, phase=phase)

    def fail_terminal_sync(job_id, status, *, runtime=None):
        if status.get("state") == "FAILED":
            raise PanelJobError("FINALIZATION_FAILED")
        return original_sync(job_id, status, runtime=runtime)

    registry.transition = fail_terminal_transition
    registry.sync = fail_terminal_sync
    service = PortfolioPanelService(tmp_path, path, registry=registry, finalists_reader=lambda *_: finalists)
    payload = {"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest}
    first = service.submit_campaign({**payload})
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and first["job_id"] in service._threads:
        time.sleep(0.01)
    assert first["job_id"] not in service._threads
    assert service.job(first["job_id"])["status"] == "INTERRUPTED"
    second = service.submit_campaign({**payload})
    assert second["job_id"] != first["job_id"]


def test_stage1_success_publishes_exact_workbook_after_commit(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    finalists = [_finalist(strategy_name="fixture")]
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: finalists, variant_generator=_profiled_variants)
    result = service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and service.job(result["job_id"])["status"] in {"QUEUED", "RUNNING"}:
        time.sleep(0.01)
    job = service.job(result["job_id"])
    assert job["status"] == "SUCCEEDED", (job, service.registry.runtime(result["job_id"]))
    assert "workbook_path" in service.registry.runtime(result["job_id"]), service.registry.runtime(result["job_id"]).keys()
    assert service.registry.runtime(result["job_id"])["summary"]["prepared"] == 1
    workbook_bytes = service.workbook(result["campaign_id"])
    workbook_path = tmp_path / "stage1.xlsx"
    workbook_path.write_bytes(workbook_bytes)
    workbook = load_workbook(workbook_path, data_only=False)
    assert workbook.sheetnames == ["Итог", "Варианты", "Состав", "Финалисты", "Исключено", "Metadata"]
    assert workbook["Metadata"].sheet_state == "hidden"
    assert [cell.value for cell in workbook["Финалисты"][2]] == [result["campaign_id"], 7, 11, "BTCUSDT", "LONG", "FINALIST", 1, 1, "SELECTED", "WITHIN_MAXIMUM"]
    assert workbook["Варианты"].max_row == 2
    assert workbook["Варианты"][2][3].value == 2
    assert workbook["Варианты"][2][5].value is None
    assert workbook["Состав"].max_row == 3


def test_http_results_serializes_fractional_required_bank_as_json(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    payloads = (
        _executable_payload("BTCUSDT", "LONG"),
        _executable_payload("ETHUSDT", "SHORT"),
    )
    for payload in payloads:
        payload["facts"]["B"] = "10000.125"
    candidate = _weighted_executable_candidate(
        payloads,
        metrics={"required_bank_usdt": Decimal("10000.125")},
    )
    service = PortfolioPanelService(
        tmp_path,
        path,
        finalists_reader=lambda *_: [_finalist()],
        variant_generator=lambda *_: (candidate,),
    )
    result = service.submit_campaign({
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "20000", "max_candidates": 1}],
        "expected_config_digest": digest,
    })
    assert _wait_stage1(service, result)["status"] == "SUCCEEDED"

    config = tmp_path / "config.local.json"
    config.write_text("{}", encoding="utf-8")
    controller = PanelController(tmp_path, config)
    controller._portfolio_service = service
    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
    try:
        connection.request("GET", f"/api/v2/portfolio/campaigns/{result['campaign_id']}/results")
        response = connection.getresponse()
        payload = json.loads(response.read())
        assert response.status == 200
        assert payload["summary"]["B required USDT"] == "10000.125"
    finally:
        connection.close()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_stage1_executables_are_digest_bound_deterministic_and_restart_loadable(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    payloads = (
        _executable_payload("BTCUSDT", "LONG", name="PORTFOLIO_BTCUSDT_7_11"),
        _executable_payload("ETHUSDT", "SHORT", name="PORTFOLIO_ETHUSDT_8_12"),
    )
    reader_calls: list[str] = []
    generator_calls: list[str] = []

    def reader(*_args):
        reader_calls.append("read")
        return [_finalist()]

    def generator(*_args):
        generator_calls.append("generate")
        return (
            _weighted_executable_candidate(payloads, candidate_id="candidate-l1", identity="candidate-l1", limiter_L=1),
            _weighted_executable_candidate(payloads),
        )

    service = PortfolioPanelService(tmp_path, path, finalists_reader=reader, variant_generator=generator)
    result = service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})
    assert _wait_stage1(service, result)["status"] == "SUCCEEDED"

    artifact_path = tmp_path / ".portfolio-results" / result["campaign_id"] / "stage1-executables.json"
    artifact_bytes = artifact_path.read_bytes()
    artifact = json.loads(artifact_bytes)
    assert artifact["schema_version"] == 1
    assert artifact["campaign_id"] == result["campaign_id"]
    assert artifact["input_digest"] == result["input_digest"]
    assert artifact["config_digest"] == result["config_digest"]
    assert len(artifact["candidates"]) == 1
    assert artifact["candidates"][0]["candidate_id"] == "candidate-executable"
    assert artifact["candidates"][0]["order"] == 1
    assert artifact["candidates"][0]["members"] == [
        {"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 1, "result_id": 101},
        {"symbol": "ETHUSDT", "side": "SHORT", "strategy_id": 2, "result_id": 102},
    ]
    assert artifact["candidates"][0]["strategy_payloads"] == list(payloads)
    assert artifact["candidates"][0]["limiter_L"] == 0
    assert artifact["candidates"][0]["pretest_period"] == {
        "start_utc": "2026-01-01T00:00:00Z",
        "end_utc": "2026-01-15T00:00:00Z",
    }
    for payload in artifact["candidates"][0]["strategy_payloads"]:
        assert payload["account"]["open_positions_limiter"] == 0
        assert payload["strategy"]["mrs"]["position_priority"] == 1
    candidate = artifact["candidates"][0]
    candidate_body = {key: value for key, value in candidate.items() if key != "candidate_digest"}
    candidate_canonical = json.dumps(candidate_body, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)
    assert candidate["candidate_digest"] == hashlib.sha256(candidate_canonical.encode("utf-8")).hexdigest()
    body = {key: value for key, value in artifact.items() if key != "payload_digest"}
    canonical = json.dumps(body, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)
    assert artifact["payload_digest"] == hashlib.sha256(canonical.encode("utf-8")).hexdigest()
    runtime = service.registry.runtime(result["job_id"])
    assert runtime["executables_path"] == str(artifact_path)
    assert runtime["executables_digest"] == artifact["payload_digest"]

    def unexpected(*_args, **_kwargs):
        raise AssertionError("Stage 1 inputs must not be recomputed on artifact load")

    restarted = PortfolioPanelService(
        tmp_path,
        path,
        finalists_reader=unexpected,
        optimizer_input_preparer=unexpected,
        variant_generator=unexpected,
    )
    loaded = restarted._load_stage1_executables(result["campaign_id"], result["input_digest"], result["config_digest"], runtime["executables_digest"])
    assert loaded == artifact
    assert artifact_path.read_bytes() == artifact_bytes
    assert reader_calls == ["read"]
    assert generator_calls == ["generate"]


def test_stage1_executable_digest_is_stable_across_decimal_metric_types() -> None:
    campaign = {
        "campaign_id": "campaign-" + "a" * 32,
        "input_digest": "b" * 64,
        "config_digest": "c" * 64,
    }
    payloads = (
        _executable_payload("BTCUSDT", "LONG"),
        _executable_payload("ETHUSDT", "SHORT"),
    )
    for payload in payloads:
        payload["facts"]["B"] = "0.1"
    decimal_artifact = PortfolioPanelService._build_stage1_executables(
        campaign,
        (_weighted_executable_candidate(payloads, metrics={"required_bank_usdt": Decimal("0.1")}),),
    )
    string_artifact = PortfolioPanelService._build_stage1_executables(
        campaign,
        (_weighted_executable_candidate(payloads, metrics={"required_bank_usdt": "0.1"}),),
    )

    assert decimal_artifact == string_artifact
    assert decimal_artifact["candidates"][0]["candidate_digest"]
    assert decimal_artifact["payload_digest"]


def test_snapshot_bytes_are_canonical_deterministic_and_roundtrip() -> None:
    value = {"z": [1, True, None, 1.5], "a": {"text": "Привет"}}
    first = _snapshot_bytes(value)
    second = _snapshot_bytes({"a": {"text": "Привет"}, "z": [1, True, None, 1.5]})
    assert first == second
    assert json.loads(__import__("gzip").decompress(first[1]).decode("utf-8")) == value
    assert first[2] == hashlib.sha256(first[0]).hexdigest()


@pytest.mark.parametrize("value", ({1: "bad"}, {"bad": float("nan")}, {"bad": object()}))
def test_snapshot_bytes_reject_unsupported_values_with_typed_error(value) -> None:
    with pytest.raises(PortfolioPanelError) as error:
        _snapshot_bytes(value)
    assert error.value.code == PORTFOLIO_SNAPSHOT_UNSERIALIZABLE


def test_snapshot_path_rejects_actual_windows_junction(tmp_path: Path) -> None:
    target = tmp_path / "junction-target"
    target.mkdir()
    root = tmp_path / ".portfolio-results"
    try:
        created = subprocess.run(
            ["cmd", "/c", "mklink", "/J", str(root), str(target)],
            capture_output=True,
            text=True,
        )
    except OSError:
        pytest.skip("Windows junction creation unavailable")
    if created.returncode != 0:
        pytest.skip("Windows junction creation unavailable")
    try:
        with pytest.raises(PortfolioPanelError) as error:
            PortfolioPanelService(tmp_path)._snapshot_path("campaign-" + "a" * 32)
        assert error.value.code == "PORTFOLIO_SNAPSHOT_UNAVAILABLE"
    finally:
        os.rmdir(root)


def test_progress_reporter_keeps_units_in_memory_and_heartbeats_bounded() -> None:
    now = [0.0]
    reporter = _PortfolioProgressReporter(lambda: now[0])
    reporter.start("job")
    first, first_persist = reporter.emit("job", {"substage": "PROFILE", "unit": "profile", "completed": 0, "total": 3, "detail": "start"})
    assert first_persist is True
    now[0] = 3.0
    steady, steady_persist = reporter.emit("job", {"substage": "PROFILE", "unit": "profile", "completed": 2, "total": 3, "detail": "unit"})
    assert steady_persist is False
    assert steady["eta_seconds"] is not None
    now[0] = 9.0
    assert reporter.heartbeat("job")[1] is False
    now[0] = 10.0
    assert reporter.heartbeat("job")[1] is True
    now[0] = 19.0
    assert reporter.heartbeat("job")[1] is False
    now[0] = 20.0
    assert reporter.heartbeat("job")[1] is True
    now[0] = 21.0
    _, transition_persist = reporter.emit("job", {"substage": "BOOTSTRAP", "unit": "batch", "completed": 0, "total": 2, "detail": "transition"})
    assert transition_persist is False
    now[0] = 22.0
    _, transition_persist = reporter.emit("job", {"substage": "SOLVER", "unit": "call", "completed": 0, "total": 1, "detail": "eligible transition"})
    assert transition_persist is True


def test_progress_reporter_eta_is_unknown_until_reliable_and_after_inconsistency() -> None:
    now = [0.0]
    reporter = _PortfolioProgressReporter(lambda: now[0])
    reporter.start("job")
    now[0] = 1.0
    too_soon, _ = reporter.emit("job", {"substage": "EARLY", "completed": 2, "total": 3})
    assert too_soon["eta_seconds"] is None
    now[0] = 1.3
    missing, _ = reporter.emit("job", {"substage": "PROFILE", "completed": 0, "total": None})
    assert missing["eta_seconds"] is None
    now[0] = 3.0
    early, _ = reporter.emit("job", {"substage": "PROFILE", "completed": 1, "total": 3})
    assert early["eta_seconds"] is None
    now[0] = 3.3
    reliable, _ = reporter.emit("job", {"substage": "PROFILE", "completed": 2, "total": 3})
    assert reliable["eta_seconds"] is not None and reliable["eta_seconds"] >= 0
    now[0] = 3.6
    inconsistent, _ = reporter.emit("job", {"substage": "PROFILE", "completed": 2, "total": 4})
    assert inconsistent["total"] is None
    assert inconsistent["eta_seconds"] is None
    assert inconsistent["indeterminate"] is True


def test_progress_persistence_is_ignored_after_terminal_or_generation_invalidation(tmp_path: Path) -> None:
    class FakeRegistry:
        def __init__(self) -> None:
            self.lock = threading.RLock()
            self.jobs = {"job": {"job_id": "job", "state": "COMMITTED", "runtime": {}}}
            self.writes = 0

        def sync(self, *_args, **_kwargs):
            self.writes += 1

    registry = FakeRegistry()
    service = PortfolioPanelService(tmp_path, registry=registry)
    generation = service._activate_progress("job")
    service._persist_progress("job", generation, {"substage": "PROFILE", "completed": 1})
    assert registry.writes == 0
    registry.jobs["job"]["state"] = "RUNNING"
    generation = service._activate_progress("job")
    service._deactivate_progress("job")
    service._persist_progress("job", generation, {"substage": "PROFILE", "completed": 2})
    assert registry.writes == 0


def test_panel_controller_defers_registry_recovery_until_portfolio_startup(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    calls: list[object] = []
    original_registry = panel_module.PanelJobRegistry

    def registry_factory(path, **kwargs):
        calls.append(("registry", kwargs.get("recover_on_load")))
        return original_registry(path, **kwargs)

    class FakePortfolioService:
        def __init__(self, *_args, **kwargs):
            assert kwargs["registry"] is controller_registry[0]

        def startup_recover(self):
            calls.append("portfolio_recover")

    controller_registry: list[object] = []

    def registry_factory_with_capture(path, **kwargs):
        registry = registry_factory(path, **kwargs)
        controller_registry.append(registry)
        return registry

    monkeypatch.setattr(panel_module, "PanelJobRegistry", registry_factory_with_capture)
    monkeypatch.setattr(panel_module, "PortfolioPanelService", FakePortfolioService)
    monkeypatch.setattr(PanelController, "_reconcile_interrupted_remote_source_jobs", lambda self: None)
    monkeypatch.setattr(PanelController, "_reconcile_interrupted_tester_jobs", lambda self: None)
    PanelController(tmp_path, tmp_path / "panel-config.json")
    assert calls == [("registry", False), "portfolio_recover"]


def test_startup_migrates_scaled_legacy_campaigns_before_recovery(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    seed_root = tmp_path / "seed"
    seed_root.mkdir()
    seed_path = seed_root / "portfolio_optimizer.local.json"
    seed_digest = _write_config(seed_path)

    class IdleThread:
        def __init__(self, *args, **kwargs):
            pass

        def start(self):
            pass

    monkeypatch.setattr("mrs3.panel_portfolio.threading.Thread", IdleThread)
    seed_service = PortfolioPanelService(seed_root, seed_path, finalists_reader=lambda *_: [_finalist(
        timeframe="1h", close_ma_len=21, order_count=1, strategy_orders=[{"open_ma_len": 10}],
    )])
    seed_result = seed_service.submit_campaign({
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
        "expected_config_digest": seed_digest,
    })
    base_campaign = seed_service.registry.runtime(seed_result["job_id"])["campaign"]

    path = tmp_path / ".panel-jobs.json"
    registry = PanelJobRegistry(path, capacity=20, recover_on_load=False)

    def campaign(index: int) -> dict[str, Any]:
        value = copy.deepcopy(base_campaign)
        value["campaign_id"] = f"campaign-{index:032x}"
        return value

    jobs: list[tuple[str, dict[str, Any]]] = []
    for index in range(14):
        value = campaign(index)
        saved = registry.submit(
            "portfolio.stage1",
            {"campaign_id": value["campaign_id"], "input_digest": value["input_digest"], "config_digest": value["config_digest"]},
            f"legacy-{index}", (f"portfolio-{index}",),
            job_id=f"legacy-{index}",
        )
        registry.reserve_runtime(saved["job_id"], "campaign", value)
        jobs.append((saved["job_id"], value))
        if index < 4:
            registry.transition(saved["job_id"], "RUNNING")
            registry.transition(saved["job_id"], "COMMITTED")
        elif index < 8:
            registry.transition(saved["job_id"], "RUNNING")
            registry.transition(saved["job_id"], "FAILED")
        elif index < 10:
            registry.cancel(saved["job_id"])

    service = PortfolioPanelService(tmp_path, registry=registry)
    service.startup_recover()

    persisted = json.loads(path.read_text(encoding="utf-8"))
    assert len(persisted) == 14
    assert not (path.with_name(".panel-jobs.snapshot-migration.bak")).exists()
    assert path.stat().st_size < 256 * 1024
    for job_id, _value in jobs:
        record = persisted[job_id]
        runtime = record.get("runtime", {})
        assert "weighted_input_rows" not in record and "weighted_input_rows" not in runtime
        assert "finalists" not in record and "finalists" not in runtime
        if record["state"] == "COMMITTED":
            assert runtime["campaign_snapshot"]["state"] == "available"
            descriptor = runtime["campaign_snapshot"]
            assert service._hydrate_campaign(record, runtime, allow_terminal=True, expected_campaign_id=descriptor["campaign_id"])["campaign_id"] == descriptor["campaign_id"]
        else:
            assert "campaign" not in runtime or set(runtime["campaign"]) <= {"campaign_id", "input_digest", "config_digest"}

    restarted = PanelJobRegistry(path, capacity=20, recover_on_load=False)
    PortfolioPanelService(tmp_path, registry=restarted).startup_recover()
    assert json.loads(path.read_text(encoding="utf-8")) == persisted


def test_startup_migration_fault_restores_journal_and_retries_one_fixed_snapshot(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    seed_root = tmp_path / "seed"
    seed_root.mkdir()
    seed_path = seed_root / "portfolio_optimizer.local.json"
    seed_digest = _write_config(seed_path)

    class IdleThread:
        def __init__(self, *args, **kwargs):
            pass

        def start(self):
            pass

    monkeypatch.setattr("mrs3.panel_portfolio.threading.Thread", IdleThread)
    seed_service = PortfolioPanelService(seed_root, seed_path, finalists_reader=lambda *_: [_finalist(
        timeframe="1h", close_ma_len=21, order_count=1, strategy_orders=[{"open_ma_len": 10}],
    )])
    seed_result = seed_service.submit_campaign({
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
        "expected_config_digest": seed_digest,
    })
    campaign = seed_service.registry.runtime(seed_result["job_id"])["campaign"]
    path = tmp_path / ".panel-jobs.json"
    registry = PanelJobRegistry(path, recover_on_load=False)
    saved = registry.submit("portfolio.stage1", {}, "legacy", job_id="legacy")
    registry.reserve_runtime(saved["job_id"], "campaign", campaign)
    registry.transition(saved["job_id"], "RUNNING")
    registry.transition(saved["job_id"], "COMMITTED")
    original_journal = path.read_bytes()
    original_jobs = copy.deepcopy(registry.jobs)
    service = PortfolioPanelService(tmp_path, registry=registry)
    verify = service._verify_migrated_registry
    calls = 0

    def fail_once() -> None:
        nonlocal calls
        calls += 1
        if calls == 1:
            raise OSError("injected verification fault")
        verify()

    monkeypatch.setattr(service, "_verify_migrated_registry", fail_once)
    service.startup_recover()

    snapshot = tmp_path / ".portfolio-results" / campaign["campaign_id"] / "campaign-input.json.gz"
    backup = path.with_name(".panel-jobs.snapshot-migration.bak")
    assert calls == 1
    assert path.read_bytes() == original_journal
    assert registry.jobs == original_jobs
    assert backup.exists()
    assert snapshot.is_file()
    assert len(list((tmp_path / ".portfolio-results").rglob("campaign-input.json.gz"))) == 1

    restarted = PanelJobRegistry(path, capacity=4, recover_on_load=False)
    PortfolioPanelService(tmp_path, registry=restarted).startup_recover()
    persisted = json.loads(path.read_text(encoding="utf-8"))
    descriptor = persisted["legacy"]["runtime"]["campaign_snapshot"]
    assert descriptor["path"] == snapshot.relative_to(tmp_path).as_posix()
    assert descriptor["state"] == "available"
    assert not backup.exists()
    assert len(list((tmp_path / ".portfolio-results").rglob("campaign-input.json.gz"))) == 1


def test_startup_disk_preflight_is_atomic_one_byte_below_required_space(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / ".panel-jobs.json"
    registry = PanelJobRegistry(path, recover_on_load=False)
    campaign = {"campaign_id": "campaign-" + "a" * 32, "input_digest": "i", "config_digest": "c"}
    saved = registry.submit("portfolio.stage1", {}, "legacy", job_id="legacy")
    registry.reserve_runtime(saved["job_id"], "campaign", campaign)
    original_journal = path.read_bytes()
    original_jobs = copy.deepcopy(registry.jobs)
    original_files = {item.relative_to(tmp_path): item.read_bytes() for item in tmp_path.rglob("*") if item.is_file()}
    _raw, compressed, _digest = _snapshot_bytes(campaign)
    required = 2 * len(original_journal) + len(compressed) + 64 * 1024 * 1024
    monkeypatch.setattr("mrs3.panel_portfolio.shutil.disk_usage", lambda _path: type("Usage", (), {"free": required - 1})())

    PortfolioPanelService(tmp_path, registry=registry).startup_recover()

    assert path.read_bytes() == original_journal
    assert registry.jobs == original_jobs
    assert {item.relative_to(tmp_path): item.read_bytes() for item in tmp_path.rglob("*") if item.is_file()} == original_files
    assert not path.with_name(".panel-jobs.snapshot-migration.bak").exists()


def test_unrepairable_active_campaign_preflights_before_recovery(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / ".panel-jobs.json"
    registry = PanelJobRegistry(path, recover_on_load=False)
    campaign = {
        "campaign_id": "campaign-" + "a" * 32,
        "input_digest": "i",
        "config_digest": "c",
        "finalists": [],
        "weighted_input_rows": [{"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 7, "result_id": 11}],
    }
    saved = registry.submit("portfolio.stage1", {}, "legacy", job_id="legacy")
    registry.reserve_runtime(saved["job_id"], "campaign", campaign)
    registry.transition(saved["job_id"], "RUNNING")
    original_journal = path.read_bytes()
    original_jobs = copy.deepcopy(registry.jobs)
    original_files = {item.relative_to(tmp_path): item.read_bytes() for item in tmp_path.rglob("*") if item.is_file()}
    required = 2 * len(original_journal) + 64 * 1024 * 1024
    monkeypatch.setattr("mrs3.panel_portfolio.shutil.disk_usage", lambda _path: type("Usage", (), {"free": required - 1})())

    PortfolioPanelService(tmp_path, registry=registry).startup_recover()

    assert path.read_bytes() == original_journal
    assert registry.jobs == original_jobs
    assert {item.relative_to(tmp_path): item.read_bytes() for item in tmp_path.rglob("*") if item.is_file()} == original_files
    assert not path.with_name(".panel-jobs.snapshot-migration.bak").exists()


def test_startup_recompacts_unrepairable_active_job_after_recovery_without_snapshot(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    seed_root = tmp_path / "seed"
    seed_root.mkdir()
    seed_path = seed_root / "portfolio_optimizer.local.json"
    seed_digest = _write_config(seed_path)

    class IdleThread:
        def __init__(self, *args, **kwargs):
            pass

        def start(self):
            pass

    monkeypatch.setattr("mrs3.panel_portfolio.threading.Thread", IdleThread)
    seed_service = PortfolioPanelService(seed_root, seed_path, finalists_reader=lambda *_: [_finalist(
        timeframe="1h", close_ma_len=21, order_count=1, strategy_orders=[{"open_ma_len": 10}],
    )])
    seed_result = seed_service.submit_campaign({
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
        "expected_config_digest": seed_digest,
    })
    base_campaign = seed_service.registry.runtime(seed_result["job_id"])["campaign"]
    path = tmp_path / ".panel-jobs.json"
    registry = PanelJobRegistry(path, recover_on_load=False)
    good = copy.deepcopy(base_campaign)
    good["campaign_id"] = "campaign-" + "a" * 32
    bad = {
        "campaign_id": "campaign-" + "b" * 32,
        "input_digest": good["input_digest"],
        "config_digest": good["config_digest"],
        "finalists": [],
        "weighted_input_rows": [{"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 7, "result_id": 11}],
    }
    good_job = registry.submit("portfolio.stage1", {}, "good", job_id="good")
    registry.reserve_runtime(good_job["job_id"], "campaign", good)
    registry.transition(good_job["job_id"], "RUNNING")
    registry.transition(good_job["job_id"], "COMMITTED")
    bad_job = registry.submit("portfolio.stage1", {}, "bad", job_id="bad")
    registry.reserve_runtime(bad_job["job_id"], "campaign", bad)
    registry.transition(bad_job["job_id"], "RUNNING")

    service = PortfolioPanelService(tmp_path, registry=registry)
    monkeypatch.setattr("mrs3.panel_portfolio.shutil.disk_usage", lambda _path: type("Usage", (), {"free": 2**40})())
    service.startup_recover()
    persisted = json.loads(path.read_text(encoding="utf-8"))

    assert persisted["good"]["state"] == "COMMITTED"
    assert persisted["good"]["runtime"]["campaign_snapshot"]["state"] == "available"
    assert persisted["bad"]["state"] == "FAILED"
    assert persisted["bad"]["runtime"]["campaign"] == {"campaign_id": bad["campaign_id"], "input_digest": bad["input_digest"], "config_digest": bad["config_digest"]}
    assert "campaign_snapshot" not in persisted["bad"]["runtime"]
    assert "finalists" not in persisted["bad"]["runtime"]
    assert "weighted_input_rows" not in persisted["bad"]["runtime"]

    restarted = PanelJobRegistry(path, capacity=4, recover_on_load=False)
    PortfolioPanelService(tmp_path, registry=restarted).startup_recover()
    assert json.loads(path.read_text(encoding="utf-8")) == persisted


def test_terminal_snapshot_cleanup_removes_only_empty_campaign_directory(tmp_path: Path) -> None:
    path = tmp_path / ".panel-jobs.json"
    registry = PanelJobRegistry(path)
    saved = registry.submit("portfolio.stage1", {}, "terminal", job_id="terminal")
    registry.transition(saved["job_id"], "RUNNING")
    registry.transition(saved["job_id"], "FAILED")
    campaign_id = "campaign-" + "c" * 32
    snapshot = tmp_path / ".portfolio-results" / campaign_id / "campaign-input.json.gz"
    snapshot.parent.mkdir(parents=True)
    snapshot.write_bytes(b"snapshot")
    registry.reserve_runtime(saved["job_id"], "campaign_snapshot", {
        "schema": "portfolio-campaign-input-v1", "campaign_id": campaign_id,
        "path": snapshot.relative_to(tmp_path).as_posix(), "state": "available",
    })

    service = PortfolioPanelService(tmp_path, registry=registry)
    service._cleanup_terminal_snapshot(saved["job_id"])

    assert not snapshot.exists()
    assert not snapshot.parent.exists()


def test_terminal_snapshot_cleanup_preserves_nonempty_campaign_directory(tmp_path: Path) -> None:
    path = tmp_path / ".panel-jobs.json"
    registry = PanelJobRegistry(path)
    saved = registry.submit("portfolio.stage1", {}, "terminal", job_id="terminal")
    registry.transition(saved["job_id"], "RUNNING")
    registry.transition(saved["job_id"], "FAILED")
    campaign_id = "campaign-" + "d" * 32
    snapshot = tmp_path / ".portfolio-results" / campaign_id / "campaign-input.json.gz"
    snapshot.parent.mkdir(parents=True)
    snapshot.write_bytes(b"snapshot")
    (snapshot.parent / "stage1.xlsx").write_bytes(b"workbook")
    registry.reserve_runtime(saved["job_id"], "campaign_snapshot", {
        "schema": "portfolio-campaign-input-v1", "campaign_id": campaign_id,
        "path": snapshot.relative_to(tmp_path).as_posix(), "state": "available",
    })

    PortfolioPanelService(tmp_path, registry=registry)._cleanup_terminal_snapshot(saved["job_id"])

    assert not snapshot.exists()
    assert snapshot.parent.exists()


def test_legacy_geometry_repair_joins_exact_identity_without_mutating_source() -> None:
    finalist = {
        "symbol": "BTCUSDT", "side": "LONG", "strategy_id": 7, "result_id": 11,
        "timeframe": "3h", "close_ma_len": 55, "order_count": 1,
        "strategy_orders": ({"open_ma_len": 34, "shift_bp": 50, "lot_x": "1"},),
    }
    campaign = {"finalists": [finalist], "weighted_input_rows": [{"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 7, "result_id": 11}]}
    repaired = _repair_legacy_geometry(campaign, strict=True)
    assert repaired["weighted_input_rows"][0]["strategy_orders"] == finalist["strategy_orders"]
    assert "strategy_orders" not in campaign["weighted_input_rows"][0]


@pytest.mark.parametrize(
    "finalists",
    (
        (),
        ({"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 7, "result_id": 11, "timeframe": "3h"},
         {"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 7, "result_id": 11, "timeframe": "3h"}),
    ),
)
def test_legacy_geometry_repair_rejects_zero_or_multiple_finalists(finalists) -> None:
    campaign = {
        "finalists": finalists,
        "weighted_input_rows": [{"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 7, "result_id": 11}],
    }
    with pytest.raises(PortfolioPanelError) as error:
        _repair_legacy_geometry(campaign, strict=True)
    assert error.value.code == "PORTFOLIO_INPUT_GEOMETRY_INVALID"


def test_stage2_baseline_preparation_uses_first_persisted_candidate_and_exact_payloads(tmp_path: Path) -> None:
    payloads = (
        _stage2_payload("BTCUSDT", "LONG", name="PORTFOLIO_BTCUSDT_7_11"),
        _stage2_payload("ETHUSDT", "SHORT", name="PORTFOLIO_ETHUSDT_8_12"),
    )
    candidate = _weighted_executable_candidate(
        payloads,
        candidate_id="a" * 64,
        identity="a" * 64,
    )
    service, result = _stage2_service(tmp_path, candidate)

    prepared = service._prepare_stage2_baseline(result["campaign_id"])
    assert prepared["campaign_id"] == result["campaign_id"]
    assert prepared["input_digest"] == result["input_digest"]
    assert prepared["config_digest"] == result["config_digest"]
    assert prepared["artifact_digest"] == service.registry.runtime(result["job_id"])["executables_digest"]
    assert "executables_digest" not in prepared
    assert prepared["portfolio_name"] == "a" * 64
    assert prepared["candidate_id"] == "a" * 64
    assert prepared["expected_names"] == [payload["strategy"]["name"] for payload in payloads]
    assert prepared["pretest_period"] == {"start_utc": "2026-01-01T00:00:00Z", "end_utc": "2026-01-15T00:00:00Z"}

    config = json.loads(prepared["tester_config_json"])
    template = json.loads((Path(__file__).parents[1] / "templates/tester/mrs3/config_tester.json").read_text(encoding="utf-8"))
    expected_config = {**template, "name_comment": "a" * 64, "StartDate": "2026-01-01", "EndDate": "2026-01-14", "InitialBalance": 10000, "single_mode": False, "UpdateData": False}
    assert config == expected_config
    assert type(config["InitialBalance"]) is int
    assert config["use_runs"] is False
    assert config["parameter_mining"] == []
    assert prepared["tester_config_json"].endswith("\n")

    for payload in payloads:
        name = payload["strategy"]["name"]
        expected = json.dumps(payload["strategy"], ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False) + "\n"
        assert prepared["strategy_jsons"][name] == expected
    assert prepared["receipt"]["candidate_order"] == 0
    assert prepared["receipt"]["profile"] == "BALANCED"
    assert prepared["receipt"]["member_identities"] == [
        {"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 1, "result_id": 101},
        {"symbol": "ETHUSDT", "side": "SHORT", "strategy_id": 2, "result_id": 102},
    ]
    assert prepared["receipt"]["tester_config_sha256"] == hashlib.sha256(
        prepared["tester_config_json"].encode("utf-8")
    ).hexdigest()
    assert [item["filename"] for item in prepared["receipt"]["strategy_manifest"]] == [
        "PORTFOLIO_BTCUSDT_7_11.json", "PORTFOLIO_ETHUSDT_8_12.json",
    ]


def test_stage2_baseline_preparation_uses_candidate_after_filtered_order_zero(tmp_path: Path) -> None:
    payloads = (
        _stage2_payload("BTCUSDT", "LONG", name="PORTFOLIO_BTCUSDT_7_11"),
        _stage2_payload("ETHUSDT", "SHORT", name="PORTFOLIO_ETHUSDT_8_12"),
    )
    candidate = _weighted_executable_candidate(payloads, candidate_id="b" * 64, identity="b" * 64)
    invalid = _weighted_executable_candidate(payloads, candidate_id="c" * 64, identity="c" * 64, limiter_L=1)
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: [_finalist()], variant_generator=lambda *_: (invalid, candidate))
    result = service.submit_campaign({
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}],
        "expected_config_digest": digest,
    })
    assert _wait_stage1(service, result)["status"] == "SUCCEEDED"
    artifact = json.loads((tmp_path / ".portfolio-results" / result["campaign_id"] / "stage1-executables.json").read_text(encoding="utf-8"))
    assert artifact["candidates"][0]["order"] == 1
    assert service._prepare_stage2_baseline(result["campaign_id"])["candidate_id"] == "b" * 64


@pytest.mark.parametrize(
    "candidate_update",
    (
        {"candidate_id": "unsafe"},
        {"identity": "unsafe"},
    ),
)
def test_stage2_baseline_preparation_rejects_unsafe_candidate_binding(tmp_path: Path, candidate_update: dict) -> None:
    payloads = (
        _stage2_payload("BTCUSDT", "LONG", name="PORTFOLIO_BTCUSDT_7_11"),
        _stage2_payload("ETHUSDT", "SHORT", name="PORTFOLIO_ETHUSDT_8_12"),
    )
    candidate = _weighted_executable_candidate(payloads, candidate_id="d" * 64, identity="d" * 64)
    candidate.update(candidate_update)
    service, result = _stage2_service(tmp_path, candidate)
    with pytest.raises(PortfolioPanelError) as error:
        service._prepare_stage2_baseline(result["campaign_id"])
    assert error.value.code == "PORTFOLIO_STAGE2_INPUT_INVALID"
    assert error.value.status == 409


def test_stage2_baseline_preparation_rejects_bank_mismatch(tmp_path: Path) -> None:
    payloads = (
        _stage2_payload("BTCUSDT", "LONG", name="PORTFOLIO_BTCUSDT_7_11", bank="9999"),
        _stage2_payload("ETHUSDT", "SHORT", name="PORTFOLIO_ETHUSDT_8_12"),
    )
    candidate = _weighted_executable_candidate(payloads, candidate_id="e" * 64, identity="e" * 64)
    with pytest.raises(ValueError, match="sizing evidence"):
        _stage2_material(candidate)


def test_stage1_rejects_candidate_bank_mismatch() -> None:
    payloads = (
        _stage2_payload("BTCUSDT", "LONG", name="PORTFOLIO_BTCUSDT_7_11", bank="9999"),
        _stage2_payload("ETHUSDT", "SHORT", name="PORTFOLIO_ETHUSDT_8_12"),
    )
    candidate = _weighted_executable_candidate(payloads, candidate_id="e" * 64, identity="e" * 64)
    campaign = {"campaign_id": "campaign-" + "a" * 32, "input_digest": "b" * 64, "config_digest": "c" * 64}

    with pytest.raises(PortfolioPanelError) as error:
        PortfolioPanelService._build_stage1_executables(campaign, (candidate,))

    assert error.value.code == "PORTFOLIO_STAGE1_EXECUTABLES_UNAVAILABLE"


def test_stage2_maps_tampered_bank_mismatch_to_typed_error(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    payloads = (
        _stage2_payload("BTCUSDT", "LONG", name="PORTFOLIO_BTCUSDT_7_11"),
        _stage2_payload("ETHUSDT", "SHORT", name="PORTFOLIO_ETHUSDT_8_12"),
    )
    service, result = _stage2_service(
        tmp_path,
        _weighted_executable_candidate(payloads, candidate_id="e" * 64, identity="e" * 64),
    )
    runtime = service.registry.runtime(result["job_id"])
    artifact = service._load_stage1_executables(
        result["campaign_id"], runtime["campaign"]["input_digest"], runtime["campaign"]["config_digest"], runtime["executables_digest"]
    )
    artifact["candidates"][0]["metrics"]["required_bank_usdt"] = "9999"
    monkeypatch.setattr(service, "_load_stage1_executables", lambda *_args, **_kwargs: artifact)

    with pytest.raises(PortfolioPanelError) as error:
        service._prepare_stage2_baseline(result["campaign_id"])

    assert (error.value.code, error.value.status) == ("PORTFOLIO_STAGE2_INPUT_INVALID", 409)


@pytest.mark.parametrize("bank", ("1800", "0.1", "1234.567890123456", "0.000000000001"))
def test_stage2_initial_balance_uses_exact_required_bank(bank: str) -> None:
    payloads = (
        _stage2_payload("BTCUSDT", "LONG", name="PORTFOLIO_BTCUSDT_7_11", bank=bank),
        _stage2_payload("ETHUSDT", "SHORT", name="PORTFOLIO_ETHUSDT_8_12", bank=bank),
    )
    candidate = _weighted_executable_candidate(
        payloads,
        candidate_id="f" * 64,
        identity="f" * 64,
        metrics={"required_bank_usdt": Decimal(bank)},
    )

    material = _stage2_material(candidate)

    assert Decimal(str(json.loads(material["tester_config_json"])["InitialBalance"])) == Decimal(bank)
    assert material["required_bank_usdt"] == bank


@pytest.mark.parametrize("missing", ("C", "max_balance"))
def test_stage2_baseline_preparation_requires_complete_receipt_evidence(
    tmp_path: Path, missing: str
) -> None:
    payloads = [
        _stage2_payload("BTCUSDT", "LONG", name="PORTFOLIO_BTCUSDT_7_11"),
        _stage2_payload("ETHUSDT", "SHORT", name="PORTFOLIO_ETHUSDT_8_12"),
    ]
    if missing == "max_balance":
        del payloads[0]["strategy"]["basic"][missing]
    else:
        del payloads[0]["facts"][missing]
    service, result = _stage2_service(
        tmp_path,
        _weighted_executable_candidate(tuple(payloads), candidate_id="2" * 64, identity="2" * 64),
    )

    with pytest.raises(PortfolioPanelError) as error:
        service._prepare_stage2_baseline(result["campaign_id"])

    assert error.value.code == "PORTFOLIO_STAGE2_INPUT_INVALID"


def test_stage2_receipt_rechecks_artifact_sizing_and_max_balance(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    payloads = (
        _stage2_payload("BTCUSDT", "LONG", name="PORTFOLIO_BTCUSDT_7_11"),
        _stage2_payload("ETHUSDT", "SHORT", name="PORTFOLIO_ETHUSDT_8_12"),
    )
    service, result = _stage2_service(
        tmp_path,
        _weighted_executable_candidate(payloads, candidate_id="3" * 64, identity="3" * 64),
    )
    prepared = service._prepare_stage2_baseline(result["campaign_id"])
    changed = service._load_stage1_executables(
        prepared["campaign_id"], prepared["input_digest"], prepared["config_digest"], prepared["artifact_digest"]
    )
    changed["candidates"][0]["strategy_payloads"][0]["strategy"]["basic"]["max_balance"] = "200"
    monkeypatch.setattr(service, "_load_stage1_executables", lambda *_args, **_kwargs: changed)

    with pytest.raises(PortfolioPanelError) as error:
        service._verify_stage2_prepared(prepared)

    assert error.value.code == "PORTFOLIO_JOB_STAGE2_INPUT_CHANGED"


def test_stage2_rejects_receiptless_prepared_package(tmp_path: Path) -> None:
    payloads = (
        _stage2_payload("BTCUSDT", "LONG", name="PORTFOLIO_BTCUSDT_7_11"),
        _stage2_payload("ETHUSDT", "SHORT", name="PORTFOLIO_ETHUSDT_8_12"),
    )
    service, result = _stage2_service(
        tmp_path,
        _weighted_executable_candidate(payloads, candidate_id="4" * 64, identity="4" * 64),
    )
    prepared = service._prepare_stage2_baseline(result["campaign_id"])
    prepared.pop("receipt")

    with pytest.raises(PortfolioPanelError) as error:
        service._verify_stage2_prepared(prepared)

    assert error.value.code == "PORTFOLIO_JOB_STAGE2_INPUT_CHANGED"


def test_stage2_baseline_preparation_rejects_unsafe_strategy_name(tmp_path: Path) -> None:
    payloads = (
        _stage2_payload("BTCUSDT", "LONG", name="BAD:NAME"),
        _stage2_payload("ETHUSDT", "SHORT", name="PORTFOLIO_ETHUSDT_8_12"),
    )
    service, result = _stage2_service(tmp_path, _weighted_executable_candidate(payloads, candidate_id="1" * 64, identity="1" * 64))
    with pytest.raises(PortfolioPanelError) as error:
        service._prepare_stage2_baseline(result["campaign_id"])
    assert error.value.code == "PORTFOLIO_STAGE2_INPUT_INVALID"
    assert error.value.status == 409


def test_stage2_baseline_preparation_loads_restart_without_stage1_recomputation(tmp_path: Path) -> None:
    payloads = (
        _stage2_payload("BTCUSDT", "LONG", name="PORTFOLIO_BTCUSDT_7_11"),
        _stage2_payload("ETHUSDT", "SHORT", name="PORTFOLIO_ETHUSDT_8_12"),
    )
    service, result = _stage2_service(tmp_path, _weighted_executable_candidate(payloads, candidate_id="f" * 64, identity="f" * 64))

    def unexpected(*_args, **_kwargs):
        raise AssertionError("stage 1 inputs must not be recomputed")

    restarted = PortfolioPanelService(tmp_path, tmp_path / "portfolio_optimizer.local.json", finalists_reader=unexpected, optimizer_input_preparer=unexpected, variant_generator=unexpected)
    prepared = restarted._prepare_stage2_baseline(result["campaign_id"])
    assert prepared["candidate_id"] == "f" * 64


class _FakeStage2Tester:
    def __init__(self, root: Path, expected_names: tuple[str, ...], *, write_result: bool = True, write_report: bool = True, delayed_report: bool = False, extra_report: bool = False, report_as_directory: bool = False, invalid_fill_readback: bool = False, after_fill: Any = None, after_start: Any = None, result_names: tuple[str, ...] | None = None, result_stats: dict[str, object] | None = None, fill_error: Exception | None = None, fill_enter: threading.Event | None = None, fill_release: threading.Event | None = None, stop_error: Exception | None = None, stop_enter: threading.Event | None = None, stop_release: threading.Event | None = None) -> None:
        self.config = SimpleNamespace(
            wizard_result=root / "wizard-result.json",
            report_dir=root / "tester" / "report" / "my_test",
            poll_interval_seconds=0.001,
            stall_timeout_seconds=0.05,
            report_stability_polls=1,
            metric_tolerance=Decimal("0.01"),
        )
        self.expected_names = expected_names
        self.result_names = result_names or expected_names
        self.result_stats = result_stats or {
            "InitialBalance": 10000,
            "FinalBalance": 10100,
            "TotalPnL": 100,
            "TotalPnLPercent": 1,
            "TotalTrades": 2,
            "WinRate": 50,
            "MaxDrawdown": 10,
            "MaxDrawdownPercent": 0.1,
            "TotalFees": 1,
        }
        self.write_result = write_result
        self.write_report = write_report
        self.delayed_report = delayed_report
        self.extra_report = extra_report
        self.report_as_directory = report_as_directory
        self.invalid_fill_readback = invalid_fill_readback
        self.after_fill = after_fill
        self.after_start = after_start
        self.fill_error = fill_error
        self.fill_enter = fill_enter
        self.fill_release = fill_release
        self.stop_error = stop_error
        self.stop_enter = stop_enter
        self.stop_release = stop_release
        self.calls: list[tuple[str, object]] = []
        self.config.wizard_result.write_text("[]", encoding="utf-8")

    def fill_prebuilt(self, **kwargs: object) -> dict[str, object]:
        self.calls.append(("fill_prebuilt", kwargs))
        if self.fill_enter is not None:
            self.fill_enter.set()
        if self.fill_release is not None:
            self.fill_release.wait(timeout=2)
        if self.fill_error is not None:
            raise self.fill_error
        strategy_jsons = kwargs["strategy_jsons"]
        assert isinstance(strategy_jsons, dict)
        manifest = [
            {
                "filename": f"{name}.json",
                "size": len(payload.encode("utf-8")),
                "sha256": hashlib.sha256(payload.encode("utf-8")).hexdigest(),
            }
            for name, payload in sorted(strategy_jsons.items())
        ]
        config = kwargs["tester_config_json"]
        assert isinstance(config, str)
        readback = {
            "strategy_names": list(strategy_jsons),
            "strategy_file_manifest": manifest,
            "tester_config_hash": hashlib.sha256(config.encode("utf-8")).hexdigest(),
        }
        if self.invalid_fill_readback:
            readback["tester_config_hash"] = "0" * 64
        if self.after_fill is not None:
            self.after_fill()
        return readback

    def start(self) -> dict[str, str]:
        self.calls.append(("start", None))
        if self.write_result:
            if self.write_report:
                report = self.config.report_dir.parent / ("a" * 64) / "portfolio.html"
                report.parent.mkdir(parents=True, exist_ok=True)
                def write_report() -> None:
                    if self.report_as_directory:
                        report.mkdir()
                    else:
                        report.write_text("<html>fresh report</html>", encoding="utf-8")
                    if self.extra_report:
                        (report.parent / "second.html").write_text("<html>second</html>", encoding="utf-8")
                if self.delayed_report:
                    timer = threading.Timer(0.005, write_report)
                    timer.daemon = True
                    timer.start()
                else:
                    write_report()
            self.config.wizard_result.write_text(json.dumps([{
                "runId": "portfolio-run",
                "strategies": list(self.result_names),
                "stats": self.result_stats,
                "chartUrl": f"/tester-report/{'a' * 64}/portfolio.html",
                "period": "2026-01-01..2026-01-14",
            }]), encoding="utf-8")
        if self.after_start is not None:
            self.after_start()
        return {"state": "STARTED", "tester_status": "RUNNING"}

    def stop(self) -> dict[str, str]:
        self.calls.append(("stop", None))
        if self.stop_enter is not None:
            self.stop_enter.set()
        if self.stop_release is not None:
            self.stop_release.wait(timeout=2)
        if self.stop_error is not None:
            raise self.stop_error
        return {"state": "STOPPED"}


def test_stage2_submission_runs_one_prepared_candidate_and_stops_service(tmp_path: Path) -> None:
    payloads = (
        _stage2_payload("BTCUSDT", "LONG", name="PORTFOLIO_BTCUSDT_7_11"),
        _stage2_payload("ETHUSDT", "SHORT", name="PORTFOLIO_ETHUSDT_8_12"),
    )
    service, stage1 = _stage2_service(
        tmp_path,
        _weighted_executable_candidate(payloads, candidate_id="a" * 64, identity="a" * 64),
    )
    fake = _FakeStage2Tester(tmp_path, tuple(payload["strategy"]["name"] for payload in payloads))
    service._local_testing_service_provider = lambda: fake

    submitted = service.submit_tester_submission(
        stage1["campaign_id"],
        {"confirmed": True, "campaign_id": stage1["campaign_id"]},
    )
    job = _wait_stage1(service, submitted)

    assert job["status"] == "SUCCEEDED"
    assert [call[0] for call in fake.calls] == ["fill_prebuilt", "start", "stop"]
    runtime = service.registry.runtime(submitted["job_id"])
    assert runtime["stage2"]["receipt"]["candidate_order"] == 0
    assert runtime["stage2"]["receipt"]["member_identities"] == [
        {"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 1, "result_id": 101},
        {"symbol": "ETHUSDT", "side": "SHORT", "strategy_id": 2, "result_id": 102},
    ]
    assert runtime["stage2_result"]["candidate_id"] == "a" * 64
    assert runtime["stage2_result"]["report_folder"] == "a" * 64
    assert runtime["stage2_result"]["strategy_names"] == [payload["strategy"]["name"] for payload in payloads]
    assert runtime["stage2_result"]["pretest_period"] == {"start_utc": "2026-01-01T00:00:00Z", "end_utc": "2026-01-15T00:00:00Z"}
    evidence = runtime["stage2_result"]["report_evidence"]
    assert evidence["fingerprint"]["size"] > 0
    assert evidence["fingerprint"]["sha256"]
    assert evidence["stable_polls"] >= 2
    assert evidence["accepted_wall_ns"] >= evidence["start_wall_ns"]


def test_stage2_accepts_reverse_alphabetical_candidate_order(tmp_path: Path) -> None:
    payloads = (
        _stage2_payload("ZUSDT", "LONG", name="PORTFOLIO_ZUSDT_7_11"),
        _stage2_payload("AUSDT", "SHORT", name="PORTFOLIO_AUSDT_8_12"),
    )
    service, stage1 = _stage2_service(
        tmp_path,
        _weighted_executable_candidate(payloads, candidate_id="a" * 64, identity="a" * 64),
    )
    fake = _FakeStage2Tester(tmp_path, tuple(payload["strategy"]["name"] for payload in payloads))
    service._local_testing_service_provider = lambda: fake

    submission = service.submit_tester_submission(
        stage1["campaign_id"], {"confirmed": True, "campaign_id": stage1["campaign_id"]}
    )

    assert _wait_stage1(service, submission)["status"] == "SUCCEEDED"


def _submit_stage2_with_fake(tmp_path: Path, *, fake_kwargs: dict[str, object] | None = None) -> tuple[PortfolioPanelService, dict, _FakeStage2Tester, dict]:
    payloads = (
        _stage2_payload("BTCUSDT", "LONG", name="PORTFOLIO_BTCUSDT_7_11"),
        _stage2_payload("ETHUSDT", "SHORT", name="PORTFOLIO_ETHUSDT_8_12"),
    )
    service, stage1 = _stage2_service(
        tmp_path,
        _weighted_executable_candidate(payloads, candidate_id="a" * 64, identity="a" * 64),
    )
    names = tuple(payload["strategy"]["name"] for payload in payloads)
    fake = _FakeStage2Tester(tmp_path, names, **(fake_kwargs or {}))
    service._local_testing_service_provider = lambda: fake
    submission = service.submit_tester_submission(stage1["campaign_id"], {"confirmed": True, "campaign_id": stage1["campaign_id"]})
    return service, stage1, fake, submission


def test_stage2_submission_is_idempotent_for_campaign(tmp_path: Path) -> None:
    service, stage1, fake, first = _submit_stage2_with_fake(tmp_path)
    second = service.submit_tester_submission(stage1["campaign_id"], {"confirmed": True, "campaign_id": stage1["campaign_id"]})
    assert second["job_id"] == first["job_id"]
    assert _wait_stage1(service, first)["status"] == "SUCCEEDED"
    assert [call[0] for call in fake.calls].count("start") == 1


def _tamper_stage1_artifact(root: Path) -> None:
    artifact = next((root / ".portfolio-results").glob("campaign-*/stage1-executables.json"))
    artifact.write_bytes(artifact.read_bytes() + b" ")


def test_stage2_fill_readback_mismatch_prevents_start_and_restores(tmp_path: Path) -> None:
    service, _stage1, fake, submission = _submit_stage2_with_fake(
        tmp_path, fake_kwargs={"invalid_fill_readback": True}
    )

    job = _wait_stage1(service, submission)

    assert job["status"] == "FAILED"
    assert job["diagnostics"][0]["code"] == "PORTFOLIO_JOB_STAGE2_FILL_READBACK_INVALID"
    assert [call[0] for call in fake.calls] == ["fill_prebuilt", "stop"]


def test_stage2_revalidates_artifact_before_start(tmp_path: Path) -> None:
    service, _stage1, fake, submission = _submit_stage2_with_fake(
        tmp_path, fake_kwargs={"after_fill": lambda: _tamper_stage1_artifact(tmp_path)}
    )

    job = _wait_stage1(service, submission)
    assert job["status"] == "FAILED"
    assert job["diagnostics"][0]["code"] == "PORTFOLIO_JOB_STAGE2_INPUT_CHANGED"
    assert [call[0] for call in fake.calls] == ["fill_prebuilt", "stop"]


def test_stage2_classifies_deleted_artifact_before_start(tmp_path: Path) -> None:
    def remove_artifact() -> None:
        next((tmp_path / ".portfolio-results").glob("campaign-*/stage1-executables.json")).unlink()

    service, _stage1, fake, submission = _submit_stage2_with_fake(
        tmp_path, fake_kwargs={"after_fill": remove_artifact}
    )

    job = _wait_stage1(service, submission)
    assert job["status"] == "FAILED"
    assert job["diagnostics"][0]["code"] == "PORTFOLIO_JOB_STAGE2_INPUT_CHANGED"
    assert [call[0] for call in fake.calls] == ["fill_prebuilt", "stop"]


def test_stage2_revalidates_artifact_before_result_acceptance(tmp_path: Path) -> None:
    service, _stage1, fake, submission = _submit_stage2_with_fake(
        tmp_path, fake_kwargs={"after_start": lambda: _tamper_stage1_artifact(tmp_path)}
    )

    job = _wait_stage1(service, submission)
    assert job["status"] == "FAILED"
    assert job["diagnostics"][0]["code"] == "PORTFOLIO_JOB_STAGE2_INPUT_CHANGED"
    assert [call[0] for call in fake.calls] == ["fill_prebuilt", "start", "stop"]


def test_stage2_rejects_unchanged_preexisting_report(tmp_path: Path) -> None:
    report = tmp_path / "tester" / "report" / ("a" * 64) / "portfolio.html"
    report.parent.mkdir(parents=True)
    report.write_text("old", encoding="utf-8")
    service, _stage1, fake, submission = _submit_stage2_with_fake(
        tmp_path, fake_kwargs={"write_report": False}
    )

    job = _wait_stage1(service, submission)

    assert job["status"] == "FAILED"
    assert job["diagnostics"][0]["code"] == "PORTFOLIO_JOB_STAGE2_TIMEOUT"
    assert [call[0] for call in fake.calls][-1] == "stop"


def test_stage2_waits_when_wizard_result_precedes_report(tmp_path: Path) -> None:
    service, _stage1, fake, submission = _submit_stage2_with_fake(
        tmp_path, fake_kwargs={"delayed_report": True}
    )

    assert _wait_stage1(service, submission)["status"] == "SUCCEEDED"
    assert [call[0] for call in fake.calls] == ["fill_prebuilt", "start", "stop"]


@pytest.mark.parametrize("fake_kwargs", ({"extra_report": True}, {"report_as_directory": True}))
def test_stage2_rejects_ambiguous_or_non_file_report(
    tmp_path: Path, fake_kwargs: dict[str, object]
) -> None:
    service, _stage1, fake, submission = _submit_stage2_with_fake(
        tmp_path, fake_kwargs=fake_kwargs
    )

    job = _wait_stage1(service, submission)
    assert job["status"] == "FAILED"
    assert job["diagnostics"][0]["code"] == "PORTFOLIO_JOB_STAGE2_RESULT_INVALID"
    assert [call[0] for call in fake.calls][-1] == "stop"


def test_stage2_rejects_dangling_symlink_report_folder(tmp_path: Path) -> None:
    payloads = (
        _stage2_payload("BTCUSDT", "LONG", name="PORTFOLIO_BTCUSDT_7_11"),
        _stage2_payload("ETHUSDT", "SHORT", name="PORTFOLIO_ETHUSDT_8_12"),
    )
    service, stage1 = _stage2_service(
        tmp_path,
        _weighted_executable_candidate(payloads, candidate_id="a" * 64, identity="a" * 64),
    )
    fake = _FakeStage2Tester(tmp_path, tuple(payload["strategy"]["name"] for payload in payloads))
    report_root = fake.config.report_dir.parent
    report_root.mkdir(parents=True, exist_ok=True)
    try:
        (report_root / ("a" * 64)).symlink_to(report_root / "missing", target_is_directory=True)
    except OSError:
        pytest.skip("symbolic links are unavailable in this test environment")
    service._local_testing_service_provider = lambda: fake

    submission = service.submit_tester_submission(
        stage1["campaign_id"], {"confirmed": True, "campaign_id": stage1["campaign_id"]}
    )
    job = _wait_stage1(service, submission)

    assert job["status"] == "FAILED"
    assert job["diagnostics"][0]["code"] == "PORTFOLIO_JOB_STAGE2_RESULT_INVALID"
    assert [call[0] for call in fake.calls] == ["fill_prebuilt", "stop"]


def test_stage2_rejects_negative_drawdown(tmp_path: Path) -> None:
    stats = {
        "InitialBalance": 10000, "FinalBalance": 10100, "TotalPnL": 100,
        "TotalPnLPercent": 1, "TotalTrades": 2, "WinRate": 50,
        "MaxDrawdown": -1, "MaxDrawdownPercent": 0.1, "TotalFees": 1,
    }
    service, _stage1, fake, submission = _submit_stage2_with_fake(
        tmp_path, fake_kwargs={"result_stats": stats}
    )

    job = _wait_stage1(service, submission)

    assert job["status"] == "FAILED"
    assert job["diagnostics"][0]["code"] == "PORTFOLIO_JOB_STAGE2_RESULT_INVALID"
    assert [call[0] for call in fake.calls][-1] == "stop"


@pytest.mark.parametrize(
    "fake_kwargs",
    (
        {"result_names": ("PORTFOLIO_BTCUSDT_7_11", "WRONG")},
        {"result_stats": {"InitialBalance": "NaN"}},
    ),
)
def test_stage2_invalid_fresh_result_fails_immediately(tmp_path: Path, fake_kwargs: dict[str, object]) -> None:
    service, _stage1, fake, submission = _submit_stage2_with_fake(tmp_path, fake_kwargs=fake_kwargs)
    job = _wait_stage1(service, submission)
    assert job["status"] == "FAILED"
    assert job["diagnostics"][0]["code"] == "PORTFOLIO_JOB_STAGE2_RESULT_INVALID"
    assert [call[0] for call in fake.calls][-1] == "stop"


def test_stage2_stale_result_times_out_and_restores(tmp_path: Path) -> None:
    service, _stage1, fake, submission = _submit_stage2_with_fake(tmp_path, fake_kwargs={"write_result": False})
    job = _wait_stage1(service, submission)
    assert job["status"] == "FAILED"
    assert job["diagnostics"][0]["code"] == "PORTFOLIO_JOB_STAGE2_TIMEOUT"
    assert [call[0] for call in fake.calls][-1] == "stop"


def test_stage2_fill_failure_does_not_stop_unowned_service(tmp_path: Path) -> None:
    service, _stage1, fake, submission = _submit_stage2_with_fake(tmp_path, fake_kwargs={"fill_error": RuntimeError("fill failed")})
    assert _wait_stage1(service, submission)["status"] == "FAILED"
    assert [call[0] for call in fake.calls] == ["fill_prebuilt"]


def test_stage2_cancel_stops_after_fill_and_never_starts_tester(tmp_path: Path) -> None:
    entered, release = threading.Event(), threading.Event()
    service, _stage1, fake, submission = _submit_stage2_with_fake(tmp_path, fake_kwargs={"fill_enter": entered, "fill_release": release})
    assert entered.wait(timeout=2)
    assert service.cancel(submission["job_id"])["status"] == "CANCEL_REQUESTED"
    release.set()
    deadline = time.monotonic() + 2
    while time.monotonic() < deadline and service.job(submission["job_id"])["status"] == "CANCEL_REQUESTED":
        time.sleep(0.01)
    assert service.job(submission["job_id"])["status"] == "CANCELLED"
    assert [call[0] for call in fake.calls] == ["fill_prebuilt", "stop"]


def test_stage2_restore_failure_after_cancel_fails_job(tmp_path: Path) -> None:
    entered, release = threading.Event(), threading.Event()
    service, _stage1, fake, submission = _submit_stage2_with_fake(tmp_path, fake_kwargs={"fill_enter": entered, "fill_release": release, "stop_error": RuntimeError("restore failed")})
    assert entered.wait(timeout=2)
    service.cancel(submission["job_id"])
    release.set()
    deadline = time.monotonic() + 2
    while time.monotonic() < deadline and service.job(submission["job_id"])["status"] == "CANCEL_REQUESTED":
        time.sleep(0.01)
    assert service.job(submission["job_id"])["status"] == "FAILED"


def test_stage2_cancel_removes_result_staged_before_restore(tmp_path: Path) -> None:
    stop_enter, stop_release = threading.Event(), threading.Event()
    service, _stage1, fake, submission = _submit_stage2_with_fake(tmp_path, fake_kwargs={"stop_enter": stop_enter, "stop_release": stop_release})
    assert stop_enter.wait(timeout=2)
    assert service.cancel(submission["job_id"])["status"] == "CANCEL_REQUESTED"
    stop_release.set()
    deadline = time.monotonic() + 2
    while time.monotonic() < deadline and service.job(submission["job_id"])["status"] == "CANCEL_REQUESTED":
        time.sleep(0.01)
    assert service.job(submission["job_id"])["status"] == "CANCELLED"
    assert "stage2_result" not in service.registry.runtime(submission["job_id"])
    assert [call[0] for call in fake.calls] == ["fill_prebuilt", "start", "stop"]


def test_stage2_reserve_failure_discards_new_queued_job(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    payloads = (
        _stage2_payload("BTCUSDT", "LONG", name="PORTFOLIO_BTCUSDT_7_11"),
        _stage2_payload("ETHUSDT", "SHORT", name="PORTFOLIO_ETHUSDT_8_12"),
    )
    service, stage1 = _stage2_service(
        tmp_path,
        _weighted_executable_candidate(payloads, candidate_id="a" * 64, identity="a" * 64),
    )
    service._local_testing_service_provider = lambda: object()

    def fail_reserve(*_args: object, **_kwargs: object) -> None:
        raise OSError("journal unavailable")

    monkeypatch.setattr(service.registry, "reserve_runtime", fail_reserve)
    with pytest.raises(PortfolioPanelError) as error:
        service.submit_tester_submission(stage1["campaign_id"], {"confirmed": True, "campaign_id": stage1["campaign_id"]})
    assert error.value.code == "PORTFOLIO_JOB_START_FAILED"
    assert not any(saved.get("kind") == "portfolio.stage2" for saved in service.registry.list())


def test_stage2_cancel_race_during_commit_becomes_cancelled(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    stop_enter, stop_release = threading.Event(), threading.Event()
    service, _stage1, _fake, submission = _submit_stage2_with_fake(tmp_path, fake_kwargs={"stop_enter": stop_enter, "stop_release": stop_release})
    assert stop_enter.wait(timeout=2)
    original_sync = service.registry.sync
    triggered = False

    def sync_with_cancel(job_id: str, status: dict, *, runtime: dict | None = None) -> dict:
        nonlocal triggered
        if status.get("state") == "COMMITTED" and not triggered:
            triggered = True
            service.cancel(job_id)
        return original_sync(job_id, status, runtime=runtime)

    monkeypatch.setattr(service.registry, "sync", sync_with_cancel)
    stop_release.set()
    deadline = time.monotonic() + 2
    while time.monotonic() < deadline and service.job(submission["job_id"])["status"] in {"RUNNING", "CANCEL_REQUESTED"}:
        time.sleep(0.01)
    assert service.job(submission["job_id"])["status"] == "CANCELLED"


def test_stage2_restart_projects_nonterminal_job_as_interrupted(tmp_path: Path) -> None:
    registry = PanelJobRegistry(tmp_path / ".panel-jobs.json")
    saved = registry.submit("portfolio.stage2", {"campaign_id": "campaign"}, "stage2-restart", ("portfolio_optimizer",), job_id="stage2-restart")
    registry.reserve_runtime("stage2-restart", "stage2", {"campaign_id": "campaign", "input_digest": "i", "config_digest": "c"})
    registry.transition("stage2-restart", "RUNNING")
    restarted = PortfolioPanelService(tmp_path, registry=PanelJobRegistry(tmp_path / ".panel-jobs.json"))
    assert restarted.job(saved["job_id"])["status"] == "INTERRUPTED"


def test_stage1_executable_loader_fails_closed_for_tampered_or_missing_artifact(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    service = PortfolioPanelService(
        tmp_path,
        path,
        finalists_reader=lambda *_: [_finalist()],
        variant_generator=lambda *_: (_weighted_executable_candidate((
            _executable_payload("BTCUSDT", "LONG"),
            _executable_payload("ETHUSDT", "SHORT"),
        )),),
    )
    result = service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})
    assert _wait_stage1(service, result)["status"] == "SUCCEEDED"
    artifact_path = tmp_path / ".portfolio-results" / result["campaign_id"] / "stage1-executables.json"
    artifact_digest = service.registry.runtime(result["job_id"])["executables_digest"]
    restarted = PortfolioPanelService(tmp_path, path)

    artifact_path.write_text('{"schema_version":1}', encoding="utf-8")
    with pytest.raises(PortfolioPanelError) as tampered:
        restarted._load_stage1_executables(result["campaign_id"], result["input_digest"], result["config_digest"], artifact_digest)
    assert tampered.value.code == "PORTFOLIO_STAGE1_EXECUTABLES_UNAVAILABLE"

    artifact_path.unlink()
    with pytest.raises(PortfolioPanelError) as missing:
        restarted._load_stage1_executables(result["campaign_id"], result["input_digest"], result["config_digest"], artifact_digest)
    assert missing.value.code == "PORTFOLIO_STAGE1_EXECUTABLES_UNAVAILABLE"


def test_weighted_summary_enriches_identity_only_members_from_payload_facts() -> None:
    campaign = {
        "campaign_id": "campaign-weighted",
        "launch": {"profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "5000", "max_candidates": 1}]},
    }
    variant = {
        "candidate_id": "candidate-weighted",
        "profile": "BALANCED",
        "search_mode": CAMPAIGN_SEARCH_MODE,
        "members": ({"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 7, "result_id": 11},),
        "strategy_payloads": ({
            "side": "LONG",
            "facts": {"B": "200", "C": "500", "q": "0.25", "x": "100"},
            "strategy": {"basic": {"symbol": "BTCUSDT", "leverage": "3"}},
        },),
        "metrics": {
            "required_bank_usdt": Decimal("200"),
            "B_sat_settings_usdt": Decimal("400"),
            "historical_bank_usdt": Decimal("150"),
            "B_risk_usdt": Decimal("180"),
            "B_margin_usdt": Decimal("120"),
            "I_all_usdt": Decimal("40"),
            "M_all_usdt": Decimal("20"),
            "cdar_peak80_usdt": Decimal("10"),
            "cdar_peak90_usdt": Decimal("14"),
        },
    }

    summary = PortfolioPanelService._summary(campaign, (), (), (variant,), ())

    assert summary["Historical bank USDT"] == Decimal("150")
    assert summary["Stress bank P95 USDT"] == Decimal("180")
    assert summary["IM all USDT"] == Decimal("40")
    assert summary["CDaR peak80 %"] == Decimal("2.5")
    assert summary["CDaR peak90 %"] == Decimal("3.5")
    assert summary["Weighted members"] == [{
        "pair": "BTCUSDT",
        "direction": "LONG",
        "strategy_id": 7,
        "result_id": 11,
        "position_usdt": Decimal("100"),
        "bank_share_pct": Decimal("25.00"),
        "balance_percentage": Decimal("25.00"),
        "pair_multiplier_pct": Decimal("25.00"),
        "max_balance": None,
        "entry_order_percentages": "UNKNOWN",
            "timeframe": None,
            "user_rank": None,
            "source_pnl": None,
            "source_max_drawdown_usdt": None,
            "source_max_drawdown_pct": None,
            "order_count": None,
            "liquidity_utilization_pct": Decimal("20.0"),
            "scaled_max_drawdown_usdt": None,
            "leverage": "3",
            "capacity_usdt": "500",
    }]

    without_saturation = {
        **variant,
        "metrics": {key: value for key, value in variant["metrics"].items() if key != "B_sat_settings_usdt"},
    }
    unavailable = PortfolioPanelService._summary(campaign, (), (), (without_saturation,), ())
    assert unavailable["CDaR peak80 %"] == "UNKNOWN"
    assert unavailable["CDaR peak90 %"] == "UNKNOWN"

    uncapped = PortfolioPanelService._summary(
        {**campaign, "launch": {"profiles": [{"profile_id": "BALANCED", "bank_available_usdt": None, "max_candidates": 1}]}},
        (), (), (variant,), (),
    )
    assert uncapped["Target bank USDT"] == "UNCAPPED"
    assert uncapped["CDaR peak80 target %"] == "—"

    zero_target = PortfolioPanelService._summary(
        campaign, (), (), ({**variant, "metrics": {**variant["metrics"], "B_available_usdt": "0"}},), (),
    )
    assert zero_target["CDaR peak80 target %"] == "UNKNOWN"


def test_weighted_summary_adds_source_maxdd_and_payload_composition_facts() -> None:
    campaign = {
        "campaign_id": "campaign-weighted",
        "config_document": {"profiles": {"BALANCED": {"max_actual_equity_dd_pct": "10"}}},
        "launch": {"profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "500", "max_candidates": 1}]},
        "finalists": ({
            "symbol": "BTCUSDT", "side": "LONG", "strategy_id": 7, "result_id": 11,
            "timeframe": "3h", "user_rank": 2, "order_count": 2,
            "strategy_orders": ({"lot_x": "0.5"}, {"lot_x": "0.5"}),
            "total_pnl": "123.4", "max_drawdown": "40", "max_drawdown_pct": "4.5", "source_initial_balance": "1000",
        },),
    }
    payload = {
        "side": "LONG",
        "facts": {"B": "400", "C": "500", "q": "0.25", "x": "100"},
        "strategy": {
            "basic": {"symbol": "BTCUSDT", "leverage": "3", "balance_percentage_long": "25", "max_balance": "500"},
            "mrs3": {"ma_long": [{"lot_x": "0.5"}, {"lot_x": "0.5"}]},
        },
    }
    variant = {
        "candidate_id": "candidate-weighted", "profile": "BALANCED", "search_mode": CAMPAIGN_SEARCH_MODE,
        "members": ({"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 7, "result_id": 11},),
        "strategy_payloads": (payload,),
        "metrics": {"B_sat_settings_usdt": Decimal("400"), "cdar_peak80_usdt": Decimal("130.32"), "I_all_usdt": Decimal("40"), "M_all_usdt": Decimal("20")},
    }

    summary = PortfolioPanelService._summary(campaign, (), (), (variant,), ())

    assert summary["Profile DD limit %"] == "10"
    assert summary["Weighted members"][0]["position_usdt"] == Decimal("100")
    assert summary["MaxDD SUM USDT"] == Decimal("4")
    member = summary["Weighted members"][0]
    assert member["balance_percentage"] == "25"
    assert member["max_balance"] == "500"
    assert member["entry_order_percentages"] == (Decimal("50"), Decimal("50"))
    assert member["timeframe"] == "3h"
    assert member["user_rank"] == 2
    assert member["source_pnl"] == "123.4"
    assert member["source_max_drawdown_usdt"] == "40"
    assert member["source_max_drawdown_pct"] == "4.5"
    assert member["order_count"] == 2
    assert member["liquidity_utilization_pct"] == Decimal("20")
    assert member["scaled_max_drawdown_usdt"] == Decimal("4")
    assert summary["CDaR peak80 target %"] == Decimal("26.064")
    assert summary["CDaR peak80 saturation %"] == Decimal("32.58")
    assert summary["IM target %"] == Decimal("8")
    assert summary["MM saturation %"] == Decimal("5")
    assert _weighted_entry_order_percentages(payload, "UNKNOWN") == "UNKNOWN"
    for invalid_x in ("not-a-number", "0", "-1", "NaN"):
        assert _weighted_source_maxdd_sum(campaign, ({
            "pair": "BTCUSDT", "direction": "LONG", "strategy_id": 7, "result_id": 11,
            "position_usdt": invalid_x,
        },)) == "UNKNOWN"


def test_weighted_results_enriches_existing_committed_campaign_from_artifact(tmp_path: Path) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    payload = _executable_payload("BTCUSDT", "LONG")
    payload["facts"].update({"B": "200", "C": "500", "q": "0.25", "x": "100"})
    payload["strategy"]["basic"].update({"leverage": "3"})
    second_payload = _executable_payload("ETHUSDT", "SHORT")
    second_payload["facts"]["B"] = "200"
    candidate = _weighted_executable_candidate(
        (payload, second_payload),
        metrics={
            "required_bank_usdt": Decimal("200"),
            "historical_bank_usdt": Decimal("150"),
            "B_risk_usdt": Decimal("180"),
            "B_margin_usdt": Decimal("120"),
            "I_all_usdt": Decimal("40"),
            "M_all_usdt": Decimal("20"),
            "cdar_peak80_usdt": Decimal("10"),
            "cdar_peak90_usdt": Decimal("14"),
        },
    )
    service = PortfolioPanelService(tmp_path, path, finalists_reader=lambda *_: [_finalist()], variant_generator=lambda *_: (candidate,))
    result = service.submit_campaign({
        "pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}],
        "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "5000", "max_candidates": 1}],
        "expected_config_digest": digest,
    })
    assert _wait_stage1(service, result)["status"] == "SUCCEEDED"

    runtime = service.registry.runtime(result["job_id"])
    legacy_summary = dict(runtime["summary"])
    for key in (
        "weighted_result_schema",
        "Historical bank USDT", "Stress bank P95 USDT", "Margin-only bank USDT",
        "IM all USDT", "CDaR peak80 %", "CDaR peak90 %", "Total full notional USDT", "Weighted members",
    ):
        legacy_summary.pop(key, None)
    runtime["summary"] = legacy_summary
    service.registry.sync(result["job_id"], {"state": service.registry.get(result["job_id"])["state"]}, runtime=runtime)
    artifact_path = tmp_path / ".portfolio-results" / result["campaign_id"] / "stage1-executables.json"
    artifact_before = artifact_path.read_bytes()

    restarted = PortfolioPanelService(tmp_path, path)
    enriched = restarted.results(result["campaign_id"])

    assert enriched["summary"]["Historical bank USDT"] == "150"
    assert enriched["summary"]["weighted_result_schema"] == 1
    assert enriched["summary"]["Stress bank P95 USDT"] == "180"
    assert enriched["summary"]["Weighted members"][0]["position_usdt"] == "100"
    assert enriched["summary"]["Weighted members"][0]["bank_share_pct"] == "25.00"
    assert restarted.registry.runtime(result["job_id"])["summary"] == legacy_summary
    assert artifact_path.read_bytes() == artifact_before

    tampered_runtime = restarted.registry.runtime(result["job_id"])
    original_digest = tampered_runtime["executables_digest"]
    tampered_runtime["executables_digest"] = "0" * 64
    restarted.registry.sync(result["job_id"], {"state": restarted.registry.get(result["job_id"])["state"]}, runtime=tampered_runtime)
    assert "Weighted members" not in PortfolioPanelService(tmp_path, path).results(result["campaign_id"])["summary"]

    tampered_runtime["executables_digest"] = original_digest
    tampered_runtime["executables_path"] = str(tmp_path / "outside" / "stage1-executables.json")
    restarted.registry.sync(result["job_id"], {"state": restarted.registry.get(result["job_id"])["state"]}, runtime=tampered_runtime)
    assert "Weighted members" not in PortfolioPanelService(tmp_path, path).results(result["campaign_id"])["summary"]
    assert artifact_path.read_bytes() == artifact_before


def test_weighted_workbook_uses_operator_sheets_and_rounds_derived_values(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = PortfolioPanelService(tmp_path)
    campaign = {
        "campaign_id": "campaign-weighted",
        "input_digest": "i",
        "config_digest": "c",
        "versions": {"policy_version": "p"},
        "created_at_utc": "2000-01-01T00:00:00Z",
        "config_document": {"profiles": {"BALANCED": {"max_actual_equity_dd_pct": "10"}}},
        "launch": {"profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "5000", "max_candidates": 1}]},
        "finalists": ({
            "symbol": "BTCUSDT", "side": "LONG", "strategy_id": 7, "result_id": 11,
            "timeframe": "3h", "user_rank": 2, "order_count": 4,
            "strategy_orders": ({"lot_x": "0.1"}, {"lot_x": "0.2"}, {"lot_x": "0.3"}, {"lot_x": "0.4"}),
            "total_pnl": "123.4", "max_drawdown": "40", "max_drawdown_pct": "4.5", "source_initial_balance": "1000",
        },),
    }
    variant = {
        "candidate_id": "candidate-weighted",
        "profile": "BALANCED",
        "search_mode": CAMPAIGN_SEARCH_MODE,
        "final_pretest_rank": 1,
        "members": ({"symbol": "BTCUSDT", "side": "LONG", "strategy_id": 7, "result_id": 11},),
        "strategy_payloads": ({
            "side": "LONG",
            "facts": {"B": "200", "C": "500", "q": "0.25", "x": "100.126"},
            "strategy": {"basic": {"symbol": "BTCUSDT", "leverage": "3", "balance_percentage_long": "25", "max_balance": "500"}, "mrs3": {"ma_long": [{"lot_x": "0.1"}, {"lot_x": "0.2"}, {"lot_x": "0.3"}, {"lot_x": "0.4"}]}},
        },),
        "metrics": {
            "required_bank_usdt": Decimal("200.126"),
            "B_sat_settings_usdt": Decimal("200.126"),
            "historical_bank_usdt": Decimal("150.125"),
            "B_risk_usdt": Decimal("180.124"),
            "B_margin_usdt": Decimal("120.123"),
            "p30_common_usdt_30d": Decimal("90.126"),
            "max_drawdown_pct": Decimal("10.125"),
            "I_all_usdt": Decimal("40.125"),
            "M_all_usdt": Decimal("20.125"),
            "cdar_peak80_usdt": Decimal("10.125"),
            "cdar_peak90_usdt": Decimal("14.125"),
        },
    }
    path = tmp_path / "weighted.xlsx"
    monkeypatch.setattr(duckdb, "connect", lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("weighted export must not read PerformanceDB")))
    service._write_workbook(path, campaign, (), (), (variant,), ())

    workbook = load_workbook(path, data_only=False)
    try:
        assert workbook.sheetnames == ["Итог", "Варианты", "Состав", "Финалисты", "Исключено", "Metadata"]
        assert workbook["Metadata"].sheet_state == "hidden"
        summary = {row[0].value: row[1].value for row in workbook["Итог"].iter_rows(min_row=2)}
        assert summary["Минимальный банк для DD ≤ 10.00% на истории"] == "150.13"
        assert summary["Банк для DD ≤ 10.00% в 95% стресс-сценариев"] == "180.12"
        assert summary["CDaR худшие 20%, USDT · target/saturation"] == "10.13 USDT · 0.20% / 5.06%"
        variant_headers = {str(cell.value).replace("\n", " "): index for index, cell in enumerate(workbook["Варианты"][1])}
        variant_row = workbook["Варианты"][2]
        assert variant_row[variant_headers["Банк насыщ., USDT"]].value == "200.13"
        assert variant_row[variant_headers["Банк лимитов, USDT"]].value == "120.12"
        assert variant_row[variant_headers["IM USDT · ц/н"]].value == "40.13 USDT · 0.80% / 20.05%"
        assert variant_row[variant_headers["MM USDT · ц/н"]].value == "20.13 USDT · 0.40% / 10.06%"
        assert workbook["Варианты"].row_dimensions[1].height >= 30
        assert all(cell.alignment.wrap_text for cell in workbook["Варианты"][1])
        assert max(cell.width for cell in workbook["Варианты"].column_dimensions.values() if cell.width) <= 30
        composition_headers = {str(cell.value).replace("\n", " "): index for index, cell in enumerate(workbook["Состав"][1])}
        member_row = workbook["Состав"][2]
        assert member_row[composition_headers["Позиция, USDT"]].value == "100.13"
        assert member_row[composition_headers["Множитель пары, %"]].value == "25.00"
        assert member_row[composition_headers["Max balance, USDT"]].value == "500.00"
        assert member_row[composition_headers["TF"]].value == "3h"
        assert member_row[composition_headers["User Rank"]].value == "2"
        assert member_row[composition_headers["Source PnL, USDT"]].value == "123.40"
        assert member_row[composition_headers["Source MaxDD, USDT"]].value == "40.00"
        assert member_row[composition_headers["Source MaxDD, %"]].value == "4.50%"
        assert member_row[composition_headers["ORD_N"]].value == "4"
        assert member_row[composition_headers["X, %"]].value == "10.00%"
        assert member_row[composition_headers["Y, %"]].value == "20.00%"
        assert member_row[composition_headers["Z, %"]].value == "30.00%"
        assert member_row[composition_headers["W, %"]].value == "40.00%"
        assert member_row[composition_headers["Исп. ликв. x/C, %"]].value == "20.03%"
        assert member_row[composition_headers["Инд. MaxDD, USDT"]].value == "4.01"
        for sheet_name in ("Итог", "Варианты", "Состав", "Финалисты", "Исключено"):
            assert workbook[sheet_name].row_dimensions[1].height >= 48
            assert all(cell.alignment.wrap_text for cell in workbook[sheet_name][1])
            assert all(str(cell.value).count("\n") <= 2 for cell in workbook[sheet_name][1])
        assert all(cell.value != "UNKNOWN" for row in workbook["Состав"].iter_rows(min_row=2) for cell in row)
    finally:
        workbook.close()

    unknown_path = tmp_path / "weighted-unknown.xlsx"
    service._write_workbook(unknown_path, {**campaign, "finalists": ()}, (), (), (variant,), ())
    unknown_workbook = load_workbook(unknown_path, data_only=False)
    try:
        unknown_summary = {row[0].value: row[1].value for row in unknown_workbook["Итог"].iter_rows(min_row=2)}
        assert unknown_summary["MaxDD SUM, USDT"] == "—"
        unknown_headers = {str(cell.value).replace("\n", " "): index for index, cell in enumerate(unknown_workbook["Варианты"][1])}
        assert unknown_workbook["Варианты"][2][unknown_headers["MaxDD SUM, USDT"]].value == "—"
    finally:
        unknown_workbook.close()


def test_stage1_executable_builder_and_loader_reject_invalid_pretest_period(tmp_path: Path) -> None:
    campaign = {
        "campaign_id": "campaign-" + "a" * 32,
        "input_digest": "b" * 64,
        "config_digest": "c" * 64,
    }
    candidate = _weighted_executable_candidate((
        _executable_payload("BTCUSDT", "LONG"),
        _executable_payload("ETHUSDT", "SHORT"),
    ), pretest_period={"start_utc": "2026-01-01T00:00:00Z", "end_utc": "2026-01-01T00:00:00Z"})
    with pytest.raises(PortfolioPanelError) as build_error:
        PortfolioPanelService._build_stage1_executables(campaign, (candidate,))
    assert build_error.value.code == "PORTFOLIO_STAGE1_EXECUTABLES_UNAVAILABLE"

    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    service = PortfolioPanelService(
        tmp_path,
        path,
        finalists_reader=lambda *_: [_finalist()],
        variant_generator=lambda *_: (_weighted_executable_candidate((
            _executable_payload("BTCUSDT", "LONG"),
            _executable_payload("ETHUSDT", "SHORT"),
        )),),
    )
    result = service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})
    assert _wait_stage1(service, result)["status"] == "SUCCEEDED"
    artifact_path = tmp_path / ".portfolio-results" / result["campaign_id"] / "stage1-executables.json"
    document = json.loads(artifact_path.read_text(encoding="utf-8"))
    document["candidates"][0]["pretest_period"] = {"start_utc": "2026-01-01T00:00:00Z", "end_utc": "2026-01-01T00:00:00Z"}
    candidate_body = {key: value for key, value in document["candidates"][0].items() if key != "candidate_digest"}
    document["candidates"][0]["candidate_digest"] = hashlib.sha256(json.dumps(candidate_body, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")).hexdigest()
    body = {key: value for key, value in document.items() if key != "payload_digest"}
    document["payload_digest"] = hashlib.sha256(json.dumps(body, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")).hexdigest()
    artifact_path.write_text(json.dumps(document, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False) + "\n", encoding="utf-8")
    with pytest.raises(PortfolioPanelError) as load_error:
        service._load_stage1_executables(result["campaign_id"], result["input_digest"], result["config_digest"], document["payload_digest"], require_committed=False)
    assert load_error.value.code == "PORTFOLIO_STAGE1_EXECUTABLES_UNAVAILABLE"


@pytest.mark.parametrize(
    "candidate",
    (
        _weighted_executable_candidate((
            _executable_payload("BTCUSDT", "LONG"),
            _executable_payload("ETHUSDT", "SHORT"),
        ), limiter_L=1),
        _weighted_executable_candidate((
            _executable_payload("BTCUSDT", "LONG"),
            _executable_payload("ETHUSDT", "SHORT", x="0"),
        )),
        _weighted_executable_candidate((
            _executable_payload("BTCUSDT", "LONG"),
            _executable_payload("ETHUSDT", "SHORT"),
        ), members=(
            _candidate_member("SOLUSDT", "LONG", 1, 101),
            _candidate_member("ETHUSDT", "SHORT", 2, 102),
        )),
    ),
)
def test_stage1_with_no_eligible_off_only_executable_candidate_fails_closed(tmp_path: Path, candidate: dict) -> None:
    path = tmp_path / "portfolio_optimizer.local.json"
    digest = _write_config(path)
    service = PortfolioPanelService(
        tmp_path,
        path,
        finalists_reader=lambda *_: [_finalist()],
        variant_generator=lambda *_: (candidate,),
    )
    result = service.submit_campaign({"pairs": [{"pair": "BTCUSDT", "max_finalist_long": 1, "max_finalist_short": 0}], "profiles": [{"profile_id": "BALANCED", "bank_available_usdt": "10000", "max_candidates": 1}], "expected_config_digest": digest})

    job = _wait_stage1(service, result)
    result_dir = tmp_path / ".portfolio-results" / result["campaign_id"]
    assert job["status"] == "FAILED"
    assert job["diagnostics"][0]["code"] == "PORTFOLIO_STAGE1_EXECUTABLES_UNAVAILABLE"
    assert not (result_dir / "stage1.xlsx").exists()
    assert not (result_dir / "stage1-executables.json").exists()


def test_http_routes_use_typed_error_envelope(tmp_path: Path) -> None:
    config = tmp_path / "config.local.json"
    config.write_text("{}", encoding="utf-8")
    controller = PanelController(tmp_path, config)
    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
    try:
        body = json.dumps({"confirmed": True, "campaign_id": "missing"}).encode()
        connection.request("POST", "/api/v2/portfolio/campaigns/missing/tester-submissions", body, {"Content-Type": "application/json"})
        response = connection.getresponse()
        payload = json.loads(response.read())
        assert response.status == 404
        assert payload["error"]["code"] == "PORTFOLIO_CAMPAIGN_NOT_FOUND"
        assert payload["error"]["field_errors"] == []
    finally:
        connection.close()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


@pytest.mark.parametrize(
    ("headers", "body"),
    (({}, b"{}"), ({"Content-Type": "application/json"}, b""), ({"Content-Type": "application/json"}, b"{")),
)
def test_http_stage2_rejects_before_request_parsing(tmp_path: Path, headers: dict[str, str], body: bytes) -> None:
    config = tmp_path / "config.local.json"
    config.write_text("{}", encoding="utf-8")
    controller = PanelController(tmp_path, config)
    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
    try:
        connection.request("POST", "/api/v2/portfolio/campaigns/missing/tester-submissions", body=body, headers=headers)
        response = connection.getresponse()
        payload = json.loads(response.read())
        assert response.status == (415 if not headers or "Content-Type" not in headers else (400 if not body else 422))
        assert payload["error"]["code"] == "PORTFOLIO_CAMPAIGN_INVALID"
    finally:
        connection.close()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


@pytest.mark.parametrize(
    "endpoint",
    (
        "/api/v2/portfolio/readiness",
        "/api/v2/portfolio/settings",
        "/api/v2/portfolio/jobs/unknown",
        "/api/v2/portfolio/jobs/active",
        "/api/v2/portfolio/campaigns/unknown/results",
        "/api/v2/portfolio/campaigns/unknown/stage1.xlsx",
    ),
)
def test_http_portfolio_get_requires_local_host(tmp_path: Path, endpoint: str) -> None:
    config = tmp_path / "config.local.json"
    config.write_text("{}", encoding="utf-8")
    controller = PanelController(tmp_path, config)
    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
    try:
        connection.request("GET", endpoint, headers={"Host": "attacker.example"})
        response = connection.getresponse()
        payload = json.loads(response.read())
        assert response.status == 403
        assert payload["error"] == "local Host header required"
    finally:
        connection.close()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_http_cancel_requires_json_request_and_returns_202(tmp_path: Path) -> None:
    config = tmp_path / "config.local.json"
    config.write_text("{}", encoding="utf-8")
    controller = PanelController(tmp_path, config)
    controller._panel_jobs.submit(
        "portfolio.stage1",
        {"campaign_id": "campaign", "input_digest": "input", "config_digest": "config"},
        "portfolio-cancel",
        ("portfolio_optimizer",),
        job_id="portfolio-job",
    )
    controller._panel_jobs.reserve_runtime(
        "portfolio-job",
        "campaign",
        {"campaign_id": "campaign", "input_digest": "input", "config_digest": "config"},
    )
    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
    try:
        connection.request("POST", "/api/v2/portfolio/jobs/portfolio-job/cancel", body=b"{}", headers={})
        response = connection.getresponse()
        payload = json.loads(response.read())
        assert response.status == 415
        assert payload["error"]["code"] == "PORTFOLIO_JOB_CANCEL_INVALID"
        connection.request("POST", "/api/v2/portfolio/jobs/portfolio-job/cancel", body=b"{}", headers={"Content-Type": "application/json"})
        response = connection.getresponse()
        payload = json.loads(response.read())
        assert response.status == 202
        assert payload["job_id"] == "portfolio-job"
        assert payload["status"] in {"CANCEL_REQUESTED", "CANCELLED"}
    finally:
        connection.close()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


@pytest.mark.parametrize("failure", (PanelJobError("RESOURCE_BUSY"), OSError("registry unavailable"), TypeError("registry unavailable"), ValueError("registry unavailable")))
def test_http_cancel_runtime_failures_return_500(monkeypatch: pytest.MonkeyPatch, tmp_path: Path, failure: Exception) -> None:
    config = tmp_path / "config.local.json"
    config.write_text("{}", encoding="utf-8")
    controller = PanelController(tmp_path, config)
    controller._panel_jobs.submit("portfolio.stage1", {}, "cancel-http-runtime", ("portfolio_optimizer",), job_id="cancel-http-runtime")
    monkeypatch.setattr(controller._portfolio_service.registry, "cancel", lambda *_: (_ for _ in ()).throw(failure))
    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
    try:
        connection.request("POST", "/api/v2/portfolio/jobs/cancel-http-runtime/cancel", body=b"{}", headers={"Content-Type": "application/json"})
        response = connection.getresponse()
        payload = json.loads(response.read())
        assert response.status == 500
        assert payload["error"]["code"] == "PORTFOLIO_JOB_RUNTIME_UNAVAILABLE"
    finally:
        connection.close()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_http_active_runtime_lookup_failure_is_typed(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    config = tmp_path / "config.local.json"
    config.write_text("{}", encoding="utf-8")
    controller = PanelController(tmp_path, config)
    controller._panel_jobs.submit("portfolio.stage1", {}, "active-runtime", ("portfolio_optimizer",), job_id="active-runtime")
    controller._portfolio_service._threads["active-runtime"] = object()
    monkeypatch.setattr(controller._portfolio_service.registry, "runtime", lambda *_args: (_ for _ in ()).throw(KeyError("runtime")))
    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
    try:
        connection.request("GET", "/api/v2/portfolio/jobs/active")
        response = connection.getresponse()
        payload = json.loads(response.read())
        assert response.status == 500
        assert payload["error"]["code"] == "PORTFOLIO_JOB_RUNTIME_UNAVAILABLE"
    finally:
        connection.close()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


@pytest.mark.parametrize("runtime_mode", ("raises_key_error", "returns_non_mapping"))
def test_http_job_runtime_failures_are_typed(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, runtime_mode: str) -> None:
    config = tmp_path / "config.local.json"
    config.write_text("{}", encoding="utf-8")
    controller = PanelController(tmp_path, config)
    controller._panel_jobs.submit("portfolio.stage1", {}, "job-runtime", ("portfolio_optimizer",), job_id="job-runtime")

    if runtime_mode == "raises_key_error":
        def broken_runtime(*_args, **_kwargs):
            raise KeyError("runtime")

        monkeypatch.setattr(controller._portfolio_service.registry, "runtime", broken_runtime)
    else:
        monkeypatch.setattr(controller._portfolio_service.registry, "runtime", lambda *_args, **_kwargs: ["runtime"])

    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
    try:
        connection.request("GET", "/api/v2/portfolio/jobs/job-runtime")
        response = connection.getresponse()
        payload = json.loads(response.read())
        assert response.status == 500
        assert payload == {
            "error": {
                "code": "PORTFOLIO_JOB_RUNTIME_UNAVAILABLE",
                "message": "portfolio job runtime is unavailable",
                "field_errors": [],
            }
        }
    finally:
        connection.close()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


@pytest.mark.parametrize("endpoint", ("/api/v2/portfolio/readiness", "/api/v2/portfolio/settings"))
@pytest.mark.parametrize(
    "failure",
    (
        OSError("read failed"),
        UnicodeDecodeError("utf-8", b"\xff", 0, 1, "invalid"),
        PortfolioConfigError("invalid config"),
    ),
)
def test_http_settings_read_failures_are_typed(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    endpoint: str,
    failure: Exception,
) -> None:
    config = tmp_path / "config.local.json"
    config.write_text("{}", encoding="utf-8")
    controller = PanelController(tmp_path, config)

    def broken_settings_raw(*_args, **_kwargs):
        raise failure

    monkeypatch.setattr(controller._portfolio_service, "_settings_raw", broken_settings_raw)
    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
    try:
        connection.request("GET", endpoint)
        response = connection.getresponse()
        payload = json.loads(response.read())
        assert response.status == 500
        assert payload["error"]["code"] == "PORTFOLIO_SETTINGS_UNAVAILABLE"
        assert payload["error"]["field_errors"] == []
    finally:
        connection.close()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_http_results_partial_runtime_is_typed(tmp_path: Path) -> None:
    config = tmp_path / "config.local.json"
    config.write_text("{}", encoding="utf-8")
    controller = PanelController(tmp_path, config)
    controller._panel_jobs.submit("portfolio.stage1", {}, "partial-runtime", ("portfolio_optimizer",), job_id="partial-runtime")
    controller._panel_jobs.reserve_runtime("partial-runtime", "campaign", {"campaign_id": "campaign"})
    controller._panel_jobs.transition("partial-runtime", "RUNNING")
    controller._panel_jobs.transition("partial-runtime", "COMMITTED")
    server = create_panel_server("127.0.0.1", 0, controller)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
    try:
        connection.request("GET", "/api/v2/portfolio/campaigns/campaign/results")
        response = connection.getresponse()
        payload = json.loads(response.read())
        assert response.status == 500
        assert payload["error"]["code"] == "PORTFOLIO_JOB_RUNTIME_UNAVAILABLE"
    finally:
        connection.close()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)
